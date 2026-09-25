"""
Verifier - verifier.py - the whole engine, in one file. For every reviewer.

Reads a model's methodology, its package of code and data, and its documentation into numbered units;
links each unit of the package to the units it takes something from and gives something to; and asks the
organisation's language model, through the chat() of cell 2, to explain each unit of code, to find the
chunks of the methodology that bear on it, and to flag where the code may depart from them. One
deliverable: Output.xlsm. Everything a run does is recorded in Audit_Log.xlsx, beside it and the three input folders.

The file is one piece of engineering in three parts, in dependency order:
  the contracts, the reading floor and the front door
  reading the methodology, the documentation and the model package
  the run: its folder, its record, the organisation's model and Output.xlsm
"""
import bz2
import collections
import concurrent.futures
import csv
import dataclasses
import datetime
import email
import functools
import getpass
import random
import gzip
import hashlib
import html.entities
import html.parser
import io
import json
import lzma
import os
import subprocess
import re
import shutil
import struct
import sys
import tarfile
import tempfile
import threading
import time
import uuid
import traceback
import unicodedata
import xml.etree.ElementTree as ElementTree
try:
    import yaml
except ImportError:                 # cell 1 imports the engine to install what is missing, yaml among them
    yaml = None
import zipfile
import zlib
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Callable, Optional
from xml.sax.saxutils import escape


# ================================================================================================
# the contracts, the reading floor and the front door
# ================================================================================================
# ---------------------------------------------------------------- vocabulary, and the words the tool may never use

UNDECIDED_REASONS = ("the equation is an image", "the equation could not be read")   # why an equation was not read

KIND_FUNCTION, KIND_FORMULA, KIND_TOPLEVEL = "Function", "Formula statement", "Top-level statement"
KIND_TEST, KIND_TABLE, KIND_OBJECT = "Test block", "Parameter table", "Parameter object"
KIND_ROXYGEN, KIND_VIGNETTE = "Roxygen block", "Vignette text"
KIND_COMPILED, KIND_NOT_READ, KIND_OTHER = "Compiled code", "File not read", "Other file"
UNIT_KINDS = (KIND_FUNCTION, KIND_FORMULA, KIND_TOPLEVEL, KIND_TEST, KIND_TABLE, KIND_OBJECT,
              KIND_ROXYGEN, KIND_VIGNETTE, KIND_COMPILED, KIND_NOT_READ, KIND_OTHER)
CHUNK_KINDS = ("Paragraph", "Table", "Figure", "Equation")

# This one assignment has to name the words the tool may never use; nothing else in the engine may.
# These terminologies CAN BE USED ONLY by HUMAN reviewers/validators. Machines cannot make these determinations.
BANNED_WORDING_PATTERNS = (
    r"\bfindings?\b", r"\berrors?\b", r"\bseverity\b", r"\bsevere\b", r"\bcritical\b",
    r"\bmajor\b", r"\bminor\b",
    r"\b(high|medium|low)[\s-]+(risk|priority|rating|impact)\b",
    r"\b(risk|priority|rating|impact)\s*[:=]?\s*(high|medium|low)\b")
_BANNED_RE = re.compile("|".join(BANNED_WORDING_PATTERNS), re.IGNORECASE)

def has_banned_wording(text):
    """Return the first word in `text` that the tool's own wording may not use, or "". Enforces: R1"""
    found = _BANNED_RE.search(text or "")
    return found.group(0) if found else ""

# ---------------------------------------------------------------- data contracts: the records (plan 2.7)
@dataclass(frozen=True)
class TableData:
    """A table kept whole: header cells, body rows, and which column identifies a row."""
    header: tuple = (); rows: tuple = (); row_key: str = ""

@dataclass(frozen=True)
class EquationData:
    """An equation as found: its source form, whether the tool could read it, and its tree."""
    source_form: str = ""; linear: str = ""; readable: bool = False
    not_readable_reason: str = ""; image_sha256: str = ""

@dataclass(frozen=True)
class Chunk:
    """One citable unit of a document: a paragraph, a table, a figure or an equation."""
    ref: str; corner: str; source_file: str; kind: str; level: int; heading_chain: tuple
    numbering: str; text: str; locator: str; content_hash: str
    table: Optional[TableData] = None; equation: Optional[EquationData] = None
    numbering_reconstructed: bool = False
    caption: str = ""; not_read_reason: str = ""; para_label: str = ""   # para_label: the number the document itself gives ("36.")

@dataclass(frozen=True)
class CodeDetail:
    """What the R reader learned about a function or a statement, without running it."""
    formals: tuple = (); calls: tuple = (); symbols_written: tuple = (); exported: Optional[bool] = None
    reads_data: tuple = ()

@dataclass(frozen=True)
class ParameterDataDetail:
    """The profile of one stored data object."""
    object_name: str; container_file: str; dims: tuple = (); columns: tuple = ()
    assessable: bool = True; not_assessable_reason: Optional[str] = None

@dataclass(frozen=True)
class RoxygenDetail:
    """A roxygen block: what it documents, and its tags with their lines."""
    documents_ref: Optional[str] = None; documents_name: str = ""; tags: tuple = ()


@dataclass(frozen=True)
class ModelUnit:
    """One citable unit of the package (see UNIT_KINDS)."""
    ref: str; kind: str; file: str; lines: Optional[tuple]; name: str; inside: str; text: str
    parent_ref: Optional[str]; file_sha256: str; content_hash: str
    code: Optional[CodeDetail] = None; data: Optional[ParameterDataDetail] = None
    roxygen: Optional[RoxygenDetail] = None
    read_problem: Optional[str] = None

@dataclass(frozen=True)
class Provenance:
    """Which run and step produced a record, and from which AI exchange if any."""
    run_id: str; step_id: str; step: str
    prompt_hash: Optional[str] = None; response_hash: Optional[str] = None



@dataclass
class StepContext:
    """What every step function receives. It never contains the access token."""
    settings: dict; options: dict; read: Callable
    work_dir: str; note: Callable; provenance: Optional[Provenance] = None

@dataclass
class StepResult:
    """What every step function returns: records by kind, counts, and plain notes."""
    records: dict = field(default_factory=dict); counts: dict = field(default_factory=dict)
    messages: list = field(default_factory=list); finished: bool = True

# ---------------------------------------------------------------- canonical JSON and hashes
def to_plain(value):
    """Turn dataclasses, tuples and Decimals into plain JSON-ready values."""
    if dataclasses.is_dataclass(value) and not isinstance(value, type):
        return {f.name: to_plain(getattr(value, f.name)) for f in dataclasses.fields(value)}
    if isinstance(value, dict):
        return {str(k): to_plain(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [to_plain(v) for v in value]
    if isinstance(value, Decimal):
        return format(value, "f")
    return value

def canonical_json(value):
    """One JSON text per content: sorted keys, no spare white space. Enforces: R5"""
    return json.dumps(to_plain(value), sort_keys=True, ensure_ascii=False, separators=(",", ":"))
def sha256_bytes(data):
    """SHA-256 of bytes, in hexadecimal."""
    return hashlib.sha256(data).hexdigest()
def sha256_text(text):
    """SHA-256 of a text in UTF-8, in hexadecimal."""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()
def swhid_content(data):
    """The ISO/IEC 18670 content identifier of a file; the same value Git computes."""
    return "swh:1:cnt:" + hashlib.sha1(b"blob %d\0" % len(data) + data).hexdigest()

def normalise_text(text):
    """Unicode NFC, one kind of line ending, runs of white space collapsed to one space."""
    text = unicodedata.normalize("NFC", text or "").replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"\s+", " ", text).strip()

def content_hash(text):
    """The hash that makes a citation re-verifiable. It ignores white-space and
    line-ending differences and nothing else. Enforces: R4"""
    return sha256_text(normalise_text(text))

def make_ref(prefix, number):
    """make_ref("C", 9) gives "C-0009". Numbers follow reading order."""
    return "%s-%04d" % (prefix, number)



# ---------------------------------------------------------------- numbers
# A number as prose writes it. Thousands may be grouped with commas in groups of exactly three
# ("1,000", "250,000.5"): that grouping is unambiguous, so it is read as one number. A comma
# between other digit counts ("1,5") is not a grouping and stays two numbers, because in some
# writing it is a decimal comma and guessing which would be guessing. Space grouping ("1 000") is
# not read either, since "section 3 100 samples" would fuse. Enforces: R3
_NUMBER_RE = re.compile(
    r"(?<![\w.])(?P<num>[-\u2212]?(?:\d{1,3}(?:,\d{3})+(?!\d)|\d+)(?:\.\d+)?(?:[eE][-+]?\d+)?)"
    r"(?P<unit>\s?%|\s?(?:basis points?|bps?)\b|(?:st|nd|rd|th)\s+percentile\b)?")
_LABEL_BEFORE_RE = re.compile(
    r"(?:table|section|sections|equation|eq\.|figure|fig\.|annex|appendix|chapter|paragraph|"
    r"page|step|version|level)\s*\(?$", re.IGNORECASE)

def parse_number(as_written, unit=""):
    """Bring one written number to a decimal value and keep how it was written.
    Percent, basis points, scientific notation and "99.9th percentile" are converted
    first; `decimals` counts decimals after that conversion, because the value rule
    (plan 2.8) compares at the precision the methodology states."""
    text = as_written.replace("\u2212", "-").strip()
    if re.fullmatch(r"-?\d{1,3}(?:,\d{3})+(?:\.\d+)?(?:[eE][-+]?\d+)?", text):
        text = text.replace(",", "")                 # the thousands grouping is how it was written, not part of the value
    try:
        value = Decimal(text)
    except InvalidOperation:
        return None
    mantissa = re.split(r"[eE]", text)[0]
    decimals = len(mantissa.split(".")[1]) if "." in mantissa else 0
    exponent = int(re.split(r"[eE]", text)[1]) if re.search(r"[eE]", text) else 0
    kind = unit.strip().lower()
    shift = 2 if (kind == "%" or kind.endswith("percentile")) else 0
    if kind.startswith("basis") or kind.startswith("bp"):
        shift, kind = 4, "bp"
    if kind.endswith("percentile"):
        kind = "percentile"
    return {"value": plain_decimal(value.scaleb(-shift)), "as_written": as_written.strip() + unit.rstrip(),
            "decimals": max(0, decimals - exponent + shift), "unit": kind}

def plain_decimal(value):
    """A Decimal as the shortest plain decimal text: no exponent, no trailing zeros."""
    text = format(value.normalize(), "f")
    return "0" if text in ("-0", "") else text

def find_numbers(text):
    """All numbers written in a piece of prose, with their position. Numbers that only
    label something ("Table 3", "section 4.2"), years and list numbering are left out."""
    found = []
    for match in _NUMBER_RE.finditer(text or ""):
        before = text[:match.start()]
        number = match.group("num")
        if _LABEL_BEFORE_RE.search(before[-24:]):
            continue
        if re.fullmatch(r"(19|20)\d\d", number) and not match.group("unit"):
            continue
        if not before.strip() and re.match(r"\)|\.?\s+[A-Z]", text[match.end():match.end() + 4]):
            continue                      # "3.1 Floors" or "2) ..." at the start: numbering
        parsed = parse_number(number, match.group("unit") or "")
        if parsed is not None:
            parsed["position"] = match.start()
            found.append(parsed)
    return found

def plain_number(value, digits=6):
    """How a computed value is shown to an analyst: whole numbers stay whole, other
    values show at most six significant digits, never an exponent. Enforces: R10"""
    if isinstance(value, str):
        return value
    if value != value or value in (float("inf"), float("-inf")):
        return "not a number" if value != value else ("infinity" if value > 0 else "minus infinity")
    if float(value) == int(value) and abs(value) < 1e15:
        return str(int(value))
    text = "%.*g" % (digits, float(value))
    return plain_decimal(Decimal(text)) if "e" in text else text

# ---------------------------------------------------------------- symbols
GREEK = {"alpha": "\u03b1", "beta": "\u03b2", "gamma": "\u03b3", "delta": "\u03b4",
         "epsilon": "\u03b5", "theta": "\u03b8", "kappa": "\u03ba", "lambda": "\u03bb",
         "mu": "\u03bc", "nu": "\u03bd", "pi": "\u03c0", "rho": "\u03c1", "sigma": "\u03c3",
         "tau": "\u03c4", "phi": "\u03c6", "omega": "\u03c9", "Phi": "\u03a6", "Sigma": "\u03a3",
         "Delta": "\u0394", "Omega": "\u03a9"}
_GREEK_BY_CHAR = {char: name for name, char in GREEK.items()}
_SUBSCRIPT_CHARS = str.maketrans("\u2080\u2081\u2082\u2083\u2084\u2085\u2086\u2087\u2088\u2089"
                                 "\u1d62\u2c7c\u2096\u2099\u209c", "0123456789ijknt")

def normalise_symbol(symbol):
    """One form for a symbol however it was written: the Greek letter by name, by
    character or in LaTeX form; a subscript written PD_i, PD[i], PD_{i} or with a
    subscript character. Case is kept for single letters and folded for words."""
    text = unicodedata.normalize("NFC", symbol or "").strip().strip("$`")
    text = re.sub(r"\\([A-Za-z]+)", r"\1", text)
    subscript = ""
    match = re.fullmatch(r"(.+?)(?:_\{?([^{}]+)\}?|\[([^\]]+)\])", text)
    if match:
        text, subscript = match.group(1), match.group(2) or match.group(3)
    else:
        tail = re.search(r"[\u2080-\u2089\u1d62\u2c7c\u2096\u2099\u209c]+$", text)
        if tail and tail.start() > 0:
            text, subscript = text[:tail.start()], tail.group(0).translate(_SUBSCRIPT_CHARS)
    parts = []
    for part in (text, subscript):
        part = "".join(_GREEK_BY_CHAR.get(char, char) for char in part)
        if len(part) > 1 and part not in GREEK:
            part = part.lower()
        parts.append(part)
    return parts[0] + ("_" + parts[1] if parts[1] else "")

# ---------------------------------------------------------------- the expression tree
@dataclass(frozen=True)
class Expr:
    """the tool's neutral tree for a formula. op is one of: num, sym, add, sub, mul, div,
    pow, neg, call, cmp, piecewise, eq. `name` is a symbol, a neutral function name
    (call) or a comparison sign (cmp); `value` is a decimal text (num)."""
    op: str; name: Optional[str] = None; value: Optional[str] = None; args: tuple = ()
    span: Optional[tuple] = None


_INFIX = {"add": (" + ", 1), "sub": (" - ", 1), "mul": (" * ", 2), "div": (" / ", 2), "pow": ("^", 4)}

def expr_to_text(expr, parent_rank=0):
    """The tree in the tool's linear notation, the form shown to analysts and to the AI."""
    if expr.op == "num":
        return expr.value
    if expr.op == "sym":
        return expr.name
    if expr.op == "neg":
        text, rank = "-" + expr_to_text(expr.args[0], 3), 3
    elif expr.op in _INFIX:
        sign, rank = _INFIX[expr.op]
        left = expr_to_text(expr.args[0], rank + (1 if expr.op == "pow" else 0))
        text = left + sign + expr_to_text(expr.args[1], rank + (0 if expr.op == "pow" else 1))
    elif expr.op == "call":
        return "%s(%s)" % (expr.name, ", ".join(expr_to_text(arg) for arg in expr.args))
    elif expr.op == "cmp":
        text, rank = (" %s " % expr.name).join(expr_to_text(arg, 1) for arg in expr.args), 0
    elif expr.op == "eq":
        return "%s = %s" % (expr_to_text(expr.args[0]), expr_to_text(expr.args[1]))
    elif expr.op == "piecewise":
        return "piecewise(%s)" % ", ".join(expr_to_text(arg) for arg in expr.args)
    else:
        return "?"
    return "(%s)" % text if rank < parent_rank else text


# ================================================================================================



# ---------------------------------------------------------------- element helpers
def local_name(tag):
    """'{namespace}oMath' and 'm:oMath' both become 'omath'."""
    if not isinstance(tag, str):
        return ""
    return tag.rsplit("}", 1)[-1].rsplit(":", 1)[-1].lower()

def attribute(element, name):
    """The value of an attribute, whatever namespace prefix it carries."""
    for key, value in element.attrib.items():
        if local_name(key) == name:
            return value
    return ""

def child_named(element, name):
    """The first child with this local tag name, or None."""
    for child in element:
        if local_name(child.tag) == name:
            return child
    return None

# ---------------------------------------------------------------- baseline slicing: what a document's shape says
def attribute_text(element, names, digits_too=False):
    """The first of the named attributes that holds text worth reading, with its name. A bare
    number is no heading, so it is passed over unless digits_too."""
    for name in names:
        value = (element.get(name) or "").strip()
        if value and (digits_too or not value.isdigit()):
            return value, name
    return "", ""

BLOCK_FAMILIES = ("heading", "container", "list_container", "paragraph", "list_item", "table")

def written_numbering(element, rules):
    """The numbering an element carries in an attribute, exactly as the document wrote it
    (num="36." gives "36."): more faithful than any count the tool could make, skipped numbers included."""
    return attribute_text(element, rules["numbering_attributes"], digits_too=True)[0]

def table_rows(element, rules):
    """The rows of a table: the children that hold cells, looked for directly below the table
    and below the wrappers the rules know (thead, tbody, tgroup)."""
    known = rules["family_of"]
    holders = [element] + [part for part in element if known.get(local_name(part.tag)) == "table_part"]
    return [row for holder in holders for row in holder
            if local_name(row.tag) and len(row) and known.get(local_name(row.tag)) in (None, "row")]

def discover_table_shape(element, rules, is_table):
    """Give a family to every tag used inside a table, whatever the tags are called. Everything
    in a row is a cell. Two kinds are told apart by name, since only the name says what they are
    for: a tag holding the table's number or title (<tablenumber>, <tabletitle>) is the caption,
    one holding a column heading (<tablecolhead>) a header cell. A tag nobody anticipated
    (<tablesub>, <tabletext>) is simply a cell and keeps its text.
    An element the rules already name as a table (is_table) needs no more than that: real tables
    have rows of differing width (a note below, a heading that spans). An element NOT known to be
    a table has to make the case by its shape, and two guards keep the skeleton of a document,
    which repeats twice over as a table does, from being read as one: a cell holds words and
    never a block, and most rows are the same width. Returns (row tags, {tag: family}) or None."""
    known, rows = rules["family_of"], table_rows(element, rules)
    cells = [cell for row in rows for cell in row if local_name(cell.tag)]
    if not cells:
        return None
    if not is_table:
        widths = [sum(1 for cell in row if local_name(cell.tag)) for row in rows]
        usual = max(set(widths), key=widths.count)
        if (len(rows) < 2 or len({local_name(row.tag) for row in rows}) > 1 or usual < 2 or widths.count(usual) * 2 <= len(rows)
                or any(known.get(local_name(cell.tag)) for cell in cells)
                or any(known.get(local_name(below.tag)) in BLOCK_FAMILIES for cell in cells for below in cell)):
            return None
    families = {}
    for name in {local_name(cell.tag) for cell in cells}:
        if any(word in name for word in rules["caption_tag_words"]):
            families[name] = "caption"
        else:
            families[name] = "header_cell" if any(word in name for word in rules["header_tag_words"]) else "cell"
    return {local_name(row.tag) for row in rows}, families

def discover_families(root, rules, report):
    """Work out a family for each tag this document uses that the rules do not name, from the
    way the tag behaves here. The rules always win, so a schema the tool already knows is read
    exactly as before; discovery only speaks where they are silent. It looks, in order, for: a
    table; an element carrying its own heading in an attribute; one holding other blocks (a
    container); one holding text (a paragraph). Every decision is recorded with its reason in
    plain words, shown on Model_Package_Info, and can be overridden by the project's tag_rules.yaml, beside its three input folders."""
    known, found = rules["family_of"], {}

    def note(tag, family, reason):
        if tag not in known and tag not in found:
            found[tag] = family
            report[tag] = {"family": family, "reason": reason, "count": 1 if tag == name else 0}

    # A tag that sits in running text, never holding a block of its own, is read inline: its
    # words belong to the sentence around it, not to a paragraph of their own.
    def readable_elements(element):
        """Every element the walker will actually read. Discovery stops where the walker
        stops: the inside of an equation or a figure is read by its own reader, and what the
        rules ignore is never read at all, so neither is catalogued here."""
        yield element
        if known.get(local_name(element.tag)) in ("equation", "figure", "ignore"):
            return
        for child in element:
            if local_name(child.tag):
                yield from readable_elements(child)

    elements = list(readable_elements(root))
    inline_looking = set()
    for element in elements:
        for child in element:
            if (child.tail or "").strip() or (element.text or "").strip():
                inline_looking.add(local_name(child.tag))

    for element in elements:
        name = local_name(element.tag)
        if not name:
            continue
        if name in report:
            report[name]["count"] += 1
        # Table shape is looked for under a known table too: the element may be named in the
        # rules while the row and cell tags inside it are not.
        if known.get(name) == "table" or name not in known:
            shape = discover_table_shape(element, rules, known.get(name) == "table")
            if shape:
                row_tags, families = shape
                note(name, "table", "holds rows of cells")
                for row_tag in sorted(row_tags):
                    note(row_tag, "row", "holds the cells of <%s>" % name)
                reasons = {"caption": "holds the number or the title of <%s>" % name,
                           "header_cell": "holds a column heading of <%s>" % name,
                           "cell": "sits in a row of <%s>" % name}
                for tag in sorted(families):
                    note(tag, families[tag], reasons[families[tag]])
                continue
        if name in known or name in found:
            continue
        blocks_below = any(local_name(child.tag) and known.get(local_name(child.tag)) not in ("inline", "ignore")
                           for child in element)
        heading, attribute = attribute_text(element, rules["heading_attributes"])
        if heading and blocks_below:
            note(name, "container", "carries its own heading in the %s attribute" % attribute)
        elif blocks_below:
            note(name, "container", "holds other blocks")
        elif name in inline_looking and not len(element):
            note(name, "inline", "appears inside running text")
        elif (element.text or "").strip() or len(element):
            note(name, "paragraph", "holds text")
        else:
            note(name, "paragraph", "empty")
    return found

def first_numbering(text, rules):
    """The numbering at the start of a heading as written, and the name of its scheme."""
    for scheme in rules["numbering_schemes"]:
        match = re.match(scheme["pattern"], text)
        if match:
            return match.group(0).strip(), scheme["name"]
    return "", ""

def infer_levels(blocks, rules):
    """Give every heading its level. When the file nests its sections, the nesting decides.
    When nesting is flat, numbering decides: a dotted number gives its depth directly
    (relative to the level of plain numbers); any other scheme seen for the first time is one
    level deeper than the heading before it, and a scheme seen before returns to its level."""
    headings = [b for b in blocks if b["type"] == "heading"]
    hints = sorted({b["level_hint"] for b in headings if b["level_hint"] is not None})
    nested = len(hints) > 1
    scheme_level, current, run_hint, run_schemes = {}, 0, None, {}
    for block in headings:
        written, scheme = first_numbering(block["text"], rules)
        if block["numbering"] and not scheme:
            written, scheme = first_numbering(block["numbering"] + " ", rules)
        block["numbering"] = block["numbering"] or written
        if nested:                                       # a flat-numbered stretch inside a nested file (an annex)
            if block["level_hint"] != run_hint:
                run_hint, run_schemes = block["level_hint"], {}
            if scheme and scheme != "dotted":
                run_schemes.setdefault(scheme, len(run_schemes))
            block["level"] = hints.index(block["level_hint"]) + 1 + run_schemes.get(scheme, 0)
        elif scheme == "dotted":
            depth = block["numbering"].strip(".").count(".") + 1
            block["level"] = scheme_level.get("number", 1) + depth - 1
        elif scheme:
            if scheme not in scheme_level:
                scheme_level[scheme] = current + 1
            block["level"] = scheme_level[scheme]
        else:
            block["level"] = max(1, current) if current else 1
        current = block["level"]
    return blocks

def without_page_furniture(lines, pages, state, file_name):
    """Leave out what a page carries only because it is a page: the running title, the footer
    with its date and page number, the logo. Such a line sits in the top or bottom margin and
    comes back on most pages, the same but for its digits. It is not part of what the document
    says, and left in it cuts a list or a sentence in two wherever a page ends. What was left
    out is reported, so nothing goes missing unseen: this is the "declared drop" class of the
    content account, and the note is what declares it. Enforces: R2, R13"""
    same = lambda entry: (entry["edge"], re.sub(r"\d+", "#", entry["text"]))
    on_pages = {}
    for entry in lines:
        if entry["edge"] and "text" in entry:
            on_pages.setdefault(same(entry), set()).add(entry["page"])
    furniture = {key for key, found in on_pages.items() if pages >= 3 and len(found) >= max(3, pages // 2)}
    for entry in lines:                              # a declared drop: named, counted, and shown in the account
        if entry.get("edge") and "text" in entry and same(entry) in furniture:
            state.dropped.append(entry["text"])
    for edge, text in sorted(furniture):
        state.notes.append("%s: left out as a page header or footer, because it repeats in the %s margin of %d of %d pages: '%s'"
                           % (file_name, edge, len(on_pages[(edge, text)]), pages, text))
    return [entry for entry in lines if not (entry["edge"] and "text" in entry and same(entry) in furniture)]

# ---------------------------------------------------------------- the fidelity ledger
# Enforces: R13. A reader may decide HOW a file is sliced. It may not add a word the file does
# not contain, and it may not lose a word the file does contain. Neither is taken on trust: the
# smallest countable pieces of the file are counted independently of the reader that read it,
# and every one of them has to end in a named class.


# An atom found in one of these places is not expected word for word in a unit, and why.
# "rewritten" is the class R1 added to the plan's four: an equation is not lost and not carried
# either - it is READ into the tool's linear notation, which is a third thing and is named as one.
EXPLAINED_BY_PLACE = {
    "equation": ("rewritten", "read into the tool's linear notation; what it was read from is named in the unit's equation"),
    "attribute": ("declared drop", "an attribute the rules do not read: an identifier, a style, a namespace or a file name"),
    "style or script": ("declared drop", "the content of a style or script element, which is not what the document says"),
    "page without a text layer": ("not read", "a page that carries no text layer, which the tool cannot count without reading the picture"),
    "tracked change": ("declared drop", "an unaccepted deletion, which is not what the document says: it is what somebody proposed the document should stop saying"),
    "header": ("declared drop", "the running header of a page, which the page carries because it is a page"),
    "footer": ("declared drop", "the running footer of a page, which the page carries because it is a page"),
    "comment": ("declared drop", "a comment somebody left on the document, which is not what the document says"),
    "whole file not read": ("not read", "a file in a format the tool does not read, named with its reason and a next step")}

# Below this many bytes a file may honestly hold no text; above it, a reader that found none at
# all and said nothing has probably not opened it. The account refuses to close over that.
VACUOUS_BYTES = 1024

def tokens(text):
    """The countable pieces of a text: runs without white space, after the same normalisation
    every chunk's text goes through. White space is not counted, because folding a list joins
    its items with a space and collapsing runs of space is a declared transform."""
    return [piece for piece in normalise_text(text or "").split(" ") if piece]

def bag(text):
    """The tokens of a text as counts, so that a word appearing twice must be found twice."""
    found = {}
    for piece in tokens(text):
        found[piece] = found.get(piece, 0) + 1
    return found

def bag_of(texts):
    """The tokens of several texts as one set of counts."""
    total = {}
    for text in texts:
        for piece, count in bag(text).items():
            total[piece] = total.get(piece, 0) + count
    return total

def minus(left, right):
    """What is left of `left` after taking away as much of `right` as it holds."""
    rest = {}
    for piece, count in left.items():
        keep = count - right.get(piece, 0)
        if keep > 0:
            rest[piece] = keep
    return rest

def atom(place, locator, text):
    """One countable piece of an input: where in the file it came from, and what it says."""
    return {"place": place, "locator": locator, "text": text}

MARKUP_METADATA = ("style", "script")

def atoms_of_markup(root, file_name, rules=None):
    """Every text node and every tail under every element, and every attribute value, each with
    the place it was found in. Place matters to the account: text inside an equation is read
    into the tool's linear notation rather than kept word for word, an attribute the rules do not
    read is metadata about the document rather than something the document says, and the
    content of a style or script element is not prose at all. Naming the place is what lets
    each of those be explained by a rule instead of counted as a loss."""
    families = (rules or {}).get("family_of", {})
    carriers = set((rules or {}).get("heading_attributes", ())) | set((rules or {}).get("numbering_attributes", ()))
    found, position = [], 0
    def walk(element, inside):
        nonlocal position
        position += 1
        name = local_name(element.tag)
        family = families.get(name) or inside
        where = "%s element %d <%s>" % (file_name, position, name)
        place = "equation" if family == "equation" else "style or script" if name in MARKUP_METADATA else "body"
        if (element.text or "").strip():
            found.append(atom(place, where, element.text))
        for child in element:
            walk(child, family if family in ("equation", "figure") else "")
            if (child.tail or "").strip():
                found.append(atom(place, where + " (tail)", child.tail))
        for attribute_name, value in sorted(element.attrib.items()):
            if value.strip():
                found.append(atom("attribute" if local_name(attribute_name) not in carriers else "body",
                                  "%s @%s" % (where, local_name(attribute_name)), value))
    walk(root, "")
    return found

DOCX_PARTS = (("word/document.xml", "body"), ("word/footnotes.xml", "footnote"), ("word/endnotes.xml", "endnote"),
              ("word/comments.xml", "comment"))
WORD_TEXT = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
WORD_DELETED = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}delText"

def atoms_of_docx(archive, parse, file_name):
    """Every run of text in a Word file, wherever Word put it. Text boxes sit inside the body
    part and are counted with it; footnotes, endnotes, comments, headers and footers are parts
    of their own and are counted here even where the reader does not yet read them, so that
    they show as unaccounted rather than vanishing without a word."""
    found = []
    named = list(DOCX_PARTS) + [(name, "header" if "header" in name else "footer")
                                for name in sorted(archive.namelist())
                                if name.startswith("word/header") or name.startswith("word/footer")]
    for part, place in named:
        if part not in archive.namelist():
            continue
        try:
            root = parse(archive.read(part))
        except Exception:                                # a part that will not parse is counted and named here
            found.append(atom(place, "%s %s" % (file_name, part), ""))
            continue
        for position, node in enumerate(root.iter(), start=1):
            if node.tag in (WORD_TEXT, WORD_DELETED) and (node.text or "").strip():
                kind = "tracked change" if node.tag == WORD_DELETED else place
                found.append(atom(kind, "%s %s run %d" % (file_name, part, position), node.text))
    return found

def atoms_of_pdf(document, file_name):
    """Every word the text layer of every page yields. A page with no text layer is counted as
    one atom with no text, so that a scanned page is visible in the account as a page that
    carries something the tool cannot count rather than as a page that carries nothing."""
    found = []
    for number, page in enumerate(document.pages, start=1):
        text = page.extract_text() or ""
        found.append(atom("body" if text.strip() else "page without a text layer", "%s p.%d" % (file_name, number), text))
    return found

def atoms_of_plain_text(text, file_name):
    """Every non-blank line of a plain text file."""
    return [atom("body", "%s line %d" % (file_name, number), line)
            for number, line in enumerate(text.split("\n"), start=1) if line.strip()]

# Transforms that put a word into a unit without taking it from an atom, or move it out of a
# unit's text into another of its fields. Each is named here, once. The account closes BECAUSE
# of this list: take an entry away and the words it explains are reported as injected or lost.
# Enforces: R13

def kept_and_relocated(chunks):
    """Two bags: what the units say in their text, and what they keep in their other fields.

    A relocated text is CARRIED, not copied: one heading of the file appears in the heading
    chain of every unit below it, and the file holds it once. So the fields that repeat by
    design - the heading chain and the section numbering - are counted once for each distinct
    text, while a caption, a table and an equation belong to their own unit and are counted
    every time they occur."""
    kept, moved, carried = [], [], set()
    for chunk in chunks:
        kept.append(chunk.get("text") or "")
        carried.update(chunk.get("heading_chain") or ())
        carried.add(chunk.get("numbering") or "")
        moved.append(chunk.get("caption") or "")
        moved.append(chunk.get("para_label") or "")
        table = chunk.get("table")
        if table:
            moved.extend(str(cell) for row in (list(table.get("header") or []),) + tuple(table.get("rows") or []) for cell in row)
        equation = chunk.get("equation")
        if equation:
            moved.append(equation.get("source_form") or "")
    return bag_of(kept), bag_of(moved + sorted(carried))


def marks_of_rendering(chunks):
    """Marks every reader makes, whatever the format: the form the tool renders a table, a figure or
    an equation in; the markup an equation was read from, whose tags are not words the document
    says; a number a document did not write that the tool counted back; and the place label the tool
    gives a paragraph of a PDF ("p.4 2") so that a person can find it again. Enforces: R13"""
    found = []
    for chunk in chunks:
        if chunk.get("kind") in ("Table", "Figure", "Equation"):
            found.append(chunk.get("text") or "")
        equation = chunk.get("equation")
        if equation:
            found.append(equation.get("source_form") or "")
            found.append(equation.get("linear") or "")
        found.append(chunk.get("numbering") or "")
        found.append(chunk.get("para_label") or "")
        if chunk.get("numbering_reconstructed"):
            found.extend(chunk.get("heading_chain") or ())
    return found

def account(file_name, atoms, chunks, dropped=(), marks=(), file_bytes=0):
    """The content account of one file. Every atom ends in exactly one class, and every token a
    unit shows comes from an atom or from a named mark. What a unit does not show word for word
    is explained by WHERE it was found (EXPLAINED_BY_PLACE) before it is called a loss, so that
    an equation read into linear notation and an identifier in an attribute are each named
    rather than swept into the same silence. Returns the counts, what could not be placed with
    where it was found, and what could not be explained. Enforces: R13"""
    source = bag_of([one["text"] for one in atoms if one["place"] not in EXPLAINED_BY_PLACE])
    kept, moved = kept_and_relocated(chunks)
    dropped_bag = bag_of(dropped)
    found = {"file": file_name, "in unit text": 0, "relocated": 0, "rewritten": 0, "declared drop": 0, "not read": 0}
    left = dict(source)
    for name, taken in (("in unit text", kept), ("relocated", moved), ("declared drop", dropped_bag)):
        used = minus(left, minus(left, taken))
        found[name] = sum(used.values())
        left = minus(left, used)
    for one in atoms:                                # what its place explains, and under which rule
        if one["place"] in EXPLAINED_BY_PLACE:
            name = EXPLAINED_BY_PLACE[one["place"]][0]
            found[name] = found[name] + max(1, len(tokens(one["text"])))
    refused = any(one["place"] == "whole file not read" for one in atoms)
    for chunk in chunks:
        if chunk.get("not_read_reason") and not refused:
            found["not read"] += 1
    placed = dict(kept)
    for extra_bag in (moved, dropped_bag):
        for piece, count in extra_bag.items():
            placed[piece] = placed.get(piece, 0) + count
    added = minus(minus(placed, source), bag_of(marks))
    found["atoms"] = sum(source.values()) + sum(found[name] for name in ("rewritten", "declared drop", "not read"))
    found["unaccounted"] = sum(left.values())
    found["injected"] = sum(added.values())
    found["where unaccounted"] = sorted(one["locator"] for one in atoms
                                        if one["place"] not in EXPLAINED_BY_PLACE
                                        and any(piece in left for piece in tokens(one["text"])))[:12]
    found["what unaccounted"] = sorted(left)[:24]
    found["what injected"] = sorted(added)[:24]
    # A file that is large, yielded no text at all, and was neither refused nor said to be
    # unreadable has probably not been opened. Its account would balance, because an account over
    # nothing balances: the one place the identity holds while everything is lost, so the one
    # place it refuses to close. Enforces: R13
    found["vacuous"] = (not refused and file_bytes > VACUOUS_BYTES and not sum(source.values())
                        and not any(chunk.get("not_read_reason") for chunk in chunks))
    found["file_bytes"] = file_bytes
    found["closed"] = found["unaccounted"] == 0 and found["injected"] == 0 and not found["vacuous"]
    return found

def account_lines(found):
    """The account of one file in the plain words an analyst reads on Model_Package_Info.
    Enforces: R10, R13"""
    lines = ["%d smallest pieces of text counted: %d kept in a unit, %d kept in a unit's other fields, "
             "%d read into another form, %d left out under a named rule, %d in a part that could not be read."
             % (found["atoms"], found["in unit text"], found["relocated"], found["rewritten"],
                found["declared drop"], found["not read"])]
    if found.get("vacuous"):
        lines.append("The file is %d KB and yielded no text at all, and no reader said it could not read it: "
                     "it has probably not been opened. Check the file itself." % max(1, found["file_bytes"] // 1024))
    if found.get("held as text only"):
        lines.append("%d of the lines inside a unit are held only as running text, in files the tool could not read as "
                     "code, data or a help page: they are kept, and nothing in them can be linked or checked."
                     % found["held as text only"])
    if found["closed"]:
        lines.append("Content account closed: nothing was lost and nothing was added.")
    if found["unaccounted"]:
        lines.append("%d piece(s) of text were read from the file but are in no unit and under no rule, "
                     "first at: %s." % (found["unaccounted"], "; ".join(found["where unaccounted"][:3]) or "not located"))
    if found["injected"]:
        lines.append("%d piece(s) of text are shown in a unit but were not found in the file: %s."
                     % (found["injected"], ", ".join(found["what injected"][:6])))
    return lines


def account_of_package(files, units, refused, is_text_file, dropped=()):
    """The content account of a package tarball. The atom of a package is a line: every
    non-blank line of every member that holds text has to lie inside a unit, be refused with a
    reason, or be named as not read. A member the tool cannot read as text (a compiled object, a
    picture, stored data in a binary form) is counted as one atom of its own, because its lines
    cannot be counted without reading it. Extends the line coverage that read_package already
    kept for parsed R files to every member of the tarball. A member left out on purpose (dropped: a help
    page, generated from the roxygen comments in the R files, which are read; a file of the package's tests)
    is a declared drop, every line of it, or one piece when it is not text. Enforces: R13"""
    inside, not_read, atoms, unaccounted, fenced, text_only, declared = 0, 0, 0, [], 0, 0, 0
    covered = {}
    for unit in units:
        lines = unit.get("lines")
        if unit.get("file") and lines:
            covered.setdefault(unit["file"], set()).update(range(int(lines[0]), int(lines[1]) + 1))
    for path in sorted(files):
        if path in dropped:                           # not read on purpose, under a named rule
            lines = 1 if not is_text_file(path) else sum(1 for line in files[path].decode("utf-8", "replace").split("\n") if line.strip())
            atoms, declared = atoms + lines, declared + lines
            continue
        if not is_text_file(path):
            atoms += 1
            if any(unit.get("file") == path for unit in units):
                inside += 1
            else:
                not_read += 1
            continue
        try:
            lines = files[path].decode("utf-8", "replace").split("\n")
        except Exception:
            atoms += 1
            not_read += 1
            continue
        numbers = [number for number, line in enumerate(lines, start=1) if line.strip()]
        # A fence that opens or closes a block of code in a vignette separates content from
        # content and says nothing itself. It is a declared drop, named here, rather than a
        # line that lies in no unit for no stated reason. Enforces: R13
        fences = {number for number in numbers if lines[number - 1].strip().startswith("```")}
        numbers = [number for number in numbers if number not in fences]
        atoms += len(numbers) + len(fences)
        fenced += len(fences)
        held = covered.get(path, set())
        whole_file = any(unit.get("file") == path and not unit.get("lines") for unit in units)
        if any(unit.get("file") == path and unit.get("kind") == "Other file" for unit in units):
            text_only += len(numbers)
        for number in numbers:
            if number in held or whole_file:
                inside += 1
            elif any(unit.get("file") == path and unit.get("read_problem") for unit in units):
                not_read += 1
            else:
                unaccounted.append("%s line %d" % (path, number))
    found = {"file": "the package", "atoms": atoms + len(refused), "in unit text": inside, "relocated": 0,
             "rewritten": 0, "declared drop": len(refused) + fenced + declared, "not read": not_read,
             "unaccounted": len(unaccounted), "injected": 0, "where unaccounted": sorted(unaccounted)[:12],
             "what unaccounted": [], "what injected": [], "held as text only": text_only}
    found["closed"] = found["unaccounted"] == 0
    return found


# ================================================================================================
# ---------------------------------------------------------------- what a file is
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
IMAGE_MAGIC = (b"\x89PNG", b"\xff\xd8\xff", b"GIF8", b"BM", b"II*\x00", b"MM\x00*", b"RIFF")
MARKDOWN_NAMES = (".md", ".markdown", ".mdown", ".mkd")
DELIMITED_NAMES = (".csv", ".tsv", ".tab")
LATEX_NAMES = (".tex", ".ltx")

# A format the tool does not read, what it is in plain words, and what to do about it.
NOT_READ = {
    "pptx": "a slide deck, which the tool does not read; save it as PDF and put the PDF in its place",
    "odf": "an OpenDocument file, which the tool does not read; save it as .docx and put that in its place",
    "epub": "an e-book, which the tool does not read; save it as PDF and put the PDF in its place",
    "zip": "an archive of other files; unpack it into the folder so that each file is read on its own",
    "gzip": "a compressed file; unpack it into the folder so that the file inside it is read",
    "ole": "in an old Microsoft Office format (.doc, .xls or .ppt), which the tool does not read; "
           "save it as .docx, .xlsx or PDF and put that in its place",
    "image": "a picture, and the tool does not read words from a picture on its own; if it holds text, "
             "save it as a PDF with a text layer",
    "binary": "binary data rather than a document"}

# The name a format goes by in what the analyst reads.
FORMAT_NAMES = {"pdf": "PDF", "docx": "Word", "xlsx": "spreadsheet", "svg": "SVG picture", "mhtml": "web archive", "html": "web page",
                "xml": "XML", "markdown": "Markdown", "delimited": "delimited rows", "rtf": "RTF",
                "latex": "LaTeX", "text": "plain text"}

def looks_binary(data):
    """Bytes that are not text: a NUL byte, or more than one in ten bytes a control character."""
    head = data[:4096]
    if b"\x00" in head:
        return True
    control = sum(1 for byte in head if byte < 32 and byte not in (9, 10, 12, 13))
    return bool(head) and control * 10 > len(head)

def sniff_zip(data):
    """What a ZIP archive holds, which is what it is. Enforces: R6"""
    try:
        names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
    except (zipfile.BadZipFile, OSError):
        return "binary"
    if "word/document.xml" in names:
        return "docx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    if "ppt/presentation.xml" in names:
        return "pptx"
    if "content.xml" in names:
        return "odf"
    if "META-INF/container.xml" in names:
        return "epub"
    return "zip"

def detect_format(data, file_name=""):
    """The format of a file from its content, whatever its name says. For the three text formats
    whose content cannot be told apart for certain - Markdown, delimited rows and LaTeX - the name
    is taken as a hint and the content has to agree with it. Enforces: R6"""
    head = data[:4096].lstrip(b"\xef\xbb\xbf \t\r\n")
    if head.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        return sniff_zip(data)
    if data.startswith(OLE_MAGIC):
        return "ole"
    if head.startswith(b"{\\rtf"):
        return "rtf"
    if data.startswith(b"\x1f\x8b"):
        return "gzip"
    if data.startswith(IMAGE_MAGIC):
        return "image"
    if looks_binary(data):
        return "binary"
    lowered, name = head.lower(), file_name.lower()
    if name.endswith(MARKDOWN_NAMES):
        return "markdown"                            # Markdown may open with an HTML comment or table; the name decides
    if lowered.startswith((b"mime-version:", b"from:", b"content-type:")) or b"multipart/related" in lowered[:1024]:
        return "mhtml"
    if lowered.startswith(b"<"):
        if re.match(rb"<(!doctype\s+html|html)\b", lowered):
            return "html"
        if re.search(rb"<svg[\s>]", lowered[:4096]) and not re.search(rb"<(body|para|section|document|p)\b", lowered[:4096]):
            return "svg"
        return "xml"
    sample = decode_text(data[:16384])
    if name.endswith(LATEX_NAMES) and re.search(r"\\[a-zA-Z]+", sample) or re.search(r"\\documentclass|\\begin\{document\}", sample):
        return "latex"
    if len(re.findall(r"(?m)^#{1,6} \S", sample)) >= 2:
        return "markdown"
    if name.endswith(DELIMITED_NAMES) and delimiter_of(sample):
        return "delimited"
    return "text"

def decode_text(data):
    """Bytes to text: a byte-order mark or a declared encoding decides, then UTF-8, then Latin-1."""
    declared = re.search(rb'(?:encoding|charset)=["\']?([\w-]+)', data[:2048])
    for encoding in ([declared.group(1).decode("ascii")] if declared else []) + ["utf-8-sig", "utf-8"]:
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("latin-1")

# ---------------------------------------------------------------- which files of a folder are documents
LEFT_BEHIND = re.compile(r"^(~\$.*|\.~lock\..*#|thumbs\.db|desktop\.ini|ehthumbs\.db|.*\.tmp|.*\.bak|.*\.swp|\.ds_store)$", re.I)

def input_files(folder):
    """The documents of one input folder, walking into its folders, in a fixed order; and every
    file left out, with why. An operating system's and an editor's leavings (a Word lock file,
    Thumbs.db) are left out, and so is the support folder a browser writes beside a saved web
    page, whose pictures and style sheets are part of the page and not documents of their own.
    Nothing is left out without being named. Enforces: R2"""
    found, skipped = [], []
    if not os.path.isdir(folder):
        return found, skipped
    for root, folders, names in os.walk(folder):
        pages = {os.path.splitext(name)[0].lower() for name in names if name.lower().endswith((".htm", ".html", ".mht", ".mhtml"))}
        for inner in sorted(folders):
            stem = re.sub(r"(_files|\.files|-files)$", "", inner, flags=re.I).lower()
            if stem != inner.lower() and stem in pages:
                skipped.append((os.path.relpath(os.path.join(root, inner), folder), "the support folder of a saved web page"))
        folders[:] = sorted(inner for inner in folders if not inner.startswith((".", "__"))
                            and not (re.sub(r"(_files|\.files|-files)$", "", inner, flags=re.I).lower() in pages
                                     and re.search(r"(_files|\.files|-files)$", inner, re.I)))
        for name in sorted(names):
            relative = os.path.relpath(os.path.join(root, name), folder)
            if name == "README.txt" and root == folder:
                continue
            if name.startswith(".") or LEFT_BEHIND.match(name):
                skipped.append((relative, "a file an operating system or an editor leaves behind"))
            else:
                found.append(os.path.join(root, name))
    return sorted(found, key=lambda path: os.path.relpath(path, folder).lower()), sorted(skipped)

# ---------------------------------------------------------------- formats the tool reads by converting them
# Each converter returns (markup, words): the markup is what the walker reads, and the words are
# the file with its syntax taken away by a rule named here, which is what the content account
# counts. The two are made by DIFFERENT code - the markup by reading the structure, the words by
# removing syntax - so that the account still sees the walker lose or add a word. It does trust
# the syntax rule, and says so. Enforces: R13

def element(tag, text):
    """One element of markup, its text escaped so that a < in a document is never a tag."""
    return "<%s>%s</%s>" % (tag, escape(text), tag)

def markdown_inline(text):
    """A line of Markdown without its inline marks: emphasis, code marks, and the target of a link
    or a picture, whose words are kept and whose address is not what the document says."""
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"(\*\*|__)(?=\S)(.+?)(?<=\S)\1", r"\2", text)
    text = re.sub(r"(?<![\w*])\*(?=\S)(.+?)(?<=\S)\*(?![\w*])", r"\1", text)
    return text.replace("`", "")

TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
LIST_MARK = re.compile(r"^\s*([-*+]|\d+[.)])\s+")

HTML_COMMENT = re.compile(r"<!--.*?-->", re.S)

def markdown_to_markup(text):
    """Markdown as markup: # headings, paragraphs, lists, tables, fenced code and quotations. An
    HTML comment is left out by the same rule in the markup and in the words: it is a note to
    whoever edits the file, not what the document says."""
    text = HTML_COMMENT.sub("", text)
    out, paragraph, rows, fenced, items = [], [], [], None, []
    def flush():
        if paragraph:
            out.append(element("p", " ".join(paragraph)))
            paragraph.clear()
        if items:
            out.append("<ul>%s</ul>" % "".join(element("li", item) for item in items))
            items.clear()
        if rows:
            cells = [[markdown_inline(cell.strip()) for cell in row.strip().strip("|").split("|")] for row in rows]
            out.append("<table>%s</table>" % "".join(
                "<tr>%s</tr>" % "".join(element("th" if number == 0 else "td", cell) for cell in row)
                for number, row in enumerate(cells)))
            rows.clear()
    lines = text.split("\n")
    for position, line in enumerate(lines):
        if fenced is not None:
            if line.strip().startswith("```") or line.strip().startswith("~~~"):
                out.append(element("pre", "\n".join(fenced)))
                fenced = None
            else:
                fenced.append(line)
            continue
        if line.strip().startswith(("```", "~~~")):
            flush()
            fenced = []
            continue
        heading = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", line)
        underline = position + 1 < len(lines) and re.match(r"^(=+|-+)\s*$", lines[position + 1]) and line.strip() and not rows
        if heading or underline:
            flush()
            level = len(heading.group(1)) if heading else (1 if lines[position + 1].strip()[0] == "=" else 2)
            out.append(element("h%d" % min(level, 4), markdown_inline((heading.group(2) if heading else line).strip())))
            continue
        if re.match(r"^(=+|-+)\s*$", line) and position and out and out[-1].startswith("<h"):
            continue                                 # the underline of a heading already read
        if "|" in line and (rows or position + 1 < len(lines) and TABLE_RULE.match(lines[position + 1])):
            if not rows:
                flush()
            if not TABLE_RULE.match(line):
                rows.append(line)
            continue
        if LIST_MARK.match(line):
            if paragraph or rows:
                flush()
            items.append(markdown_inline(LIST_MARK.sub("", line).strip()))
            continue
        if not line.strip():
            flush()
            continue
        if items and line.startswith((" ", "\t")):
            items[-1] += " " + markdown_inline(line.strip())
            continue
        if rows or items:
            flush()
        paragraph.append(markdown_inline(re.sub(r"^\s*>\s?", "", line).strip()))
    if fenced is not None:
        out.append(element("pre", "\n".join(fenced)))
    flush()
    return "<document>%s</document>" % "".join(out)

def markdown_words(text):
    """The words of a Markdown file, with its syntax taken away by rule rather than by reading its
    structure: heading marks, list marks, table rules and bars, fence lines, quotation marks."""
    kept = []
    for line in HTML_COMMENT.sub("", text).split("\n"):
        if TABLE_RULE.match(line) or line.strip().startswith(("```", "~~~")) or re.match(r"^(=+|-+)\s*$", line):
            continue
        line = re.sub(r"^\s*#{1,6}\s+|\s+#+\s*$", " ", line)
        line = LIST_MARK.sub(" ", line)
        line = re.sub(r"^\s*>\s?", " ", line)
        kept.append(markdown_inline(line.replace("|", " ")))
    return "\n".join(kept)

def delimiter_of(text):
    """The separator of a file of delimited rows, where the first rows agree on one."""
    rows = [line for line in text.split("\n") if line.strip()][:6]
    if len(rows) < 2:
        return ""
    for separator in ("\t", ",", ";", "|"):
        counts = {row.count(separator) for row in rows}
        if len(counts) == 1 and counts.pop() >= 1:
            return separator
    return ""

def delimited_rows(text):
    """The rows of a file of delimited values, read by the rules of that format."""
    separator = delimiter_of(text) or ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=separator) if any(cell.strip() for cell in row)]

def delimited_to_markup(text, file_name):
    """Rows of values as a table, the first row naming the columns."""
    rows = delimited_rows(text)
    body = "".join("<tr>%s</tr>" % "".join(element("th" if number == 0 else "td", cell.strip()) for cell in row)
                   for number, row in enumerate(rows))
    return "<document><table><caption>%s</caption>%s</table></document>" % (escape(os.path.basename(file_name)), body)

def delimited_words(text, file_name):
    """The words of a file of delimited values: its name, which heads its table, and its cells."""
    return os.path.basename(file_name) + "\n" + "\n".join(" ".join(cell.strip() for cell in row) for row in delimited_rows(text))

def spreadsheet_to_markup(data):
    """A spreadsheet as one heading and one table for every sheet that holds anything. What a cell
    shows is its stored value: a formula is never worked out, only the value the spreadsheet saved
    with it is read. Enforces: R7"""
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts, words = [], []
    for sheet in book.worksheets:
        rows = [[("" if value is None else str(value)).strip() for value in row] for row in sheet.iter_rows(values_only=True)]
        rows = [row for row in rows if any(row)]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        parts.append(element("h1", sheet.title))
        parts.append("<table>%s</table>" % "".join(
            "<tr>%s</tr>" % "".join(element("th" if number == 0 else "td", cell) for cell in row)
            for number, row in enumerate(rows)))
        words.append(sheet.title)
        words.extend(" ".join(row) for row in rows)
    return "<document>%s</document>" % "".join(parts), "\n".join(words)

def rtf_to_text(text):
    """The words of an RTF file: its tables of fonts, colours and styles, its pictures and its
    control words taken away, a paragraph mark read as a new paragraph and an escaped character as
    that character."""
    text = re.sub(r"\{\\\*[^{}]*(?:\{[^{}]*\}[^{}]*)*\}", "", text)
    for group in ("fonttbl", "colortbl", "stylesheet", "info", "pict", "listtable", "listoverridetable"):
        text = re.sub(r"\{\\%s[^{}]*(?:\{[^{}]*(?:\{[^{}]*\}[^{}]*)*\}[^{}]*)*\}" % group, "", text)
    text = re.sub(r"\\'([0-9a-fA-F]{2})", lambda found: bytes([int(found.group(1), 16)]).decode("cp1252", "replace"), text)
    text = re.sub(r"\\u(-?\d+)\??", lambda found: chr(int(found.group(1)) % 65536), text)
    text = re.sub(r"\\(par|line|sect|page)\b ?", "\n\n", text)
    text = re.sub(r"\\tab\b ?", " ", text)
    text = re.sub(r"\\([{}\\])", lambda found: "\x00" + found.group(1), text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    text = text.replace("{", "").replace("}", "").replace("\x00", "")
    return re.sub(r"\n{3,}", "\n\n", text).strip()

LATEX_HEADINGS = (("part", 1), ("chapter", 1), ("section", 1), ("subsection", 2), ("subsubsection", 3), ("paragraph", 4))

def latex_body(text):
    """A LaTeX file without its comments and its preamble. What comes before \\begin{document} -
    the class of the document and the packages it loads - sets the document up and is not what it
    says, so it is left out by this rule, the same rule for the markup and for the words."""
    text = re.sub(r"(?<!\\)%.*", "", text)
    body = re.search(r"\\begin\{document\}(.*?)(\\end\{document\}|$)", text, re.S)
    return body.group(1) if body else text

def latex_words(text, whole=True):
    """The words of a LaTeX file: comments, environments and the names of commands taken away,
    the arguments of a command kept. Mathematics is kept as its words and symbols; it is not read
    as an equation here."""
    text = latex_body(text) if whole else text
    text = re.sub(r"\\(begin|end)\{[^}]*\}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+\*?(\[[^\]]*\])?", " ", text)
    return re.sub(r"[{}$]", " ", text).replace("\\\\", " ")

def latex_to_markup(text):
    """LaTeX as markup: its sectioning commands as headings, its items as a list, and everything
    between as paragraphs of its words."""
    text = latex_body(text)
    pattern = r"\\(%s)\*?\{([^}]*)\}" % "|".join(name for name, _ in LATEX_HEADINGS)
    levels, out, cursor = dict(LATEX_HEADINGS), [], 0
    for found in list(re.finditer(pattern, text)) + [None]:
        chunk = text[cursor:found.start() if found else len(text)]
        for part in re.split(r"\n\s*\n", chunk):
            items = re.split(r"\\item\b", part)
            lead = latex_words(items[0], whole=False).strip()
            if lead:
                out.append(element("p", re.sub(r"\s+", " ", lead)))
            if len(items) > 1:
                out.append("<ul>%s</ul>" % "".join(element("li", re.sub(r"\s+", " ", latex_words(item, whole=False)).strip())
                                                  for item in items[1:] if latex_words(item, whole=False).strip()))
        if found:
            out.append(element("h%d" % levels[found.group(1)], re.sub(r"\s+", " ", latex_words(found.group(2), whole=False)).strip()))
            cursor = found.end()
    return "<document>%s</document>" % "".join(out)


def svg_texts(root):
    """Every piece of text an SVG holds as text, with where it is drawn: (x, y, text). Three places:
    a <text> (its <tspan>s split it only where they carry positions of their own; a <textPath> or an
    <a> inside it is read with it); and a <foreignObject>, where drawing tools put HTML - divs and
    paragraphs - instead of SVG text. A transform is not applied, so a moved group keeps the order
    it was written in."""
    found = []
    def number(value, fallback=0.0):
        try:
            return float(re.split(r"[ ,]", (value or "").strip())[0])
        except (ValueError, IndexError):
            return fallback
    for node in root.iter():
        name = local_name(node.tag)
        if name == "text":
            x, y = number(node.get("x")), number(node.get("y"))
            placed = [span for span in node.iter() if local_name(span.tag) == "tspan" and (span.get("x") or span.get("y") or span.get("dy"))]
            if not placed:
                words = normalise_text("".join(node.itertext()))
                if words:
                    found.append((x, y, words))
                continue
            lead = normalise_text(node.text or "")
            if lead:
                found.append((x, y, lead))
            line_y = y
            for span in placed:
                line_y = number(span.get("y"), line_y + number(span.get("dy"), 0.0))
                words = normalise_text("".join(span.itertext()))
                if words:
                    found.append((number(span.get("x"), x), line_y, words))
        elif name == "foreignobject":                    # local_name lower-cases: foreignObject
            words = normalise_text(" ".join(part for part in node.itertext() if part.strip()))
            if words:
                found.append((number(node.get("x")), number(node.get("y")), words))
    return found

def svg_embedded_pictures(root):
    """The raster pictures an SVG carries inside itself as data: URIs - what a chart exported as an
    image and wrapped in SVG looks like. Their words can only be read by OCR."""
    import base64
    pictures = []
    for node in root.iter():
        if local_name(node.tag) != "image":
            continue
        link = next((value for key, value in node.attrib.items() if local_name(key) == "href"), "")
        found = re.match(r"data:image/(png|jpe?g|gif|bmp|webp);base64,(.*)", link.strip(), re.S)
        if found:
            try:
                pictures.append(base64.b64decode(re.sub(r"\s", "", found.group(2))))
            except (ValueError, TypeError):
                continue
    return pictures

def svg_why_no_text(root):
    """Why an SVG gave no text, in words an analyst can act on. Enforces: R2"""
    if svg_embedded_pictures(root):
        return "the chart is a picture stored inside the SVG, so its words can only be read by OCR"
    glyphs = sum(1 for node in root.iter() if local_name(node.tag) == "use"
                 and any("glyph" in (value or "").lower() for key, value in node.attrib.items() if local_name(key) == "href"))
    shapes = sum(1 for node in root.iter() if local_name(node.tag) in ("path", "use"))
    if glyphs or shapes > 40:
        return ("its letters are drawn as shapes rather than stored as text, so they cannot be read as text; "
                "export the chart with its text kept as text, or add its data as a table in the methodology")
    return "it holds no text"

def svg_rows(texts, tolerance=None):
    """Texts grouped into rows by their y position, each row sorted by x. The tolerance is a
    share of the median line height, so a chart's labels and a table's cells both group."""
    if not texts:
        return []
    ys = sorted({round(y, 1) for _, y, _ in texts})
    gaps = [b - a for a, b in zip(ys, ys[1:]) if b - a > 0.5]
    tolerance = tolerance or (max(2.0, sorted(gaps)[len(gaps) // 2] * 0.4) if gaps else 2.0)
    rows, current, last_y = [], [], None
    for x, y, words in sorted(texts, key=lambda t: (t[1], t[0])):
        if last_y is not None and y - last_y > tolerance:
            rows.append(sorted(current))
            current = []
        current.append((x, y, words))
        last_y = y
    if current:
        rows.append(sorted(current))
    return [[words for _, _, words in row] for row in rows]

def svg_to_markup(data, file_name):
    """An SVG as markup the walker reads: its <title> and <desc> as the caption; its text, where
    it stands in a grid of at least two rows of the same width, as a table with the first row
    as the header; otherwise as a figure whose words are the labels of the picture in reading
    order. An SVG holds text as text, so nothing is read from a picture by guesswork: every
    word comes from a <text> element. Enforces: R7, R13"""
    import xml.etree.ElementTree as ElementTree
    root = ElementTree.fromstring(data)
    caption = " ".join(normalise_text("".join(node.itertext())) for node in root
                       if local_name(node.tag) in ("title", "desc") and normalise_text("".join(node.itertext())))
    rows = svg_rows(svg_texts(root))
    widths = {len(row) for row in rows}
    words = [caption] + [" ".join(row) for row in rows]
    if len(rows) >= 2 and len(widths) == 1 and widths.pop() >= 2:
        body = "".join("<tr>%s</tr>" % "".join(element("th" if number == 0 else "td", cell) for cell in row)
                       for number, row in enumerate(rows))
        markup = "<document><table><caption>%s</caption>%s</table></document>" % (escape(caption or os.path.basename(file_name)), body)
        return markup, "\n".join(words), "a picture whose text stands in a grid, read as a table"
    labels = " ".join(" ".join(row) for row in rows)
    # no src: this markup IS the picture, so it must not send the reader looking for itself beside itself
    markup = "<document><figure alt=\"%s\"><caption>%s</caption></figure></document>" % (escape(labels), escape(caption))
    return markup, "\n".join(words), "a picture, read by its own text: %d label(s)" % sum(len(row) for row in rows)

def converted(found, data, file_name):
    """A file in a format read by converting it: (markup, words, what was done, in plain words)."""
    if found == "svg":
        return svg_to_markup(data, file_name)
    if found == "xlsx":
        markup, words = spreadsheet_to_markup(data)
        return markup, words, "read sheet by sheet, each sheet a heading over a table of its stored values"
    text = decode_text(data)
    if found == "markdown":
        return markdown_to_markup(text), markdown_words(text), "read as Markdown: its headings, lists and tables kept as such"
    if found == "delimited":
        return (delimited_to_markup(text, file_name), delimited_words(text, file_name),
                "read as a table of values, one row per line, the first line naming the columns")
    if found == "latex":
        return latex_to_markup(text), latex_words(text), "read as LaTeX: its sections as headings; its mathematics kept as words, not read as equations"
    plain = rtf_to_text(text)
    return ("<document>%s</document>" % "".join(element("p", part.strip()) for part in re.split(r"\n\s*\n", plain) if part.strip()),
            plain, "read as RTF: its control words, font tables and pictures left out")

CONVERTED = ("xlsx", "markdown", "delimited", "latex", "rtf", "svg")

def reason_for(problem):
    """Why a reader failed on a file, in words an analyst can act on. Enforces: R2"""
    name, said = type(problem).__name__, str(problem).lower()
    if isinstance(problem, IsADirectoryError):
        return "it is a folder, not a file"
    if isinstance(problem, PermissionError):
        return "the tool is not allowed to open it; check who may read the file"
    if "password" in said or "encrypt" in said or "decrypt" in said or "pdfminer" in name.lower():
        return "it is protected by a password; save an unprotected copy and put that in its place"
    if isinstance(problem, zipfile.BadZipFile) or name in ("ReadError", "CompressionError", "HeaderError", "EOFError"):
        return "it is damaged, or is not the archive its name says; save it again and put the new copy in its place"
    return "a reader failed on it (%s); the rest of the run went on without it" % name

# ---------------------------------------------------------------- the three input folders of a project
INPUT_FOLDERS = (
    ("methodology", "1_Methodology", "Put the canonical methodology here: XML (also inside a .txt), .mhtml, .docx, "
     ".pdf, Markdown, .csv, .xlsx, .rtf or .tex. Folders are read too, in name order; anything the tool cannot read is named."),
    ("package", "2_Model_Package", "Put the R package here: its tarball (.tar.gz), a .zip of it, or its source folder."),
    ("documentation", "3_Model_Documentation", "Put the model documentation here (.docx is preferred; .pdf, .mhtml, "
     "XML, Markdown, .csv, .xlsx, .rtf and .tex are read too). Folders are read too, in name order."))

def list_input_files(inputs_dir):
    """The input files of a project, by corner, walking into folders; the folder each corner is
    read from, so that a file can be named by its path inside it; and every file each corner left
    out, with why. The optional glossary and tag rules sit beside the three folders. Enforces: R2"""
    inputs = {"glossary": None, "tag_rules": None, "roots": {}, "skipped": {}}
    for corner, folder, _ in INPUT_FOLDERS:
        inputs["roots"][corner] = os.path.join(inputs_dir, folder)
        inputs[corner], inputs["skipped"][corner] = input_files(inputs["roots"][corner])
    for key, name in (("glossary", "glossary.xlsx"), ("tag_rules", "tag_rules.yaml")):
        if os.path.exists(os.path.join(inputs_dir, name)):
            inputs[key] = os.path.join(inputs_dir, name)
    return inputs


# ================================================================================================
# reading the methodology, the documentation and the model package
# ================================================================================================
# ---------------------------------------------------------------- the methodology and the documentation, read into units
class NotReadable(Exception):
    """A formula or a file that the tool cannot read. The message is a plain reason for the analyst."""

# ---------------------------------------------------------------- the linear-notation parser
SUPERSCRIPTS = {"\u207b\u00b9": "^-1", "\u00b2": "^2", "\u00b3": "^3", "\u00b9": "^1"}
SIGNS = {"\u2212": "-", "\u00d7": "*", "\u00b7": "*", "\u22c5": "*", "\u2217": "*", "\u00f7": "/",
         "\u2264": "<=", "\u2265": ">=", "**": "^", "\u2061": "", "\u2062": "*", "\u2009": " "}
_TOKEN_RE = re.compile(
    r"\s*(?:(?P<number>\d+(?:\.\d+)?(?:[eE][-+]?\d+)?)(?P<percent>\s?%)?"
    r"|(?P<name>[^\W\d_][\w]*(?:\.[^\W\d_]\w*)*(?:_\{[^{}]+\}|\[[^\[\]]+\])?)"
    r"|(?P<sign><=|>=|[-+*/^(),=<>\u221a]))")

# ---------------------------------------------------------------- reference data of the reader
# The tag rules: which tag of a document is what (heading, paragraph, table, ...), the numbering
# schemes. An analyst's tag_rules.yaml, beside the project's input folders, is laid
# over these for one project and always wins. Kept as YAML text and parsed on every call, so that a
# caller that changes the rules it was given changes only its own copy.
TAG_RULES_YAML = r'''# tag_rules.yaml - which tag belongs to which family when the tool reads XML or HTML.
# Reviewer 1 owns this file. A project can override any part of it with its own tag_rules.yaml
# (same layout; a family given there replaces the family given here).
# Tag names are compared in lower case and without their namespace prefix.
# A tag that is in no family is read as a paragraph (or, when it only wraps other blocks,
# as a container) and is reported on Model_Package_Info, so that it can be added here.
families:
  heading:        [h1, h2, h3, h4, h5, h6, title, heading, head-line, sectiontitle]
  container:      [html, body, div, section, sect, sect1, sect2, sect3, sect4, sect5, chapter, part,
                   article, document, doc, annex, appendix, subsection, subsubsection, main, document-root]
  paragraph:      [p, para, paragraph, text, blockquote, pre, dd, dt]
  list_container: [ul, ol, list, itemizedlist, orderedlist, dl]
  list_item:      [li, item, listitem]
  table:          [table, informaltable, tbl]
  table_part:     [thead, tbody, tfoot, tgroup, colgroup, col]
  row:            [tr, row]
  header_cell:    [th]
  cell:           [td, entry, cell]
  figure:         [img, image, figure, graphic, mediaobject, svg, object, chart, imagedata]
  equation:       [math, equation, omath, omathpara, formula, informalequation]
  caption:        [caption, figcaption, legend]
  inline:         [a, b, i, u, em, strong, span, font, sub, sup, br, code, tt, small, big, emphasis, xref, o:p]
  ignore:         [head, script, style, meta, link, toc, index, nav, xml]
# Attributes that hold the numbering of a heading as written ("3.1").
numbering_attributes: [number, num, label, n]
# Numbering schemes, tried in this order. The first one that matches the start of a heading
# names its scheme. "dotted" numbers give their depth directly; every other scheme gets the
# level at which it first appeared (one deeper than the heading before it).
numbering_schemes:
  - {name: annex,          pattern: '^(?:Annex|Appendix|Annexe|Anhang)\s+([A-Z0-9]+)[.:]?(?=\s|$)'}
  - {name: dotted,         pattern: '^(\d+(?:\.\d+)+)\.?(?=\s|$)'}
  - {name: number,         pattern: '^(\d+)[.)]?(?=\s|$)'}
  - {name: upper_roman,    pattern: '^([IVXLC]+)\.(?=\s|$)'}
  - {name: upper_letter,   pattern: '^([A-Z])[.)](?=\s|$)'}
  - {name: bracket_roman,  pattern: '^\(([ivxlc]+)\)(?=\s|$)'}
  - {name: bracket_letter, pattern: '^\(([a-z])\)(?=\s|$)'}
# Words that start a cross-reference as written ("see Table 3", "section 4.2").
cross_reference_labels: [Table, Figure, Section, Sections, Equation, Annex, Appendix, Paragraph, Chapter]
'''

# The notation of R's mathematical functions: how each is written in the tool's linear notation and
# which of its arguments are the operands. Shared with review.py, which reads it for its word lists.
R_FUNCTION_MAP_YAML = r'''# r_function_map.yaml - how names in R code and in written formulas map to the tool's neutral
# function names. Reviewers 1, 2 and 4 read this file.
#
# notation: names that the tool reads as a FUNCTION when a written formula shows name( ... ).
#   Any other name followed by "(" could be a product or a function; the tool does not guess and
#   marks the formula "could not be read".
notation:
  functions:
    N: normal_cdf
    "\u03a6": normal_cdf
    Phi: normal_cdf
    ln: log
    log: log
    exp: exp
    sqrt: sqrt
    max: max
    min: min
    abs: abs
    sum_over: sum_over
    piecewise: piecewise
  inverse:                      # name^-1( ... ) is read as the inverse function
    normal_cdf: normal_inverse
# r_functions: R function -> neutral name, with the argument order the tool relies on and the
#   domain in which the neutral function can be evaluated (used to draw valid sample points).
r_functions:
  pnorm:  {neutral: normal_cdf,     arguments: [q],      only_defaults: [mean, sd, lower.tail, log.p]}
  qnorm:  {neutral: normal_inverse, arguments: [p],      only_defaults: [mean, sd, lower.tail, log.p]}
  exp:    {neutral: exp,            arguments: [x]}
  log:    {neutral: log,            arguments: [x],      only_defaults: [base]}
  log1p:  {neutral: log1p,          arguments: [x]}
  expm1:  {neutral: expm1,          arguments: [x]}
  sqrt:   {neutral: sqrt,           arguments: [x]}
  abs:    {neutral: abs,            arguments: [x]}
  max:    {neutral: max,            variadic: true}
  min:    {neutral: min,            variadic: true}
  pmax:   {neutral: max,            variadic: true}
  pmin:   {neutral: min,            variadic: true}
  ifelse: {neutral: piecewise,      arguments: [test, "yes", "no"]}
  sum:    {neutral: sum_over,       variadic: true}
  round:  {neutral: round,          arguments: [x, digits], optional: 1}
  floor:  {neutral: floor,          arguments: [x]}
  ceiling: {neutral: ceiling,       arguments: [x]}
# domains: where each neutral function has a value (open intervals; null means unbounded).
domains:
  normal_cdf:     {argument: [null, null]}
  normal_inverse: {argument: [0, 1]}
  log:            {argument: [0, null]}
  log1p:          {argument: [-1, null]}
  sqrt:           {argument: [0, null]}
  exp:            {argument: [null, 50]}
'''

def load_notation():
    """The names the tool reads as functions in a written formula."""
    return yaml.safe_load(R_FUNCTION_MAP_YAML)["notation"]

def tokenize_formula(text):
    """Cut a written formula into numbers, names and signs. Anything else makes it unreadable."""
    for written, plain in list(SUPERSCRIPTS.items()) + list(SIGNS.items()):
        text = text.replace(written, plain)
    text = re.sub(r"_\{([^{}]*)\}", lambda found: "_" + re.sub(r"\W", "", found.group(1)), text)   # x_{i,j} -> x_ij
    tokens, position = [], 0
    text = text.strip()
    while position < len(text):
        match = _TOKEN_RE.match(text, position)
        if not match or match.end() == position:
            raise NotReadable("it contains '%s', which the tool does not read in a formula" % text[position:position + 12].strip())
        if match.group("number"):
            value = Decimal(match.group("number"))
            tokens.append(("number", plain_decimal(value / 100 if match.group("percent") else value)))
        elif match.group("name"):
            tokens.append(("name", match.group("name")))
        else:
            tokens.append(("sign", match.group("sign")))
        position = match.end()
    return tokens

class FormulaReader:
    """A small recursive-descent reader over the tokens of one formula. Order of strength, from
    weakest: equals, comparison, plus and minus, times and divide, a leading minus, power.
    Two terms side by side are read as a product only when `implicit_product` is set, which
    is done for equation markup (its extent is exact) and never for running text."""
    def __init__(self, tokens, notation, implicit_product=False):
        self.tokens, self.position, self.notation = tokens, 0, notation
        self.functions, self.implicit_product = notation.get("functions", {}), implicit_product

    def peek(self, offset=0):
        """The token at the reading position (or further on), without taking it."""
        index = self.position + offset
        return self.tokens[index] if index < len(self.tokens) else ("end", "")

    def take(self, sign=None):
        """Take the next token; with `sign`, insist that it is that sign."""
        token = self.peek()
        if sign is not None and token != ("sign", sign):
            raise NotReadable("a '%s' was expected where '%s' stands" % (sign, token[1] or "the end"))
        self.position += 1
        return token

    def statement(self):
        """A whole formula: an expression, or `left = right`."""
        left = self.comparison()
        if self.peek() == ("sign", "="):
            self.take()
            left = Expr("eq", args=(left, self.comparison()))
            if self.peek() == ("sign", "="):
                raise NotReadable("it has more than one equals sign")
        return left

    def comparison(self):
        """An expression, optionally compared with another one."""
        left = self.sum()
        if self.peek()[0] == "sign" and self.peek()[1] in ("<", ">", "<=", ">="):
            sign = self.take()[1]
            left = Expr("cmp", name=sign, args=(left, self.sum()))
        return left

    def sum(self):
        """Terms joined by + and -, from left to right."""
        left = self.product()
        while self.peek() in (("sign", "+"), ("sign", "-")):
            op = "add" if self.take()[1] == "+" else "sub"
            left = Expr(op, args=(left, self.product()))
        return left

    def product(self):
        """Factors joined by * and /, from left to right; juxtaposition only where the source allows it."""
        left = self.unary()
        while True:
            if self.peek() in (("sign", "*"), ("sign", "/")):
                op = "mul" if self.take()[1] == "*" else "div"
                left = Expr(op, args=(left, self.unary()))
            elif self.implicit_product and self.term_follows():
                left = Expr("mul", args=(left, self.power()))
            else:
                return left

    def term_follows(self):
        """Does another factor start here without a sign in between?"""
        kind, text = self.peek()
        return kind in ("number", "name") or (kind == "sign" and text in ("(", "\u221a"))

    def unary(self):
        """A leading minus or plus. It binds less tightly than a power, as in mathematics and in R."""
        if self.peek() == ("sign", "-"):
            self.take()
            return Expr("neg", args=(self.unary(),))
        if self.peek() == ("sign", "+"):
            self.take()
            return self.unary()
        return self.power()

    def power(self):
        """A base with an optional power (right-associative), an inverse-function mark or a percent sign."""
        base = self.atom()
        if self.peek() == ("sign", "^"):
            self.take()
            base = Expr("pow", args=(base, self.unary()))
        if self.term_follows() and not self.implicit_product:
            raise NotReadable("two terms stand side by side without a sign between them, which could "
                              "mean a product or something else; the tool does not guess")
        return base

    def minus_one_follows(self):
        """Is the next thing ^-1 or ^(-1) followed by an opening bracket? Returns tokens to skip."""
        shapes = ([("sign", "^"), ("sign", "-"), ("number", "1"), ("sign", "(")],
                  [("sign", "^"), ("sign", "("), ("sign", "-"), ("number", "1"), ("sign", ")"), ("sign", "(")])
        for shape in shapes:
            if [self.peek(i) for i in range(len(shape))] == shape:
                return len(shape) - 1
        return 0

    def atom(self):
        """A number, a symbol, a function call or a bracketed expression."""
        kind, text = self.take()
        if kind == "number":
            return Expr("num", value=text)
        if kind == "sign" and text == "(":
            inner = self.statement()
            self.take(")")
            return inner
        if kind == "sign" and text == "\u221a":
            return Expr("call", name="sqrt", args=(self.atom(),))
        if kind != "name":
            raise NotReadable("'%s' stands where a number or a symbol was expected" % (text or "the end"))
        function = self.functions.get(text) or self.functions.get(normalise_symbol(text))
        skip = self.minus_one_follows() if function else 0
        if skip:
            inverse = self.notation.get("inverse", {}).get(function)
            if not inverse:
                raise NotReadable("the inverse of '%s' is not a function the tool knows" % text)
            self.position += skip
            function = inverse
        if self.peek() == ("sign", "("):
            if not function:
                raise NotReadable("'%s(' could be a product or a function; the tool does not guess" % text)
            self.take()
            arguments = [self.statement()]
            while self.peek() == ("sign", ","):
                self.take()
                arguments.append(self.statement())
            self.take(")")
            return Expr("call", name=function, args=tuple(arguments))
        return Expr("sym", name=normalise_symbol(text))

def parse_formula(text, notation, implicit_product=False):
    """Read a formula written in linear notation into the tool's expression tree. Raises NotReadable
    with a plain reason; it never guesses and never executes anything. Enforces: R7"""
    tokens = tokenize_formula(text)
    if not tokens:
        raise NotReadable("it is empty")
    reader = FormulaReader(tokens, notation, implicit_product)
    tree = reader.statement()
    if reader.peek()[0] != "end":
        raise NotReadable("'%s' stands where the formula should have ended" % reader.peek()[1])
    return tree

def read_equation(source_form, linear, notation, image_sha256=""):
    """Build the EquationData of a chunk: readable with its tree, or not readable with the reason."""
    linear = normalise_text(linear)
    if not linear:
        reason = UNDECIDED_REASONS[0] if image_sha256 or source_form == "image" else "no formula text was found"
        return EquationData(source_form, "", False, reason, image_sha256)
    try:
        tree = parse_formula(linear, notation, implicit_product=source_form in ("omml", "mathml", "latex"))
    except NotReadable as problem:
        return EquationData(source_form, linear, False, str(problem), image_sha256)
    return EquationData(source_form, expr_to_text(tree), True, "", image_sha256)

_INLINE_FORMULA_RE = re.compile(r"(?<![\w.])([^\W\d_][\w.\[\]{}]*)\s*=\s*([^=;]+)")

def inline_formula(text, notation):
    """A formula written inside running text, such as "K = LGD * N(x)". The right-hand side ends
    at a semicolon, at ", where", or at the end of the sentence. When the notation is ambiguous
    the paragraph simply stays a paragraph (plan 2.8). Returns EquationData or None."""
    for match in _INLINE_FORMULA_RE.finditer(text or ""):
        right = re.split(r",?\s+(?:where|with|and where|for)\b|\.\s+[A-Z]|\.$|:\s", match.group(2))[0]
        right = right.strip().rstrip(".,")
        if not re.search(r"[-+*/^(\u00d7\u00b7\u2212]", right):
            continue                                  # "x = 5" or "a = b": a value, not a formula
        try:
            tree = parse_formula("%s = %s" % (match.group(1), right), notation)
        except NotReadable:
            continue
        return EquationData("inline", expr_to_text(tree), True, "", "")
    return None

# ---------------------------------------------------------------- equation markup -> linear notation
def bracketed(text):
    """Put brackets around a part unless it is one number, one symbol or one call."""
    text = text.strip()
    if re.fullmatch(r"[\w.]+|[\w.]+\([^()]*\)", text) or (text.startswith("(") and text.endswith(")") and
                                                         text.count("(") == 1):
        return text
    return "(%s)" % text

NARY_NAMES = {"\u2211": "sum_over", "\u220f": "product_over", "\u222b": "integral_over"}
MATH_PROPERTIES = ("rpr", "ctrlpr", "fpr", "dpr", "narypr", "radpr", "ssuppr", "ssubpr", "ssubsuppr", "funcpr",
                   "annotation", "omathparapr", "begchr", "endchr")

def math_to_linear(element):
    """Office Math (OMML) and MathML to linear notation, by the local names of the elements:
    fractions, powers, subscripts, roots, brackets and function application. A sum is marked
    as sum_over(...) and never expanded. Unknown elements contribute their text."""
    name = local_name(element.tag)
    def part(child_name):
        child = child_named(element, child_name)
        return math_children(child) if child is not None else ""
    if name in ("t", "mi", "mn", "mtext"):
        return (element.text or "").strip()
    if name == "mo":
        return " %s " % (element.text or "").strip()
    if name == "f":                                               # OMML fraction
        return "%s/%s" % (bracketed(part("num")), bracketed(part("den")))
    if name == "ssup":
        return "%s^%s" % (bracketed(part("e")), bracketed(part("sup")))
    if name == "ssub":
        return "%s_{%s}" % (part("e").strip(), part("sub").strip())
    if name == "ssubsup":
        return "%s_{%s}^%s" % (part("e").strip(), part("sub").strip(), bracketed(part("sup")))
    if name == "rad":
        degree = part("deg").strip()
        return "sqrt(%s)" % part("e") if degree in ("", "2") else "%s^(1/%s)" % (bracketed(part("e")), bracketed(degree))
    if name == "d":                                               # OMML brackets
        return "(%s)" % ", ".join(math_children(child) for child in element if local_name(child.tag) == "e")
    if name == "func":
        return "%s(%s)" % (part("fname").strip(), strip_outer_brackets(part("e")))
    if name == "nary":
        properties = child_named(element, "narypr")
        sign = child_named(properties, "chr") if properties is not None else None
        kind = NARY_NAMES.get(attribute(sign, "val"), "sum_over") if sign is not None else "integral_over"
        return "%s(%s, %s, %s)" % (kind, part("sub").strip() or "0", part("sup").strip() or "0", part("e"))
    children = list(element)
    if name == "mfrac" and len(children) == 2:
        return "%s/%s" % (bracketed(math_to_linear(children[0])), bracketed(math_to_linear(children[1])))
    if name == "msup" and len(children) == 2:
        return "%s^%s" % (bracketed(math_to_linear(children[0])), bracketed(math_to_linear(children[1])))
    if name == "msub" and len(children) == 2:
        return "%s_{%s}" % (math_to_linear(children[0]).strip(), math_to_linear(children[1]).strip())
    if name == "msubsup" and len(children) == 3:
        return "%s_{%s}^%s" % (math_to_linear(children[0]).strip(), math_to_linear(children[1]).strip(),
                               bracketed(math_to_linear(children[2])))
    if name == "msqrt":
        return "sqrt(%s)" % math_children(element)
    if name == "mroot" and len(children) == 2:
        return "%s^(1/%s)" % (bracketed(math_to_linear(children[0])), bracketed(math_to_linear(children[1])))
    if name == "mfenced":
        return "(%s)" % ", ".join(math_to_linear(child) for child in children)
    if name == "munderover" and len(children) == 3 and (children[0].text or "").strip() in NARY_NAMES:
        return "%s(%s, %s, " % (NARY_NAMES[children[0].text.strip()], math_to_linear(children[1]).strip(),
                                math_to_linear(children[2]).strip())
    if name in MATH_PROPERTIES:
        return ""                                                 # properties, not content
    return math_children(element)

def math_children(element):
    """The linear text of all children, in order. An open "sum_over(a, b, " from MathML takes
    the part that follows it as its body and is then closed."""
    pieces = [(element.text or "").strip() if not len(element) else ""]
    for child in element:
        piece = math_to_linear(child)
        if pieces[-1].endswith(", ") and pieces[-1].startswith(tuple(NARY_NAMES.values())):
            pieces[-1] = pieces[-1] + piece.strip() + ")"
        else:
            pieces.append(piece)
    return re.sub(r"\s+", " ", "".join(pieces)).strip()

def strip_outer_brackets(text):
    """Remove one pair of brackets that encloses the whole text."""
    text = text.strip()
    return text[1:-1] if text.startswith("(") and text.endswith(")") and text.count("(") == 1 else text

LATEX_WORDS = {"cdot": "*", "times": "*", "div": "/", "le": "<=", "leq": "<=", "ge": ">=", "geq": ">=",
               "ln": "ln", "log": "log", "exp": "exp", "max": "max", "min": "min", "left": "", "right": "",
               "quad": " ", "qquad": " ", ",": " ", ";": " ", "!": "", " ": " ", "displaystyle": ""}
LATEX_WRAPPERS = ("text", "mathrm", "mathit", "mathbf", "operatorname", "mbox", "textit", "textbf", "code")

def latex_group(source, position):
    """The content of the {...} group that starts at `position`, and the position after it."""
    if position >= len(source):
        raise NotReadable("a LaTeX command lacks its argument")
    if source[position] != "{":
        return source[position], position + 1
    depth, start = 0, position
    while position < len(source):
        depth += {"{": 1, "}": -1}.get(source[position], 0)
        position += 1
        if depth == 0:
            return source[start + 1:position - 1], position
    raise NotReadable("a LaTeX group is never closed")

def latex_to_linear(source):
    """The LaTeX subset found in roxygen \\eqn{} and \\deqn{} and in some XML, to linear notation:
    \\frac, \\sqrt, ^{}, _{}, Greek letters, \\cdot, \\times, \\left, \\right, text wrappers.
    An unknown command makes the formula unreadable; it is never skipped silently."""
    source, output, position = source.strip().strip("$"), [], 0
    while position < len(source):
        char = source[position]
        if char == "\\":
            match = re.match(r"\\([A-Za-z]+|.)", source[position:])
            word, position = match.group(1), position + match.end()
            if word == "frac":
                top, position = latex_group(source, position)
                bottom, position = latex_group(source, position)
                output.append("(%s)/(%s)" % (latex_to_linear(top), latex_to_linear(bottom)))
            elif word == "sqrt":
                degree = ""
                if source[position:position + 1] == "[":
                    end = source.index("]", position)
                    degree, position = source[position + 1:end], end + 1
                inner, position = latex_group(source, position)
                inner = latex_to_linear(inner)
                output.append("sqrt(%s)" % inner if degree in ("", "2") else "(%s)^(1/(%s))" % (inner, degree))
            elif word in LATEX_WRAPPERS:
                inner, position = latex_group(source, position)
                output.append(re.sub(r"\s+", "_", inner.strip()))
            elif word in GREEK:
                output.append(GREEK[word])
            elif word in LATEX_WORDS:
                output.append(LATEX_WORDS[word])
            else:
                raise NotReadable("it uses the LaTeX command '\\%s', which the tool does not read" % word)
        elif char == "^":
            inner, position = latex_group(source, position + 1)
            output.append("^(%s)" % latex_to_linear(inner))
        elif char == "_":
            inner, position = latex_group(source, position + 1)
            output.append("_{%s}" % latex_to_linear(inner).replace(" ", ""))
        elif char in "{}":
            output.append("(" if char == "{" else ")")
            position += 1
        else:
            output.append(char)
            position += 1
    return re.sub(r"\s+", " ", "".join(output)).strip()

# ---------------------------------------------------------------- format from content, and repairs
XML_ENTITIES = ("amp", "lt", "gt", "quot", "apos")

def repair_markup(text, file_name, repairs):
    """The repairs the tool makes before strict parsing. Each one is recorded with its position, its
    kind, and the text before and after, so that a reviewer can see exactly what was changed."""
    def record(kind, position, before, after):
        repairs.append({"file": file_name, "position": position, "kind": kind,
                        "before": before[:200], "after": after[:200]})
    def fix(pattern, kind, replacement, flags=0):
        nonlocal text
        def change(found):
            after = replacement(found) if callable(replacement) else replacement
            record(kind, found.start(), found.group(0), after)
            return after
        text = re.sub(pattern, change, text, flags=flags)
    fix(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "control character removed", "")
    fix(r"<!DOCTYPE[^>\[]*(\[.*?\])?\s*>", "document-type declaration removed", "", re.S | re.I)
    def entity(found):
        name = found.group(1)
        if name in XML_ENTITIES or name not in html.entities.name2codepoint:
            return found.group(0)
        return "&#%d;" % html.entities.name2codepoint[name]
    before_entities = text
    text = re.sub(r"&([A-Za-z][A-Za-z0-9]*);", entity, text)
    if text != before_entities:
        named = sorted(set(re.findall(r"&([A-Za-z][A-Za-z0-9]*);", before_entities)) - set(XML_ENTITIES))
        record("named characters replaced by their numbers", 0, ", ".join("&%s;" % n for n in named), "")
    fix(r'=""([^"<>=]*)""', "doubled quotes in an attribute", lambda found: '="%s"' % found.group(1))
    body = re.sub(r"<\?xml[^>]*\?>|<!--.*?-->", "", text, flags=re.S)
    roots, depth = 0, 0
    for tag in re.finditer(r"<(/?)([^\s<>/!?][^<>]*?)(/?)>", body):
        if tag.group(1):
            depth -= 1
        elif not tag.group(3):
            roots += 1 if depth == 0 else 0
            depth += 1
        else:
            roots += 1 if depth == 0 else 0
    if roots > 1:
        declaration = re.match(r"\s*<\?xml[^>]*\?>", text)
        start = declaration.end() if declaration else 0
        text = text[:start] + "<document-root>" + text[start:] + "</document-root>"
        record("several top elements wrapped in one", start, "%d top elements" % roots, "<document-root>")
    return text

VOID_TAGS = ("img", "br", "hr", "meta", "link", "input", "col", "area", "base", "wbr")
BLOCK_STARTS = ("p", "h1", "h2", "h3", "h4", "h5", "h6", "table", "ul", "ol", "div", "li", "tr")
IMPLIED_END = dict({tag: ("p",) for tag in BLOCK_STARTS}, li=("li", "p"), tr=("tr", "td", "th", "p"), td=("td", "th", "p"), th=("td", "th", "p"))

class TolerantReader(html.parser.HTMLParser):
    """Builds the same kind of element tree as the strict XML parser, from start, end and text
    events of the standard library's HTML parser. It forgives what real exports contain: tags
    never closed, tags closed in the wrong order, attributes without quotes. Equation markup
    that Word's web export hides inside conditional comments is read too. (This is the one
    place where the tool subclasses: the standard parser offers no other way to receive events.)"""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.builder, self.open_tags = ElementTree.TreeBuilder(), []
        self.builder.start("document-root", {})

    def handle_starttag(self, tag, attrs):
        """Open an element, first closing what HTML closes implicitly."""
        while self.open_tags and self.open_tags[-1] in IMPLIED_END.get(tag, ()):       # <p> ends an open <p>, as browsers read it
            self.builder.end(self.open_tags.pop())
        self.builder.start(tag, {name: value or "" for name, value in attrs})
        if tag in VOID_TAGS:
            self.builder.end(tag)
        else:
            self.open_tags.append(tag)

    def handle_startendtag(self, tag, attrs):
        """An element written as empty: opened and closed at once."""
        self.builder.start(tag, {name: value or "" for name, value in attrs})
        self.builder.end(tag)

    def handle_endtag(self, tag):
        """Close the nearest open element of this name; a stray end tag is recorded as a repair."""
        if tag not in self.open_tags:
            return                                       # an end tag that closes nothing is ignored
        while self.open_tags:
            closing = self.open_tags.pop()
            self.builder.end(closing)
            if closing == tag:
                return

    def handle_data(self, data):
        """Text between tags."""
        self.builder.data(data)

    def handle_comment(self, data):
        """Keep equation markup that Word's web export hides in conditional comments; drop other comments."""
        preserved = re.match(r"\[if[^\]]*msEquation[^\]]*\]>(.*)<!\[endif\]", data, re.S)
        if preserved and "omath" in preserved.group(1).lower():
            self.builder.start("verifier-preserved-equation", {})
            self.feed_inner(preserved.group(1))
            self.builder.end("verifier-preserved-equation")

    def feed_inner(self, markup):
        """Read preserved equation markup found inside a comment into the tree being built."""
        inner = TolerantReader()
        inner.feed(markup)
        for child in inner.finish():
            self.builder.start(child.tag, child.attrib)
            self.copy_children(child)
            self.builder.end(child.tag)

    def copy_children(self, element):
        """Copy an element read elsewhere into the tree being built."""
        if element.text:
            self.builder.data(element.text)
        for child in element:
            self.builder.start(child.tag, child.attrib)
            self.copy_children(child)
            self.builder.end(child.tag)
            if child.tail:
                self.builder.data(child.tail)

    def finish(self):
        """Close whatever is still open and return the root element."""
        self.close()
        while self.open_tags:
            self.builder.end(self.open_tags.pop())
        self.builder.end("document-root")
        return self.builder.close()

def parse_markup(text, file_name, repairs, tolerant_only=False):
    """Strict XML parsing first; when that still fails after the repairs, the tolerant reader.
    The path taken is recorded. Returns the root element."""
    repaired = repair_markup(text, file_name, repairs)
    if not tolerant_only:
        try:
            return ElementTree.fromstring(repaired.encode("utf-8"))
        except ElementTree.ParseError as problem:
            repairs.append({"file": file_name, "position": 0, "kind": "read with the tolerant reader",
                            "before": "strict reading stopped: %s" % problem, "after": ""})
    reader = TolerantReader()
    reader.feed(repaired)
    return reader.finish()

# ---------------------------------------------------------------- tag rules and the block walker
def load_tag_rules(override_path=None):
    """The shipped tag rules, with any part replaced by the project's own tag_rules.yaml."""
    rules = yaml.safe_load(TAG_RULES_YAML)
    rules["shipped_tags"] = sorted(str(tag).lower() for tags in rules["families"].values() for tag in tags)
    analyst_families = {}
    if override_path:
        with open(override_path, encoding="utf-8") as handle:
            override = yaml.safe_load(handle) or {}
        analyst_families = override.get("families") or {}
        for key, value in override.items():
            if key == "families":
                rules["families"].update(value or {})
            else:
                rules[key] = value
    rules["analyst_tags"] = sorted(str(tag).lower() for tags in analyst_families.values() for tag in tags)
    rules["family_of"] = {}
    for family, tags in rules["families"].items():
        for tag in tags:
            rules["family_of"][str(tag).lower()] = family
    rules.setdefault("heading_attributes", ["name", "title", "heading", "label", "caption"])
    rules.setdefault("caption_tag_words", ["title", "number", "caption", "legend"])
    rules.setdefault("header_tag_words", ["head", "header"])
    return rules

# ---------------------------------------------------------------- discovering an unfamiliar schema
BLOCKS_INSIDE = ("paragraph", "list_item", "list_container", "heading", "container", "row", "cell", "header_cell")

def element_text(element, rules, skip=("figure", "equation", "ignore", "caption")):
    """The running text of an element without the text of figures, equations and captions in it.

    A child that is a block of its own - a list item inside a table cell, say - is separated by
    a space rather than run straight onto what came before it. Without that, two items of a
    list in one cell arrive as one word that the document does not contain ("renewable twice" +
    "no fine" giving "twiceno"), which is text the tool made up. Enforces: R13"""
    pieces = [element.text or ""]
    for child in element:
        if rules["family_of"].get(local_name(child.tag)) not in skip:
            if rules["family_of"].get(local_name(child.tag)) in BLOCKS_INSIDE and "".join(pieces).strip():
                pieces.append(" ")
            pieces.append(element_text(child, rules, skip))
        pieces.append(child.tail or "")
    return "".join(pieces)

def new_block(kind, text="", locator="", **more):
    """One block of a document before numbering: kind, text, where it was found, and what its kind needs."""
    block = {"type": kind, "text": normalise_text(text), "locator": locator, "level_hint": None,
             "numbering": "", "caption": "", "table": None, "equation": None, "reconstructed": False,
             "not_read_reason": ""}
    block.update(more)
    return block

def not_read_block(file_name, reason):
    """A whole file, or a part, that could not be read still becomes one block: the "not read"
    class of the content account, never a silent gap. Enforces: R2, R13"""
    return new_block("paragraph", "", file_name, not_read_reason=reason)

@dataclass
class WalkState:
    """What the walker carries along: the rules, the notation, images by name, the report of
    tags it met that are in no family, and what discovery made of those tags in this document."""
    rules: dict; notation: dict; images: dict; unknown_tags: dict; blocks: list
    skip_next_image: bool = False
    notes: list = field(default_factory=list)      # what was left out or read in a fallback way, in plain words
    atoms: list = field(default_factory=list)      # the smallest pieces of text the file holds, counted from the file itself
    dropped: list = field(default_factory=list)    # text left out under a named rule, kept so the account can show it
    lent_numbering: str = ""                       # a number a container carries for the heading inside it
    settings: dict = field(default_factory=dict)
    folder: str = ""                               # where the file being read stands, so a picture beside it can be found
    svgs: dict = field(default_factory=dict)       # the corner's SVG files, by lower-case name and by name without .svg
    consumed: set = field(default_factory=set)     # SVGs already read in place by a document of the corner (shared by the corner)
    file_name: str = ""
    lists: list = field(default_factory=list)      # the lists the walker is inside of: [numbered?, items so far]

    def __post_init__(self):
        """Each file reads with its own view of the rules, so a tag discovered in one file
        never changes how the next file is read."""
        self.rules = dict(self.rules)
        self.rules["family_of"] = dict(self.rules["family_of"])

    def family(self, tag):
        """The family of a tag: what the rules say, else what discovery made of it here."""
        return self.rules["family_of"].get(tag)

def walk_element(element, path, depth, state):
    """Turn one element and everything below it into blocks, in reading order."""
    name = local_name(element.tag)
    family = state.family(name)
    here = "%s/%s" % (path, name)
    if name == "verifier-preserved-equation":
        state.blocks.append(equation_block(element, here, state))
        state.skip_next_image = True                     # the picture that follows shows the same equation
        return
    if family == "ignore" or not name:
        return
    linked, named_by = (svg_reference(element, state), element) if state.svgs else (None, None)
    if not linked and state.svgs and state.family(name) == "figure":
        # a figure may name its picture one level down: <chart><file>Chart2.svg</file></chart>
        linked, named_by = next(((found, below) for below in element.iter() if below is not element
                                 for found in [svg_reference(below, state)] if found), (None, None))
    if linked:                                           # a chart or table drawn as an SVG beside the document: read it here, in place
        if named_by is not None and len(named_by) == 0 and (named_by.text or "").strip():
            state.dropped.append(named_by.text)          # the file name written as text is a reference, not what the document says
        caption = next((normalise_text(element_text(n, state.rules, skip=())) for n in element.iter()
                        if state.rules["family_of"].get(local_name(n.tag)) == "caption"), "") or element.get("alt") or element.get("title") or ""
        with open(linked, "rb") as handle:
            state.blocks.extend(svg_blocks(handle.read(), os.path.basename(linked), here, state, caption))
        state.consumed.add(linked)
        state.notes.append("%s: read in place, where %s refers to it." % (os.path.basename(linked), state.file_name or "the document"))
        return
    if family is None:                                   # discovery names every tag it reaches; this is the net under it
        family = "container" if len(element) else "paragraph"
    numbering = written_numbering(element, state.rules)
    if family == "heading":
        text = element_text(element, state.rules)
        digit = re.fullmatch(r"h([1-6])", name)
        state.blocks.append(new_block("heading", text, here, numbering=numbering or state.lent_numbering,
                                      level_hint=int(digit.group(1)) if digit else depth))
        state.lent_numbering = ""
    elif family in ("container", "list_container", "inline"):
        # A container that carries its own heading in an attribute (<section name="4. Market">)
        # gives that heading a block of its own, so the chain below it is not lost.
        heading, _ = attribute_text(element, state.rules["heading_attributes"]) if family == "container" else ("", "")
        if heading:
            state.blocks.append(new_block("heading", heading, here, numbering=numbering, level_hint=depth))
        if family == "list_container":                   # <ol>, or <list type="numbered">: its items are counted
            kind = " ".join([name] + [element.get(a) or "" for a in ("type", "style", "numeration", "class")]).lower()
            state.lists.append([bool(re.search(r"\bol\b|order|num|decimal|arabic|alpha|roman", kind)), 0])
        # A container often carries the number of the section (<section num="2.">) while the
        # heading it labels is a child of it (<title>). The number is lent to the first heading
        # the container produces, so that "2." belongs to "2. Volume per bed" and not to
        # nothing at all. Enforces: R13
        lent, state.lent_numbering = state.lent_numbering, numbering if family == "container" and numbering and not heading else ""
        walk_mixed(element, here, depth + (1 if family == "container" else 0), state)
        state.lent_numbering = lent
        if family == "list_container":
            state.lists.pop()
    elif family in ("paragraph", "list_item"):
        marker = ""
        if family == "list_item":
            inside = state.lists[-1] if state.lists else [False, 0]
            inside[1] += 1
            marker = "  " * max(0, len(state.lists) - 1) + ("%d. " % inside[1] if inside[0] else LIST_MARKER)
        walk_mixed(element, here, depth, state, own_kind="list_item" if family == "list_item" else "paragraph",
                   numbering=numbering, marker=marker)
    elif family == "table":
        state.blocks.append(table_block(element, here, state))
    elif family == "figure":
        if state.skip_next_image:
            state.skip_next_image = False
            return
        found = figure_block(element, here, state)
        state.blocks.extend(found if isinstance(found, list) else [found])   # an SVG beside the document may give a table
    elif family == "equation":
        state.blocks.append(equation_block(element, here, state))
    elif family == "caption" and state.blocks and state.blocks[-1]["type"] in ("figure", "table", "equation"):
        state.blocks[-1]["caption"] = normalise_text(element_text(element, state.rules, skip=()))

def walk_mixed(element, here, depth, state, own_kind=None, numbering="", marker=""):
    """An element that may hold both running text and blocks. Its own text becomes one
    paragraph; a formula that fills the paragraph alone becomes an Equation block instead."""
    block_families = ("heading", "container", "list_container", "paragraph", "list_item", "table",
                      "figure", "equation", "caption", None)
    children = [c for c in element if local_name(c.tag) and
                (state.family(local_name(c.tag)) in block_families or local_name(c.tag) == "verifier-preserved-equation")]
    inline_only = [c for c in children if state.family(local_name(c.tag)) is None and not len(c)
                   and own_kind]
    children = [c for c in children if c not in inline_only]
    text = normalise_text(element_text(element, state.rules, skip=("figure", "equation", "ignore", "caption",
                                 "table", "list_container")) if own_kind or not children else (element.text or ""))
    equations = [c for c in children if state.family(local_name(c.tag)) == "equation"]
    if own_kind and text and equations:                  # text around a formula: show the formula in place
        text = normalise_text(text + " " + " ".join(math_to_linear(c) for c in equations))
        children = [c for c in children if c not in equations]
    if text and (own_kind or not children):
        state.blocks.append(new_block(own_kind or "paragraph", text, here, numbering=numbering, marker=marker))
    for position, child in enumerate(children, start=1):
        if own_kind and state.family(local_name(child.tag)) in ("paragraph", "inline"):
            continue                                     # already part of the paragraph's own text
        walk_element(child, "%s[%d]" % (here, position), depth, state)
        if not own_kind and child.tail and child.tail.strip():
            state.blocks.append(new_block("paragraph", child.tail, here))

def table_block(element, here, state):
    """A table is always one block: header cells, body rows and its caption stay together.
    Whatever sits in a row is a cell unless it is the caption or is ignored, so a kind of cell
    that first turns up in the fortieth row still keeps its words. Tags holding the table's
    number and title beside empty cells are the caption, kept in the order written; the rows
    they leave empty are dropped, so the first row with content is the header. A table whose
    rows cannot be found keeps its words as running text and says so: a table that is present
    and empty misleads more than none."""
    family = lambda node: state.family(local_name(node.tag))
    found = [node for node in element.iter() if family(node) == "row"] or table_rows(element, state.rules)
    rows, captions = [], []
    in_a_row = {id(node) for row in found for node in row.iter()}
    for node in element.iter():
        # A caption may be named as one, or may simply sit beside the rows rather than in them
        # (<gridcaption> under <gridholder>). Either way its words belong to the table and are
        # kept: a table whose title is lost cannot be cited by its number. Enforces: R13
        beside = (node is not element and id(node) not in in_a_row and family(node) not in ("row", "table_part", "ignore")
                  and not any(child is not None and family(child) == "row" for child in node))
        part = normalise_text(element_text(node, state.rules, skip=())) if family(node) == "caption" or beside else ""
        if part and part not in captions:
            captions.append(part)
    for row in found:
        cells = [normalise_text(element_text(cell, state.rules, skip=())) for cell in row
                 if local_name(cell.tag) and family(cell) not in ("caption", "ignore")]
        if any(cells):
            rows.append(cells)
    block = table_from_rows(rows, here, ". ".join(captions))
    words = "" if rows else normalise_text(element_text(element, state.rules, skip=("ignore", "caption")))
    if words:
        block["display"] = words
        state.notes.append("The rows of the table at %s could not be told apart, so its words are shown as running text." % here)
    return block

def table_from_rows(rows, locator, caption=""):
    """The one-cell display form of a table. Short values are shown as a grid: cells joined by
    "; ", one row per line, header first. Sentences would be unreadable that way, so each cell
    goes on its own line under the heading of its column ("Very Strong: Airport that ..."); a row
    with one filled cell (a sub-heading, a note) is shown as it stands. The grid is kept either way."""
    width = max((len(row) for row in rows), default=0)
    rows = [list(row) + [""] * (width - len(row)) for row in rows]
    header, body = (rows[0], rows[1:]) if rows else ([], [])
    filled = [cell for row in body if sum(1 for cell in row if cell) > 1 for cell in row if cell]   # sub-headings and notes do not vote
    if filled and sum(len(cell) for cell in filled) > 60 * len(filled):
        lines = ["; ".join(cell for cell in header if cell)]
        for row in body:
            cells = [(header[i], cell) for i, cell in enumerate(row) if cell]
            lines.extend(("%s: %s" % pair if pair[0] and len(cells) > 1 else pair[1]) for pair in cells)
        display = "\n".join(lines)
    else:
        display = "\n".join("; ".join(row[:max((i for i, cell in enumerate(row) if cell), default=0) + 1]) for row in rows)
    first = [row[0] for row in body if row[0]] if width else []     # a first column of words keys the rows
    numeric = first and all(find_numbers(v) and len(find_numbers(v)) == 1 and
                            len(re.sub(r"[\d.,%\s+-]|bps?|basis points?", "", v)) == 0 for v in first)
    row_key = header[0] if header and width and not numeric else ""
    table = TableData(tuple(header), tuple(tuple(row) for row in body), row_key)
    return new_block("table", "", locator, table=table, caption=caption, display=display)

PICTURE_READER = []                     # the OCR engine, looked for once: [engine] or [None]
OCR_NOTE = "Words read from the picture by OCR (a machine reading: check it against the picture itself):"

def read_picture(data, state):
    """The words in a picture, read by OCR, as lines to show under the Figure; "" when there are
    none or no OCR package is installed (rapidocr-onnxruntime is optional; its models come inside
    the package, so nothing is fetched when it runs). The words help a person find the
    picture. They are never evidence: a machine misreads digits, so a Figure still ends
    "for manual review" whatever was read. A missing reader is said once per file. Enforces: R2"""
    if not PICTURE_READER:
        try:
            from rapidocr_onnxruntime import RapidOCR
            PICTURE_READER.append(RapidOCR())
        except Exception:                                # not installed, or its models cannot be loaded
            PICTURE_READER.append(None)
    missing = "The words inside pictures were not read: the optional OCR package rapidocr-onnxruntime is not installed."
    if PICTURE_READER[0] is None and missing not in state.notes:
        state.notes.append(missing)
    if PICTURE_READER[0] is None or len(data) < 2048 or not state.rules.get("read_pictures", True):
        return ""                                        # under 2 kB is an icon or a rule, not a picture with words
    try:
        found = sorted(PICTURE_READER[0](data)[0] or [], key=lambda item: (round(item[0][0][1] / 14), item[0][0][0]))
    except Exception:                                    # a form the reader cannot open (.emf, .wmf)
        unread = "The words inside one or more pictures were not read: the picture is in a form the reader cannot open."
        if unread not in state.notes:
            state.notes.append(unread)                   # the picture is still a unit; only its words are missing (R2)
        return ""
    rows = {}
    for box, words, _ in found:                          # what stands on one line of the picture stays on one line
        rows.setdefault(round(box[0][1] / 14), []).append(re.sub(r"(?<=[a-z])(?=[A-Z])", " ", words))
    return "\n%s\n%s" % (OCR_NOTE, "\n".join("  ".join(row) for row in rows.values())) if rows else ""

def picture_words(page, box, state):
    """PDF: the part of the page that a picture covers, drawn at 150 dpi and read by OCR."""
    if box[2] - box[0] < 40 or box[3] - box[1] < 40:
        return ""
    try:
        drawn = io.BytesIO()
        area = (max(box[0], 0), max(box[1], 0), min(box[2], page.width), min(box[3], page.height))
        page.crop(area).to_image(resolution=150).original.save(drawn, "PNG")
    except Exception:
        return ""
    return read_picture(drawn.getvalue(), state)

def figure_block(element, here, state):
    """A figure: never read, kept with its caption or alternative text and the fingerprint of the image."""
    source = element.get("src") or element.get("href") or element.get("fileref") or ""
    inner = next((n for n in element.iter() if n is not element and (n.get("src") or n.get("fileref"))), None)
    if not source and inner is not None:
        source = inner.get("src") or inner.get("fileref") or ""
    label = element.get("alt") or element.get("title") or (inner.get("alt") if inner is not None else "") or ""
    caption = next((normalise_text(element_text(n, state.rules, skip=())) for n in element.iter()
                    if state.rules["family_of"].get(local_name(n.tag)) == "caption"), "")
    fingerprint = state.images.get(source) or state.images.get(source.replace("cid:", "")) or ""
    return new_block("figure", label or caption or source, here, caption=caption, image_sha256=fingerprint,
                     source=source)

LINK_ATTRIBUTE = re.compile(r"src|href|ref|file|data|path|url|image|graphic", re.I)

def svg_reference(element, state):
    """The SVG of this corner an element refers to, or None. A reference is either an attribute
    whose name says it links (src, href, xlink:href, fileref, data, a tool's own graphic= or
    image=), or an element's whole text when it names a file ending in .svg. A plain name is
    matched to a file of the corner in any case, and without .svg only when it came from a link
    attribute; a path is followed only if it stays inside the document's folder. Found on review:
    matching any short text without .svg took the heading <title>Floors</title> for floors.svg and
    replaced the heading with the chart. Enforces: R6"""
    values = [value for key, value in element.attrib.items() if LINK_ATTRIBUTE.search(local_name(key))]
    text = (element.text or "").strip() if len(element) == 0 else ""
    if text.lower().endswith(".svg") and len(text) < 200:
        values.append(text)
    for value in values:
        written = (value or "").strip().replace("\\", "/").split("?")[0].split("#")[0]
        if not written or ".." in written.split("/") or "://" in written or written.startswith("/"):
            continue
        if "/" in written and state.folder:
            resolved = os.path.normpath(os.path.join(state.folder, written)).lower()
            found = next((path for path in state.svgs.values() if os.path.normpath(path).lower() == resolved), None)
            if found:
                return found
        name = os.path.basename(written).lower()
        for key in ((name,) if name.endswith(".svg") else (name, name + ".svg")):
            if key in state.svgs:
                return state.svgs[key]
    return None

def svg_blocks(data, name, here, state, caption=""):
    """The units of one SVG, read by the text it holds: a table where its text stands in a grid, a
    figure of its labels otherwise. Where it holds no text, a picture stored inside it is read by
    OCR when that is installed; and a figure that still gives no words says why, instead of standing
    empty. The words counted are the SVG's own, so the content account closes. Enforces: R2, R13"""
    import xml.etree.ElementTree as ElementTree
    try:
        root = ElementTree.fromstring(data)
    except ElementTree.ParseError:
        return [not_read_block(name, "the SVG file is damaged and could not be opened")]
    markup, words, _ = svg_to_markup(data, name)
    own_caption, labels = (words.split("\n", 1) + [""])[:2]
    state.atoms.extend(atoms_of_plain_text(labels if caption else words, name))
    blocks = walk_svg_markup(parse_markup(markup, name, [], False), here, state)
    for block in blocks:
        block["caption"] = caption or own_caption
        block["source"], block["image_sha256"] = name, sha256_bytes(data)
        if block["type"] == "figure" and not block["text"].strip():
            seen = "".join(read_picture(picture, state) for picture in svg_embedded_pictures(root)).strip()
            if seen:
                block["text"] = seen                     # words read from a picture: shown, never evidence
            else:
                block["not_read_reason"] = "the picture gave no words: %s" % svg_why_no_text(root)
    return blocks

def walk_svg_markup(root, here, state):
    """The blocks the SVG's markup gives: a table block, or one figure block carrying the
    picture's labels as its text."""
    blocks = []
    for node in root:
        family = state.rules["family_of"].get(local_name(node.tag))
        if family == "table":
            block = table_block(node, here, state)
            if block is not None:
                blocks.append(block)
        elif family == "figure":
            blocks.append(new_block("figure", node.get("alt") or "", here,
                                    caption=next((normalise_text("".join(c.itertext())) for c in node if local_name(c.tag) == "caption"), "")))
    return blocks


def equation_block(element, here, state):
    """An equation element: MathML or Office Math is converted; LaTeX or linear text is read as
    written; an equation that is only a picture stays an Equation chunk that could not be read."""
    markup = next((n for n in element.iter() if local_name(n.tag) in ("math", "omath")), None)
    picture = next((n for n in element.iter() if local_name(n.tag) in ("img", "image", "graphic", "imagedata")), None)
    if markup is not None:
        form = "mathml" if local_name(markup.tag) == "math" else "omml"
        equation = read_equation(form, math_to_linear(markup), state.notation)
    elif picture is not None and not normalise_text(element_text(element, state.rules)):
        source = picture.get("src") or picture.get("fileref") or ""
        equation = read_equation("image", "", state.notation, state.images.get(source, "") or sha256_text(source))
    else:
        written = normalise_text(element_text(element, state.rules, skip=("caption", "ignore")))
        try:
            linear = latex_to_linear(written) if "\\" in written else written
            equation = read_equation("latex" if "\\" in written else "inline", linear, state.notation)
        except NotReadable as problem:
            equation = EquationData("latex", written, False, str(problem), "")
    label = picture.get("alt") if picture is not None and picture.get("alt") else ""
    text = equation.linear or label or "Equation shown as a picture"
    return new_block("equation", text, here, equation=equation)

def blocks_from_markup(text, file_name, state, repairs, tolerant_only=False):
    """XML or HTML text to blocks: parse (repairing where needed), then walk the tree by the tag rules."""
    root = parse_markup(text, file_name, repairs, tolerant_only)
    discovered = discover_families(root, state.rules, state.unknown_tags)
    state.rules["family_of"].update(discovered)
    walk_element(root, "", 0, state)
    return state.blocks

# ---------------------------------------------------------------- MHTML (a web page saved as one file)
def blocks_from_mhtml(data, file_name, state, repairs):
    """Parts are read with the standard `email` package. The HTML part goes through the tolerant
    reader; every image part is fingerprinted so that a Figure chunk can name its content."""
    message = email.message_from_bytes(data)
    page = None
    for part in message.walk():
        content_type = part.get_content_type()
        payload = part.get_payload(decode=True)
        if payload is None:
            continue
        if content_type == "text/html" and page is None:
            page = payload.decode(part.get_content_charset() or "utf-8", "replace")
        elif content_type.startswith("image/"):
            fingerprint = sha256_bytes(payload)
            for key in (part.get("Content-Location", ""), (part.get("Content-ID", "") or "").strip("<>")):
                if key:
                    state.images[key] = fingerprint
                    state.images[key.rsplit("/", 1)[-1]] = fingerprint
    if page is None:
        return [not_read_block(file_name, "the file holds no web page part")]
    blocks = blocks_from_markup(page, file_name, state, repairs, tolerant_only=True)
    return [block for block in title_block(page, file_name) + blocks]

def title_block(page, file_name):
    """The <title> of a web page, as a heading above everything in it. In a Word file exported
    to the web it is often the only place the document's own name survives, because the visible
    heading may be a styled paragraph carrying no heading level. Without this the title reaches
    no unit and can be cited by nothing. Enforces: R13"""
    found = re.search(r"<title[^>]*>(.*?)</title>", page, re.IGNORECASE | re.DOTALL)
    text = normalise_text(re.sub(r"<[^>]+>", " ", found.group(1))) if found else ""
    return [new_block("heading", text, "%s title" % file_name, level_hint=1)] if text else []

# ---------------------------------------------------------------- .docx
WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
           "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
           "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
           "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
           "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}

def word_value(element, path, attribute_name="val"):
    """The value of a Word property such as a style id or an outline level, or None."""
    found = element.find(path, WORD_NS)
    return found.get("{%s}%s" % (WORD_NS["w"], attribute_name)) if found is not None else None

def docx_paragraph_facts(paragraph, styles):
    """Heading level (from the style name or the outline level, following based-on styles) and
    whether Word numbers this paragraph automatically."""
    style_id = word_value(paragraph, "w:pPr/w:pStyle") or ""
    level = word_value(paragraph, "w:pPr/w:outlineLvl")
    numbering = paragraph.find("w:pPr/w:numPr", WORD_NS)
    name, hops = "", 0
    while style_id in styles and hops < 5:
        style = styles[style_id]
        name = name or (word_value(style, "w:name") or "")
        level = level if level is not None else word_value(style, "w:pPr/w:outlineLvl")
        numbering = numbering if numbering is not None else style.find("w:pPr/w:numPr", WORD_NS)
        style_id, hops = word_value(style, "w:basedOn") or "", hops + 1
    numbered = (word_value(numbering, "w:numId") or "", word_value(numbering, "w:ilvl") or "0") if numbering is not None else None
    if numbered and numbered[0] == "0":                  # Word writes numId 0 to switch numbering off
        numbered = None
    heading = re.fullmatch(r"[Hh]eading\s*(\d)", name)
    if heading:
        return int(heading.group(1)), numbered, name
    if level is not None and level.isdigit() and int(level) < 9:
        return int(level) + 1, numbered, name
    return None, numbered, name

def docx_number_formats(archive):
    """How each numbering of a Word file shows its items, by numbering id and level: "bullet",
    "decimal", "lowerLetter" ... Word keeps this apart from the text, in word/numbering.xml."""
    if "word/numbering.xml" not in archive.namelist():
        return {}
    root, key = safe_xml(archive.read("word/numbering.xml")), "{%s}" % WORD_NS["w"]
    shapes = {a.get(key + "abstractNumId"): {lvl.get(key + "ilvl"): word_value(lvl, "w:numFmt") for lvl in a.findall("w:lvl", WORD_NS)}
              for a in root.findall("w:abstractNum", WORD_NS)}
    return {n.get(key + "numId"): shapes.get(word_value(n, "w:abstractNumId"), {}) for n in root.findall("w:num", WORD_NS)}

def docx_paragraph_parts(paragraph):
    """The text of a paragraph with its formulas in place, its formulas, and its pictures."""
    pieces, formulas, pictures = [], [], []
    for node in paragraph.iter():
        name = local_name(node.tag)
        if name == "t" and node.tag.startswith("{%s}" % WORD_NS["w"]):
            pieces.append(node.text or "")
        elif name in ("tab", "br") and node.tag.startswith("{%s}" % WORD_NS["w"]):
            pieces.append(" ")
        elif name == "omath":
            formulas.append(node)
            pieces.append(" %s " % math_to_linear(node))
        elif name in ("drawing", "pict"):
            pictures.append(node)
    return "".join(pieces), formulas, pictures

def docx_figure(picture, related, locator, state):
    """A picture in a Word file as a figure block with the fingerprint of the embedded image,
    and the words in it where OCR is installed."""
    description = next((n.get("descr") or n.get("title") or n.get("name") or "" for n in picture.iter()
                        if local_name(n.tag) == "docpr"), "")
    fingerprint, words = "", ""
    for node in picture.iter():
        for key, value in node.attrib.items():
            if key.startswith("{%s}" % WORD_NS["r"]) and value in related:
                fingerprint, words = sha256_bytes(related[value]), read_picture(related[value], state)
    fingerprint = fingerprint or sha256_bytes(ElementTree.tostring(picture))
    said = (description or "Picture without a description") + words
    return new_block("figure", said, locator, image_sha256=fingerprint, display=said)

def safe_xml(data):
    """Parse one XML part of an Office file. Any document-type declaration is removed first,
    so no entity can be defined inside the file (no entity expansion, no outside fetch)."""
    text = re.sub(rb"<!DOCTYPE[^>\[]*(\[.*?\])?\s*>", b"", data, flags=re.S | re.I)
    return ElementTree.fromstring(text)

def note_blocks(archive, file_name):
    """The footnotes and endnotes of a Word file, which are parts of their own and are not in
    the run of paragraphs a plain reader walks. Real documentation puts definitions, caveats and
    parameter values in a footnote routinely, and a reader that takes the body and not the notes
    reports a statement as undocumented when its documentation is two lines below the text.
    They are read after the body, each saying which note it is. Enforces: R13"""
    found = []
    for part, kind in (("word/footnotes.xml", "footnote"), ("word/endnotes.xml", "endnote")):
        if part not in archive.namelist():
            continue
        try:
            root = safe_xml(archive.read(part))
        except ElementTree.ParseError:
            continue
        for note in root:
            if (note.get("{%s}type" % WORD_NS["w"]) or "") in ("separator", "continuationSeparator", "continuationNotice"):
                continue                                 # the rule Word draws above a footnote, not a note
            number = note.get("{%s}id" % WORD_NS["w"]) or "?"
            for position, paragraph in enumerate(note.findall("w:p", WORD_NS), start=1):
                text = normalise_text("".join(run.text or "" for run in paragraph.iter("{%s}t" % WORD_NS["w"])))
                if text:
                    found.append(new_block("paragraph", text, "%s %s %s paragraph %d" % (file_name, kind, number, position)))
    return found

def blocks_from_docx(data, file_name, state):
    """Body elements in document order, so that tables stay where they are. Heading numbers that
    Word produces automatically are not stored in the file; they are reconstructed by counting
    and marked as reconstructed."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        body = safe_xml(archive.read("word/document.xml")).find("w:body", WORD_NS)
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError):
        return [not_read_block(file_name, "the file is not a Word document that can be opened")]
    styles, related = {}, {}
    if "word/styles.xml" in archive.namelist():
        for style in safe_xml(archive.read("word/styles.xml")).findall("w:style", WORD_NS):
            styles[style.get("{%s}styleId" % WORD_NS["w"])] = style
    if "word/_rels/document.xml.rels" in archive.namelist():
        for relation in safe_xml(archive.read("word/_rels/document.xml.rels")):
            target = "word/" + relation.get("Target", "").lstrip("/")
            if relation.get("TargetMode") != "External" and target in archive.namelist():
                related[relation.get("Id")] = archive.read(target)
    blocks, counters, pending_caption = [], [0] * 9, ""
    notes = note_blocks(archive, file_name)
    formats, counts, page, on_page = docx_number_formats(archive), {}, 1, {}
    paged = b"lastRenderedPageBreak" in archive.read("word/document.xml")     # Word noted where its pages ended
    for position, element in enumerate(body, start=1):
        locator, name = "body element %d" % position, local_name(element.tag)
        if name == "tbl":
            rows = []
            for row in element.findall("w:tr", WORD_NS):
                rows.append([normalise_text(docx_paragraph_parts(cell)[0]) for cell in row.findall("w:tc", WORD_NS)])
            blocks.append(table_from_rows(rows, locator, pending_caption))
            pending_caption = ""
            continue
        if name != "p":
            continue
        level, numbered, style_name = docx_paragraph_facts(element, styles)
        text, formulas, pictures = docx_paragraph_parts(element)
        text = normalise_text(text)
        if style_name.lower() == "caption":
            if blocks and blocks[-1]["type"] in ("figure", "table", "equation") and not blocks[-1]["caption"]:
                blocks[-1]["caption"] = text
            else:
                pending_caption = text
            continue
        for picture in pictures:
            blocks.append(docx_figure(picture, related, locator, state))
        if level is not None and text:
            block = new_block("heading", text, locator, level_hint=level)
            if numbered and not first_numbering(text, state.rules)[0]:
                counters[level - 1] += 1
                counters[level:] = [0] * (9 - level)
                block["numbering"] = ".".join(str(c) for c in counters[:level] if c)
                block["reconstructed"] = True
            blocks.append(block)
        elif formulas and not normalise_text(re.sub(r"\s+", " ", text.replace(math_to_linear(formulas[0]), ""))):
            equation = read_equation("omml", math_to_linear(formulas[0]), state.notation)
            blocks.append(new_block("equation", equation.linear or text, locator, equation=equation))
        elif text:
            page += sum(1 for node in element.iter() if local_name(node.tag) == "lastrenderedpagebreak")
            shape = formats.get(numbered[0], {}).get(numbered[1], "") if numbered else ""
            if numbered and shape != "bullet" and "bullet" not in style_name.lower():
                counts[numbered] = counts.get(numbered, 0) + 1            # Word counts the items; the file does not hold the numbers
                counts.update({key: 0 for key in counts if key[0] == numbered[0] and key[1] > numbered[1]})
                written = chr(96 + counts[numbered]) if shape == "lowerLetter" and counts[numbered] < 27 else str(counts[numbered])
            if numbered and (shape == "bullet" or "list" in style_name.lower() or "bullet" in style_name.lower()):
                marker = "  " * int(numbered[1]) + (LIST_MARKER if shape == "bullet" or "bullet" in style_name.lower() else written + ". ")
                blocks.append(new_block("list_item", text, locator, marker=marker))
                continue
            on_page[page] = on_page.get(page, 0) + 1
            label = written + "." if numbered else "p.%d \u00b6%d" % (page, on_page[page]) if paged else ""
            blocks.append(new_block("paragraph", text, locator, numbering=label, reconstructed=bool(numbered)))
    if notes:
        blocks.append(new_block("heading", NOTES_HEADING, "%s notes" % file_name, level_hint=1))
        blocks.extend(notes)
    return blocks

# ---------------------------------------------------------------- .pdf
MATH_CHARACTERS = set("=+\u2212\u00d7\u00f7\u2211\u221a\u222b^/\u2264\u2265\u2248()") | set(GREEK.values())

BULLET = re.compile(r"^(?:[\u2022\u25aa\u25cf\u25e6\u2023\u2043\u2013\u2014*o-]|\(cid:\d+\))\s+(?=\S)")

def pdf_lines(document, file_name):
    """Every line, table and picture of a PDF in reading order, each with its page, its place on
    the page, and whether it sits in the top or bottom margin ("edge")."""
    lines = []
    for number, page in enumerate(document.pages, start=1):
        edge = lambda top, bottom: "top" if bottom < 0.12 * page.height else "bottom" if top > 0.88 * page.height else ""
        tables = page.find_tables()
        for table in tables:
            rows = [[normalise_text(cell or "") for cell in row] for row in table.extract()]
            lines.append({"page": number, "top": table.bbox[1], "table": rows, "edge": ""})
        for count, image in enumerate(page.images, start=1):
            mark = "%s page %d picture %d %s" % (file_name, number, count, image.get("srcsize"))
            lines.append({"page": number, "top": image["top"], "figure": sha256_text(mark), "text": "picture %s" % (image.get("srcsize"),),
                          "edge": edge(image["top"], image["bottom"]), "box": (image["x0"], image["top"], image["x1"], image["bottom"])})
        for line in page.extract_text_lines():
            inside = any(t.bbox[0] <= line["x0"] and line["top"] >= t.bbox[1] and line["bottom"] <= t.bbox[3] for t in tables)
            if inside or not line["text"].strip():
                continue
            sizes = sorted(char["size"] for char in line["chars"])
            bold = sum(1 for char in line["chars"] if "bold" in char.get("fontname", "").lower()) > len(line["chars"]) / 2
            lines.append({"page": number, "top": line["top"], "bottom": line["bottom"], "text": line["text"].strip(),
                          "size": sizes[len(sizes) // 2], "bold": bold, "edge": edge(line["top"], line["bottom"])})
    return sorted(lines, key=lambda entry: (entry["page"], entry["top"]))

def blocks_from_pdf(data, file_name, state):
    """PDF keeps no structure, so this reader is the weakest (the manual recommends .docx where
    both exist). With pdfplumber it works from where each line sits and how it is set: page
    headers and footers are left out; a heading is a short line set larger or bolder than the
    body, numbered or not; a line opening with a bullet is a list item; a sentence running over a
    page break is joined again; tables are single blocks; pictures are Figure blocks, read by OCR
    where installed; lines dense in mathematical characters are unread Equation blocks. A paragraph
    is labelled with its page and place ("p.4 \u00b62"), which is how a person finds it in a PDF.
    With only pypdf: text and numbering alone. With neither: one block that could not be read."""
    try:
        import pdfplumber
    except ImportError:
        return blocks_from_pdf_text_only(data, file_name, state)
    blocks, open_block, on_page = [], None, {}
    with pdfplumber.open(io.BytesIO(data)) as document:
        lines = without_page_furniture(pdf_lines(document, file_name), len(document.pages), state, file_name)
        sizes = sorted(entry["size"] for entry in lines if "size" in entry)
        body_size = sizes[len(sizes) // 2] if sizes else 10.0
        numbered = any("size" in entry and first_numbering(entry["text"], state.rules)[0] for entry in lines)
        for entry in lines:
            locator = "page %d" % entry["page"]
            if "table" in entry:
                blocks.append(table_from_rows(entry["table"], locator))
            elif "figure" in entry:
                words = picture_words(document.pages[entry["page"] - 1], entry["box"], state)
                said = "Picture on page %d" % entry["page"]
                blocks.append(new_block("figure", said + words, locator, image_sha256=entry["figure"], display=said + words))
            if "size" not in entry:
                open_block = None
                continue
            text, bullet = entry["text"], BULLET.match(entry["text"])
            dense = sum(1 for char in text if char in MATH_CHARACTERS) / max(1, len(text.replace(" ", "")))
            stands_out = entry["size"] > body_size * 1.08 or (entry["bold"] and entry["size"] >= body_size * 0.97)
            title = stands_out and not bullet and len(text) < 160 and (
                first_numbering(text, state.rules)[0] or (len(text.split()) <= 14 and text[-1] not in ".,;:"))
            near = open_block and entry["page"] == open_block["page"] and entry["top"] - open_block["bottom"] < 0.7 * entry["size"]
            over_the_page = open_block and entry["page"] == open_block["page"] + 1 and not title and not bullet \
                and open_block["block"]["type"] != "heading" and open_block["block"]["text"][-1:] not in ".:;?!" and text[:1].islower()
            if title and near and open_block["block"]["type"] == "heading":       # a heading set over two lines
                kind = None
            elif title:
                kind, more = "heading", {"level_hint": None if numbered else -int(round(entry["size"] * 2)) + (0 if entry["size"] > body_size * 1.08 else 1)}
            elif dense > 0.25 and len(text.split()) <= 12:
                kind, more = "equation", {"equation": EquationData("pdf", text, False, UNDECIDED_REASONS[1], "")}
            elif bullet:
                kind, more, text = "list_item", {"marker": LIST_MARKER}, text[bullet.end():]
            elif (near and open_block["block"]["type"] in ("paragraph", "list_item")) or over_the_page:
                kind = None
            else:
                on_page[entry["page"]] = on_page.get(entry["page"], 0) + 1
                kind, more = "paragraph", {"numbering": "p.%d \u00b6%d" % (entry["page"], on_page[entry["page"]])}
            if kind is None:
                open_block["block"]["text"] += " " + text
            else:
                blocks.append(new_block(kind, text, locator, **more))
                open_block = {"block": blocks[-1]} if kind != "equation" else None
            if open_block:
                open_block.update(page=entry["page"], bottom=entry["bottom"])
    return blocks

def blocks_from_pdf_text_only(data, file_name, state):
    """The fallback PDF reader: page texts as paragraphs, when the layout-aware reader cannot open the file."""
    try:
        import pypdf
    except ImportError:
        return [not_read_block(file_name, "no PDF reader is installed, so this file was not read")]
    blocks = []
    for page_number, page in enumerate(pypdf.PdfReader(io.BytesIO(data)).pages, start=1):
        for paragraph in re.split(r"\n\s*\n", page.extract_text() or ""):
            text = normalise_text(paragraph)
            if text:
                kind = "heading" if first_numbering(text, state.rules)[0] and len(text) < 100 else "paragraph"
                blocks.append(new_block(kind, text, "page %d" % page_number))
    return blocks

# ---------------------------------------------------------------- levels, references, chunks



KIND_OF_BLOCK = {"paragraph": "Paragraph", "list_item": "Paragraph", "table": "Table", "figure": "Figure",
                 "equation": "Equation"}
NOTES_HEADING = "Notes"                 # the heading the tool puts above a Word file's footnotes and endnotes
LIST_MARKER = "- "                      # how an item of a bulleted list is shown; a numbered one shows its number

def fold_lists(blocks):
    """A list belongs to the paragraph that introduces it: "We apply the following principles:"
    and its three bullets are one thought, and a bullet alone cannot be traced to anything. So
    the items of a list that follows a paragraph are folded into it, each on its own line behind
    its marker, and are not units. Done here, where blocks become chunks, so that it holds alike
    for XML, Word and PDF. A list that follows anything else (a heading, a table) has no
    paragraph to belong to; its items are whole statements and stay units of their own."""
    folded = []
    for block in blocks:
        into = folded[-1] if folded else None
        if block["type"] != "list_item" or into is None or into["type"] != "paragraph" or into["not_read_reason"]:
            folded.append(block)
        elif block["text"]:
            into["display"] = "%s\n%s%s" % (into.get("display") or into["text"], block.get("marker") or LIST_MARKER, block["text"])
            into["text"] = "%s %s" % (into["text"], block["text"])
    return folded

def blocks_to_chunks(blocks, corner, source_file, first_number, state):
    """Blocks to chunks. A heading is not a chunk of its own: it becomes part of the heading
    chain of everything below it. Paragraph numbers restart under every heading. Enforces: R2, R4"""
    prefix = "C" if corner == "canon" else "D"
    chunks, chain, levels, section_numbering = [], [], [], ""
    carried, empty = [], []                          # (heading block, anything under it yet), and those with nothing
    for block in fold_lists(infer_levels(blocks, state.rules)):
        if block["type"] == "heading":
            while levels and levels[-1] >= block["level"]:
                levels.pop()
                chain.pop()
                # A heading carries its words to the units below it. Where one is popped off the
                # chain with no unit ever placed under it, those words reach nothing at all, and
                # a file read as headings alone would say nothing. Such a heading becomes a unit
                # of its own, so that what it says is still there to be cited. Enforces: R13
                was, under = carried.pop()
                if not under:
                    empty.append(was)
            levels.append(block["level"])
            chain.append(block["text"])
            carried.append((block, False))
            section_numbering = block["numbering"]
            # The number belongs to the heading whether the document wrote it (num="1.") or Word
            # left it to be counted back. Without it a citation to "section 2" can be resolved
            # against nothing, and the number itself reaches no unit at all. Enforces: R13
            if block["numbering"] and not block["text"].startswith(block["numbering"]):
                chain[-1] = "%s %s" % (block["numbering"], block["text"])
            continue
        carried = [(was, True) for was, _ in carried]
        kind = KIND_OF_BLOCK[block["type"]]
        text = block.get("display") or block["text"]
        equation = block["equation"]
        if kind == "Paragraph" and equation is None:
            equation = inline_formula(text, state.notation)
        chunks.append(Chunk(
            ref=make_ref(prefix, first_number + len(chunks)), corner=corner, source_file=source_file,
            kind=kind, level=levels[-1] if levels else 0, heading_chain=tuple(chain), numbering=section_numbering,
            text=text, locator=block["locator"], content_hash=content_hash(text + block.get("image_sha256", "")),
            table=block["table"], equation=equation,
            numbering_reconstructed=bool(chain) and block_is_under_reconstructed(blocks, block),
            caption=block["caption"], not_read_reason=block["not_read_reason"],
            para_label=block["numbering"] if kind == "Paragraph" else ""))
    while carried:                                   # whatever is still on the chain when the file ends
        was, under = carried.pop()
        if not under:
            empty.append(was)
    for block in empty:
        chunks.append(Chunk(
            ref=make_ref(prefix, first_number + len(chunks)), corner=corner, source_file=source_file,
            kind="Paragraph", level=block["level"], heading_chain=(), numbering=block["numbering"],
            text=block["text"], locator=block["locator"],
            content_hash=content_hash(block["text"]), table=None, equation=None,
            numbering_reconstructed=False, caption="", not_read_reason="",
            para_label=block["numbering"]))
    return chunks

def block_is_under_reconstructed(blocks, block):
    """Was the numbering of the heading directly above this block reconstructed by counting?"""
    above = None
    for candidate in blocks:
        if candidate is block:
            break
        if candidate["type"] == "heading":
            above = candidate
    return bool(above and above["reconstructed"])


# ---------------------------------------------------------------- the two steps
def atoms_of_file(data, found, file_name, state, repairs):
    """The smallest pieces of text the file holds, counted straight from the file and NOT from
    the blocks the reader made of it. That independence is the whole point: an account drawn
    from the reader's own output could never show what the reader missed. Enforces: R13"""
    try:
        if found == "docx":
            return atoms_of_docx(zipfile.ZipFile(io.BytesIO(data)), safe_xml, file_name)
        if found == "pdf":
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as document:
                return atoms_of_pdf(document, file_name)
        if found == "mhtml":
            message = email.message_from_bytes(data)
            for part in message.walk():
                payload = part.get_payload(decode=True)
                if part.get_content_type() == "text/html" and payload is not None:
                    text = payload.decode(part.get_content_charset() or "utf-8", "replace")
                    return atoms_of_markup(parse_markup(repair_markup(text, file_name, []), file_name, [], tolerant_only=True), file_name, state.rules)
            return []
        if found in ("xml", "html"):
            text = repair_markup(decode_text(data), file_name, [])
            return atoms_of_markup(parse_markup(text, file_name, [], tolerant_only=found == "html"), file_name, state.rules)
        return atoms_of_plain_text(decode_text(data), file_name)
    except Exception:                          # an account that cannot be drawn is written down, never a stopped run
        state.notes.append("%s: the content account could not be drawn for this file." % file_name)
        return []

def read_file_blocks(path, file_name, state, repairs, max_bytes):
    """One input file to blocks, by the format found in its content. Enforces: R6"""
    if os.path.getsize(path) > max_bytes:
        return "too large", [not_read_block(file_name, "the file is larger than the size limit for one input file")]
    with open(path, "rb") as handle:
        data = handle.read()
    state.folder = os.path.dirname(os.path.abspath(path))
    found = detect_format(data, file_name)
    if found in NOT_READ:                    # said in plain words, with a next step; never decoded as text
        state.atoms = [atom("whole file not read", file_name, "")]
        return found, [not_read_block(file_name, "the file is " + NOT_READ[found])]
    if found == "svg":                               # read by the text it holds, as a table or a figure
        state.atoms = []
        return found, svg_blocks(data, file_name, file_name, state)
    if found in CONVERTED:                   # read through markup the walker already reads
        markup, words, how = converted(found, data, file_name)
        state.atoms = atoms_of_plain_text(words, file_name)
        state.notes.append("%s: %s." % (file_name, how))
        return found, blocks_from_markup(markup, file_name, state, repairs)
    state.atoms = atoms_of_file(data, found, file_name, state, repairs)
    if found == "pdf":
        return found, blocks_from_pdf(data, file_name, state)
    if found == "docx":
        return found, blocks_from_docx(data, file_name, state)
    if found == "mhtml":
        return found, blocks_from_mhtml(data, file_name, state, repairs)
    if found in ("xml", "html"):
        return found, blocks_from_markup(decode_text(data), file_name, state, repairs, tolerant_only=found == "html")
    blocks = [new_block("paragraph", part, "paragraph %d" % number)
              for number, part in enumerate(re.split(r"\n\s*\n", decode_text(data)), start=1) if part.strip()]
    return "plain text", blocks

def read_corner(ctx, corner, input_key, label):
    """Read every file of one corner, in file-name order, into chunks numbered in reading order."""
    options = ctx.options
    rules = load_tag_rules(options["inputs"].get("tag_rules"))
    notation = load_notation()
    chunks, repairs, info_rows, accounts, read_as_what = [], [], [], [], []
    root = (options["inputs"].get("roots") or {}).get(input_key)
    for left_out, why in (options["inputs"].get("skipped") or {}).get(input_key, []):
        info_rows.append({"group": label, "item": "%s: left out of the folder" % left_out, "value": "Not read: %s." % why})
    paths = options["inputs"][input_key]
    svgs = {}                                        # every SVG of the corner, by name and by name without .svg
    for path in paths:
        if path.lower().endswith(".svg"):
            svgs.setdefault(os.path.basename(path).lower(), path)
            svgs.setdefault(os.path.splitext(os.path.basename(path))[0].lower(), path)
    consumed = set()
    # documents first, pictures last: a chart an XML refers to is read in place, where the XML puts it,
    # and is not read a second time on its own; one nothing refers to is still read, on its own
    for path in [p for p in paths if not p.lower().endswith(".svg")] + [p for p in paths if p.lower().endswith(".svg")]:
        file_name = os.path.relpath(path, root).replace(os.sep, "/") if root else os.path.basename(path)
        if path in consumed:
            info_rows.append({"group": label, "item": file_name, "value": "Read in place, as part of the document that refers to it."})
            continue
        state = WalkState(dict(rules, read_pictures=ctx.settings.get("read_pictures", True)), notation, {}, {}, [])
        state.settings = ctx.settings
        state.svgs, state.consumed, state.file_name = svgs, consumed, file_name
        try:
            found, blocks = read_file_blocks(path, file_name, state, repairs, int(ctx.settings["max_file_mb"] * 1024 * 1024))
        except Exception as problem:                 # a file that breaks a reader is named, never a stopped run (R2)
            found, blocks = "unreadable", [not_read_block(file_name, reason_for(problem))]
            state.atoms, state.blocks = [atom("whole file not read", file_name, "")], []
        read_as_what.append((file_name, found, blocks[0].get("not_read_reason", "") if len(blocks) == 1 else ""))
        new_chunks = blocks_to_chunks(blocks, corner, file_name, len(chunks) + 1, state)
        chunks.extend(new_chunks)
        plain = [to_plain(chunk) for chunk in new_chunks]
        found_account = account(file_name, state.atoms, plain, state.dropped,
                                        marks_of_rendering(plain)
                                        + [LIST_MARKER] * (len(plain) + 1) + [NOTES_HEADING],
                                        os.path.getsize(path) if os.path.isfile(path) else 0)
        accounts.append(found_account)
        info_rows.extend({"group": label, "item": "%s: content account" % file_name, "value": line}
                         for line in account_lines(found_account))
        counts = {}
        for chunk in new_chunks:
            counts[chunk.kind] = counts.get(chunk.kind, 0) + 1
        summary = ", ".join("%d %s" % (counts[kind], kind.lower() + ("s" if counts[kind] != 1 else ""))
                            for kind in CHUNK_KINDS if kind in counts) or "nothing could be read"
        info_rows.append({"group": label, "item": file_name, "value": "Read as %s: %d units (%s)"
                          % (FORMAT_NAMES.get(found, found), len(new_chunks), summary)})
        info_rows.extend({"group": label, "item": "%s: reading note" % file_name, "value": note} for note in state.notes)
        for tag in sorted(state.unknown_tags):
            seen = state.unknown_tags[tag]
            because = ", because it %s" % seen["reason"] if seen.get("reason") else ""
            info_rows.append({"group": label, "item": "%s: unrecognised tag" % file_name, "value":
                              "Unrecognised tag '%s', %d times, read as %s%s. "
                              "It can be added to the project's tag_rules.yaml."
                              % (tag, seen["count"], seen.get("family") or seen.get("read_as"), because)})
    kind = "chunks_canon" if corner == "canon" else "chunks_doc"
    unreadable = sum(1 for chunk in chunks if chunk.kind == "Equation" and not chunk.equation.readable)
    messages = ["%d units read from %d file(s)." % (len(chunks), len(options["inputs"][input_key]))]
    read_as = {}
    for _, found, _ in read_as_what:
        if found in FORMAT_NAMES:
            read_as[FORMAT_NAMES[found]] = read_as.get(FORMAT_NAMES[found], 0) + 1
    if read_as:                                      # what each file was read as, and what was not read and why
        messages.append("Read as: %s." % ", ".join("%d %s" % (count, name) for name, count in sorted(read_as.items())))
    messages.extend("Not read: %s - %s." % (name, why) for name, found, why in read_as_what
                    if why and (found in NOT_READ or found == "unreadable"))
    skipped = (options["inputs"].get("skipped") or {}).get(input_key, [])
    if skipped:
        messages.append("%d file(s) in the folder were left out; Model_Package_Info lists them." % len(skipped))
    if unreadable:
        messages.append("%d equation(s) could not be read." % unreadable)
    open_accounts = [one for one in accounts if not one["closed"]]
    if open_accounts:
        messages.append("The content account is open on %d file(s); Model_Package_Info says what could not be placed."
                        % len(open_accounts))
    return StepResult({kind: chunks, "read_repairs": repairs, "info_rows": info_rows},
                             {"units": len(chunks), "repairs": len(repairs),
                              "content account open on": len(open_accounts)}, messages)

def read_methodology(ctx):
    """Step 02, read-inputs, first part: the canonical methodology into chunks C-0001, C-0002, ..."""
    return read_corner(ctx, "canon", "methodology", "Methodology files")

def read_documentation(ctx):
    """Step 02, read-inputs, second part: the model documentation into chunks D-0001, D-0002, ..."""
    return read_corner(ctx, "doc", "documentation", "Documentation files")


# ================================================================================================
# ---------------------------------------------------------------- the model package, read into units
TEXT_MEMBERS = (".r", ".txt", ".md", ".rd", ".rmd", ".csv", ".tsv", ".yaml", ".yml", ".json", ".html")
PARSER_NAME = "flowR"

class NotParsed(Exception):
    """One R expression that the tool's reader could not read. The message is a plain reason."""

# ---------------------------------------------------------------- safe unpacking and inventory
ARCHIVE_NAMES = (".tar.gz", ".tgz", ".tar", ".tar.bz2", ".tbz2", ".tar.xz", ".txz", ".zip")

def is_archive(path):
    """Whether a file of the package folder is an archive holding the package, rather than one
    file of a package that was put there unpacked."""
    return path.lower().endswith(ARCHIVE_NAMES) or (os.path.getsize(path) > 0 and tarfile.is_tarfile(path))

def unpack_zip(zip_path, max_member_bytes):
    """A package delivered as a ZIP, under exactly the limits a tarball is read under: no absolute
    path, no path that climbs out, no link, nothing over the cap, and nothing protected by a
    password. The size checked is the one the archive declares, before anything is expanded.
    Enforces: R6"""
    files, refused = {}, []
    with zipfile.ZipFile(zip_path) as archive:
        for member in archive.infolist():
            name = member.filename.replace("\\", "/")
            if name.endswith("/"):
                continue
            if name.startswith("/") or re.match(r"^[A-Za-z]:", name) or ".." in name.split("/"):
                refused.append((name, "its path leaves the package folder"))
            elif (member.external_attr >> 16) & 0o170000 == 0o120000:
                refused.append((name, "it is a link"))
            elif member.file_size > max_member_bytes:
                refused.append((name, "it is larger than the size limit for one file"))
            elif member.flag_bits & 0x1:
                refused.append((name, "it is protected by a password"))
            else:
                files[name] = archive.read(member)
    return files, refused

def loose_package(paths, root, max_member_bytes):
    """A package put in the folder unpacked - its source folder rather than a built tarball -
    read file by file, under the same size limit. People very often have the one and not the
    other."""
    files, refused = {}, []
    for path in paths:
        name = os.path.relpath(path, root).replace(os.sep, "/") if root else os.path.basename(path)
        if os.path.getsize(path) > max_member_bytes:
            refused.append((name, "it is larger than the size limit for one file"))
        else:
            with open(path, "rb") as handle:
                files[name] = handle.read()
    return files, refused

def unpack_package(tar_path, max_member_bytes):
    """Read the tarball member by member, into memory, never onto disk by path. A member with
    an absolute path, a path that climbs out with "..", a link, or a size over the cap is
    refused and reported; everything else is returned as {path: bytes}. Enforces: R6"""
    if tar_path.lower().endswith(".zip") and zipfile.is_zipfile(tar_path):
        return unpack_zip(tar_path, max_member_bytes)
    files, refused = {}, []
    with tarfile.open(tar_path, "r:*") as archive:
        for member in archive:
            name = member.name.replace("\\", "/")
            if name.startswith("/") or re.match(r"^[A-Za-z]:", name) or ".." in name.split("/"):
                refused.append((name, "its path leaves the package folder"))
            elif member.issym() or member.islnk():
                refused.append((name, "it is a link"))
            elif member.isfile() and member.size > max_member_bytes:
                refused.append((name, "it is larger than the size limit for one file"))
            elif member.isfile():
                files[name] = archive.extractfile(member).read()
            elif not member.isdir():
                refused.append((name, "it is not a regular file"))
    return files, refused

def strip_top_folder(files):
    """R tarballs hold one top folder named after the package; paths are shown without it."""
    tops = {path.split("/", 1)[0] for path in files if "/" in path}
    if len(tops) == 1 and all("/" in path for path in files):
        return {path.split("/", 1)[1]: data for path, data in files.items()}
    return dict(files)

def read_description(text):
    """The fields of DESCRIPTION (name: value, continuation lines start with white space)."""
    fields, current = {}, None
    for line in text.split("\n"):
        if line[:1] in (" ", "\t") and current:
            fields[current] += " " + line.strip()
        elif ":" in line:
            current, value = line.split(":", 1)
            fields[current.strip()] = value.strip()
    return fields


def is_exported(name, namespace):
    """Is `name` exported by this NAMESPACE (by name, by pattern or as an S3 method)? None when there is no NAMESPACE."""
    if namespace is None:
        return None
    return name in namespace["exports"] or any(re.search(pattern, name) for pattern in namespace["patterns"])

# ---------------------------------------------------------------- the expression tree's node
@dataclass
class Node:
    """One node of the syntax tree. kind is one of: num, str, name, call, index, binary,
    unary, dollar, ns, function, block, if, for, while, repeat, paren, missing.
    For a call, args[0] is what is called and `names` holds the argument names ("" when the
    argument is positional). For a function, `names` are the formals, args[:-1] their
    defaults (kind "missing" when there is none) and args[-1] the body."""
    kind: str; value: str = ""; args: tuple = (); names: tuple = (); line: int = 0; end_line: int = 0

ASSIGNMENT_SIGNS = ("<-", "<<-", "=", "->", "->>")


# ---------------------------------------------------------------- R code, read by flowR
# flowR reads the package's R code; its syntax tree is turned into the tool's Node, the shape every reader of
# R code here expects. One flowR run reads a whole package, and every source is read once per session.
PARSED = {}                               # R source text -> its top-level expressions

def flowr_answer(folder, target, queries, work, *flags):
    """flowR's answer to `queries` about `target` - an R file or a folder of them - read from a file: a pipe
    can lose the end of a large answer. An answer too large to read safely is refused before it is read."""
    answer = os.path.join(work, "answer.json")
    with open(answer, "w", encoding="utf-8") as sink:
        subprocess.run([os.path.join(folder, "flowr"), "--no-ansi", *flags, "--default-engine", "tree-sitter",
                        "--engine.r-shell.disabled", "--engine.tree-sitter.wasm-path", os.path.join(folder, "tree-sitter-r.wasm"),
                        "--engine.tree-sitter.tree-sitter-wasm-path", os.path.join(folder, "tree-sitter.wasm"), "--execute",
                        ":query* %s file://%s" % (json.dumps(queries, separators=(",", ":")), target)], stdout=sink, stderr=subprocess.DEVNULL, cwd=work, timeout=900)
    if os.path.getsize(answer) > FLOWR_ANSWER_MAX:
        raise ValueError("flowR's answer is %d MB, too large to read safely" % (os.path.getsize(answer) // 2 ** 20))
    with open(answer, encoding="utf-8") as handle:
        said = handle.read()
    if "{" not in said:
        raise RuntimeError("flowR gave no answer for the package's R code.")
    return json.JSONDecoder().raw_decode(said[said.index("{"):])[0]

def flowr_files(found):
    """The syntax tree of every file in a flowR answer, by path."""
    tree = found["normalized-ast"].get("normalized", found["normalized-ast"]).get("ast")
    return {os.path.normpath(entry["filePath"]): entry["root"] for entry in tree.get("files", [])}

def top_level(root, text):
    """The top-level expressions of one file; one flowR could not give in the tool's shape is 'not read'. Enforces: R2"""
    found = []
    for child in (root or {}).get("children", []):
        try:
            found.append(flowr_node(child))
        except (KeyError, ValueError, TypeError, IndexError) as problem:
            span = (child.get("info") or {}).get("fullRange") or child.get("location") or [1, 0, text.count("\n") + 1, 0]
            found.append(("not read", span[0], span[2], "its syntax tree could not be used (%s)" % problem))
    return found

def flowr_package(files, folder=None):
    """ONE flowR run over the package's R files and NAMESPACE: every R file's syntax tree is kept for
    parse_r_source, and the NAMESPACE is returned as flowR reads it - exported names and export patterns -
    or None when the package has none."""
    folder = folder or flowr_ready()
    with tempfile.TemporaryDirectory(prefix="pkg-") as work:
        root, texts = os.path.join(work, "package"), {}
        for path, data in files.items():
            if path == "NAMESPACE" or path.lower().endswith(".r"):
                target = os.path.normpath(os.path.join(root, path))
                os.makedirs(os.path.dirname(target), exist_ok=True)
                texts[target] = decode_text(data)
                with open(target, "w", encoding="utf-8") as handle:
                    handle.write(texts[target])
        try:
            found = flowr_answer(folder, root, [{"type": "files"}, {"type": "normalized-ast"}], work)
        except (RuntimeError, ValueError, subprocess.TimeoutExpired):  # too large for one answer: each source is read when asked for
            for path in list(texts):
                if path.lower().endswith(".r"):
                    os.remove(path)
            found = flowr_answer(folder, root, [{"type": "files"}], work)
            found["normalized-ast"] = {"normalized": {"ast": {"files": []}}}
    for path, tree in flowr_files(found).items():
        if path in texts and path.lower().endswith(".r"):
            PARSED[texts[path]] = top_level(tree, texts[path])
    namespace = next((entry["content"]["current"] for entry in found["files"]["files"] if "namespace" in entry.get("roles", [])), None)
    return {"exports": set(namespace["exportedSymbols"]), "patterns": list(namespace["exportedPatterns"])} if namespace else None

def parse_r_sources(sources, folder=None):
    """Read in ONE flowR run every source not read yet this session."""
    todo = [text for text in dict.fromkeys(sources) if text not in PARSED]
    if not todo:
        return
    folder = folder or flowr_ready()
    with tempfile.TemporaryDirectory(prefix="src-") as work:
        root = os.path.join(work, "sources")
        os.makedirs(root)
        names = {}
        for number, text in enumerate(todo):
            names[os.path.normpath(os.path.join(root, "s%05d.R" % number))] = text
            with open(os.path.join(root, "s%05d.R" % number), "w", encoding="utf-8") as handle:
                handle.write(text)
        try:
            trees = flowr_files(flowr_answer(folder, root, [{"type": "normalized-ast"}], work))
        except (RuntimeError, ValueError, subprocess.TimeoutExpired):
            if len(todo) == 1:
                raise
            for text in todo:                         # too large for one answer: one source at a time
                parse_r_sources([text], folder)
            return
    for path, text in names.items():
        PARSED[text] = top_level(trees[path], text) if path in trees else [("not read", 1, text.count("\n") + 1, "flowR gave no syntax tree for it")]

def parse_r_source(source):
    """The top-level expressions of one R source, as flowR reads them; one that cannot be read becomes
    ("not read", first line, last line, reason) and the rest are still read. Enforces: R2, R7"""
    parse_r_sources([source])
    return PARSED[source]

ARITHMETIC_SIGNS = ("+", "-", "*", "/", "^", "%%", "%/%")
COMPARISON_SIGNS = ("==", "!=", "<", ">", "<=", ">=")

def flowr_node(n):
    """One flowR node as the tool's Node. Kinds and argument order are those Node documents."""
    kind = n["type"]
    token = kind in ("RSymbol", "RLogical", "RNumber", "RString", "RBreak", "RNext")
    rng = (n.get("location") if token else None) or (n.get("info") or {}).get("fullRange") or n.get("location") or [0, 0, 0, 0]
    line, end = rng[0], rng[2]
    conv = flowr_node
    if kind == "RExpressionList":
        grouping = [g for g in (n.get("grouping") or []) if isinstance(g, dict)]
        children = tuple(conv(c) for c in n.get("children", []))
        if grouping:
            opening, closing = grouping[0], grouping[-1]
            line, end = opening["location"][0], closing["location"][2]
            if opening.get("lexeme") == "(":
                return Node("paren", args=children[:1], line=line, end_line=end)
            return Node("block", args=children, line=line, end_line=end)
        return children[0] if len(children) == 1 else Node("block", args=children, line=line, end_line=end)
    if kind == "RSymbol":
        content = n.get("content", n.get("lexeme", ""))
        if isinstance(content, list):                    # pkg::name comes as [name, package, internal]
            return namespaced(content[1], ":::" if len(content) > 2 and content[2] else "::", str(content[0]), line, end)
        name = str(content)
        written = str(n.get("lexeme", ""))
        if n.get("namespace") or "::" in written:
            return namespaced(n.get("namespace") or written.split(":")[0], ":::" if ":::" in written else "::", name, line, end)
        return Node("name", name, line=line, end_line=end)
    if kind == "RLogical":
        return Node("name", "TRUE" if n.get("content") else "FALSE", line=line, end_line=end)
    if kind == "RNumber":
        return Node("num", n.get("lexeme", ""), line=line, end_line=end)
    if kind == "RString":
        return Node("str", (n.get("content") or {}).get("str", ""), line=line, end_line=end)
    if kind in ("RBreak", "RNext"):
        return Node("name", "break" if kind == "RBreak" else "next", line=line, end_line=end)
    if kind in ("RBinaryOp", "RPipe"):
        op = "|>" if kind == "RPipe" else "^" if n["operator"] == "**" else n["operator"]
        left, right = conv(n["lhs"]), conv(n["rhs"])
        return Node("binary", op, (left, right), line=left.line, end_line=max(end, right.end_line))
    if kind == "RUnaryOp":
        return Node("unary", n["operator"], (conv(n["operand"]),), line=line, end_line=end)
    if kind == "RFunctionDefinition":
        names, defaults = [], []
        for p in n.get("parameters", []):
            names.append(str(p["name"].get("content", p["name"].get("lexeme", ""))))
            default = p.get("defaultValue")
            defaults.append(conv(default) if default else Node("missing", line=p["name"]["location"][0]))
        return Node("function", args=tuple(defaults) + (conv(n["body"]),), names=tuple(names), line=line, end_line=end)
    if kind in ("RFunctionCall", "RAccess"):
        arguments, names = [], []
        for a in (n.get("arguments") if kind == "RFunctionCall" else n.get("access")) or []:
            if not isinstance(a, dict) or a.get("type") != "RArgument" or a.get("value") is None:
                arguments.append(Node("missing", line=line)); names.append(
                    str(a["name"].get("content", a["name"].get("lexeme", ""))) if isinstance(a, dict) and a.get("name") else "")
                continue
            arguments.append(conv(a["value"]))
            names.append(str(a["name"].get("content", a["name"].get("lexeme", ""))) if a.get("name") else "")
        if kind == "RAccess":
            target, op = conv(n["accessed"]), n["operator"]
            if op in ("$", "@"):
                field = arguments[0] if arguments else Node("missing")
                field = Node(field.kind, field.value, line=field.line) if field.kind in ("name", "str") else field
                return Node("dollar", op, (target, field), line=target.line, end_line=end)
            return Node("index", op, (target,) + tuple(arguments), tuple(names), line=target.line, end_line=end)
        if n.get("infixSpecial"):
            return Node("binary", n["functionName"]["content"], tuple(arguments), line=arguments[0].line,
                        end_line=max([end] + [a.end_line for a in arguments]))
        if n.get("named"):
            written = str(n.get("lexeme", ""))
            callee = conv(n["functionName"])
            if "::" in written and callee.kind == "name":
                callee = namespaced(written.split(":")[0], ":::" if ":::" in written else "::", callee.value, callee.line, callee.end_line)
        else:
            callee = conv(n["calledFunction"])
        return Node("call", "(", (callee,) + tuple(arguments), tuple(names), line=callee.line, end_line=end)
    if kind == "RIfThenElse":
        parts = [conv(n["condition"]), conv(n["then"])] + ([conv(n["otherwise"])] if n.get("otherwise") else [])
        return Node("if", args=tuple(parts), line=line, end_line=end)
    if kind == "RForLoop":
        return Node("for", str(n["variable"].get("content", "")), (conv(n["vector"]), conv(n["body"])), line=line, end_line=end)
    if kind == "RWhileLoop":
        return Node("while", args=(conv(n["condition"]), conv(n["body"])), line=line, end_line=end)
    if kind == "RRepeatLoop":
        return Node("repeat", args=(conv(n["body"]),), line=line, end_line=end)
    raise ValueError("flowR node %s is not one the tool reads" % kind)


def namespaced(package, sign, name, line, end):
    """pkg::name as the tool's parser made it: the package, then the name, which carries no end line."""
    return Node("ns", sign, (Node("name", package, line=line, end_line=line), Node("name", name, line=line)), line=line, end_line=end)

def walk(node):
    """Every node of a tree, parents first."""
    yield node
    for child in node.args:
        yield from walk(child)

def callee_name(node):
    """The name of the function a call node calls: f(...) and pkg::f(...) both give "f"."""
    target = node.args[0]
    if target.kind == "ns":
        target = target.args[1]
    return target.value if target.kind in ("name", "str") else ""

def assignment_parts(node):
    """(target, value) when the node is an assignment, whichever arrow it uses; else None."""
    if node.kind != "binary" or node.value not in ASSIGNMENT_SIGNS:
        return None
    left, right = node.args
    return (right, left) if node.value in ("->", "->>") else (left, right)


def unparse(node):
    """The tree back as short R-like text: used to show defaults and conditions as written."""
    kind, args = node.kind, node.args
    if kind in ("num", "name"):
        return node.value
    if kind == "str":
        return '"%s"' % node.value
    if kind == "missing":
        return ""
    if kind == "paren":
        return "(%s)" % unparse(args[0])
    if kind == "unary":
        return node.value + unparse(args[0])
    if kind in ("binary", "dollar", "ns"):
        tight = kind != "binary" or node.value in ("^", ":")
        return ("%s%s%s" if tight else "%s %s %s") % (unparse(args[0]), node.value, unparse(args[1]))
    if kind in ("call", "index"):
        shown = [("%s = " % name if name else "") + unparse(arg) for name, arg in zip(node.names, args[1:])]
        closer = {"(": ")", "[": "]", "[[": "]]"}[node.value]
        return "%s%s%s%s" % (unparse(args[0]), node.value, ", ".join(shown), closer)
    if kind == "if":
        return "if (%s) %s%s" % (unparse(args[0]), unparse(args[1]), " else " + unparse(args[2]) if len(args) == 3 else "")
    if kind == "function":
        return "function(%s) ..." % ", ".join(node.names)
    return "{ ... }" if kind == "block" else kind

def has_arithmetic(node, function_map):
    """Does this code compute something: arithmetic, a mathematical function, a comparison with
    a number, or a number? Every number counts, whatever its value; only a position inside [ ]
    does not. What this decides is a label - Formula statement or Top-level statement - and both
    are read, shown, linked and asked about alike."""
    positions = set()
    for inner in walk(node):
        if inner.kind == "index":
            positions.update(id(n) for arg in inner.args[1:] for n in walk(arg) if n.kind == "num")
    for inner in walk(node):
        if inner.kind == "function" and inner is not node:
            continue
        if inner.kind == "binary" and inner.value in ARITHMETIC_SIGNS:
            return True
        if inner.kind == "unary" and inner.value == "-" and inner.args[0].kind != "num":
            return True
        if inner.kind == "call" and callee_name(inner) in function_map["r_functions"]:
            return True
        if inner.kind == "binary" and inner.value in COMPARISON_SIGNS and "num" in (inner.args[0].kind, inner.args[1].kind):
            return True
        if inner.kind == "num" and id(inner) not in positions:
            return True
    return False

def code_facts(node, skip_inner_functions=True):
    """The functions a piece of code calls and the names it assigns, in order of appearance."""
    calls, written = [], []
    def visit(inner, top):
        if inner.kind == "function" and not top and skip_inner_functions:
            return
        parts = assignment_parts(inner)
        if parts and parts[0].kind in ("name", "str"):
            written.append(parts[0].value)
            visit(parts[1], False)
            return
        if inner.kind == "call":
            name = callee_name(inner)
            if name:
                calls.append(name)
            for arg in (inner.args[1:] if name else inner.args):
                visit(arg, False)
            return
        if inner.kind in ("dollar", "ns"):
            visit(inner.args[0], False) if inner.kind == "dollar" else None
            return
        for child in inner.args:
            visit(child, False)
    visit(node, True)
    unique = lambda values: tuple(dict.fromkeys(values))
    return {"calls": unique(calls), "symbols_written": unique(written)}


# ---------------------------------------------------------------- units from R source
def draft(kind, path, lines, name, text, **more):
    """A unit before it has its reference. References are given at the end, in reading order."""
    unit = {"kind": kind, "file": path, "lines": lines, "name": name, "inside": "", "text": text,
            "parent_key": None, "key": "%s:%s:%s:%s" % (path, lines[0] if lines else 0, kind, name),
            "code": None, "data": None, "roxygen": None, "read_problem": None, "node": None}
    unit.update(more)
    return unit

def code_detail(node, function_map, settings, **more):
    """The facts about a piece of code that later steps use: its calls, what it assigns, what it reads."""
    facts = code_facts(node)
    return dict(facts, formals=(), exported=None, reads_data=(), **more)


def function_units(name, function_node, lines, path, source_lines, context, inside="", parent_key=None):
    """A function, as one unit: what is written inside it is part of it."""
    function_map = context["function_map"]
    formals = tuple((formal, unparse(default)) for formal, default in zip(function_node.names, function_node.args[:-1]))
    detail = code_detail(function_node, function_map, context["settings"], )
    detail.update(formals=formals, exported=is_exported(name, context["namespace"]) if not inside else False)
    unit = draft(KIND_FUNCTION, path, lines, name, "\n".join(source_lines[lines[0] - 1:lines[1]]),
                 inside=inside, parent_key=parent_key, code=detail, node=function_node)
    return [unit]

def statement_units(node, path, source_lines, context, in_tests):
    """The unit(s) of one top-level expression."""
    function_map, lines = context["function_map"], (node.line, node.end_line)
    text = "\n".join(source_lines[lines[0] - 1:lines[1]])
    parts = assignment_parts(node)
    if parts and parts[1].kind == "function" and parts[0].kind in ("name", "str"):
        return function_units(parts[0].value, parts[1], lines, path, source_lines, context)
    if node.kind == "call" and callee_name(node) in ("test_that", "it", "describe") and len(node.args) >= 2:
        label = node.args[1].value if node.args[1].kind == "str" else "test"
        return [draft(KIND_TEST, path, lines, label, text, code=code_detail(node, function_map, context["settings"]), node=node)]
    detail = code_detail(node, function_map, context["settings"])
    if parts and parts[0].kind == "name" and has_arithmetic(parts[1], function_map):
        return [draft(KIND_FORMULA, path, lines, parts[0].value, text, code=detail, node=node)]
    kind = KIND_TEST if in_tests and node.kind == "call" and callee_name(node).startswith("expect_") else KIND_TOPLEVEL
    name = parts[0].value if parts and parts[0].kind in ("name", "str") else (callee_name(node) if node.kind == "call" else "")
    return [draft(kind, path, lines, name, text, code=detail, node=node)]

def units_from_r_source(path, source, context, line_offset=0):
    """All units of one R file: roxygen blocks, functions, formula statements, test blocks,
    top-level statements, and a "File not read" unit for each expression that could not be
    parsed. Comment lines are attached to the unit that follows them (or, at the end of the
    file, to the unit before), so every non-blank line lies inside a unit. Enforces: R2"""
    source_lines = source.split("\n")
    units, in_tests = [], path.startswith("tests/")
    parsed = parse_r_source(source)
    for result in parsed:
        if isinstance(result, tuple):
            _, first, last, reason = result
            units.append(draft(KIND_NOT_READ, path, (first, last), "", "\n".join(source_lines[first - 1:last]),
                               read_problem="the tool's R reader could not read this expression: %s" % reason))
        else:
            units.extend(statement_units(result, path, source_lines, context, in_tests))
    units.extend(roxygen_units(path, source_lines, parsed, context))
    units.sort(key=lambda unit: (unit["lines"][0], 0 if unit["kind"] == KIND_ROXYGEN else 1, unit["lines"][1] * -1))
    covered = set()
    for unit in units:
        covered.update(range(unit["lines"][0], unit["lines"][1] + 1))
    top_level = [unit for unit in units if not unit["parent_key"]]
    for number, line in enumerate(source_lines, start=1):
        if line.strip() and number not in covered and top_level:
            after = [unit for unit in top_level if unit["lines"][0] > number]
            owner = after[0] if after else top_level[-1]
            owner["lines"] = (min(owner["lines"][0], number), max(owner["lines"][1], number))
            covered.add(number)
    if not units and source.strip():
        units.append(draft(KIND_TOPLEVEL, path, (1, len(source_lines)), "", source, code=None))
    units = self_contained(units)
    for unit in units:
        if unit["parent_key"] is None or unit["kind"] == KIND_FUNCTION:
            unit["text"] = "\n".join(source_lines[unit["lines"][0] - 1:unit["lines"][1]])
        unit["lines"] = (unit["lines"][0] + line_offset, unit["lines"][1] + line_offset)
    return units

def self_contained(units):
    """Rows that never overlap, each a whole piece of code: what is written inside a function is the
    function's; a roxygen block is one row with what it documents; and two rows over the same lines are
    one row."""
    kept = []
    for unit in sorted((u for u in units if not u["parent_key"]), key=lambda u: (u["lines"][0], -u["lines"][1])):
        last = kept[-1] if kept else None
        if last and last["kind"] == KIND_ROXYGEN and unit["kind"] not in (KIND_ROXYGEN, KIND_NOT_READ):
            unit["roxygen"], unit["lines"] = last["roxygen"], (last["lines"][0], unit["lines"][1])
            kept[-1] = unit
        elif last and unit["lines"][0] <= last["lines"][1]:
            if unit["data"] and not last["data"]:            # a stored object keeps what it is
                last["kind"], last["name"] = unit["kind"], unit["name"]
            last["lines"] = (last["lines"][0], max(last["lines"][1], unit["lines"][1]))
            for part in ("data", "roxygen", "code"):
                last[part] = last[part] or unit[part]
        else:
            kept.append(unit)
    return kept

# ---------------------------------------------------------------- roxygen blocks and help pages

def roxygen_units(path, source_lines, parsed, context):
    """Consecutive #' lines form one block. It documents the object that follows: a function,
    a quoted name (a data object), or NULL together with an @name tag."""
    units, number = [], 0
    while number < len(source_lines):
        if not re.match(r"\s*#'", source_lines[number]):
            number += 1
            continue
        first = number
        while number < len(source_lines) and re.match(r"\s*#'", source_lines[number]):
            number += 1
        block_lines = [re.sub(r"^\s*#' ?", "", line) for line in source_lines[first:number]]
        tags, current = [], {"tag": "description", "name": "", "text": "", "line": first + 1}
        for offset, line in enumerate(block_lines):
            found = re.match(r"\s*@(\w+)\s*(.*)", line)
            if found:
                tags.append(current)
                current = {"tag": found.group(1), "name": "", "text": found.group(2).strip(), "line": first + 1 + offset}
                if current["tag"] == "param":
                    pieces = current["text"].split(None, 1)
                    current["name"], current["text"] = (pieces + ["", ""])[0], (pieces + ["", ""])[1]
            else:
                current["text"] = (current["text"] + "\n" + line).strip("\n")
        tags.append(current)
        tags = [tag for tag in tags if tag["tag"] != "description" or tag["text"].strip()]
        follower = next((r for r in parsed if not isinstance(r, tuple) and r.line >= number + 1), None)
        documents = ""
        if follower is not None:
            parts = assignment_parts(follower)
            if parts and parts[0].kind in ("name", "str"):
                documents = parts[0].value
            elif follower.kind == "str":
                documents = follower.value
        documents = documents or next((tag["text"].split()[0] for tag in tags if tag["tag"] == "name" and tag["text"]), "")
        detail = {"documents_ref": None, "documents_name": documents, "tags": tuple(tags)}
        units.append(draft(KIND_ROXYGEN, path, (first + 1, number), documents, "\n".join(source_lines[first:number]),
                           roxygen=detail))
    return units


def vignette_units(path, text, context):
    """A vignette: prose becomes "Vignette text" units, one per stretch between code chunks;
    code chunks are read as R code (and never run)."""
    units, lines, prose_start, number = [], text.split("\n"), 0, 0
    def close_prose(end):
        prose = "\n".join(lines[prose_start:end]).strip()
        if prose:
            units.append(draft(KIND_VIGNETTE, path, (prose_start + 1, end), "", prose))
    while number < len(lines):
        if re.match(r"^```+\s*\{r\b", lines[number]):
            close_prose(number)
            end = number + 1
            while end < len(lines) and not lines[end].startswith("```"):
                end += 1
            units.extend(units_from_r_source(path, "\n".join(lines[number + 1:end]), context, line_offset=number + 1))
            number = prose_start = end + 1
        else:
            number += 1
    close_prose(len(lines))
    return units

# ---------------------------------------------------------------- stored parameter data
COMPRESSION = ((b"\x1f\x8b", "gzip"), (b"BZh", "bzip2"), (b"\xfd7zXZ\x00", "xz"))

def expand(data, cap):
    """Undo gzip, bzip2 or xz compression, recognised from the first bytes and never from the
    extension, and stop when the content grows beyond the cap."""
    method = next((name for magic, name in COMPRESSION if data.startswith(magic)), "none")
    if method == "none":
        return data, method
    opener = {"gzip": lambda: zlib.decompressobj(31), "bzip2": bz2.BZ2Decompressor, "xz": lzma.LZMADecompressor}[method]()
    expanded = opener.decompress(data, cap + 1)
    if len(expanded) > cap:
        raise NotParsed("it expands to more than %d MB, so it was not read; it looks like a dataset" % (cap // 1048576))
    return expanded, method

def cell_text(value):
    """One stored value as its canonical text: numbers as the shortest decimal that gives the
    same stored value back, missing values as NA, logical values as TRUE/FALSE, dates in ISO form."""
    import numpy
    import pandas
    if value is None or value is pandas.NA or value is pandas.NaT:
        return "NA"
    if isinstance(value, (bool, numpy.bool_)):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, numpy.integer)):
        return str(int(value))
    if isinstance(value, (float, numpy.floating)):
        number = float(value)
        if number != number:
            return "NA"
        if number in (float("inf"), float("-inf")):
            return "Inf" if number > 0 else "-Inf"
        return plain_decimal(Decimal(repr(number)))
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.decode("utf-8", "replace")
    return str(value)

def table_of(value):
    """A decoded object as a table: (header, rows of canonical cell texts, column types, R-side
    description), or None when the object has no tabular meaning. A data frame keeps its
    columns; a named vector becomes name and value; a matrix keeps its row and column names;
    a list of short values becomes rows of path and value."""
    import numpy
    import pandas
    if isinstance(value, pandas.DataFrame):
        header = [str(column) for column in value.columns]
        rows = [[cell_text(cell) for cell in row] for row in value.itertuples(index=False, name=None)]
        index = list(value.index)
        if index not in (list(range(len(index))), list(range(1, len(index) + 1))):
            header, rows = ["(row name)"] + header, [[cell_text(name)] + row for name, row in zip(index, rows)]
        return header, rows, "data frame"
    if hasattr(value, "dims") and hasattr(value, "coords"):                  # a named vector or matrix
        array = numpy.asarray(value)
        names = [[cell_text(n) for n in numpy.asarray(value.coords[d])] if d in value.coords else None for d in value.dims]
        if array.ndim == 1:
            labels = names[0] or [str(i + 1) for i in range(array.shape[0])]
            return ["name", "value"], [[label, cell_text(cell)] for label, cell in zip(labels, array.tolist())], "named vector"
        if array.ndim == 2:
            columns = names[1] or ["V%d" % (i + 1) for i in range(array.shape[1])]
            labels = names[0] or [str(i + 1) for i in range(array.shape[0])]
            return ["(row name)"] + columns, [[label] + [cell_text(c) for c in row] for label, row in zip(labels, array.tolist())], "matrix"
    if isinstance(value, numpy.ndarray) and value.ndim == 1:
        if value.shape[0] == 1:
            return ["value"], [[cell_text(value.tolist()[0])]], "single value"
        return ["position", "value"], [[str(i + 1), cell_text(c)] for i, c in enumerate(value.tolist())], "vector"
    if isinstance(value, numpy.ndarray) and value.ndim == 2:
        header = ["(row)"] + ["V%d" % (i + 1) for i in range(value.shape[1])]
        return header, [[str(i + 1)] + [cell_text(c) for c in row] for i, row in enumerate(value.tolist())], "matrix"
    if isinstance(value, (dict, list)):
        rows = []
        def flatten(path, inner):
            if isinstance(inner, dict):
                for key in inner:
                    flatten(path + [str(key)], inner[key])
            elif isinstance(inner, list):
                for position, item in enumerate(inner, start=1):
                    flatten(path + [str(position)], item)
            elif isinstance(inner, numpy.ndarray) and inner.ndim == 1 and inner.shape[0] <= 20:
                rows.append(["/".join(path), ", ".join(cell_text(c) for c in inner.tolist())])
            elif isinstance(inner, (str, int, float, bool)) or inner is None:
                rows.append(["/".join(path), cell_text(inner)])
            else:
                raise ValueError("not a short value")
        try:
            flatten([], value)
        except ValueError:
            return None
        return ["path", "value"], rows, "list"
    if isinstance(value, (str, int, float, bool)):
        return ["value"], [[cell_text(value)]], "single value"
    return None


def table_display(header, rows):
    """A stored table as the workbook shows it: its header and every row, one line each, cells joined by "; ".
    Nothing is left out; a table too long for one row of Chunks_Model takes several (model_rows). Enforces: R13"""
    return "\n".join(["; ".join(header)] + ["; ".join(row) for row in rows])

def data_object_unit(name, value, path, settings):
    """One stored object to a unit: a table shown whole (table_display), or, when it has no tabular meaning, described
    in a sentence - either opening with the object's name, its file and its size, so that every link to it can be
    recognised. Too large to be a parameter table, it is marked not assessable: a dataset (asked_rows)."""
    found = table_of(value)
    if found is None:
        described = "An object of Python type %s after decoding; it has no tabular meaning, so it is described and not compared." % type(value).__name__
        detail = {"object_name": name, "container_file": path, "dims": (), "columns": (), "assessable": False,
                  "not_assessable_reason": "it is not a table, a vector or a list of short values"}
        return draft(KIND_OBJECT, path, None, name, "%s - %s: %s" % (name, path, described), data=detail), None
    header, rows, shape = found
    n_cells = len(rows) * len(header)
    too_large = n_cells > settings["max_parameter_cells"] or len(header) > settings["max_parameter_columns"]
    reason = "too large to be a parameter table (%d cells); it looks like a dataset" % n_cells if too_large else None
    detail = {"object_name": name, "container_file": path, "dims": (len(rows), len(header)), "columns": tuple(header),
              "assessable": not too_large, "not_assessable_reason": reason}
    kind = KIND_OBJECT if shape == "list" else KIND_TABLE
    named = "%s - %s: %d row%s, %d column%s" % (name, path, len(rows), "" if len(rows) == 1 else "s",
                                                len(header), "" if len(header) == 1 else "s")
    unit = draft(kind, path, None, name, named + "\n" + table_display(header, rows), data=detail)
    return unit

def decode_data_file(path, data, settings):
    """Decode one stored-data file without R. The file is parsed by `rdata`, which evaluates
    nothing; a function or another language object inside it is described, never called.
    Returns (units, the fact for Model_Package_Info). Enforces: R7"""
    cap = int(settings["max_file_mb"] * 1024 * 1024)
    stem = os.path.splitext(os.path.basename(path))[0]
    if path.lower().endswith((".csv", ".tsv")):
        delimiter = "\t" if path.lower().endswith(".tsv") else ","
        table = [row for row in csv.reader(io.StringIO(decode_text(data)), delimiter=delimiter) if row]
        frame_header, frame_rows = (table[0], table[1:]) if table else ([], [])
        return [data_object_unit_from_rows(stem, frame_header, frame_rows, path, settings)], "%s: text table" % path
    try:
        raw, method = expand(data, cap)
        import rdata
        container = "several objects (.rda)" if raw[:3] in (b"RDX", b"RDA") else "one object (.rds)"
        body = raw[5:] if raw[:3] in (b"RDX", b"RDA") else raw
        layout = {b"X": "binary", b"A": "ASCII", b"B": "native binary"}.get(body[:1], "unknown")
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            converted = rdata.conversion.convert(rdata.parser.parse_data(raw))
        decoded_by = "rdata %s" % rdata.__version__
    except Exception as problem:                                   # any failure means: not decoded, never lost
        reason = str(problem) if isinstance(problem, NotParsed) else "the file could not be decoded as R data"
        unit = draft(KIND_NOT_READ, path, None, stem, "", read_problem="This data file was not read: %s." % reason)
        return [unit], "%s: not decoded (%s)" % (path, type(problem).__name__ if not isinstance(problem, NotParsed) else reason)
    objects = converted if container.startswith("several") and isinstance(converted, dict) else {stem: converted}
    units = [data_object_unit(str(name), objects[name], path, settings) for name in objects]
    fact = "%s: %s, %s layout, %s compression, %d object(s), decoded by %s" % (path, container, layout, method, len(units), decoded_by)
    return units, fact

def data_object_unit_from_rows(name, header, rows, path, settings):
    """A unit for a table given as header and rows (delimited text files under data/ or inst/extdata/)."""
    import pandas
    width = len(header)
    frame = pandas.DataFrame([list(row) + [""] * (width - len(row)) for row in rows], columns=header)
    return data_object_unit(name, frame, path, settings)

# ---------------------------------------------------------------- which code reads which data
def data_reads(function_node, formals, data_names, data_files):
    """Where a function reads stored data, and how the code shows it: a bare symbol that is
    neither an argument nor assigned in the function; data("x"); readRDS(...) or load(...) with
    a literal file name; pkg::x; get("x"). Column and row key are recorded when the syntax
    shows them, as in floors$lgd_floor or floors[floors$segment == "Retail", "lgd_floor"]."""
    written = set(code_facts(function_node)["symbols_written"]) | set(formals)
    reads = {}
    def note(name, how, column="", row_key=""):
        entry = reads.setdefault(name, {"object": name, "how": how, "column": "", "row_key": ""})
        entry["column"], entry["row_key"] = entry["column"] or column, entry["row_key"] or row_key
    for node in walk(function_node):
        if node.kind == "name" and node.value in data_names and node.value not in written:
            note(node.value, "bare symbol")
        elif node.kind == "ns" and node.args[1].value in data_names:
            note(node.args[1].value, "package::name")
        elif node.kind == "call" and callee_name(node) in ("data", "get", "readRDS", "load"):
            for inner in walk(node):
                if inner.kind in ("str", "name") and inner is not node.args[0]:
                    stem = os.path.splitext(os.path.basename(inner.value))[0]
                    if inner.value in data_names:
                        note(inner.value, '%s("%s")' % (callee_name(node), inner.value))
                    elif inner.kind == "str" and stem in data_files:
                        for name in data_files[stem]:
                            note(name, "%s() of the file %s" % (callee_name(node), os.path.basename(inner.value)))
        if node.kind == "dollar" and node.args[0].kind == "name" and node.args[0].value in reads:
            note(node.args[0].value, "", column=node.args[1].value)
        if node.kind == "index" and node.args[0].kind == "name" and node.args[0].value in reads:
            literals = [inner.value for arg in node.args[1:2] for inner in walk(arg) if inner.kind == "str"]
            columns = [arg.value for arg in node.args[2:3] if arg.kind == "str"]
            note(node.args[0].value, "", column=columns[0] if columns else "", row_key=literals[0] if literals else "")
    return tuple(reads[name] for name in sorted(reads))

# ---------------------------------------------------------------- reading the package
DATA_EXTENSIONS = (".rda", ".rdata", ".rds")

def is_data_file(path):
    """Does this path name a stored-data file that the tool decodes?"""
    lowered = path.lower()
    if lowered.endswith(DATA_EXTENSIONS):
        return True
    return lowered.endswith((".csv", ".tsv")) and lowered.split("/")[0] in ("data", "inst")

def is_test_path(path):
    """Is this member one of the package's tests - a file anywhere under tests/: testthat scripts, their helpers
    and setup, fixtures, snapshots? They check the model rather than compute it, so for now they are left out:
    not read into units, and every line of them counted as left out under a named rule (account_of_package).
    Enforces: R13"""
    return path.split("/")[0].lower() == "tests"


def is_parsed_r_file(path):
    """Is this an R source file in a folder whose code the tool parses?"""
    return path.lower().endswith(".r") and path.lower().split("/")[0] in ("r", "tests", "data", "inst", "data-raw", "demo")


def file_units(path, data, context, facts, reader=""):
    """The units of one file of the package, by where it lies and what it is. A reader chosen
    for this member overrides only where the built-in tests give none."""
    if is_test_path(path):                       # left out for now: the package's tests check the model, not compute it
        facts["tests"].append(path)
        return []
    lowered = path.lower()
    if reader and not is_data_file(path) and not lowered.startswith("src/"):
        text = decode_text(data) if b"\x00" not in data[:4096] else None
        if text is not None:
            if reader == "r-source":
                return units_from_r_source(path, text, context)
            if reader == "help-page":                 # generated from the roxygen comments, which are read
                facts["help pages"].append(path)
                return []
            if reader == "vignette":
                return vignette_units(path, text, context)
            if reader in ("r-data", "table-file"):
                units, fact = decode_data_file(path, data, context["settings"])
                facts["data"].append(fact)
                return units
    if is_data_file(path):
        units, fact = decode_data_file(path, data, context["settings"])
        facts["data"].append(fact)
        return units
    if lowered.startswith("src/"):
        return [draft(KIND_COMPILED, path, None, os.path.basename(path), "",
                      read_problem="Compiled code is not read by the tool; it needs a manual review.")]
    text = decode_text(data) if b"\x00" not in data[:4096] else None
    if text is None:
        return [draft(KIND_NOT_READ, path, None, os.path.basename(path), "",
                      read_problem="A binary file of a kind the tool does not know; it needs a manual review.")]
    if is_parsed_r_file(path):
        return units_from_r_source(path, text, context)
    if lowered.endswith(".rd") and lowered.startswith("man/"):   # a help page repeats the roxygen comments in the R files
        facts["help pages"].append(path)
        return []
    if lowered.endswith((".rmd", ".rnw", ".qmd")):
        return vignette_units(path, text, context)
    return [draft(KIND_OTHER, path, (1, text.count("\n") + 1), os.path.basename(path), text)]

def link_documentation_units(units):
    """Tie each roxygen block to the object it documents."""
    by_name = {}
    for unit in units:
        if unit["kind"] in (KIND_FUNCTION, KIND_TABLE, KIND_OBJECT) and not unit["inside"]:
            by_name.setdefault(unit["name"], unit)
    for unit in units:
        if unit["roxygen"]:                               # a block, or the object it is one row with
            target = by_name.get(unit["roxygen"]["documents_name"])
            unit["roxygen"]["documents_ref"] = target["key"] if target else None

def finalise_units(drafts, file_hashes):
    """Give every draft its reference, in reading order, and turn keys into references."""
    refs = {unit["key"]: make_ref("M", number) for number, unit in enumerate(drafts, start=1)}
    units = []
    for unit in drafts:
        if unit["roxygen"] and unit["roxygen"]["documents_ref"]:
            unit["roxygen"]["documents_ref"] = refs.get(unit["roxygen"]["documents_ref"])
        units.append(ModelUnit(
            ref=refs[unit["key"]], kind=unit["kind"], file=unit["file"], lines=unit["lines"], name=unit["name"],
            inside=unit["inside"], text=unit["text"], parent_ref=refs.get(unit["parent_key"]),
            file_sha256=file_hashes[unit["file"]], content_hash=content_hash(unit["text"]),
            code=CodeDetail(**unit["code"]) if unit["code"] else None,
            data=ParameterDataDetail(**unit["data"]) if unit["data"] else None,
            roxygen=RoxygenDetail(**unit["roxygen"]) if unit["roxygen"] else None,
            read_problem=unit["read_problem"]))
    return units, refs

def package_rows(description, namespace, units, facts, refused):
    """The package lines of Model_Package_Info."""
    rows = [("Package", "Name", description.get("Package", "not stated")),
            ("Package", "Version", description.get("Version", "not stated")),
            ("Package", "Title", description.get("Title", "")), ("Package", "Parser", PARSER_NAME)]
    counts = {}
    for unit in units:
        counts[unit.kind] = counts.get(unit.kind, 0) + 1
    rows.extend(("Package", "Units of kind %s" % kind, counts[kind]) for kind in UNIT_KINDS if kind in counts)
    exported = sorted(u.name for u in units if u.kind == KIND_FUNCTION and u.code and u.code.exported)
    rows.append(("Package", "Exported functions", ", ".join(exported) or "none found"))
    if facts["help pages"]:
        rows.append(("Package", "Help pages not read", "%d in man/: generated from the roxygen comments in the R files, which are read"
                     % len(facts["help pages"])))
    if facts["tests"]:
        rows.append(("Package", "Tests not read", "%d in tests/: the package's tests check the model rather than compute it, so "
                     "for now they are left out - not model chunks, not linked, not asked about; every line of them is "
                     "counted as left out" % len(facts["tests"])))
    rows.extend(("Package data", "Data file", fact) for fact in facts["data"])
    rows.extend(("Package data", "Not assessed", "%s (%s): %s" % (u.ref, u.name, u.data.not_assessable_reason))
                for u in units if u.data and not u.data.assessable)
    rows.extend(("Package", "Could not be read", "%s %s lines %s" % (u.ref, u.file, "-".join(str(n) for n in u.lines or ())))
                for u in units if u.kind == KIND_NOT_READ)
    rows.extend(("Package", "Member of the tarball refused", "%s: %s" % entry) for entry in refused)
    return [{"group": group, "item": item, "value": value} for group, item, value in rows if value != ""]


def read_package(ctx):
    """Step 02, read-inputs, third part: the package. Files are taken in a fixed order (DESCRIPTION, NAMESPACE, R/,
    data, man/, tests/, vignettes/, the rest; by name inside each), so references are stable
    for an unchanged tarball. Enforces: R2, R5"""
    tarballs = ctx.options["inputs"]["package"]
    if not tarballs:
        return StepResult({}, {"units": 0}, ["No package was found in 2_Model_Package."])
    limit, root = int(ctx.settings["max_file_mb"] * 1024 * 1024), (ctx.options["inputs"].get("roots") or {}).get("package")
    archives = [path for path in tarballs if is_archive(path)]
    try:
        files, refused = unpack_package(archives[0], limit) if archives else loose_package(tarballs, root, limit)
    except Exception as problem:                     # a package that cannot be opened is named, never a stopped run (R2)
        return StepResult({}, {"units": 0}, ["The package could not be opened: %s." % reason_for(problem)])
    tarballs = archives or tarballs
    files = strip_top_folder(files)
    function_map = yaml.safe_load(R_FUNCTION_MAP_YAML)
    description = read_description(decode_text(files["DESCRIPTION"])) if "DESCRIPTION" in files else {}
    namespace = flowr_package(files)
    context = {"function_map": function_map, "notation": function_map["notation"], "namespace": namespace,
               "settings": ctx.settings}
    order = ("description", "namespace", "r", "data", "inst", "man", "tests", "vignettes")
    def rank(path):
        top = path.split("/")[0].lower()
        return (order.index(top) if top in order else len(order), path)
    chosen, plan_notes = {}, []
    drafts, facts = [], {"data": [], "help pages": [], "tests": []}
    for path in sorted(files, key=rank):
        drafts.extend(file_units(path, files[path], context, facts, chosen.get(path, "")))
    data_names = {unit["name"] for unit in drafts if unit["data"]}
    data_files = {}
    for unit in drafts:
        if unit["data"]:
            data_files.setdefault(os.path.splitext(os.path.basename(unit["file"]))[0], []).append(unit["name"])
    for unit in drafts:
        if unit["kind"] == KIND_FUNCTION:
            formals = [name for name, _ in unit["code"]["formals"]]
            unit["code"]["reads_data"] = data_reads(unit["node"], formals, data_names, data_files)
    link_documentation_units(drafts)
    hashes = {path: sha256_bytes(data) for path, data in files.items()}
    units, refs = finalise_units(drafts, hashes)
    inventory = [{"file": path, "bytes": len(files[path]), "sha256": hashes[path], "swhid": swhid_content(files[path])}
                 for path in sorted(files)]
    for entry in inventory:                              # for identity part 3: the lines that must lie inside a unit
        if is_parsed_r_file(entry["file"]) and not is_test_path(entry["file"]):
            lines = decode_text(files[entry["file"]]).split("\n")
            entry["nonblank_lines"] = [number for number, line in enumerate(lines, start=1) if line.strip()]
    info = {"name": description.get("Package", ""), "version": description.get("Version", ""), "parser": PARSER_NAME,
            "tarball": os.path.basename(tarballs[0]), "files": inventory,
            "rows": package_rows(description, namespace, units, facts, refused)}
    messages = ["%d units read from %d files of the package." % (len(units), len(files))]
    if facts["tests"]:
        messages.append("%d files in tests/ - the package's tests - are left out for now, and counted as such." % len(facts["tests"]))
    r_files = sum(1 for path in files if path.lower().endswith(".r"))
    other_code = sum(1 for path in files if path.lower().endswith((".py", ".sas", ".m", ".jl", ".scala", ".java", ".cpp", ".c")))
    is_r_package = ("DESCRIPTION" in files and "Package:" in decode_text(files["DESCRIPTION"])) or (r_files and r_files >= other_code)
    if not is_r_package:                             # a DESCRIPTION naming the package, or more R than any other code
        # Found on a real run: an R package that keeps its code in one file beside many help pages
        # was called "not an R package" because R files were under a tenth of all its files.
        # the tool's reader of code reads R and nothing else. A model in Python, SAS or MATLAB comes
        # through as files of running text, fully accounted for and impossible to check - so the
        # analyst is told plainly, rather than left to wonder why nothing was linked. Enforces: R2
        note = ("This does not look like an R package (%s). The tool reads the code of R packages only: files in any "
                "other language are kept as running text and nothing in them can be linked or checked."
                % ("it has no DESCRIPTION file" if "DESCRIPTION" not in files else "%d of its %d files are R code" % (r_files, len(files))))
        messages.append(note)
        info["rows"].append({"group": "The package", "item": "what kind of package", "value": note})
    if len(tarballs) > 1:
        messages.append("More than one tarball was found; only %s was read." % os.path.basename(tarballs[0]))
    plain_units = [to_plain(unit) for unit in units]
    account = account_of_package(files, plain_units, refused, lambda path: is_parsed_r_file(path) or os.path.splitext(path)[1].lower() in TEXT_MEMBERS or "/" not in path,
                                 dropped=set(facts["help pages"]) | set(facts["tests"]))
    info["rows"].extend({"group": "The package", "item": "content account", "value": line}
                        for line in account_lines(account))
    if not account["closed"]:
        messages.append("The content account of the package is open; Model_Package_Info says what could not be placed.")
    info["rows"].extend({"group": "The package", "item": "how it was read", "value": note} for note in plan_notes)
    return StepResult({"model_units": units, "package_info": [info],
                              },
                             {"units": len(units), "files": len(files), "members refused": len(refused)}, messages)

# ---------------------------------------------------------------- flowR: the program R code is read with
# flowR (Sihler and Tichy, Ulm University; GPLv3) parses R with tree-sitter and tells, for every name, the
# definition it reads, for every call, the function it calls, and for every argument, the parameter it becomes.
# It is fetched once, verified against the pinned SHA-256, and run in one-shot mode: no R process, no server,
# no open port, no network. Its answer is written to a file - through a pipe it stops at 128 KiB.
FLOWR_VERSION = "2.15.8"
FLOWR_SHA256 = "39e1b9e5e4fab67f76204dbf6856d32e9e7f2a36132b1323e02cc8e3392436e4"
FLOWR_URL = "https://github.com/flowr-analysis/flowr-r-adapter/releases/download/flowr-v{0}/flowr-{0}-linux-x64.tar.gz"
READS, CALLS = 1, 4                     # flowR's edge bits: reads, calls; any other bit is ignored
FLOWR_ANSWER_MAX = 200 * 1024 * 1024    # flowR answers ~64 KB of JSON per line of R, and Python needs ~8x that to read it

FLOWR_FOLDERS = {}                      # once ready in a session, ready for the rest of it

def flowr_ready(source=""):
    """flowR's folder, fetched once from source - a staged archive, such as one in a Volume - or from the
    pinned release, refused unless its SHA-256 is the pinned one, and run once: ready means it runs here.
    The folder is this system user's own: a shared cluster runs every notebook session as a user of its
    own, and a folder another session made can be neither read nor written. Enforces: R12"""
    import tarfile, urllib.request
    if source in FLOWR_FOLDERS:
        return FLOWR_FOLDERS[source]
    folder = os.path.join(tempfile.gettempdir(), "flowr-%s-%d" % (FLOWR_VERSION, os.getuid()))
    binary = os.path.join(folder, "flowr")
    if not all(os.path.exists(os.path.join(folder, name)) for name in ("flowr", "tree-sitter.wasm", "tree-sitter-r.wasm")):
        os.makedirs(folder, mode=0o700, exist_ok=True)
        archive = os.path.join(folder, "flowr.tar.gz")
        shutil.copyfile(source, archive) if source else urllib.request.urlretrieve(FLOWR_URL.format(FLOWR_VERSION), archive)
        if file_sha256(archive) != FLOWR_SHA256:
            os.remove(archive)
            raise RuntimeError("The flowR archive is not the pinned release %s: its checksum differs." % FLOWR_VERSION)
        with tarfile.open(archive) as bundle:
            bundle.extractall(folder, members=[m for m in bundle.getmembers() if m.isfile() and "/" not in m.name.strip("./")])
        os.chmod(binary, 0o700)
    try:                                            # one line of R, read the way the review reads: ready means it ran
        flowr_read(folder, "x <- 1\n")
    except Exception as problem:
        raise RuntimeError("flowR is here but could not run on this cluster: %s" % problem)
    FLOWR_FOLDERS[source] = folder
    return folder

def flowr_read(folder, text, prefix=""):
    """flowR's syntax tree of one R source, indexed by id, and its edges as {id: [(target, bits)]}; every
    id carries the prefix, so that the answers for several sources can stand side by side. An answer too
    large to read safely is refused before it is read: loading it would take the driver's memory."""
    with tempfile.TemporaryDirectory() as work:
        source = os.path.join(work, "package.R")
        with open(source, "w", encoding="utf-8") as handle:
            handle.write(text)
        found = flowr_answer(folder, source, [{"type": "dataflow"}, {"type": "normalized-ast"}], work, "--no-fs")
    tree = found["normalized-ast"].get("normalized", found["normalized-ast"]).get("ast")
    edges = {prefix + str(s): [(prefix + str(t), e["types"]) for t, e in targets] for s, targets in found["dataflow"]["graph"]["edgeInformation"]}
    nodes, stack = {}, [(tree, None)]
    while stack:                                   # every node once, with its parent; no recursion, so no depth limit
        node, up = stack.pop()
        if isinstance(node, dict):
            if "type" in node and "info" in node:
                node["info"]["id"] = prefix + str(node["info"]["id"])
                node["up"], up = up, node["info"]["id"]
                nodes[up] = node
            stack += [(value, up) for key, value in node.items() if key not in ("info", "location", "lexeme", "up")]
        elif isinstance(node, list):
            stack += [(value, up) for value in node]
    return nodes, edges

@dataclass
class FunctionReadings:
    """flowR's readings of a package's functions, each read alone: the function each prefix stands for, the
    functions themselves, those flowR could not read and why, and the syntax tree and edges of the rest, every id
    carrying its reading's prefix. The chunk links add the readings of the package's scripts to the same tree."""
    functions: dict; unit_of: dict = field(default_factory=dict); unread: dict = field(default_factory=dict)
    tree: dict = field(default_factory=dict); edges: dict = field(default_factory=dict)


def read_functions(units, folder):
    """Every function of the package read by flowR, one at a time: memory is then bounded by the largest function,
    not by the package - the whole package at once once took a driver down. Enforces: R7"""
    readings = FunctionReadings({u["name"]: u for u in units if u["kind"] == KIND_FUNCTION and not u.get("inside")})
    for number, unit in enumerate(sorted(readings.functions.values(), key=lambda u: u["ref"])):
        prefix = "%d:" % number
        readings.unit_of[prefix] = unit["name"]
        try:
            tree, edges = flowr_read(folder, unit["text"], prefix)
        except (ValueError, KeyError, OSError, subprocess.SubprocessError) as problem:
            readings.unread[unit["name"]] = "flowR could not read it: %s" % problem
            continue
        readings.tree.update(tree)
        readings.edges.update(edges)
    return readings


# Which units of Chunks_Model a unit takes something from - a function it calls, a variable, a stored
# table or a file another unit defines - and which take something from it. R looks a name up inside the
# function first, then in the script it runs in, then in the package, and so does this. flowR resolves
# what it can: inside a function, a parameter or a value set there; across the statements of one script
# (a test file, a vignette, the top-level code of an R file), the statement that set a variable before it
# was read. What flowR leaves unresolved is what the unit takes from outside, and the package names it:
# its functions, the variables its R files set at top level, and its stored data. The package's own name
# comes before a base function of the same name, as it does in R. Enforces: R2, R4, R7
SCRIPT_KINDS = (KIND_FORMULA, KIND_TOPLEVEL, KIND_TEST)
ASSIGNING, ASSIGNING_RIGHT = ("<-", "<<-", "=", ":="), ("->", "->>")
BY_NAME_CALLS = ("do.call", "match.fun", "get", "get0", "exists", "mget", "data")
BASE_OPERATORS = ("%%", "%/%", "%in%", "%o%", "%*%", "%x%")


def node_id(node):
    return ((node or {}).get("info") or {}).get("id")


def symbol_read(tree, node, package):
    """The name a flowR syntax node reads or calls, and whether it names this package outright
    (package::name), or None when the node reads nothing: an assignment's target, a parameter, the name
    of a named argument, a field after $ or @, a brace, or a name of another package."""
    if node["type"] in ("RBinaryOp", "RUnaryOp"):     # a %op% of the package's own, used in place
        operator = node.get("operator") or ""
        if operator.startswith("%") and operator.endswith("%") and operator not in BASE_OPERATORS:
            return operator, False
        return None
    if node["type"] == "RString":                     # do.call("f"), get("x"), data("x"): a name as text
        up = tree.get(node.get("up") or "", {})
        call = tree.get(up.get("up") or "", {})
        if up.get("type") == "RArgument" and call.get("type") == "RFunctionCall" and \
                (call.get("functionName") or {}).get("lexeme") in BY_NAME_CALLS:
            return (node.get("lexeme") or "").strip("\"'`"), False
        return None
    if node["type"] != "RSymbol":
        return None
    name, me, up = (node.get("lexeme") or "").strip("`"), node_id(node), tree.get(node.get("up") or "", {})
    if not name or name in ("{", "}", "(", ")"):
        return None
    content = node.get("content")
    if isinstance(content, list) and len(content) > 1 and content[1]:
        return (name, True) if content[1] == package else None
    kind = up.get("type")
    if kind == "RParameter" or (kind == "RArgument" and node_id(up.get("name")) == me):
        return None
    if kind == "RBinaryOp" and ((up.get("operator") in ASSIGNING and node_id(up.get("lhs")) == me) or
                                (up.get("operator") in ASSIGNING_RIGHT and node_id(up.get("rhs")) == me)):
        return None
    if kind == "RArgument" and tree.get(up.get("up") or "", {}).get("type") == "RAccess" and \
            tree[up["up"]].get("operator") in ("$", "@"):
        return None
    return name, False


def visible(definer, user):
    """Can code in `user` see what `definer` defines? The package's R files and its stored data are seen
    everywhere; what a test or vignette sets is seen later in the same file only, and a test helper by
    every test."""
    if definer["ref"] == user["ref"]:
        return False
    home = definer["file"]
    if home.startswith("R/") or (definer.get("data") or {}).get("object_name"):
        return True
    if home == user["file"]:
        return (definer.get("lines") or [0])[0] < (user.get("lines") or [0])[0]
    return home.startswith("tests/") and user["file"].startswith("tests/") and \
        os.path.basename(home).startswith(("helper", "setup"))


# Stored data, followed further. Beyond the names flowR leaves unresolved, code reaches a stored object in three
# ways its names alone do not show: through a helper that takes the object's name as text and hands it to get(),
# data(), readRDS() or load() - at any depth of helpers; through an environment the package keeps, where one unit
# stores a value and another reads it back; and by a name put together at run time. The first two become links like
# any other; the third names the stored objects the name could be, each marked "(possible)", or, when nothing of
# the name is known, says so. flowR gives, within each reading, the syntax tree and which definition each symbol
# reads - a parameter, a local variable, or nothing there; the engine binds a call's arguments to a package
# function's parameters by R's rules, follows helpers to a fixed point, puts paste0() and its kind together from
# known text, and matches names against the objects decoded from the package's .rda and .rds files. Enforces: R2, R7
NAME_READERS = {"get": ("x",), "get0": ("x",), "exists": ("x",), "mget": ("x",), "data": ("list",),
                "readRDS": ("file",), "load": ("file",)}
FILE_READERS = ("readRDS", "load")
PACKAGE_ENVIRONMENTS = ("topenv", "asNamespace", "getNamespace", "parent.env", "environment", "globalenv")
TEMPLATE_DEPTH, TEMPLATE_PARTS = 6, 16
POSSIBLE = " (possible)"
UNKNOWN_NAME = "reads stored data by a name known only when it runs"
READ_BY_NO_CODE = "None: no code of the package reads it."


def call_name(node):
    """The function a flowR call node calls, by name, or "" when the call is not by a plain name."""
    function = (node or {}).get("functionName") or {}
    return (function.get("lexeme") or "").strip("`") if function.get("type") == "RSymbol" else ""


def call_arguments(node):
    """A call's arguments, in order, as (name or None, value node); an empty argument is left out."""
    out = []
    for argument in (node or {}).get("arguments") or []:
        if isinstance(argument, dict) and argument.get("type") == "RArgument" and argument.get("value"):
            name = argument.get("name")
            out.append(((name.get("lexeme") or "").strip("`") if isinstance(name, dict) else None, argument["value"]))
    return out


def merged(parts):
    """Adjacent pieces of known text joined, and runs of the unknown made one; too long a name becomes unknown."""
    out = []
    for kind, value in parts:
        if out and kind == "text" and out[-1][0] == "text":
            out[-1] = ("text", out[-1][1] + value)
        elif not (out and kind == "any" and out[-1][0] == "any"):
            out.append((kind, value))
    return out if len(out) <= TEMPLATE_PARTS else [("any", "")]


def template_of(tree, edges, node, depth=0):
    """What a name given to a reader is made of: ("text", s) known text, ("param", p) a parameter of the function it
    is in, ("any", "") what is known only at run time. flowR's edges say what a symbol reads: a parameter, or a local
    variable whose value is followed in turn. paste0(), paste(), sprintf(), file.path() and system.file() are put
    together from their parts; nothing is run. Enforces: R7"""
    if not node or depth > TEMPLATE_DEPTH:
        return [("any", "")]
    kind = node.get("type")
    if kind == "RString":
        return [("text", (node.get("lexeme") or "").strip("\"'"))]
    if kind == "RNumber":
        return [("text", node.get("lexeme") or "")]
    if kind == "RSymbol":
        for target, bits in edges.get(node_id(node), []):
            if not bits & READS or target not in tree:
                continue
            definition = tree[target]
            holder = tree.get(definition.get("up") or "", {})
            if holder.get("type") == "RParameter":
                return [("param", (definition.get("lexeme") or "").strip("`"))]
            if holder.get("type") == "RBinaryOp" and holder.get("operator") in ASSIGNING and node_id(holder.get("lhs")) == target:
                return template_of(tree, edges, holder.get("rhs"), depth + 1)
        return [("any", "")]
    if kind != "RFunctionCall":
        return [("any", "")]
    name, arguments = call_name(node), call_arguments(node)
    def joined(values, separator):
        parts = []
        for number, value in enumerate(values):
            parts += (separator if number else []) + template_of(tree, edges, value, depth + 1)
        return merged(parts)
    if name == "paste0":
        return joined([v for n, v in arguments if n not in ("collapse", "recycle0")], [])
    if name == "paste":
        separator = next((v for n, v in arguments if n == "sep"), None)
        return joined([v for n, v in arguments if n not in ("sep", "collapse", "recycle0")],
                      template_of(tree, edges, separator, depth + 1) if separator else [("text", " ")])
    if name in ("file.path", "system.file"):
        dropped = ("fsep",) if name == "file.path" else ("package", "lib.loc", "mustWork")
        return joined([v for n, v in arguments if n not in dropped], [("text", "/")])
    if name == "sprintf" and arguments:
        form = template_of(tree, edges, arguments[0][1], depth + 1)
        if len(form) == 1 and form[0][0] == "text" and "%" not in form[0][1].replace("%s", ""):
            pieces, values, parts = form[0][1].split("%s"), [v for _, v in arguments[1:]], []
            for number, piece in enumerate(pieces):
                parts.append(("text", piece))
                if number < len(pieces) - 1:
                    parts += template_of(tree, edges, values[number], depth + 1) if number < len(values) else [("any", "")]
            return merged(parts)
    return [("any", "")]


def reader_calls(nodes):
    """The calls among `nodes` that read stored data by a name or by a file's name: (reader, the node giving it).
    data() takes its names as text or in list =, and mget() and data() a vector of them, c("a", "b")."""
    for node in nodes:
        reader = call_name(node) if node.get("type") == "RFunctionCall" else ""
        if reader not in NAME_READERS:
            continue
        arguments = call_arguments(node)
        if reader == "data":
            values = [v for n, v in arguments if n in (None, "list") and v.get("type") != "RSymbol"]
        else:
            named = [v for n, v in arguments if n in NAME_READERS[reader]]
            values = named[:1] or [v for n, v in arguments if n is None][:1]
        for value in values:
            if value.get("type") == "RFunctionCall" and call_name(value) == "c":
                for _, inner in call_arguments(value):
                    yield reader, inner
            else:
                yield reader, value


def access_member(access):
    """The member an access names - floors in cache$floors or cache[["floors"]] - or "" for anything else."""
    items = [a.get("value") for a in access.get("access") or [] if isinstance(a, dict) and a.get("type") == "RArgument"]
    if len(items) != 1 or not items[0]:
        return ""
    if access.get("operator") == "$" and items[0].get("type") == "RSymbol":
        return (items[0].get("lexeme") or "").strip("`")
    return (items[0].get("lexeme") or "").strip("\"'") if items[0].get("type") == "RString" else ""


def bound_arguments(arguments, formals):
    """Which argument of a call each parameter of the function called receives, as R matches them: by exact name,
    then by position up to ..., after which only by name."""
    bound, positional = {}, []
    for name, value in arguments:
        if name is not None and name in formals:
            bound[name] = value
        elif name is None:
            positional.append(value)
    for formal in formals:
        if formal == "...":
            break
        if formal not in bound and positional:
            bound[formal] = positional.pop(0)
    return bound


def name_candidates(reader, parts, data_names, data_files):
    """The stored objects a name partly known could be: those whose name - or, for readRDS() and load(), whose file's
    name - fits the known text, the rest being anything. None when too little of it is known to choose any."""
    if reader in FILE_READERS:                         # only the file's own name counts, not the folders before it
        last = max((i for i, (kind, value) in enumerate(parts) if kind == "text" and "/" in value), default=None)
        if last is not None:
            parts = [("text", parts[last][1].rsplit("/", 1)[1])] + parts[last + 1:]
    known = "".join(value for kind, value in parts if kind == "text")
    if len(re.sub(r"(?i)\.(rds|rda|rdata)$", "", known).strip("/._- ")) < 3:
        return None
    pattern = re.compile("".join(re.escape(value) if kind == "text" else ".*" for kind, value in parts), re.S)
    pool = data_files if reader in FILE_READERS else data_names
    return sorted({ref for name, refs in pool.items() if pattern.fullmatch(name) for ref in refs})


def unit_links(units, parsed, folder, package):
    """The immediate upstream and downstream units of every unit, as records of kind unit_links, each with the
    names that make each link, and the units flowR could not read for their links. `parsed` holds flowR's
    readings of the functions (read_functions); the scripts are read here."""
    by_ref = {u["ref"]: u for u in units}
    definers = {}                                          # name -> the units that define it
    for u in units:
        names = set()
        if u["kind"] == KIND_FUNCTION and not u.get("inside"):
            names.add(u["name"].strip("`"))                # `%||%` is defined in backticks and used without
        elif u["kind"] in (KIND_FORMULA, KIND_TOPLEVEL):
            names |= set((u.get("code") or {}).get("symbols_written") or ())
        if (u.get("data") or {}).get("object_name"):
            names.add(u["data"]["object_name"])
        for name in names:
            definers.setdefault(name, []).append(u["ref"])
    readings, unread = [], set()                           # (prefix, [(first line, last line, ref)])
    for prefix, name in parsed.unit_of.items():             # every function: flowR read it already, alone
        unit = parsed.functions[name]
        if name in parsed.unread:
            unread.add(unit["ref"])
        else:
            readings.append((prefix, [(1, unit["text"].count("\n") + 1, unit["ref"])]))
    scripts = {}
    for u in units:
        if u["kind"] in SCRIPT_KINDS:
            scripts.setdefault(u["file"], []).append(u)
    def read(members, prefix):
        text, spans, line = "", [], 1
        for u in members:
            size = u["text"].count("\n") + 1
            spans.append((line, line + size - 1, u["ref"]))
            text, line = text + u["text"] + "\n\n", line + size + 1
        tree, edges = flowr_read(folder, text, prefix)
        parsed.tree.update(tree)
        parsed.edges.update(edges)
        readings.append((prefix, spans))
    for number, (file, members) in enumerate(sorted(scripts.items())):
        members.sort(key=lambda u: (u.get("lines") or [0])[0])
        try:                                               # one script, its statements in order: one reading
            read(members, "s%d:" % number)
        except (ValueError, KeyError, OSError, subprocess.SubprocessError):
            for place, u in enumerate(members):            # a statement flowR cannot read stops only itself
                try:
                    read([u], "s%d.%d:" % (number, place))
                except (ValueError, KeyError, OSError, subprocess.SubprocessError):
                    unread.add(u["ref"])
    data_files = {}                                       # a data file's name, as code may write it -> its units
    for u in units:
        if u["kind"] in (KIND_TABLE, KIND_OBJECT):
            data_files.setdefault(os.path.basename(u["file"]), []).append(u["ref"])
    data_names = {}                                       # a stored object's name -> its units
    for u in units:
        if (u.get("data") or {}).get("object_name"):
            data_names.setdefault(u["data"]["object_name"], []).append(u["ref"])
    by_prefix = {}
    for key, node in parsed.tree.items():
        by_prefix.setdefault(key.split(":", 1)[0] + ":", []).append(node)
    def locator(spans):
        def unit_at(node):
            while node is not None and not node.get("location"):
                node = parsed.tree.get(node.get("up") or "")
            line = node["location"][0] if node else 0
            return next((ref for first, last, ref in spans if first <= line <= last), None)
        return unit_at
    nodes_of = {}                                          # ref -> the nodes of its code, from every reading
    for prefix, spans in readings:
        unit_at = locator(spans)
        for node in by_prefix.get(prefix, []):
            ref = unit_at(node)
            if ref:
                nodes_of.setdefault(ref, []).append(node)
    tree, edges = parsed.tree, parsed.edges
    def local(symbol):                                     # flowR finds, in this reading, the definition it reads: a
        return any(bits & READS and (tree.get(target) or {}).get("type") == "RSymbol"   # variable or a parameter,
                   for target, bits in edges.get(node_id(symbol), []))                  # not a replacement's own access
    env_writes, env_reads, namespace_writes = environment_members(nodes_of, by_ref, definers, tree, edges, local)
    for name, writers in sorted(namespace_writes.items()): # a name stored in the package's own environment
        for writer in sorted(writers):
            if writer not in definers.setdefault(name, []):
                definers[name].append(writer)
    via = {}                                               # (definer, user) -> names
    def link(definer, user, name):
        via.setdefault((definer, user), set()).add(name)
    for prefix, spans in readings:
        unit_at = locator(spans)
        for node in by_prefix.get(prefix, []):
            if node["type"] == "RString":                  # a data file named in the code: read.csv(system.file(...))
                written = (node.get("lexeme") or "").strip("\"'")
                user = unit_at(node)
                for definer in data_files.get(os.path.basename(written), []) if user else ():
                    if by_ref[definer]["file"].endswith(written) and definer != user:
                        link(definer, user, os.path.basename(written))
            found = symbol_read(parsed.tree, node, package)
            user = unit_at(node) if found else None
            if user is None:
                continue
            name, named_package = found
            up = parsed.tree.get(node.get("up") or "", {})
            ends = [t for t, bits in parsed.edges.get(node_id(node), []) if bits & (CALLS if node["type"] in ("RBinaryOp", "RUnaryOp")
                                                                                 else READS) and t in parsed.tree]
            if up.get("type") == "RFunctionCall" and node_id(up.get("functionName")) == node_id(node):
                ends += [t for t, bits in parsed.edges.get(node_id(up), []) if bits & CALLS and t in parsed.tree]
            if ends and not named_package:                 # flowR found where it is defined: in this reading
                for end in ends:
                    definer = unit_at(parsed.tree[end])
                    if definer and definer != user:
                        link(definer, user, name)
                continue
            for definer in definers.get(name, []):         # not defined here: the package names it
                if visible(by_ref[definer], by_ref[user]):
                    link(definer, user, name)
    for u in units:                                        # stored data the code reads by file or by data()
        for read_data in (u.get("code") or {}).get("reads_data") or ():
            name = read_data.get("object") if isinstance(read_data, dict) else read_data
            for definer in definers.get(name, []):
                if (by_ref[definer].get("data") or {}).get("object_name") == name and definer != u["ref"]:
                    link(definer, u["ref"], name)
    possible, notes = follow_stored_data(units, by_ref, nodes_of, tree, edges, data_names, data_files, link)
    for key, readers in sorted(env_reads.items()):         # a member of an environment, read where it was stored
        for writer in sorted(env_writes.get(key, ())):
            for reader in sorted(readers):
                if writer != reader:
                    link(writer, reader, "%s$%s" % key)
    upstream, downstream, maybe_up, maybe_down = {}, {}, {}, {}
    for definer, user in via:
        upstream.setdefault(user, set()).add(definer)
        downstream.setdefault(definer, set()).add(user)
    for definer, user in possible:
        if (definer, user) not in via:
            maybe_up.setdefault(user, set()).add(definer)
            maybe_down.setdefault(definer, set()).add(user)
    return [{"record_type": "unit_links", "ref": u["ref"], "upstream": sorted(upstream.get(u["ref"], ())),
             "downstream": sorted(downstream.get(u["ref"], ())),
             "possible_upstream": sorted(maybe_up.get(u["ref"], ())), "possible_downstream": sorted(maybe_down.get(u["ref"], ())),
             "notes": sorted(notes.get(u["ref"], ())),
             "via": {definer: sorted(names) for (definer, user), names in sorted(via.items()) if user == u["ref"]},
             "not_read_by_flowr": u["ref"] in unread} for u in units]


def environment_members(nodes_of, by_ref, definers, tree, edges, local):
    """Who stores a member of an environment the package keeps, and who reads it: cache$floors <- x,
    cache[["floors"]] <- x or assign("floors", x, envir = cache), read back as cache$floors, cache[["floors"]] or
    get("floors", envir = cache). The environment is a variable the package defines at top level, which flowR finds
    no definition of inside the reading. Also the names a function stores in the package's own environment -
    assign("x", ..., envir = topenv()) and the like, or x <<- ... - which then define x for the whole package.
    Returns (writes, reads, namespace writes): {(environment, member): refs}, the same, and {name: refs}."""
    writes, reads, namespace = {}, {}, {}
    def key_of(holder, member):
        name = (holder.get("lexeme") or "").strip("`") if holder.get("type") == "RSymbol" and not local(holder) else ""
        return (name, member) if name in definers and member else None
    for ref, nodes in sorted(nodes_of.items()):
        for node in nodes:
            kind = node.get("type")
            if kind == "RBinaryOp" and node.get("operator") in ("<-", "=", "<<-"):
                lhs = node.get("lhs") or {}
                if lhs.get("type") == "RAccess" and lhs.get("operator") in ("$", "[["):
                    key = key_of(lhs.get("accessed") or {}, access_member(lhs))
                    if key:
                        writes.setdefault(key, set()).add(ref)
                elif node.get("operator") == "<<-" and lhs.get("type") == "RSymbol" and by_ref[ref]["kind"] == KIND_FUNCTION:
                    namespace.setdefault((lhs.get("lexeme") or "").strip("`"), set()).add(ref)
            elif kind == "RAccess" and node.get("operator") in ("$", "[["):
                holder = tree.get(node.get("up") or "", {})
                if holder.get("type") == "RBinaryOp" and node_id(holder.get("lhs")) == node_id(node):
                    continue                               # the target of an assignment: a write, above
                key = key_of(node.get("accessed") or {}, access_member(node))
                if key:
                    reads.setdefault(key, set()).add(ref)
            elif kind == "RFunctionCall" and call_name(node) in ("assign", "get", "get0", "exists"):
                arguments = call_arguments(node)
                named = {n: v for n, v in arguments if n}
                positional = [v for n, v in arguments if n is None]
                first = named.get("x") or (positional[0] if positional else None)
                parts = template_of(tree, edges, first) if first is not None else [("any", "")]
                member = parts[0][1] if len(parts) == 1 and parts[0][0] == "text" else ""
                where = named.get("envir") or named.get("pos") or \
                    (positional[2] if call_name(node) == "assign" and len(positional) > 2 else None)
                if not member or where is None:
                    continue
                if where.get("type") == "RFunctionCall" and call_name(where) in PACKAGE_ENVIRONMENTS:
                    if call_name(node) == "assign":
                        namespace.setdefault(member, set()).add(ref)
                    continue
                key = key_of(where, member)
                if key:
                    (writes if call_name(node) == "assign" else reads).setdefault(key, set()).add(ref)
    return writes, reads, namespace


def follow_stored_data(units, by_ref, nodes_of, tree, edges, data_names, data_files, link):
    """Stored data reached by a name given as text (2.12): at each call of get(), data(), readRDS() or load(), what
    the name is made of (template_of). A name wholly known links the stored object it names. A name that holds a
    parameter makes its function a helper, and every call of a helper - by a function or a script - puts its
    arguments in the parameters' place, as R binds them, until no more is learnt: a name known in the end links the
    object to the caller that gave it and to the helper whose code reads it. A name partly known at run time names
    the objects it could be, as possible links (name_candidates); a name not known at all is a note on its row.
    Returns (possible, notes): {(object, unit): patterns} and {unit: notes}."""
    possible, notes, helpers, resolved = {}, {}, {}, set()
    functions = {u["name"]: u for u in units if u["kind"] == KIND_FUNCTION and not u.get("inside")}
    formals = {name: [f for f, _ in (u.get("code") or {}).get("formals") or []] for name, u in functions.items()}
    defaults = {}
    for name, u in functions.items():
        for node in nodes_of.get(u["ref"], []):
            if node.get("type") == "RParameter":
                parameter = ((node.get("name") or {}).get("lexeme") or "").strip("`")
                defaults.setdefault(name, {}).setdefault(parameter, node.get("defaultValue"))
    def unknown(parts):
        return merged([("any", "") if kind == "param" else (kind, value) for kind, value in parts])
    def settle(user, reader, parts, owner):
        if any(kind == "param" for kind, _ in parts):
            unit = by_ref[user]
            if unit["kind"] == KIND_FUNCTION and not unit.get("inside"):
                entry = (reader, tuple(parts), owner)
                if entry in helpers.setdefault(unit["name"], []):
                    return False
                helpers[unit["name"]].append(entry)
                return True
            parts = unknown(parts)
        if all(kind == "text" for kind, _ in parts):
            written = "".join(value for _, value in parts)
            named = data_files.get(written.rsplit("/", 1)[-1], []) if reader in FILE_READERS else data_names.get(written, [])
            for definer in named:
                for reading in sorted({user, owner}):
                    if definer != reading:
                        link(definer, reading, written)
                if owner != user:
                    resolved.add(owner)
            return False
        found = name_candidates(reader, parts, data_names, data_files)
        if found is None:
            notes.setdefault(user, set()).add(UNKNOWN_NAME)
        for definer in found or ():
            if definer != user:
                possible.setdefault((definer, user), set()).add("".join(v if k == "text" else "*" for k, v in parts))
        return False
    for ref, nodes in sorted(nodes_of.items()):            # every call that reads stored data by a name
        for reader, value in reader_calls(nodes):
            settle(ref, reader, template_of(tree, edges, value), ref)
    calls = {ref: [n for n in nodes if n.get("type") == "RFunctionCall"] for ref, nodes in nodes_of.items()}
    for _ in range(len(functions) + 1):                    # every call of a helper, until no more is learnt
        grown = False
        for ref, nodes in sorted(calls.items()):
            me = by_ref[ref]["name"] if by_ref[ref]["kind"] == KIND_FUNCTION else None
            for node in nodes:
                callee = call_name(node)
                if callee not in helpers or callee == me or \
                        any(bits & CALLS and target in tree for target, bits in edges.get(node_id(node), [])):
                    continue                               # not a helper, itself, or a function defined in place
                bound = bound_arguments(call_arguments(node), formals.get(callee, []))
                for reader, parts, owner in list(helpers[callee]):
                    filled = []
                    for kind, value in parts:
                        if kind != "param":
                            filled.append((kind, value))
                        elif value in bound:
                            filled += template_of(tree, edges, bound[value])
                        else:
                            default = (defaults.get(callee) or {}).get(value)
                            filled += unknown(template_of(tree, edges, default)) if default else [("any", "")]
                    grown = settle(ref, reader, merged(filled), owner) or grown
        if not grown:
            break
    for name, entries in sorted(helpers.items()):          # a helper no caller names an object to: its own reads
        for reader, parts, owner in entries:
            if owner == functions[name]["ref"] and owner not in resolved:
                settle(owner, reader, unknown(parts), owner)
    return possible, notes


def link_chunks(ctx):
    """Step 03, link-chunks: for every unit of Chunks_Model, the units it takes something from and the units that
    take something from it, read by flowR - each function alone, each script whole - and resolved as R resolves a
    name. Enforces: R2, R4, R7"""
    units = ctx.read("model_units")
    folder = flowr_ready()
    package = ((ctx.read("package_info") or [{}])[0] or {}).get("name", "")
    links = unit_links(units, read_functions(units, folder), folder, package)
    unread = [r["ref"] for r in links if r["not_read_by_flowr"]]
    count = sum(len(r["upstream"]) for r in links)
    maybe = sum(len(r["possible_upstream"]) for r in links)
    unknown = [r["ref"] for r in links if r["notes"]]
    data = {u["ref"] for u in units if u["kind"] in (KIND_TABLE, KIND_OBJECT)}
    idle = [r["ref"] for r in links if r["ref"] in data and not r["downstream"] and not r["possible_downstream"]]
    notes = [("flowR could not read %s, so %s no links of %s own." % (", ".join(unread), "it has" if len(unread) == 1 else "they have",
                                                                     "its" if len(unread) == 1 else "their"))] if unread else []
    if maybe:
        notes.append("%d possible links, where code reads stored data by a name put together when it runs: marked "
                     "(possible) beside the stored objects the name could be." % maybe)
    if unknown:
        notes.append("%s read%s stored data by a name known only when it runs." % (", ".join(unknown), "s" if len(unknown) == 1 else ""))
    if idle:
        notes.append("No code of the package reads %d of its %d stored objects: %s." % (len(idle), len(data), ", ".join(idle)))
    return StepResult({"unit_links": links}, {"chunk links": count, "possible links": maybe, "not read by flowR": len(unread),
                                              "stored objects no code reads": len(idle)},
                      ["%d links between the units of Chunks_Model, read by flowR." % count] + notes)


# ================================================================================================
# ---------------------------------------------------------------- a fault inside the tool
class EngineFault(Exception):
    """The tool found itself inconsistent. Never raised about the model under review. Enforces: R2"""


# ---------------------------------------------------------------- the run: its folder, its record and its two deliverables
ENGINE_DIR = os.path.dirname(os.path.abspath(__file__))


# ---------------------------------------------------------------- settings (allow-list)
DEFAULT_SETTINGS = {
    "max_parameter_cells": 5000, "max_parameter_columns": 50, "protect_sheets": True,
    "max_file_mb": 200.0, "reviewer_id": "", "read_pictures": True,
    "parallel_chats": 256,
    "chat_token_limit": 40000, "methodology_batch_tokens": 12000}

def make_settings(overrides=None):
    """The settings of a run. Only names on the allow-list above exist, so a new setting
    can never leak into the manifest by default, and no setting can hold the token.
    Enforces: R8"""
    settings = json.loads(json.dumps(DEFAULT_SETTINGS))
    for name, value in (overrides or {}).items():
        if name not in DEFAULT_SETTINGS:
            raise ValueError("'%s' is not a setting the tool knows" % name)
        settings[name] = value
    return settings

# ---------------------------------------------------------------- paths and project setup
PROJECT_NAME_RE = re.compile(r"^[A-Za-z0-9_-]{1,24}$")
OUTPUT_FILE = "Output.xlsm"               # the deliverable: a workbook with one macro (the workbook's macro)
LEGACY_OUTPUT_FILE = "Output.xlsx"        # what the tool wrote before its workbook held the macro
LONGEST_AUDIT_NAME = OUTPUT_FILE
PATH_BUDGET = 100

@dataclass
class RunPaths:
    """Where a project's run lives: its two files, Output.xlsm and Audit_Log.xlsx, sit in the project folder
    beside its three input folders - run_dir, outputs_dir, audit_dir and inputs_dir are that one folder, each name kept so that
    every reader says which file it means - and local_dir is scratch space on the driver. previous is the
    manifest of the run this one replaced, if any; opened says in plain words which run this is, and why."""
    projects_dir: str; project: str; project_dir: str; inputs_dir: str
    run_id: str; run_dir: str; outputs_dir: str; audit_dir: str; local_dir: str
    previous: Optional[dict] = None; opened: str = ""


class OutputsEdited(Exception):
    """The project's Output.xlsm was changed after the tool wrote it: a person's work, never replaced."""

def check_project_name(project):
    """Return "" when the project name - a model id will do - is usable as a folder's name, otherwise a plain
    sentence saying why not."""
    if PROJECT_NAME_RE.match(project or ""):
        return ""
    return ("The project name may hold at most 24 characters from letters, digits, hyphen and "
            "underscore. Please shorten or change '%s'." % project)

LEGACY_INPUTS = "Inputs"                   # where the three folders were kept before they stood beside Output.xlsm
BESIDE_FOLDERS = ("glossary.xlsx", "tag_rules.yaml")   # the optional files that sit beside the three folders


def holds_nothing(path):
    """Does a folder hold nothing but the tool's README.txt and hidden files?"""
    return not [n for n in os.listdir(path) if n != "README.txt" and not n.startswith(".")]


def lift_inputs(project_dir):
    """A project laid out before - its three folders inside Inputs/ - brought to the present layout, once: each folder,
    and each optional file beside them, moved up into the project's folder whole, not a file of it changed. A folder
    above that holds only its README.txt gives way to the one it replaces; one that holds more is never overwritten:
    then nothing of that folder is moved, and the run cannot start until a person keeps one. Inputs is removed when
    nothing but hidden files is left in it. Returns (what was done, what stops the run), in plain words. Enforces: R6"""
    legacy = os.path.join(project_dir, LEGACY_INPUTS)
    if not os.path.isdir(legacy):
        return [], []
    done, blocking = [], []
    for name in [folder for _, folder, _ in INPUT_FOLDERS] + list(BESIDE_FOLDERS):
        old, new = os.path.join(legacy, name), os.path.join(project_dir, name)
        if not os.path.exists(old):
            continue
        if os.path.isdir(new) and os.path.isdir(old) and holds_nothing(new):
            shutil.rmtree(new)                                 # only the tool's README: the older folder takes its place
        if os.path.exists(new):
            blocking.append("%s is both in %s and beside it, and both hold files. Keep one: move what the one in Inputs holds "
                            "into %s, or delete it, then remove Inputs and run cell 3 again." % (name, legacy, new))
            continue
        try:
            os.rename(old, new)
            done.append(name)
        except OSError as problem:
            blocking.append("%s could not be moved up out of %s (%s). Move it beside Output.xlsm yourself, then run cell 3 "
                            "again." % (name, legacy, problem))
    left = [n for n in os.listdir(legacy) if not n.startswith(".")]
    said = (["Moved up out of Inputs, into the project's folder beside Output.xlsm: %s; no file in them was changed."
             % ", ".join(done)] if done else [])
    if not left and not blocking:
        shutil.rmtree(legacy)
        said.append("Inputs, now empty, was removed.")
    elif left and not blocking:
        said.append("Inputs still holds %s, which the tool does not read; left as it is." % ", ".join(sorted(left)))
    return said, blocking


def setup_project(projects_dir, project):
    """Create the project skeleton and say what is still missing: the three input folders, each with its README.txt,
    in the project's folder beside Output.xlsm and Audit_Log.xlsx. A project laid out before, with its folders inside
    Inputs/, is first brought to this layout (lift_inputs). Returns (project folder, what stops the run, what was
    done). The content of an input is never touched. Enforces: R6"""
    problem = check_project_name(project)
    if problem:
        raise ValueError(problem)
    project_dir = os.path.join(projects_dir, project)          # the project's folder: its three input folders, and a run's two files
    os.makedirs(project_dir, exist_ok=True)
    said, missing = lift_inputs(project_dir)
    for _, folder, readme in INPUT_FOLDERS:
        path = os.path.join(project_dir, folder)
        os.makedirs(path, exist_ok=True)
        readme_path = os.path.join(path, "README.txt")
        if not os.path.exists(readme_path):
            with open(readme_path, "w", encoding="utf-8") as handle:
                handle.write(readme + "\n")
        if holds_nothing(path):
            missing.append("%s is still empty. %s" % (folder, readme))
    return project_dir, missing, said

def new_run_id(taken, now=None):
    """A run is named by the minute it started, <date>_<HHMM>, with a letter added while that name is taken:
    the run it replaces, or a run this driver has opened already."""
    base = (now or datetime.datetime.now()).strftime("%Y-%m-%d_%H%M")
    for suffix in [""] + list("bcdefghijklmnopqrstuvwxyz"):
        if not taken(base + suffix):
            return base + suffix
    raise ValueError("Too many runs were started in the same minute; please wait a minute.")

def recorded_manifest(project_dir):
    """The manifest of the run the project's Audit_Log.xlsx records, read from its Run sheet alone; None when
    there is none, or the workbook cannot be read."""
    target = os.path.join(project_dir, AUDIT_FILE)
    if not os.path.exists(target):
        return None
    import openpyxl
    try:
        book = openpyxl.load_workbook(target, read_only=True)
        parts = [row[1] or "" for row in book["Run"].iter_rows(min_row=2, values_only=True)
                 if re.sub(r" \(part \d+\)$", "", str(row[0] or "")) == "run_manifest"]
        book.close()
        return json.loads("".join(parts)) if parts else None
    except Exception:
        return None

def run_changes(record, inputs_dir, settings):
    """Why the run a project records cannot be carried on, in plain words, or "" when it can: it is carried on
    while its input files, the engine and the settings are the ones it started with. Who runs it does not count."""
    if not record or not record.get("run_id") or "inputs" not in record:
        return "no earlier run"
    before = {f["file"]: f["sha256"] for f in record.get("inputs", [])}
    now = {f["file"]: f["sha256"] for f in input_fingerprints(inputs_dir)}
    if before != now:
        return "the input files changed"
    if record.get("engine_files") != engine_file_hashes():
        return "the engine changed"
    ignore = lambda values: {k: v for k, v in (values or {}).items() if k != "reviewer_id"}
    if ignore(record.get("settings")) != ignore(json.loads(json.dumps(settings))):
        return "the settings changed"
    return ""

def pick_scratch_root(preferred=""):
    """The first folder on the driver the tool can actually write in, tried in order.

    A run is built on the driver's own disk and copied whole into the Workspace afterwards
    (R12), so this folder is needed before anything else can happen. A cluster is shared, and
    a scratch folder made by one user cannot be written into by another; the folders tried
    here therefore carry the user's own name. Writing is tested, not assumed, because a
    folder can exist and still refuse."""
    user = re.sub(r"[^A-Za-z0-9_.-]", "_", getpass.getuser() or "user")
    refused = []
    for root in ([preferred] if preferred else []) + [
            os.path.join(tempfile.gettempdir(), "verifier_scratch_" + user),
            os.path.join("/local_disk0", "verifier_scratch_" + user)]:
        try:
            os.makedirs(root, exist_ok=True)
            probe = os.path.join(root, ".verifier_write_test")
            with open(probe, "w") as handle:
                handle.write("x")
            os.remove(probe)
            return root
        except OSError as problem:
            refused.append("%s (%s)" % (root, problem.strerror or problem))
    raise PermissionError(
        "the tool builds each run on the driver's own disk and copies the finished files into the "
        "Workspace, so it needs one folder it may write in. These were refused: %s. Put a "
        "folder you can write in into the 'Scratch folder' widget, or ask for one on this "
        "cluster." % "; ".join(refused))

def open_run(projects_dir, project, scratch_root="", now=None, settings=None):
    """The project's run, and its local scratch folder. A project holds one run: its two files, Output.xlsm and
    Audit_Log.xlsx, sit beside its three input folders. The run the audit log records is carried on - by this session
    or a later one - while its inputs, the engine and the settings are the ones it started with; otherwise a new
    run starts, replaces both files, and records what changed since the run before. An Output.xlsm a person has
    changed since the tool wrote it is never replaced: OutputsEdited says so. Inputs are never touched.
    Enforces: R6, R12"""
    project_dir, _, _ = setup_project(projects_dir, project)
    inputs_dir = project_dir                                # the three input folders stand in the project's folder
    record = recorded_manifest(project_dir)
    why = run_changes(record, inputs_dir, settings or make_settings({}))
    scratch_root = pick_scratch_root(scratch_root)
    place = sha256_text(os.path.abspath(project_dir))[:8]   # two Projects folders never share scratch space
    local_of = lambda run: os.path.join(scratch_root, "%s_%s_%s" % (project, run, place))
    previous = None
    if not why:
        run_id = record["run_id"]
        opened = "Carrying on run %s, which Audit_Log.xlsx records: a finished step is not repeated." % run_id
    else:
        for name in (OUTPUT_FILE, LEGACY_OUTPUT_FILE):
            workbook = os.path.join(project_dir, name)
            if os.path.exists(workbook) and file_sha256(workbook) != (record or {}).get("last_workbook_sha256"):
                raise OutputsEdited(
                    "%s in %s has been changed since the tool wrote it, and a new run would replace it (%s). "
                    "Move it to another folder or rename it, then run cell 3 again." % (name, project_dir, why))
        legacy = os.path.join(project_dir, LEGACY_OUTPUT_FILE)
        replaced = os.path.exists(legacy)
        if replaced:                                      # the tool's own, unchanged: Output.xlsm takes its place
            os.remove(legacy)
        previous = record if record and record.get("inputs") else None
        run_id = new_run_id(lambda run: run == (record or {}).get("run_id") or os.path.exists(local_of(run)), now)
        opened = ("A new run, %s: %s since run %s, so Output.xlsm and Audit_Log.xlsx are replaced." % (run_id, why, record["run_id"])
                  if record and record.get("run_id") else "A new run, %s." % run_id)
        if replaced:
            opened += (" %s, which the tool wrote before, is replaced by %s: the same sheets, and a click on a reference "
                       "shows only the chunks it names." % (LEGACY_OUTPUT_FILE, OUTPUT_FILE))
    local_dir = local_of(run_id)
    paths = RunPaths(projects_dir, project, project_dir, inputs_dir, run_id, project_dir,
                     project_dir, project_dir, local_dir, previous, opened)
    longest = os.path.join(paths.project_dir, LONGEST_AUDIT_NAME)
    relative = os.path.relpath(longest, os.path.dirname(os.path.abspath(projects_dir)))
    if len(relative) > PATH_BUDGET:
        raise ValueError("The folder path is %d characters long and the limit is %d, so that "
                         "Excel can still open downloaded files. Please use a shorter project name "
                         "or Projects folder." % (len(relative), PATH_BUDGET))
    os.makedirs(paths.local_dir, exist_ok=True)
    if why:                                                 # a new run: the record of the one it replaces is not read back
        OPEN_STORES[(paths.local_dir, paths.audit_dir)] = AuditStore(paths.local_dir, paths.audit_dir, loaded=True)
    return paths

# ---------------------------------------------------------------- live values (the token)
@dataclass
class LiveValues:
    """The three values chat() reads when it is CALLED: endpoint, token, user id. They live
    in memory only. `generation` goes up whenever a different token arrives, which is how
    paused workers learn that a fresh one was pasted. Enforces: R8"""
    values: dict = field(default_factory=dict); generation: int = 0; set_at: float = 0.0
    recent_tokens: list = field(default_factory=list)
    lock: threading.Lock = field(default_factory=threading.Lock)

    def update(self, llm_endpoint="", llm_token="", llm_user_id=""):
        """Take the latest widget values; a different token starts a new generation."""
        with self.lock:
            if llm_token and llm_token != self.values.get("llm_token"):
                self.generation += 1
                self.set_at = time.time()
                self.recent_tokens = ([llm_token] + self.recent_tokens)[:5]
            self.values = {"llm_endpoint": llm_endpoint, "llm_token": llm_token, "llm_user_id": llm_user_id}

    def get(self, name):
        """One of the three values, read now. The user id is one value under two names: the widget is
        called reviewer_id, because the same id records who made a decision and is sent to the LLM."""
        with self.lock:
            return self.values.get("llm_user_id" if name == "reviewer_id" else name, "")

    def token_age_minutes(self):
        """Minutes since the current token was pasted."""
        return (time.time() - self.set_at) / 60.0 if self.set_at else 0.0

    def redact(self, text):
        """Remove the current and recent token strings from any text before it is kept."""
        for token in list(self.recent_tokens):
            if token:
                text = text.replace(token, "[token removed]")
        return text

# ---------------------------------------------------------------- the audit store
AUDIT_OBJECTS = ("run_manifest", "package_info")
AUDIT_FILE = "Audit_Log.xlsx"            # the record of a run: one workbook, beside Output.xlsm and the three input folders
CELL_LIMIT = 30000                       # Excel holds 32,767 characters in a cell; longer text is written in parts

def copy_whole(source, target):
    """Copy one whole file: to a temporary name, then replace; a plain copy if the file
    system does not support replace (probe P-5). Enforces: R12"""
    os.makedirs(os.path.dirname(target), exist_ok=True)
    temporary = target + ".part"
    shutil.copyfile(source, temporary)
    try:
        os.replace(temporary, target)
    except OSError:
        shutil.copyfile(source, target)
        os.remove(temporary)

@dataclass
class AuditStore:
    """The record of a run, as one workbook a person can open: Audit_Log.xlsx, in the project folder.

        Run            the run's manifest and account, one line per entry
        Steps          every step that ran, what it counted and what it said
        Records        every record of every kind, in the order written, each as its own JSON
        Model_Calls    every exchange with the model: the question, the answer and the outcome

    Records are held in memory while a cell runs and the workbook is written whole at each
    sync, beside and then swapped in, so a reader never sees it half written. Text longer than
    a cell holds is written in numbered parts and joined again when read, so nothing is cut.
    Enforces: R4, R5, R12"""
    local_dir: str; remote_dir: str
    rows: list = field(default_factory=list)       # {"kind", "record"}, in the order written
    calls: list = field(default_factory=list)
    account: dict = field(default_factory=dict)    # the manifest and the other single-object kinds
    loaded: bool = False

    def target(self):
        return os.path.join(self.remote_dir, AUDIT_FILE)

    def read(self, kind):
        """Every record of one kind, in the order written."""
        self.restore()
        if kind == "llm_calls":
            return list(self.calls)
        if kind in AUDIT_OBJECTS:
            found = self.account.get(kind)
            return [found] if found is not None else []
        return [row["record"] for row in self.rows if row["kind"] == kind]

    def append(self, kind, records):
        """Add records of one kind; the single-object kinds are replaced whole."""
        self.restore()
        if kind in AUDIT_OBJECTS:
            self.account[kind] = to_plain(records[-1])
            return
        self.rows += [{"kind": kind, "record": to_plain(record)} for record in records]

    def append_calls(self, records):
        self.restore()
        self.calls += [to_plain(record) for record in records]

    def read_calls(self):
        return self.read("llm_calls")

    def sync(self):
        """Write the whole workbook, beside and then swapped in. Returns what was written."""
        import openpyxl
        os.makedirs(self.remote_dir, exist_ok=True)
        book = openpyxl.Workbook()
        book.remove(book.active)
        run = book.create_sheet("Run")
        run.append(["Entry", "Value"])
        for key in sorted(self.account):
            for line, part in parts_of(canonical_json(self.account[key])):
                run.append([key if line == 1 else "%s (part %d)" % (key, line), part])
        steps = book.create_sheet("Steps")
        steps.append(["Step", "Name", "Seconds", "What it counted", "What it said"])
        for record in self.read("step_records"):
            steps.append([record.get("step_id", ""), record.get("step", ""), record.get("seconds", ""), canonical_json(record.get("counts") or {})[:CELL_LIMIT],
                          "\n".join(record.get("messages") or [])[:CELL_LIMIT]])
        records = book.create_sheet("Records")
        records.append(["Kind", "Number", "Part", "Record (JSON)"])
        for number, row in enumerate(self.rows, start=1):
            for part_number, part in parts_of(canonical_json(row["record"])):
                records.append([row["kind"], number, part_number, part])
        calls = book.create_sheet("Model_Calls")
        calls.append(["Number", "Question", "Step", "Question type", "Attempt", "Outcome", "Part", "Exchange (JSON)"])
        for number, record in enumerate(self.calls, start=1):    # a question asked twice has two records: each its own number
            for part_number, part in parts_of(canonical_json(record)):
                calls.append([number, record.get("question_id", ""), record.get("step", ""), record.get("question_type", ""),
                              record.get("attempt", ""), record.get("outcome", ""), part_number, part])
        for sheet in book.worksheets:
            sheet.freeze_panes = "A2"
        partial = os.path.join(self.local_dir, AUDIT_FILE + ".writing")
        os.makedirs(self.local_dir, exist_ok=True)
        book.save(partial)
        copy_whole(partial, self.target())
        os.remove(partial)
        return [AUDIT_FILE]

    def restore(self):
        """On resume: read the project's Audit_Log.xlsx back, once."""
        if self.loaded:
            return
        self.loaded = True
        if not os.path.exists(self.target()):
            return
        import openpyxl
        book = openpyxl.load_workbook(self.target(), read_only=True)
        joined = {}
        for row in book["Records"].iter_rows(min_row=2, values_only=True):
            kind, number, _, part = row[0], row[1], row[2], row[3] or ""
            joined.setdefault((kind, number), []).append(part)
        self.rows = [{"kind": kind, "record": json.loads("".join(parts))}
                     for (kind, number), parts in sorted(joined.items(), key=lambda pair: pair[0][1])]
        whole = {}
        for row in book["Model_Calls"].iter_rows(min_row=2, values_only=True):
            whole.setdefault(row[0], []).append((row[6] or 1, row[7] or ""))
        self.calls = [json.loads("".join(part for _, part in sorted(parts))) for _, parts in sorted(whole.items())]
        entries = {}
        for row in book["Run"].iter_rows(min_row=2, values_only=True):
            key = re.sub(r" \(part \d+\)$", "", str(row[0] or ""))
            entries.setdefault(key, []).append(row[1] or "")
        self.account = {key: json.loads("".join(parts)) for key, parts in entries.items() if key}

def parts_of(text):
    """One long text as numbered parts, each short enough for a cell."""
    return [(number + 1, text[at:at + CELL_LIMIT]) for number, at in enumerate(range(0, max(len(text), 1), CELL_LIMIT))]

OPEN_STORES = {}                         # one store per run: two of them would overwrite each other's records

def open_store(paths, settings):
    """The audit store of a run, read back from the project's Audit_Log.xlsx when there is one and the run is
    carried on. The same store is returned for the same run, so everything a run records goes into one account."""
    key = (paths.local_dir, paths.audit_dir)
    if key not in OPEN_STORES:
        OPEN_STORES[key] = AuditStore(paths.local_dir, paths.audit_dir)
    return OPEN_STORES[key]

# ---------------------------------------------------------------- the organisation's model: what one question may hold
# chat() holds a question and its answer together up to chat_token_limit tokens (40,000 by default). The model's own
# tokenizer is not at hand, so the tool counts tokens its own way, on the high side: measured on prose, R code, tables
# of numbers, mathematics and JSON against six tokenizers (three of OpenAI's, Llama's, Mistral's and Claude's), its count
# was never below theirs. Every question is built to fit what is left once its answer's share is set aside.
CODE_QUESTION = "code interpretation"
METHODOLOGY_SEARCH = "methodology search"
METHODOLOGY_COMPARISON = "methodology comparison"
ANSWER_TOKENS = {CODE_QUESTION: 6000, METHODOLOGY_SEARCH: 4000, METHODOLOGY_COMPARISON: 10000}   # kept for the answer
QUESTION_MARGIN = 1000       # tokens kept for the gateway's own wrapping of a question
TOKEN_PIECES = re.compile(r"[A-Za-z]+|[0-9]|\n|[^\S\n]+|[^A-Za-z0-9\s]")


def token_costs(text):
    """Each piece of a text with the tokens it is counted as: a run of letters one for every six letters; a digit, a
    sign or a line break one; a space one unless it leads into a word; a character outside ASCII one for each of its
    UTF-8 bytes. Yields (where the piece ends, its tokens)."""
    size = len(text)
    for match in TOKEN_PIECES.finditer(text):
        piece, end = match.group(0), match.end()
        first = piece[0]
        if first.isascii() and first.isalpha():
            yield end, -(-len(piece) // 6)
        elif first.isspace() and first != "\n":
            leads = len(piece) == 1 and end < size and text[end].isascii() and text[end].isalpha()
            yield end, 0 if leads else 1
        elif first.isascii():
            yield end, 1
        else:
            yield end, len(first.encode("utf-8"))


def estimate_tokens(text):
    """About how many tokens a text takes a model, erring high (see token_costs)."""
    return sum(cost for _, cost in token_costs(text or ""))


def question_room(settings, question_type):
    """How many tokens a question of this type may take: the limit, less the share kept for its answer and a margin
    for the gateway's own wrapping. A low limit keeps a quarter of itself for the answer at most."""
    limit = int((settings or {}).get("chat_token_limit") or DEFAULT_SETTINGS["chat_token_limit"])
    return limit - min(ANSWER_TOKENS[question_type], limit // 4) - min(QUESTION_MARGIN, limit // 20)


def cut_to_tokens(text, tokens):
    """The longest start of `text` that takes at most `tokens`, ended at a line end when one is near; and whether
    anything was cut."""
    total, end = 0, 0
    for at, cost in token_costs(text):
        if total + cost > tokens:
            break
        total, end = total + cost, at
    else:
        return text, False
    if end == 0:                                         # one piece longer than the room: cut inside it
        end = max(1, tokens)
    kept = text[:end]
    line_end = kept.rfind("\n")
    if line_end > len(kept) * 0.8:
        kept = kept[:line_end]
    return kept, True




def shown_cut(text, tokens):
    """A text cut to `tokens`, saying so at its end when it was cut."""
    if estimate_tokens(text) <= tokens:
        return text
    kept, _ = cut_to_tokens(text, max(1, tokens - 20))
    return kept + "\n[Only the first %d characters are shown.]" % len(kept)


# ---------------------------------------------------------------- the organisation's model: many questions at once
# Steps 04 and 05 put their questions to the chat() of cell 2, as many at once as parallel_chats allows. A question is
# exact - its id is the hash of what was sent - so an answer already recorded in this run is never asked for twice.
# Every answer is kept, in memory, the moment it arrives, by the thread that received it; only the step's own thread
# writes records, in an order fixed by the units and never by which answer came back first. A cell interrupted, or a
# step stopped because the model stopped answering, therefore loses nothing: running cell 3 again finds what arrived
# and asks only for the rest. As many questions are out at once as the gateway takes, up to parallel_chats: the
# number starts at CHAT_START and doubles every round while no call fails; a failed call halves it, once a round,
# and it then grows by one a round, so that it settles just under what the gateway takes - the congestion control
# of TCP. Enforces: R2, R3, R5, R8
CHAT_ATTEMPTS = 3            # tries per question: a failed call, an empty answer or an answer turned down ask again
CHAT_BACKOFF = 2             # seconds before the second try of a failed call; before the third, its square; each with jitter
CHAT_START = 32              # questions out at once when asking starts, doubled every round while no call fails
CHAT_REFRESH_SECONDS = 1     # how often the step's own thread reads the widgets again while it asks
CHAT_WORKER = threading.local()
CHAT_STOP_AFTER = 4          # questions in a row without an answer, beyond those already out: the model has stopped answering
CHAT_STOP_WAIT = 120         # seconds a stopped step waits for the questions still out
CHAT_PROGRESS_SECONDS = 60   # how often a long step says how far it has got
ANSWERS = {}                 # run (its scratch folder) -> {question id: (question, what came back)}, until written
ANSWERS_LOCK = threading.Lock()


class NoAnswer(Exception):
    """What chat() returned held no answer: a dictionary without "answer", or with an empty one, or not a dictionary."""


def unwelcome_words(text):
    """The words of an answer that no cell of the workbook may hold, as they were written. Enforces: R1, R10"""
    found = [match.group(0) for match in PYTHON_TRACES.finditer(text)]
    found += [match.group(0) for match in _BANNED_RE.finditer(text)]
    return sorted(set(found), key=str.lower)


def own_words(text):
    """A text without what it quotes inside curly double quotes: the words it says itself."""
    return re.sub(r"\u201c.*?\u201d", "", text or "", flags=re.S)


def check_words(answer, last):
    """Step 04's check of an answer: words the workbook cannot hold are asked about again, and kept on the last try.
    Returns (what was read, what to ask again or None, a note)."""
    unwelcome = unwelcome_words(answer)
    if not unwelcome:
        return None, None, ""
    if last:
        return None, None, "the answer still used words the workbook cannot hold"
    return None, {"plain": "the answer used words the workbook cannot hold (%s)" % ", ".join(unwelcome),
                  "ask": "Your last description used words this workbook cannot hold (%s). Write it again without "
                         "them." % ", ".join(unwelcome)}, ""


def ask_model(chat, system, main, check=None, halt=None):
    """One question, asked on a worker thread: up to CHAT_ATTEMPTS tries. A failed call - one that raises, or returns
    anything but a dictionary with an answer under "answer": an error from the gateway, an expired token - waits a
    little and tries again; an answer `check` turns down - words the workbook cannot hold, or not the form asked for -
    is asked for again, saying what was wrong. Once `halt` is set - the cell that asked has ended, or was stopped - no
    further try is made: the question stays open for the next time cell 3 runs. Returns what happened, in plain words
    and in technical ones, and what check read from the answer; it never raises and never writes, because only the
    step's own thread writes. Enforces: R5, R8"""
    check = check or check_words
    halt = halt or threading.Event()
    CHAT_WORKER.active = True
    started, plain, technical, question, failed_calls = time.time(), [], [], main, 0
    for attempt in range(1, CHAT_ATTEMPTS + 1):
        if halt.is_set():
            plain.append("try %d: not made, because the cell had ended" % attempt)
            break
        try:
            reply = chat(system, question)
            answer = str(reply.get("answer") or "").strip() if isinstance(reply, dict) else ""
            if not answer:
                raise NoAnswer("chat() returned no \"answer\": %s" % " ".join(repr(reply).split())[:300])
        except Exception as problem:
            failed_calls += 1
            plain.append("try %d: the call did not return an answer" % attempt)
            technical.append("try %d: %s" % (attempt, problem if isinstance(problem, NoAnswer) else "%s: %s" % (type(problem).__name__, problem)))
            if attempt < CHAT_ATTEMPTS:
                halt.wait(CHAT_BACKOFF ** attempt * random.uniform(0.5, 1.5))   # a busy gateway is given a moment; the
                                                 # jitter keeps many refused calls from coming back all at once
            continue
        try:
            reading, problem, note = check(answer, attempt == CHAT_ATTEMPTS)
        except Exception as fault:                       # a check that fails is the tool's fault: said, never raised
            reading, problem, note = None, {"plain": "the answer could not be checked", "ask": ""}, ""
            technical.append("try %d: checking the answer: %s: %s" % (attempt, type(fault).__name__, fault))
        if problem is None:
            if note:
                plain.append("try %d: %s" % (attempt, note))
            return {"answer": answer, "reading": reading, "attempt": attempt, "plain": plain, "technical": technical,
                    "failed calls": failed_calls, "seconds": time.time() - started}
        plain.append("try %d: %s" % (attempt, problem["plain"]))
        technical.append("try %d: the answer turned down began: %s" % (attempt, " ".join(answer.split())[:600]))
        question = main + ("\n\n" + problem["ask"] if problem["ask"] else "")
    return {"answer": "", "reading": None, "attempt": CHAT_ATTEMPTS, "plain": plain, "technical": technical,
            "failed calls": failed_calls, "seconds": time.time() - started}


def held_answers(ctx, types):
    """The answers of this run that arrived but are not yet written, as call records without their provenance: those
    of a cell that was interrupted, for instance. Read, not taken: ask_all takes them when it hands them over."""
    held = ANSWERS.get(ctx.options["paths"].local_dir) or {}
    with ANSWERS_LOCK:
        entries = list(held.values())
    return [dict(question, question_id=question["id"], question_type=question["type"], reading=result["reading"],
                 outcome="answered" if result["answer"] else "not answered")
            for question, result in entries if question["type"] in types]


def ask_all(ctx, chat, work, build, check_of, after=None, label="questions", types=()):
    """Ask what `work` holds - and what `after` adds as answers come back - through chat(), as many at once as
    the gateway takes, up to parallel_chats (see the section's comment). build(item) makes the question, on this thread, when it is sent; check_of(question) is the
    check its answers pass; after(question, result) returns the next items, which go before the rest. Every answer is
    held in ANSWERS by the thread that received it, so an interrupted cell loses none. When the questions already out,
    and CHAT_STOP_AFTER more, all come back without an answer, no more are sent: the model has stopped answering.
    Returns (every held result of these types, taken, as (question, result)), why it stopped or "", and the most
    questions that were out at once. Enforces: R2, R5, R8"""
    held = ANSWERS.setdefault(ctx.options["paths"].local_dir, {})
    most = max(1, int(ctx.settings.get("parallel_chats") or 1))
    waiting, running = collections.deque(work), {}
    window, threshold = float(min(CHAT_START, most)), float(most)      # how many may be out; where doubling gives way to adding
    answered = failed = in_a_row = out_when_failing = completed = calm_after = peak = 0
    stopped, stopped_at, said, refreshed, failures, last_problem = "", 0.0, time.time(), time.time(), [], ""
    halt = threading.Event()                                  # set when this cell stops asking, however it ends

    def keep(question, future):
        if future.cancelled():
            return
        try:
            result = future.result()
        except BaseException:
            return
        with ANSWERS_LOCK:
            held[question["id"]] = (question, result)

    def progress(final=False):
        parts = ["%d answered" % answered] + (["%d not answered" % failed] if failed else [])
        if not final:
            parts.append("%d still to ask, %d out at once" % (len(waiting) + len(running), len(running)))
        else:
            parts.append("at most %d out at once" % peak)
        print("  %s: %s" % (label, ", ".join(parts)))

    if waiting:
        pool = concurrent.futures.ThreadPoolExecutor(max_workers=most)
        live("llm_token")                                    # the values the workers use, read here, now
        try:
            while waiting or running:
                while waiting and len(running) < int(window) and not stopped:
                    question = build(waiting.popleft())
                    kept = {key: value for key, value in question.items() if key not in ("system", "main")}
                    future = pool.submit(ask_model, chat, question["system"], question["main"], check_of(question), halt)
                    future.add_done_callback(functools.partial(keep, kept))
                    running[future] = kept
                peak = max(peak, len(running))
                if not running:
                    break
                done, _ = concurrent.futures.wait(list(running), timeout=5, return_when=concurrent.futures.FIRST_COMPLETED)
                if time.time() - refreshed >= CHAT_REFRESH_SECONDS:
                    live("llm_token")                        # a token pasted meanwhile reaches the next calls
                    refreshed = time.time()
                for future in done:
                    question, result = running.pop(future), future.result()
                    completed += 1
                    if result["failed calls"]:                   # the gateway refused or failed a call: slow down,
                        if completed >= calm_after:              # once a round - those out were sent at the old pace
                            threshold = max(1.0, window / 2)
                            window, calm_after = threshold, completed + len(running)
                    elif result["answer"]:                       # doubling each round up to the threshold, then one a round
                        window = min(float(most), window + (1.0 if window < threshold else 1.0 / window))
                    if result["answer"]:
                        answered, in_a_row = answered + 1, 0
                    else:
                        failed, in_a_row = failed + 1, in_a_row + 1
                        out_when_failing = len(running) + 1 if in_a_row == 1 else out_when_failing
                        failures.append(bool(result["failed calls"]))
                        last_problem = re.sub(r"^try \d+: ", "", result["technical"][-1]) if result["technical"] else last_problem
                    follow = after(question, result) if after else []
                    if not stopped and in_a_row >= out_when_failing + CHAT_STOP_AFTER:
                        stopped, stopped_at = stop_message(failures[-in_a_row:], answered,
                                                           (NOTEBOOK["live"] or LiveValues()).redact(last_problem)), time.time()
                        waiting.clear()
                    if follow and not stopped:
                        waiting.extendleft(reversed(follow))
                if stopped and running and time.time() - stopped_at > CHAT_STOP_WAIT:
                    break                                    # what is still out is held when it arrives
                if time.time() - said >= CHAT_PROGRESS_SECONDS:
                    progress()
                    said = time.time()
        finally:
            halt.set()                                       # no question is tried again once the cell has stopped asking;
            pool.shutdown(wait=False, cancel_futures=True)   # what is still out is held when it arrives
        progress(final=True)
    with ANSWERS_LOCK:
        taken = [held.pop(key) for key in [key for key, (question, _) in held.items() if question["type"] in types]]
    return taken, stopped, peak


def stop_message(failures, answered, last_problem=""):
    """Why a step stopped asking, in plain words: the calls failed - the token may have run out, or the gateway be
    down - or the answers came back in a form that could not be read. `last_problem` is what the last failed call
    returned, the token removed."""
    if all(failures):
        return ("The model stopped answering: the last %d questions got no answer, so no more were sent.%s If the access token "
                "has run out, paste a new one into widget 02; if the gateway is down, wait until it is back. Then run cell 3 "
                "again - or every cell, from cell 1: the %d answers received in this cell are kept, with every answer received "
                "before, and only the questions still open are asked." % (
                    len(failures), " The last call returned: %s." % last_problem.rstrip(".")[:300] if last_problem else "", answered))
    return ("The last %d answers of the model could not be read in the form asked for, so no more questions were sent. "
            "Run cell 3 again to carry on: the %d answers received in this cell are kept; what the model wrote is in "
            "run_log.txt." % (len(failures), answered))


def call_record(ctx, asked, result, redact, **extra):
    """The record of one exchange with the model, as the audit log's Model_Calls sheet holds it: the token removed
    from everything the model wrote, the technical account of what went wrong left to run_log.txt. `extra` may hold
    the question itself, as step 04 keeps it. Enforces: R4, R8, R10"""
    answer = redact(result["answer"])
    record = {"run_id": ctx.provenance.run_id, "step_id": ctx.provenance.step_id, "step": ctx.provenance.step,
              "question_id": asked["id"], "question_type": asked["type"], "unit_ref": asked["unit_ref"]}
    record.update(extra)
    record.update({"attempt": result["attempt"], "outcome": "answered" if answer else "not answered",
                   "tokens": asked.get("tokens", 0), "answer": answer,
                   "reading": redacted(result["reading"], redact),
                   "prompt_hash": asked["id"], "response_hash": sha256_text(answer) if answer else "",
                   "what happened": [redact(line) for line in result["plain"]], "seconds": round(result["seconds"], 3)})
    return record


def redacted(value, redact):
    """Every text inside a value, the token removed. Enforces: R8"""
    if isinstance(value, str):
        return redact(value)
    if isinstance(value, list):
        return [redacted(item, redact) for item in value]
    if isinstance(value, dict):
        return {key: redacted(item, redact) for key, item in value.items()}
    return value


def log_technical(ctx, step_id, lines):
    """Technical text of a step - what a failed call said - goes to run_log.txt only. Enforces: R10"""
    if lines:
        with open(os.path.join(os.path.dirname(ctx.work_dir), "run_log.txt"), "a", encoding="utf-8") as handle:
            handle.write("".join("step %s, %s\n" % (step_id, line) for line in lines))


# ---------------------------------------------------------------- step 04: interpret-code
# Step 04 puts every unit of Chunks_Model in front of the organisation's model, with the units it takes from and gives
# to as context, and asks it to explain the credit concepts the code implements, for a CFA-level analyst checking it
# against the methodology. The model's words go to one column, headed "(by LLM)", and to the audit log; nothing the
# code reads, maps or checks depends on them. Enforces: R3, R5, R8
CODE_SYSTEM_PROMPT = (
    "You help financial analysts review how a credit model is implemented in R. Your reader is a credit analyst at "
    "CFA level who is checking, piece by piece, whether the model's code implements its intended methodology, and "
    "whether it is conceptually sound; they read financial concepts fluently, but not R. You are given one piece of "
    "the model's R package - a function, a statement, a test, stored data or another file - exactly as written, "
    "followed, for context, by the pieces of the package it takes something from and the pieces that take something "
    "from it. Explain the one piece in the language of credit: the financial concept it implements (a probability of "
    "default, a loss given default, an exposure, a correlation, a capital requirement, a rating, a score, and so "
    "on); the inputs it takes and what each means financially; how it computes its result, as a formula an analyst "
    "would recognise; the parameters, floors, caps, thresholds and constants it fixes, and what they stand for; and "
    "what it returns and where that goes in the model. Say which methodological choices and assumptions the code "
    "makes, and where it follows or departs from a standard credit convention (the Basel formulas, for example), "
    "stated as fact, so that the analyst can set them against the methodology. Use the context only to understand "
    "this piece, and describe it only as far as it explains this one. Decide yourself how much detail the analyst "
    "needs: a sentence or two for a simple piece, as many paragraphs as an involved calculation deserves. Describe "
    "only what the code shows; where it relies on something the code does not show, say so, and do not guess. Do "
    "not rate, grade or recommend, and do not say whether the code is right: the analyst decides that. Write prose, "
    "with formulas in plain text where they help; no bullet points, headings or code blocks. Never use the words "
    "finding, error, severity, severe, critical, major or minor, and never call anything high, medium or low in "
    "risk, rating, priority or impact: say riskier or safer, stronger or weaker, near the top or the bottom of the "
    "scale, or give the number. Where the code stops with a message, say that it stops with a message.")
CODE_TOKENS_MAX = 16000      # tokens of one unit's text in its question; a longer text is cut, and the cut is said
CONTEXT_PIECE_TOKENS = 2000  # tokens of one context piece's text
NO_ANSWER = "No interpretation: the model gave no answer. Run cell 3 again to ask again."
NOT_ASKED = "Not asked: nothing was read from this file."


def askable(unit):
    """Does this unit have text worth putting to the model? A file that could not be read has not."""
    return unit["kind"] != KIND_NOT_READ and bool(unit["text"].strip())


def unit_place(unit):
    """Where a unit is: its kind, its file and its lines."""
    return "%s, %s%s" % (unit["kind"], unit["file"], ", lines %d-%d" % tuple(unit["lines"]) if unit.get("lines") else "")


def linked_units(unit, links, by_ref):
    """The units a unit takes something from, each with the names it takes, and the units that take something from it:
    step 03's links, as [(unit, names)] twice."""
    linked = links.get(unit.get("unit_ref", unit["ref"])) or {}
    upstream = [(by_ref[ref], (linked.get("via") or {}).get(ref, [])) for ref in linked.get("upstream") or () if ref in by_ref]
    downstream = [(by_ref[ref], []) for ref in linked.get("downstream") or () if ref in by_ref]
    return upstream, downstream


def context_block(upstream, downstream, room):
    """For context, the units a unit takes something from - each with the names it takes - and the units that take
    something from it, each cut to CONTEXT_PIECE_TOKENS and all of them to `room` tokens. A piece past that is named,
    not shown; once even the names do not fit, how many more there are is said instead. Returns the two blocks."""
    blocks = []
    for title, pieces in (("CONTEXT: THE PIECES IT TAKES SOMETHING FROM", upstream),
                          ("CONTEXT: THE PIECES THAT TAKE SOMETHING FROM IT", downstream)):
        shown, unnamed = [title], 0
        for piece, names in pieces:
            head = "[%s] %s%s" % (piece["ref"], unit_place(piece), " - it takes: %s" % ", ".join(names) if names else "")
            body = shown_cut(piece["text"], CONTEXT_PIECE_TOKENS)
            cost = estimate_tokens(head) + estimate_tokens(body) + 3
            if cost <= room:
                room -= cost
                shown.append(head + "\n" + body)
                continue
            named = head + "\n[Not shown: the context is already long.]"
            if estimate_tokens(named) + 3 <= room - 40:
                room -= estimate_tokens(named) + 3
                shown.append(named)
            else:
                unnamed += 1
        if unnamed:
            shown.append("[%d more are not shown or named: the context is already long.]" % unnamed)
        blocks.append("\n\n".join(shown if len(shown) > 1 else shown + ["(none in the package)"]))
    return blocks


def code_question(unit, upstream=(), downstream=(), room=None):
    """The question about one unit, exactly as sent: (system half, main half, question id). The main half is the
    unit itself, cut to CODE_TOKENS_MAX, then, for context, the units it takes something from - each with the names
    it takes - and the units that take something from it (step 03's links), in whatever room the question has left."""
    room = room or question_room(DEFAULT_SETTINGS, CODE_QUESTION)
    text = shown_cut(unit["text"], min(CODE_TOKENS_MAX, room // 2))
    first = "THE PIECE TO EXPLAIN\nKind: %s\nFile: %s%s\nName: %s%s\n\n%s" % (
        unit["kind"], unit["file"], ", lines %d-%d" % tuple(unit["lines"]) if unit.get("lines") else "", unit.get("name") or "-",
        "\n" + part_note(unit) if part_note(unit) else "", text)
    left = room - estimate_tokens(CODE_SYSTEM_PROMPT) - estimate_tokens(first) - 200
    main = "\n\n\n".join([first] + context_block(upstream, downstream, left))
    return CODE_SYSTEM_PROMPT, main, sha256_text(CODE_SYSTEM_PROMPT + "\n\n" + main)


def interpret_code(ctx):
    """Step 04, interpret-code: every unit of Chunks_Model explained by the organisation's model in the credit
    concepts it implements, for a CFA-level analyst, through the chat() of cell 2 - with the units it takes from and
    the units that take from it (step 03's links) as context, and at the length the model judges it needs. Questions
    go out parallel_chats at a time (ask_all); a question already answered in this run is not asked again, and while
    any is left unanswered the step does not finish: running cell 3 again asks only for those. Enforces: R2, R3, R5, R8"""
    whole = ctx.read("model_units")
    units, by_ref = asked_rows(whole, ctx.settings), {u["ref"]: u for u in whole}
    datasets = [unit["ref"] for unit in whole if is_dataset(unit)]
    links = {r["ref"]: r for r in ctx.read("unit_links")}
    recorded = {call["question_id"] for call in ctx.read("llm_calls")
                if call.get("question_type") == CODE_QUESTION and call.get("outcome") == "answered"}
    arrived = {call["question_id"] for call in held_answers(ctx, (CODE_QUESTION,)) if call["outcome"] == "answered"}
    room = question_room(ctx.settings, CODE_QUESTION)
    questions, wanted = {}, []
    for unit in units:
        if askable(unit):
            upstream, downstream = linked_units(unit, links, by_ref)
            system, main, question_id = code_question(unit, upstream, downstream, room)
            question = questions[question_id] = {
                "id": question_id, "type": CODE_QUESTION, "unit_ref": unit["ref"], "system": system, "main": main,
                "tokens": estimate_tokens(system) + estimate_tokens(main),
                "context": {"upstream": [u["ref"] for u, _ in upstream], "downstream": [u["ref"] for u, _ in downstream]}}
            if question_id not in recorded and question_id not in arrived:
                wanted.append(question)
    counts = {"code blocks": len(units), "interpreted": len(recorded & set(questions)), "asked now": len(wanted),
              "not answered": 0, "not asked": sum(1 for unit in units if not askable(unit))}
    chat = NOTEBOOK["chat"]
    if wanted and chat is None:
        counts["not answered"] = len(wanted)
        return StepResult({}, counts, ["The code interpretations are written by your chat(): run cell 2, then cell 3 "
                                       "again."], finished=False)
    taken, stopped, peak = ask_all(ctx, chat, wanted, lambda question: question, lambda question: None,
                                   label="code interpretations", types=(CODE_QUESTION,))
    redact = (NOTEBOOK["live"] or LiveValues()).redact
    order = {u["ref"]: number for number, u in enumerate(units)}
    records, technical = [], []
    for question, result in sorted(taken, key=lambda pair: (order.get(pair[0]["unit_ref"], 0), pair[0]["id"])):
        if question["id"] in recorded:
            continue                                     # written already, by an earlier cell
        whole = questions.get(question["id"]) or {"system": "", "main": ""}
        records.append(call_record(ctx, question, result, redact, context=question.get("context"),
                                   question={"system": whole["system"], "main": whole["main"]}))
        technical += ["%s, %s" % (question["unit_ref"], redact(line)) for line in result["technical"]]
    log_technical(ctx, "04", technical)
    answered = recorded | {record["question_id"] for record in records if record["outcome"] == "answered"}
    missing = sum(1 for question_id in questions if question_id not in answered)
    counts.update({"interpreted": len(questions) - missing, "not answered": missing, "most at once": peak})
    messages = [stopped] if stopped else []
    sliced = sorted({row["unit_ref"] for row in units if row["parts"] > 1})
    if datasets:
        messages.append("%d stored dataset%s, too large to be a parameter table, %s asked about once, from a view of its "
                        "columns and first %d rows; Chunks_Model shows it whole: %s." % (len(datasets), "" if len(datasets) == 1 else "s",
                                                                                        "is" if len(datasets) == 1 else "are", DATASET_VIEW_ROWS,
                                                                                         ", ".join(datasets[:10]) + (" and more" if len(datasets) > 10 else "")))
    if sliced:
        messages.append("%d pieces of the package too long for one row are shown, and asked about, in parts - rows %s-1, %s-2 and "
                        "so on: %s." % (len(sliced), sliced[0], sliced[0], ", ".join(sliced[:10]) + (" and more" if len(sliced) > 10 else "")))
    if missing and not stopped:
        messages.append("The model gave no interpretation for %d of %d code blocks. Run cell 3 again to ask for those "
                        "again; what went wrong is in run_log.txt." % (missing, len(questions)))
    return StepResult({"llm_calls": records}, counts, messages, finished=not missing)


# ---------------------------------------------------------------- step 05: search-methodology
# Step 05 finds, for every unit of Chunks_Model, the chunks of the methodology that describe, explain or inform it, and
# then asks where the code departs from them. The methodology is the canonical description of the code. Two rounds:
#   search      every unit is put to the model once for each batch of the methodology, until every chunk has been
#               searched for it - the question shows the batch, then the unit's code and its interpretation (step 04);
#   comparison  once a unit's search is complete, the chunks found are put to the model together, with the unit, the
#               units it takes from and gives to, and its interpretation, for every potential deviation.
# Comparing in a round of its own lets the model see every chunk that bears on the unit at once: a floor stated in one
# batch and the formula it applies to in another are compared together, not flagged apart. The model's words go to two
# columns headed "(... by LLM ...)" and to the audit log; code checks that every chunk it names was shown to it.
# Enforces: R2, R3, R5, R8
UNIT_TOKENS_MAX = 12000         # tokens of one unit's code and interpretation in a question of step 05
COMPARE_CONTEXT_TOKENS = 5000   # tokens of context code in a comparison; a context piece past that is named, not shown
NEIGHBOURS_NAMED = 20           # units named on the line of what a unit takes from, and on the line of what it gives to
DEVIATION_KINDS = {"differs": "The code differs", "omits": "Not done in the code",
                   "adds": "Not described in the methodology", "ambiguous": "The methodology can be read more than one way"}
DEVIATION_PARTS = (("methodology", "Methodology"), ("code", "Code"), ("why", "Why it potentially deviates"), ("effect", "Effect"),
                   ("example", "Concrete example"))
NO_EXAMPLE = "The model gave no example."
FLAGGED_DECISIONS = ("True Positive", "False Positive", "True Negative", "False Negative", "For further discussion",
                     "Other Case (see notes)")
NOT_SEARCHED = "Not searched: nothing was read from this file."
NOT_COMPARED = "Not compared: nothing was read from this file."
SEARCH_SYSTEM_PROMPT = (
    "You help credit analysts check whether an R package implements the methodology it was built from. The methodology "
    "is the canonical description of what the code should do. You are given part of the methodology, cut into chunks in "
    "reading order, each headed by its reference in square brackets (such as [C-0012]), the headings it sits under and "
    "its type; then one piece of the model's R package - a function, a statement, a test, stored data or another file - "
    "exactly as written, with an explanation of it that a language model wrote. The rest of the methodology is put to "
    "you in other questions: judge only the chunks shown. Pick out every chunk the analyst needs to read to check this "
    "piece against the methodology: a chunk that describes what the piece computes or does - its formula, its steps, its "
    "rules and conditions; a chunk that explains it - what a quantity means, why a step is taken, what an assumption "
    "stands for; and a chunk that informs it - a definition, a parameter value, a floor, a cap, a threshold, a table of "
    "values, a data source, a segment, a unit or a convention that the piece uses, or that the methodology says it "
    "should use. Judge by meaning, not by shared words: the methodology may name a quantity in words or by a symbol "
    "where the code uses a variable name or an abbreviation - the probability of default, PD and pd_1y can be one "
    "quantity - and a chunk that only shares a word with the piece, or speaks of the same concept at a point of the "
    "model the piece does not touch, does not count. For a test, pick the chunks that state what the test checks; for "
    "stored data, the chunks that define its values or give them. The explanation can be wrong: where it and the code "
    "disagree, the code counts. Go through the chunks from the first to the last; a part of the methodology often holds "
    "only a few chunks that count, and often none. Answer with one JSON object and nothing else, in this form: "
    "{\"relevant\": [{\"ref\": \"C-0012\", \"relation\": \"describes\", \"why\": \"gives the formula this function "
    "computes\"}]}. The relation is describes, explains or informs, and why says in fifteen words at most, without "
    "quotation marks, what the chunk gives the piece. Name only chunks shown, each once. If no chunk counts, answer "
    "{\"relevant\": []}.")
COMPARE_SYSTEM_PROMPT = (
    "You help credit analysts check whether an R package implements the methodology it was built from. The methodology "
    "is the canonical description of what the code should do, and the code may deviate from it. You are given the "
    "chunks of the methodology found to bear on one piece of the model's R package, each headed by its reference in "
    "square brackets (such as [C-0012]), the headings it sits under and its type; then, for context, the pieces of the "
    "package this piece takes something from and the pieces that take something from it; then the piece itself, exactly "
    "as written, with an explanation of it that a language model wrote. Find every potential deviation between what "
    "these chunks say and what this piece of code does, and explain each one so that a CFA-level credit analyst "
    "understands, without reading the code, exactly what differs, why it is a deviation, and what it changes in the "
    "results. Compare everything the chunks say that bears on the piece: formulas and the order of their steps; "
    "constants and parameter values; floors, caps and thresholds, and whether a boundary value is included; the "
    "treatment of missing, zero, negative or out-of-range values; units, scales and conventions - a percentage or a "
    "fraction, basis points, annual or monthly figures, signs; rounding and precision; the inputs used, their sources, "
    "filters, segments and level of aggregation; defaults and fallbacks; what the chunks require that the piece does "
    "not do; and what the piece does to its result that the chunks do not describe. Where a chunk can be read in more "
    "than one way and the code follows one reading, explain both readings. A requirement met by a piece shown as "
    "context is not a deviation of this piece; where it is met in neither, or you cannot tell, list it and say which. "
    "For each deviation give: a title - one pointed sentence naming exactly what differs and where, with the values or "
    "steps on both sides, such as \u201cThe monthly PD is the annual PD divided by 12, where the methodology converts it "
    "by compounding\u201d; the methodology - what the chunks require, quoting them word for word where the wording "
    "matters; the code - what the piece does instead, naming its variables, functions and line numbers and quoting it "
    "where that helps; why - the precise mechanism by which the code's result departs from what the methodology "
    "prescribes, step by step where it takes several; the effect - which inputs or cases are affected, in which "
    "direction the result moves, and by how much where the code shows it; and a concrete example - one case followed "
    "through in clear prose, with numbers where the chunks and the code give them: what the methodology gives, what "
    "the code gives instead, and why the two part. Be specific and direct: name the quantities, give the numbers, say which cases are "
    "affected. Write nothing that would fit any code, such as \u201cthis may affect the results\u201d or \u201cthis "
    "should be reviewed\u201d, and do not repeat yourself. State what the code does and what follows from it as fact; "
    "where the code leaves something undetermined, say exactly what. Do not rate how important a deviation is, "
    "recommend a change, or say which of the two is right: the analyst decides that. Compare only with the chunks "
    "shown, and rest every deviation on the references of the chunks it concerns. Quote inside curly double quotes "
    "\u201clike this\u201d - never straight ones, which would break the JSON. The explanation of the piece can be wrong: "
    "where it and the code disagree, the code counts. Answer with one JSON object and nothing else, in this form: "
    "{\"deviations\": [{\"refs\": [\"C-0012\"], \"kind\": \"differs\", \"title\": \"...\", \"methodology\": \"...\", "
    "\"code\": \"...\", \"why\": \"...\", \"effect\": \"...\", \"example\": \"...\"}]}. The kind is differs (the piece does "
    "what the chunks "
    "describe, differently), omits (the chunks require something the piece does not do), adds (the piece does something "
    "to its result that the chunks do not describe) or ambiguous (the chunks can be read in more than one way, and the "
    "code follows one reading). The depth wanted, shown on another model: {\"refs\": [\"C-0047\"], \"kind\": "
    "\"differs\", \"title\": \"The monthly PD is the annual PD divided by 12, where the methodology converts it by "
    "compounding\", \"methodology\": \"C-0047 derives the monthly PD from the annual one as \u201c1 - (1 - PD)^(1/12)"
    "\u201d.\", \"code\": \"monthly_pd() returns \u201cpd_annual / 12\u201d (line 4).\", \"why\": \"Dividing by 12 "
    "spreads defaults evenly over the year; the methodology compounds survival month by month, which gives a higher "
    "monthly PD for the same annual PD.\", \"effect\": \"Every monthly PD is lower than the methodology's, and more so "
    "as the annual PD rises: at an annual PD of 20% the code gives 1.667% a month, the methodology 1.842%.\", "
    "\"example\": \"Take a borrower with an annual PD of 20%. The methodology gives a monthly PD of 1 - (1 - 0.20)^(1/12) = "
    "1.842%; the code gives 0.20 / 12 = 1.667%. Compounded over the twelve months, the code\u2019s figure is an annual PD of "
    "18.3%, not 20%: the borrower\u2019s chance of default comes out 1.7 points lower than the methodology intends.\"}. If "
    "the "
    "piece does what the chunks say, answer {\"deviations\": []}. Outside quotations, never use the words finding, "
    "error, severity, severe, critical, major or minor, and never call anything high, medium or low in risk, rating, "
    "priority or impact.")


def search_shares(settings):
    """(tokens of the methodology, tokens of the unit) in one search question: methodology_batch_tokens, or less
    when the question would not otherwise leave the unit its room."""
    room = question_room(settings, METHODOLOGY_SEARCH) - estimate_tokens(SEARCH_SYSTEM_PROMPT) - 300
    unit = min(UNIT_TOKENS_MAX, room // 2)
    wanted = int(settings.get("methodology_batch_tokens") or DEFAULT_SETTINGS["methodology_batch_tokens"])
    return max(200, min(wanted, room - unit)), unit


def compare_shares(settings):
    """(tokens of the methodology's chunks, of the unit, of the context) in one comparison question."""
    room = question_room(settings, METHODOLOGY_COMPARISON) - estimate_tokens(COMPARE_SYSTEM_PROMPT) - 300
    unit, context = min(UNIT_TOKENS_MAX, room // 2), min(COMPARE_CONTEXT_TOKENS, room // 6)
    return max(200, room - unit - context), unit, context


def piece_cap(settings):
    """The most tokens one piece of the methodology may take in a question of step 05: small enough for a batch of
    its own, and for a comparison."""
    batch, _ = search_shares(settings)
    chunk_room, _, _ = compare_shares(settings)
    return min(batch, chunk_room)


def piece_slices(text, head_tokens, tokens, chars):
    """How a chunk's text is cut into rows: whole when it fits `chars` characters and, beside a head line of
    `head_tokens`, `tokens` tokens; otherwise text_slices, with room for a part's head line."""
    if len(text) <= chars and head_tokens + estimate_tokens(text) + 1 <= tokens:
        return [(text, False)]
    return text_slices(text, max(100, tokens - head_tokens - 30), chars)


def chunk_where(chunk):
    """The headings a chunk of the methodology sits under, and its type, as a question shows them."""
    return "%s | %s" % (" > ".join(chunk.get("heading_chain") or ()) or "(no heading)", chunk["kind"])


def methodology_slices(chunk, cap):
    """The rows a chunk of the methodology takes - on Chunks_Methodology and in the questions of step 05 alike: whole,
    or cut at line ends into parts small enough for a cell of Excel and for a batch of their own. Enforces: R2, R13"""
    return piece_slices(chunk["text"] or "", estimate_tokens("[%s] %s" % (chunk["ref"], chunk_where(chunk))), cap, SLICE_CHARS)


def methodology_pieces(chunks, cap):
    """Every chunk of the methodology as the model is shown it: its ref, the headings it sits under, its type and its
    text - whole, or, when too long, in the parts it takes as rows of Chunks_Methodology, each headed as the part it is
    ([C-0045-2, part 2 of 3 of C-0045]), so that nothing of a long chunk goes unsearched and a part is the same row in
    the sheet and in a question. The parts stay one chunk: what the model finds in any of them counts for the whole.
    Enforces: R2"""
    pieces = []
    for chunk in chunks:
        slices = methodology_slices(chunk, cap)
        for number, (part, _) in enumerate(slices, start=1):
            head = ("[%s] %s" % (chunk["ref"], chunk_where(chunk)) if len(slices) == 1 else
                    "[%s-%d, part %d of %d of %s] %s" % (chunk["ref"], number, number, len(slices), chunk["ref"], chunk_where(chunk)))
            pieces.append({"ref": chunk["ref"], "part": number, "parts": len(slices), "text": head + "\n" + part})
    for piece in pieces:
        piece["tokens"] = estimate_tokens(piece["text"])
    return pieces


def methodology_batches(pieces, cap):
    """The pieces in reading order, packed into batches of at most cap tokens: every piece in exactly one batch, and
    the batches the same for every unit, so a batch comes to the model as the same words each time."""
    batches, current, used = [], [], 0
    for piece in pieces:
        if current and used + piece["tokens"] + 2 > cap:
            batches.append(current)
            current, used = [], 0
        current.append(piece)
        used += piece["tokens"] + 2
    return batches + ([current] if current else [])


def methodology_plan(chunks, settings):
    """The methodology as step 05 puts it to the model: (its pieces, the batches they go in)."""
    batch, _ = search_shares(settings)
    pieces = methodology_pieces(chunks, piece_cap(settings))
    return pieces, methodology_batches(pieces, batch)


def unit_label(unit):
    """A unit named on one line: its ref and its name, or its kind when it has none."""
    return "%s %s" % (unit["ref"], unit["name"]) if unit.get("name") else "%s (%s)" % (unit["ref"], unit["kind"])


def neighbours_line(upstream, downstream):
    """The units a unit takes something from and gives something to, named: what orients a search."""
    lines = []
    for title, pieces in (("It takes something from", upstream), ("It gives something to", downstream)):
        if pieces:
            named = "; ".join(unit_label(unit) for unit, _ in pieces[:NEIGHBOURS_NAMED])
            more = " and %d more" % (len(pieces) - NEIGHBOURS_NAMED) if len(pieces) > NEIGHBOURS_NAMED else ""
            lines.append("%s: %s%s" % (title, named, more))
    return "\n".join(lines)


SLICE_CHARS = 30000          # the most characters one row of Chunks_Model shows; a cell of Excel holds 32,767


def slice_limits(settings):
    """(tokens, characters) one row of Chunks_Model may hold: few enough characters for a cell of Excel, and few
    enough tokens for every question of steps 04 and 05 to show the row whole, beside what the model wrote about
    it, which takes at most a third of the unit's room (unit_part)."""
    _, search_unit = search_shares(settings)
    _, compare_unit, _ = compare_shares(settings)
    return max(500, min(search_unit, compare_unit) * 2 // 3 - 600), SLICE_CHARS


def text_slices(text, tokens, chars):
    """A text in consecutive slices of at most `tokens` and `chars` each, cut at line ends; a line too long for a slice
    of its own is cut inside. Returns [(slice, joined)]: joined where a slice ends inside a line, which the next one
    continues. Joined back - a line break after each slice that is not joined - the slices are the text exactly."""
    slices, part, size, cost = [], [], 0, 0
    for line in text.split("\n"):
        line_cost, line_size = estimate_tokens(line) + 1, len(line) + 1
        if part and (cost + line_cost > tokens or size + line_size > chars):
            slices.append(("\n".join(part), False))
            part, size, cost = [], 0, 0
        while line_cost > tokens or line_size > chars:
            head = cut_to_tokens(line, tokens - 1)[0][:chars - 1] or line[:1]
            slices.append((head, True))
            line = line[len(head):]
            line_cost, line_size = estimate_tokens(line) + 1, len(line) + 1
        part.append(line)
        size, cost = size + line_size, cost + line_cost
    slices.append(("\n".join(part), False))
    return slices


def joined_rows(rows):
    """The text of a unit, rebuilt from its rows in order."""
    return "".join(row["text"] + ("" if row.get("joined") or number == len(rows) - 1 else "\n")
                   for number, row in enumerate(rows))


def model_rows(units, settings):
    """The rows of Chunks_Model. A unit is one row, or - when its text is too long for a cell of Excel, or for the
    questions of steps 04 and 05 to show whole - several: M-0003-1, M-0003-2 and so on, cut at line ends, each with
    its own lines. They stay one analytical chunk: they are sliced so that nothing of the piece is cut, in the
    workbook or in a question, and every character of it is in exactly one of its rows. Enforces: R2, R13"""
    tokens, chars = slice_limits(settings)
    rows = []
    for unit in units:
        text = unit["text"] or ""
        if len(text) <= chars and estimate_tokens(text) <= tokens:
            rows.append(dict(unit, unit_ref=unit["ref"], part=1, parts=1, joined=False, unit_lines=unit.get("lines")))
            continue
        slices, first = text_slices(text, tokens, chars), (unit.get("lines") or [None])[0]
        line, mine = first, []
        for number, (piece, joined) in enumerate(slices, start=1):
            lines = [line, line + piece.count("\n")] if first is not None else None
            mine.append(dict(unit, ref="%s-%d" % (unit["ref"], number), unit_ref=unit["ref"], part=number, parts=len(slices),
                             text=piece, lines=lines, joined=joined, unit_lines=unit.get("lines")))
            line = lines[1] + (0 if joined else 1) if lines else None
        if joined_rows(mine) != text:
            raise EngineFault("The rows of %s do not rebuild its text. This is a defect in the tool, not in the model under "
                              "review." % unit["ref"])
        rows += mine
    return rows


DATASET_VIEW_ROWS = 50       # the rows of a stored dataset a question shows; the workbook shows every row


def is_dataset(unit):
    """A stored table too large to be a parameter table - more cells than max_parameter_cells, or more columns than
    max_parameter_columns: a dataset."""
    data = unit.get("data") or {}
    return data.get("assessable") is False and bool(data.get("dims"))


def dataset_view(unit, tokens):
    """What a question shows of a stored dataset: what it is, its size, and its columns and first rows - as many of the
    first DATASET_VIEW_ROWS as fit in `tokens` - saying that this is a view, and that the workbook shows it whole."""
    count, width = (list(unit["data"]["dims"]) + [0, 0])[:2]
    head = ("A stored dataset of %d rows and %d columns - too large to be a parameter table, so it is shown here as a "
            "view: its columns and first rows. Chunks_Model shows it whole." % (count, width))
    named, *lines = unit["text"].split("\n")                 # the first line names the object; then the table
    head += "\n" + named
    shown = []
    for line in lines[:DATASET_VIEW_ROWS + 1]:
        if estimate_tokens(head + "\n" + "\n".join(shown + [line])) > tokens:
            break
        shown.append(line)
    rest = len(lines) - len(shown)
    return head + "\n" + "\n".join(shown) + ("\n[%d more rows are not in this view.]" % rest if rest > 0 else "")


def asked_rows(units, settings):
    """The rows steps 04 and 05 ask about: every part of every unit, as model_rows cuts it - so that no piece of code
    is cut in a question - except a stored dataset, which is asked about once, from dataset_view, the question saying
    that it is a view: the model is told what the data is, not asked to read every row. What it says stands for the
    whole unit (rows_model_units). Enforces: R3, R13"""
    tokens, _ = slice_limits(settings)
    rows = []
    for unit in units:
        if is_dataset(unit):
            rows.append(dict(unit, unit_ref=unit["ref"], part=1, parts=1, joined=False, unit_lines=unit.get("lines"),
                             text=dataset_view(unit, tokens)))
        else:
            rows += model_rows([unit], settings)
    return rows


def part_note(row, comparing=False):
    """What a question says of a row that is one part of a long unit, or "" for a whole one."""
    if row.get("parts", 1) < 2:
        return ""
    whole = " (lines %d-%d)" % tuple(row["unit_lines"]) if row.get("unit_lines") else ""
    note = ("Part %d of %d of %s%s, which is too long to show whole: each part is asked about on its own. Where this "
            "part relies on something set in another part, say so." % (row["part"], row["parts"], row["unit_ref"], whole))
    return note + (" A requirement this part does not meet may be met in another part: say so, rather than list it "
                   "as not done." if comparing else "")


def unit_part(unit, said, neighbours, room, comparing=False):
    """One unit as a question of step 05 shows it: what and where it is, what it takes from and gives to, its text as
    written, and what step 04's model wrote about it - cut to `room` tokens, the code keeping two thirds of it or more."""
    head = "THE PIECE: %s, %s%s" % (unit["ref"], unit_place(unit), " - %s" % unit["name"] if unit.get("name") else "")
    head += "\n" + part_note(unit, comparing) if part_note(unit) else ""
    head += "\n" + neighbours if neighbours else ""
    said, code = said or "(none)", unit["text"]
    left = room - estimate_tokens(head) - 40
    code_tokens, said_tokens = estimate_tokens(code), estimate_tokens(said)
    if code_tokens + said_tokens > left:
        said_share = min(said_tokens, max(left // 3, left - code_tokens))
        code, said = shown_cut(code, left - said_share), shown_cut(said, said_share)
    return head + "\n\n" + code + "\n\n\nWHAT A LANGUAGE MODEL WROTE ABOUT THIS PIECE\n" + said


def question_of(kind, unit, pieces, system, parts, **extra):
    """A question of step 05 as sent - its two halves, its id, the pieces of the methodology it shows - and the
    tokens it takes, counted part by part, each join counted too. A question over its room is a fault of the tool."""
    main = "\n\n\n".join(text for text, _ in parts)
    tokens = estimate_tokens(system) + sum(count for _, count in parts) + 3 * len(parts)
    question = {"type": kind, "unit_ref": unit["ref"], "pieces": [[p["ref"], p["part"], p["parts"]] for p in pieces],
                "system": system, "main": main, "id": sha256_text(system + "\n\n" + main), "tokens": tokens}
    question.update(extra)
    return question


def search_question(unit, said, neighbours, batch, room):
    """The search question about one unit and one batch of the methodology: the batch first, then the unit."""
    title = "THE METHODOLOGY: A PART OF IT, IN READING ORDER\n\n"
    chunks = title + "\n\n".join(piece["text"] for piece in batch)
    piece = unit_part(unit, said, neighbours, room)
    ask = ("Which chunks of the methodology above does the analyst need to read to check this piece? Answer with the "
           "JSON object only.")
    return question_of(METHODOLOGY_SEARCH, unit, batch, SEARCH_SYSTEM_PROMPT, [
        (chunks, estimate_tokens(title) + sum(p["tokens"] + 2 for p in batch)), (piece, estimate_tokens(piece)),
        (ask, estimate_tokens(ask))])


def compare_question(unit, said, upstream, downstream, chosen, room, context_room):
    """The comparison question about one unit: the chunks found to bear on it, the units it takes from and gives to,
    then the unit and what step 04's model wrote about it."""
    title = "THE METHODOLOGY: THE CHUNKS FOUND TO BEAR ON THIS PIECE\n\n"
    chunks = title + "\n\n".join(piece["text"] for piece in chosen)
    blocks = context_block(upstream, downstream, context_room)
    piece = unit_part(unit, said, "", room, comparing=True)
    ask = ("List every potential deviation between the chunks of the methodology above and this piece. Answer with the "
           "JSON object only.")
    return question_of(METHODOLOGY_COMPARISON, unit, chosen, COMPARE_SYSTEM_PROMPT,
                       [(chunks, estimate_tokens(title) + sum(p["tokens"] + 2 for p in chosen))] +
                       [(block, estimate_tokens(block)) for block in blocks] +
                       [(piece, estimate_tokens(piece)), (ask, estimate_tokens(ask))],
                       context={"upstream": [u["ref"] for u, _ in upstream], "downstream": [u["ref"] for u, _ in downstream]})


UNREADABLE = {"plain": "the answer was not the JSON object asked for",
              "ask": "Your last answer could not be read as the JSON object the instructions ask for. Answer again with "
                     "that JSON object alone, with no other words around it."}


def json_object(answer):
    """The one JSON object an answer holds, whatever fences or words stand around it; None when it holds none. Parsed,
    never evaluated. Enforces: R7"""
    decoder, start = json.JSONDecoder(), answer.find("{")
    for _ in range(20):                                  # the first brace may open words, not the object
        if start < 0:
            return None
        try:
            found, _ = decoder.raw_decode(answer, start)
            if isinstance(found, dict):
                return found
        except ValueError:
            pass
        start = answer.find("{", start + 1)
    return None


def chunk_ref(value):
    """A chunk's reference as the model wrote it, made regular: [C-0012], c-12 and C-0012 are all C-0012."""
    found = re.search(r"\bC\s*-?\s*(\d{1,6})\b", str(value or ""), re.I)
    return "C-%04d" % int(found.group(1)) if found else str(value or "").strip()


def plain_text(value):
    """A text field of an answer, on one line."""
    return " ".join(str(value or "").split()) if isinstance(value, (str, int, float)) else ""


def search_check(question):
    """The check a search answer passes: the JSON object asked for, naming only chunks the question showed. On the last
    try, chunks it named that were not shown are left out, and said so. Enforces: R3"""
    shown = [ref for ref, _, _ in question["pieces"]]

    def check(answer, last):
        found = json_object(answer)
        items = found.get("relevant") if found else None
        if not isinstance(items, list):
            return None, UNREADABLE, ""
        relevant, unknown = {}, []
        for item in items:
            ref = chunk_ref(item.get("ref") if isinstance(item, dict) else item)
            if ref not in shown:
                unknown.append(ref or "(empty)")
            elif ref not in relevant:
                detail = item if isinstance(item, dict) else {}
                relevant[ref] = {"ref": ref, "relation": plain_text(detail.get("relation")).lower(), "why": plain_text(detail.get("why"))}
        if unknown and not last:
            return None, {"plain": "the answer named chunks that were not shown (%s)" % ", ".join(unknown),
                          "ask": "Your last answer named chunks that are not in this part of the methodology (%s). Name only "
                                 "chunks shown above, and answer with the JSON object alone." % ", ".join(unknown)}, ""
        reading = {"relevant": [relevant[ref] for ref in dict.fromkeys(shown) if ref in relevant]}
        return reading, None, "the chunks it named that were not shown were left out (%s)" % ", ".join(unknown) if unknown else ""
    return check


def compare_check(question):
    """The check a comparison answer passes: the JSON object asked for, every deviation resting on chunks the question
    showed, and no words outside quotations that the workbook cannot hold. On the last try, chunks it named that were
    not shown are left out, and said so. Enforces: R1, R3, R10"""
    order = {ref: number for number, (ref, _, _) in enumerate(question["pieces"])}

    def check(answer, last):
        found = json_object(answer)
        items = found.get("deviations") if found else None
        if not isinstance(items, list):
            return None, UNREADABLE, ""
        deviations, unknown = [], []
        for item in items:
            if not isinstance(item, dict):
                continue
            named = item.get("refs", item.get("ref", []))
            named = [chunk_ref(ref) for ref in (named if isinstance(named, list) else [named])]
            unknown += [ref or "(empty)" for ref in named if ref not in order]
            kind = plain_text(item.get("kind")).lower()
            entry = {"refs": sorted(set(ref for ref in named if ref in order), key=order.get),
                     "kind": kind if kind in DEVIATION_KINDS else "", "title": plain_text(item.get("title"))}
            entry.update({part: plain_text(item.get(part)) for part, _ in DEVIATION_PARTS})
            if any(entry[part] for part in ("title",) + tuple(part for part, _ in DEVIATION_PARTS)) and entry not in deviations:
                deviations.append(entry)
        if unknown and not last:
            return None, {"plain": "the answer named chunks that were not shown (%s)" % ", ".join(unknown),
                          "ask": "Your last answer named chunks that are not among those shown (%s). Rest every deviation "
                                 "on chunks shown above, and answer with the JSON object alone." % ", ".join(unknown)}, ""
        unwelcome = unwelcome_words(own_words(deviation_lines(deviations)))
        if unwelcome and not last:
            return None, {"plain": "the answer used words the workbook cannot hold (%s)" % ", ".join(unwelcome),
                          "ask": "Your last answer used words this workbook cannot hold outside quotations (%s). Write it "
                                 "again without them; inside curly quotes \u201c \u201d they may stay." % ", ".join(unwelcome)}, ""
        notes = (["the chunks it named that were not shown were left out (%s)" % ", ".join(unknown)] if unknown else []) + \
                (["the answer still used words the workbook cannot hold"] if unwelcome else [])
        return {"deviations": deviations}, None, "; ".join(notes)
    return check


def sentence(text):
    """A field of an answer as one sentence of a cell: closed with a full stop, and never opening with a bare None."""
    text = "none" + text[4:] if text.startswith("None ") else text
    return text if not text or text[-1] in ".!?" or text[-2:] in (".\u201d", "!\u201d", "?\u201d", ".)") else text + "."


DEVIATION_WITHHELD = "The model's words for this one cannot be shown in plain words; they are in the audit log (Model_Calls)."


def deviation_lines(deviations, gated=False):
    """Deviations as the workbook shows them: one block each, numbered, blocks apart by a blank line. The first line
    names the chunks of the methodology the deviation rests on, its kind and its title; then what the methodology
    requires, what the code does, why that is a deviation and what it changes, each on a line of its own. Gated, a
    deviation whose own words - outside quotations - the workbook cannot hold is replaced by a notice that keeps its
    chunks, so that one such answer never hides the others of its cell. Enforces: R1, R2, R10"""
    blocks = []
    for number, item in enumerate(deviations, start=1):
        refs = ", ".join(item["refs"]) or "No chunk named"
        head = " ".join(s for s in (DEVIATION_KINDS.get(item.get("kind"), "") + "." if item.get("kind") in DEVIATION_KINDS else "",
                                    sentence(item.get("title", ""))) if s)
        lines = ["%d. %s - %s" % (number, refs, head or "A deviation.")]
        lines += ["%s: %s" % (label, sentence(item[part])) for part, label in DEVIATION_PARTS if item.get(part)]
        block = "\n".join(lines)
        blocks.append("%d. %s - %s" % (number, refs, DEVIATION_WITHHELD) if gated and unwelcome_words(own_words(block)) else block)
    return "\n\n".join(blocks)


def methodology_account(units, chunks, calls, settings):
    """How far step 05 has got with each analytical chunk of Chunks_Model - a unit, in one row or several - worked out
    from the recorded answers alone, and those arrived but not yet written, so that the step and the workbook always
    agree. A unit is one chunk however many rows it takes: each row is searched against every piece of the methodology;
    a chunk of the methodology any row finds counts for the whole unit; once every row is searched, each row is
    compared with every chunk found for the unit, in all its parts; the deviations named for any row are the unit's, in
    the order of its rows. A question answered twice counts once. Enforces: R2, R3"""
    pieces, batches = methodology_plan(chunks, settings)
    position = {(piece["ref"], piece["part"]): number for number, piece in enumerate(pieces)}
    parts_of = {}
    for piece in pieces:
        parts_of.setdefault(piece["ref"], []).append((piece["ref"], piece["part"]))
    every, batch_keys = [(p["ref"], p["part"]) for p in pieces], [{(p["ref"], p["part"]) for p in batch} for batch in batches]
    chunk_of, order = {row["ref"]: row["unit_ref"] for row in units}, {row["ref"]: number for number, row in enumerate(units)}
    searched, relevant, compared, named, seen = {}, {}, {}, {}, set()
    for call in calls:
        kind = call.get("question_type")
        if kind not in (METHODOLOGY_SEARCH, METHODOLOGY_COMPARISON) or call.get("outcome") != "answered" \
                or call["question_id"] in seen or call.get("unit_ref") not in chunk_of:
            continue
        seen.add(call["question_id"])
        row, keys = call["unit_ref"], [(ref, part) for ref, part, _ in call.get("pieces") or ()]
        chunk, reading = chunk_of[row], call.get("reading") or {}
        if kind == METHODOLOGY_SEARCH:
            searched.setdefault(row, set()).update(keys)
            shown = {ref for ref, _ in keys}
            for item in reading.get("relevant") or ():
                if item["ref"] in shown:
                    relevant.setdefault(chunk, {}).setdefault(item["ref"], item)
        else:
            compared.setdefault(chunk, set()).update((row, key) for key in keys)
            first = min([position.get(key, len(pieces)) for key in keys] or [len(pieces)])
            named.setdefault(chunk, []).append((order[row], first, call["question_id"], reading.get("deviations") or []))
    by_chunk, account = {}, {}
    for row in units:
        by_chunk.setdefault(row["unit_ref"], []).append(row)
    for chunk, rows in by_chunk.items():
        asked = [row for row in rows if askable(row)]
        unsearched = {row["ref"]: [key for key in every if key not in searched.get(row["ref"], ())] for row in asked}
        missing = {key for keys in unsearched.values() for key in keys}
        found = relevant.get(chunk, {})
        refs = sorted(found, key=lambda ref: position.get(parts_of.get(ref, [(ref, 1)])[0], len(pieces)))
        done = compared.get(chunk, set())
        deviations = []
        for _, _, _, items in sorted(named.get(chunk, []), key=lambda entry: entry[:3]):
            deviations += [item for item in items if item not in deviations]
        account[chunk] = {"askable": bool(asked), "batches": len(batches), "chunks": len(parts_of),
                          "open batches": [number for number, keys in enumerate(batch_keys) if not keys.isdisjoint(missing)],
                          "unsearched": sorted({ref for ref, _ in missing}), "unsearched keys": unsearched,
                          "relevant": [found[ref] for ref in refs],
                          "to compare": [(row["ref"], key) for row in asked for ref in refs for key in parts_of.get(ref, [])
                                         if (row["ref"], key) not in done],
                          "deviations": deviations}
    return {"pieces": pieces, "batches": batches, "units": account}


def methodology_cells(state):
    """A unit's two cells of step 05 on Chunks_Model: the chunks of the methodology found to bear on it, as refs
    joined with "; "; and how many items it has on Flagged_Items - 0 when nothing was flagged or nothing found to
    compare, None before the methodology is searched in full, and, while comparisons are open, the number so far and
    which chunks are still to compare. What is not finished says so, and says what to do. Enforces: R2, R10"""
    refs, _ = methodology_texts(state)
    if not state["askable"] or state["open batches"]:
        return refs, None
    found = len(state["deviations"])
    if state["relevant"] and state["to compare"]:
        open_refs = "; ".join(dict.fromkeys(key[0] for _, key in state["to compare"]))
        return refs, "%d so far - not compared in full with %s; run cell 3 again" % (found, open_refs)
    return refs, found


def methodology_texts(state):
    """The chunks found, as methodology_cells shows them, and what the items flagged say in words (the words check)."""
    if not state["askable"]:
        return NOT_SEARCHED, NOT_COMPARED
    refs = "; ".join(dict.fromkeys(item["ref"] for item in state["relevant"]))
    if state["open batches"]:
        left = state["unsearched"]
        return ("Not searched in full: %d of %d chunks of the methodology searched so far%s.%s Run cell 3 again for the rest."
                % (state["chunks"] - len(left), state["chunks"], "; still to search: %s" % ", ".join(left) if len(left) <= 5 else "",
                   " Found so far: %s." % refs if refs else ""),
                "Not compared yet: the methodology has not been searched in full.")
    if not refs:
        return "None found", "Nothing to compare: no chunk of the methodology was found to bear on this piece."
    lines = deviation_lines(state["deviations"], gated=True)
    if state["to compare"]:
        open_refs = "; ".join(dict.fromkeys(key[0] for _, key in state["to compare"]))
        return refs, (lines + "\n\n" if lines else "") + ("Not compared in full: %s not yet compared with this piece. "
                                                         "Run cell 3 again." % open_refs)
    return refs, lines or "None flagged against %s." % refs


def search_methodology(ctx):
    """Step 05, search-methodology: for every unit of Chunks_Model, the chunks of the methodology that describe, explain
    or inform it - searched batch by batch until every chunk has been searched for it - and then the potential
    deviations of the code from those chunks, all through the chat() of cell 2, parallel_chats questions at a time
    (ask_all). A unit's comparison is asked as soon as its search is complete. A question the model does not answer
    is asked again split in two, when it shows more than one piece of the methodology; a batch taken up again asks
    only for its pieces not yet searched. A question already answered in this run is not
    asked again, and while any unit is not searched and compared in full the step does not finish:
    running cell 3 again asks only for what is open. Enforces: R2, R3, R5, R8"""
    whole, chunks, calls = ctx.read("model_units"), ctx.read("chunks_canon"), ctx.read("llm_calls")
    units = asked_rows(whole, ctx.settings)
    types = (METHODOLOGY_SEARCH, METHODOLOGY_COMPARISON)
    recorded = {call["question_id"] for call in calls if call.get("question_type") in types and call.get("outcome") == "answered"}
    start = methodology_account(units, chunks, calls + held_answers(ctx, types), ctx.settings)
    pieces, batches, states = start["pieces"], start["batches"], start["units"]
    position = {(piece["ref"], piece["part"]): number for number, piece in enumerate(pieces)}
    by_key = {(piece["ref"], piece["part"]): piece for piece in pieces}
    by_ref = {unit["ref"]: unit for unit in units}
    wholes = {unit["ref"]: unit for unit in whole}
    links = {r["ref"]: r for r in ctx.read("unit_links")}
    said = {call["unit_ref"]: call["answer"] for call in calls
            if call.get("question_type") == CODE_QUESTION and call.get("outcome") == "answered"}
    _, unit_room = search_shares(ctx.settings)
    chunk_room, compare_room, context_room = compare_shares(ctx.settings)
    rows_of, parts_of = {}, {}
    for row in units:
        rows_of.setdefault(row["unit_ref"], []).append(row)
    for piece in pieces:
        parts_of.setdefault(piece["ref"], []).append((piece["ref"], piece["part"]))
    to_search = {row: set(keys) for state in states.values() for row, keys in state["unsearched keys"].items()}
    found = {chunk: {item["ref"]: item for item in state["relevant"]} for chunk, state in states.items()}

    def comparisons(pairs):
        """Comparison questions for (row, piece of the methodology) pairs: for each row, its pieces in reading order,
        as many to a question as fit."""
        items = []
        for row in dict.fromkeys(row for row, _ in pairs):
            groups, used = [[]], 0
            for key in sorted((key for r, key in pairs if r == row), key=position.get):
                if groups[-1] and used + by_key[key]["tokens"] + 2 > chunk_room:
                    groups.append([])
                    used = 0
                groups[-1].append(key)
                used += by_key[key]["tokens"] + 2
            items += [(METHODOLOGY_COMPARISON, row, group) for group in groups if group]
        return items

    def compare_whole(chunk):
        """A unit, once every row of it is searched: each row against every chunk of the methodology found for any of
        them, in all its parts - the unit and those chunks each treated as a whole."""
        refs = sorted(found[chunk], key=lambda ref: position.get(parts_of[ref][0], 0))
        return comparisons([(row["ref"], key) for row in rows_of[chunk] if askable(row) for ref in refs for key in parts_of[ref]])

    work = [item for state in states.values() if state["askable"] and not state["open batches"]
            for item in comparisons(state["to compare"])]
    for batch in batches:
        keys = [(p["ref"], p["part"]) for p in batch]
        work += [(METHODOLOGY_SEARCH, row["ref"], [key for key in keys if key in to_search[row["ref"]]])
                 for row in units if row["ref"] in to_search and any(key in to_search[row["ref"]] for key in keys)]
    neighbours = {}

    def build(item):
        kind, ref, keys = item
        row = by_ref[ref]
        upstream, downstream = linked_units(row, links, wholes)
        chosen = [by_key[key] for key in keys]
        if kind == METHODOLOGY_SEARCH:
            if row["unit_ref"] not in neighbours:
                neighbours[row["unit_ref"]] = neighbours_line(upstream, downstream)
            question = search_question(row, said.get(ref, ""), neighbours[row["unit_ref"]], chosen, unit_room)
        else:
            question = compare_question(row, said.get(ref, ""), upstream, downstream, chosen, compare_room, context_room)
        if question["tokens"] > question_room(ctx.settings, kind):
            raise EngineFault("A question of step 05 would take %d tokens where chat() holds %d with its answer. This is a "
                              "defect in the tool, not in the model under review." % (question["tokens"], question_room(ctx.settings, kind)))
        return question

    def after(question, result):
        """What an answer leads to: a unit whose rows are now all searched is compared as a whole, and a question left
        without an answer that shows more than one piece of the methodology is asked again in two halves - one piece
        may be what the gateway refuses, or the question too long for it. Splitting costs nothing when the model
        answers nothing at all, as when the token has run out: ask_all stops after the same number of questions either
        way."""
        ref, keys = question["unit_ref"], [(p[0], p[1]) for p in question["pieces"]]
        if not result["answer"]:
            half = len(keys) // 2
            return [(question["type"], ref, keys[:half]), (question["type"], ref, keys[half:])] if half else []
        if question["type"] != METHODOLOGY_SEARCH:
            return []
        chunk, shown = by_ref[ref]["unit_ref"], {r for r, _ in keys}
        for item in result["reading"]["relevant"]:
            if item["ref"] in shown:
                found[chunk].setdefault(item["ref"], item)
        mine = [row["ref"] for row in rows_of[chunk] if row["ref"] in to_search]
        was_open = any(to_search[r] for r in mine)
        to_search[ref] -= set(keys)
        return compare_whole(chunk) if was_open and not any(to_search[r] for r in mine) else []

    chat = NOTEBOOK["chat"]
    if work and chat is None:
        return StepResult({}, {"questions to ask": len(work)}, ["The methodology is searched by your chat(): run cell 2, "
                                                                 "then cell 3 again."], finished=False)
    taken, stopped, peak = ask_all(ctx, chat, work, build, lambda question: (search_check if question["type"] == METHODOLOGY_SEARCH
                                                                       else compare_check)(question),
                             after, label="methodology search and comparison", types=types)
    redact = (NOTEBOOK["live"] or LiveValues()).redact
    order = {unit["ref"]: number for number, unit in enumerate(units)}
    records, technical = [], []
    for question, result in sorted(taken, key=lambda pair: (order.get(pair[0]["unit_ref"], 0), types.index(pair[0]["type"]),
                                                            min([position.get((p[0], p[1]), 0) for p in pair[0]["pieces"]] or [0]),
                                                            pair[0]["id"])):
        if question["id"] in recorded:
            continue                                     # written already, by an earlier cell
        records.append(call_record(ctx, question, result, redact, pieces=question["pieces"], context=question.get("context")))
        technical += ["%s, %s" % (question["unit_ref"], redact(line)) for line in result["technical"]]
    log_technical(ctx, "05", technical)
    every = calls + records
    final = [state for state in methodology_account(units, chunks, every, ctx.settings)["units"].values() if state["askable"]]
    open_units = sum(1 for state in final if state["open batches"] or state["to compare"])
    counts = {"pieces of code": len(final), "chunks of the methodology": len(chunks), "batches": len(batches),
              "searches answered": sum(1 for c in every if c.get("question_type") == METHODOLOGY_SEARCH and c.get("outcome") == "answered"),
              "comparisons answered": sum(1 for c in every if c.get("question_type") == METHODOLOGY_COMPARISON and c.get("outcome") == "answered"),
              "chunks found relevant": sum(len({item["ref"] for item in state["relevant"]}) for state in final),
              "deviations flagged": sum(len(state["deviations"]) for state in final),
              "not answered now": sum(1 for record in records if record["outcome"] != "answered"),
              "pieces not finished": open_units, "most at once": peak}
    messages = [stopped] if stopped else []
    unexplained = sum(1 for unit in units if askable(unit) and unit["ref"] not in said)
    if unexplained:
        messages.append("%d pieces of code had no interpretation from step 04, so the model searched and compared them on "
                        "their code alone." % unexplained)
    if not chunks:
        messages.append("The methodology gave no chunks, so there was nothing to search.")
    elif open_units and not stopped:
        messages.append("%d pieces of code are not yet searched and compared in full: the model gave no answer to %d "
                        "questions. Run cell 3 again to ask them again; what went wrong is in run_log.txt."
                        % (open_units, counts["not answered now"]))
    unsearched = sorted({ref for state in final for ref in state["unsearched"]})
    if unsearched and len(unsearched) <= 5 and counts["searches answered"]:
        messages.append("No question showing %s has been answered, though questions showing the rest of the methodology "
                        "were: the gateway may refuse what %s. The rows of Chunks_Model say which pieces of code "
                        "are affected." % (", ".join(unsearched), "that chunk holds" if len(unsearched) == 1 else "those chunks hold"))
    elif not open_units:
        messages.append("Each of the %d pieces of code was searched against all %d chunks of the methodology, in %d "
                        "batch%s of up to %d tokens; %d chunks were found to bear on them, and %d potential deviations "
                        "flagged." % (len(final), len(chunks), len(batches), "" if len(batches) == 1 else "es",
                                      search_shares(ctx.settings)[0], counts["chunks found relevant"], counts["deviations flagged"]))
    return StepResult({"llm_calls": records}, counts, messages, finished=not open_units)


# ---------------------------------------------------------------- the pipeline runner
PIPELINE = (                       # the steps, in order, each carried out by one function of STEP_FUNCTIONS. Enforces: R11
    {"id": "01", "name": "prepare-run", "carried_out_by": "prepare_run"},
    {"id": "02", "name": "read-inputs", "carried_out_by": "read_inputs"},
    {"id": "03", "name": "link-chunks", "carried_out_by": "link_chunks"},
    {"id": "04", "name": "interpret-code", "carried_out_by": "interpret_code"},
    {"id": "05", "name": "search-methodology", "carried_out_by": "search_methodology"})

def update_manifest(store, changes):
    """Change fields of the run manifest and write it back."""
    manifest = (store.read("run_manifest") or [{}])[0]
    manifest.update(changes)
    store.append("run_manifest", [manifest])
    return manifest


def run_pipeline(paths, settings, stop_after=""):
    """Run, or resume, the pipeline. Each finished step leaves a step record; called again,
    the run continues at the first step without one. A step that says it did not finish - step 04 or 05 with
    questions still unanswered - stops the run there, and is carried out again when cell 3 runs again.
    Returns {"state", "message", "steps_run"}."""
    store = open_store(paths, settings)
    done = {record["step_id"] for record in store.read("step_records") if record.get("finished", True)}
    steps_run = []
    for step in PIPELINE:
        if stop_after and step["id"] > stop_after:
            break
        if step["id"] in done:
            continue
        finished = run_step(step, store, paths, settings)
        steps_run.append(step["name"])
        if not finished:
            message = "Step %s (%s) did not finish. Run cell 3 again to carry on from it." % (step["id"], step["name"])
            rebuild_outputs(store, paths, settings, message)
            return {"state": "unfinished", "message": message, "steps_run": steps_run}
        if stop_after and step["id"] == stop_after:
            break
    message = "Every step has run." if not stop_after else "Stopped after step %s as asked." % stop_after
    rebuild_outputs(store, paths, settings, message)
    return {"state": "finished", "message": message, "steps_run": steps_run}

def run_step(step, store, paths, settings):
    """Build the context, call the step function, write what it returns, record the step,
    rebuild the outputs and sync. Steps never touch the store themselves."""
    notes = []
    provenance = Provenance(paths.run_id, step["id"], step["name"])
    options = dict(step.get("with") or {})
    options.update({"inputs": list_input_files(paths.inputs_dir), "paths": paths,
                    "run": {"project": paths.project, "run_id": paths.run_id}})
    work_dir = os.path.join(paths.local_dir, "work")
    os.makedirs(work_dir, exist_ok=True)
    context = StepContext(settings, options, store.read, work_dir, notes.append, provenance)
    function = STEP_FUNCTIONS[step["carried_out_by"]]
    started = time.time()
    try:
        result = function(context)
    except Exception as problem:                     # a step that fails is written down, never a stopped run (R2)
        result = StepResult({}, {"step did not finish": 1}, [step_failure(step, problem, work_dir)])
    for kind in sorted(result.records):
        if kind == "llm_calls":                      # exchanges with the model: the audit log's Model_Calls sheet
            store.append_calls(result.records[kind])
        else:
            store.append(kind, result.records[kind])
    result.messages = list(result.messages) + notes
    record_step(store, step, result, time.time() - started)
    rebuild_outputs(store, paths, settings, "")
    store.sync()
    return result.finished

def step_failure(step, problem, work_dir):
    """What the analyst is told when a step could not finish, and where the details are kept for
    whoever maintains the tool. The run goes on: the steps after this one work with what there is, and
    each says what it could not do. The details go to the run's work folder, not to the evidence
    pack, because they are about the tool and not about the model under review. Enforces: R2"""
    with open(os.path.join(work_dir, "step_%s_did_not_finish.txt" % step["id"]), "w", encoding="utf-8") as handle:
        handle.write("".join(traceback.format_exception(type(problem), problem, problem.__traceback__)))
    return ("Step %s (%s) could not finish, because of a fault inside the tool (%s). The steps after it ran on what there "
            "was, and say what they could not do. The details are in the run's work folder, for whoever maintains the tool."
            % (step["id"], step["name"], type(problem).__name__))

def record_step(store, step, result, seconds):
    """Leave the step record that makes a finished step visible and resume possible."""
    produced = {kind: len(records) for kind, records in sorted(result.records.items())}
    store.append("step_records", [{
        "step_id": step["id"], "name": step["name"],
        "carried_out_by": step["carried_out_by"], "produced": produced, "counts": result.counts,
        "messages": result.messages, "finished": result.finished,
        "finished_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "seconds": round(seconds, 3)}])

def log_line(store, text):
    """Technical text (exception messages, Python names) belongs in run_log.txt only."""
    with open(os.path.join(store.local_dir, "run_log.txt"), "a", encoding="utf-8") as handle:
        handle.write("%s  %s\n" % (datetime.datetime.now().isoformat(timespec="seconds"), text))

# ---------------------------------------------------------------- step 01: prepare-run
PACKAGES_RECORDED = ("PyYAML", "openpyxl", "python-docx", "numpy", "scipy", "sympy", "rdata",
                     "pdfplumber", "pypdf", "pyreadr")

def fingerprint_file(path, corner, inputs_dir):
    """Name, corner, size, SHA-256 and content identifier of one input file."""
    with open(path, "rb") as handle:
        data = handle.read()
    return {"corner": corner, "file": os.path.relpath(path, inputs_dir).replace(os.sep, "/"),
            "bytes": len(data), "sha256": sha256_bytes(data), "swhid": swhid_content(data)}

def engine_file_hashes():
    """SHA-256 of every file that makes up the engine - its code, which holds its steps, prompts and reference
    data, and its requirements - so that an evidence pack names exactly the code that produced it."""
    import glob
    found = {}
    for pattern in ("*.py", "requirements.txt"):
        for path in sorted(glob.glob(os.path.join(ENGINE_DIR, pattern), recursive=True)):
            if os.path.isfile(path):
                found["engine/" + os.path.relpath(path, ENGINE_DIR).replace(os.sep, "/")] = file_sha256(path)
    return found

def prepare_run(ctx):
    """Step 01. Fingerprint every input, record the environment and the allow-listed
    settings, and say what changed since the run of the project this one replaced."""
    import importlib.metadata
    paths, inputs = ctx.options["paths"], ctx.options["inputs"]
    fingerprints = input_fingerprints(paths.inputs_dir, inputs)
    versions = {"python": sys.version.split()[0]}
    for name in PACKAGES_RECORDED:
        try:
            versions[name] = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:
            versions[name] = "not installed"
    previous, changes = paths.previous, []
    if previous is not None:
        before = {f["file"]: f["sha256"] for f in previous["inputs"]}
        now = {f["file"]: f["sha256"] for f in fingerprints}
        changes = ["%s: changed" % n for n in sorted(now) if n in before and before[n] != now[n]]
        changes += ["%s: new" % n for n in sorted(now) if n not in before]
        changes += ["%s: no longer present" % n for n in sorted(before) if n not in now]
    manifest = dict(ctx.options["run"], versions=versions, engine_files=engine_file_hashes(),
                    settings=ctx.settings, inputs=fingerprints, changes_since_previous_run=changes,
                    previous_run=previous["run_id"] if previous else "",
                    started_at=datetime.datetime.now().isoformat(timespec="seconds"))
    return StepResult({"run_manifest": [manifest]}, {"input files": len(fingerprints)},
                             ["%d input files fingerprinted." % len(fingerprints)])

def input_fingerprints(inputs_dir, inputs=None):
    """The fingerprint of every input file of a project, corner by corner, in the order they are read."""
    inputs = inputs or list_input_files(inputs_dir)
    found = []
    for corner, _, _ in INPUT_FOLDERS:
        found.extend(fingerprint_file(path, corner, inputs_dir) for path in inputs[corner])
    for key in ("glossary", "tag_rules"):
        if inputs[key]:
            found.append(fingerprint_file(inputs[key], key, inputs_dir))
    return found

# ---------------------------------------------------------------- the workbook's macro
# Output.xlsm carries one macro, in the workbook's own module: a reference in Chunks_Model, clicked, shows only the
# chunks it names. Its source is WORKBOOK_MACRO below, and vba_project() packs it into the part Excel reads macros
# from, xl/vbaProject.bin, written here from Microsoft's published formats - [MS-OVBA] for the project, [MS-CFB] for
# the compound file that holds it. The project holds the source alone, no compiled code, as [MS-OVBA] asks of a
# writer: Excel compiles it when it opens the workbook. It is the same for every run, and nothing in it comes from an
# input. Without macros, each link still leads to the first chunk it names. Enforces: R5, R7

LINK_COLUMNS = {("Chunks_Model", "methodology_refs"): ("Chunks_Methodology", "Click: Chunks_Methodology shows only these chunks"),
                ("Chunks_Model", "upstream"): ("Chunks_Model", "Click: Chunks_Model shows only these chunks and this one"),
                ("Chunks_Model", "downstream"): ("Chunks_Model", "Click: Chunks_Model shows only these chunks and this one"),
                ("Chunks_Model", "flagged_count"): ("Flagged_Items", "Click: Flagged_Items shows only this chunk's items"),
                ("Flagged_Items", "location"): ("Chunks_Model", "Click: Chunks_Model shows only this chunk"),
                ("Flagged_Items", "methodology_refs"): ("Chunks_Methodology", "Click: Chunks_Methodology shows only these chunks")}
LINK_KEYS = {"Chunks_Methodology": "ref", "Chunks_Model": "ref", "Flagged_Items": "location"}   # what a link finds a row by
REF_LIST = re.compile(r"^([CDM]-\d{4,})(?: \(possible\))?(?:; (?:[CDM]-\d{4,}(?: \(possible\))?|%s))*$" % re.escape(UNKNOWN_NAME))
                                            # a cell of references - some possible - and at most the note of an unknown name
LINK_FONT = "0563C1"                                            # the blue Excel gives a hyperlink

WORKBOOK_MACRO = """Option Explicit

' Verifier: a reference, clicked, shows only the rows it names.
'   Chunks_Model, Relevant Chunks in Methodology (searched by LLM): Chunks_Methodology, filtered to the chunks listed.
'   Chunks_Model, Immediate Upstream and Immediate Downstream Model Chunk: Chunks_Model, filtered to the chunks listed
'   and the row clicked, the cell clicked staying the active cell.
'   Chunks_Model, Count of Flagged Items (by LLM): Flagged_Items, filtered to the items of the row clicked.
'   Flagged_Items, Location of Flagged Item: Chunks_Model, filtered to that chunk.
'   Flagged_Items, Ref in Chunks_Methodology: Chunks_Methodology, filtered to the chunks listed.
' A chunk shown in several rows (C-0012-1, C-0012-2 ...) is shown whole. Nothing else is changed:
' Data > Clear shows every row again. Without macros, a link still leads to the first chunk it names.

Private Sub Workbook_SheetFollowHyperlink(ByVal Sh As Object, ByVal Target As Hyperlink)
    Dim clicked As Range, heading As String, refs As String, rowRef As String
    On Error GoTo Finish
    Set clicked = Target.Range.Cells(1, 1)
    heading = CStr(Sh.Cells(1, clicked.Column).Value)
    refs = RefsIn(CStr(clicked.Value))
    rowRef = UnitOf(CStr(Sh.Cells(clicked.Row, 1).Value))
    If Sh.Name = "Chunks_Model" Then
        If heading = "Relevant Chunks in Methodology (searched by LLM)" Then
            ShowOnly ThisWorkbook.Worksheets("Chunks_Methodology"), refs, 1
        ElseIf heading = "Immediate Upstream Model Chunk" Or heading = "Immediate Downstream Model Chunk" Then
            ShowOnly Sh, refs & "|" & rowRef, 1, clicked
        ElseIf heading = "Count of Flagged Items (by LLM)" Then
            ShowOnly ThisWorkbook.Worksheets("Flagged_Items"), rowRef, 2
        End If
    ElseIf Sh.Name = "Flagged_Items" Then
        If heading = "Location of Flagged Item" Then
            ShowOnly ThisWorkbook.Worksheets("Chunks_Model"), refs, 1
        ElseIf heading = "Ref in Chunks_Methodology" Then
            ShowOnly ThisWorkbook.Worksheets("Chunks_Methodology"), refs, 1
        End If
    End If
Finish:
End Sub

Private Function RefsIn(ByVal cellText As String) As String
    ' The references a cell names, such as C-0012 or M-0003, joined with |.
    Dim position As Long, letter As String, word As String, found As String
    cellText = cellText & " "
    For position = 1 To Len(cellText)
        letter = Mid$(cellText, position, 1)
        If InStr(1, "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789-", letter, vbBinaryCompare) > 0 Then
            word = word & letter
        Else
            If IsRef(word) Then found = found & "|" & word
            word = ""
        End If
    Next position
    If Len(found) > 0 Then RefsIn = Mid$(found, 2)
End Function

Private Function IsRef(ByVal word As String) As Boolean
    ' C-0012, D-0003 or M-0017: C, D or M, a hyphen, then four digits or more.
    Dim position As Long
    If Len(word) < 6 Then Exit Function
    If InStr(1, "CDM", Left$(word, 1), vbBinaryCompare) = 0 Then Exit Function
    If Mid$(word, 2, 1) <> "-" Then Exit Function
    For position = 3 To Len(word)
        If InStr(1, "0123456789", Mid$(word, position, 1), vbBinaryCompare) = 0 Then Exit Function
    Next position
    IsRef = True
End Function

Private Function UnitOf(ByVal rowRef As String) As String
    ' A row's reference without its part: M-0003-2 is a part of M-0003.
    Dim cut As Long
    cut = InStr(3, rowRef, "-")
    If cut > 0 Then UnitOf = Left$(rowRef, cut - 1) Else UnitOf = rowRef
End Function

Private Sub ShowOnly(ByVal onSheet As Worksheet, ByVal refs As String, ByVal byColumn As Long, Optional ByVal returnTo As Range)
    ' Filter the sheet on one column - Ref, or Flagged_Items' Location - to the chunks named, each with every row it takes.
    ' A filter of the sheet clicked on ends on the cell clicked (returnTo); a filter of another sheet, on its first row.
    Dim wanted As Variant, shown As String, cellRef As String, last As Long, rowNumber As Long, i As Long, locked As Boolean
    If Len(refs) = 0 Then Exit Sub
    wanted = Split(refs, "|")
    last = onSheet.UsedRange.Row + onSheet.UsedRange.Rows.Count - 1
    For rowNumber = 2 To last
        cellRef = CStr(onSheet.Cells(rowNumber, byColumn).Value)
        For i = LBound(wanted) To UBound(wanted)
            If Len(wanted(i)) > 0 Then
                If cellRef = wanted(i) Or Left$(cellRef, Len(wanted(i)) + 1) = wanted(i) & "-" Then
                    shown = shown & "|" & cellRef
                    Exit For
                End If
            End If
        Next i
    Next rowNumber
    If Len(shown) = 0 Then Exit Sub
    locked = onSheet.ProtectContents
    On Error GoTo Restore
    If locked Then onSheet.Unprotect
    On Error Resume Next
    onSheet.ShowAllData                                ' clears any filter; with none on, there is nothing to clear
    On Error GoTo Restore
    If Not onSheet.AutoFilterMode Then onSheet.Range("A1").CurrentRegion.AutoFilter
    onSheet.AutoFilter.Range.AutoFilter Field:=byColumn, Criteria1:=Split(Mid$(shown, 2), "|"), Operator:=xlFilterValues
    onSheet.Activate
    If returnTo Is Nothing Then
        Application.Goto onSheet.Range("A1"), True
    Else
        returnTo.Select
    End If
Restore:
    If locked Then onSheet.Protect DrawingObjects:=False, Contents:=True, Scenarios:=False, AllowFormattingColumns:=True, AllowFiltering:=True
End Sub
"""

VBA_WORKBOOK_CLASS = "0{00020819-0000-0000-C000-000000000046}"   # the classes of a workbook's and a sheet's module
VBA_SHEET_CLASS = "0{00020820-0000-0000-C000-000000000046}"
VBA_STDOLE = b"*\\G{00020430-0000-0000-C000-000000000046}#2.0#0#C:\\WINDOWS\\system32\\stdole2.tlb#OLE Automation"


def ovba_compress(data):
    """[MS-OVBA] 2.4.1 compression: chunks of 4,096 bytes, each a run of tokens - a literal byte, or a copy of bytes
    met earlier in the chunk. A chunk that would not shrink is kept as it is."""
    out = bytearray(b"\x01")
    for start in range(0, len(data), 4096):
        chunk = data[start:start + 4096]
        body, position, seen = bytearray(), 0, {}
        while position < len(chunk):
            flag_at, flags = len(body), 0
            body.append(0)
            for bit in range(8):
                if position >= len(chunk):
                    break
                bits = max((position - 1).bit_length(), 4) if position else 16
                longest, offset = 0, 0
                most = min((0xFFFF >> bits) + 3, len(chunk) - position) if position else 0
                for earlier in reversed(seen.get(bytes(chunk[position:position + 3]), [])) if most >= 3 else ():
                    length = 0
                    while length < most and chunk[earlier + length] == chunk[position + length]:
                        length += 1
                    if length > longest:
                        longest, offset = length, position - earlier
                        if length == most:
                            break
                step = longest if longest >= 3 else 1
                if longest >= 3:
                    body += struct.pack("<H", ((offset - 1) << (16 - bits)) | (longest - 3))
                    flags |= 1 << bit
                else:
                    body.append(chunk[position])
                for index in range(position, position + step):
                    seen.setdefault(bytes(chunk[index:index + 3]), []).append(index)
                position += step
            body[flag_at] = flags
        if len(body) > 4096:
            if len(chunk) < 4096:
                raise ValueError("a short chunk that does not compress")
            out += struct.pack("<H", 0x3FFF) + chunk
        else:
            out += struct.pack("<H", 0xB000 | (len(body) - 1)) + body
    return bytes(out)


def ovba_encrypt(project_id, data, seed):
    """[MS-OVBA] 2.4.3.2 data encryption, for the project's protection, password and visibility: a seed, the version,
    the project's key - the sum of its id's bytes - then the data's length and the data, each byte folded into the
    ones before it. Returned as the hexadecimal text the PROJECT stream holds."""
    key = sum(project_id.encode("latin-1")) & 0xFF
    version_enc, key_enc = seed ^ 2, seed ^ key
    out = [seed, version_enc, key_enc]
    plain1, enc1, enc2 = key, key_enc, version_enc
    for byte in [7] * ((seed & 6) // 2) + list(struct.pack("<I", len(data)) + data):
        byte_enc = byte ^ ((enc2 + plain1) & 0xFF)
        out.append(byte_enc)
        enc2, enc1, plain1 = enc1, byte_enc, byte
    return bytes(out).hex().upper()


def vba_record(record_id, payload=b""):
    """One record of the project's dir stream: its id, its size, what it holds."""
    return struct.pack("<HI", record_id, len(payload)) + payload


@functools.lru_cache(maxsize=None)
def vba_project(sheets):
    """The bytes of xl/vbaProject.bin: the workbook's module holding WORKBOOK_MACRO, and a module for each sheet,
    `sheets` giving their code names in order. [MS-OVBA] 2.2 and 2.3; the same bytes on every call."""
    project_id = "{%s}" % str(uuid.UUID(bytes=hashlib.sha256(b"Verifier workbook macro").digest()[:16])).upper()
    attributes = ('Attribute VB_Name = "%s"\r\nAttribute VB_Base = "%s"\r\nAttribute VB_GlobalNameSpace = False\r\n'
                  'Attribute VB_Creatable = False\r\nAttribute VB_PredeclaredId = True\r\nAttribute VB_Exposed = True\r\n'
                  'Attribute VB_TemplateDerived = False\r\nAttribute VB_Customizable = True\r\n')
    modules = [("ThisWorkbook", attributes % ("ThisWorkbook", VBA_WORKBOOK_CLASS) + WORKBOOK_MACRO.replace("\n", "\r\n"))]
    modules += [(name, attributes % (name, VBA_SHEET_CLASS)) for name in sheets]
    wide = lambda name: name.encode("utf-16-le")
    directory = (vba_record(0x01, struct.pack("<I", 1)) + vba_record(0x02, struct.pack("<I", 0x409))
                 + vba_record(0x14, struct.pack("<I", 0x409)) + vba_record(0x03, struct.pack("<H", 1252))
                 + vba_record(0x04, b"VBAProject") + vba_record(0x05) + vba_record(0x40) + vba_record(0x06)
                 + vba_record(0x3D) + vba_record(0x07, struct.pack("<I", 0)) + vba_record(0x08, struct.pack("<I", 0))
                 + struct.pack("<HIIH", 0x09, 4, 1, 0) + vba_record(0x0C) + vba_record(0x3C)
                 + vba_record(0x16, b"stdole") + vba_record(0x3E, wide("stdole"))
                 + vba_record(0x0D, struct.pack("<I", len(VBA_STDOLE)) + VBA_STDOLE + b"\0" * 6)
                 + vba_record(0x0F, struct.pack("<H", len(modules))) + vba_record(0x13, struct.pack("<H", 0xFFFF)))
    for name, _ in modules:
        encoded = name.encode("latin-1")
        directory += (vba_record(0x19, encoded) + vba_record(0x47, wide(name)) + vba_record(0x1A, encoded)
                      + vba_record(0x32, wide(name)) + vba_record(0x1C) + vba_record(0x48)
                      + vba_record(0x31, struct.pack("<I", 0)) + vba_record(0x1E, struct.pack("<I", 0))
                      + vba_record(0x2C, struct.pack("<H", 0xFFFF)) + vba_record(0x22) + vba_record(0x2B))
    directory += vba_record(0x10)
    project = ('ID="%s"\r\n' % project_id + "".join("Document=%s/&H00000000\r\n" % name for name, _ in modules)
               + 'Name="VBAProject"\r\nHelpContextID="0"\r\nVersionCompatible32="393222000"\r\n'
               + 'CMG="%s"\r\nDPB="%s"\r\nGC="%s"\r\n\r\n' % (ovba_encrypt(project_id, b"\0\0\0\0", 0x3E),
                                                              ovba_encrypt(project_id, b"\0", 0x9A),
                                                              ovba_encrypt(project_id, b"\xff", 0x51))
               + "[Host Extender Info]\r\n&H00000001={3832D640-CF90-11CF-8E43-00A0C911005A};VBE;&H00000000\r\n\r\n"
               + "[Workspace]\r\n" + "".join("%s=0, 0, 0, 0, C\r\n" % name for name, _ in modules))
    streams = {"PROJECT": project.encode("latin-1"),
               "PROJECTwm": b"".join(name.encode("latin-1") + b"\0" + wide(name) + b"\0\0" for name, _ in modules) + b"\0\0",
               "VBA/_VBA_PROJECT": b"\xcc\x61\xff\xff\x00\x00\x00",
               "VBA/dir": ovba_compress(directory)}
    streams.update({"VBA/" + name: ovba_compress(source.encode("latin-1")) for name, source in modules})
    return cfb_file(streams)


def cfb_file(streams):
    """A compound file, [MS-CFB] version 3, holding `streams` - a path of names, "VBA/dir", to its bytes; the storages
    are those the paths name. Streams under 4,096 bytes live in the mini stream, in sectors of 64 bytes; the rest,
    and the file's own tables, in sectors of 512. The entries of each storage form a balanced tree in the order the
    format sets - shorter names first, then by their capitals - coloured so that it is a red-black tree."""
    SECTOR, MINI, CUTOFF, FREE, END, FATSECT = 512, 64, 4096, 0xFFFFFFFF, 0xFFFFFFFE, 0xFFFFFFFD
    nodes = [{"name": "Root Entry", "type": 5, "children": {}, "data": b""}]
    for path in streams:
        parent = nodes[0]
        *storages, name = path.split("/")
        for storage in storages:
            if storage not in parent["children"]:
                nodes.append({"name": storage, "type": 1, "children": {}, "data": b""})
                parent["children"][storage] = len(nodes) - 1
            parent = nodes[parent["children"][storage]]
        nodes.append({"name": name, "type": 2, "children": {}, "data": streams[path]})
        parent["children"][name] = len(nodes) - 1
    for node in nodes:
        node.update(left=FREE, right=FREE, child=FREE, colour=1, start=0, size=len(node["data"]))
    for node in nodes:                                            # each storage's entries: a balanced tree
        order = sorted(node["children"].values(), key=lambda i: (len(nodes[i]["name"]), nodes[i]["name"].upper()))
        depths = {}
        def grow(members, depth):
            if not members:
                return FREE
            middle = len(members) // 2
            index = members[middle]
            depths[index] = depth
            nodes[index]["left"], nodes[index]["right"] = grow(members[:middle], depth + 1), grow(members[middle + 1:], depth + 1)
            return index
        node["child"] = grow(order, 0)
        if depths and len(depths) != 2 ** (max(depths.values()) + 1) - 1:
            for index, depth in depths.items():                   # an incomplete last level is red
                nodes[index]["colour"] = 0 if depth == max(depths.values()) else 1
    sectors, fat = [], []
    def allocate(data):
        if not data:
            return END
        count, first = -(-len(data) // SECTOR), len(sectors)
        for number in range(count):
            sectors.append(data[number * SECTOR:(number + 1) * SECTOR].ljust(SECTOR, b"\0"))
            fat.append(first + number + 1 if number < count - 1 else END)
        return first
    mini, minifat = bytearray(), []
    for node in nodes:
        if node["type"] == 2 and node["size"] < CUTOFF:
            count = -(-node["size"] // MINI)
            node["start"] = len(minifat) if count else END
            minifat += [node["start"] + number + 1 if number < count - 1 else END for number in range(count)]
            mini += node["data"].ljust(count * MINI, b"\0")
    for node in nodes:
        if node["type"] == 2 and node["size"] >= CUTOFF:
            node["start"] = allocate(node["data"])
    nodes[0]["start"], nodes[0]["size"] = allocate(bytes(mini)), len(mini)
    minifat_sectors = -(-len(minifat) // (SECTOR // 4))
    minifat_start = allocate(struct.pack("<%dI" % (minifat_sectors * SECTOR // 4), *(minifat + [FREE] * (minifat_sectors * SECTOR // 4 - len(minifat)))))
    entries = b""
    for node in nodes:
        name = node["name"].encode("utf-16-le") + b"\0\0"
        entries += struct.pack("<64sHBBIII16sIQQIQ", name, len(name), node["type"], node["colour"], node["left"],
                               node["right"], node["child"], b"\0" * 16, 0, 0, 0, node["start"], node["size"])
    entries += struct.pack("<64sHBBIII16sIQQIQ", b"", 0, 0, 0, FREE, FREE, FREE, b"\0" * 16, 0, 0, 0, 0, 0) * (-len(nodes) % 4)
    directory_start = allocate(entries)
    fat_sectors = 1
    while fat_sectors * (SECTOR // 4) < len(sectors) + fat_sectors:
        fat_sectors += 1
    fat_start = len(sectors)
    table = fat + [FATSECT] * fat_sectors
    table += [FREE] * (fat_sectors * (SECTOR // 4) - len(table))
    sectors += [struct.pack("<%dI" % (SECTOR // 4), *table[number * (SECTOR // 4):(number + 1) * (SECTOR // 4)]) for number in range(fat_sectors)]
    if fat_sectors > 109:
        raise ValueError("a compound file this large needs more than its header's table of tables")
    header = struct.pack("<8s16sHHHHH6sIIIIIIIII", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", b"\0" * 16, 0x3E, 3, 0xFFFE, 9, 6,
                         b"\0" * 6, 0, fat_sectors, directory_start, 0, CUTOFF, minifat_start if minifat else END,
                         minifat_sectors, END, 0)
    header += struct.pack("<109I", *([fat_start + number for number in range(fat_sectors)] + [FREE] * (109 - fat_sectors)))
    return header + b"".join(sectors)


# ---------------------------------------------------------------- Output.xlsm
CELL_WITHHELD = "This text could not be shown in plain words; the technical text is in the audit records."
CUT_NOTE = " ... (cut here; the full text is in the audit files)"
PYTHON_TRACES = re.compile(r"Traceback|\b\w+(Err" r"or|Exception)\b|<class |object at 0x|\bnan\b|\bverifier\d_\w+|"
                           r"[:=(\[]\s*None\b|\{'|\['|^None$")


def plain_cell(value, input_text, store):
    """The last gate before a cell is written. the tool's own words must be free of Python
    traces and of words the wording rule rejects; otherwise the text goes to run_log.txt
    and the cell gets one fixed, plain sentence. Enforces: R1, R10"""
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value == int(value) else plain_number(value)
    text = str(value)
    if not input_text:
        said = own_words(text)
        if PYTHON_TRACES.search(said) or has_banned_wording(said):
            log_line(store, "cell text withheld: " + text)
            text = CELL_WITHHELD
    return text if len(text) <= 32000 else text[:31900] + CUT_NOTE


def run_identity(store, paths):
    """What ties a workbook to its run: also written into the workbook's properties."""
    return {"project": paths.project, "run_id": paths.run_id}

def stored_data_rows(units, links):
    """Model_Package_Info's account of the stored objects, once step 03 has linked them: how many the code reads, and
    those no code reads - each of which says so in its own row of Immediate Downstream Model Chunk. Enforces: R2"""
    if not links:
        return []
    linked = {record["ref"]: record for record in links}
    data = [u for u in units if u["kind"] in (KIND_TABLE, KIND_OBJECT)]
    unread = [u for u in data if not (linked.get(u["ref"]) or {}).get("downstream")
              and not (linked.get(u["ref"]) or {}).get("possible_downstream")]
    rows = [("Stored objects read by the code", "%d of %d" % (len(data) - len(unread), len(data)))] if data else []
    if unread:
        rows.append(("Stored objects no code reads", "; ".join("%s %s (%s)" % (u["ref"], u["name"], u["file"]) for u in unread)
                     + ". The model's code reads them by no name, file, helper or environment the tool can see: each "
                       "says so in Immediate Downstream Model Chunk."))
    return rows


def rows_package_info(store, paths, settings, progress, split=None):
    """The rows of Model_Package_Info: identity, inputs, what was read, and the repairs made while reading."""
    identity, rows = run_identity(store, paths), []
    def add(group, item, value):
        rows.append({"group": group, "item": item, "value": value})
    add("Identity", "Project", identity["project"])
    add("Identity", "Run", identity["run_id"])
    add("Identity", "Run progress", progress)
    manifest = (store.read("run_manifest") or [{}])[0]
    if manifest.get("engine_files"):
        add("Identity", "Engine files fingerprint", sha256_text(canonical_json(manifest["engine_files"]))[:16] +
            " (%d files; the full list is in the run manifest)" % len(manifest["engine_files"]))
    for entry in manifest.get("inputs", []):
        add("Inputs", entry["file"], "SHA-256 %s (%d bytes)" % (entry["sha256"], entry["bytes"]))
    for change in manifest.get("changes_since_previous_run", []):
        add("Inputs", "Changed since run %s" % manifest.get("previous_run", ""), change)
    for info in store.read("package_info"):
        for row in info.get("rows", []):
            add(row["group"], row["item"], row["value"])
    for row in store.read("info_rows"):
        add(row["group"], row["item"], row["value"])
    stored = stored_data_rows(store.read("model_units"), store.read("unit_links"))
    at = max((number for number, row in enumerate(rows) if row["group"] == "Package data"), default=len(rows) - 1) + 1
    rows[at:at] = [{"group": "Package data", "item": item, "value": value} for item, value in stored]
    for group, several in (split or {}).items():          # beside the other rows of their group: what takes several rows
        if several:
            at = max((number for number, row in enumerate(rows) if row["group"] == group), default=len(rows) - 1) + 1
            code = group == "Package"
            rows.insert(at, {"group": group, "item": "Pieces shown in several rows" if code else "Chunks shown in several rows",
                             "value": "; ".join("%s in %d rows, %s-1 to %s-%d" % (ref, n, ref, ref, n) for ref, n in several) +
                                      (". Each is one piece of the package, too long for one row, and is treated as one wherever "
                                       "it is used: what the model says of it stands together on its first row." if code else
                                       ". Each is one chunk, too long for one row, and is treated as one wherever it is used.")})
    repairs = {}
    for repair in store.read("read_repairs"):
        repairs.setdefault(repair["file"], []).append(repair["kind"])
    for name in sorted(repairs):
        kinds = sorted(set(repairs[name]))
        add("Repairs made while reading", name, "; ".join("%s (%d)" % (k, repairs[name].count(k)) for k in kinds))
    return rows


def rows_chunks(chunks, slicer=None):
    """The rows of Chunks_Methodology and Chunks_Documentation: a chunk in one row or, too long for one, in rows
    C-0045-1, C-0045-2 and so on as `slicer` cuts it, each with the chunk's section, type and file. Enforces: R2, R13"""
    rows = []
    for c in chunks:
        slices = slicer(c) if slicer else [(c["text"], False)]
        for number, (part, joined) in enumerate(slices, start=1):
            rows.append({"ref": c["ref"] if len(slices) == 1 else "%s-%d" % (c["ref"], number), "chunk_ref": c["ref"],
                         "joined": joined, "section": " > ".join(c["heading_chain"]), "kind": c["kind"], "text": part,
                         "source_file": c["source_file"]})
    return rows


def interpretation_of(parts, said, unanswered, asked):
    """What the organisation's model says of a unit: its answer, or, for a unit asked about in parts, the answers of
    all its parts together, each under the part and the lines it explains."""
    texts = []
    for part in parts:
        answer = said.get(part["ref"]) or (NO_ANSWER if part["ref"] in unanswered else "")
        texts.append(answer or (NOT_ASKED if asked and not askable(part) else ""))
    if len(parts) == 1:
        return texts[0]
    return "\n\n".join("Part %d of %d%s: %s" % (part["part"], part["parts"], ", lines %d-%d" % tuple(part["lines"]) if part.get("lines") else "", answer)
                        for part, answer in zip(parts, texts) if answer)


def spread_rows(ref, parts, columns):
    """The rows a unit takes on Chunks_Model: its code down its parts, and each of its own columns from the first row
    down, continued on the rows below where longer than a cell, never cut. A unit takes as many rows as the longest of
    these needs; with more than one, they are numbered M-0003-1, M-0003-2 and so on. Enforces: R2, R13"""
    cut = {field: ([value] if value is not None else []) if not isinstance(value, str) else
                  [piece for piece, _ in text_slices(value, float("inf"), SLICE_CHARS)] if value else []
           for field, value in columns.items()}
    count = max([len(parts)] + [len(pieces) for pieces in cut.values()])
    rows = []
    for number in range(count):
        part = parts[number] if number < len(parts) else None
        lines = part.get("lines") if part else None
        row = {"ref": ref if count == 1 else "%s-%d" % (ref, number + 1), "unit_ref": ref, "code_row": part is not None,
               "joined": part["joined"] if part else False, "kind": parts[0]["kind"], "file": parts[0]["file"],
               "text": part["text"] if part else "", "lines": "%d-%d" % tuple(lines) if lines else ""}
        row.update({field: pieces[number] if number < len(pieces) else "" for field, pieces in cut.items()})
        rows.append(row)
    return rows


def rows_model_units(units, calls=(), links=None, methodology=None, asked=None):
    """The rows of Chunks_Model. `units` are the rows model_rows makes: a unit whole, or the parts of a long one. A unit
    is one analytical chunk however many rows it takes: the units it takes something from and gives something to
    (step 03), what the organisation's model says of it (step 04 - its parts' explanations together), and the chunks
    of the methodology found for it with the potential deviations flagged (step 05) stand once, from its first row
    down, and its code runs down its rows (spread_rows). Enforces: R2, R3"""
    links, methodology = links or {}, (methodology or {}).get("units")
    said, unanswered = {}, set()
    for call in calls:
        if call.get("question_type") == CODE_QUESTION:
            if call.get("outcome") == "answered":
                said[call["unit_ref"]] = call["answer"]
                unanswered.discard(call["unit_ref"])
            elif call["unit_ref"] not in said:
                unanswered.add(call["unit_ref"])
    was_asked, by_unit, questioned, rows = bool(said or unanswered), {}, {}, []
    for row in units:
        by_unit.setdefault(row["unit_ref"], []).append(row)
    for row in asked if asked is not None else units:       # the rows the model was asked about: a dataset once
        questioned.setdefault(row["unit_ref"], []).append(row)
    for ref, parts in by_unit.items():
        linked = links.get(ref) or {}
        refs, count = methodology_cells(methodology[ref]) if methodology and ref in methodology else ("", None)
        rows += spread_rows(ref, parts, {"interpretation": interpretation_of(questioned.get(ref, parts), said, unanswered, was_asked),
                                         "methodology_refs": refs, "flagged_count": count,
                                         "upstream": link_cell(linked, "upstream", parts[0]),
                                         "downstream": link_cell(linked, "downstream", parts[0])})
    return rows

def link_cell(linked, side, unit):
    """A unit's cell of Immediate Upstream or Immediate Downstream Model Chunk: the units the code shows it takes from
    or gives to, then those a name known only in part could be, each marked (possible), then what could not be known.
    A stored object no code reads says so. Empty before step 03 has run. Enforces: R2"""
    if not linked:
        return ""
    proven = list(linked.get(side) or ())
    items = proven + [ref + POSSIBLE for ref in linked.get("possible_" + side) or () if ref not in proven]
    if side == "upstream":
        items += list(linked.get("notes") or ())
    elif not items and unit.get("kind") in (KIND_TABLE, KIND_OBJECT):
        items = [READ_BY_NO_CODE]
    return "; ".join(items)


def rows_flagged_items(unit_rows, methodology):
    """Flagged_Items: one row per item the model flagged, numbered F-0001, F-0002 ... in the order of Chunks_Model and,
    within a piece, in the order the model gave them; with the piece it was found in, the chunks of the methodology
    it rests on, each of its parts in a column of its own, and the reviewer's decision and notes, left empty. The
    items of a piece shown in several rows, or compared in several questions, are one list. Enforces: R2, R3"""
    states = (methodology or {}).get("units") or {}          # methodology_account: its state of each piece
    rows = []
    for ref in dict.fromkeys(row["unit_ref"] for row in unit_rows):
        for item in (states.get(ref) or {}).get("deviations") or ():
            head, title = DEVIATION_KINDS.get(item.get("kind"), "Flagged"), item.get("title") or ""
            rows.append({"ref": "F-%04d" % (len(rows) + 1), "location": ref, "methodology_refs": "; ".join(item.get("refs") or ()),
                         "type": "%s: %s" % (head, title) if title else head + ".",
                         "methodology": item.get("methodology") or "", "code": item.get("code") or "",
                         "why": item.get("why") or "", "effect": item.get("effect") or "",
                         "example": item.get("example") or NO_EXAMPLE, "decision": "", "notes": ""})
    return rows


def several_rows(rows, key):
    """The chunks a sheet shows in more than one row, with how many, in order: [(ref, rows)]."""
    counts = {}
    for row in rows:
        counts[row[key]] = counts.get(row[key], 0) + 1
    return [(ref, count) for ref, count in counts.items() if count > 1]


def sheet_rows(store, paths, settings, progress):
    """The rows of all five sheets, by sheet name. A chunk too long for one row takes several, on every sheet."""
    links = {r["ref"]: r for r in store.read("unit_links")}
    whole, calls, canon = store.read("model_units"), store.read("llm_calls"), store.read("chunks_canon")
    parts, asked = model_rows(whole, settings), asked_rows(whole, settings)
    searched = any(record.get("step_id") == "05" for record in store.read("step_records")) or any(
        call.get("question_type") in (METHODOLOGY_SEARCH, METHODOLOGY_COMPARISON) for call in calls)
    methodology = methodology_account(asked, canon, calls, settings) if searched else None
    unit_rows = rows_model_units(parts, calls, links, methodology, asked)
    cap = piece_cap(settings)
    canon_rows = rows_chunks(canon, lambda chunk: methodology_slices(chunk, cap))
    doc_rows = rows_chunks(store.read("chunks_doc"), lambda chunk: piece_slices(chunk["text"] or "", 0, float("inf"), SLICE_CHARS))
    split = {"Package": several_rows(unit_rows, "unit_ref"), "Methodology files": several_rows(canon_rows, "chunk_ref"),
             "Documentation files": several_rows(doc_rows, "chunk_ref")}
    return {"Model_Package_Info": rows_package_info(store, paths, settings, progress, split),
            "Chunks_Methodology": canon_rows, "Chunks_Documentation": doc_rows, "Chunks_Model": unit_rows,
            "Flagged_Items": rows_flagged_items(unit_rows, methodology)}

def check_written_totals(rows, store):
    """The identity of the workbook: every unit read is one row of its sheet, or several, and no row is anything else;
    a unit's rows, joined back, are its text exactly, character for character. Raised as a fault of the tool, never as
    a remark about the model. Enforces: R2, R13"""
    for kind, sheet, key in (("model_units", "Chunks_Model", "unit_ref"), ("chunks_doc", "Chunks_Documentation", "chunk_ref"),
                             ("chunks_canon", "Chunks_Methodology", "chunk_ref")):
        read, written, texts = {record["ref"]: record for record in store.read(kind)}, rows[sheet], {}
        for row in written:
            if row.get("code_row", True):
                texts.setdefault(row[key], []).append(row)
        refs = [row["ref"] for row in written]
        if {row[key] for row in written} != set(read) or len(set(refs)) != len(refs) or \
                any(joined_rows(texts.get(ref, [])) != (read[ref]["text"] or "") for ref in read):
            raise EngineFault(
                "The workbook does not hold exactly what was read: %d unit(s) read for %s and %d row(s) written. "
                "This is a defect in the tool, not in the model under review." % (len(read), sheet, len(written)))
    flagged, counted, pieces = {}, {}, {row["unit_ref"] for row in rows["Chunks_Model"]}
    for row in rows.get("Flagged_Items") or ():
        flagged[row["location"]] = flagged.get(row["location"], 0) + 1
    for row in rows["Chunks_Model"]:
        value = row.get("flagged_count")
        if isinstance(value, int) or (isinstance(value, str) and value[:1].isdigit()):
            counted[row["unit_ref"]] = int(str(value).split()[0])
    ids = [row["ref"] for row in rows.get("Flagged_Items") or ()]
    if set(flagged) - pieces or set(flagged) - set(counted) or len(set(ids)) != len(ids) or \
            any(flagged.get(ref, 0) != number for ref, number in counted.items()):
        raise EngineFault("The workbook's flagged items do not match the counts of Chunks_Model: %d item(s) on "
                          "Flagged_Items. This is a defect in the tool, not in the model under review." % len(ids))

# ---------------------------------------------------------------- the workbook layout
# Every sheet of Output.xlsm in order, and every column of each: its header, its colour group, the
# field of the row it shows, its width, and whether it is typed by a person (input_text). One line
# per column. Parsed on every call, so a caller's change stays its own. Enforces: R10
WORKBOOK_LAYOUT_YAML = r'''colours: {identity: D9E1F2, code_text: E2EFDA, methodology: FCE4D6, documentation: E4DFEC, assessments: DDEBF7, model: FFF2CC, links: D0E0E3, reviewer: C6EFCE}
sheets:
- name: Model_Package_Info
  columns:
  - {header: Group, group: identity, field: group, width: 26}
  - {header: Item, group: identity, field: item, width: 40}
  - {header: Value, group: assessments, field: value, width: 90}
- name: Chunks_Methodology
  columns:
  - {header: Ref, group: identity, field: ref, width: 10}
  - {header: Section (heading chain), group: identity, field: section, width: 44}
  - {header: Type, group: identity, field: kind, width: 11}
  - {header: Text, group: methodology, field: text, width: 90, input_text: true}
  - {header: Source file, group: identity, field: source_file, width: 28}
- name: Chunks_Documentation
  columns:
  - {header: Ref, group: identity, field: ref, width: 10}
  - {header: Section (heading chain), group: identity, field: section, width: 44}
  - {header: Type, group: identity, field: kind, width: 11}
  - {header: Text, group: documentation, field: text, width: 90, input_text: true}
  - {header: Source file, group: identity, field: source_file, width: 28}
- name: Chunks_Model
  columns:
  - {header: Ref, group: identity, field: ref, width: 10}
  - {header: Kind, group: identity, field: kind, width: 20}
  - {header: File, group: identity, field: file, width: 30}
  - {header: Lines, group: identity, field: lines, width: 10}
  - {header: Text, group: code_text, field: text, width: 80, input_text: true}
  - {header: Immediate Upstream Model Chunk, group: links, field: upstream, width: 22}
  - {header: Immediate Downstream Model Chunk, group: links, field: downstream, width: 22}
  - {header: Code Interpretation (by LLM), group: model, field: interpretation, width: 80}
  - {header: Relevant Chunks in Methodology (searched by LLM), group: model, field: methodology_refs, width: 24}
  - {header: Count of Flagged Items (by LLM), group: model, field: flagged_count, width: 14}
- name: Flagged_Items
  columns:
  - {header: Ref, group: identity, field: ref, width: 10}
  - {header: Location of Flagged Item, group: links, field: location, width: 14}
  - {header: Ref in Chunks_Methodology, group: links, field: methodology_refs, width: 20}
  - {header: Type, group: model, field: type, width: 36}
  - {header: Methodology Says, group: model, field: methodology, width: 50}
  - {header: Code Does, group: model, field: code, width: 50}
  - {header: Why Potential Flagged Item, group: model, field: why, width: 50}
  - {header: Effect, group: model, field: effect, width: 40}
  - {header: Concrete Example of Potential Deviation, group: model, field: example, width: 55}
  - {header: Decision (by Human Reviewer), group: reviewer, field: decision, width: 22, editable: true, choices: ['True Positive', 'False Positive', 'True Negative', 'False Negative', 'For further discussion', 'Other Case (see notes)']}
  - {header: "Human Reviewer's Notes", group: reviewer, field: notes, width: 45, editable: true}
'''

def load_layout():
    """The workbook layout: every sheet and every column of Output.xlsm."""
    return yaml.safe_load(WORKBOOK_LAYOUT_YAML)

def write_sheet(sheet, sheet_layout, rows, colours, settings, store):
    """One generic writer for every sheet: header row and first column frozen, filter on
    the header, wrapped text, no merged cells; a column a person fills in (editable) unlocked and shaded in the
    reviewer's colour, with a dropdown of its choices where the layout gives them."""
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    columns = sheet_layout["columns"]
    wrap = Alignment(wrap_text=True, vertical="top")
    for number, column in enumerate(columns, start=1):
        header = column["header"]
        cell = sheet.cell(row=1, column=number, value=header)
        cell.font, cell.alignment = Font(bold=True), wrap
        cell.fill = PatternFill("solid", start_color=colours[column["group"]])
        sheet.column_dimensions[get_column_letter(number)].width = column.get("width", 20)
    for row_number, row in enumerate(rows, start=2):
        for number, column in enumerate(columns, start=1):
            value = row.get(column["field"])
            cell = sheet.cell(row=row_number, column=number, value=plain_cell(value, column.get("input_text"), store))
            cell.alignment = wrap
    for number, column in enumerate(columns, start=1):     # a column a person fills in: unlocked, shaded, its choices
        if not column.get("editable"):
            continue
        letter = get_column_letter(number)
        for row_number in range(2, len(rows) + 2):
            cell = sheet.cell(row=row_number, column=number)
            cell.protection = Protection(locked=False)
            cell.fill = PatternFill("solid", start_color=colours[column["group"]])
        if column.get("choices") and rows:
            check = DataValidation(type="list", formula1='"%s"' % ",".join(column["choices"]), allow_blank=True,
                                   showErrorMessage=True, errorTitle=column["header"], error="Choose one of the listed values.")
            check.add("%s2:%s%d" % (letter, letter, len(rows) + 1))
            sheet.add_data_validation(check)
    last = get_column_letter(len(columns))
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = "A1:%s%d" % (last, max(1, len(rows) + 1))
    if settings["protect_sheets"]:
        sheet.protection.sheet = True
        sheet.protection.autoFilter = False          # False = not locked: filtering stays possible
        sheet.protection.formatColumns = False

def build_workbook(store, paths, settings, progress, target):
    """Build Output.xlsm on local disk from the record of the run. All five sheets always
    exist; a sheet whose step has not run shows its header only."""
    import openpyxl
    layout = load_layout()
    rows = sheet_rows(store, paths, settings, progress)
    check_written_totals(rows, store)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet_layout in layout["sheets"]:
        sheet = workbook.create_sheet(sheet_layout["name"])
        write_sheet(sheet, sheet_layout, rows[sheet_layout["name"]], layout["colours"], settings, store)
    link_references(workbook, rows)
    workbook.code_name = "ThisWorkbook"
    for number, sheet in enumerate(workbook.worksheets, start=1):
        sheet.sheet_properties.codeName = "Sheet%d" % number
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as parts:                 # openpyxl takes a macro from a workbook it read:
        parts.writestr("xl/vbaProject.bin", vba_project(tuple(s.sheet_properties.codeName for s in workbook.worksheets)))
        parts.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        parts.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')
    workbook.vba_archive = zipfile.ZipFile(archive)              # the two empty parts say it declared nothing else
    workbook.properties.title = "the tool Output"
    workbook.properties.description = canonical_json(run_identity(store, paths))
    workbook.save(target)
    return rows


def link_references(workbook, rows):
    """A hyperlink on each cell that names chunks (LINK_COLUMNS): on Chunks_Model, Relevant Chunks in Methodology,
    Immediate Upstream and Immediate Downstream Model Chunk, and the count of a piece's flagged items; on
    Flagged_Items, the location of each item and the chunks of the methodology it rests on. Each leads to the first
    row it names, where the workbook's macro then shows only the rows named (WORKBOOK_MACRO); a count leads to the
    piece's first item. One link to a cell: Excel holds no more. Enforces: R2"""
    from openpyxl.styles import Font
    from openpyxl.worksheet.hyperlink import Hyperlink
    first_row = {}
    for name, key in LINK_KEYS.items():
        at = first_row.setdefault(name, {})
        for number, row in enumerate(rows.get(name) or (), start=2):
            at.setdefault(re.sub(r"^([CDM]-\d+)-\d+$", r"\1", str(row.get(key) or "")), number)
    column_of = {s["name"]: {c["field"]: n for n, c in enumerate(s["columns"], start=1)} for s in load_layout()["sheets"]}
    for (name, field), (target, tip) in LINK_COLUMNS.items():
        sheet = workbook[name]
        for number, row in enumerate(rows.get(name) or (), start=2):
            value = row.get(field)
            if field == "flagged_count":                    # a count leads to the piece's items, if it has any
                shown = int(str(value).split()[0]) if isinstance(value, int) or (isinstance(value, str) and value[:1].isdigit()) else 0
                first = row.get("unit_ref") if shown else None
            else:
                matched = REF_LIST.match(value) if isinstance(value, str) else None
                first = matched.group(1) if matched else None
            if not first or first not in first_row[target]:
                continue
            cell = sheet.cell(row=number, column=column_of[name][field])
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location="'%s'!A%d" % (target, first_row[target][first]), tooltip=tip)
            cell.font = Font(color=LINK_FONT, underline="single")


def file_sha256(path):
    """SHA-256 of a file's bytes."""
    with open(path, "rb") as handle:
        return sha256_bytes(handle.read())


def progress_text(store, waiting_message):
    """Where the run stands, in one or two plain sentences."""
    records = [record for record in store.read("step_records") if record.get("finished", True)]
    if not records:
        return "The run has been opened; no step has finished yet."
    last = records[-1]
    text = "Step %s (%s) is the last finished step." % (last["step_id"], last["name"])
    return text + (" " + waiting_message if waiting_message else "")

def rebuild_outputs(store, paths, settings, waiting_message):
    """Rebuild Output.xlsm (and the report once flagged items exist) on local disk and copy
    them whole into the project folder. The guard: a workbook there that differs from the last
    one the tool wrote is a reviewer's work in progress and is never overwritten before it has
    been read in. Afterwards the project folder holds the two files beside its three input folders. Enforces: R6, R12"""
    work = os.path.join(store.local_dir, "work")
    os.makedirs(work, exist_ok=True)
    progress = progress_text(store, waiting_message)
    manifest = (store.read("run_manifest") or [{}])[0]
    target = os.path.join(paths.outputs_dir, OUTPUT_FILE)
    guarded = (os.path.exists(target) and manifest.get("last_workbook_sha256")
               and file_sha256(target) not in (manifest["last_workbook_sha256"], manifest.get("last_upload_sha256")))
    local_workbook = os.path.join(work, OUTPUT_FILE)
    build_workbook(store, paths, settings, progress, local_workbook)
    if guarded:
        log_line(store, "Output.xlsm in the project folder was edited and not yet read in: left untouched")
    else:
        copy_whole(local_workbook, target)
        if manifest:
            update_manifest(store, {"last_workbook_sha256": file_sha256(target)})
    store.sync()
    return not guarded

# ---------------------------------------------------------------- verification, and the steps' allow-list
def contents_of(path):
    """Every byte string a file holds, looking INSIDE the compressed ones: the call log is gzip,
    and the workbook and the report are ZIP archives, so a token written into any of them would
    be invisible to a search of the raw bytes. Yields the raw bytes first, then each member of a
    ZIP and the decompressed stream of a gzip. Enforces: R8"""
    with open(path, "rb") as handle:
        raw = handle.read()
    yield raw
    try:
        if raw.startswith(b"\x1f\x8b"):
            yield gzip.decompress(raw)
        elif raw.startswith(b"PK\x03\x04"):
            import zipfile
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                for member in archive.namelist():
                    yield archive.read(member)
    except Exception:
        return                                       # a damaged archive is caught by the other lines of the verification

def verify_evidence_pack(paths, settings, live=None):
    """Works from the project's two files and its three input folders alone. Returns rows (what was checked,
    "Confirmed" or "Not confirmed", detail). Re-reading the inputs is reperformance: every
    content hash is computed again from the input files and compared with the record."""
    store, rows = AuditStore(paths.audit_dir, paths.audit_dir), []        # read the pack itself, not the scratch copy of this driver
    def line(what, good, detail=""):
        """good is True, False, or None for a check the run has not reached yet - which is not a failure
        and must not read as one. Found on a real run that stopped after step 06: the coverage line said
        Not confirmed, as if the pack had been tampered with."""
        rows.append((what, "Confirmed" if good else "Not reached yet" if good is None else "Not confirmed", detail))
    manifest = (store.read("run_manifest") or [{}])[0]
    changed = [e["file"] for e in manifest.get("inputs", []) if not os.path.exists(os.path.join(paths.inputs_dir, e["file"]))
               or file_sha256(os.path.join(paths.inputs_dir, e["file"])) != e["sha256"]]
    line("Input files have the recorded fingerprints", not changed, ", ".join(changed))
    installed, recorded = engine_file_hashes(), manifest.get("engine_files", {})
    other = sorted(name for name in set(installed) | set(recorded) if installed.get(name) != recorded.get(name))
    line("The engine files that produced this run are the ones installed here", not other, ", ".join(other[:5]))
    options = {"inputs": list_input_files(paths.inputs_dir)}
    context = StepContext(settings, options, lambda kind: [], None, paths.local_dir, lambda text: None)
    for kind, function in (("chunks_canon", read_methodology), ("chunks_doc", read_documentation),
                           ("model_units", read_package)):
        fresh = {to_plain(r)["ref"]: to_plain(r)["content_hash"] for r in function(context).records.get(kind, [])}
        recorded = {r["ref"]: r["content_hash"] for r in store.read(kind)}
        differing = sorted(ref for ref in set(fresh) | set(recorded) if fresh.get(ref) != recorded.get(ref))
        sheet = {"chunks_canon": "Chunks_Methodology", "chunks_doc": "Chunks_Documentation", "model_units": "Chunks_Model"}.get(kind, kind)
        line("Re-reading the inputs gives the recorded content hashes (%s)" % sheet, not differing, ", ".join(differing[:5]))
    written = {kind: len(store.read(kind)) for kind in ("chunks_canon", "chunks_doc", "model_units")}
    line("Every unit read is in the record", all(written.values()),
         ", ".join("%s: %d" % pair for pair in sorted(written.items())))
    identity = run_identity(store, paths)
    try:
        import openpyxl
        found = json.loads(openpyxl.load_workbook(os.path.join(paths.outputs_dir, OUTPUT_FILE), read_only=True).properties.description)
        line("The workbook carries this run's ids and fingerprints", all(found.get(k) == v for k, v in identity.items()))
    except Exception:
        line("The workbook carries this run's ids and fingerprints", False, "Output.xlsm could not be opened")
    secrets = [value.encode("utf-8") for value in (list(live.recent_tokens) if live else []) if value]
    leaked = []
    for name in (OUTPUT_FILE, AUDIT_FILE):                # the tool's own files; the input folders are the project's
        path = os.path.join(paths.outputs_dir, name)
        if os.path.exists(path):
            leaked += [name for data in contents_of(path) for value in secrets if value in data]
    line("No access token was written into Output.xlsm or Audit_Log.xlsx", not leaked, ", ".join(sorted(set(leaked))))
    return rows

def combine(ctx, *parts):
    """One step that carries out several: each part in order, everything they record kept, their counts
    and their messages side by side. A later part reads what the earlier ones have just recorded, as it
    would if each were still a step of its own - the records of a step reach the store only when the step
    ends. Enforces: R2"""
    records, counts, messages = {}, {}, []
    for part in parts:
        stored = ctx.read
        so_far = {kind: [to_plain(row) for row in rows] for kind, rows in records.items()}   # as the store would hand them back
        reading = dataclasses.replace(ctx, read=lambda kind, stored=stored, so_far=so_far: list(stored(kind)) + so_far.get(kind, []))
        found = part(reading)
        for kind, rows in (found.records or {}).items():
            records.setdefault(kind, []).extend(rows)
        for name, value in (found.counts or {}).items():     # two parts that count the same thing add up
            counts[name] = counts.get(name, 0) + value if isinstance(value, (int, float)) and isinstance(counts.get(name, 0), (int, float)) else value
        messages += found.messages or []
    return StepResult(records, counts, messages)

def read_inputs(ctx):
    """Step 02, read-inputs: the methodology, the model documentation and the model package, each read
    into units, in that order. Enforces: R2, R7"""
    return combine(ctx, read_methodology, read_documentation, read_package)

STEP_FUNCTIONS = {        # the function that carries out each step of PIPELINE. Enforces: R11
    "read_methodology": read_methodology,
    "read_documentation": read_documentation,
    "read_package": read_package,
    "interpret_code": interpret_code,
    "search_methodology": search_methodology,
    "prepare_run": prepare_run,
    "read_inputs": read_inputs,
    "link_chunks": link_chunks}

# ---------------------------------------------------------------- the notebook: four cells, each one call
# Everything the notebook does is here, so that it holds no code of its own but the organisation's chat():
# cell 1 is setup(dbutils), cell 2 check_chat(chat), cell 3 review(), cell 4 verify(). The session - the
# widgets, the chat() that answered, the run being worked on - is kept in NOTEBOOK, not in the notebook.
REQUIRED_PACKAGES = ("yaml", "openpyxl", "numpy", "rdata")   # what the engine imports; installed only if missing
WIDGETS = (("llm_endpoint", "", "01 LLM endpoint"), ("llm_token", "", "02 LLM token"),
           ("project_name", "", "03 Project Name (can be a model ID)"))
OLD_WIDGETS = ("llm_user_id", "reviewer_role", "run", "projects_dir", "scratch_dir", "concept_subject", "flowr_archive",
               "reviewer_id", "jfrog_index_url", "concurrency_limit", "token_cap", "model_id", "project")
PYPI = "https://pypi.org/simple/"          # where every package comes from: named here, so no pip setting of the cluster redirects it
NOTEBOOK = {"dbutils": None, "home": "", "projects": "", "user": "", "live": None, "chat": None, "paths": None, "result": None}

def databricks_user(dbutils):
    """Who runs the notebook, as Databricks knows them: the notebook context's user name; on a cluster that keeps
    that context from Python, Spark's current_user(); outside Databricks, the system user."""
    try:
        return dbutils.notebook.entry_point.getDbutils().notebook().getContext().userName().get()
    except Exception:
        pass
    try:
        from pyspark.sql import SparkSession
        session = SparkSession.getActiveSession()
        if session is not None:
            return session.sql("SELECT current_user()").first()[0]
    except Exception:
        pass
    return getpass.getuser()

def live(name):
    """The endpoint and token, read from the widgets at the moment chat() calls - a token pasted into widget 02
    while a run works is used by its next call - and the user id (asked for as "reviewer_id"), from Databricks. Enforces: R8"""
    session = NOTEBOOK["live"] = NOTEBOOK["live"] or LiveValues()
    if getattr(CHAT_WORKER, "active", False):       # a worker of step 04: the values its step last read
        return session.get(name)
    try:
        widgets = NOTEBOOK["dbutils"].widgets
        session.update(widgets.get("llm_endpoint"), widgets.get("llm_token"), NOTEBOOK["user"])
    except Exception:
        pass                                        # no notebook, or a widget read failed: the last values stand
    return session.get(name)

def notebook_folder(dbutils):
    try:
        path = dbutils.notebook.entry_point.getDbutils().notebook().getContext().notebookPath().get()
        return os.path.dirname(path if path.startswith("/Workspace") else "/Workspace" + path)
    except Exception:
        return os.getcwd()

def setup(dbutils, home=None, projects=None):
    """Cell 1: the widgets, the packages the engine needs (installed only if one is missing), flowR, and
    what to do next. Safe to run any number of times; run it again after anything restarts Python."""
    import importlib.util
    widgets = dbutils.widgets
    carried = {}                                      # a project name typed into an earlier notebook's 03 Model ID stays
    try:
        carried["project_name"] = widgets.get("model_id")
    except Exception:
        pass
    for name, default, label in WIDGETS:              # made before anything is installed: a bare cluster works
        try:
            widgets.get(name)
        except Exception:
            widgets.text(name, carried.get(name, default), label)
    for name in OLD_WIDGETS:                          # widgets of an earlier notebook, no longer used
        try:
            widgets.remove(name)
        except Exception:
            pass
    home = home or notebook_folder(dbutils)
    NOTEBOOK.update(dbutils=dbutils, home=home, projects=projects or os.path.join(home, "Projects"), user=databricks_user(dbutils))
    missing = [name for name in REQUIRED_PACKAGES if importlib.util.find_spec(name) is None]
    if missing:
        install(missing, dbutils, os.path.join(os.path.dirname(os.path.abspath(__file__)), "requirements.txt"))
        return
    print("Folder:", home, "| Python", sys.version.split()[0])
    for name in REQUIRED_PACKAGES + ("pdfplumber", "pypdf"):
        found = importlib.util.find_spec(name)
        print("  %-11s %s" % (name, "installed" if found else "not installed" + ("" if name in ("pdfplumber", "pypdf") else " - run this cell again")))
    token = live("llm_token")
    print("Endpoint set:", bool(live("llm_endpoint")), "| token:", ("%d characters" % len(token)) if token else "not set")
    staged = os.path.join(home, FLOWR_URL.format(FLOWR_VERSION).rsplit("/", 1)[-1])
    try:                                              # flowR reads the R code: fetched once, checked against its SHA-256
        print("flowR %s ready in %s" % (FLOWR_VERSION, flowr_ready(staged if os.path.exists(staged) else "")))
    except Exception as problem:
        print("flowR IS NOT READY: %s" % problem)
        if problem.__class__.__module__.startswith("urllib"):   # only a failed download is helped by putting it here
            print("The cluster could not download it. Download %s on an approved machine and put it in %s, next to this"
                  " notebook." % (FLOWR_URL.format(FLOWR_VERSION), home))
    print("\nTHE ENGINE IS READY. What happens next:")
    print("  Cell 2  paste your organisation's chat(), check it answers, and see where to put your files.")
    print("  Cell 3  read the inputs and run the review; it prints what each step did.")
    print("  Cell 4  check the project's Output.xlsm and Audit_Log.xlsx against their own record.")
    print("Widgets 01 and 02 carry the endpoint and the token; 03 the project's name, which names its folder in Projects.")
    print("Your user id is taken from Databricks: %s. Paste a fresh token into widget 02 at any time - chat() reads it" % NOTEBOOK["user"])
    print("at the moment it calls.")
    print("Next: cell 2.")

def pip_said(stderr):
    """What pip said, without any credentials of an index URL, and what its most common failure means."""
    text = re.sub(r"//[^/@\s]+@", "//...@", stderr or "")[-1500:]
    missing = re.findall(r"satisfies the requirement ([\w.\-\[\]]+)", stderr or "")
    if "from versions: none" in (stderr or ""):
        text += ("\nPyPI offered no version of %s that fits this cluster's Python, or the cluster cannot reach %s: its network "
                 "has to let the cluster reach PyPI." % (missing[0] if missing else "that package", PYPI))
    elif "ResolutionImpossible" in (stderr or "") or "conflicting dependencies" in (stderr or ""):
        text += "\nThe packages asked for cannot be installed together; the lines above say which."
    return text

def install(missing, dbutils, requirements):
    """Install what the engine needs from PyPI, the runtime's own packages pinned as they are, and restart Python
    only if those still import together afterwards."""
    import importlib.metadata
    hide = lambda text: re.sub(r"//[^/@\s]+@", "//...@", text or "")    # no credentials of an index URL are shown
    pins = []
    for name in ("numpy", "pandas", "pyarrow", "scipy"):
        try:
            pins.append("%s==%s" % (name, importlib.metadata.version(name)))
        except importlib.metadata.PackageNotFoundError:
            pass
    constraints = os.path.join(tempfile.mkdtemp(prefix="verifier_"), "constraints.txt")
    with open(constraints, "w") as handle:
        handle.write("\n".join(pins) + "\n")
    command = [sys.executable, "-m", "pip", "install", "-r", requirements, "-c", constraints, "--index-url", PYPI]
    print("Installing", ", ".join(missing), "| kept as the runtime has them:", ", ".join(pins) or "none found")
    done = subprocess.run(command, capture_output=True, text=True)
    if done.returncode:
        print("The install did not finish. What pip said:\n" + pip_said(done.stderr))
        print("Nothing the runtime needs was changed.")
        return
    with open(requirements, encoding="utf-8") as handle:
        wanted = handle.read().split("# --- optional ---")
    if len(wanted) > 1:
        spare = os.path.join(tempfile.gettempdir(), "optional.txt")
        with open(spare, "w", encoding="utf-8") as handle:
            handle.write(wanted[1])
        extra = subprocess.run(command[:4] + ["-r", spare] + command[6:], capture_output=True, text=True)
        print("Optional packages (the PDF readers):", "installed." if not extra.returncode else "not installed; a PDF will say it could not be read.")
    probe = subprocess.run([sys.executable, "-c", "import numpy, pandas, pyarrow"], capture_output=True, text=True)
    if probe.returncode:
        print("STOPPED BEFORE RESTARTING PYTHON: the runtime's own packages no longer import together (%s)." % hide((probe.stderr or "").strip())[-300:])
        print("Restarting now would crash this notebook session. Detach this notebook from the cluster and attach it again to undo the install.")
        return
    still = subprocess.run([sys.executable, "-c", "import " + ", ".join(missing)], capture_output=True, text=True)
    if still.returncode:
        print("STOPPED BEFORE RESTARTING PYTHON: %s is still missing after the install, so restarting would only bring "
              "this cell back here." % ", ".join(missing))
        print("What Python said:\n" + hide(still.stderr)[-600:])
        print("Check that engine/requirements.txt names it, and that PyPI carries it for this cluster's Python.")
        return
    print("Installed. Restarting Python; then run this cell once more.")
    dbutils.library.restartPython()

def notebook_settings():
    return make_settings({"reviewer_id": NOTEBOOK["user"]})

def check_chat(chat):
    """Cell 2: ask the organisation's chat() one question and, once it answers, keep it for cell 3's steps 04 and 05,
    make the project's three input folders and say what belongs in each."""
    if NOTEBOOK["dbutils"] is None:
        print("Run cell 1 first.")
        return
    NOTEBOOK["chat"] = chat
    try:
        reply = chat("Reply with the single word OK.", "Reply with the single word OK.")["answer"]
        print("chat() answered:", str(reply)[:60])
        print("Cell 3 sends each piece of the model's code to this chat(), for the column Code Interpretation (by LLM);")
        print("then the methodology, batch by batch, with each piece, for the columns Relevant Chunks in Methodology")
        print("(searched by LLM), its count of flagged items, and the sheet Flagged_Items, one row per item.")
    except Exception as problem:
        print("chat() did not answer (%s: %s). Check widgets 01 and 02 - and that the gateway knows your Databricks user id, %s -"
              " then run this cell again." % (type(problem).__name__, problem, NOTEBOOK["user"]))
        return
    widgets = NOTEBOOK["dbutils"].widgets
    project_dir, missing, said = setup_project(NOTEBOOK["projects"], widgets.get("project_name"))
    for line in said:
        print("\n" + line)
    print("\nPUT YOUR FILES IN THESE THREE FOLDERS, in the project's folder beside Output.xlsm, then run cell 3:")
    for _, folder, note in INPUT_FOLDERS:
        print("  %s\n      %s" % (os.path.join(project_dir, folder), note))
    print("\nOne methodology file, the model package as it ships (.tar.gz, .zip or the unpacked folder), and the")
    print("model documentation. A folder with nothing in it stops the run and says so, rather than reading around it.")
    print("\nBefore cell 3: " + " ".join(missing) if missing else "\nAll three folders have files in them. Next: cell 3.")

def open_current():
    """The run of the project the widgets name, carried on or started anew as open_run decides - asked each time
    cell 3 runs, so that an input changed meanwhile is noticed - and said in plain words when it is not the run
    this session worked on already. None, with what to do, while an input folder is empty or while a new run
    would replace an Output.xlsm a person has changed."""
    widgets = NOTEBOOK["dbutils"].widgets
    project = widgets.get("project_name")
    _, missing, said = setup_project(NOTEBOOK["projects"], project)
    for line in said:
        print(line)
    if missing:
        print("\n".join(missing))
        print("Put the files in, then run cell 3.")
        return None
    try:
        paths = open_run(NOTEBOOK["projects"], project, settings=notebook_settings())
    except OutputsEdited as problem:
        print(problem)
        return None
    before = NOTEBOOK["paths"]
    if before is None or (before.project_dir, before.run_id) != (paths.project_dir, paths.run_id):
        print(paths.opened)
    NOTEBOOK["paths"] = paths
    return paths

def review():
    """Cell 3: read the inputs, link the units of the package and put them to the organisation's model, in this
    cell; then say what each step did. A step already finished is never repeated."""
    if NOTEBOOK["dbutils"] is None:
        print("Run cell 1 first.")
        return
    paths = open_current()
    if paths is None:
        return
    settings = notebook_settings()
    result = NOTEBOOK["result"] = run_pipeline(paths, settings)
    print(result["message"])
    store = open_store(paths, settings)
    print("\nWhat each step did:")
    latest = {}
    for record in store.read("step_records"):                 # a step tried again says what its latest attempt did;
        latest[record["step_id"]] = record                     # the audit log keeps every attempt
    for record in (latest[step_id] for step_id in sorted(latest)):
        print("  step %s %-14s %s" % (record["step_id"], record["name"], ", ".join("%s: %s" % item for item in sorted((record["counts"] or {}).items()))))
        for message in record["messages"]:
            print("      " + message)
    print("\nProject folder:", paths.project_dir)
    print("Open Output.xlsm there, beside the three input folders: the three Chunks sheets show everything that was read, and Chunks_Model also what the")
    print("organisation's model says of each piece, the chunks of the methodology it found for it, and how many items it")
    print("flagged; Flagged_Items lists every item, one to a row, with a column for your decision and one for your notes.")
    print("A click on a reference shows only the rows it names, once Excel lets the workbook's macro run;")
    print("if it blocks it, unblock the file first (the manual, section 7: Letting the macro run).")
    print("Then run cell 4 to check the project's two files against their own record.")

def verify():
    """Cell 4: check the project's two files against their own record."""
    paths = NOTEBOOK["paths"]
    if paths is None:
        print("Cell 3 has not read the inputs yet. Run cell 3 first.")
        return
    print("Verifying the evidence pack:")
    for what, verdict, detail in verify_evidence_pack(paths, notebook_settings(), live=NOTEBOOK["live"] or LiveValues()):
        print("  %-62s %-16s %s" % (what, verdict, detail))
    print("\nProject folder:", paths.project_dir, "- Output.xlsm is the deliverable; Audit_Log.xlsx beside it is the")
    print("record of the run: every step, every record, and every exchange with the model.")
