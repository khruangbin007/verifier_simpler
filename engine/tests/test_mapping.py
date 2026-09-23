"""Tests of verifier3_mapping: the ledger, the deterministic search signals, questions and validators."""
import math
import unittest

import yaml

import os

import helpers
import verifier
core = review = runner = verifier   # the engine is one module now
SETTINGS = runner.make_settings()
PROVENANCE = core.Provenance("Run_X", "05", "build-graph", "0.0.1")


def edge(source, target, kind="contains"):
    return core.Edge(source, target, kind, core.HOW_PARSED, PROVENANCE)


class Ledger(unittest.TestCase):
    def ledger(self):
        records = [review.node_record(ref, "Function", "model") for ref in ("M-0002", "M-0001", "M-0003")]
        return review.ledger_records([], records + [edge("M-0001", "M-0002"), edge("M-0001", "M-0003", "calls")])

    def test_the_same_content_gives_the_same_graph_version_whatever_the_order(self):
        first = self.ledger()
        records = [edge("M-0001", "M-0003", "calls"), edge("M-0001", "M-0002")] + [review.node_record(ref, "Function", "model") for ref in ("M-0003", "M-0001", "M-0002")]
        self.assertEqual(review.graph_version_id(first), review.graph_version_id(review.ledger_records([], records)))

    def test_an_edited_a_removed_and_a_reordered_record_are_each_detected(self):
        good = self.ledger()
        self.assertTrue(review.verify_ledger(good)[0])
        edited = [dict(r) for r in good]
        edited[3]["target"] = "M-0003"
        removed = good[:2] + good[3:]
        reordered = good[:3] + [good[4], good[3]]
        for tampered in (edited, removed, reordered):
            self.assertFalse(review.verify_ledger(tampered)[0])

    def test_a_later_record_is_appended_and_never_replaces_an_earlier_one(self):
        first = self.ledger()
        more = review.ledger_records(first, [edge("M-0002", "M-0003", "calls")])
        self.assertTrue(review.verify_ledger(first + more)[0])
        self.assertEqual(more[0]["prev_hash"], first[-1]["record_hash"])
        self.assertFalse([name for name in dir(review) if name.startswith(("delete_", "remove_", "update_edge", "edit_"))])

    def test_paths_are_found_over_typed_edges(self):
        graph = review.load_graph(self.ledger())
        path = review.find_path(graph, "M-0002", lambda ref: ref == "M-0003")
        self.assertEqual([reached for _, reached in path], ["M-0001", "M-0003"])
        self.assertIsNone(review.find_path(graph, "M-0002", lambda ref: ref == "M-0003", allowed_kinds=("contains",)))


class Words(unittest.TestCase):
    def test_identifiers_are_split_and_stemmed(self):
        self.assertEqual(review.split_words("computeRiskWeights cond_pd applied.floors", {"of"}), ["comput", "risk", "weight", "cond", "pd", "apply", "floor"])
        self.assertEqual(review.stem("applies"), review.stem("applied"))

    def test_bm25_by_hand(self):
        index = review.build_index({"A": {"fields": {"body": ["floor", "rate"]}}, "B": {"fields": {"body": ["rate", "rate", "cap"]}}})
        scores = review.bm25_scores(index, {"floor": 1.0}, 1.2, 0.75)
        idf = math.log(1 + (2 - 1 + 0.5) / (1 + 0.5))
        expected = idf * 1 * 2.2 / (1 + 1.2 * (1 - 0.75 + 0.75 * 2 / 2.5))
        self.assertEqual(list(scores), ["A"])
        self.assertAlmostEqual(scores["A"][0], expected, places=12)

    def test_field_weights_count_a_name_more_than_the_body(self):
        index = review.build_index({"A": {"fields": {"name": ["floor"], "body": []}}, "B": {"fields": {"name": [], "body": ["floor"]}}})
        scores = review.bm25_scores(index, {"floor": 1.0}, 1.2, 0.75)
        self.assertGreater(scores["A"][0], scores["B"][0])

    def test_bridge_entries_are_harvested_from_the_inputs_own_definitions(self):
        patterns = {"definition_verbs": ["denotes", "represents"]}
        text = "The loss given default (LGD) applies, see (see Table 3), where R denotes the asset correlation. Let K be the capital requirement."
        found = {(term, phrase) for term, phrase, _, _ in review.harvest_text(text, "C-0001", patterns)}
        self.assertEqual(found, {("LGD", "loss given default"), ("R", "asset correlation"), ("K", "capital requirement")})


class Search(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.paths, cls.settings, _ = helpers.run_sample("F_capital", stop_after="07c")     # the search runs after the concepts
        cls.store = runner.open_store(cls.paths, cls.settings)

    def test_every_searched_unit_has_a_search_record_with_words_an_analyst_can_read(self):
        units = [u for u in self.store.read("model_units") if review.is_searched(u, self.settings)]
        records = {(r["unit_ref"], r["target_corner"]) for r in self.store.read("search_records")}
        for unit in units:
            self.assertIn((unit["ref"], "canon"), records)
        cond_pd = next(u["ref"] for u in self.store.read("model_units") if u["kind"] == "Function" and u["name"] == "cond_pd")
        text = next(r["searched_text"] for r in self.store.read("search_records") if r["unit_ref"] == cond_pd and r["target_corner"] == "canon")
        self.assertTrue(text.startswith("Searched the methodology for: cond, pd"))

    def test_the_gold_passage_is_proposed_and_every_proposal_gives_its_reason(self):
        cond_pd = next(u["ref"] for u in self.store.read("model_units") if u["kind"] == "Function" and u["name"] == "cond_pd")
        candidates = [c for c in self.store.read("candidates") if c["unit_ref"] == cond_pd and c["target_corner"] == "canon"]
        self.assertIn("C-0012", [c["target_ref"] for c in candidates][:5])
        self.assertTrue(all(c["reason"].startswith("Proposed") for c in candidates))
        self.assertLessEqual(len(candidates), self.settings["k_candidates"])

    def test_an_anchor_that_too_many_units_mention_is_dropped(self):
        mentions = {("symbol", "x"): ["U%d" % n for n in range(40)], ("number", "0.999"): ["U1", "U2"], ("number", "7"): ["U3"]}
        weights = review.anchor_weights(mentions, dict(self.settings, anchor_max_share=0.10))
        self.assertEqual(list(weights), [("number", "0.999")])

    def test_the_walk_agrees_with_an_independent_personalised_pagerank(self):
        try:
            import networkx
        except ImportError:
            self.skipTest("networkx is only present in the build environment")
        mentions = {("number", "0.999"): ["M-1", "C-1", "C-4"], ("symbol", "rho"): ["M-1", "C-1", "C-2"], ("term", "floor"): ["M-2", "C-3", "C-2"], ("number", "0.03"): ["M-2", "C-3"]}
        weights = review.anchor_weights(mentions, dict(self.settings, anchor_max_share=1.0))
        ours = review.restart_walk(mentions, weights, ["M-1"], dict(self.settings, walk_rounds=200))["M-1"]
        graph = networkx.Graph()
        for anchor, refs in mentions.items():
            for ref in refs:
                graph.add_edge(ref, "anchor:%s:%s" % anchor, weight=weights[anchor])
        theirs = networkx.pagerank(graph, alpha=1 - self.settings["walk_restart"], personalization={"M-1": 1.0}, weight="weight", tol=1e-12, max_iter=1000)
        for ref in ("C-1", "C-2", "C-3", "C-4"):
            self.assertAlmostEqual(ours.get(ref, 0.0), theirs[ref], places=6)
        self.assertEqual(sorted(ours, key=lambda ref: -ours[ref])[0], "C-1")

    def test_fusion_is_by_rank_only_and_ties_break_by_reference(self):
        fused = review.fuse({"fields": ["C-2", "C-1"], "anchors": ["C-1", "C-2"]}, self.settings)
        self.assertEqual([ref for ref, _, _ in fused], ["C-1", "C-2"])
        self.assertEqual(fused[0][1], fused[1][1])

    def test_signatures_reorder_and_never_propose_alone(self):
        fused = review.fuse({"fields": ["C-1"], "signatures": ["C-9", "C-1"]}, self.settings)
        self.assertEqual([ref for ref, _, _ in fused], ["C-1"])


class QuestionsAndValidators(unittest.TestCase):
    def question(self):
        prompt = core.load_prompt("judge-unit-to-canon")
        passages = [("C-0009", "3.1.1 Floor, paragraph 1", "The probability of default is never taken below 0.03%."),
                    ("C-0008", "3.1, paragraph 1", "The probability of default is estimated from internal ratings and is reviewed every year."),
                    ("C-0099", "9 Other, paragraph 1", "All parcels are weighed at the counter before they are priced.")]
        return review.assemble_question("judge-unit-to-canon", "M-0021", [("UNIT (function floor_pd)", "floor_pd <- function(pd, floor = 0.0003) {\n  pmax(pd, floor)\n}")],
                                         passages, prompt, SETTINGS, planted=["C-0099"], more={"target_corner": "canon"})

    def test_letters_come_from_a_hash_not_from_the_score_and_the_question_id_is_the_prompt_hash(self):
        first, second = self.question(), self.question()
        self.assertEqual(first["question_id"], second["question_id"])
        self.assertEqual(first["letters"], second["letters"])
        self.assertEqual(sorted(first["letters"].values()), ["C-0008", "C-0009", "C-0099"])
        self.assertEqual(first["planted"], [letter for letter, ref in first["letters"].items() if ref == "C-0099"])
        self.assertNotIn("C-0009", first["main_prompt"])

    def test_the_bad_answer_corpus(self):
        question = self.question()
        letter = {ref: letter for letter, ref in question["letters"].items()}
        with open(helpers.os.path.join(helpers.TESTS_DIR, "bad_answers", "judge_answers.yaml"), encoding="utf-8") as handle:
            corpus = yaml.safe_load(handle)["answers"]
        self.assertGreaterEqual(len(corpus), 20)
        for case in corpus:
            text = case["text"].replace('"GOOD"', '"%s"' % letter["C-0009"]).replace('"PLANTED"', '"%s"' % letter["C-0099"])
            outcome, answer = review.validate_answer(question, text)
            expected = "accepted" if case["expected"] == "accepted" else "rejected: " + core.REJECTION_REASONS[case["expected"]]
            self.assertEqual(outcome, expected, case["name"])
            self.assertEqual(answer is not None, expected == "accepted", case["name"])

    def test_decoys_share_no_anchor_with_the_unit_and_are_chosen_by_hash(self):
        pool = {"C-1": {"text": "x" * 50}, "C-2": {"text": "y" * 50}, "C-3": {"text": "z" * 50}, "C-4": {"text": "short"}}
        anchors = {"M-1": [("number", "0.5")], "C-1": [("number", "0.5")], "C-2": [("number", "9")], "C-3": []}
        first = review.choose_decoys("M-1", ["C-9"], set(), pool, anchors, 1)
        self.assertEqual(first, review.choose_decoys("M-1", ["C-9"], set(), pool, anchors, 1))
        self.assertIn(first[0], ("C-2", "C-3"))

    def test_text_written_by_the_model_is_filtered_before_it_is_shown(self):
        self.assertEqual(review.shown_ai_text("This is a " + "crit" + "ical gap."), core.AI_WORDING_NOT_SHOWN)
        self.assertEqual(review.shown_ai_text("The passage states another rate."), "The passage states another rate.")

    def test_a_rejected_answer_never_becomes_a_link(self):
        paths, settings, _ = helpers.run_sample("A_minimal", chat=lambda SystemPrompt, MainPrompt: {"answer": "I cannot answer in JSON."}, stop_after="08")
        store = runner.open_store(paths, settings)
        links = [r for r in store.read("graph_ledger") if r["record_type"] == "edge" and r["kind"] == "corresponds"]
        self.assertEqual(links, [])
        self.assertTrue(store.read("judgement_problems"))



class Concepts(unittest.TestCase):
    """Concepts, model first: the model's names and how its own roxygen and help pages describe them are
    the basis; the documents are searched for those and nothing else; what a unit shows is the words
    it writes, exactly; and the model's guesses of synonyms and acronyms are kept apart, on the
    Concepts sheet only."""

    MODEL = [{"ref": "M-0001", "kind": "Function", "name": "capital_k", "text": "capital_k <- function(pd, sacp) pd * sacp",
              "code": {"formals": [["pd", ""], ["sacp", ""]]}},
             {"ref": "M-0002", "kind": "Roxygen block", "name": "capital_k", "text": "#' @param pd probability of default",
              "roxygen": {"documents_name": "capital_k", "tags": [{"tag": "param", "name": "pd", "text": "probability of default"}]}}]

    @classmethod
    def setUpClass(cls):
        cls.paths, cls.settings, _ = helpers.run_sample("F_capital")
        cls.store = runner.open_store(cls.paths, cls.settings)
        cls.concepts, cls.per_unit = review.latest_concepts(cls.store.read)

    def registry(self, canon_text, guessed=()):
        sources = [({"ref": "C-0001", "text": canon_text, "heading_chain": []}, "canon")] + [(unit, "model") for unit in self.MODEL]
        return review.concept_registry(sources, guessed)

    def test_every_concept_is_the_models_and_says_how_the_model_names_it(self):
        self.assertTrue(self.concepts)
        for concept in self.concepts:
            self.assertTrue(concept["identifiers"], "a concept with no name in the code is not a concept")
        pd = [c for c in self.concepts if "pd" in c["identifiers"]][0]
        self.assertIn("probability of default", pd["described"], "the roxygen @param text, word for word")

    def test_a_term_of_the_documents_with_no_likely_equivalent_in_the_model_is_never_extracted(self):
        _, per_unit = self.registry("The Enterprise Risk Profile and the probability of default decide it.")
        self.assertEqual(per_unit[0]["as_written"], ["probability of default"], "not 'Enterprise Risk Profile'")

    def test_what_a_unit_shows_is_its_own_words_exactly(self):
        _, per_unit = self.registry("Both the PD and the Probability-Of-Default are shown.")
        self.assertEqual(per_unit[0]["as_written"], ["PD", "Probability-Of-Default"], "the source's capitals and hyphens, not the model's")

    def test_a_guess_is_kept_as_the_unit_writes_it_and_apart_from_what_code_proved(self):
        concepts, per_unit = self.registry("The stand-alone credit profile applies.",
                                           guessed=[("C-0001", "stand-alone credit profile", "sacp")])
        sacp = [c for c in concepts if c["key"] == "sacp"][0]
        self.assertEqual(sacp["found"]["canon"], {}, "code did not prove it")
        self.assertEqual(sacp["guessed"]["canon"], {"stand-alone credit profile": ["C-0001"]}, "the model guessed it")
        self.assertEqual(per_unit[0]["as_written"], ["stand-alone credit profile"])

    def test_a_guess_must_be_the_passages_words_and_a_listed_concept(self):
        question = {"question_type": "match-concepts", "unit_texts": {"C-0001": "The stand-alone credit profile applies."},
                    "concept_ids": ["K-0001", "K-0002"]}
        review.validate_concepts(question, {"units": {"C-0001": [{"concept": "K-0002", "words": "Stand-Alone Credit Profile"}]}})
        self.assertEqual(review.verbatim("Stand-Alone Credit Profile", question["unit_texts"]["C-0001"]), "stand-alone credit profile",
                         "what is kept is the passage's words, not the model's rendering of them")
        for bad in ({"units": {"C-0001": [{"concept": "K-0009", "words": "credit profile"}]}},
                    {"units": {"C-0001": [{"concept": "K-0002", "words": "standalone profile"}]}},
                    {"units": {"C-0009": []}}, {"units": []}):
            with self.assertRaises(review.Rejected):
                review.validate_concepts(question, bad)

    def test_the_chunks_sheets_show_only_the_units_words_under_their_header(self):
        import openpyxl
        book = openpyxl.load_workbook(os.path.join(self.paths.run_dir, "Output.xlsx"), read_only=True)
        header = "Extracted concepts with candidate equivalent in model"
        rows = list(book["Chunks_Canon"].iter_rows(values_only=True))
        for name in ("Chunks_Canon", "Chunks_Doc", "Chunks_Model"):
            self.assertIn(header, next(book[name].iter_rows(max_row=1, values_only=True)), name)
        column, text_column = rows[0].index(header), rows[0].index("Text")
        for row in rows[1:]:
            for words in filter(None, str(row[column] or "").split("; ")):
                self.assertIn(words, row[text_column], "every extracted concept is the unit's own words")
        self.assertEqual(book["Concepts"].max_row - 1, len(self.concepts), "one row per model concept")

    def test_a_shared_concept_puts_a_passage_on_the_shortlist_with_its_reason(self):
        shared = [c for c in self.store.read("candidates") if "concepts" in (c.get("signals") or {})]
        self.assertTrue(shared)
        self.assertTrue(all("shares the concept" in c["reason"] for c in shared))


class InterpretingTheCode(unittest.TestCase):
    """Step 07a: what each piece of code does, in plain words, for the column LLM Interpretation."""

    @classmethod
    def setUpClass(cls):
        import standin_chat
        cls.paths, cls.settings, cls.outcome = helpers.run_sample("A_minimal", chat=standin_chat.make_chat(misbehave=False), stop_after="07a")
        cls.store = runner.open_store(cls.paths, cls.settings)
        cls.units = {u["ref"]: u for u in cls.store.read("model_units")}
        cls.records = cls.store.read("interpretations")
        cls.calls = [c for c in cls.store.read_calls() if c["question_type"] == "interpret-code"]

    def test_every_piece_of_code_is_asked_about_once_and_nothing_else_is(self):
        asked = sorted(r["unit_ref"] for r in self.records)
        code = sorted(ref for ref, u in self.units.items() if u["kind"] in review.INTERPRETED_KINDS)
        self.assertEqual(asked, code)
        self.assertEqual(len(asked), len(set(asked)))
        self.assertFalse([r for r in self.records if self.units[r["unit_ref"]]["kind"] in ("Roxygen block", "Help page", "Parameter table")])

    def test_a_function_is_shown_whole_with_its_statements_and_where_it_sits(self):
        """A statement is no longer a row of its own: the model reads it as part of its function."""
        function = next(ref for ref, u in self.units.items() if u["kind"] == "Function" and u["name"] == "parcel_price")
        prompt = next(c["main_prompt"] for c in self.calls if c["unit_ref"] == function)
        piece, about = prompt.split("WHERE IT SITS IN THE PACKAGE")
        self.assertIn("price <- base + rate * cw", piece)
        self.assertIn("The whole package", about)
        self.assertIn('base <- zone_rates[zone_rates$zone == zone, "base_rate"]', piece, "the whole function, in its one row")
        self.assertIn("Base rate plus rate per kilogram times chargeable weight", piece, "its roxygen block, in the same row")
        self.assertIn("Within this package it calls: volume_discount.", about)
        self.assertIn("It reads the stored data: zone_rates.", about)
        self.assertIn("Package parcelcost 1.2.0: Parcel Pricing", about)
        self.assertIn("R/weights.R: dim_weight(l, w, h, divisor) - Dimensional weight", about, "every file of the package, not only its own")

    def test_a_function_is_told_what_calls_it(self):
        called = next(ref for ref, u in self.units.items() if u["kind"] == "Function" and u["name"] == "volume_discount")
        prompt = next(c["main_prompt"] for c in self.calls if c["unit_ref"] == called)
        self.assertIn("It is called by: parcel_price.", prompt)

    def test_an_answer_that_quotes_code_that_is_not_there_is_refused(self):
        question = {"question_type": "interpret-code", "code_text": "  price <- base + rate * cw", "strip_patterns": []}
        good = '{"interpretation": "It adds the base rate to the rate per kilogram times the weight.", "quote_from_unit": "base + rate * cw"}'
        self.assertEqual(review.validate_answer(question, good)[0], "accepted")
        invented = '{"interpretation": "It adds the base rate to the rate per kilogram times the weight.", "quote_from_unit": "price <- base * 2"}'
        self.assertEqual(review.validate_answer(question, invented), ("rejected: " + core.REJECTION_REASONS[2], None))
        empty = '{"interpretation": "", "quote_from_unit": "base + rate * cw"}'
        self.assertEqual(review.validate_answer(question, empty)[0], "rejected: " + core.REJECTION_REASONS[0])
        rambling = '{"interpretation": "%s", "quote_from_unit": "base + rate * cw"}' % ("word " * 200)
        self.assertEqual(review.validate_answer(question, rambling)[0], "rejected: " + core.REJECTION_REASONS[0])

    def test_the_column_shows_the_models_words_as_a_quotation_and_a_refusal_as_a_plain_note(self):
        units = [self.units[r["unit_ref"]] for r in self.records[:2]]
        told = [dict(self.records[0], interpretation='It adds \u201cbase\u201d to the rest.', note=""),
                dict(self.records[1], interpretation="", note="The AI's answer could not be used: it quoted words that are not in the text.")]
        rows = runner.rows_model_units(units, told)
        self.assertEqual(rows[0]["llm_interpretation"], '\u201cIt adds "base" to the rest.\u201d')
        self.assertEqual(rows[1]["llm_interpretation"], "The AI's answer could not be used: it quoted words that are not in the text.")
        self.assertEqual(runner.rows_model_units(units)[0]["llm_interpretation"], "", "before the step has run the column is empty")

    def test_a_model_that_speaks_of_an_error_in_the_code_does_not_get_its_cell_withheld(self):
        """Code that calls stop() will be described with the very word the tool never uses of its own
        results. The words are the model's, shown as a quotation, so the wording gate lets them by."""
        row = runner.rows_model_units([self.units[self.records[0]["unit_ref"]]],
                                   [dict(self.records[0], interpretation="It stops with an error when the zone is unknown.", note="")])[0]
        self.assertEqual(runner.plain_cell(row["llm_interpretation"], False, self.store), row["llm_interpretation"])
        self.assertEqual(runner.plain_cell("It stops with an error.", False, self.store), runner.CELL_WITHHELD, "the tool's own words are still held to the rule")

    def test_an_interpretation_is_outside_the_accounting(self):
        """It is an aid to reading: no status, no flagged item, no part in the coverage identity."""
        import standin_chat
        with_it = helpers.run_sample("A_minimal", chat=standin_chat.make_chat(misbehave=False))
        without = helpers.run_sample("A_minimal", chat=standin_chat.make_chat(misbehave=False), settings={"interpret_code": False})
        kept = lambda item: {k: v for k, v in item.items() if k not in ("provenance", "created_at", "item_id")}
        first, second = (runner.open_store(p, s) for p, s, _ in (with_it, without))
        self.assertEqual([kept(i) for i in first.read("flagged_items")], [kept(i) for i in second.read("flagged_items")])
        self.assertEqual(helpers.without_times(first.read("coverage")), helpers.without_times(second.read("coverage")))
        self.assertEqual(second.read("interpretations"), [], "switched off, the step asks nothing")

if __name__ == "__main__":
    unittest.main()
