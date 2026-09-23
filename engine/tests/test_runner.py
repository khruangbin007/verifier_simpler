"""Tests of runner.py: the wrapper around chat(), the store, paths, the runner, Output.xlsx."""
import json
import os
import threading
import time
import tempfile
import unittest

import helpers
import verifier
core = runner = verifier   # the engine is one module now
import failing_chat
import standin_chat


def simple_validate(question, response_text):
    try:
        return "accepted", json.loads(response_text)
    except ValueError:
        return "rejected: " + core.REJECTION_REASONS[0], None


def make_questions(count):
    questions = []
    for number in range(count):
        prompt = "QUESTION TYPE: self-test\nquestion number %03d" % number
        questions.append({"question_id": core.sha256_text("system" + prompt), "question_type": "self-test",
                          "unit_ref": "M-%04d" % number, "system_prompt": "system", "main_prompt": prompt})
    return questions


def fresh_store():
    folder = helpers.scratch()
    return runner.AuditStore(os.path.join(folder, "local"), os.path.join(folder, "remote")), folder


def quick_settings(**overrides):
    base = {"retry_wait_seconds": 0.0, "sync_every_calls": 7, "token_lifetime_minutes": 1000.0}
    base.update(overrides)
    return runner.make_settings(base)


class TokenRefresh(unittest.TestCase):
    """Plan, Phase 0: workers pause and resume; no question is lost or asked twice; results do
    not depend on the number of workers; the token never reaches the disk."""

    def run_with_workers(self, workers):
        store, folder = fresh_store()
        os.makedirs(store.local_dir)
        live, issuer = runner.LiveValues(), failing_chat.TokenIssuer(lifetime_seconds=0.2)
        first_token = issuer.issue()
        live.update("https://gateway.example", first_token, "analyst-1")
        chat, calls = failing_chat.make_failing_chat(live, issuer, jitter=0.01)
        issued = [first_token]
        time.sleep(0.25)                           # the first token is stale before the step starts

        def paste_fresh_tokens():                  # the helper thread that "pastes" new tokens
            for _ in range(60):
                time.sleep(0.1)
                issued.append(issuer.issue())
                live.update("https://gateway.example", issued[-1], "analyst-1")
        helper = threading.Thread(target=paste_fresh_tokens, daemon=True)
        helper.start()
        ask = runner.make_asker(chat, live, store, quick_settings(concurrency_limit=workers), simple_validate,
                             sleep=lambda seconds: time.sleep(min(seconds, 0.01)))
        questions = make_questions(45)
        answers = ask(questions)
        return store, folder, questions, answers, calls, issued

    def test_pause_resume_nothing_lost_nothing_twice_no_token_on_disk(self):
        outcomes = {}
        for workers in (1, 4, 16):
            store, folder, questions, answers, calls, issued = self.run_with_workers(workers)
            self.assertEqual(len(answers), len(questions), "a question was lost")
            self.assertTrue(all(record["outcome"] == "accepted" for record in answers.values()))
            self.assertTrue(all(count == 1 for count in calls.values()), "a question was asked twice")
            failed = [r for r in store.read_calls() if r["outcome"].startswith("failed: authentication")]
            self.assertTrue(failed, "the test never met an expired token")
            outcomes[workers] = {qid: record["answer"] for qid, record in answers.items()}
            for root, _, names in os.walk(folder):
                for name in names:
                    with open(os.path.join(root, name), "rb") as handle:
                        data = handle.read()
                    if name.endswith(".gz"):
                        import gzip
                        data = gzip.decompress(data)
                    for token in issued:
                        self.assertNotIn(token.encode(), data, "a token reached %s" % name)
        self.assertEqual(outcomes[1], outcomes[4])
        self.assertEqual(outcomes[4], outcomes[16])

    def test_all_three_failure_shapes_are_classified_as_authentication(self):
        live, issuer = runner.LiveValues(), failing_chat.TokenIssuer(lifetime_seconds=0.0)
        live.update("e", issuer.issue(), "u")
        chat, _ = failing_chat.make_failing_chat(live, issuer)
        seen = [runner.call_chat(chat, "s", "m", live) for _ in range(3)]
        self.assertEqual([failure for _, failure, _ in seen], ["authentication"] * 3)
        self.assertTrue(all("TESTTOKEN" not in text for _, _, text in seen))

    def gateway_reply(self, answer="OK.", finish_reason="stop", **more):
        """A reply shaped like the real gateway's: the answer at the top, the OpenAI-shaped
        part underneath, and the prompts echoed back."""
        reply = {"query": "Say OK.", "documents": [], "thread_id": "default", "app": "sparkair",
                 "flow_name": "secure_ai_chat", "maxtoken": 4096, "temperature": 0.01, "top_k": 1,
                 "defaultprompt": "You are a connectivity test.", "chat_id": "6aad41b6361d41df126b3b4c",
                 "user_id": "mel_lorenzo", "datetime": "2026-09-18 13:50:48.697807+00:00",
                 "history": [], "intent": "LLM",
                 "source": [{"doc_name": "Referred Prompt", "source_content": "You are a connectivity test."}],
                 "answer": answer,
                 "response": {"id": "chatcmpl-3729b8fc", "object": "chat.completion",
                              "model": "google/gemma-4-26B-A4B-it",
                              "choices": [{"index": 0, "finish_reason": finish_reason,
                                           "message": {"role": "assistant", "content": answer}}],
                              "usage": {"prompt_tokens": 308, "completion_tokens": 3, "total_tokens": 311}}}
        reply.update(more)
        return reply

    def test_the_real_gateway_reply_gives_its_answer_and_its_record(self):
        live = runner.LiveValues(); live.update("e", "TESTTOKEN-secret", "u")
        outcome = runner.call_chat(lambda s, m, history=[]: self.gateway_reply(), "s", "m", live)
        answer, failure, seen = outcome
        self.assertEqual((answer, failure, seen), ("OK.", "", ""))
        self.assertEqual(outcome.meta["model"], "google/gemma-4-26B-A4B-it")
        self.assertEqual(outcome.meta["total_tokens"], 311)
        self.assertEqual(outcome.meta["chat_id"], "6aad41b6361d41df126b3b4c")
        self.assertEqual(outcome.meta["finish_reason"], "stop")

    def test_history_is_sent_empty_and_only_when_chat_takes_it(self):
        live = runner.LiveValues(); live.update("e", "TESTTOKEN-0001-secret", "u")
        seen_history = []

        def three_argument_chat(SystemPrompt, MainPrompt, history=[]):
            seen_history.append(history)
            return self.gateway_reply()

        def two_argument_chat(SystemPrompt, MainPrompt):
            return {"answer": "OK."}

        self.assertEqual(runner.call_chat(three_argument_chat, "s", "m", live)[0], "OK.")
        self.assertEqual(seen_history, [[]])
        self.assertTrue(runner.accepts_history(three_argument_chat))
        self.assertFalse(runner.accepts_history(two_argument_chat))
        self.assertEqual(runner.call_chat(two_argument_chat, "s", "m", live)[0], "OK.")

    def test_an_answer_cut_off_at_the_token_limit_is_a_failure_not_an_answer(self):
        live = runner.LiveValues(); live.update("e", "TESTTOKEN-0001-secret", "u")
        answer, failure, seen = runner.call_chat(
            lambda s, m, history=[]: self.gateway_reply(answer='{"matches": [', finish_reason="length"),
            "s", "m", live)
        self.assertIsNone(answer)
        self.assertEqual(failure, "truncated")
        self.assertIn("cut off", seen)

    def test_the_answer_is_read_from_the_nested_part_when_the_top_one_is_empty(self):
        live = runner.LiveValues(); live.update("e", "TESTTOKEN-0001-secret", "u")
        reply = self.gateway_reply(answer="the text")
        reply["answer"] = ""
        self.assertEqual(runner.call_chat(lambda s, m, history=[]: reply, "s", "m", live)[0], "the text")
        self.assertEqual(runner.call_chat(lambda s, m: "plain text", "s", "m", live)[0], "plain text")

    def test_a_failed_reply_is_recorded_without_the_echoed_prompts(self):
        live = runner.LiveValues(); live.update("e", "TESTTOKEN-secret", "u")
        reply = self.gateway_reply()
        reply["answer"], reply["response"] = "", {}
        reply["status"], reply["message"] = 503, "service overloaded"
        reply["defaultprompt"] = "x" * 5000
        answer, failure, seen = runner.call_chat(lambda s, m, history=[]: reply, "s", "m", live)
        self.assertIsNone(answer)
        self.assertEqual(failure, "overload")
        self.assertNotIn("x" * 100, seen)
        self.assertNotIn("Say OK.", seen)
        self.assertIn("503", seen)

    def test_a_scratch_folder_that_refuses_is_passed_over_for_one_that_does_not(self):
        """On a shared cluster /tmp/verifier_scratch may already belong to another user and refuse
        this one. Refusal is tested here with a file standing where the folder should be,
        which refuses every user alike, including root."""
        refuses = os.path.join(helpers.scratch(), "refuses")
        with open(refuses, "w") as handle:
            handle.write("not a folder")
        chosen = runner.pick_scratch_root(refuses)
        self.assertNotEqual(chosen, refuses, "the folder that refuses is not the one used")
        probe = os.path.join(chosen, ".still_writable")
        with open(probe, "w") as handle:
            handle.write("x")
        os.remove(probe)

    def test_a_scratch_folder_that_works_is_the_one_used(self):
        wanted = os.path.join(helpers.scratch(), "wanted")
        self.assertEqual(runner.pick_scratch_root(wanted), wanted)

    def test_stop_mode_pauses_the_run_instead_of_waiting(self):
        store, _ = fresh_store()
        os.makedirs(store.local_dir)
        live, issuer = runner.LiveValues(), failing_chat.TokenIssuer(lifetime_seconds=0.0)
        live.update("e", issuer.issue(), "u")
        chat, _ = failing_chat.make_failing_chat(live, issuer)
        ask = runner.make_asker(chat, live, store, quick_settings(token_wait="stop"), simple_validate, sleep=lambda s: None)
        with self.assertRaises(runner.RunPaused):
            ask(make_questions(3))


class ResumeAndBreaker(unittest.TestCase):
    def test_breaker_opens_and_a_second_call_finishes_without_repeating(self):
        store, _ = fresh_store()
        os.makedirs(store.local_dir)
        live, issuer = runner.LiveValues(), failing_chat.TokenIssuer(lifetime_seconds=1000)
        live.update("e", issuer.issue(), "u")
        outage = {"on": False}
        counted = {}

        def chat(SystemPrompt, MainPrompt):
            if outage["on"]:
                raise RuntimeError("HTTP 503 service overloaded")
            counted[MainPrompt] = counted.get(MainPrompt, 0) + 1
            if len(counted) == 12:
                outage["on"] = True                 # the gateway goes down in the middle of the step
            return {"answer": "{}"}
        settings = quick_settings(concurrency_limit=1, breaker_after_failures=4, max_attempts=10, sync_every_calls=5)
        questions = make_questions(30)
        with self.assertRaises(runner.RunPaused):
            runner.make_asker(chat, live, store, settings, simple_validate, sleep=lambda s: None)(questions)
        answered_before = len(counted)
        outage["on"] = False                        # "closes on request": simply run again
        answers = runner.make_asker(chat, live, store, settings, simple_validate, sleep=lambda s: None)(questions)
        self.assertEqual(len(answers), 30)
        self.assertGreaterEqual(answered_before, 12)
        self.assertTrue(all(count == 1 for count in counted.values()), "a finished question was repeated")

    def test_a_question_that_keeps_failing_ends_as_failed_and_does_not_stop_the_run(self):
        store, _ = fresh_store()
        os.makedirs(store.local_dir)
        live = runner.LiveValues()
        live.update("e", "t", "u")
        def chat(SystemPrompt, MainPrompt):
            if "number 001" in MainPrompt:
                return {"no_answer_here": True}
            return {"answer": "{}"}
        answers = runner.make_asker(chat, live, store, quick_settings(), simple_validate, sleep=lambda s: None)(make_questions(3))
        self.assertEqual(sorted(r["outcome"] for r in answers.values()), ["accepted", "accepted", "failed: other"])


class Store(unittest.TestCase):
    """The record of a run is one workbook a person can open, and everything written to it comes back
    exactly, however long the text."""

    def store(self):
        folder = tempfile.mkdtemp(prefix="store_")
        return runner.AuditStore(os.path.join(folder, "local"), os.path.join(folder, "run", "_audit"))

    def test_everything_written_comes_back_from_the_workbook(self):
        store = self.store()
        store.append("chunks_canon", [{"ref": "C-0001", "text": "a paragraph", "numbers": [1, 2.5]}])
        store.append("run_manifest", [{"run_id": "R-1", "outline_confirmed_by": "analyst.one"}])
        store.append_calls([{"question_id": "q1", "question_type": "judge-unit-to-canon", "outcome": "accepted",
                             "main_prompt": "x" * 70000}])                     # longer than a cell holds
        store.sync()
        self.assertEqual(os.listdir(os.path.dirname(store.target())), ["Audit_Log.xlsx"])
        again = runner.AuditStore(store.local_dir, store.remote_dir)
        self.assertEqual(again.read("chunks_canon"), [{"ref": "C-0001", "text": "a paragraph", "numbers": [1, 2.5]}])
        self.assertEqual(again.read("run_manifest")[0]["outline_confirmed_by"], "analyst.one")
        self.assertEqual(len(again.read_calls()[0]["main_prompt"]), 70000, "text longer than a cell is written in parts")

    def test_the_workbook_is_never_seen_half_written(self):
        store = self.store()
        store.append("chunks_canon", [{"ref": "C-0001", "text": "first"}])
        store.sync()
        import openpyxl
        store.append("chunks_canon", [{"ref": "C-0002", "text": "second"}])
        store.sync()
        book = openpyxl.load_workbook(store.target(), read_only=True)
        self.assertEqual(book["Records"].max_row - 1, 2)
        self.assertEqual(sorted(book.sheetnames), ["Model_Calls", "Records", "Run", "Steps"])

    def test_one_store_for_one_run_folder(self):
        """Two stores on one run folder would overwrite each other's records."""
        folder = tempfile.mkdtemp(prefix="store_")
        paths = type("P", (), {"local_dir": os.path.join(folder, "local"), "audit_dir": os.path.join(folder, "_audit")})()
        runner.OPEN_STORES.clear()
        self.assertIs(runner.open_store(paths, {}), runner.open_store(paths, {}))

class RunnerAndWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.projects = os.path.join(helpers.scratch(), "Projects")
        cls.paths = runner.open_run(cls.projects, "DEMO", scratch_root=helpers.scratch())
        cls.settings = runner.make_settings({})
        cls.first = runner.run_pipeline(cls.paths, cls.settings, chat=standin_chat.chat)

    def test_a_run_goes_from_end_to_end_and_repeats_no_finished_step(self):
        """There is no person in the middle of a run: it reads, maps, asks the model and judges, and a
        second call carries on where the first stopped rather than doing anything twice."""
        self.assertEqual(self.first["state"], "finished", self.first)
        self.assertEqual(self.first["steps_run"],
                         ["prepare-run", "read-inputs", "build-map", "read-with-ai", "link-units"])
        second = runner.run_pipeline(self.paths, self.settings, chat=standin_chat.chat)
        self.assertEqual(second["steps_run"], [], "a finished step was repeated")

    def test_pipeline_refuses_an_unknown_function_and_an_unversioned_step(self):
        """pipeline.yaml is the one place a step is named, versioned and mapped to its function.
        A function the engine does not offer is refused, so nothing is ever loaded by path; and a
        step with no version is refused, so every record's provenance carries one."""
        import tempfile
        folder = tempfile.mkdtemp(prefix="pipeline_")
        with open(os.path.join(runner.ENGINE_DIR, "pipeline.yaml"), encoding="utf-8") as handle:
            text = handle.read()
        with open(os.path.join(folder, "pipeline.yaml"), "w", encoding="utf-8") as handle:
            handle.write(text.replace("carried_out_by: prepare_run", "carried_out_by: os.system", 1))
        with self.assertRaises(ValueError) as refused:
            runner.load_pipeline(folder)
        self.assertIn("not a function this engine offers", str(refused.exception))
        with open(os.path.join(folder, "pipeline.yaml"), "w", encoding="utf-8") as handle:
            handle.write(__import__("re").sub(r"  version: [^\n]*\n(  carried_out_by: prepare_run)", r"\1", text, count=1))
        with self.assertRaises(ValueError) as refused:
            runner.load_pipeline(folder)
        self.assertIn("has no 'version'", str(refused.exception))

    def test_workbook_format(self):
        import openpyxl
        layout = runner.load_layout()
        workbook = openpyxl.load_workbook(os.path.join(self.paths.outputs_dir, "Output.xlsx"))
        self.assertEqual(workbook.sheetnames, [sheet["name"] for sheet in layout["sheets"]])
        self.assertEqual(len(workbook.sheetnames), 7)
        for sheet_layout in layout["sheets"]:
            sheet = workbook[sheet_layout["name"]]
            self.assertLessEqual(len(sheet.title), 31)
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(headers, [column["header"] for column in sheet_layout["columns"]])
            self.assertEqual(sheet.freeze_panes, "B2")
            self.assertTrue(sheet.auto_filter.ref.startswith("A1:"))
            self.assertEqual(list(sheet.merged_cells.ranges), [])
            self.assertTrue(all(cell.alignment.wrap_text for cell in sheet[1]))
            self.assertTrue(sheet.protection.sheet)
        self.assertNotIn("Flagged_Items", workbook.sheetnames, "the tool raises nothing for a person to decide")
        info = {row[1]: row[2] for row in workbook["Model_Package_Info"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(info["Model ID"], "DEMO")
        self.assertIn("Run progress", info)
        self.assertEqual(json.loads(workbook.properties.description)["run_id"], self.paths.run_id)

    def test_a_run_writes_only_inside_its_own_folder(self):
        project_entries = sorted(os.listdir(self.paths.project_dir))
        self.assertEqual(project_entries, ["Inputs", self.paths.run_id])

    def test_plain_cell_withholds_python_text_and_banned_wording_but_not_quotations(self):
        store = runner.AuditStore(self.paths.local_dir, self.paths.audit_dir)
        self.assertEqual(runner.plain_cell("KeyError: 'x'", False, store), runner.CELL_WITHHELD)
        self.assertEqual(runner.plain_cell("a critical point", False, store), runner.CELL_WITHHELD)
        quoted = "The methodology says " + runner.quoted("the standard error is small", "C-0007")
        self.assertEqual(runner.plain_cell(quoted, False, store), quoted)
        self.assertEqual(runner.plain_cell("None accepted: nothing fits", False, store), "None accepted: nothing fits")
        self.assertEqual(runner.plain_cell(5.0, False, store), 5)


if __name__ == "__main__":
    unittest.main()
