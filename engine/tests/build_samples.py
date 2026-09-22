"""build_samples.py - builds the sample projects under engine/tests/sample_projects/.

Usage: python tools/build_samples.py [A_minimal F_capital F_capital_known D_dosing]
The generated input files are checked in, so the tests never depend on this tool; run it only
when a sample has to change. Everything a sample contains is written down in this file, so a
reviewer can read what the engine is tested against. Samples are invented; none is real.
"""
import io
import csv
import os
import sys
import tarfile
import zipfile
import warnings

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SAMPLES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_projects")
FOLDERS = {"canon": "1_Methodology", "package": "2_Model_Package", "doc": "3_Model_Documentation"}


# ------------------------------------------------------------------ writers
def write_bytes(sample, corner, name, data):
    folder = os.path.join(SAMPLES, sample, "Inputs", FOLDERS[corner])
    os.makedirs(folder, exist_ok=True)
    with open(os.path.join(folder, name), "wb") as handle:
        handle.write(data if isinstance(data, bytes) else data.encode("utf-8"))


def tarball(package_name, files):
    """A .tar.gz with one top folder, fixed times and owners, so the same content gives the same bytes."""
    raw = io.BytesIO()
    with tarfile.open(fileobj=raw, mode="w", format=tarfile.GNU_FORMAT) as archive:
        for path in sorted(files):
            data = files[path] if isinstance(files[path], bytes) else files[path].encode("utf-8")
            member = tarfile.TarInfo("%s/%s" % (package_name, path))
            member.size, member.mtime, member.mode = len(data), 1767225600, 0o644
            archive.addfile(member, io.BytesIO(data))
    import gzip
    packed = io.BytesIO()
    with gzip.GzipFile(filename="", mode="wb", fileobj=packed, mtime=0) as handle:
        handle.write(raw.getvalue())
    return packed.getvalue()


def fixed_gzip(data):
    """A gzip stream with no build time and no file name in its header, so the same content gives
    the same bytes. The writer of stored R data stamps both, and every sample's tarball changed on
    every rebuild because of it."""
    import gzip
    if not data.startswith(b"\x1f\x8b"):
        return data
    packed = io.BytesIO()
    with gzip.GzipFile(filename="", mode="wb", fileobj=packed, mtime=0) as handle:
        handle.write(gzip.decompress(data))
    return packed.getvalue()


def fixed_zip(data):
    """A ZIP archive - a Word file is one - with every member stamped with the same time, in the
    same order and compressed the same way, so the same content gives the same bytes. Word files
    carry the moment they were saved on every member, which made every sample's documentation
    change on every rebuild though not one word of it had."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for member in source.infolist():
            fixed = zipfile.ZipInfo(member.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type, fixed.external_attr, fixed.create_system = member.compress_type, member.external_attr, 0
            target.writestr(fixed, source.read(member))
    return out.getvalue()


def r_data(objects, single=False):
    """Stored R data written without R, by the writer of the `rdata` package (fixture route c)."""
    import rdata
    path = os.path.join(ROOT, "tmp_rdata.bin")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        if single:
            rdata.write_rds(path, objects)
        else:
            rdata.write_rda(path, objects)
    with open(path, "rb") as handle:
        data = handle.read()
    os.remove(path)
    return fixed_gzip(data)


def docx_file(parts):
    """parts: ("h", level, text) | ("p", text) | ("li", text) | ("table", caption, rows) | ("eq", omml xml)
    | ("image", caption)"""
    import docx
    from docx.oxml import parse_xml
    document = docx.Document()
    for part in parts:
        if part[0] == "h":
            document.add_heading(part[2], level=part[1])
        elif part[0] == "p":
            document.add_paragraph(part[1])
        elif part[0] == "li":
            document.add_paragraph(part[1], style="List Bullet")
        elif part[0] == "table":
            document.add_paragraph(part[1], style="Caption")
            table = document.add_table(rows=len(part[2]), cols=len(part[2][0]))
            table.style = "Table Grid"
            for row, cells in zip(table.rows, part[2]):
                for cell, text in zip(row.cells, cells):
                    cell.text = text
        elif part[0] == "eq":
            paragraph = document.add_paragraph()
            paragraph._p.append(parse_xml(part[1]))
        elif part[0] == "image":
            document.add_picture(io.BytesIO(tiny_png()))
            document.add_paragraph(part[1], style="Caption")
    raw = io.BytesIO()
    document.save(raw)
    return fixed_zip(raw.getvalue())


def pdf_file(parts):
    """parts as for docx_file; headings are bold and larger, tables are drawn with a grid."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    styles, story, raw = getSampleStyleSheet(), [], io.BytesIO()
    for part in parts:
        if part[0] == "h":
            story.append(Paragraph(part[2], styles["Heading%d" % min(3, part[1])]))
        elif part[0] in ("p", "li"):
            story.append(Paragraph(part[1], styles["BodyText"]))
        elif part[0] == "table":
            story.append(Paragraph(part[1], styles["BodyText"]))
            story.append(Table(part[2], style=TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)])))
        story.append(Spacer(1, 8))
    SimpleDocTemplate(raw, pagesize=A4, invariant=1).build(story)
    return raw.getvalue()


def tiny_png():
    import base64
    return base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def omml(inner):
    return ('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">%s</m:oMath>' % inner)


def run(text):
    return "<m:r><m:t>%s</m:t></m:r>" % text


# ------------------------------------------------------------------ sample A_minimal (neutral flavour)
A_METHODOLOGY = """<?xml version="1.0" encoding="UTF-8"?>
<document>
 <title>Parcel pricing method</title>
 <section><title>1. Purpose and scope</title>
  <p>This method sets out how the price of a parcel is determined from its size, its weight and its destination zone.</p>
  <p>It applies to every parcel handed in at a counter. Pallets and letters are outside its scope.</p>
  <p>The method is reviewed once a year by the pricing committee.</p>
 </section>
 <section><title>2. Definitions</title>
  <list>
   <item>Actual weight (AW) is the weight of the parcel in kilograms as shown by the scale.</item>
   <item>Dimensional weight (DW) is a weight derived from the volume of the parcel.</item>
   <item>Chargeable weight (CW) is the weight on which the price is based.</item>
   <item>The zone is the destination area of the parcel: Local, National or Abroad.</item>
  </list>
 </section>
 <section><title>3. Weights</title>
  <section><title>3.1 Dimensional weight</title>
   <p>The dimensional weight is obtained from the length L, the width W and the height H of the parcel, measured in centimetres.</p>
   <equation><math xmlns="http://www.w3.org/1998/Math/MathML"><mi>DW</mi><mo>=</mo><mfrac><mrow><mi>L</mi><mo>&#xD7;</mo><mi>W</mi><mo>&#xD7;</mo><mi>H</mi></mrow><mn>5000</mn></mfrac></math></equation>
   <p>The divisor of 5000 is fixed and is not varied by zone.</p>
  </section>
  <section><title>3.2 Chargeable weight</title>
   <p>The chargeable weight is the larger of the actual weight and the dimensional weight: CW = max(AW, DW).</p>
   <p>The chargeable weight is never taken below 0.5 kilograms.</p>
   <p>Weights are not rounded before the price is calculated.</p>
  </section>
 </section>
 <section><title>4. Price</title>
  <section><title>4.1 Zone rates</title>
   <p>Each zone has a base rate and a rate per kilogram, as set out in Table 1.</p>
   <table><caption>Table 1. Rates by zone</caption>
    <tr><th>Zone</th><th>Base rate</th><th>Rate per kg</th></tr>
    <tr><td>Local</td><td>3.20</td><td>0.85</td></tr>
    <tr><td>National</td><td>4.90</td><td>1.10</td></tr>
    <tr><td>Abroad</td><td>9.50</td><td>2.75</td></tr>
   </table>
   <p>Rates are stated in the currency of the counter and exclude tax.</p>
  </section>
  <section><title>4.2 Price before surcharge</title>
   <p>The price before surcharge is the base rate plus the rate per kilogram times the chargeable weight: P = B + R * CW.</p>
   <p>B and R are taken from Table 1 for the zone of the parcel.</p>
  </section>
  <section><title>4.3 Fuel surcharge</title>
   <p>A fuel surcharge of 12.5% is added to the price before surcharge: S = P * (1 + 0.125).</p>
   <p>The surcharge rate is set by the pricing committee and is the same for all zones.</p>
  </section>
  <section><title>4.4 Volume discount</title>
   <p>A customer who hands in n parcels at once receives a discount rate of 1% per parcel, capped at 20%: D = min(0.2, 0.01 * n).</p>
   <p>The discount is applied to the price after surcharge.</p>
   <figure><img src="discount_curve.png" alt="Discount rate against number of parcels"/><caption>Figure 1. Discount rate by number of parcels</caption></figure>
  </section>
  <section><title>4.5 Minimum price</title>
   <p>The final price is never below 4.50, whatever the zone and the discount.</p>
   <p>Final prices are rounded to 2 decimals.</p>
  </section>
 </section>
 <section><title>5. Governance</title>
  <p>Changes to Table 1 are approved by the pricing committee and recorded in the change log.</p>
  <p>Questions about this method are answered by the pricing office.</p>
 </section>
</document>
"""

A_PACKAGE = {
    "DESCRIPTION": "Package: parcelcost\nTitle: Parcel Pricing\nVersion: 1.2.0\nDescription: Prices parcels by zone\n    and chargeable weight.\nLicense: MIT\nEncoding: UTF-8\n",
    "NAMESPACE": "export(dim_weight)\nexport(chargeable_weight)\nexport(parcel_price)\nexport(volume_discount)\n",
    "R/weights.R": """#' Dimensional weight
#'
#' Weight derived from the volume of a parcel: \\deqn{DW = \\frac{L \\cdot W \\cdot H}{5000}}
#' @param l length in centimetres
#' @param w width in centimetres
#' @param h height in centimetres
#' @param divisor the fixed divisor, 5000
#' @export
dim_weight <- function(l, w, h, divisor = 5000) {
  stopifnot(is.numeric(l), is.numeric(w), is.numeric(h))
  l * w * h / divisor
}

#' Chargeable weight
#'
#' The larger of actual weight and dimensional weight, never below 0.5 kilograms.
#' @param aw actual weight in kilograms
#' @param dw dimensional weight in kilograms
#' @export
chargeable_weight <- function(aw, dw) {
  cw <- pmax(aw, dw)
  pmax(cw, 0.5)
}
""",
    "R/price.R": """#' Price of a parcel
#'
#' Base rate plus rate per kilogram times chargeable weight, with the fuel surcharge of 12.5%,
#' the volume discount and the minimum price of 4.50.
#' @param zone destination zone: Local, National or Abroad
#' @param cw chargeable weight in kilograms
#' @param n number of parcels handed in at once
#' @export
parcel_price <- function(zone, cw, n = 1) {
  if (!zone %in% zone_rates$zone) stop("unknown zone")
  base <- zone_rates[zone_rates$zone == zone, "base_rate"]
  rate <- zone_rates[zone_rates$zone == zone, "rate_per_kg"]
  price <- base + rate * cw
  with_fuel <- price * (1 + 0.125)
  discounted <- with_fuel * (1 - volume_discount(n))
  round(pmax(discounted, 4.50), 2)
}

#' Volume discount
#'
#' Discount rate of 1% per parcel, capped at 20%.
#' @param n number of parcels handed in at once
#' @export
volume_discount <- function(n) {
  pmin(0.2, 0.01 * n)
}

# Helper used when printing a price list.
format_price <- function(x) {
  paste0(format(x, nsmall = 2), " per parcel")
}
""",
    "R/data.R": """#' Rates by zone
#'
#' Base rate and rate per kilogram for each destination zone.
#' @format A data frame with 3 rows and 3 columns:
#' \\describe{
#'   \\item{zone}{destination zone}
#'   \\item{base_rate}{base rate per parcel}
#'   \\item{rate_per_kg}{rate per kilogram of chargeable weight}
#' }
"zone_rates"
""",
    "man/dim_weight.Rd": """% Generated by roxygen2: do not edit by hand
% Please edit documentation in R/weights.R
\\name{dim_weight}
\\alias{dim_weight}
\\title{Dimensional weight}
\\usage{
dim_weight(l, w, h, divisor = 5000)
}
\\arguments{
\\item{l}{length in centimetres}
\\item{w}{width in centimetres}
\\item{h}{height in centimetres}
\\item{divisor}{the fixed divisor, 5000}
}
\\description{
Weight derived from the volume of a parcel.
}
""",
    "man/chargeable_weight.Rd": """% Generated by roxygen2: do not edit by hand
% Please edit documentation in R/weights.R
\\name{chargeable_weight}
\\alias{chargeable_weight}
\\title{Chargeable weight}
\\usage{
chargeable_weight(aw, dw)
}
\\arguments{
\\item{aw}{actual weight in kilograms}
\\item{dw}{dimensional weight in kilograms}
}
""",
    "man/parcel_price.Rd": """% Generated by roxygen2: do not edit by hand
% Please edit documentation in R/price.R
\\name{parcel_price}
\\alias{parcel_price}
\\title{Price of a parcel}
\\usage{
parcel_price(zone, cw, n = 1)
}
\\arguments{
\\item{zone}{destination zone: Local, National or Abroad}
\\item{cw}{chargeable weight in kilograms}
\\item{n}{number of parcels handed in at once}
}
""",
    "man/volume_discount.Rd": """% Generated by roxygen2: do not edit by hand
% Please edit documentation in R/price.R
\\name{volume_discount}
\\alias{volume_discount}
\\title{Volume discount}
\\usage{
volume_discount(n)
}
\\arguments{
\\item{n}{number of parcels handed in at once}
}
""",
    "tests/testthat.R": "library(testthat)\nlibrary(parcelcost)\ntest_check(\"parcelcost\")\n",
    "tests/testthat/test-price.R": """test_that("dimensional weight follows the divisor", {
  expect_equal(dim_weight(50, 40, 30), 12)
})

test_that("the minimum price holds", {
  expect_equal(parcel_price("Local", 0.5), 4.50)
})
""",
}

A_DOCUMENTATION = [
    ("h", 1, "1 Overview"),
    ("p", "This document describes how the parcelcost package implements the parcel pricing method."),
    ("p", "The package is maintained by the pricing office and released once a year."),
    ("h", 1, "2 Weights"),
    ("p", "The function dim_weight derives the dimensional weight from length, width and height with the fixed divisor of 5000."),
    ("p", "The function chargeable_weight returns the larger of actual weight and dimensional weight and never returns less than 0.5 kilograms."),
    ("h", 1, "3 Price"),
    ("p", "The function parcel_price looks up the base rate and the rate per kilogram of the zone in the table zone_rates."),
    ("table", "Table A. Rates held in the package", [["Zone", "Base rate", "Rate per kg"], ["Local", "3.20", "0.85"],
                                                     ["National", "4.90", "1.10"], ["Abroad", "9.50", "2.75"]]),
    ("p", "The price before surcharge equals the base rate plus the rate per kilogram times the chargeable weight."),
    ("p", "A fuel surcharge of 12.5% is added to the price before surcharge."),
    ("p", "The volume discount is 1% per parcel and is capped at 20%."),
    ("p", "The final price is never below 4.50 and is rounded to 2 decimals."),
    ("h", 1, "4 Testing"),
    ("p", "Unit tests cover the dimensional weight and the minimum price."),
    ("p", "The tests are run before every release."),
    ("h", 1, "5 Contacts"),
    ("p", "Questions are answered by the pricing office."),
    ("image", "Figure A. Release calendar"),
]


def build_a_minimal():
    import pandas
    sample = "A_minimal"
    write_bytes(sample, "canon", "parcel_pricing_method.xml", A_METHODOLOGY)
    rates = pandas.DataFrame({"zone": ["Local", "National", "Abroad"], "base_rate": [3.20, 4.90, 9.50],
                              "rate_per_kg": [0.85, 1.10, 2.75]})
    files = dict(A_PACKAGE)
    files["data/zone_rates.rda"] = r_data({"zone_rates": rates})
    write_bytes(sample, "package", "parcelcost_1.2.0.tar.gz", tarball("parcelcost", files))
    write_bytes(sample, "doc", "parcelcost_documentation.docx", docx_file(A_DOCUMENTATION))



# ------------------------------------------------------------------ sample F_capital (finance flavour, invented)
MATHML_COND_PD = ('<math xmlns="http://www.w3.org/1998/Math/MathML"><mi>PDc</mi><mo>=</mo><mi>N</mi><mo>&#x2061;</mo><mfenced><mfrac>'
                  '<mrow><msup><mi>N</mi><mrow><mo>-</mo><mn>1</mn></mrow></msup><mo>&#x2061;</mo><mfenced><mi>PD</mi></mfenced><mo>+</mo>'
                  '<msqrt><mi>&#x3C1;</mi></msqrt><mo>&#x2062;</mo><msup><mi>N</mi><mrow><mo>-</mo><mn>1</mn></mrow></msup><mo>&#x2061;</mo>'
                  '<mfenced><mn>0.999</mn></mfenced></mrow><msqrt><mrow><mn>1</mn><mo>-</mo><mi>&#x3C1;</mi></mrow></msqrt></mfrac></mfenced></math>')

F_METHODOLOGY = """<?xml version="1.0" encoding="UTF-8"?>
<methodology>
 <title>Capital requirement methodology</title>
 <part><title>A. Framework</title>
  <section><title>1. Scope</title>
   <p>This methodology sets out how the capital requirement for unexpected losses is determined for each exposure.</p>
   <p>It applies to all exposures of the lending book. Trading positions are outside its scope.</p>
  </section>
  <section><title>2. Definitions</title>
   <p>The probability of default (PD) is the likelihood that an obligor defaults within one year.</p>
   <p>The loss given default (LGD) is the share of the exposure that is lost when the obligor defaults.</p>
   <p>The exposure at default (EAD) is the amount outstanding when the obligor defaults.</p>
   <p>The requirement uses the normal distribution function N and its inverse, where &rho; denotes the asset correlation.</p>
   <table><caption>Table 1. Symbols</caption>
    <tr><th>Symbol</th><th>Description</th></tr>
    <tr><td>PDc</td><td>conditional probability of default</td></tr>
    <tr><td>K</td><td>capital requirement per unit of exposure</td></tr>
    <tr><td>M</td><td>effective maturity in years</td></tr>
   </table>
  </section>
  <section><title>3. Risk components</title>
   <section><title>3.1 Probability of default</title>
    <p>The probability of default is estimated from internal ratings and is reviewed every year.</p>
    <section><title>3.1.1 Floor</title>
     <p>The probability of default is never taken below 0.03%.</p>
     <p>The floor applies before any other step of the calculation.</p>
    </section>
   </section>
   <section><title>3.2 Conditional probability of default</title>
    <p>The conditional probability of default is evaluated at the 99.9th percentile of the systematic factor.</p>
    <equation>""" + MATHML_COND_PD + """</equation>
    <p>The confidence level of 0.999 is fixed and is set out again in Annex A.</p>
   </section>
   <section><title>3.3 Asset correlation</title>
    <p>The asset correlation falls with the probability of default: R = 0.12 * (1 - exp(-50 * PD)) / (1 - exp(-50)) + 0.24 * (1 - (1 - exp(-50 * PD)) / (1 - exp(-50))).</p>
    <p>The correlation therefore lies between 0.12 and 0.24.</p>
   </section>
   <section><title>3.4 Loss given default</title>
    <p>The loss given default is at least the floor of its segment, as set out in Table 3.</p>
    <table><caption>Table 3. Floors by segment</caption>
     <tr><th>Segment</th><th>LGD floor</th></tr>
     <tr><td>Retail</td><td>0.15</td></tr>
     <tr><td>Corporate</td><td>0.25</td></tr>
     <tr><td>Bank</td><td>0.45</td></tr>
    </table>
   </section>
  </section>
 </part>
 <part><title>B. Capital</title>
  <section><title>4. Capital requirement</title>
   <section><title>4.1 Requirement per unit of exposure</title>
    <p>The unexpected loss rate is the conditional probability of default less the probability of default: UL = PDc - PD.</p>
    <p>K is the product of LGD and UL.</p>
   </section>
   <section><title>4.2 Risk-weighted amount</title>
    <equation><img src="rwa_formula.png" alt="Formula for the risk-weighted amount"/></equation>
    <p>The requirement is scaled by 12.5 and multiplied by the exposure at default.</p>
   </section>
   <section><title>4.3 Effective maturity</title>
    <p>The effective maturity is at least 1 year.</p>
    <p>The effective maturity is capped at 5 years.</p>
   </section>
  </section>
 </part>
 <annex>
  <heading>Annex A Parameters</heading>
  <p>This annex restates the fixed parameters of the methodology.</p>
  <heading>I. Confidence level</heading>
  <p>The confidence level is 0.999 for every segment.</p>
  <heading>II. Floors</heading>
  <p>Floors are reviewed once a year by the model owner.</p>
  <heading>(a) Probability of default</heading>
  <p>The floor of the probability of default is 0.03%, as stated in section 3.1.1.</p>
  <heading>(b) Loss given default</heading>
  <p>The floors of the loss given default are those of Table 3.</p>
 </annex>
</methodology>
"""

def rd_page(name, title, arguments, source):
    usage = "%s(%s)" % (name, ", ".join(a if d is None else "%s = %s" % (a, d) for a, d, _ in arguments))
    items = "\n".join("\\item{%s}{%s}" % (a, text) for a, _, text in arguments)
    return ("%% Generated by roxygen2: do not edit by hand\n%% Please edit documentation in %s\n\\name{%s}\n\\alias{%s}\n\\title{%s}\n"
            "\\usage{\n%s\n}\n\\arguments{\n%s\n}\n" % (source, name, name, title, usage, items))

F_R_PD = """#' Floor of the probability of default
#'
#' The probability of default is never taken below 0.03%.
#' @param pd probability of default
#' @param floor the floor, defaults to 0.0003
#' @export
floor_pd <- function(pd, floor = 0.0003) {
  pmax(pd, floor)
}

#' Conditional probability of default
#'
#' Evaluated at the 99.9th percentile of the systematic factor:
#' \\deqn{PDc = \\Phi\\left(\\frac{\\Phi^{-1}(PD) + \\sqrt{\\rho} \\cdot \\Phi^{-1}(0.999)}{\\sqrt{1 - \\rho}}\\right)}
#' @param pd probability of default
#' @param rho asset correlation
#' @param q confidence level, defaults to 0.999
#' @export
cond_pd <- function(pd, rho, q = 0.999) {
  a <- qnorm(pd)
  b <- sqrt(rho) * qnorm(q)
  pnorm((a + b) / sqrt(1 - rho))
}
"""
F_R_CORRELATION = """#' Asset correlation
#'
#' The asset correlation falls with the probability of default and lies between 0.12 and 0.24.
#' @param pd probability of default
#' @export
asset_correlation <- function(pd) {
  w <- (1 - exp(-50 * pd)) / (1 - exp(-50))
  0.12 * w + 0.24 * (1 - w)
}
"""
F_R_CAPITAL = """#' Capital requirement per unit of exposure
#'
#' The product of the loss given default, floored by segment, and the unexpected loss rate.
#' @param pd probability of default
#' @param lgd loss given default
#' @param segment segment of the exposure: Retail, Corporate or Bank
#' @export
capital_k <- function(pd, lgd, segment) {
  check_inputs(pd, lgd)
  lgd_f <- pmax(lgd, lgd_floors[lgd_floors$segment == segment, "lgd_floor"])
  pdc <- cond_pd(floor_pd(pd), asset_correlation(pd))
  ul <- pdc - pd
  lgd_f * ul
}

#' Risk-weighted amount
#'
#' The requirement scaled by 12.5 and multiplied by the exposure at default.
#' @param k capital requirement per unit of exposure
#' @param ead exposure at default
#' @export
rwa <- function(k, ead) {
  k * 12.5 * ead
}

#' Effective maturity
#'
#' At least 1 year and capped at 5 years.
#' @param m maturity in years
#' @export
effective_maturity <- function(m) {
  pmin(pmax(m, 1), 5)
}
"""
F_R_UTILS = """# Argument checks shared by the exported functions.
check_inputs <- function(pd, lgd) {
  if (!is.numeric(pd)) stop("pd must be numeric")
  if (!is.numeric(lgd)) stop("lgd must be numeric")
  invisible(TRUE)
}

`%||%` <- function(a, b) if (is.null(a)) b else a
"""
F_R_DATA = """#' Floors of the loss given default by segment
#'
#' @format A data frame with 3 rows and 2 columns:
#' \\describe{
#'   \\item{segment}{segment of the exposure}
#'   \\item{lgd_floor}{floor of the loss given default}
#' }
#' @source Table 3 of the methodology
"lgd_floors"
"""
F_VIGNETTE = """---
title: "Using capreq"
---

This vignette shows how the capital requirement of one exposure is obtained.

```{r}
library(capreq)
k <- capital_k(0.01, 0.40, "Corporate")
rwa(k, 1000)
```

The floors by segment are stored in the object lgd_floors.
"""

def f_package(known):
    """The files of package capreq; `known` plants the six seeded differences of the brief."""
    import pandas
    pd_source, capital, data = F_R_PD, F_R_CAPITAL, F_R_DATA
    floors = pandas.DataFrame({"segment": ["Retail", "Corporate", "Bank"], "lgd_floor": [0.15, 0.25, 0.45]})
    if known:
        pd_source = pd_source.replace("pnorm((a + b) / sqrt(1 - rho))", "pnorm((a - b) / sqrt(1 - rho))")      # 1. a flipped sign
        pd_source = pd_source.replace("  pmax(pd, floor)\n", "  pd\n")                                           # 3. a removed floor
        pd_source = pd_source.replace("#' @param rho asset correlation", "#' @param rho_a asset correlation")    # 5. roxygen out of step
        floors.loc[0, "lgd_floor"] = 0.10                                                                        # 4. a changed cell
    correlation = F_R_CORRELATION.replace("0.12 * w", "0.13 * w") if known else F_R_CORRELATION                  # 2. a changed constant
    confidence = pandas.DataFrame({"parameter": ["confidence_level"], "value": [0.999]})
    files = {
        "DESCRIPTION": "Package: capreq\nTitle: Capital Requirement\nVersion: 0.9.1\nDescription: Capital requirement for unexpected losses.\nLicense: MIT\nEncoding: UTF-8\nVignetteBuilder: knitr\n",
        "NAMESPACE": "\n".join("export(%s)" % n for n in ("floor_pd", "cond_pd", "asset_correlation", "capital_k", "rwa", "effective_maturity")) + "\n",
        "R/pd.R": pd_source, "R/correlation.R": correlation, "R/capital.R": capital, "R/utils.R": F_R_UTILS, "R/data.R": data,
        "data/lgd_floors.rda": r_data({"lgd_floors": floors}), "inst/extdata/confidence.rds": r_data(confidence, single=True),
        "man/floor_pd.Rd": rd_page("floor_pd", "Floor of the probability of default", [("pd", None, "probability of default"), ("floor", "0.0003", "the floor, defaults to 0.0003")], "R/pd.R"),
        # the page of cond_pd is stale on purpose: it was generated before the argument q was added
        "man/cond_pd.Rd": rd_page("cond_pd", "Conditional probability of default", [("pd", None, "probability of default"), ("rho", None, "asset correlation")], "R/pd.R"),
        "man/asset_correlation.Rd": rd_page("asset_correlation", "Asset correlation", [("pd", None, "probability of default")], "R/correlation.R"),
        "man/capital_k.Rd": rd_page("capital_k", "Capital requirement per unit of exposure", [("pd", None, "probability of default"), ("lgd", None, "loss given default"), ("segment", None, "segment of the exposure: Retail, Corporate or Bank")], "R/capital.R"),
        "man/rwa.Rd": rd_page("rwa", "Risk-weighted amount", [("k", None, "capital requirement per unit of exposure"), ("ead", None, "exposure at default")], "R/capital.R"),
        "man/effective_maturity.Rd": rd_page("effective_maturity", "Effective maturity", [("m", None, "maturity in years")], "R/capital.R"),
        "tests/testthat.R": "library(testthat)\nlibrary(capreq)\ntest_check(\"capreq\")\n",
        "tests/testthat/test-pd.R": "test_that(\"the floor of the probability of default holds\", {\n  expect_equal(floor_pd(0.0001), 0.0003)\n})\n\n"
                                    "test_that(\"the conditional probability exceeds the probability\", {\n  expect_true(cond_pd(0.01, 0.2) > 0.01)\n})\n",
        "vignettes/capreq.Rmd": F_VIGNETTE}
    return files

def f_documentation(known):
    level = "99.5th percentile" if known else "99.9th percentile"                                              # 6. a stale value in the documentation
    cond = omml(run("PDc=N") + "<m:d><m:e><m:f><m:num>" + "<m:sSup><m:e>" + run("N") + "</m:e><m:sup>" + run("-1") + "</m:sup></m:sSup><m:d><m:e>" + run("PD")
                + "</m:e></m:d>" + run("+") + "<m:rad><m:radPr><m:degHide m:val=\"1\"/></m:radPr><m:deg/><m:e>" + run("&#x3C1;") + "</m:e></m:rad>" + run("&#xD7;")
                + "<m:sSup><m:e>" + run("N") + "</m:e><m:sup>" + run("-1") + "</m:sup></m:sSup><m:d><m:e>" + run("0.999") + "</m:e></m:d></m:num><m:den>"
                + "<m:rad><m:radPr><m:degHide m:val=\"1\"/></m:radPr><m:deg/><m:e>" + run("1-&#x3C1;") + "</m:e></m:rad></m:den></m:f></m:e></m:d>")
    return [
        ("h", 1, "1 Purpose"),
        ("p", "This document describes how the package capreq implements the capital requirement methodology."),
        ("p", "It is written for model validators and for the developers who maintain the package."),
        ("h", 1, "2 Risk components"),
        ("h", 2, "2.1 Probability of default"),
        ("p", "The function floor_pd applies the floor of 0.03% to the probability of default."),
        ("h", 2, "2.2 Conditional probability of default"),
        ("p", "The function cond_pd evaluates the conditional probability of default at the %s of the systematic factor." % level),
        ("eq", cond),
        ("h", 2, "2.3 Asset correlation"),
        ("p", "The function asset_correlation returns a correlation between 0.12 and 0.24 that falls with the probability of default."),
        ("h", 2, "2.4 Loss given default"),
        ("p", "The floors of the loss given default are held in the object lgd_floors and are shown in Table D1."),
        ("table", "Table D1. Floors by segment", [["Segment", "LGD floor"], ["Retail", "0.15"], ["Corporate", "0.25"], ["Bank", "0.45"]]),
        ("h", 1, "3 Capital"),
        ("p", "The function capital_k multiplies the floored loss given default by the unexpected loss rate."),
        ("p", "The function rwa scales the requirement by 12.5 and multiplies it by the exposure at default."),
        ("p", "The function effective_maturity keeps the maturity between 1 year and 5 years."),
        ("h", 1, "4 Change history"),
        ("p", "The package is released after each yearly review of the methodology."),
        ("image", "Figure D1. Release process"),
    ]

F_USER_GUIDE = [
    ("h", 1, "1 Installing the package"),
    ("p", "The package capreq is installed from the internal repository and needs no other package."),
    ("h", 1, "2 Calculating a requirement"),
    ("p", "Call capital_k with the probability of default, the loss given default and the segment of the exposure."),
    ("p", "The risk-weighted amount is obtained with rwa, which scales the requirement by 12.5."),
    ("table", "Floors used by the package", [["Segment", "LGD floor"], ["Retail", "0.15"], ["Corporate", "0.25"], ["Bank", "0.45"]]),
]

def build_f_capital(sample="F_capital", known=False):
    write_bytes(sample, "canon", "capital_methodology.txt", F_METHODOLOGY)          # XML inside a .txt file, on purpose
    write_bytes(sample, "package", "capreq_0.9.1.tar.gz", tarball("capreq", f_package(known)))
    write_bytes(sample, "doc", "capreq_model_documentation.docx", docx_file(f_documentation(known)))
    write_bytes(sample, "doc", "capreq_user_guide.pdf", pdf_file(F_USER_GUIDE))

def build_f_capital_known():
    build_f_capital("F_capital_known", known=True)

# ------------------------------------------------------------------ sample D_dosing (another field; keeps the engine honest about rule R9)
D_OMML = ('<m:oMath><m:r><m:t>D=W&#215;R</m:t></m:r></m:oMath>')
D_HTML = """<html xmlns:m="http://schemas.microsoft.com/office/2004/12/omml"><head><title>Dosing method</title><style>p{margin:0}</style></head><body>
<h1>1 Purpose</h1>
<p>This method sets out how the daily dose of the medicine is determined for an adult patient.
<p>It applies to oral treatment only.
<h1>2 Dose by body weight</h1>
<p>The body weight (W) is measured in kilograms, and the dose rate (R) is 15 milligrams per kilogram.
<p><!--[if gte msEquation 12]>""" + D_OMML + """<![endif]--><![if !msEquation]><img src="image001.png" alt="Dose formula"><![endif]></p>
<p>The daily dose is never above 4000 milligrams.
<h1>3 Adjustment for kidney function</h1>
<p>The dose is multiplied by the factor of the clearance band of the patient, as set out in Table 1.
<table><caption>Table 1. Factors by clearance band</caption>
<tr><th>Band</th><th>Factor</th></tr><tr><td>Normal</td><td>1.00</td></tr><tr><td>Reduced</td><td>0.75</td></tr><tr><td>Low</td><td>0.50</td></tr></table>
<p><img src="image002.png" alt="Adjusted dose formula"></p>
<p>The adjusted dose is rounded to 10 milligrams.
</body></html>
"""

def mhtml_file(page, images):
    boundary = "----=_NextPart_AIVA_SAMPLE"
    parts = ["MIME-Version: 1.0\nContent-Type: multipart/related; boundary=\"%s\"\n" % boundary,
             "--%s\nContent-Location: file:///C:/method.htm\nContent-Transfer-Encoding: 8bit\nContent-Type: text/html; charset=\"utf-8\"\n\n%s" % (boundary, page)]
    import base64
    for name in images:
        parts.append("--%s\nContent-Location: %s\nContent-Transfer-Encoding: base64\nContent-Type: image/png\n\n%s\n"
                     % (boundary, name, base64.b64encode(tiny_png() + name.encode("ascii")).decode("ascii")))
    return "\n".join(parts) + "\n--%s--\n" % boundary

D_R = """#' Dose by body weight
#'
#' The body weight times the dose rate, never above 4000 milligrams.
#' @param weight body weight in kilograms
#' @param rate dose rate in milligrams per kilogram, defaults to 15
#' @export
dose_by_weight <- function(weight, rate = 15) {
  dose <- weight * rate
  pmin(dose, 4000)
}

#' Dose adjusted for kidney function
#'
#' The dose times the factor of the clearance band, rounded to 10 milligrams.
#' @param dose daily dose in milligrams
#' @param band clearance band: Normal, Reduced or Low
#' @export
adjusted_dose <- function(dose, band) {
  factor <- band_factors[band_factors$band == band, "factor"]
  round(dose * factor / 10) * 10
}
"""
D_GUIDE = [
    ("h", 1, "1 Overview"),
    ("p", "The package dosecalc implements the dosing method for adult patients."),
    ("h", 1, "2 Functions"),
    ("p", "The function dose_by_weight multiplies the body weight by the dose rate of 15 milligrams per kilogram and never returns more than 4000 milligrams."),
    ("p", "The function adjusted_dose multiplies the dose by the factor of the clearance band."),
    ("table", "Factors held in the package", [["Band", "Factor"], ["Normal", "1.00"], ["Reduced", "0.75"], ["Low", "0.50"]]),
]

def build_d_dosing():
    import pandas
    sample = "D_dosing"
    write_bytes(sample, "canon", "dosing_method.mhtml", mhtml_file(D_HTML, ["image001.png", "image002.png"]))
    factors = pandas.DataFrame({"band": ["Normal", "Reduced", "Low"], "factor": [1.0, 0.75, 0.5]})
    files = {"DESCRIPTION": "Package: dosecalc\nTitle: Dose Calculation\nVersion: 0.2.0\nLicense: MIT\nEncoding: UTF-8\n",
             "NAMESPACE": "export(dose_by_weight)\nexport(adjusted_dose)\n", "R/dose.R": D_R,
             "R/data.R": "#' Factors by clearance band\n#'\n#' @format A data frame with 3 rows and 2 columns:\n#' \\describe{\n#'   \\item{band}{clearance band}\n#'   \\item{factor}{factor applied to the dose}\n#' }\n\"band_factors\"\n",
             "inst/extdata/band_factors.rds": r_data(factors, single=True),
             "man/dose_by_weight.Rd": rd_page("dose_by_weight", "Dose by body weight", [("weight", None, "body weight in kilograms"), ("rate", "15", "dose rate in milligrams per kilogram, defaults to 15")], "R/dose.R"),
             "man/adjusted_dose.Rd": rd_page("adjusted_dose", "Dose adjusted for kidney function", [("dose", None, "daily dose in milligrams"), ("band", None, "clearance band: Normal, Reduced or Low")], "R/dose.R")}
    write_bytes(sample, "package", "dosecalc_0.2.0.tar.gz", tarball("dosecalc", files))
    write_bytes(sample, "doc", "dosecalc_guide.pdf", pdf_file(D_GUIDE))


BUILDERS = {"A_minimal": build_a_minimal, "F_capital": build_f_capital, "F_capital_known": build_f_capital_known, "D_dosing": build_d_dosing}



# ================================================================== the hard reading samples (plan 0.0.2, R6)
# G, H and I exist to be READ badly by the rules AIVA ships, and to say so out loud where they
# are. Each stresses one corner: G an XML schema whose tags are named nothing the rules know,
# H a PDF laid out in two columns with running headers and footnotes, I a Word file whose
# headings are bold paragraphs and whose words hide in text boxes and tracked changes. Every
# distinctive phrase is listed in the sample's gold_reading.csv and must end in some unit.
# R9 holds: the flavour is invented and carries no domain concept.

# ------------------------------------------------------------------ G_schema: tags nobody anticipated
G_METHODOLOGY = """<?xml version="1.0" encoding="utf-8"?>
<lendingrules xmlns="urn:example:lending:1">
  <ruleblock idx="1.">
    <blockcaption>Loan periods</blockcaption>
    <statementbody>A standard loan runs for fourteen days from the day of issue.</statementbody>
    <statementbody>A short loan runs for two days and may not be renewed.</statementbody>
    <gridholder>
      <gridcaption>Table 1. Loan period by item class</gridcaption>
      <gridline>
        <gridcolhead>Item class</gridcolhead><gridcolhead>Days</gridcolhead><gridcolhead>Conditions</gridcolhead>
      </gridline>
      <gridline>
        <gridcell>Standard</gridcell><gridcell>14</gridcell>
        <gridcell><bullets><point>renewable twice</point><point>no fine in the first two days</point></bullets></gridcell>
      </gridline>
      <gridline>
        <gridcell>Short</gridcell><gridcell>2</gridcell>
        <gridcell><bullets><point>not renewable</point><point>fine from the first day</point></bullets></gridcell>
      </gridline>
      <gridline>
        <gridcell>Reference</gridcell><gridcell>0</gridcell>
        <gridcell>consulted on the premises only</gridcell>
      </gridline>
    </gridholder>
  </ruleblock>
  <ruleblock idx="2.">
    <blockcaption>Fines</blockcaption>
    <statementbody>The fine is the number of days overdue multiplied by the daily rate for the item class.</statementbody>
    <statementbody>The fine for one item is capped at eight units.</statementbody>
    <statementbody>No fine is charged where the overdue period is two days or fewer for a standard loan.</statementbody>
    <sidenote>Librarians may waive a fine; a waiver is recorded but does not change the calculation.</sidenote>
  </ruleblock>
  <ruleblock idx="3.">
    <blockcaption>Renewals</blockcaption>
    <statementbody>A standard loan may be renewed twice, each renewal running for a further fourteen days.</statementbody>
    <statementbody>A loan with a reservation against it may not be renewed.</statementbody>
  </ruleblock>
</lendingrules>
"""

G_R = '''#\' Loan period in days for an item class
#\'
#\' @param item_class one of "Standard", "Short" or "Reference"
#\' @return the loan period in days
#\' @export
loan_days <- function(item_class) {
  periods <- c(Standard = 14, Short = 2, Reference = 0)
  unname(periods[item_class])
}

#\' Fine for an overdue item
#\'
#\' @param days_overdue whole number of days overdue
#\' @param item_class one of "Standard", "Short" or "Reference"
#\' @param daily_rate rate charged per day
#\' @return the fine, capped at eight units
#\' @export
fine_due <- function(days_overdue, item_class, daily_rate) {
  free <- if (item_class == "Standard") 2 else 0
  charged <- max(0, days_overdue - free)
  min(charged * daily_rate, 8)
}

#\' Whether a loan may be renewed
#\'
#\' @param item_class one of "Standard", "Short" or "Reference"
#\' @param renewals_so_far how many renewals have already been made
#\' @param reserved whether another borrower has reserved the item
#\' @export
may_renew <- function(item_class, renewals_so_far, reserved) {
  item_class == "Standard" && renewals_so_far < 2 && !reserved
}
'''

G_DOCUMENTATION = [
    ("h", 1, "Lending calculator"),
    ("p", "The package works out loan periods, fines and whether a loan may be renewed."),
    ("h", 2, "Loan periods"),
    ("p", "A standard loan runs for fourteen days. A short loan runs for two days."),
    ("table", "Table 1. Loan period by item class", [["Item class", "Days"], ["Standard", "14"], ["Short", "2"], ["Reference", "0"]]),
    ("h", 2, "Fines"),
    ("p", "The fine is the days overdue times the daily rate, capped at eight units. A standard loan has two free days."),
    ("h", 2, "Renewals"),
    ("p", "A standard loan may be renewed twice unless another borrower has reserved the item."),
]

# expected_level is the DEPTH of the heading a unit sits under, not the number the document
# prints beside it. Every block of G is a top-level section, so its statements sit at depth 1.
G_GOLD = [
    ("A standard loan runs for fourteen days", "a statement in a tag the rules do not know", "1"),
    ("Loan periods", "a heading in a tag the rules do not know", "1"),
    ("renewable twice", "a list item inside a table cell", ""),
    ("no fine in the first two days", "a second list item inside the same table cell", ""),
    ("consulted on the premises only", "a table cell holding plain text beside cells holding lists", ""),
    ("capped at eight units", "a statement in the second block", "1"),
    ("Librarians may waive a fine", "a side note in a tag that appears only once", ""),
    ("A loan with a reservation against it may not be renewed", "the last statement of the last block", "1"),
]


def build_g_schema():
    import pandas
    sample = "G_schema"
    write_bytes(sample, "canon", "lending_rules.xml", G_METHODOLOGY.encode("utf-8"))
    rates = pandas.DataFrame({"item_class": ["Standard", "Short", "Reference"], "daily_rate": [0.20, 0.50, 0.00]})
    files = {"DESCRIPTION": "Package: lendcalc\nTitle: Lending Calculator\nVersion: 0.3.0\nLicense: MIT\nEncoding: UTF-8\n",
             "NAMESPACE": "export(loan_days)\nexport(fine_due)\nexport(may_renew)\n",
             "R/lending.R": G_R,
             "R/data.R": "#' Daily rate by item class\n#'\n#' @format A data frame with 3 rows and 2 columns.\n\"daily_rates\"\n",
             "data/daily_rates.rda": r_data({"daily_rates": rates}),
             "man/loan_days.Rd": rd_page("loan_days", "Loan period in days", [("item_class", None, "one of Standard, Short or Reference")], "R/lending.R"),
             "man/fine_due.Rd": rd_page("fine_due", "Fine for an overdue item", [("days_overdue", None, "whole number of days overdue"), ("item_class", None, "the item class"), ("daily_rate", None, "rate charged per day")], "R/lending.R")}
    write_bytes(sample, "package", "lendcalc_0.3.0.tar.gz", tarball("lendcalc", files))
    write_bytes(sample, "doc", "lendcalc_documentation.docx", docx_file(G_DOCUMENTATION))
    write_gold(sample, G_GOLD)


# ------------------------------------------------------------------ H_twocolumn: a PDF laid out the hard way
def two_column_pdf(title, running_header, footer, columns, annex):
    """A PDF in two columns with a running header and footer on every page and a footnote at the
    foot of the first, followed by an annex whose heading carries no number."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    note = ParagraphStyle("note", parent=body, fontSize=7, leading=9)
    head = ParagraphStyle("hd", parent=styles["Heading2"], spaceBefore=6)
    raw = io.BytesIO()

    def furniture(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(20 * mm, A4[1] - 12 * mm, running_header)
        canvas.drawString(20 * mm, 10 * mm, "%s - page %d" % (footer, document.page))
        canvas.restoreState()

    document = BaseDocTemplate(raw, pagesize=A4, invariant=1,
                               leftMargin=20 * mm, rightMargin=20 * mm, topMargin=22 * mm, bottomMargin=18 * mm)
    width = (document.width - 8 * mm) / 2
    frames = [Frame(document.leftMargin, document.bottomMargin, width, document.height, id="left"),
              Frame(document.leftMargin + width + 8 * mm, document.bottomMargin, width, document.height, id="right")]
    document.addPageTemplates([PageTemplate(id="two", frames=frames, onPage=furniture)])
    story = [Paragraph(title, styles["Heading1"]), Spacer(1, 4)]
    for heading, paragraphs in columns:
        story.append(Paragraph(heading, head))
        for text in paragraphs:
            story.append(Paragraph(text, body))
            story.append(Spacer(1, 3))
    story.append(Spacer(1, 6))
    story.append(Paragraph("__________", note))
    story.append(Paragraph("1 " + annex["footnote"], note))
    story.append(Spacer(1, 8))
    story.append(Paragraph(annex["heading"], head))
    for text in annex["paragraphs"]:
        story.append(Paragraph(text, body))
        story.append(Spacer(1, 3))
    document.build(story)
    return raw.getvalue()


H_COLUMNS = [
    ("1. Watering windows", [
        "Watering runs only between four and seven in the morning, when evaporation is lowest.",
        "A bed is watered at most once a day. A second run on the same day is refused.",
        "Where the forecast gives more than five millimetres of rain, the run is skipped.<sup>1</sup>"]),
    ("2. Volume per bed", [
        "The volume for a bed is its area in square metres times the depth in millimetres set for its crop group.",
        "The depth is four millimetres for leaf crops, six for root crops and nine for fruiting crops.",
        "The volume for one bed is capped at two hundred litres in a single run."]),
    ("3. Sensors", [
        "A bed with a moisture reading above sixty per cent is skipped whatever the schedule says.",
        "Where a sensor has not reported for two days its bed is watered on the schedule alone."]),
]

H_ANNEX = {
    "footnote": "Rain is taken from the forecast issued at midnight, not from the gauge, because the run is planned before dawn.",
    "heading": "Annex: how a skipped run is recorded",
    "paragraphs": [
        "A skipped run is written to the log with the reason, so that a dry bed can be told from an unwatered one.",
        "The log keeps skipped runs for ninety days."],
}

H_R = '''#\' Volume of water for one bed
#\'
#\' @param area_m2 the area of the bed in square metres
#\' @param crop_group one of "leaf", "root" or "fruiting"
#\' @return the volume in litres, capped at two hundred
#\' @export
bed_volume <- function(area_m2, crop_group) {
  depths <- c(leaf = 4, root = 6, fruiting = 9)
  min(area_m2 * unname(depths[crop_group]), 200)
}

#\' Whether a bed is watered on a given run
#\'
#\' @param moisture_pct the latest moisture reading, as a percentage
#\' @param forecast_mm rain in the forecast, in millimetres
#\' @param watered_today whether the bed has already been watered today
#\' @export
water_bed <- function(moisture_pct, forecast_mm, watered_today) {
  !watered_today && moisture_pct <= 60 && forecast_mm <= 5
}
'''

H_METHODOLOGY = """<?xml version="1.0" encoding="utf-8"?>
<document>
  <section num="1."><title>Watering windows</title>
    <para>Watering runs only between four and seven in the morning.</para>
    <para>A bed is watered at most once a day.</para>
    <para>Where the forecast gives more than five millimetres of rain, the run is skipped.</para>
  </section>
  <section num="2."><title>Volume per bed</title>
    <para>The volume for a bed is its area in square metres times the depth in millimetres set for its crop group.</para>
    <para>The depth is four millimetres for leaf crops, six for root crops and nine for fruiting crops.</para>
    <para>The volume for one bed is capped at two hundred litres in a single run.</para>
  </section>
  <section num="3."><title>Sensors</title>
    <para>A bed with a moisture reading above sixty per cent is skipped whatever the schedule says.</para>
  </section>
</document>
"""

# The guide's own title sits at depth 1, so every numbered section of it sits at depth 2. The
# methodology says some of the same sentences in the same words, so the rows about the PDF name
# the PDF: without that, the measurement found the methodology's unit first and scored the PDF's
# heading depth against a file that has no title above its sections.
H_GOLD = [
    ("evaporation is lowest", "the first column of a two-column page", "2", "irrigate_field_guide.pdf"),
    ("capped at two hundred litres", "the second column of a two-column page", "2", "irrigate_field_guide.pdf"),
    ("watered on the schedule alone", "the last paragraph before the footnote", "2", "irrigate_field_guide.pdf"),
    ("not from the gauge", "a footnote at the foot of the page", ""),
    ("Annex: how a skipped run is recorded", "a heading carrying no number at all", ""),
    ("a dry bed can be told from an unwatered one", "a paragraph under the unnumbered annex heading", ""),
    ("keeps skipped runs for ninety days", "the last paragraph of the annex", ""),
]


def build_h_twocolumn():
    import pandas
    sample = "H_twocolumn"
    write_bytes(sample, "canon", "irrigation_method.xml", H_METHODOLOGY.encode("utf-8"))
    depths = pandas.DataFrame({"crop_group": ["leaf", "root", "fruiting"], "depth_mm": [4.0, 6.0, 9.0]})
    files = {"DESCRIPTION": "Package: irrigate\nTitle: Irrigation Schedule\nVersion: 0.4.0\nLicense: MIT\nEncoding: UTF-8\n",
             "NAMESPACE": "export(bed_volume)\nexport(water_bed)\n",
             "R/irrigation.R": H_R,
             "R/data.R": "#' Watering depth by crop group\n#'\n#' @format A data frame with 3 rows and 2 columns.\n\"crop_depths\"\n",
             "data/crop_depths.rda": r_data({"crop_depths": depths}),
             "man/bed_volume.Rd": rd_page("bed_volume", "Volume of water for one bed", [("area_m2", None, "area of the bed in square metres"), ("crop_group", None, "one of leaf, root or fruiting")], "R/irrigation.R")}
    write_bytes(sample, "package", "irrigate_0.4.0.tar.gz", tarball("irrigate", files))
    write_bytes(sample, "doc", "irrigate_field_guide.pdf",
                two_column_pdf("Irrigation field guide", "Irrigation field guide", "Greenhouse operations", H_COLUMNS, H_ANNEX))
    write_gold(sample, H_GOLD)


# ------------------------------------------------------------------ I_wordtraps: a Word file that hides things
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def bold_paragraph(document, text):
    """A heading made of nothing but bold: no outline level, no heading style, no number. This
    is how a great many real documents mark a section, and AIVA cannot see it as a heading from
    the style, only from the way the paragraph is set."""
    paragraph = document.add_paragraph()
    run_in = paragraph.add_run(text)
    run_in.bold = True
    run_in.font.size = None
    return paragraph


def text_box(document, text):
    """A paragraph holding a text box. Word puts the words inside w:txbxContent, which sits in
    the body part but not in the run of paragraphs a plain reader walks."""
    from docx.oxml import parse_xml
    paragraph = document.add_paragraph()
    paragraph._p.append(parse_xml(
        '<w:r xmlns:w="%s"><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml" style="width:200pt;height:40pt">'
        '<v:textbox><w:txbxContent><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:txbxContent></v:textbox>'
        '</v:shape></w:pict></w:r>' % (W, text)))
    return paragraph


def tracked_paragraph(document, kept, inserted, deleted):
    """A paragraph carrying an unaccepted insertion and an unaccepted deletion."""
    from docx.oxml import parse_xml
    paragraph = document.add_paragraph()
    paragraph._p.append(parse_xml('<w:r xmlns:w="%s"><w:t xml:space="preserve">%s </w:t></w:r>' % (W, kept)))
    paragraph._p.append(parse_xml(
        '<w:ins xmlns:w="%s" w:id="901" w:author="A" w:date="2026-01-01T00:00:00Z">'
        '<w:r><w:t xml:space="preserve">%s </w:t></w:r></w:ins>' % (W, inserted)))
    paragraph._p.append(parse_xml(
        '<w:del xmlns:w="%s" w:id="902" w:author="A" w:date="2026-01-01T00:00:00Z">'
        '<w:r><w:delText xml:space="preserve">%s</w:delText></w:r></w:del>' % (W, deleted)))
    return paragraph


def word_traps_docx():
    import docx
    document = docx.Document()
    document.add_heading("Shift pay calculator", level=1)
    document.add_paragraph("The package works out what a courier is paid for a shift.")
    bold_paragraph(document, "Base pay")
    document.add_paragraph("A courier is paid a base rate of nine units for every hour of the shift.")
    document.add_paragraph("A shift shorter than two hours is paid as two hours.")
    text_box(document, "Rounding note: hours are rounded up to the nearest quarter of an hour before pay is worked out.")
    bold_paragraph(document, "Distance pay")
    document.add_paragraph("Distance pay is the kilometres ridden times the distance rate for the zone.")
    document.add_table(rows=1, cols=1)
    tracked_paragraph(document,
                      "The distance rate is",
                      "zero point four units a kilometre in the outer zone and",
                      "zero point two units a kilometre everywhere.")
    document.add_paragraph("Distance pay for one shift is capped at forty units.")
    bold_paragraph(document, "Night premium")
    document.add_paragraph("A shift beginning after ten at night earns a premium of one quarter of the base pay.")
    footnote_like = document.add_paragraph()
    footnote_like.add_run("The premium is worked out on base pay only and never on distance pay.").italic = True
    raw = io.BytesIO()
    document.save(raw)
    return fixed_zip(raw.getvalue())


I_R = '''#\' Base pay for a shift
#\'
#\' @param hours length of the shift in hours
#\' @param base_rate pay for one hour
#\' @return the base pay, with a minimum of two hours
#\' @export
base_pay <- function(hours, base_rate) {
  max(hours, 2) * base_rate
}

#\' Distance pay for a shift
#\'
#\' @param km kilometres ridden
#\' @param zone one of "inner" or "outer"
#\' @return the distance pay, capped at forty units
#\' @export
distance_pay <- function(km, zone) {
  rates <- c(inner = 0.2, outer = 0.4)
  min(km * unname(rates[zone]), 40)
}

#\' Night premium on a shift
#\'
#\' @param base the base pay for the shift
#\' @param starts_after_ten whether the shift began after ten at night
#\' @export
night_premium <- function(base, starts_after_ten) {
  if (starts_after_ten) base * 0.25 else 0
}
'''

I_METHODOLOGY = """<?xml version="1.0" encoding="utf-8"?>
<document>
  <section num="1."><title>Base pay</title>
    <para>A courier is paid a base rate of nine units for every hour of the shift.</para>
    <para>A shift shorter than two hours is paid as two hours.</para>
  </section>
  <section num="2."><title>Distance pay</title>
    <para>Distance pay is the kilometres ridden times the distance rate for the zone.</para>
    <para>The distance rate is zero point four units a kilometre in the outer zone.</para>
    <para>Distance pay for one shift is capped at forty units.</para>
  </section>
  <section num="3."><title>Night premium</title>
    <para>A shift beginning after ten at night earns a premium of one quarter of the base pay.</para>
  </section>
</document>
"""

I_GOLD = [
    ("Base pay", "a heading made of nothing but bold", "1"),
    ("base rate of nine units", "the first paragraph under a bold-only heading", "1"),
    ("rounded up to the nearest quarter of an hour", "a text box, which is not in the run of paragraphs", ""),
    ("zero point four units a kilometre in the outer zone", "an unaccepted insertion", ""),
    ("zero point two units a kilometre everywhere", "an unaccepted deletion, which says the opposite", ""),
    ("capped at forty units", "a paragraph after a tracked change", ""),
    ("Night premium", "the third bold-only heading", "1"),
    ("never on distance pay", "an italic paragraph standing in for a footnote", ""),
]


def build_i_wordtraps():
    import pandas
    sample = "I_wordtraps"
    write_bytes(sample, "canon", "shift_pay_method.xml", I_METHODOLOGY.encode("utf-8"))
    rates = pandas.DataFrame({"zone": ["inner", "outer"], "rate_per_km": [0.20, 0.40]})
    files = {"DESCRIPTION": "Package: shiftpay\nTitle: Shift Pay Calculator\nVersion: 0.5.0\nLicense: MIT\nEncoding: UTF-8\n",
             "NAMESPACE": "export(base_pay)\nexport(distance_pay)\nexport(night_premium)\n",
             "R/shiftpay.R": I_R,
             "R/data.R": "#' Distance rate by zone\n#'\n#' @format A data frame with 2 rows and 2 columns.\n\"zone_rates\"\n",
             "data/zone_rates.rda": r_data({"zone_rates": rates}),
             "man/base_pay.Rd": rd_page("base_pay", "Base pay for a shift", [("hours", None, "length of the shift in hours"), ("base_rate", None, "pay for one hour")], "R/shiftpay.R"),
             "man/distance_pay.Rd": rd_page("distance_pay", "Distance pay for a shift", [("km", None, "kilometres ridden"), ("zone", None, "one of inner or outer")], "R/shiftpay.R")}
    write_bytes(sample, "package", "shiftpay_0.5.0.tar.gz", tarball("shiftpay", files))
    write_bytes(sample, "doc", "shiftpay_documentation.docx", word_traps_docx())
    write_gold(sample, I_GOLD)


def write_gold(sample, rows):
    """gold_reading.csv: a phrase that must end in some unit, where it sits in the file, the
    heading level the unit should carry where the file makes that plain, and - where the same
    words sit in more than one file - which file the row is about. It is READING gold: it
    says nothing about links, checks or statuses."""
    path = os.path.join(SAMPLES, sample, "gold_reading.csv")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(("must_appear", "where_it_sits", "expected_level", "in_file"))
        writer.writerows(tuple(row) + ("",) * (4 - len(row)) for row in rows)

BUILDERS.update({"G_schema": build_g_schema, "H_twocolumn": build_h_twocolumn, "I_wordtraps": build_i_wordtraps})

if __name__ == "__main__":
    for name in (sys.argv[1:] or list(BUILDERS)):
        BUILDERS[name]()
        print("built", name)
