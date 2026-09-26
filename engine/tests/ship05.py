import os, sys; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__))); from paths import TESTS, ENGINE, REPO, SAMPLES, NOTEBOOK, FLOWR_ARCHIVES
import os, sys, json, shutil, glob, zipfile, threading, io, contextlib
import tempfile
home = tempfile.mkdtemp(prefix="ship05-")         # the notebook's folder for this run: a fresh one each time
shutil.rmtree(home, ignore_errors=True); os.makedirs(home)
repo = REPO
shutil.copytree(os.path.join(repo, "engine"), os.path.join(home, "engine"), dirs_exist_ok=True); shutil.copy(os.path.join(repo, "Verifier.ipynb"), home)
for staged in glob.glob(os.path.join(FLOWR_ARCHIVES or "/nonexistent", "flowr-*.tar.gz")): shutil.copy(staged, home)
os.chdir(home)
os.environ["ENGINE"] = os.path.join(home, "engine")
sys.path.insert(0, TESTS)
from harness import Gateway, TOKEN, DBUtils, SETTINGS
import verifier
assert verifier.__file__.startswith(home), verifier.__file__
verifier.notebook_folder = lambda d: home
cells = ["".join(c["source"]) for c in json.load(open("Verifier.ipynb"))["cells"] if c["cell_type"] == "code"]
dbutils = DBUtils(); dbutils.widgets.values.update({"project_name": "SHIP"})
class Chain:
    def __getattr__(self, name): return self
    def __call__(self, *a): return self
    def get(self): return "reviewer@example.com"
class Library:
    def restartPython(self): print("(restartPython)")
dbutils.notebook, dbutils.library = Chain(), Library()
namespace = {"dbutils": dbutils}
print("=== cell 1"); exec(cells[0], namespace)
SETTINGS[0] = verifier.notebook_settings()
gateway = Gateway(delay=(0.01, 0.05))
print("=== cell 2"); verifier.check_chat(gateway)
shutil.copytree(SAMPLES + "/F_capital/Inputs", os.path.join(home, "Projects", "SHIP"), dirs_exist_ok=True)
print("=== cell 3"); exec(cells[2], namespace)
print("=== cell 4"); exec(cells[3], namespace)
paths = verifier.NOTEBOOK["paths"]
found = []
for path in glob.glob(os.path.join(paths.run_dir, "**", "*"), recursive=True):
    if os.path.isfile(path):
        data = open(path, "rb").read()
        blobs = [data] + ([zipfile.ZipFile(path).read(n) for n in zipfile.ZipFile(path).namelist()] if zipfile.is_zipfile(path) else [])
        found += [path for b in blobs if TOKEN.encode() in b]
print("\nTOKEN in the run folder, inside archives too:", found or "nowhere")
print("widget reads from threads other than the main one:", sorted(t for t in dbutils.widgets.threads if t != "MainThread") or "none")
print("most chat() calls at once:", gateway.most, "| questions over their room:", len(gateway.over_room))
import openpyxl
book = openpyxl.load_workbook(os.path.join(paths.run_dir, verifier.OUTPUT_FILE), read_only=True)
rows = list(book["Chunks_Model"].iter_rows(values_only=True))
withheld = [(r[0], i) for r in rows[1:] for i, v in enumerate(r) if v == verifier.CELL_WITHHELD]
print("sheets:", book.sheetnames, "| Chunks_Model columns:", len(rows[0]), "| withheld cells:", withheld or "none")
print("audit log size: %.0f KB" % (os.path.getsize(os.path.join(paths.audit_dir, "Audit_Log.xlsx")) / 1024))
