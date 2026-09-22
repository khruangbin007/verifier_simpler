"""Tests of verifier2_package: safe unpacking, the R reader, documentation units and stored data."""
import io
import os
import tarfile
import unittest

import helpers
import core
import reading
import review
import runner
import build_samples

R_SNIPPETS = (          # what real R code looks like; every snippet must parse without a "not read" part
    "f <- function(x, ...) UseMethod('f')", "x <- if (a > b) a else b", "for (i in seq_len(n)) { s <- s + i }",
    "while (TRUE) { break }", "repeat { next }", "y <- x %>% f() %>% g(2)", "y <- x |> f()", "`my var` <- 3", "a$b$c <- d[[1]][2, 'k']",
    "f <- function(a = 1, b = c(1, 2)) { a + b[1] }", "z <- -x^2 + 3e-4 * 0x1F - 1L", "g <- \\(x) x + 1", "s <- r\"(raw \\ text)\"",
    "m <- x %*% t(x) %in% y", "if (is.null(x) && !y || z) stop('no')", "k <- function(x) {\n  # a comment\n  x[x > 0 & !is.na(x)]\n}",
    "lst <- list(a = 1, `b c` = 2)[['a']]", "pkg::fun(x)@slot <- pkg:::hidden(y)", "x[-1] <- NA_real_; y = NULL", "function(x) -x",
    "h <- function(x) {\n  y <- x * 2\n  z <- y +\n    3\n  z / 4\n}", "~ a + b | c", "y ~ x", "a <<- b -> c", "t <- tryCatch(f(x), finally = cat('done'))")


class UnpackingSafely(unittest.TestCase):
    def tarball_with(self, members):
        raw = io.BytesIO()
        with tarfile.open(fileobj=raw, mode="w:gz") as archive:
            for name, kind, data in members:
                info = tarfile.TarInfo(name)
                info.type, info.size = kind, len(data)
                if kind in (tarfile.SYMTYPE, tarfile.LNKTYPE):
                    info.linkname = "/etc/passwd"
                archive.addfile(info, io.BytesIO(data))
        path = helpers.os.path.join(helpers.scratch(), "pkg.tar.gz")
        with open(path, "wb") as handle:
            handle.write(raw.getvalue())
        return path

    def test_paths_that_leave_the_package_links_and_devices_are_refused_and_nothing_touches_disk(self):
        path = self.tarball_with([("pkg/R/a.R", tarfile.REGTYPE, b"x <- 1\n"), ("pkg/../../evil.R", tarfile.REGTYPE, b"y <- 2\n"),
                                  ("/abs/evil.R", tarfile.REGTYPE, b"y <- 2\n"), ("pkg/link", tarfile.SYMTYPE, b""), ("pkg/hard", tarfile.LNKTYPE, b"")])
        files, refused = reading.unpack_package(path, 10_000_000)
        self.assertEqual(sorted(files), ["pkg/R/a.R"])
        self.assertEqual(len(refused), 4)

    def test_a_member_larger_than_the_cap_is_refused(self):
        path = self.tarball_with([("pkg/R/big.R", tarfile.REGTYPE, b"x" * 5000)])
        files, refused = reading.unpack_package(path, 1000)
        self.assertEqual((files, len(refused)), ({}, 1))


class ReadingR(unittest.TestCase):
    def test_snippet_corpus_parses_completely(self):
        for snippet in R_SNIPPETS:
            parts = reading.parse_r_source(snippet)
            self.assertFalse([p for p in parts if isinstance(p, tuple)], "not read: %r" % snippet)

    def test_a_broken_expression_is_kept_as_not_read_and_the_rest_is_still_parsed(self):
        units, _ = helpers.units_of({"R/a.R": "good <- function(x) x * 2.5\n\nbad <- function(x) { x +* }\n\nalso_good <- function(y) y / 4.5\n"})
        kinds = [(u["kind"], u["name"]) for u in units if u["file"] == "R/a.R"]
        self.assertIn((core.KIND_FUNCTION, "good"), kinds)
        self.assertIn((core.KIND_FUNCTION, "also_good"), kinds)
        self.assertTrue(any(u["kind"] == core.KIND_NOT_READ for u in units))

    def test_every_non_blank_line_lies_inside_a_unit(self):
        units, result = helpers.units_of(build_samples.f_package(False))
        info = core.to_plain(result.records["package_info"][0])
        for entry in info["files"]:
            covered = {n for u in units if u["file"] == entry["file"] and u.get("lines") for n in range(u["lines"][0], u["lines"][1] + 1)}
            self.assertFalse(set(entry.get("nonblank_lines", [])) - covered, entry["file"])

    def test_straight_line_functions_are_composed_so_that_a_sign_above_the_return_is_seen(self):
        units, _ = helpers.units_of({"R/a.R": "f <- function(pd, rho, q = 0.999) {\n  a <- qnorm(pd)\n  b <- sqrt(rho) * qnorm(q)\n  pnorm((a - b) / sqrt(1 - rho))\n}\n"})
        function = next(u for u in units if u["kind"] == core.KIND_FUNCTION)
        text = core.expr_to_text(core.expr_from_dict(function["code"]["composed"]))
        self.assertEqual(text, "f = normal_cdf((normal_inverse(pd) - sqrt(rho) * normal_inverse(q)) / sqrt(1 - rho))")
        self.assertEqual([u["name"] for u in units if u["kind"] == core.KIND_FORMULA], ["a", "b", "f"])

    def test_supporting_code_by_syntax_never_holds_arithmetic_or_a_non_trivial_number(self):
        units, _ = helpers.units_of({"R/a.R": "check <- function(x) {\n  stopifnot(is.numeric(x))\n  invisible(TRUE)\n}\n\n"
                                              "keep <- function(pd, floor = 0.0003) {\n  pd\n}\n\nrate <- function(x) {\n  x * 1.5\n}\n"})
        plumbing = {u["name"]: bool(u["code"]["plumbing"]) for u in units if u["kind"] == core.KIND_FUNCTION}
        self.assertEqual(plumbing, {"check": True, "keep": False, "rate": False})

    def test_roxygen_blocks_help_pages_tests_and_vignettes_become_units(self):
        units, _ = helpers.units_of(build_samples.f_package(False))
        kinds = {kind: sum(1 for u in units if u["kind"] == kind) for kind in core.UNIT_KINDS}
        self.assertEqual((kinds[core.KIND_ROXYGEN], kinds[core.KIND_HELP], kinds[core.KIND_TEST], kinds[core.KIND_VIGNETTE]), (7, 6, 2, 2))
        block = next(u for u in units if u["kind"] == core.KIND_ROXYGEN and u["name"] == "cond_pd")
        self.assertEqual([t["name"] for t in block["roxygen"]["tags"] if t["tag"] == "param"], ["pd", "rho", "q"])
        self.assertTrue(block["roxygen"]["formulas"][0]["readable"])
        target = next(u for u in units if u["ref"] == block["roxygen"]["documents_ref"])
        self.assertEqual((target["kind"], target["name"], target["code"]["exported"]), (core.KIND_FUNCTION, "cond_pd", True))
        stale = next(u for u in units if u["kind"] == core.KIND_HELP and u["name"] == "cond_pd")
        self.assertIs(stale["helppage"]["in_step_with_source"], False)


class StoredData(unittest.TestCase):
    def test_rda_and_rds_are_decoded_without_r_and_profiled(self):
        units, result = helpers.units_of(build_samples.f_package(False))
        tables = {core.to_plain(t)["object_name"]: core.to_plain(t) for t in result.records["parameter_tables"]}
        floors = tables["lgd_floors"]
        self.assertEqual(floors["header"], ["segment", "lgd_floor"])
        self.assertEqual([row[1] for row in floors["rows"]], ["0.15", "0.25", "0.45"])
        self.assertEqual(floors["row_key"], ["segment"])
        self.assertIn("confidence", tables)

    def test_the_value_hash_ignores_how_the_file_was_written(self):
        import pandas
        frame = pandas.DataFrame({"zone": ["A", "B"], "rate": [0.5, 1.25]})
        first, _ = helpers.units_of({"data/t.rda": build_samples.r_data({"t": frame})})
        second, _ = helpers.units_of({"inst/extdata/t.rds": build_samples.r_data(frame, single=True)})
        hashes = [next(u for u in units if u.get("data"))["data"]["canonical_value_hash"] for units in (first, second)]
        self.assertEqual(hashes[0], hashes[1])

    def test_reads_of_stored_data_are_recognised(self):
        units, _ = helpers.units_of(build_samples.f_package(False))
        function = next(u for u in units if u["kind"] == core.KIND_FUNCTION and u["name"] == "capital_k")
        self.assertEqual([r["object"] for r in function["code"]["reads_data"]], ["lgd_floors"])

    def test_a_data_file_that_is_not_r_data_is_recorded_as_not_assessed(self):
        units, _ = helpers.units_of({"data/odd.rda": b"this is not R data at all"})
        odd = next(u for u in units if u["file"] == "data/odd.rda")
        self.assertTrue(odd["kind"] == core.KIND_NOT_READ or not odd["data"]["assessable"])


class ImplementationMapSample(unittest.TestCase):
    """Stage 1 of the Model Implementation Map: a package shaped like the real ones - one R file among
    many help pages, dplyr pipelines, a purrr lambda, do.call, recursion, dead code - and an answer key
    that says what a complete map must hold. The baseline says what today's reader already recovers."""

    @classmethod
    def setUpClass(cls):
        import yaml
        import build_samples
        with open(os.path.join(helpers.SAMPLES_DIR, "J_pipeline", "gold_map.yaml"), encoding="utf-8") as handle:
            cls.gold = yaml.safe_load(handle)
        cls.source, cls.methodology = build_samples.J_R_SOURCE, build_samples.J_METHODOLOGY

    def test_every_name_in_the_answer_key_is_written_in_the_code(self):
        gold = self.gold
        names = set(gold["final_outputs"] + gold["not_reached"]) | set(gold["calls"]) | {c for cs in gold["calls"].values() for c in cs}
        names |= {v for flows in gold["flows"].values() for v in flows} | set(gold["columns"])
        names |= {s for flows in gold["flows"].values() for ss in flows.values() for s in ss if "/" not in s}
        names |= {s for ss in gold["columns"].values() for s in ss}
        raw = gold["raw_inputs"]
        names |= set(raw["argument"] + raw["column_of_an_argument"] + raw["stored_data"] + raw["hard_coded_number"])
        for name in sorted(names):
            self.assertIn(name, self.source, "the answer key names '%s', which the code does not write" % name)
        for gap in gold["gaps"]:
            self.assertIn(gap["code"], self.source)
        for heading in gold["methodology_not_implemented"]:
            self.assertIn(heading, self.methodology)

    def test_the_sample_is_read_as_an_r_package_with_every_function(self):
        path = os.path.join(helpers.SAMPLES_DIR, "J_pipeline", "Inputs", "2_Model_Package", "harbourscore_1.2.0.tar.gz")
        result = reading.read_package(helpers.context_for({"package": [path]}))
        self.assertFalse([m for m in result.messages if "R package" in m])
        functions = {unit.name for unit in result.records["model_units"] if unit.kind == core.KIND_FUNCTION}
        expected = set(self.gold["final_outputs"] + self.gold["not_reached"]) | set(self.gold["calls"]) | {c for cs in self.gold["calls"].values() for c in cs}
        self.assertEqual(functions, expected)

    def test_the_traced_flow_holds_the_whole_answer_key(self):
        """Stage 2 is done when code alone recovers everything the answer key asks of it: the final output
        and the functions no output reaches, every call (the recursive one too), every value and column
        with what it is computed from, every raw input, and both gaps named with their code. What the map
        says of the documents is its own work, stage 5."""
        import develop
        for part, (got, total) in develop.map_measure("J_pipeline", record=False).items():
            if "stage 5" not in part:
                self.assertEqual(got, total, part)

    def flow(self, sample):
        paths, settings, _ = helpers.run_sample(sample, stop_after="05a")
        store = runner.open_store(paths, settings)
        return store.read("dataflow"), {u["ref"]: u["text"] for u in store.read("model_units")}

    def test_capital_k_resolves_to_raw_inputs_with_no_gap(self):
        """The first done-when of stage 2. pdc is set by a call into cond_pd, whose pd is what floor_pd
        returns there, whose rho is what asset_correlation returns, and whose q is its default."""
        flow, _ = self.flow("F_capital")
        leaves, loops, reached = reading.walk_dataflow(flow, "capital_k")
        nodes = {r["node"]: r for r in flow if r["record_type"] == "node"}
        self.assertEqual(reached, {"capital_k", "cond_pd", "floor_pd", "asset_correlation"})
        self.assertEqual({nodes[leaf]["kind"] for leaf in leaves}, {"argument", "stored data", "number"})
        self.assertEqual({nodes[leaf]["name"] for leaf in leaves if nodes[leaf]["kind"] == "argument"}, {"pd", "lgd", "segment"})
        self.assertFalse(loops)
        self.assertFalse([g for g in flow if g["record_type"] == "gap" and g["function"] in reached])
        call = next(r for r in nodes.values() if r["kind"] == "call" and r["callee"] == "cond_pd" and r["function"] == "capital_k")
        self.assertEqual([nodes[s]["callee"] for s in call["bindings"]["pd"] + call["bindings"]["rho"]], ["floor_pd", "asset_correlation"])
        self.assertEqual(call["bindings"]["q"], ["default"])

    def test_every_source_is_a_node_and_every_line_of_code_is_as_written(self):
        for sample in ("J_pipeline", "F_capital"):
            flow, text = self.flow(sample)
            nodes = {r["node"] for r in flow if r["record_type"] == "node"}
            for record in flow:
                if record["record_type"] == "node":
                    for source in record["from"] + (record.get("default_from") or []):
                        self.assertIn(source, nodes, "%s: %s comes from %s, which is not a node" % (sample, record["node"], source))
                if record.get("code") and record.get("function_ref"):
                    self.assertIn(record["code"], text[record["function_ref"]], "not as written")

    def test_the_pronouns_of_dplyr_are_read_as_dplyr_reads_them(self):
        """Packages written for CRAN name columns as .data$x, to keep the checks quiet, and values of the
        function as .env$x. Without this, .data$throughput would be read as a column called .data."""
        source = ('f <- function(d, limit) {\n  d %>% dplyr::mutate(ok = .data$throughput > .env$limit, '
                  'wide = .data$berth_utilisation * 2)\n}')
        unit = {"ref": "M-0001", "kind": "Function", "name": "f", "text": source, "lines": [1, 3], "code": {"exported": True}}
        records = reading.Dataflow([unit]).run()
        nodes = {r["node"]: r for r in records if r["record_type"] == "node"}
        self.assertEqual(sorted(nodes["column:ok"]["from"]), ["column:throughput", "f:arg:limit"])
        self.assertIn("column:berth_utilisation", nodes["column:wide"]["from"])
        self.assertNotIn("column:.data", nodes)

    def test_a_join_key_only_matches_and_a_loop_keeps_what_it_is_given(self):
        """Found building stage 2: left_join(anchor_table, by = "score_band") made score_band look computed
        from the table, though only anchor is added; and the walk stopped at notch_down's recursive call
        without the arguments it is given, losing the stored table rating_scale."""
        flow, _ = self.flow("J_pipeline")
        nodes = {r["node"]: r for r in flow if r["record_type"] == "node"}
        self.assertNotIn("data:anchor_table", nodes["column:score_band"]["from"])
        self.assertIn("data:anchor_table", nodes["column:anchor"]["from"])
        leaves, loops, _ = reading.walk_dataflow(flow, "harbour_rating")
        self.assertIn("data:rating_scale", leaves)
        self.assertEqual([nodes[loop]["callee"] for loop in loops], ["notch_down"])


class FinalOutputs(unittest.TestCase):
    """Stage 3 of the implementation map: code proposes the final outputs; a person decides, in the
    yellow column of Model_Implementation_Map, read back by function name when the outline is confirmed."""

    def run_to_cell_3(self):
        import openpyxl
        projects = helpers.scratch()
        helpers.copy_sample("J_pipeline", projects, "J", "2026-09-22")
        settings = runner.make_settings({})
        paths = runner.open_run(projects, "J", "2026-09-22", scratch_root=helpers.scratch())
        runner.run_pipeline(paths, settings, stop_after="06")
        return paths, settings, os.path.join(paths.run_dir, "Output.xlsx"), openpyxl

    def decide(self, path, openpyxl, choices):
        book = openpyxl.load_workbook(path)
        sheet = book["Model_Implementation_Map"]
        header = [cell.value for cell in sheet[1]]
        at_name, at_word = header.index("Function"), header.index("Final output (your decision)")
        for row in range(2, sheet.max_row + 1):
            name = sheet.cell(row=row, column=at_name + 1).value
            if name in choices:          # .value, not cell(value=None): openpyxl's cell() ignores None, so it cannot empty a cell
                sheet.cell(row=row, column=at_word + 1).value = choices[name]
        book.save(path)

    def test_the_sheet_lists_every_function_with_what_code_proposes_and_why(self):
        _, _, path, openpyxl = self.run_to_cell_3()
        rows = list(openpyxl.load_workbook(path, read_only=True)["Model_Implementation_Map"].iter_rows(values_only=True))
        header, body = rows[0], rows[1:]
        by_name = {row[header.index("Function")]: dict(zip(header, row)) for row in body}
        self.assertEqual(len(body), 10, "every package function, once")
        self.assertEqual(body[0][header.index("Function")], "harbour_rating", "the final outputs come first")
        self.assertEqual(by_name["harbour_rating"]["Map ID"], "01")
        self.assertIn("called by a test or vignette", by_name["harbour_rating"]["Why"])
        self.assertEqual(by_name["legacy_score"]["Role"], "Not reached from any final output")
        self.assertEqual(by_name["notch_down"]["Called by"], "cap_rating")

    def test_a_decision_overrides_the_proposal_is_chained_and_an_emptied_cell_withdraws_it(self):
        paths, settings, path, openpyxl = self.run_to_cell_3()
        self.decide(path, openpyxl, {"harbour_rating": "no", "legacy_score": "yes", "notch_down": "maybe"})
        said = runner.confirm_outline(paths, settings, "analyst.one")
        self.assertIn("Final outputs: legacy_score (your decision).", said)
        self.assertIn("notch_down: 'maybe' is not one of yes / no and was ignored", said)
        store = runner.open_store(paths, settings)
        records = store.read("output_decisions")
        self.assertEqual({(r["function"], r["decision"], r["reviewer_id"]) for r in records},
                         {("harbour_rating", "no", "analyst.one"), ("legacy_score", "yes", "analyst.one")})
        self.assertTrue(core.verify_chain(records)[0], "the decisions are a hash chain, like the determinations")
        self.decide(path, openpyxl, {"harbour_rating": None, "legacy_score": None})
        said = runner.confirm_outline(paths, settings, "analyst.one")
        self.assertIn("Final outputs: harbour_rating (proposed by code", said, "emptied cells give code's proposal back")

    def test_without_a_proposal_or_a_decision_every_function_nothing_calls_stands_in(self):
        records = [{"record_type": "roots", "proposed": [], "not_reached": [], "why": {"f": "nothing in the package calls it",
                                                                                          "g": "nothing in the package calls it"}}]
        outputs, how, _ = reading.decided_outputs(records, {})
        self.assertEqual(outputs, ["f", "g"])
        self.assertIn("no final output was proposed or decided", how["f"])


class MapAgents(unittest.TestCase):
    """Stage 4, the skill map-implementation: the Tracer works only on the gaps on the path from a final
    output, one action a turn from a fixed list, every link copying the code; the Namer names the steps;
    the Auditor is code. Each turn is a question of its own, so a run replays without a model."""

    @classmethod
    def setUpClass(cls):
        import standin_chat
        cls.paths, cls.settings, _ = helpers.run_sample("J_pipeline", chat=standin_chat.chat_well_behaved, stop_after="07d")
        cls.store = runner.open_store(cls.paths, cls.settings)

    def custom_run(self, trace_answer, **settings):
        """J_pipeline to step 07d with a chat that answers trace-gap as given and everything else as the stand-in."""
        import json
        import standin_chat
        def chat(system_prompt, main_prompt, history=()):
            if "QUESTION TYPE: trace-gap" in main_prompt:
                return {"answer": json.dumps(trace_answer(main_prompt))}
            return standin_chat.chat_well_behaved(system_prompt, main_prompt)
        paths, run_settings, _ = helpers.run_sample("J_pipeline", chat=chat, settings=settings, stop_after="07d")
        store = runner.open_store(paths, run_settings)
        return store.read("map_traces"), store.read("map_audit")[0], store.read("step_names")

    def test_the_tracer_resolves_the_gap_on_the_path_and_dead_code_costs_no_call(self):
        traces, audit = self.store.read("map_traces"), self.store.read("map_audit")[0]
        self.assertEqual([t["gap"]["function"] for t in traces], ["adjust_score"], "only the gap a final output reaches")
        trace = traces[0]
        self.assertTrue(trace["status"].startswith("traced"))
        self.assertEqual([(e["value"], sorted(e["from"])) for e in trace["edges"]], [("steps", ["adjustment_table", "adjustments"])])
        unit = next(u for u in self.store.read("model_units") if u["ref"] == trace["gap"]["function_ref"])
        self.assertIn(trace["edges"][0]["quote"], unit["text"], "the link copies the code word for word")
        self.assertEqual([g["function"] for g in audit["gaps not reached"]], ["combine_results"])

    def test_every_step_on_the_path_gets_a_plain_name(self):
        audit, names = self.store.read("map_audit")[0], self.store.read("step_names")
        self.assertEqual(len(names), audit["steps"])
        for record in names:
            self.assertLessEqual(len(record["name"].split()), 12)
            self.assertEqual(core.has_banned_wording(record["name"]), "")

    def test_each_bad_action_is_refused_with_its_reason(self):
        flow, units = self.store.read("dataflow"), self.store.read("model_units")
        gap = next(g for g in flow if g["record_type"] == "gap" and g["function"] == "adjust_score")
        taken = {"action": "callers_of", "args": {"function": "adjust_score"}, "shown": "", "question_id": ""}
        question = review.trace_question({"gap": gap, "hops": [taken], "edges": [], "inputs": [], "status": "", "opened": []},
                                         review.MapTools(flow, units), self.settings)
        good = "steps <- purrr::map_dbl(adjustments, function(name) adjustment_table$points[adjustment_table$adjustment == name])"
        cases = [({"action": "guess", "args": {}}, 0), ({"action": "open_unit", "args": "M-0001"}, 0),
                 ({"action": "open_unit", "args": {"ref": "M-9999"}}, 1), ({"action": "return_of", "args": {"function": "invented"}}, 1),
                 ({"action": "columns_of", "args": {"table": "invented_table"}}, 1),
                 ({"action": "declare_edge", "args": {"value": "steps", "from": ["adjustments"], "quote": "steps <- adjustments * 2"}}, 2),
                 ({"action": "declare_edge", "args": {"value": "steps", "from": ["rating_scale"], "quote": good}}, 1),
                 ({"action": "declare_edge", "args": {"value": "steps", "from": [], "quote": good}}, 0),
                 ({"action": "declare_input", "args": {"name": "adjustments", "kind": "a guess", "quote": good}}, 0),
                 ({"action": "callers_of", "args": {"function": "adjust_score"}}, 5), ({"action": "done", "args": {}}, 0)]
        import json
        for answer, reason in cases:
            self.assertEqual(review.validate_answer(question, json.dumps(answer))[0], "rejected: " + core.REJECTION_REASONS[reason], answer)
        accepted = {"action": "declare_edge", "args": {"value": "steps", "from": ["adjustment_table", "adjustments"], "quote": good}}
        self.assertEqual(review.validate_answer(question, json.dumps(accepted))[0], "accepted")

    def test_a_run_replays_from_its_record(self):
        replayed = runner.replay_chat(self.store.read_calls())
        paths, settings, _ = helpers.run_sample("J_pipeline", chat=replayed, stop_after="07d")
        again = runner.open_store(paths, settings)
        keep = lambda traces: [(t["gap"]["function"], t["status"], [(h["action"], h["args"]) for h in t["hops"]], t["edges"]) for t in traces]
        self.assertEqual(keep(again.read("map_traces")), keep(self.store.read("map_traces")))
        self.assertEqual([(n["node"], n["name"]) for n in again.read("step_names")], [(n["node"], n["name"]) for n in self.store.read("step_names")])

    def test_the_turn_limit_ends_a_trace_that_never_ends(self):
        counter = iter(range(1000))
        traces, audit, _ = self.custom_run(lambda prompt: {"action": "statements_setting", "args": {"function": "adjust_score", "name": "x%d" % next(counter)}},
                                           map_hops_max=3)
        self.assertEqual(traces[0]["status"], "stopped at the limit of 3 turns")
        self.assertEqual(len(traces[0]["hops"]), 3)
        self.assertEqual(audit["gaps traced"], 0)

    def test_a_refused_answer_ends_the_trace_and_the_gap_stays_open(self):
        traces, audit, _ = self.custom_run(lambda prompt: {"action": "guess_the_answer", "args": {}})
        self.assertEqual(traces[0]["status"], "the AI's answer could not be used: " + core.REJECTION_REASONS[0])
        self.assertEqual([g["function"] for g in audit["gaps open"]], ["adjust_score"])

    def test_without_the_model_the_gaps_stay_named_and_nothing_is_asked(self):
        traces, audit, names = self.custom_run(lambda prompt: {}, map_with_ai=False)
        self.assertEqual(traces[0]["status"], "not asked: no model")
        self.assertEqual((audit["questions"], names), (0, []))


if __name__ == "__main__":
    unittest.main()
