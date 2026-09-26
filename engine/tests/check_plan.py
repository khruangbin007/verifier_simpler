import os, sys; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__))); from paths import TESTS, ENGINE, REPO, SAMPLES, NOTEBOOK, FLOWR_ARCHIVES
PLAN = sys.argv[1] if len(sys.argv) > 1 else os.environ.get("VERIFIER_PLAN", "Verifier_Build_Plan.md")
import sys, re, json, os; sys.path.insert(0, TESTS)
from harness import *
verifier.CHAT_BACKOFF = 0.01
plan = open(PLAN).read()
def sect(title):
    return plan.split("\n## " + title)[1].split("\n## ")[0].split("\n# ")[0]
# the pipeline and the allow-list
block = re.search(r"PIPELINE = \((.*?)\)\n```", sect("2.6 The pipeline and its runner"), re.S).group(1)
steps = [json.loads(s) for s in re.findall(r"\{[^}]*\}", block)]
print("PIPELINE as in the plan:", [{k: s[k] for k in ("id", "name", "carried_out_by")} for s in verifier.PIPELINE] == steps)
listed = re.search(r"`STEP_FUNCTIONS` is the allow-list of functions a step may name: (.*?)\. Nothing", plan).group(1)
print("STEP_FUNCTIONS as in the plan:", sorted(re.findall(r"`(\w+)`", listed)) == sorted(verifier.STEP_FUNCTIONS))
# a finished run: its record kinds, sheets, headers, info groups, checks
settings = verifier.make_settings({})
paths = open_sample("F_capital", settings)
attach(Gateway(delay=(0, 0.002)))
verifier.run_pipeline(paths, settings)
store = verifier.open_store(paths, settings)
written = {row["kind"] for row in store.rows} | set(store.account) | ({"llm_calls"} if store.calls else set())
kinds = set(re.findall(r"^\| `(\w+)` \|", sect("2.7 Records and identifiers"), re.M))
print("record kinds of the plan:", sorted(kinds), "| written by a run and not in the plan:", sorted(written - kinds) or "none", "| in the plan, never written:", sorted(kinds - written) or "none")
import openpyxl
book = openpyxl.load_workbook(os.path.join(paths.run_dir, verifier.OUTPUT_FILE), read_only=True)
table = {name: cols for name, cols in re.findall(r"^\| (Model_Package_Info|Chunks_\w+|Flagged_Items) \| (.+) \|$", sect("2.16 `Output.xlsm`"), re.M)}
for ws in book.worksheets:
    head = [c for c in next(ws.iter_rows(values_only=True))]
    print("  sheet %-22s headers as in the plan: %s" % (ws.title, ", ".join(head) == table.get(ws.title)))
print("sheets in order:", book.sheetnames == list(table))
groups_plan = re.findall(r"^\| ([A-Z][^|]+?) \| ", sect("2.16 `Output.xlsm`").split("Its rows fall into these groups")[1].split("**Every cell")[0], re.M)
groups_run = list(dict.fromkeys(r[0] for r in list(book["Model_Package_Info"].iter_rows(values_only=True))[1:]))
print("info groups of this run, all in the plan and in its order:", all(g in groups_plan for g in groups_run) and sorted(groups_run, key=groups_plan.index) == groups_run, groups_run)
checks = [what for what, _, _ in verifier.verify_evidence_pack(paths, settings, live=verifier.NOTEBOOK["live"])]
rows = re.findall(r"^\| (\d) \| (.+?) \| ", sect("2.17 Verification: cell 4"), re.M)
same = [w for (_, w) in rows]
print("checks, as many and labelled as in the plan:", len(checks) == len(rows), [c for c in checks if c not in same and not c.startswith("Re-reading")] or "all")
# settings and widgets
srows = {n: d for n, d in re.findall(r"^\| `(\w+)` \| (.+?) \| ", sect("2.18 Settings"), re.M)}
print("settings of the plan equal the allow-list:", sorted(srows) == sorted(verifier.DEFAULT_SETTINGS))
print("defaults differing:", {n: (d, verifier.DEFAULT_SETTINGS[n]) for n, d in srows.items() if n in verifier.DEFAULT_SETTINGS and d.strip("`") not in (str(verifier.DEFAULT_SETTINGS[n]), "(from Databricks)")} or "none")
wrows = re.findall(r"^\| \d\d \| `(\w+)` \| (.+?) \|$", sect("2.4 The notebook, cell by cell"), re.M)
print("widgets as in the plan:", [(n, l) for n, _, l in verifier.WIDGETS] == wrows)
kinds_plan = [k for k in re.findall(r"^\| ([A-Z][a-z]+(?:[ -][a-z]+)*) \| ", sect("2.11 Reading the model package").split("UNIT_KINDS")[1], re.M) if k != "Kind"]
print("unit kinds as in the plan:", kinds_plan == list(verifier.UNIT_KINDS), kinds_plan)
named = re.findall(r"(?i)khruangbin|verifier_simpler", open(PLAN).read())
print("the plan never names the repository:", not named)

source = open(os.path.join(ENGINE, "verifier.py")).read().split("\n")
first = next(i for i, l in enumerate(source) if l.startswith("# What cell 1 installs"))
last = next(i for i, l in enumerate(source) if l.startswith("OPTIONAL_REQUIREMENTS = "))
appendix = plan[plan.index("# Appendix C."):]
block = appendix[appendix.index("```python\n") + len("```python\n"):appendix.index("```", appendix.index("```python\n") + 10)]
print("Appendix C as REQUIREMENTS and OPTIONAL_REQUIREMENTS in verifier.py:", block.strip() == "\n".join(source[first:last + 1]).strip())
