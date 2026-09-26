"""Step 05 test harness: runs the pipeline on the sample projects with a stand-in chat() whose answers are a pure function
of the question, so the two new columns can be predicted independently, while faults are injected around it."""
import os, sys; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__))); from paths import TESTS, ENGINE, REPO, SAMPLES, NOTEBOOK, FLOWR_ARCHIVES
import os, sys, re, json, time, random, shutil, tempfile, threading, hashlib, _thread, collections
ENGINE = os.environ.get("ENGINE", ENGINE)
sys.path.insert(0, ENGINE)
import verifier, tiktoken
CL100K = tiktoken.get_encoding("cl100k_base")
try:
    from tokenizers import Tokenizer
    LLAMA = Tokenizer.from_pretrained("hf-internal-testing/llama-tokenizer")
except Exception:
    LLAMA = None

TOKEN = "SECRET-TOKEN-4417"


def h(*parts):
    return int(hashlib.sha256("|".join(parts).encode()).hexdigest(), 16)


def relevant(unit, ref):            # the oracle: which chunks bear on which unit
    return h("rel", unit, ref) % 3 == 0


def deviates(unit, ref):
    return h("dev", unit, ref) % 2 == 0


class Widgets:
    def __init__(self):
        self.values = {"llm_endpoint": "https://gateway.example", "llm_token": TOKEN, "project_name": "S"}
        self.threads = set()
    def get(self, name):
        self.threads.add(threading.current_thread().name)
        return self.values[name]
    def text(self, *a): pass
    def remove(self, *a): pass


class DBUtils:
    def __init__(self): self.widgets = Widgets()


class Gateway:
    """The stand-in for the organisation's chat(): a pure answer for each question, with faults injected around it."""
    def __init__(self, fail_rate=0.0, malformed=0.0, invent=0.0, words=0.0, expire_after=None, interrupt_after=None,
                 too_large=None, delay=(0.0, 0.02), seed=7, refuse=None, limit=None):
        self.rng, self.lock = random.Random(seed), threading.Lock()
        self.fail_rate, self.malformed, self.invent, self.words = fail_rate, malformed, invent, words
        self.expire_after, self.interrupt_after, self.too_large, self.delay = expire_after, interrupt_after, too_large, delay
        self.calls, self.active, self.most = 0, 0, 0
        self.answered = collections.Counter()      # question text hash -> answers given
        self.expired_token = None
        self.refuse, self.refused = refuse, 0
        self.limit, self.throttled = limit, 0
        self.first_calls = set()
        self.over_room = []
        self.seen_tokens = set()

    def __call__(self, system, main, history=()):
        token = verifier.live("llm_token")
        with self.lock:
            self.calls += 1
            number = self.calls
            self.active += 1
            self.most = max(self.most, self.active)
            self.seen_tokens.add(token)
            roll = [self.rng.random() for _ in range(4)]
            original = main.split("\n\nYour last")[0]
            first = original not in self.first_calls       # the question's very first call, whatever happens to it
            self.first_calls.add(original)
            if self.expire_after is not None and number > self.expire_after and self.expired_token is None:
                self.expired_token = token
        try:
            if self.limit and self.active > self.limit:            # a gateway that takes so many calls at once, no more
                with self.lock: self.throttled += 1
                raise RuntimeError("429 Client Error: Too Many Requests for url: https://gateway.example")
            time.sleep(self.rng.uniform(*self.delay))
            if system.startswith("Reply with"):
                return {"answer": "OK"}
            if self.interrupt_after is not None and number == self.interrupt_after:
                _thread.interrupt_main()
            if self.expired_token is not None and token == self.expired_token:
                raise RuntimeError("401 Unauthorized: token %s has expired" % token)
            kind = next(k for k, p in ((verifier.CODE_QUESTION, verifier.CODE_SYSTEM_PROMPT),
                                       (verifier.METHODOLOGY_SEARCH, verifier.SEARCH_SYSTEM_PROMPT),
                                       (verifier.METHODOLOGY_COMPARISON, verifier.COMPARE_SYSTEM_PROMPT)) if system == p)
            real = len(CL100K.encode(system + main))
            room = verifier.question_room(SETTINGS[0], kind)
            if real > room or (LLAMA and len(LLAMA.encode(system + main).ids) > room):
                self.over_room.append((kind, real, room))
            if self.too_large and kind != verifier.CODE_QUESTION and real > self.too_large:
                raise RuntimeError("413 request too large: %d tokens" % real)
            if self.refuse and kind != verifier.CODE_QUESTION:
                unit_ref = re.search(r"^THE PIECE: (M-\d{4})", main, re.M).group(1)
                shown_refs = list(dict.fromkeys(re.findall(r"^\[(C-\d{4})", main.split("\n\n\n")[0], re.M)))
                if self.refuse(kind, unit_ref, shown_refs):
                    with self.lock: self.refused += 1
                    raise RuntimeError("400 the gateway refused this request")
            if roll[0] < self.fail_rate:
                raise RuntimeError("503 gateway busy")
            retry = "Your last" in main
            if roll[1] < self.malformed and not retry:
                return {"answer": "Here are the chunks: C-0001 and maybe others."}
            if kind == verifier.CODE_QUESTION:
                unit = re.search(r"^Kind: (.*)$", main, re.M).group(1)
                return {"answer": "This %s computes a value. [%s]" % (unit.lower(), hashlib.sha256(main.encode()).hexdigest()[:8])}
            unit = re.search(r"^THE PIECE: (M-\d{4})", main, re.M).group(1)
            methodology = main.split("\n\n\n")[0]
            shown = list(dict.fromkeys(re.findall(r"^\[(C-\d{4})", methodology, re.M)))
            if kind == verifier.METHODOLOGY_SEARCH:
                found = [{"ref": ref, "relation": "describes", "why": "gives what %s computes" % unit} for ref in shown if relevant(unit, ref)]
                if roll[2] < self.invent and first:
                    found.append({"ref": "C-9999", "relation": "informs", "why": "invented"})
                answer = json.dumps({"relevant": found})
                return {"answer": "```json\n%s\n```" % answer if number % 5 == 0 else answer}
            deviations = [{"refs": [ref], "kind": "differs",
                           "title": "The floor of %s is 0.05%% in %s, where the methodology sets 0.03%%" % (ref, unit),
                           "methodology": "The chunk %s sets the floor at 0.03%% and says \u201cthe standard error is ignored\u201d" % ref,
                           "code": "%s floors it at 0.05%% (line %d)" % (unit, int(ref[2:]) % 40 + 1),
                           "why": "The code's floor is higher than the methodology's",
                           "effect": "Every PD below 0.05% is raised further than the methodology raises it",
                           "example": "Take a PD of 0.04%% in %s: the methodology leaves it at 0.04%%, the code raises it to 0.05%%" % unit}
                          for ref in shown if deviates(unit, ref)]
            if roll[3] < self.words and first and deviations:
                deviations[0]["code"] += " which is a major departure"
            if roll[2] < self.invent and first:
                deviations.append({"refs": ["C-9999"], "kind": "adds", "methodology": "x", "code": "y"})
            return {"answer": json.dumps({"deviations": deviations}, ensure_ascii=False)}
        finally:
            with self.lock:
                self.active -= 1


SETTINGS = [verifier.make_settings({})]


def open_sample(sample, settings):
    projects, scratch = tempfile.mkdtemp(prefix="s05-"), tempfile.mkdtemp(prefix="s05-")
    shutil.copytree(SAMPLES + "/%s/Inputs" % sample, os.path.join(projects, "S"), dirs_exist_ok=True)
    paths = verifier.open_run(projects, "S", scratch_root=scratch)
    SETTINGS[0] = settings
    return paths


def attach(gateway):
    verifier.NOTEBOOK.update(dbutils=DBUtils(), chat=gateway, user="tester@example.com", live=None)
    verifier.live("llm_token")


def expected_columns(store, settings):
    """Chunks_Model's two columns of step 05, and every row of Flagged_Items, as they must come out, worked out from
    the oracle and the methodology alone: {ref: (relevant chunks, count, [items])}."""
    units, chunks = store.read("model_units"), store.read("chunks_canon")
    out, number = {}, 0
    for unit in sorted(units, key=lambda u: u["ref"]):
        if not verifier.askable(unit):
            out[unit["ref"]] = (verifier.NOT_SEARCHED, None, [])
            continue
        refs = [c["ref"] for c in chunks if relevant(unit["ref"], c["ref"])]
        if not refs:
            out[unit["ref"]] = ("None found", 0, [])
            continue
        items = []
        for ref in [r for r in refs if deviates(unit["ref"], r)]:
            number += 1
            items.append(("F-%04d" % number, unit["ref"], ref,
                          "The code differs: The floor of %s is 0.05%% in %s, where the methodology sets 0.03%%" % (ref, unit["ref"]),
                          "The chunk %s sets the floor at 0.03%% and says \u201cthe standard error is ignored\u201d" % ref,
                          "%s floors it at 0.05%% (line %d)" % (unit["ref"], int(ref[2:]) % 40 + 1),
                          "The code's floor is higher than the methodology's",
                          "Every PD below 0.05% is raised further than the methodology raises it",
                          "Take a PD of 0.04%% in %s: the methodology leaves it at 0.04%%, the code raises it to 0.05%%" % unit["ref"]))
        out[unit["ref"]] = ("; ".join(refs), len(items), items)
    return out


def workbook_columns(paths):
    """The same, read from Output.xlsm: Chunks_Model's two columns of step 05 on each piece's first row, and the rows
    of Flagged_Items under the piece they name."""
    import openpyxl
    book = openpyxl.load_workbook(os.path.join(paths.run_dir, verifier.OUTPUT_FILE), read_only=True)
    rows = list(book["Chunks_Model"].iter_rows(values_only=True))
    head = rows[0]
    a, c = head.index("Relevant Chunks in Methodology (searched by LLM)"), head.index("Count of Flagged Items (by LLM)")
    assert "Flagged Items (by LLM, subject to human review)" not in head, "the flagged items are still on Chunks_Model"
    items = {}
    flagged = list(book["Flagged_Items"].iter_rows(values_only=True))
    assert list(flagged[0]) == ["Ref", "Location of Flagged Item", "Ref in Chunks_Methodology", "Type", "Methodology Says",
                                "Code Does", "Why Potential Flagged Item", "Effect", "Concrete Example of Potential Deviation",
                                "Decision (by Human Reviewer)", "Human Reviewer's Notes"], flagged[0]
    for row in flagged[1:]:
        items.setdefault(row[1], []).append(tuple(row[:9]))
    return {row[0]: (row[a], row[c], items.get(row[0], [])) for row in rows[1:] if "-" not in row[0][2:]}

def _run_until_done_inner(paths, settings, rounds=12, between=None):
    results = []
    for number in range(rounds):
        try:
            result = verifier.run_pipeline(paths, settings)
        except KeyboardInterrupt:
            results.append("interrupted")
            continue
        results.append(result["state"])
        if between:
            between(number, result)
        if result["state"] == "finished":
            break
    return results


def assert_no_fault(store):
    faults = [m for r in store.read("step_records") for m in r.get("messages") or () if "fault inside the tool" in m]
    assert not faults, faults


def run_until_done(paths, settings, *args, **kwargs):
    result = _run_until_done_inner(paths, settings, *args, **kwargs)
    assert_no_fault(verifier.open_store(paths, settings))
    return result
