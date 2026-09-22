"""Tests of verifier1_documents: reading methodology and documentation files of every supported form."""
import io
import textwrap
import unittest

import helpers
import core
import reading
import build_samples

UNFAMILIAR_SCHEMA = """<table>
<tablerow><tablecell>Joseph J Pezzimenti</tablecell><tablecell>+ 1 (212) 438 2038</tablecell><tablecell>New York</tablecell></tablerow>
<tablerow><tablecell>Bhavini Patel, CFA</tablecell><tablecell>(1) 416-507-2558</tablecell><tablecell>Toronto</tablecell></tablerow>
</table>
<section name=""OVERVIEW AND SCOPE"">
<paranum num=""1."" style=""Pleading Paragraph"">
These criteria apply to transportation infrastructure enterprises globally.</paranum>
<paranum num=""2."">
We generally apply the following principles when determining whether to utilize our framework under these criteria:
</paranum>
<list type=""bullet""><listitem>Mission: The entity has a public policy objective.</listitem><listitem>Motive: The entity is not commercially motivated.</listitem><listitem>Distributions: Gains may be generated, but the entity generally retains earnings.</listitem></list><paranum num=""3."" style=""Pleading Paragraph"">
If each of the above characteristics is met, a TIE is typically assessed through these criteria.</paranum>
</section>
<section name=""b) Industry Risk"">
<paranum num=""36."" style=""Pleading Paragraph"">
The moderate weight assigned to industry risk is because such enterprises have higher barriers to entry.</paranum>
<paranum num=""37."">
Industry risk is assessed by applying &#34;<marked_data hyperlink=""yes"" rev_id=""18"" type=""article"">Methodology: Industry Risk</marked_data>.&#34; See Appendix III.</paranum>
<paranum num=""38."" style=""Pleading Paragraph"">
This paragraph has been deleted.</paranum>
<paranum num=""42."" style=""Pleading Paragraph"">
The blended industry risk score is a weighted average of those industry scores.</paranum>
</section>
<section name=""c) Market Position (60% weighting) "">
<paranum num=""43."" style=""Pleading Paragraph"">
Our market position assessment focuses on the role a provider plays within its industry.</paranum>
<paranum num=""44."" style=""Pleading Paragraph"">
Table 2 includes guidance on the characteristics we typically expect to see for each assessment level by asset class.</paranum>
<table>
<tablerow>
<tablenumber width=""208"">
Table 2</tablenumber>
<tablecell width=""204""> </tablecell>
<tablecell width=""204""> </tablecell>
</tablerow>
<tablerow>
<tabletitle>
Market Position Assessment By Asset Class</tabletitle>
<tablecell> </tablecell>
<tablecell> </tablecell>
</tablerow>
<tablerow>
<tablecolhead align=""left"">
Extremely Strong</tablecolhead>
<tablecolhead align=""left"">
Adequate</tablecolhead>
<tablecolhead align=""left"">
Highly Vulnerable</tablecolhead>
</tablerow>
<tablerow>
<tablesub>
Airports</tablesub>
<tablecell> </tablecell>
<tablecell> </tablecell>
</tablerow>
<tablerow>
<tabletext>
Airport that provides essential air service with no apparent constraints on increasing rates. </tabletext>
<tabletext>
Airport with an adequate competitive position. </tabletext>
<tabletext>
Airport that functions as a major connecting hub with extremely high air carrier concentration. </tabletext>
</tablerow>
<tablerow>
<tablefootnote>
TIE--Transportation infrastructure enterprise.</tablefootnote>
</tablerow>
</table>
</section>
"""


NOTATION = reading.load_notation()


class FormulaNotation(unittest.TestCase):
    def test_precedence_and_associativity(self):
        for text, expected in (("y = a - b - c", "y = a - b - c"), ("y = a/b/c", "y = a / b / c"), ("y = -x^2", "y = -x^2"),
                               ("y = 2^3^2", "y = 2^3^2"), ("y = a*(b + c)", "y = a * (b + c)"), ("y = 12.5% * x", "y = 0.125 * x")):
            self.assertEqual(core.expr_to_text(reading.parse_formula(text, NOTATION)), expected)
        tree = reading.parse_formula("y = a - (b - c)", NOTATION)
        self.assertEqual(core.expr_to_text(tree), "y = a - (b - c)")

    def test_functions_and_inverse(self):
        tree = reading.parse_formula("K = LGD * N((N^-1(PD) + sqrt(R) * N^-1(0.999)) / sqrt(1 - R))", NOTATION)
        names = [node.name for node in core.expr_walk(tree) if node.op == "call"]
        self.assertEqual(sorted(set(names)), ["normal_cdf", "normal_inverse", "sqrt"])

    def test_ambiguous_notation_is_not_guessed(self):
        for text in ("y = 2x", "y = (a+b)(c+d)", "y = f(x)", "y = a +"):
            with self.assertRaises(reading.NotReadable, msg=text):
                reading.parse_formula(text, NOTATION)

    def test_juxtaposition_is_a_product_only_for_markup_sources(self):
        tree = reading.parse_formula("y = 2 x", NOTATION, implicit_product=True)
        self.assertEqual(core.expr_to_text(tree), "y = 2 * x")

    def test_mathml_and_latex_conversion(self):
        import xml.etree.ElementTree as ET
        math = ET.fromstring('<math xmlns="http://www.w3.org/1998/Math/MathML"><mi>y</mi><mo>=</mo><mfrac><mrow><mi>a</mi><mo>+</mo><mi>b</mi></mrow>'
                             '<msqrt><mi>c</mi></msqrt></mfrac></math>')
        linear = reading.math_to_linear(math)
        self.assertEqual(core.expr_to_text(reading.parse_formula(linear, NOTATION, implicit_product=True)), "y = (a + b) / sqrt(c)")
        latex = reading.latex_to_linear(r"K = \frac{a + b}{\sqrt{1 - \rho}} \cdot \Phi^{-1}(0.999)")
        tree = reading.parse_formula(latex, NOTATION, implicit_product=True)
        self.assertIn("normal_inverse", [n.name for n in core.expr_walk(tree)])
        self.assertIn("rho", core.expr_symbols(tree))


class UnfamiliarSchema(unittest.TestCase):
    """A schema nobody listed in tag_rules.yaml, read by the shape of its own markup."""

    def chunks(self):
        return helpers.chunks_of("method.txt", UNFAMILIAR_SCHEMA)

    def test_a_table_is_read_without_its_tags_being_listed_anywhere(self):
        chunks, _ = self.chunks()
        tables = [c for c in chunks if c["kind"] == "Table"]
        self.assertEqual(len(tables), 2, "each table is one chunk")
        self.assertIn("Bhavini Patel, CFA; (1) 416-507-2558; Toronto", tables[0]["text"])
        tables = tables[1:]
        table = tables[0]["table"]
        self.assertEqual(list(table["header"]), ["Extremely Strong", "Adequate", "Highly Vulnerable"])
        self.assertEqual([list(r) for r in table["rows"]],
                         [["Airports", "", ""],
                          ["Airport that provides essential air service with no apparent constraints on increasing rates.",
                           "Airport with an adequate competitive position.",
                           "Airport that functions as a major connecting hub with extremely high air carrier concentration."],
                          ["TIE--Transportation infrastructure enterprise.", "", ""]])
        self.assertEqual(tables[0]["caption"], "Table 2. Market Position Assessment By Asset Class")
        self.assertIn("Extremely Strong", tables[0]["text"], "the table's words are in the chunk, not only in its parts")

    def test_the_number_the_document_gives_a_paragraph_is_the_number_shown(self):
        chunks, _ = self.chunks()
        labels = [c["para_label"] for c in chunks if c["kind"] == "Paragraph"]
        self.assertEqual(labels, ["1.", "2.", "3.", "36.", "37.", "38.", "42.", "43.", "44."])

    def test_a_heading_carried_in_an_attribute_is_not_lost(self):
        chunks, _ = self.chunks()
        chains = {c["heading_chain"][-1] for c in chunks if c["heading_chain"]}
        self.assertIn("OVERVIEW AND SCOPE", chains)
        self.assertIn("b) Industry Risk", chains)
        self.assertIn("c) Market Position (60% weighting)", chains)

    def test_a_tag_inside_running_text_stays_in_its_sentence(self):
        chunks, _ = self.chunks()
        paragraph = next(c for c in chunks if c["para_label"] == "37.")
        self.assertIn("Methodology: Industry Risk", paragraph["text"])
        self.assertEqual(sum(1 for c in chunks if "Methodology: Industry Risk" in c["text"]), 1)

    def test_every_inferred_tag_is_reported_with_its_reason(self):
        _, result = self.chunks()
        said = " ".join(row["value"] for row in result.records["info_rows"] if "unrecognised" in row["item"])
        for tag in ("paranum", "tablerow", "tablecell", "tabletitle", "tablecolhead", "tabletext"):
            self.assertIn("'%s'" % tag, said)
        self.assertIn("because it", said, "each one says why it was read that way")

    def test_one_row_of_another_width_does_not_empty_the_table(self):
        """The real Table 2 came out EMPTY. An earlier table had taught the tool that <tablecell> is a
        cell; then one row of a different width made the shape test turn Table 2 down, so it was
        read with <tablecell> alone - the blank spacers - and every word in it was lost."""
        chunks, _ = self.chunks()
        table = [c for c in chunks if c["kind"] == "Table"][1]
        self.assertIn("Airport with an adequate competitive position.", table["text"])
        self.assertIn("TIE--Transportation infrastructure enterprise.", table["text"], "the row of another width is kept too")

    def test_a_table_of_sentences_is_shown_under_its_column_headings(self):
        chunks, _ = self.chunks()
        lines = [c for c in chunks if c["kind"] == "Table"][1]["text"].split("\n")
        self.assertEqual(lines[0], "Extremely Strong; Adequate; Highly Vulnerable")
        self.assertEqual(lines[1], "Airports", "a row with one filled cell is shown as it stands")
        self.assertEqual(lines[3], "Adequate: Airport with an adequate competitive position.")
        short = [c for c in chunks if c["kind"] == "Table"][0]["text"]
        self.assertEqual(short.split("\n")[0], "Joseph J Pezzimenti; + 1 (212) 438 2038; New York", "short values stay a grid")

    def test_list_items_are_folded_into_the_paragraph_that_introduces_them(self):
        chunks, _ = self.chunks()
        two = next(c for c in chunks if c["para_label"] == "2.")
        self.assertEqual(two["text"].split("\n")[1:], ["- Mission: The entity has a public policy objective.",
                                                        "- Motive: The entity is not commercially motivated.",
                                                        "- Distributions: Gains may be generated, but the entity generally retains earnings."])
        self.assertFalse([c for c in chunks if c["text"].startswith("Mission:")], "a bullet is not a unit of its own")
        after = chunks[chunks.index(two) + 1]
        self.assertEqual(after["para_label"], "3.", "paragraph 3 follows paragraph 2, with no 3, 4, 5 counted in between")

    def test_a_list_with_no_paragraph_before_it_keeps_its_items(self):
        """Under a heading, the items of a list are whole statements - definitions - and there is
        no paragraph to fold them into. They stay units, so the sample's references do not move."""
        chunks, _ = helpers.chunks_of("method.xml", build_samples.A_METHODOLOGY)
        self.assertEqual(len(chunks), 27)
        self.assertTrue(chunks[3]["text"].startswith("Actual weight (AW)"))

    def test_every_tag_that_was_inferred_says_how_often_it_occurs(self):
        _, result = self.chunks()
        said = " ".join(row["value"] for row in result.records["info_rows"] if "unrecognised" in row["item"])
        self.assertIn("'paranum', 9 times", said)
        self.assertNotIn(", 0 times", said)

    def test_the_skeleton_of_a_document_is_never_mistaken_for_a_table(self):
        """<methodology> holding <part>s holding <section>s repeats twice over exactly as a
        table does. What separates them is that a cell holds words and a section holds blocks."""
        chunks, _ = helpers.chunks_of("method.txt", build_samples.F_METHODOLOGY)
        self.assertEqual(sum(1 for c in chunks if c["kind"] == "Table"), 2)
        self.assertGreater(sum(1 for c in chunks if c["kind"] == "Paragraph"), 10)


def picture_with_words():
    """A PNG big enough to count as a picture (over 2 kB), as a chart in a model document is."""
    from PIL import Image, ImageDraw
    image = Image.new("RGB", (640, 300), "white")
    draw = ImageDraw.Draw(image)
    for row in range(6):
        draw.rectangle((20, 20 + row * 44, 40 + row * 97, 52 + row * 44), fill=(60 + row * 30, 110, 180))
        draw.text((480, 28 + row * 44), "level %d" % (row + 1), fill="black")
    data = io.BytesIO()
    image.save(data, "PNG")
    return data.getvalue()


def model_document_pdf():
    """A PDF laid out as a real model document is: a running title, a logo and a dated, numbered
    footer on every page; headings set in bold without numbers; a sentence at the foot of page 1
    whose bullets open page 2; a sentence that runs over the break to page 3; a chart."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    width, height = letter
    data, picture = io.BytesIO(), picture_with_words()
    page = canvas.Canvas(data, pagesize=letter)

    def furniture(number):
        page.setFont("Helvetica", 8)
        page.drawString(72, height - 40, "Credit Scoring Model (R)")
        page.drawImage(ImageReader(io.BytesIO(picture)), width - 168, height - 52, width=96, height=24)
        page.drawString(72, 36, "www.example.com/ratingsdirect November 8, 2023")
        page.drawRightString(width - 72, 36, str(number))

    def paragraph(y, text, font="Helvetica", size=10):
        page.setFont(font, size)
        for line in textwrap.wrap(text, 92):
            page.drawString(72, y, line)
            y -= 13
        return y - 7

    furniture(1)
    y = paragraph(height - 100, "Model Implementation Document", "Helvetica-Bold", 18) - 12
    y = paragraph(y, "Enterprise Profile", "Helvetica-Bold", 12)
    y = paragraph(y, "The enterprise profile is assessed first and is then combined with the financial profile of the enterprise.")
    paragraph(96, "Specifically, the model uses the following factors and weightings to assess the enterprise profile:")
    page.showPage()
    furniture(2)
    y = height - 100
    for mark, item in (("-", "Economic Assessment (10%)"), ("-", "Industry Risk Assessment (20%)"), ("\u2022", "Market Position Assessment (60%)")):
        y = paragraph(y, "%s %s" % (mark, item))
    y = paragraph(y, "Financial Profile", "Helvetica-Bold", 12)
    paragraph(y, "The model derives initial assessments from quantitative and qualitative inputs, which may be adjusted to arrive at the")
    page.showPage()
    furniture(3)
    y = paragraph(height - 100, "final assessment for that factor, as shown in table 1.")
    page.drawImage(ImageReader(io.BytesIO(picture)), 72, y - 200, width=420, height=196)
    paragraph(y - 220, "The chart above restates the weights used by the model.")
    page.showPage()
    page.save()
    return data.getvalue()


class StandInReader:
    """Stands in for the OCR engine: same call, same shape of answer, no package needed."""

    def __call__(self, data):
        return [[[[20, 20], [200, 20], [200, 40], [20, 40]], "Enterprise Risk Weights", 0.99],
                [[[20, 60], [60, 60], [60, 80], [20, 80]], "60%", 0.98],
                [[[300, 61], [500, 61], [500, 81], [300, 81]], "MarketPosition", 0.97]], None


class ModelDocuments(unittest.TestCase):
    """PDF and Word files laid out the way real model documentation is."""

    @classmethod
    def setUpClass(cls):
        cls.kept = list(reading.PICTURE_READER)
        reading.PICTURE_READER[:] = [StandInReader()]
        cls.pdf, cls.pdf_result = helpers.chunks_of("MID.PDF", model_document_pdf(), corner="documentation")

    @classmethod
    def tearDownClass(cls):
        reading.PICTURE_READER[:] = cls.kept

    def test_what_a_page_carries_only_because_it_is_a_page_is_left_out_and_said(self):
        text = " ".join(c["text"] for c in self.pdf)
        self.assertNotIn("ratingsdirect", text)
        self.assertNotIn("Credit Scoring Model (R)", text)
        self.assertEqual(sum(1 for c in self.pdf if c["kind"] == "Figure"), 1, "the logo on every page is no figure; the chart is")
        notes = [row["value"] for row in self.pdf_result.records["info_rows"] if "reading note" in row["item"]]
        self.assertTrue(any("ratingsdirect" in note and "3 of 3 pages" in note for note in notes), "nothing is left out unseen")

    def test_a_heading_is_known_by_how_it_is_set_when_it_has_no_number(self):
        chains = [c["heading_chain"] for c in self.pdf]
        self.assertEqual(chains[0], ["Model Implementation Document", "Enterprise Profile"])
        self.assertEqual(chains[-1], ["Model Implementation Document", "Financial Profile"])
        self.assertEqual({c["level"] for c in self.pdf}, {2}, "the larger title is level 1, the bold headings level 2")

    def test_bullets_that_open_a_page_go_to_the_sentence_that_closed_the_page_before(self):
        introducing = next(c for c in self.pdf if c["text"].startswith("Specifically"))
        self.assertEqual(introducing["text"].split("\n")[1:], ["- Economic Assessment (10%)", "- Industry Risk Assessment (20%)",
                                                                "- Market Position Assessment (60%)"])
        self.assertFalse([c for c in self.pdf if "Economic Assessment" in c["text"] and c is not introducing])

    def test_a_sentence_that_runs_over_a_page_break_is_one_sentence(self):
        joined = next(c for c in self.pdf if c["text"].startswith("The model derives"))
        self.assertTrue(joined["text"].endswith("to arrive at the final assessment for that factor, as shown in table 1."))

    def test_a_paragraph_is_labelled_with_its_page_and_its_place_on_it(self):
        labels = [c["para_label"] for c in self.pdf if c["kind"] == "Paragraph"]
        self.assertEqual(labels, ["p.1 \u00b61", "p.1 \u00b62", "p.2 \u00b61", "p.3 \u00b61"])

    def test_the_words_in_a_picture_are_shown_as_a_machine_reading_and_the_figure_stays_a_figure(self):
        figure = next(c for c in self.pdf if c["kind"] == "Figure")
        lines = figure["text"].split("\n")
        self.assertEqual(lines[:2], ["Picture on page 3", reading.OCR_NOTE])
        self.assertEqual(lines[2:], ["Enterprise Risk Weights", "60%  Market Position"], "what stands on one line of the picture stays on one line")

    def test_without_the_ocr_package_a_picture_is_still_a_figure_and_the_file_says_so_once(self):
        reading.PICTURE_READER[:] = [None]
        try:
            chunks, result = helpers.chunks_of("MID.PDF", model_document_pdf(), corner="documentation")
        finally:
            reading.PICTURE_READER[:] = [StandInReader()]
        self.assertEqual([c["text"] for c in chunks if c["kind"] == "Figure"], ["Picture on page 3"])
        notes = [row["value"] for row in result.records["info_rows"] if "rapidocr-onnxruntime" in row["value"]]
        self.assertEqual(len(notes), 1)

    def test_word_tells_bullets_from_numbered_items_and_reads_its_pictures(self):
        import docx
        from docx.shared import Inches
        file = docx.Document()
        file.add_heading("Enterprise Profile", level=1)
        file.add_paragraph("The model uses the following factors:")
        for item in ("Economic Assessment (10%)", "Market Position Assessment (60%)"):
            file.add_paragraph(item, style="List Bullet")
        file.add_paragraph("The assessment proceeds in this order:")
        for item in ("Score each factor from 1 to 6.", "Weight the scores."):
            file.add_paragraph(item, style="List Number")
        file.add_picture(io.BytesIO(picture_with_words()), width=Inches(4))
        data = io.BytesIO()
        file.save(data)
        chunks, _ = helpers.chunks_of("MID.docx", data.getvalue(), corner="documentation")
        self.assertEqual([c["kind"] for c in chunks], ["Paragraph", "Paragraph", "Figure"])
        self.assertEqual(chunks[0]["text"].split("\n")[1:], ["- Economic Assessment (10%)", "- Market Position Assessment (60%)"])
        self.assertEqual(chunks[1]["text"].split("\n")[1:], ["1. Score each factor from 1 to 6.", "2. Weight the scores."])
        self.assertIn(reading.OCR_NOTE, chunks[2]["text"])


class ReadingFiles(unittest.TestCase):
    def test_format_is_found_from_content_not_from_the_extension(self):
        self.assertEqual(core.detect_format(b"%PDF-1.4 ..."), "pdf")
        self.assertEqual(core.detect_format(b"<?xml version='1.0'?><a/>"), "xml")
        self.assertEqual(core.detect_format(b"MIME-Version: 1.0\nContent-Type: multipart/related"), "mhtml")
        self.assertEqual(core.detect_format(b"  <html><body>x</body></html>"), "html")
        self.assertEqual(core.detect_format(b"Just some words."), "text")

    def test_xml_in_a_txt_file_nested_levels_tables_equations(self):
        chunks, result = helpers.chunks_of("method.txt", build_samples.F_METHODOLOGY)
        kinds = [c["kind"] for c in chunks]
        self.assertEqual((kinds.count("Table"), kinds.count("Equation")), (2, 2))
        self.assertEqual(max(c["level"] for c in chunks), 5)
        floor = next(c for c in chunks if "never taken below" in c["text"])
        self.assertEqual(floor["heading_chain"][-1], "3.1.1 Floor")
        readable = [c["equation"]["readable"] for c in chunks if c["kind"] == "Equation"]
        self.assertEqual(readable, [True, False])
        image = next(c for c in chunks if c["kind"] == "Equation" and not c["equation"]["readable"])
        self.assertEqual(image["equation"]["not_readable_reason"], "the equation is an image")
        annex = {c["numbering"]: c["level"] for c in chunks if c["heading_chain"] and c["heading_chain"][0].startswith("Annex") or c["numbering"] in ("I.", "(a)")}
        self.assertLess(annex["I."], annex["(a)"])

    def test_one_table_is_one_chunk_and_inline_formulas_are_found(self):
        chunks, _ = helpers.chunks_of("method.txt", build_samples.F_METHODOLOGY)
        table = next(c for c in chunks if c["kind"] == "Table" and "Retail" in c["text"])
        self.assertEqual(table["table"]["header"], ["Segment", "LGD floor"])
        self.assertEqual(len(table["table"]["rows"]), 3)
        inline = next(c for c in chunks if c["text"].startswith("The asset correlation falls"))
        self.assertTrue(inline["equation"]["readable"])
        self.assertEqual(inline["equation"]["source_form"], "inline")

    def test_malformed_markup_is_repaired_and_the_repairs_are_recorded(self):
        broken = "<doc><section><title>1 Scope</title><p>Costs & prices are set here.<p>A second one with a bad &entity; inside.</section></doc>"
        chunks, result = helpers.chunks_of("broken.xml", broken)
        self.assertEqual(len(chunks), 2)
        self.assertIn("Costs & prices", chunks[0]["text"])
        self.assertTrue(result.records["read_repairs"])

    def test_unknown_tags_are_read_and_reported(self):
        chunks, result = helpers.chunks_of("odd.xml", "<doc><sect><ttl>1 Scope</ttl><blurb>The rate is 5% for all.</blurb></sect></doc>")
        self.assertTrue(any("5%" in c["text"] for c in chunks))
        self.assertTrue(any("unrecognised tag" in row["item"].lower() for row in result.records["info_rows"]))

    def test_mhtml_keeps_preserved_equation_markup_and_skips_its_fallback_image(self):
        chunks, _ = helpers.chunks_of("method.mhtml", build_samples.mhtml_file(build_samples.D_HTML, ["image001.png", "image002.png"]))
        equations = [c for c in chunks if c["kind"] == "Equation"]
        self.assertEqual(len(equations), 1)
        self.assertEqual(equations[0]["text"], "D = W * R")
        self.assertEqual([c["kind"] for c in chunks].count("Figure"), 1)
        self.assertEqual(len(chunks), 9)

    def test_docx_headings_numbering_tables_equations_and_figures(self):
        chunks, _ = helpers.chunks_of("doc.docx", build_samples.docx_file(build_samples.f_documentation(False)), corner="documentation")
        kinds = [c["kind"] for c in chunks]
        self.assertEqual((kinds.count("Table"), kinds.count("Equation"), kinds.count("Figure")), (1, 1, 1))
        equation = next(c for c in chunks if c["kind"] == "Equation")
        self.assertTrue(equation["equation"]["readable"])
        self.assertIn("normal_inverse", equation["text"])
        self.assertEqual(next(c for c in chunks if "floor_pd" in c["text"])["heading_chain"], ["2 Risk components", "2.1 Probability of default"])

    def test_pdf_text_and_tables(self):
        chunks, _ = helpers.chunks_of("guide.pdf", build_samples.pdf_file(build_samples.F_USER_GUIDE), corner="documentation")
        self.assertTrue(any(c["kind"] == "Table" and "Retail" in c["text"] for c in chunks))
        self.assertTrue(any("scales the requirement by 12.5" in c["text"] for c in chunks))

    def test_a_file_that_cannot_be_read_gives_a_not_read_unit_and_never_stops_the_run(self):
        chunks, _ = helpers.chunks_of("broken.docx", b"PK\x03\x04 this is not a real archive", corner="documentation")
        self.assertEqual(len(chunks), 1)
        self.assertTrue(chunks[0]["not_read_reason"])

    def test_entities_defined_inside_an_office_file_are_never_expanded(self):
        bomb = b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;">]><d>&b;</d>'
        with self.assertRaises(Exception):
            reading.safe_xml(bomb)
        self.assertEqual(reading.safe_xml(b'<?xml version="1.0"?><!DOCTYPE d><d>fine</d>').text, "fine")

    def test_checkable_statements_and_cross_references(self):
        chunks, _ = helpers.chunks_of("m.xml", "<doc><section><title>1 Rates</title><p>This part gives the background of the method.</p>"
                                               "<p>The rate is at least 3% as set out in Table 2 and section 4.1.</p></section></doc>", corner="documentation")
        self.assertEqual([c["checkable"] for c in chunks], [False, True])
        self.assertEqual(sorted(chunks[1]["refs_out"]), ["Table 2", "section 4.1"])

    def test_same_bytes_give_the_same_chunks(self):
        first, _ = helpers.chunks_of("method.txt", build_samples.F_METHODOLOGY)
        second, _ = helpers.chunks_of("method.txt", build_samples.F_METHODOLOGY)
        self.assertEqual([c["content_hash"] for c in first], [c["content_hash"] for c in second])
        self.assertEqual([c["ref"] for c in first], ["C-%04d" % n for n in range(1, len(first) + 1)])


if __name__ == "__main__":
    unittest.main()
