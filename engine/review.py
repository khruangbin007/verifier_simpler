"""
Verifier 0.0.3 - review.py - what corresponds to what, and what differs. For Reviewer 3
(mapping) and Reviewer 4 (checks).

WHAT THIS FILE DOES
  Mapping. Every unit becomes a node of one graph. Candidate passages for each unit are found by
  deterministic signals (fields, the bridge vocabulary, the graph itself) and shown to the model,
  which chooses among them and quotes them; a quotation that is not word for word in the passage,
  or a reference that was not shown, is refused, so no link rests on the model's say-so. Every
  link is a ledger record with the hash of the one before it. What each piece of code does is
  asked in plain words for the column LLM Interpretation, outside the accounting.
  Checks. Formulas are compared as expression trees, symbolically and then numerically with a
  counterexample when they differ; values in stored data are compared with values stated in the
  methodology and the documentation; stated rules are checked against the code that should carry
  them; roxygen blocks and help pages against the functions they describe. Every unit ends with
  exactly one status, every status that is not clean becomes a flagged item, and the four-part
  coverage identity is checked on every run so that nothing can fall between the parts.

WHAT IT TAKES IN AND PRODUCES
  In: the units of all three corners, the references, a chat() for the judged and checked steps.
  Out: graph_ledger, candidates, search_records, interpretations, math_checks, value_checks,
  rule_checks, package_doc_checks, unit_status, flagged_items and coverage records.

WHICH SHEETS SHOW ITS RESULTS
  Model_Implementation_Map (each model unit in its place in the computation, with its links, checks
  and status), Chunks_Doc (a documentation passage with its links and checks), Chunks_Model,
  Mapping_Coverage, Flagged_Items; and the coverage identity on Model_Package_Info.

DESIGN RULES ENFORCED HERE
  R1  a status is a plain observation; a flagged item is a question for a person, never a grade
  R3  every answer is validated by code; a quotation must be verbatim; nothing is repaired
  R4  the graph ledger is a hash chain and is only ever appended to
  R5  the same inputs and the same recorded answers give the same graph version
  R7  a formula is compared by the tool's own tree code, never by evaluating text
  R8  every flagged item cites the units it rests on, and every citation resolves

HOW TO SANITY-CHECK IT
  Run tests/test_mapping.py and tests/test_checks.py. Run the harness (cell 5, APPENDIX =
  "harness"): every seeded difference in a sample should end flagged in an expected category, and
  the clean baseline should flag none of the units listed as clean.
"""

from decimal import ROUND_HALF_EVEN, ROUND_HALF_UP, Decimal
import hashlib
import json
import math
import os
import random
import re

import yaml

import core
import reading

# ================================================================================================
# ---------------------------------------------------------------- mapping: what corresponds to what
LEDGER_VOLATILE = ("created_at", "run_id")
CORNER_NAMES = {"canon": "the methodology", "doc": "the documentation", "model": "the package"}

# ---------------------------------------------------------------- the ledger and the graph in memory
def node_record(ref, node_kind, corner):
    """The ledger record of one node."""
    return {"record_type": "node", "ref": ref, "node_kind": node_kind, "corner": corner}

def ledger_records(existing, new_records):
    """Chain new records onto the ledger. Nodes come first (by reference), then edges sorted
    by source, target and kind, never in the order threads finished, so the same content
    always gives the same graph version id. There is no function that edits or removes a
    record: a later contradiction is a NEW edge. Enforces: R4, R5"""
    plain = [core.to_plain(record) for record in new_records]
    nodes = sorted((r for r in plain if r["record_type"] == "node"), key=lambda r: r["ref"])
    edges = sorted((r for r in plain if r["record_type"] == "edge"),
                   key=lambda r: (r["source"], r["target"], r["kind"], r.get("relation", ""), r.get("how", "")))
    return core.chain_records(core.chain_head(existing), nodes + edges, LEDGER_VOLATILE)

def verify_ledger(records):
    """True when no record of the ledger was edited, removed or re-ordered."""
    return core.verify_chain(records, LEDGER_VOLATILE)

def graph_version_id(records):
    """G- and the first twelve characters of the ledger's head hash."""
    return "G-" + core.chain_head(records)[:12]

def load_graph(records):
    """The ledger as two adjacency dictionaries: outgoing and incoming edges per node."""
    graph = {"nodes": {}, "out": {}, "in": {}}
    for record in records:
        if record["record_type"] == "node":
            graph["nodes"][record["ref"]] = record
        else:
            graph["out"].setdefault(record["source"], []).append(record)
            graph["in"].setdefault(record["target"], []).append(record)
    return graph

def find_path(graph, start, is_goal, allowed_kinds=None, max_hops=4):
    """Breadth-first walk over typed edges in either direction, at most `max_hops` long.
    Returns the list of (edge, node reached) hops of the first shortest path, or None."""
    frontier, seen = [(start, [])], {start}
    for _ in range(max_hops):
        following = []
        for node, path in frontier:
            steps = [(e, e["target"]) for e in graph["out"].get(node, [])] + [(e, e["source"]) for e in graph["in"].get(node, [])]
            for edge, reached in sorted(steps, key=lambda step: (step[1], step[0]["kind"])):
                if reached in seen or (allowed_kinds and edge["kind"] not in allowed_kinds):
                    continue
                seen.add(reached)
                if is_goal(reached):
                    return path + [(edge, reached)]
                following.append((reached, path + [(edge, reached)]))
        frontier = following
    return None

def links_of(graph, ref, corner_prefix):
    """The `corresponds` edges of a unit into one corner ("C", "D" or "M"), in ledger order."""
    edges = [e for e in graph["out"].get(ref, []) if e["kind"] == "corresponds" and e["target"].startswith(corner_prefix)]
    return edges + [dict(e, target=e["source"]) for e in graph["in"].get(ref, [])
                    if e["kind"] == "corresponds" and e["source"].startswith(corner_prefix)]

# ---------------------------------------------------------------- words: splitting, stemming, word lists
# ---------------------------------------------------------------- reference data of the search
# Words too common to say anything about which passage corresponds to which, and the patterns that
# bridge a code name and a written term (rho_a and "asset correlation"). No domain word may appear
# in either: the layout lint checks them. Enforces: R9
STOPWORDS_TEXT = r'''# stopwords.txt - generic function words left out of the word index. One per line.
# Rule R9: no word of any field of business belongs here.
a an the and or of to in is are be by for with as at on it this that these those from each which
was were been being has have had do does did not no nor if then than so such any all some its their
there here where when while who whom whose what how why can could may might must shall should will would
into per under over between within without about above below after before during through up down out off
also only other more most less least very same own both either neither one two
function return returns returned value values given using used use uses see set sets get gets
'''
BRIDGE_PATTERNS_YAML = r'''# bridge_patterns.yaml - generic words the tool uses to recognise where a document ties a short
# name to a longer phrase. Reviewer 3 owns this file. Rule R9: no word of any field of business.
definition_verbs: [denotes, represents, stands for, is defined as, means]
symbol_headers: [symbol, variable, notation, parameter, name, term, abbreviation, column, field]
description_headers: [description, definition, meaning, explanation, stands for, content]
glossary_columns: {term: Term, also: Also written as}
# words that tie a function name to neutral mathematical words (used for the called-functions field)
function_words:
  normal_cdf: [cumulative, normal, distribution]
  normal_inverse: [inverse, normal, quantile]
  max: [maximum, larger, floor, least]
  min: [minimum, smaller, cap, capped, most]
  exp: [exponential]
  log: [logarithm]
  sqrt: [square, root]
  round: [rounded, decimals]
  piecewise: [condition, otherwise]
  sum_over: [sum, total]
'''

def load_word_lists():
    """Stop words, bridge patterns and the neutral words of mathematical functions. The lint
    checks every one of these lists for domain words. Enforces: R9"""
    stop = {word for line in STOPWORDS_TEXT.split("\n") if not line.startswith("#") for word in line.split()}
    return {"stop": stop, "patterns": yaml.safe_load(BRIDGE_PATTERNS_YAML),
            "function_map": yaml.safe_load(reading.R_FUNCTION_MAP_YAML)}

def stem(word):
    """A light, rule-based stemmer: plural endings, -ing, -ed, and a doubled last letter.
    Deliberately weak: a wrong merge costs more than a missed one."""
    for ending, replacement in (("ies", "y"), ("ied", "y"), ("sses", "ss"), ("ing", ""), ("ed", ""), ("es", "e"), ("s", "")):
        if word.endswith(ending) and len(word) - len(ending) >= 3 and not word.endswith(("ss", "us", "is")):
            word = word[:len(word) - len(ending)] + replacement
            break
    if len(word) > 3 and word[-1] == word[-2] and word[-1] not in "lsz":
        word = word[:-1]
    return word[:-1] if len(word) > 4 and word.endswith("e") else word

AS_WRITTEN = {}      # stem -> the first word seen with that stem; used only to show words to analysts

def split_words(text, stop):
    """Text or identifiers to index words: split at anything that is not a letter or digit,
    at underscores, dots and capital letters inside a name; lower-case; drop stop words,
    single characters and pure numbers; stem."""
    words = []
    for piece in re.findall(r"[^\W_]+", text or ""):
        for part in re.findall(r"[A-Z]+(?![a-z])|[A-Z]?[a-z]+|[^\W\d_]+|\d+", piece):
            lowered = part.lower()
            if len(lowered) > 1 and not lowered.isdigit() and lowered not in stop:
                words.append(stem(lowered))
                AS_WRITTEN.setdefault(words[-1], lowered)
    return words

def shown(words):
    """Index words as an analyst should see them: as first written, not as stems."""
    return [AS_WRITTEN.get(word, word) for word in words]

def symbols_in(text):
    """Short symbols a passage uses: single letters and Greek names standing alone, short
    capital abbreviations, and names with a subscript."""
    found = re.findall(r"(?<![\w.])([A-Za-z\u0370-\u03ff]{1,4}(?:_\{?\w+\}?|\[\w+\])?)(?![\w(])", text or "")
    keep = []
    for symbol in found:
        plain = core.normalise_symbol(symbol)
        if symbol.isupper() or len(symbol) == 1 or "_" in symbol or plain in core.GREEK or symbol in core.GREEK.values():
            if symbol.lower() not in ("a", "i"):
                keep.append(plain)
    return list(dict.fromkeys(keep))

def numbers_in(text, trivial):
    """The non-trivial numbers of a text, as normalised values ("12.5%" gives 0.125)."""
    return list(dict.fromkeys(n["value"] for n in core.find_numbers(text or "") if n["value"] not in trivial))

# ---------------------------------------------------------------- what is indexed for each unit
FIELD_WEIGHTS = {"name": 3, "heading": 3, "about": 2, "caption": 2, "symbols": 2, "body": 1, "calls": 1}

def chunk_fields(chunk, lists, trivial):
    """The named fields of a methodology or documentation chunk (plan 2.6)."""
    table = chunk.get("table") or {}
    table_words = " ".join(table.get("header", [])) + " " + " ".join(row[0] for row in table.get("rows", []) if row)
    return {"fields": {"heading": split_words(" ".join(chunk["heading_chain"][-2:]), lists["stop"]),
                       "caption": split_words(chunk.get("caption", "") + " " + table_words, lists["stop"]),
                       "body": split_words(chunk["text"], lists["stop"])},
            "symbols": symbols_in(chunk["text"]) + list((chunk.get("equation") or {}).get("expression") and
                       sorted(core.expr_symbols(core.expr_from_dict(chunk["equation"]["expression"]))) or []),
            "numbers": numbers_in(chunk["text"], trivial), "identifiers": re.findall(r"\b[a-z]+(?:[_.][a-z0-9]+)+\b", chunk["text"])}

def unit_fields(unit, units_by_ref, documented_by, lists, trivial):
    """The named fields of a model unit: name words, what the package says about it (roxygen
    of the object or of the function it sits in), comments, symbols, numbers, the neutral
    words of the mathematical functions it calls, and string literals."""
    code = unit.get("code") or {}
    about_unit = documented_by.get(unit["ref"])
    about = about_unit["text"] if about_unit else ""
    parent_block = documented_by.get(unit.get("parent_ref") or "")
    if parent_block and not about_unit:                  # a statement takes only the @param lines of the symbols it uses
        used = set(code.get("symbols_read", ())) | set(code.get("symbols_written", ()))
        about = " ".join(tag["text"] for tag in parent_block["roxygen"]["tags"] if tag["tag"] == "param" and tag["name"] in used)
    if unit.get("roxygen") or unit.get("helppage") or unit["kind"] == core.KIND_VIGNETTE:
        about = unit["text"]
    comments = " ".join(re.findall(r"#(?!')\s*(.*)", unit["text"])) if code else ""
    neutral = [word for call in code.get("calls", ()) for word in
               lists["patterns"]["function_words"].get((lists["function_map"]["r_functions"].get(call) or {}).get("neutral", ""), [])]
    data = unit.get("data") or {}
    columns = " ".join(name for name, _ in data.get("columns", ()))
    body = comments + " " + " ".join(code.get("strings", ())) + " " + (unit["text"] if data else "")
    symbols = [core.normalise_symbol(s) for s in tuple(code.get("symbols_read", ())) + tuple(code.get("symbols_written", ()))]
    return {"fields": {"name": split_words(unit["name"] + " " + unit.get("inside", "") + " " + columns, lists["stop"]),
                       "about": split_words(re.sub(r"#'|@\w+|\\\w+", " ", about), lists["stop"]),
                       "body": split_words(body, lists["stop"]), "calls": [stem(word) for word in neutral]},
            "symbols": list(dict.fromkeys(symbols)),
            "numbers": list(dict.fromkeys([n["value"] for n in code.get("numbers", ()) if n["value"] not in trivial] +
                                          numbers_in(about + " " + (unit["text"] if data else ""), trivial))),
            "identifiers": [unit["name"]] if code and unit["name"] else []}

# ---------------------------------------------------------------- S1: field-aware BM25 (the only text-ranking formula)
def build_index(documents):
    """documents: {ref: {"fields": {field: [words]}}}. Term frequencies are weighted by field;
    the length of a document is its weighted number of words."""
    index = {"tf": {}, "length": {}, "df": {}, "n": len(documents)}
    for ref in sorted(documents):
        weighted = {}
        for field_name, words in documents[ref]["fields"].items():
            for word in words:
                weighted[word] = weighted.get(word, 0) + FIELD_WEIGHTS.get(field_name, 1)
        index["tf"][ref], index["length"][ref] = weighted, sum(weighted.values())
        for word in weighted:
            index["df"][word] = index["df"].get(word, 0) + 1
    index["average_length"] = (sum(index["length"].values()) / len(documents)) if documents else 0.0
    return index

def bm25_scores(index, query_weights, k1, b):
    """BM25 with the usual constants k1 and b. query_weights: {word: weight}; a bridged word
    seen only once in the inputs counts half. Returns {ref: (score, matched words)}."""
    scores = {}
    for word in sorted(query_weights):
        df = index["df"].get(word)
        if not df:
            continue
        idf = math.log(1 + (index["n"] - df + 0.5) / (df + 0.5))
        for ref, weighted in index["tf"].items():
            tf = weighted.get(word)
            if tf:
                norm = tf + k1 * (1 - b + b * index["length"][ref] / (index["average_length"] or 1.0))
                score, words = scores.get(ref, (0.0, []))
                scores[ref] = (score + query_weights[word] * idf * tf * (k1 + 1) / norm, words + [word])
    return scores

def ranked(scores):
    """References by falling score; ties break by reference. Enforces: R5"""
    return [ref for ref, _ in sorted(scores.items(), key=lambda item: (-item[1][0] if isinstance(item[1], tuple) else -item[1], item[0]))]

# ---------------------------------------------------------------- S2: the bridge vocabulary
def initials_match(short, words):
    """Do the letters of an abbreviation appear, in order, as initials of the long form?
    Guards against reading "(see Table 3)" as an abbreviation."""
    letters = [c for c in short.lower() if c.isalpha()]
    initials = [w[0].lower() for w in words if w]
    position = 0
    for letter in letters:
        while position < len(initials) and initials[position] != letter:
            position += 1
        if position == len(initials):
            return False
        position += 1
    return bool(letters)

def harvest_text(text, source, patterns):
    """Bridge entries from one text: "long form (ABBR)", "ABBR (long form)", "where X denotes
    ...", "let X be ...". Returns (term, phrase, source, pattern name) tuples."""
    entries = []
    for found in re.finditer(r"((?:[A-Za-z][\w-]*\s+){1,6}[A-Za-z][\w-]*)\s+\(([A-Za-z\u0370-\u03ff][\w]{0,9})\)", text):
        words = found.group(1).split()
        for start in range(len(words) - 1, -1, -1):
            if initials_match(found.group(2), words[start:]) and len(words[start:]) <= len(found.group(2)) + 2:
                entries.append((found.group(2), " ".join(words[start:]), source, "long form (short form)"))
                break
    for found in re.finditer(r"\b([A-Z][A-Za-z0-9_]{1,9})\s+\(([a-z][^()]{3,60})\)", text):
        if initials_match(found.group(1), found.group(2).split()):
            entries.append((found.group(1), found.group(2), source, "short form (long form)"))
    verbs = "|".join(re.escape(verb) for verb in patterns["definition_verbs"])
    for found in re.finditer(r"(?:\bwhere|,|\band)\s+(\S{1,12})\s+(?:%s|is)\s+(?:the\s+|an?\s+)?([^,.;()]{3,60})" % verbs, text):
        if len(found.group(1)) <= 4 or "_" in found.group(1):
            entries.append((found.group(1), found.group(2).strip(), source, "where ... denotes"))
    for found in re.finditer(r"\b[Ll]et\s+(\S{1,12})\s+be\s+(?:the\s+|an?\s+)?([^,.;()]{3,60})", text):
        entries.append((found.group(1), found.group(2).strip(), source, "let ... be"))
    return entries

def harvest_bridge(canon, doc, units, lists, glossary_path):
    """The bridge vocabulary of this project, harvested from its own inputs: documents,
    symbol tables, roxygen @param and @return lines, column descriptions of data blocks,
    comments of the form "# x: phrase", and the optional Inputs/glossary.xlsx. Every entry
    records where it was seen and how often; entries seen once count half."""
    patterns, raw = lists["patterns"], []
    for chunk in canon + doc:
        raw.extend(harvest_text(chunk["text"], chunk["ref"], patterns))
        table = chunk.get("table") or {}
        header = [cell.lower() for cell in table.get("header", [])]
        if len(header) >= 2 and header[0] in patterns["symbol_headers"] and any(h in patterns["description_headers"] for h in header[1:]):
            described = next(i for i, h in enumerate(header) if h in patterns["description_headers"])
            raw.extend((row[0], row[described], chunk["ref"], "table of symbols") for row in table["rows"] if row[0] and row[described])
    for unit in units:
        roxygen = unit.get("roxygen")
        if roxygen:
            for tag in roxygen["tags"]:
                if tag["tag"] == "param" and tag["name"] and tag["text"]:
                    raw.extend((name, tag["text"], "%s@param:%s" % (unit["ref"], name), "roxygen @param") for name in tag["name"].split(","))
                elif tag["tag"] == "return" and tag["text"] and roxygen["documents_name"]:
                    raw.append((roxygen["documents_name"], tag["text"], unit["ref"] + "@return", "roxygen @return"))
            for found in re.finditer(r"\\item\{([^{}]+)\}\{([^{}]+)\}", unit["text"]):
                raw.append((found.group(1), found.group(2), unit["ref"], "column description"))
        if unit.get("code"):
            for found in re.finditer(r"#(?!')\s*([A-Za-z][\w.]{0,20})\s*:\s+([^\n]{3,60})", unit["text"]):
                raw.append((found.group(1), found.group(2).strip(), unit["ref"], "code comment"))
    if glossary_path:
        import openpyxl
        sheet = openpyxl.load_workbook(glossary_path, read_only=True).worksheets[0]
        for row in list(sheet.iter_rows(values_only=True))[1:]:
            if row and row[0] and len(row) > 1 and row[1]:
                raw.extend([(str(row[0]), str(row[1]), "glossary.xlsx", "glossary"), (str(row[1]), str(row[0]), "glossary.xlsx", "glossary")])
    merged = {}
    for term, phrase, source, pattern in raw:
        key = (core.normalise_symbol(term.strip()), " ".join(split_words(phrase, lists["stop"])))
        if not key[1]:
            continue
        entry = merged.setdefault(key, {"term": key[0], "term_as_written": term.strip(), "phrase": core.normalise_text(phrase)[:80],
                                        "words": key[1].split(), "sources": [], "patterns": [], "count": 0})
        entry["count"] += 1
        entry["sources"].append(source)
        entry["patterns"].append(pattern)
    return [merged[key] for key in sorted(merged)]

def expansions_for(representation, bridge_by_term):
    """The bridge entries that apply to a unit: by its symbols and by its name words."""
    found = []
    for term in list(representation["symbols"]) + representation["fields"].get("name", []) + representation.get("identifiers", []):
        for entry in bridge_by_term.get(core.normalise_symbol(term), []):
            if entry not in found:
                found.append(entry)
    return found

# ---------------------------------------------------------------- S3: explicit references as written
def resolve_reference(reference, chunks):
    """"Table 3", "section 4.2", "Annex B" as written, resolved among `chunks` (one corner):
    tables, figures and equations by the label at the start of their caption, sections by
    their numbering as written. Several matches are all returned; none gives an empty list."""
    label, _, number = reference.partition(" ")
    label, number = label.lower().rstrip("s"), number.strip("().").lower()
    matches = []
    for chunk in chunks:
        caption = (chunk.get("caption") or "").lower()
        if label in ("table", "figure", "equation") and chunk["kind"].lower() == label:
            if re.match(r"%s\s+\(?%s\)?(?!\w)(?!\.\d)" % (label, re.escape(number)), caption):
                matches.append(chunk["ref"])
        elif label in ("section", "paragraph", "chapter", "annex", "appendix"):
            written = (chunk.get("numbering") or "").lower().rstrip(".")
            if written == number or written in ("%s %s" % (label, number), "annex %s" % number, "appendix %s" % number):
                matches.append(chunk["ref"])
    return matches

# ---------------------------------------------------------------- S4: rare shared anchors and the restart random walk
def anchors_of(representation, heading_words, scope):
    """The anchors one unit or chunk mentions, as (kind, key) pairs. A symbol is scoped by
    `scope` (symbol -> key of the section that defines it) where a definition is known."""
    anchors = [("number", value) for value in representation["numbers"]]
    anchors += [("symbol", scope.get(symbol, symbol)) for symbol in representation["symbols"]]
    anchors += [("identifier", name) for name in representation.get("identifiers", [])]
    anchors += [("term", word) for word in representation.get("terms", [])]
    if heading_words:
        anchors.append(("heading", " ".join(heading_words)))
    return list(dict.fromkeys(anchors))

def anchor_weights(mentions, settings):
    """Weight of an anchor = 1 / log(1 + number of units that mention it). An anchor that
    more than `anchor_max_share` of all units mention is dropped; an anchor only one unit
    mentions joins nothing and is dropped too; heading anchors are capped so that a section
    title shared by many passages cannot dominate."""
    total = len({ref for refs in mentions.values() for ref in refs})
    weights = {}
    for anchor, refs in mentions.items():
        if len(refs) < 2 or len(refs) > max(2, settings["anchor_max_share"] * total):
            continue
        weight = 1.0 / math.log(1 + len(refs))
        weights[anchor] = min(weight, settings["heading_anchor_cap"]) if anchor[0] == "heading" else weight
    return weights

def restart_walk(mentions, weights, starts, settings):
    """A random walk over the two-sided graph of units and anchors that keeps restarting at the
    unit (personalised PageRank): fixed restart probability and a fixed number of rounds, so
    it is deterministic. Returns {start: {ref: closeness}}."""
    import numpy
    from scipy import sparse
    refs = sorted({ref for anchor in weights for ref in mentions[anchor]})
    anchors = sorted(weights)
    position = {ref: i for i, ref in enumerate(refs)}
    rows, columns, values = [], [], []
    for column, anchor in enumerate(anchors):
        for ref in mentions[anchor]:
            rows.append(position[ref])
            columns.append(len(refs) + column)
            values.append(weights[anchor])
    size = len(refs) + len(anchors)
    if not values:
        return {start: {} for start in starts}
    matrix = sparse.coo_matrix((values + values, (rows + columns, columns + rows)), shape=(size, size)).tocsr()
    degree = numpy.asarray(matrix.sum(axis=0)).ravel()
    degree[degree == 0] = 1.0
    transition = matrix.multiply(1.0 / degree).tocsr()               # each column sums to one: hubs pass on little
    results, restart = {}, settings["walk_restart"]
    for start in starts:
        if start not in position:
            results[start] = {}
            continue
        home = numpy.zeros(size)
        home[position[start]] = 1.0
        state = home.copy()
        for _ in range(int(settings["walk_rounds"])):
            state = (1 - restart) * transition.dot(state) + restart * home
        results[start] = {ref: float(state[position[ref]]) for ref in refs if ref != start and state[position[ref]] > 1e-9}
    return results

# ---------------------------------------------------------------- S6: signatures (re-order only) and table shape
def formula_signature(expression):
    """What a formula is made of, whatever its symbols are called: operators, neutral
    function names with their number of arguments, and non-trivial constants."""
    parts = {}
    for node in core.expr_walk(core.expr_from_dict(expression)):
        if node.op in ("sym", "eq"):
            continue
        key = "%s/%d" % (node.name, len(node.args)) if node.op == "call" else node.value if node.op == "num" else node.op
        if key not in ("0", "1", "2"):
            parts[key] = parts.get(key, 0) + 1
    return parts

def overlap(left, right):
    """Weighted overlap of two multisets, between 0 and 1."""
    shared_count = sum(min(count, right.get(key, 0)) for key, count in left.items())
    return shared_count / max(1, max(sum(left.values()), sum(right.values())))

def table_shape_score(table, chunk_table, stop):
    """How alike two tables are: shared header words, shared row keys and shared values at
    printed precision. The package table comes from parameter_tables; the other from a chunk."""
    header_a = set(split_words(" ".join(table["header"]), stop))
    header_b = set(split_words(" ".join(chunk_table.get("header", [])), stop))
    keys_a = {row[0].strip().lower() for row in table["rows"] if row}
    keys_b = {row[0].strip().lower() for row in chunk_table.get("rows", []) if row}
    values_a = {n["value"] for row in table["rows"] for cell in row for n in core.find_numbers(cell)}
    values_b = {n["value"] for row in chunk_table.get("rows", []) for cell in row for n in core.find_numbers(cell)}
    share = lambda a, b: len(a & b) / max(1, min(len(a), len(b)))
    return share(header_a, header_b) + share(keys_a, keys_b) + share(values_a, values_b)

# ---------------------------------------------------------------- fusion and reasons
REASON_TEMPLATES = {      # every phrase the search stage can put into "How established"
    "concepts": "shares the concept {detail}",
    "fields": "shares the words {detail}",
    "bridge": "{detail}",
    "references": "it is cited as written ({detail})",
    "anchors": "shares the rare {detail}",
    "signatures": "its formula has a similar structure",
    "table shape": "its table has similar headers, row keys or values",
    "propagation": "inherited from {detail}"}
SIGNAL_ORDER = ("concepts", "references", "fields", "bridge", "anchors", "table shape", "signatures", "propagation")

def fuse(rankings, settings, cited=(), reserve_from=("concepts", "anchors", "propagation")):
    """Reciprocal rank fusion: score = sum over signals of 1 / (60 + rank). Only ranks are
    combined, so no signal's raw scale matters. Shared concepts come first: that signal counts
    concept_weight times, and it is the first to fill the reserved places, so a passage sharing a
    concept with the unit reaches the judge even where the words differ. Explicitly cited chunks
    are always included (up to three); ties break by reference. Enforces: R5"""
    constant, k = settings["rrf_constant"], int(settings["k_candidates"])
    scores, ranks = {}, {}
    for signal in SIGNAL_ORDER:
        refs = rankings.get(signal, [])
        if signal == "signatures":                       # used to re-order, never alone
            refs = [ref for ref in refs if ref in scores]
        weight = float(settings["concept_weight"]) if signal == "concepts" else 1.0
        for rank, ref in enumerate(refs, start=1):
            scores[ref] = scores.get(ref, 0.0) + weight / (constant + rank)
            ranks.setdefault(ref, {})[signal] = rank
    ordered = sorted(scores, key=lambda ref: (-scores[ref], ref))
    must = list(dict.fromkeys(list(cited)[:3]))
    reserved = [ref for signal in reserve_from for ref in rankings.get(signal, []) if "fields" not in ranks.get(ref, {})]
    must += [ref for ref in dict.fromkeys(reserved) if ref not in must][:int(settings["reserved_places"])]
    shortlist = [ref for ref in ordered if ref not in must][:max(0, k - len(must))] + must
    shortlist.sort(key=lambda ref: (-scores.get(ref, 0.0), ref))
    return [(ref, scores.get(ref, 0.0), ranks.get(ref, {})) for ref in shortlist]

def reason_text(signals, details):
    """The plain reason of one candidate, assembled from the signals that proposed it."""
    phrases = [REASON_TEMPLATES[signal].format(detail=details.get(signal, "")) for signal in SIGNAL_ORDER
               if signal in signals and (details.get(signal) or "{detail}" not in REASON_TEMPLATES[signal])]
    return "Proposed because: " + "; ".join(phrases) if phrases else "Proposed by rank only"

# ---------------------------------------------------------------- concepts: the model's own, and where the documents name them
# The model is the basis. Its concepts are the names its code gives things - functions, their
# arguments, the variables its statements set, its stored tables and their columns - together with
# how its own roxygen and help pages describe them. The methodology, the documentation and the rest of
# the model are then searched for those concepts, and nothing else: a term in the documents that has
# no likely equivalent in the model is never extracted. What a unit shows is always the words that
# unit writes, character for character. Code finds the same words; the model is asked only where code
# sees a likely candidate it cannot prove, and its guesses - synonyms, acronyms, abbreviations - are
# kept apart and shown as guesses on the Concepts sheet. Enforces: R3, R4, R9
CONCEPT_STOP = {"the", "a", "an", "of", "and", "or", "to", "for", "in", "on", "at", "by", "with", "as", "&", "is", "are"}
GENERIC_NAMES = {"data", "value", "values", "result", "results", "out", "output", "input", "df", "dt", "tmp", "temp", "obj",
                 "object", "list", "vec", "vector", "arg", "args", "id", "name", "names", "type", "flag", "idx", "index",
                 "row", "col", "file", "path", "table", "fn", "fun", "res", "ret", "val", "var", "len", "num", "count",
                 "na", "null", "true", "false", "self", "env", "call", "x", "y", "z", "i", "j", "k", "n", "m", "t", "f"}
ACRONYM = re.compile(r"(?<![\w&])([A-Z][A-Z0-9&]{1,7})(s?)(?![\w&])")
LONG_THEN_SHORT = re.compile(r"((?:[A-Za-z][\w'\u2019&/-]*\s+){0,7}[A-Za-z][\w'\u2019&/-]*)\s*\(\s*([A-Z][A-Za-z0-9&]{1,9})\s*\)")
SHORT_THEN_LONG = re.compile(r"(?<![\w&])([A-Z][A-Z0-9&]{1,9})\s*\(\s*([A-Za-z][^()]{3,90})\)")
CONCEPT_TOKEN = re.compile(r"[A-Za-z0-9]+")

def concept_singular(word):
    if len(word) > 4 and word.endswith("ies"):
        return word[:-3] + "y"
    if len(word) > 3 and word.endswith("s") and not word.endswith(("ss", "us", "is")):
        return word[:-1]
    return word

def identifier_words(name):
    """The words of a name the code gives: debt_to_net_revenue, debtToNetRevenue and debt.to.net.revenue
    all give debt, to, net, revenue."""
    return [part.lower() for part in re.split(r"[_.]+|(?<=[a-z0-9])(?=[A-Z])", name or "") if part]

def concept_key(text):
    """How forms are compared: lower case, each word singular, no article or preposition at either end."""
    words = [concept_singular(w) for w in CONCEPT_TOKEN.findall((text or "").replace("_", " ").lower())]
    while words and words[0] in CONCEPT_STOP:
        words.pop(0)
    while words and words[-1] in CONCEPT_STOP:
        words.pop()
    return " ".join(words)

def spells(short, words):
    """Do the words spell the acronym by their initials, hyphenated compounds counted word by word,
    starting at its first letter? Uses the search's own initials_match for the order."""
    parts = [part for word in words for part in re.split(r"[-/]", word) if part]
    letters = [c for c in short.lower() if c.isalpha()]
    return bool(parts and letters) and parts[0][0].lower() == letters[0] and len(parts) <= len(letters) + 3 and initials_match(short, parts)

def defined_acronyms(text):
    """Acronyms a text defines, as (acronym, what it stands for, the defining words): 'discounted cash
    flow (DCF)' and 'DCF (discounted cash flow)'. The long form is the shortest run of words beside the
    brackets that spells the acronym. Both sides are the text's own words."""
    found = []
    for match in LONG_THEN_SHORT.finditer(text):
        words, short = match.group(1).split(), re.sub(r"(?<=[A-Z0-9])s$", "", match.group(2))
        for size in range(1, min(len(words), 9) + 1):
            if spells(short, words[-size:]):
                found.append((short, " ".join(words[-size:]), match.group(0).strip()))
                break
    for match in SHORT_THEN_LONG.finditer(text):
        short, words = match.group(1), match.group(2).split()
        for size in range(len(words), 0, -1):
            if spells(short, words[:size]):
                found.append((short, " ".join(words[:size]).rstrip(".,;:"), match.group(0).strip()))
                break
    return found

def glossary_definition(unit):
    """A glossary entry written as its own heading - '108. SACP.' - over a paragraph that begins with
    what it stands for. Returns (acronym, long form, evidence) or None; both are the document's words."""
    chain = unit.get("heading_chain") or []
    short = re.sub(r"^\s*(?:\d+(?:\.\d+)*\.?|[A-Za-z][.)])\s+", "", chain[-1]).strip(" .:;") if chain else ""
    if not ACRONYM.fullmatch(short):
        return None
    words = re.split(r"[.:;]", unit.get("text") or "", maxsplit=1)[0].split()
    for size in range(1, min(len(words), 9) + 1):
        if spells(short, words[:size]):
            return short, " ".join(words[:size]), "under the heading '%s': %s" % (chain[-1], " ".join(words[:size]))
    return None

def first_clause(text):
    """The name a description gives, as written: its first line up to a comma, a colon, a stop or a bracket,
    and only if it is short enough to be a name. None otherwise."""
    clause = re.split(r"[,;:.(\n]", (text or "").strip(), maxsplit=1)[0].strip()
    return clause if clause and re.search(r"[A-Za-z]{2}", clause) and len(clause.split()) <= 8 else None

def model_concepts(units):
    """The model's concepts, by key: each with the names the code gives it and where, and how the
    model's own roxygen and help pages describe it, word for word. Names the code only calls (base
    functions) and generic names (x, data, result) are not concepts."""
    seeds = {}
    def add(name, ref, how):
        name = (name or "").strip()
        key = concept_key(" ".join(identifier_words(name)))
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_.]*", name) or len(name) < 2 or not key or key in GENERIC_NAMES:
            return
        seed = seeds.setdefault(key, {"key": key, "identifiers": {}, "described": {}, "how": []})
        if ref not in seed["identifiers"].setdefault(name, []):
            seed["identifiers"][name].append(ref)
        seed["how"].append("%s: %s in %s" % (name, how, ref))
    def describe(name, text, ref, how):
        key, phrase = concept_key(" ".join(identifier_words(name or ""))), first_clause(text)
        if key in seeds and phrase and concept_key(phrase) != key:
            if ref not in seeds[key]["described"].setdefault(phrase, []):
                seeds[key]["described"][phrase].append(ref)
            seeds[key]["how"].append("%s is described in the model as \u201c%s\u201d (%s, %s)" % (name, phrase, how, ref))
    for unit in units:
        code, data = unit.get("code") or {}, unit.get("data") or {}
        if unit["kind"] in (core.KIND_FUNCTION, core.KIND_FORMULA):
            add(unit["name"], unit["ref"], "a function" if unit["kind"] == core.KIND_FUNCTION else "a variable a statement sets")
        for formal, _ in code.get("formals") or []:
            add(formal, unit["ref"], "an argument of %s" % unit["name"])
        for written in code.get("symbols_written") or []:
            add(written, unit["ref"], "a variable the code sets")
        if data.get("object_name"):
            add(data["object_name"], unit["ref"], "a stored table")
            for column, _ in data.get("columns") or []:
                add(column, unit["ref"], "a column of %s" % data["object_name"])
    for unit in units:
        roxygen, helppage = unit.get("roxygen") or {}, unit.get("helppage") or {}
        for tag in roxygen.get("tags") or []:
            if tag.get("tag") == "param":
                describe(tag.get("name"), tag.get("text"), unit["ref"], "roxygen @param")
            elif tag.get("tag") in ("title", "description") and roxygen.get("documents_name"):
                describe(roxygen["documents_name"], tag.get("text"), unit["ref"], "roxygen")
        for name, text in helppage.get("arguments") or []:
            describe(name, text, unit["ref"], "help page")
        if helppage.get("title") and helppage.get("rd_name"):
            describe(helppage["rd_name"], helppage["title"], unit["ref"], "help page title")
    return seeds

def concept_forms(seeds):
    """Every form a model concept may be written in, as the words that must be found: the words of
    each name the code gives it, and each description the model gives it. [(words, key, how)]"""
    forms = []
    for key, seed in seeds.items():
        for name in seed["identifiers"]:
            forms.append((tuple(concept_singular(w) for w in identifier_words(name)), key, "the same words as the model's %s" % name))
        for phrase, refs in seed["described"].items():
            forms.append((tuple(concept_key(phrase).split()), key, "the words the model uses to describe %s (%s)" % (next(iter(seed["identifiers"])), refs[0])))
    return [(words, key, how) for words, key, how in forms if words]

def concept_spans(text, forms):
    """Where a text writes any form, as (the words exactly as the text writes them, key, how). Words
    are compared in lower case and singular; between them only spaces, hyphens or underscores may
    stand, so a match never runs across a full stop. The longest form wins at each place."""
    tokens = [(m.start(), m.end(), concept_singular(m.group(0).lower())) for m in CONCEPT_TOKEN.finditer(text or "")]
    starting = {}
    for words, key, how in forms:
        starting.setdefault(words[0], []).append((words, key, how))
    found, position = [], 0
    while position < len(tokens):
        best = None
        for words, key, how in starting.get(tokens[position][2], ()):
            end = position + len(words)
            if end <= len(tokens) and tuple(t[2] for t in tokens[position:end]) == words and (best is None or len(words) > len(best[0])):
                gaps = [text[tokens[j][1]:tokens[j + 1][0]] for j in range(position, end - 1)]
                if all(re.fullmatch(r"[\s\-_\u2010-\u2015]{1,3}", gap) for gap in gaps):
                    best = (words, key, how)
        if best:
            end = position + len(best[0])
            found.append((text[tokens[position][0]:tokens[end - 1][1]], best[1], best[2]))
            position = end
        else:
            position += 1
    return found

def concept_sources(ctx):
    """The units whose concepts are read: the three corners, in order."""
    return [(unit, corner) for kind, corner in (("chunks_canon", "canon"), ("chunks_doc", "doc"), ("model_units", "model"))
            for unit in ctx.read(kind)]

def concept_registry(sources, guessed=()):
    """The registry and each unit's concepts. Model first: its concepts; then every definition the
    documents give of an acronym whose one side is already a model concept's form, which adds the
    other side, with the unit that defines it; then every unit searched for every form. `guessed` are
    the model's accepted answers, (ref, words as written, key), kept apart. Returns (concepts, per unit)."""
    seeds = model_concepts([unit for unit, corner in sources if corner == "model"])
    forms = concept_forms(seeds)
    known = {words: key for words, key, _ in forms}
    for unit, corner in sources:
        if corner == "model":
            continue
        definitions = defined_acronyms(unit.get("text") or "") + ([glossary_definition(unit)] if glossary_definition(unit) else [])
        for short, long_form, evidence in definitions:
            short_words, long_words = (concept_singular(short.lower()),), tuple(concept_key(long_form).split())
            key = known.get(short_words) or known.get(long_words)
            if key:
                for words in (short_words, long_words):
                    if words not in known:
                        known[words] = key
                        forms.append((words, key, "defined in %s: \u201c%s\u201d" % (unit["ref"], evidence[:120])))
                        seeds[key]["how"].append("defined in %s: \u201c%s\u201d" % (unit["ref"], evidence[:120]))
    order = list(seeds)
    ids = {key: "K-%04d" % number for number, key in enumerate(order, start=1)}
    found = {key: {"canon": {}, "doc": {}, "model": {}} for key in seeds}
    per_unit, guessed_by_ref = [], {}
    for ref, words, key in guessed:
        guessed_by_ref.setdefault(ref, []).append((words, key))
    guesses = {key: {"canon": {}, "doc": {}} for key in seeds}
    for unit, corner in sources:
        spans = [(words, key) for words, key, _ in concept_spans(unit.get("text") or "", forms)]
        for words, key in spans:
            refs = found[key][corner].setdefault(words, [])
            if unit["ref"] not in refs:
                refs.append(unit["ref"])
        for words, key in guessed_by_ref.get(unit["ref"], []):
            if key in guesses and corner in guesses[key] and unit["ref"] not in guesses[key][corner].setdefault(words, []):
                guesses[key][corner][words].append(unit["ref"])
        spans += [pair for pair in guessed_by_ref.get(unit["ref"], []) if pair[1] in ids]
        per_unit.append({"unit_ref": unit["ref"], "corner": corner, "concepts": list(dict.fromkeys(ids[key] for _, key in spans)),
                         "as_written": list(dict.fromkeys(words for words, _ in spans))})
    concepts = [{"concept_id": ids[key], "key": key, "identifiers": seeds[key]["identifiers"], "described": seeds[key]["described"],
                 "found": found[key], "guessed": guesses[key], "how": list(dict.fromkeys(seeds[key]["how"]))} for key in order]
    return concepts, per_unit

def latest_concepts(read):
    """The registry and the units' concepts of the latest stage that wrote them: the model's, once
    judge-concepts has run, and code's before that. read is a store's or a step context's read."""
    for stage in ("ai", "code"):
        concepts = [r for r in read("concepts") if r.get("stage") == stage]
        if concepts:
            return concepts, [r for r in read("unit_concepts") if r.get("stage") == stage]
    return [], []

def extract_concepts(ctx):
    """Step 06, extract-concepts: the model's concepts, and every place the methodology, the
    documentation and the rest of the model write one of them in the same words - before any model
    call. Enforces: R3, R4"""
    concepts, per_unit = concept_registry(concept_sources(ctx))
    for record in concepts + per_unit:
        record["stage"] = "code"
    return core.StepResult({"concepts": concepts, "unit_concepts": per_unit},
                           {"concepts in the model": len(concepts),
                            "document units naming one": sum(1 for u in per_unit if u["corner"] != "model" and u["concepts"])},
                           ["%d concepts in the model; %d methodology and documentation units name one in the same words."
                            % (len(concepts), sum(1 for u in per_unit if u["corner"] != "model" and u["concepts"]))])

def likely_concepts(text, concepts, limit):
    """Model concepts a unit may name without the same words: a name whose parts begin words of the
    unit (adj_rating, 'adjusted rating'), an acronym-like name whose letters are the initials of words
    of the unit (sacp, 'stand-alone credit profile'), or a description sharing its words with the unit.
    Scored, best first; these are only candidates for the model to judge."""
    words = [w.lower() for w in re.findall(r"[A-Za-z][\w'-]*", text or "")]
    plain = {concept_singular(p) for w in words for p in re.split(r"[-']", w) if p}
    scores = {}
    for concept in concepts:
        score = 0.0
        for name in concept["identifiers"]:
            parts = identifier_words(name)
            begun = sum(1 for p in parts if p in plain or (len(p) >= 3 and any(w.startswith(p) for w in plain)))
            if len(parts) >= 2 and begun == len(parts):
                score += 2.0
            elif len(parts) >= 2:
                score += begun / (2.0 * len(parts))
            if len(parts) == 1 and 2 <= len(parts[0]) <= 6 and any(spells(parts[0], words[start:start + size])
                                                                    for start in range(len(words)) for size in range(len(parts[0]), len(parts[0]) + 3)):
                score += 2.0
        for phrase in concept["described"]:
            mine = set(concept_key(phrase).split()) - CONCEPT_STOP
            if mine:
                score += len(mine & plain) / len(mine)
        if score >= 1.0:
            scores[concept["concept_id"]] = score
    return sorted(scores, key=lambda cid: (-scores[cid], cid))[:limit]

def concept_match_question(group, shown, by_id, settings):
    """One question: some passages, and the model concepts they may name, each with the names the code
    gives it and how the model describes it. The model says which it names and copies the words."""
    subject = settings.get("concept_subject") or ""
    listing = "\n".join("[%s] %s%s" % (cid, " / ".join(by_id[cid]["identifiers"]),
                                       " \u2014 described in the model as: %s" % "; ".join(by_id[cid]["described"]) if by_id[cid]["described"] else "")
                        for cid in shown)
    passages = "\n".join("[%s]\n%s" % (ref, words) for ref, words in group)
    label = ("THE DOCUMENTS ARE ABOUT: %s\n" % subject if subject else "") + "THE MODEL'S CONCEPTS\n" + listing + "\nTHE PASSAGES"
    return narrow_question("match-concepts", group[0][0], [(label, passages)], settings,
                           more={"unit_texts": {ref: words for ref, words in group}, "concept_ids": list(shown)})

def verbatim(words, text):
    """The words exactly as the text writes them, found without regard to upper or lower case or the
    width of a space; None where the text does not hold them. What is kept is always the text's."""
    pattern = r"(?<![A-Za-z0-9])" + r"\s+".join(re.escape(w) for w in (words or "").split()) + r"(?![A-Za-z0-9])"
    found = re.search(pattern, text or "", re.I) if (words or "").strip() else None
    return found.group(0) if found else None

def validate_concepts(question, answer):
    """match-concepts: every passage and every concept must be one that was shown, and every name
    must be in its passage word for word, so no concept is ever one the model wrote. Enforces: R3"""
    units = answer.get("units")
    if not isinstance(units, dict):
        raise Rejected(core.REJECTION_REASONS[0])
    for ref, named in units.items():
        if ref not in question["unit_texts"] or not isinstance(named, list):
            raise Rejected(core.REJECTION_REASONS[1])
        for item in named:
            if not isinstance(item, dict) or item.get("concept") not in question["concept_ids"]:
                raise Rejected(core.REJECTION_REASONS[1])
            if not isinstance(item.get("words"), str) or verbatim(item["words"], question["unit_texts"][ref]) is None:
                raise Rejected(core.REJECTION_REASONS[2])

def judge_concepts(ctx):
    """Step 07b, judge-concepts: after the outline is confirmed, for the methodology and documentation
    units where code sees a likely model concept it could not prove - an abbreviation, an acronym, a
    description in other words - the model is shown the unit and those concepts and asked which it
    names, copying the words. A name not in the unit word for word refuses the whole answer. Accepted
    words are kept as the unit writes them, and on the Concepts sheet as the model's guesses, apart
    from what code proved. Without a model, code's registry stands. Enforces: R3, R4"""
    sources, settings = concept_sources(ctx), ctx.settings
    concepts, per_unit = concept_registry(sources)
    guessed = []
    if ctx.ask is not None and settings["concepts_with_ai"] and concepts:
        by_id, text_of = {c["concept_id"]: c for c in concepts}, {u["ref"]: u.get("text") or "" for u, _ in sources}
        proved = {u["unit_ref"]: set(u["concepts"]) for u in per_unit}
        wanted = [(unit["ref"], [cid for cid in likely_concepts(unit.get("text"), concepts, int(settings["concept_candidates_max"]))
                                 if cid not in proved.get(unit["ref"], set())]) for unit, corner in sources if corner != "model"]
        wanted = [(ref, cids) for ref, cids in wanted if cids]
        size, limit = max(1, int(settings["concept_batch"])), int(settings["max_passage_chars"])
        questions = []
        for start in range(0, len(wanted), size):
            batch = wanted[start:start + size]
            shown = sorted({cid for _, cids in batch for cid in cids})[:int(settings["concept_candidates_max"])]
            questions.append(concept_match_question([(ref, core.cut_text(text_of[ref], limit)) for ref, _ in batch], shown, by_id, settings))
        answers = ctx.ask([q for q in questions if not q["too_large"]]) if questions else {}
        for question in questions:
            for ref, named in (((answers.get(question["question_id"]) or {}).get("answer") or {}).get("units") or {}).items():
                for item in named:
                    words = verbatim(item["words"], text_of.get(ref, ""))
                    if words:
                        guessed.append((ref, words, by_id[item["concept"]]["key"]))
        concepts, per_unit = concept_registry(sources, guessed)
    for record in concepts + per_unit:
        record["stage"] = "ai"
    return core.StepResult({"concepts": concepts, "unit_concepts": per_unit},
                           {"concepts in the model": len(concepts), "names guessed by the model": len(guessed)},
                           ["%d concepts in the model; the model named %d more of them in other words, kept as guesses on the Concepts sheet."
                            % (len(concepts), len(guessed))])


# ---------------------------------------------------------------- the skill map-implementation: the implementation map's agents
# Code traced the data flow (step 05a); these agents work only where it stopped. The Tracer is given one
# gap on the path from a final output and a fixed list of actions; each turn it chooses one, code carries
# it out on the records and shows what it found, and every link it declares must copy the code word for
# word and use only names that code holds. Each turn is a question of its own, recorded, so a run replays
# without a model. The Namer gives each step a plain name, outside the accounting. The Auditor is code.
# The model chooses; code executes. Enforces: R3, R4, R5
MAP_ACTIONS = ("open_unit", "statements_setting", "callers_of", "return_of", "columns_of", "declare_edge", "declare_input", "done", "give_up")
INPUT_KINDS = ("argument", "stored data", "file", "hard-coded number", "from outside")

class MapTools:
    """What the Tracer can ask to see, answered by code from the traced flow and the units."""
    def __init__(self, flow, units):
        self.nodes = {r["node"]: r for r in flow if r["record_type"] == "node"}
        self.units = {u["ref"]: u for u in units}
        self.functions = {u["name"]: u for u in units if u["kind"] == core.KIND_FUNCTION and not u.get("inside")}
        self.tables = {r["name"]: r for r in self.nodes.values() if r["kind"] == "stored data"}

    def label(self, node_id):
        return self.nodes.get(node_id, {}).get("name", node_id)

    def sources(self, record):
        return ", ".join(self.label(s) for s in record["from"]) or "nothing the tool could see"

    def show(self, action, args):
        if action == "open_unit":
            return core.cut_text(self.units[args["ref"]]["text"], 3000)
        if action == "statements_setting":
            found = [r for r in self.nodes.values() if r["name"] == args["name"] and r["kind"] in ("value", "column")
                     and (r.get("function") == args["function"] or args["function"] in r.get("created_in", []))]
            return "\n".join("line %d: %s  (computed from: %s)" % (r["line"], r["code"], self.sources(r)) for r in found) or \
                "No statement of %s sets %s that the tool could see." % (args["function"], args["name"])
        if action == "callers_of":
            calls = [r for r in self.nodes.values() if r["kind"] == "call" and r["callee"] == args["function"]]
            return "\n".join("in %s, line %d: %s  (%s)" % (r["function"], r["line"], r["code"], "; ".join(
                "%s = %s" % (formal, ", ".join(self.label(s) for s in given)) for formal, given in r["bindings"].items())) for r in calls) or \
                "Nothing in the package calls %s." % args["function"]
        if action == "return_of":
            record = self.nodes.get("%s:return" % args["function"])
            return "%s returns what is computed from: %s" % (args["function"], self.sources(record)) if record else "The tool could not read %s." % args["function"]
        if action == "columns_of":
            return "%s has the columns: %s" % (args["table"], ", ".join(self.tables[args["table"]].get("columns") or []))
        return "Recorded."

def trace_question(state, tools, settings):
    """One turn of the Tracer: the gap, what the tool knows of its function, the names it may use, and
    every action taken so far with what it showed."""
    gap, function = state["gap"], state["gap"]["function"]
    known = [r for r in tools.nodes.values() if r.get("function") == function and r["kind"] in ("value", "argument", "return")]
    listing = "\n".join("  %s (%s): computed from %s" % (r["name"], r["kind"], tools.sources(r)) for r in known)
    place = ("function %s (%s), line %d: %s\ncode at this place: %s\n\nTHE WHOLE FUNCTION\n%s\n\n"
             "WHAT THE TOOL KNOWS OF THIS FUNCTION\n%s\n\nTHE PACKAGE\nfunctions: %s\nstored tables: %s") % (
        function, gap["function_ref"], gap["line"], gap["why"], gap["code"], tools.functions[function]["text"], listing or "  nothing",
        ", ".join("%s (%s)" % (name, unit["ref"]) for name, unit in sorted(tools.functions.items())), ", ".join(sorted(tools.tables)) or "none")
    history = "\n".join("%d. %s %s\n   the tool showed: %s" % (number, hop["action"], json.dumps(hop["args"], sort_keys=True), hop["shown"])
                        for number, hop in enumerate(state["hops"], start=1)) or "Nothing yet: this is the first action."
    texts = [tools.functions[function]["text"]] + [tools.units[ref]["text"] for ref in state["opened"]]
    # a block needs its label: the assembler drops a block without one, which sent the model a question with no gap in it
    return narrow_question("trace-gap", gap["function_ref"], [("THE PLACE THE TOOL COULD NOT FOLLOW", place), ("WHAT HAPPENED SO FAR", history)], settings,
                           more={"texts": texts, "refs": sorted(tools.units), "functions": sorted(tools.functions),
                                 "tables": sorted(tools.tables), "taken": [[hop["action"], hop["args"]] for hop in state["hops"]]})

def validate_trace(question, answer):
    """One action of the Tracer. It must be one of the list, with the arguments that action takes; every
    unit, function and table must be the package's; a quote must be in the code shown word for word and
    must hold every name the action declares, so no link rests on a name the model made up; and no
    action may be taken twice. Enforces: R3"""
    action, args = answer.get("action"), answer.get("args")
    if action not in MAP_ACTIONS or not isinstance(args, dict):
        raise Rejected(core.REJECTION_REASONS[0])
    if [action, args] in question["taken"]:
        raise Rejected(core.REJECTION_REASONS[5])
    text = lambda key: args.get(key) if isinstance(args.get(key), str) and args.get(key).strip() else None
    if action == "open_unit" and args.get("ref") not in question["refs"]:
        raise Rejected(core.REJECTION_REASONS[1])
    if action in ("statements_setting", "callers_of", "return_of") and args.get("function") not in question["functions"]:
        raise Rejected(core.REJECTION_REASONS[1])
    if action == "statements_setting" and not text("name"):
        raise Rejected(core.REJECTION_REASONS[0])
    if action == "columns_of" and args.get("table") not in question["tables"]:
        raise Rejected(core.REJECTION_REASONS[1])
    if action in ("declare_edge", "declare_input"):
        names = [args.get("value")] + list(args.get("from") or []) if action == "declare_edge" else [args.get("name")]
        if not text("quote") or not all(isinstance(n, str) and n.strip() for n in names) or \
           (action == "declare_edge" and not args.get("from")) or (action == "declare_input" and args.get("kind") not in INPUT_KINDS):
            raise Rejected(core.REJECTION_REASONS[0])
        squash = lambda words: re.sub(r"\s+", " ", words).strip()
        if not any(squash(args["quote"]) in squash(code) for code in question["texts"]):
            raise Rejected(core.REJECTION_REASONS[2])
        if not all(re.search(r"(?<![\w.$])%s(?![\w.])" % re.escape(n.strip()), args["quote"]) for n in names):
            raise Rejected(core.REJECTION_REASONS[1])
    if action in ("done", "give_up") and not text("because"):
        raise Rejected(core.REJECTION_REASONS[0])

def map_implementation(ctx):
    """Step 07d, map-implementation, the skill: the Tracer resolves the gaps on the path from each final
    output, turn by turn, within map_hops_max turns a gap and map_calls_max questions in all; the
    Auditor, code alone, says what is traced, what is open and what no output
    reaches. Without a model the gaps stay named and the steps unnamed. Enforces: R2, R3, R4, R5, R14"""
    flow, units, settings = ctx.read("dataflow"), ctx.read("model_units"), ctx.settings
    if not flow:
        return core.StepResult(messages=["No data flow was traced, so there is nothing to map."])
    outputs, _, not_reached = reading.decided_outputs(flow, {})
    reached = set()
    for output in outputs:
        reached |= reading.walk_dataflow(flow, output)[2]
    tools, gaps = MapTools(flow, units), [g for g in flow if g["record_type"] == "gap"]
    states = [{"gap": g, "hops": [], "edges": [], "inputs": [], "status": "", "opened": []} for g in gaps if g["function"] in reached]
    budget, asked = int(settings["map_calls_max"]), 0
    use_ai = ctx.ask is not None and settings["map_with_ai"]
    for _ in range(int(settings["map_hops_max"]) if use_ai else 0):
        turn = [state for state in states if not state["status"]][:max(0, budget - asked)]
        if not turn:
            break
        questions = [trace_question(state, tools, settings) for state in turn]
        asked += len(questions)
        answers = ctx.ask([q for q in questions if not q["too_large"]])
        for question, state in zip(questions, turn):
            final = answers.get(question["question_id"])
            if final is None or final["outcome"] != "accepted":
                state["status"] = "the AI's answer could not be used: %s" % (final or {}).get("outcome", "not asked, the question was too large").split(": ", 1)[-1]
                continue
            action, args = final["answer"]["action"], final["answer"]["args"]
            shown = tools.show(action, args)
            state["hops"].append({"action": action, "args": args, "shown": core.cut_text(shown, 600), "question_id": question["question_id"]})
            if action == "open_unit" and args["ref"] not in state["opened"]:
                state["opened"].append(args["ref"])
            elif action == "declare_edge":
                state["edges"].append({"value": args["value"], "from": list(args["from"]), "quote": args["quote"]})
            elif action == "declare_input":
                state["inputs"].append({"name": args["name"], "kind": args["kind"], "quote": args["quote"]})
            elif action == "done":
                state["status"] = "traced: %s" % args["because"]
            elif action == "give_up":
                state["status"] = "the code cannot tell: %s" % args["because"]
    for state in states:
        state["status"] = state["status"] or ("not asked: no model" if not use_ai else "stopped at the limit of %d turns" % int(settings["map_hops_max"])
                                              if len(state["hops"]) >= int(settings["map_hops_max"]) else "stopped at the limit of %d questions" % budget)
    traced = [s for s in states if s["status"].startswith("traced")]
    audit = {"final_outputs": outputs, "functions reached": sorted(reached), "not reached": not_reached,
             "gaps": len(gaps), "gaps on the path": len(states), "gaps traced": len(traced),
             "gaps open": [{"function": s["gap"]["function"], "line": s["gap"]["line"], "status": s["status"]} for s in states if s not in traced],
             "gaps not reached": [{"function": g["function"], "line": g["line"], "why": g["why"]} for g in gaps if g["function"] not in reached],
             "loops": sorted({loop for output in outputs for loop in reading.walk_dataflow(flow, output)[1]}), "questions": asked}
    traces = [{"gap": s["gap"], "hops": s["hops"], "edges": s["edges"], "inputs": s["inputs"], "status": s["status"]} for s in states]
    return core.StepResult({"map_traces": traces, "map_audit": [audit]},
                           {"gaps on the path": len(states), "gaps traced": len(traced)},
                           ["Final outputs %s: %d gaps on the path, %d traced by the AI; %d gaps in functions no output reaches."
                            % (", ".join(outputs), len(states), len(traced), len(audit["gaps not reached"]))])

# ---------------------------------------------------------------- step 05: build-graph
def structural_edges(units, provenance):
    """Edges the readers established: contains, calls, tested_by, documents, generated_from,
    reads_data. All are parsed from the files, none comes from the AI."""
    functions = {u["name"]: u["ref"] for u in units if u["kind"] == core.KIND_FUNCTION and not u["inside"]}
    data_units = {u["name"]: u["ref"] for u in units if u.get("data")}
    edges = []
    def add(source, target, kind, evidence=None):
        edges.append(core.Edge(source, target, kind, core.HOW_PARSED, provenance, evidence=evidence or {}))
    for unit in units:
        code = unit.get("code") or {}
        if unit.get("parent_ref"):
            add(unit["parent_ref"], unit["ref"], "contains")
        for name in code.get("calls", ()):
            if name in functions and functions[name] != unit["ref"] and unit["kind"] in (core.KIND_FUNCTION, core.KIND_TEST):
                add(functions[name] if unit["kind"] == core.KIND_TEST else unit["ref"],
                    unit["ref"] if unit["kind"] == core.KIND_TEST else functions[name],
                    "tested_by" if unit["kind"] == core.KIND_TEST else "calls")
        for read in code.get("reads_data", ()):
            if read["object"] in data_units:
                add(unit["ref"], data_units[read["object"]], "reads_data", dict(read))
        if unit.get("roxygen") and unit["roxygen"]["documents_ref"]:
            add(unit["ref"], unit["roxygen"]["documents_ref"], "documents")
        page = unit.get("helppage")
        if page:
            if page["generated_from_ref"]:
                add(unit["ref"], page["generated_from_ref"], "generated_from")
            target = functions.get(page["rd_name"]) or data_units.get(page["rd_name"])
            if target:
                add(unit["ref"], target, "documents")
    return edges

def build_graph(ctx):
    """Step 05, build-graph: nodes for every chunk and unit, structural edges,
    cross-references resolved within their own corner, and the bridge vocabulary."""
    canon, doc, units = ctx.read("chunks_canon"), ctx.read("chunks_doc"), ctx.read("model_units")
    lists = load_word_lists()
    records = [node_record(c["ref"], c["kind"], c["corner"]) for c in canon + doc]
    records += [node_record(u["ref"], u["kind"], "model") for u in units]
    edges, unresolved = structural_edges(units, ctx.provenance), []
    for corner_chunks in (canon, doc):
        for chunk in corner_chunks:
            same_file = [c for c in corner_chunks if c["source_file"] == chunk["source_file"]]
            for reference in chunk.get("refs_out", ()):
                targets = [ref for ref in resolve_reference(reference, same_file) if ref != chunk["ref"]]
                for target in targets:
                    edges.append(core.Edge(chunk["ref"], target, "cross_reference", core.HOW_PARSED, ctx.provenance,
                                             evidence={"as_written": reference}))
                own_caption = (chunk.get("caption") or "").lower().startswith(reference.lower())
                if not targets and not own_caption:
                    unresolved.append({"unit_ref": chunk["ref"], "reference": reference,
                                       "note": "'%s' is cited here but could not be found in %s" % (reference, chunk["source_file"])})
    bridge = harvest_bridge(canon, doc, units, lists, ctx.options["inputs"].get("glossary"))
    ledger = ledger_records(ctx.read("graph_ledger"), records + edges)
    return core.StepResult({"graph_ledger": ledger, "bridge_vocabulary": bridge, "unresolved_references": unresolved},
                             {"nodes": len(records), "edges": len(edges), "bridge entries": len(bridge)},
                             ["Graph version %s." % graph_version_id(ctx.read("graph_ledger") + ledger)])

# ---------------------------------------------------------------- step 06: find-candidates
SEARCHED_KINDS = (core.KIND_FUNCTION, core.KIND_FORMULA, core.KIND_TABLE, core.KIND_OBJECT, core.KIND_VIGNETTE)

def is_searched(unit, settings):
    """Which model units look for passages. Test blocks, files that were not read, package
    files without code, and (unless the setting says otherwise) supporting code by syntax
    are not sent to the search or to the judge; roxygen blocks and help pages take the
    tracing of the object they document."""
    if unit["kind"] not in SEARCHED_KINDS:
        return False
    if unit.get("data") and not unit["data"]["assessable"]:
        return False
    code = unit.get("code") or {}
    return not code.get("plumbing") or bool(settings["judge_supporting_code"])

def searched_text(representation, expansions, corner, proposed):
    """The "What was searched" sentence of one unit and corner."""
    words = list(dict.fromkeys(representation["fields"].get("name", []) + representation["fields"].get("heading", [])))[:8]
    if not words:
        words = list(dict.fromkeys(representation["fields"].get("body", [])))[:8]
    listed = []
    for word in words:
        entry = next((e for e in expansions if stem(e["term"].lower()) == word or e["term"] == word), None)
        listed.append("%s (%s)" % (AS_WRITTEN.get(word, word), entry["phrase"]) if entry else AS_WRITTEN.get(word, word))
    parts = ["Searched %s for: %s" % (CORNER_NAMES[corner], ", ".join(listed) or "no usable words")]
    if representation["symbols"]:
        parts.append("symbols " + ", ".join(representation["symbols"][:8]))
    if representation["numbers"]:
        parts.append("numbers " + ", ".join(representation["numbers"][:8]))
    return "; ".join(parts) + ". %d passage(s) proposed." % proposed

def search_one(source_ref, representation, corner, world, settings):
    """All signals for one unit and one target corner, fused into a shortlist of candidates."""
    targets, index, enabled = world["targets"][corner], world["index"][corner], settings["signals"]
    rankings, details = {}, {}
    mine = world["concepts_of"].get(source_ref, set())
    if "concepts" in enabled and mine:                   # shared concepts, the rarer the concept the stronger
        shared = {ref: sum(world["concept_idf"][c] for c in mine & world["concepts_of"].get(ref, set())) for ref in targets}
        rankings["concepts"] = ranked({ref: score for ref, score in shared.items() if score > 0})
        details["concepts"] = {ref: ", ".join(sorted(world["concept_names"][c] for c in mine & world["concepts_of"].get(ref, set()))[:3])
                               for ref in rankings["concepts"]}
    query = {word: 1.0 for words in representation["fields"].values() for word in words}
    expansions = expansions_for(representation, world["bridge_by_term"]) if "bridge" in enabled else []
    if "fields" in enabled:
        scores = bm25_scores(index, query, settings["bm25_k1"], settings["bm25_b"])
        rankings["fields"] = ranked(scores)
        details["fields"] = {ref: ", ".join(dict.fromkeys(shown(words[:4]))) for ref, (score, words) in scores.items()}
    if expansions:
        bridged = {word: (1.0 if e["count"] > 1 else 0.5) for e in expansions for word in e["words"] if word not in query}
        scores = bm25_scores(index, bridged, settings["bm25_k1"], settings["bm25_b"])
        rankings["bridge"] = ranked(scores)
        for ref, (score, words) in scores.items():
            entry = next(e for e in expansions if set(e["words"]) & set(words))
            details.setdefault("bridge", {})[ref] = "'%s' is described as '%s' (%s), which this passage mentions" % (
                entry["term_as_written"], entry["phrase"], entry["sources"][0])
    cited = []
    if "references" in enabled:
        for reference in representation.get("references", ()):
            for ref in resolve_reference(reference, targets.values()):
                cited.append(ref)
                details.setdefault("references", {})[ref] = reference
        rankings["references"] = list(dict.fromkeys(cited))
    if "anchors" in enabled:
        closeness = world["walk"].get(source_ref, {})
        rankings["anchors"] = ranked({ref: score for ref, score in closeness.items() if ref in targets})[:2 * int(settings["k_candidates"])]
        own = set(world["anchors"].get(source_ref, ()))
        for ref in rankings["anchors"]:
            common = sorted(own & set(world["anchors"].get(ref, ())), key=lambda a: (-world["weights"].get(a, 0), a))
            common = [a for a in common if a in world["weights"]]
            if common:
                details.setdefault("anchors", {})[ref] = " and ".join("%s %s" % (kind, key.split("@")[0]) for kind, key in common[:2])
    if "signatures" in enabled and representation.get("signature"):
        alike = {ref: overlap(representation["signature"], world["signatures"][ref]) for ref in targets if ref in world["signatures"]}
        rankings["signatures"] = ranked({ref: score for ref, score in alike.items() if score >= 0.5})
    if representation.get("table"):
        shapes = {ref: table_shape_score(representation["table"], targets[ref]["table"], world["lists"]["stop"])
                  for ref in targets if targets[ref].get("table")}
        rankings["table shape"] = ranked({ref: score for ref, score in shapes.items() if score >= 1.0})
    if "propagation" in enabled:
        inherited = world["propagated"].get((source_ref, corner), [])
        rankings["propagation"] = [ref for ref, _ in inherited]
        details["propagation"] = dict(inherited)
    fused = fuse(rankings, settings, cited)
    candidates = []
    for rank, (ref, score, signals) in enumerate(fused, start=1):
        reason = reason_text(signals, {signal: details.get(signal, {}).get(ref, "") for signal in signals})
        candidates.append(core.Candidate(source_ref, ref, corner, rank, round(score, 6), signals, reason,
                                           suggestion_only="propagation" in signals and len(signals) == 1,
                                           search_pass=world["search_pass"]))
    ranked_anywhere = sorted({ref for refs in rankings.values() for ref in refs})
    record = {"unit_ref": source_ref, "target_corner": corner, "search_pass": world["search_pass"],
              "query_fields": representation["fields"], "expansions": [(e["term"], e["words"], e["sources"][0]) for e in expansions],
              "anchors_used": [list(a) for a in world["anchors"].get(source_ref, ()) if a in world["weights"]][:20],
              "per_signal": {signal: len(refs) for signal, refs in rankings.items()}, "shortlist": [c.target_ref for c in candidates],
              "ranked_anywhere": ranked_anywhere, "searched_text": searched_text(representation, expansions, corner, len(candidates)),
              "note": "" if candidates else "Nothing in %s shares a word, a symbol, a number or a citation with this unit." % CORNER_NAMES[corner]}
    return candidates, record

def build_world(ctx, search_pass):
    """Everything the search needs, built once per step: representations of all units and
    chunks, one BM25 index per corner, the bridge vocabulary, anchors and the walk."""
    settings, lists = ctx.settings, load_word_lists()
    trivial = set(settings["trivial_numbers"])
    canon, doc, units = (ctx.read(kind) for kind in ("chunks_canon", "chunks_doc", "model_units"))
    tables = {t["unit_ref"]: t for t in ctx.read("parameter_tables")}
    by_ref = {u["ref"]: u for u in units}
    documented_by = {u["roxygen"]["documents_ref"]: u for u in units if u.get("roxygen") and u["roxygen"]["documents_ref"]}
    representations = {}
    for chunk in canon + doc:
        representations[chunk["ref"]] = dict(chunk_fields(chunk, lists, trivial), references=chunk.get("refs_out", ()),
                                             heading=split_words(" ".join(chunk["heading_chain"][-1:]), lists["stop"]))
        equation = chunk.get("equation") or {}
        if equation.get("expression"):
            representations[chunk["ref"]]["signature"] = formula_signature(equation["expression"])
    for unit in units:
        representation = unit_fields(unit, by_ref, documented_by, lists, trivial)
        about = documented_by.get(unit["ref"])
        representation["references"] = reading.cross_references(unit["text"] + " " + (about["text"] if about else ""),
                                                                        {"cross_reference_labels": ["Table", "Figure", "Section", "Equation", "Annex", "Appendix"]})
        expression = (unit.get("code") or {}).get("expression") or (unit.get("code") or {}).get("composed")
        if expression:
            representation["signature"] = formula_signature(expression)
        if unit["ref"] in tables:
            representation["table"] = tables[unit["ref"]]
        representations[unit["ref"]] = representation
    bridge_by_term = {}
    for entry in ctx.read("bridge_vocabulary"):
        bridge_by_term.setdefault(entry["term"], []).append(entry)
    definitions = {}
    for entries in bridge_by_term.values():
        for entry in entries:
            for source in entry["sources"]:
                if source[:2] in ("C-", "D-"):
                    definitions.setdefault(entry["term"], set()).add(source)
    scope = {term: "%s@%s" % (term, sorted(sources)[0]) for term, sources in definitions.items() if len(sources) == 1}
    anchors, mentions = {}, {}
    for ref, representation in representations.items():
        anchors[ref] = anchors_of(representation, representation.get("heading"), scope)
        for anchor in anchors[ref]:
            mentions.setdefault(anchor, []).append(ref)
    weights = anchor_weights(mentions, settings)
    searched = [u["ref"] for u in units if is_searched(u, settings)] + [c["ref"] for c in doc if c.get("checkable")]
    walk = restart_walk(mentions, weights, searched, settings) if "anchors" in settings["signals"] else {}
    targets = {"canon": {c["ref"]: c for c in canon}, "doc": {c["ref"]: c for c in doc},
               "model": {u["ref"]: u for u in units if is_searched(u, settings)}}
    index = {corner: build_index({ref: representations[ref] for ref in targets[corner]}) for corner in targets}
    signatures = {ref: r["signature"] for ref, r in representations.items() if r.get("signature")}
    concepts, unit_concepts = latest_concepts(ctx.read)
    concepts_of = {u["unit_ref"]: set(u["concepts"]) for u in unit_concepts}
    frequency = {}
    for found in concepts_of.values():
        for concept in found:
            frequency[concept] = frequency.get(concept, 0) + 1
    total = max(1, len(concepts_of))
    return {"concepts_of": concepts_of, "concept_idf": {c: math.log(1 + total / n) for c, n in frequency.items()},
            "concept_names": {c["concept_id"]: next(iter(c["identifiers"]), c["key"]) for c in concepts},
            "representations": representations, "targets": targets, "index": index, "bridge_by_term": bridge_by_term,
            "anchors": anchors, "weights": weights, "walk": walk, "signatures": signatures, "lists": lists,
            "units": units, "doc": doc, "canon": canon, "propagated": {}, "search_pass": search_pass}

def propagated_candidates(world, graph):
    """S5, pass 2 only. Once the judge accepted that a function corresponds to a passage, the
    statements inside it, the functions it calls, the data it reads and its tests inherit
    that passage, and the tables the passage cites, as SUGGESTIONS. Never as links."""
    inherited = {}
    names = {u["ref"]: "%s %s" % (u["kind"].lower(), u["name"]) for u in world["units"]}
    for unit in world["units"]:
        accepted = [e for e in links_of(graph, unit["ref"], "C") + links_of(graph, unit["ref"], "D") if e["relation"] in core.LINKING_RELATIONS]
        if unit["kind"] != core.KIND_FUNCTION or not accepted:
            continue
        heirs = [e["target"] for e in graph["out"].get(unit["ref"], []) if e["kind"] in ("contains", "calls", "reads_data", "tested_by")]
        for edge in accepted:
            corner = "canon" if edge["target"].startswith("C") else "doc"
            cited = [e["target"] for e in graph["out"].get(edge["target"], []) if e["kind"] == "cross_reference"]
            for heir in heirs:
                for target in [edge["target"]] + cited:
                    detail = "%s (%s), which the AI judged to correspond to %s" % (unit["ref"], names[unit["ref"]], edge["target"])
                    inherited.setdefault((heir, corner), []).append((target, detail))
    return inherited

def find_candidates(ctx):
    """Step 06 (pass 1) and step 08 (pass 2), find-candidates. Pass 1 searches for every
    searched model unit in the methodology and the documentation, and for every checkable
    documentation passage in the methodology. Pass 2 searches again, with propagation, only
    for units still without a methodology link, and looks in the package for documentation
    passages still without any link. Every search leaves a search record. Enforces: R2"""
    search_pass = int(ctx.options.get("pass", 1))
    world = build_world(ctx, search_pass)
    graph = load_graph(ctx.read("graph_ledger"))
    jobs = []
    if search_pass == 1:
        jobs += [(u["ref"], corner) for u in world["units"] if is_searched(u, ctx.settings) for corner in ("canon", "doc")]
        jobs += [(c["ref"], "canon") for c in world["doc"] if c.get("checkable")]
    else:
        world["propagated"] = propagated_candidates(world, graph)
        linked = lambda ref, prefix: any(e["relation"] in core.LINKING_RELATIONS for e in links_of(graph, ref, prefix))
        jobs += [(ref, corner) for (ref, corner) in sorted(world["propagated"]) if not linked(ref, "C" if corner == "canon" else "D")
                 and ref in world["targets"]["model"]]
        jobs += [(c["ref"], "model") for c in world["doc"] if c.get("checkable") and not linked(c["ref"], "C") and not linked(c["ref"], "M")]
    candidates, records = [], []
    for source_ref, corner in jobs:
        found, record = search_one(source_ref, world["representations"][source_ref], corner, world, ctx.settings)
        candidates.extend(found)
        records.append(record)
    return core.StepResult({"candidates": candidates, "search_records": records},
                             {"searches": len(jobs), "candidates": len(candidates),
                              "searches without any proposal": sum(1 for r in records if not r["shortlist"])}, [])

# ---------------------------------------------------------------- questions for the AI
def cut_code(text, limit):
    """Cut a long function around its formula lines: the header, then every line that
    computes something with two lines of context, until the limit is reached."""
    if len(text) <= limit:
        return text
    lines = text.split("\n")
    wanted = {0, 1, len(lines) - 1}
    for number, line in enumerate(lines):
        if re.search(r"(<-|=).*[-+*/^]|\b(pnorm|qnorm|exp|log|sqrt|pmax|pmin|max|min)\(", line):
            wanted.update(range(max(0, number - 2), min(len(lines), number + 3)))
    kept, size, previous = [], 0, -1
    for number in sorted(wanted):
        if size + len(lines[number]) > limit:
            break
        if number != previous + 1:
            kept.append("# [... lines cut ...]")
        kept.append(lines[number])
        size, previous = size + len(lines[number]) + 1, number
    return "\n".join(kept)

def passage_text(node, settings, keep_words=()):
    """How a chunk or a unit is shown as a lettered passage. A table is never pasted whole:
    its header, three rows and its dimensions."""
    table = node.get("table")
    if table and table.get("header") is not None and node.get("kind") == "Table":
        rows = table.get("rows", [])
        shown = ["; ".join(table["header"])] + ["; ".join(row) for row in rows[:3]]
        return "\n".join(shown) + "\n(table of %d rows and %d columns%s)" % (
            len(rows), len(table["header"]), "; caption: " + node["caption"] if node.get("caption") else "")
    if node.get("code") is not None or node.get("file"):
        return cut_code(node["text"], int(settings["max_passage_chars"]))
    return core.cut_text(node["text"], int(settings["max_passage_chars"]), keep_words)

def passage_label(node):
    """The heading line of a lettered passage: where it stands, never its reference."""
    if node.get("heading_chain") is not None:
        place = " > ".join(node["heading_chain"][-3:]) or node["source_file"]
        return "%s, %s%s" % (place, node["kind"].lower(), " %d" % node["para_no"] if node.get("para_no") else "")
    lines = "lines %d-%d" % tuple(node["lines"]) if node.get("lines") else ""
    return "%s `%s`, %s %s" % (node["kind"].lower(), node["name"], node["file"], lines)

def letters_for(count):
    """A, B, ... Z, AA, AB, ... for the passages of one question."""
    alphabet = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    return [alphabet[i] if i < 26 else alphabet[i // 26 - 1] + alphabet[i % 26] for i in range(count)]

def choose_decoys(unit_ref, shortlist, never, pool, anchors, count):
    """Planted control passages: chosen by a hash of the unit reference (never at random) from
    passages that share no anchor with the unit and appear nowhere in its rankings."""
    own = set(anchors.get(unit_ref, ()))
    eligible = [ref for ref in sorted(pool) if ref not in never and ref not in shortlist
                and not own & set(anchors.get(ref, ())) and len(pool[ref].get("text", "")) > 40]
    eligible.sort(key=lambda ref: hashlib.sha256((unit_ref + ref).encode("utf-8")).hexdigest())
    return eligible[:count]

def assemble_question(question_type, unit_ref, blocks, passages, prompt, settings, planted=(), more=None):
    """Put one question together. blocks: [(label line, text)]; passages: [(ref, label, text)].
    Passages are lettered in an order derived from a hash, never by score, so position
    carries no hint. The question id is the hash of the full prompt. Enforces: R5"""
    ordered = sorted(passages, key=lambda p: hashlib.sha256((unit_ref + "|" + p[0]).encode("utf-8")).hexdigest())
    letters = dict(zip(letters_for(len(ordered)), ordered))
    shown = "\n".join("[%s] %s\n<<<\n%s\n>>>" % (letter, label, text) for letter, (ref, label, text) in letters.items())
    main = prompt["main"].replace("[[PASSAGES]]", shown)
    for slot, (label, text) in zip(("[[UNIT]]", "[[ABOUT]]"), list(blocks) + [("", "")]):
        main = main.replace(slot, "%s\n<<<\n%s\n>>>" % (label, text) if label else "")
    main = re.sub(r"\n{2,}", "\n", main)
    question = {"question_type": question_type, "unit_ref": unit_ref, "system_prompt": prompt["system"], "main_prompt": main,
                "letters": {letter: ref for letter, (ref, _, _) in letters.items()},
                "planted": [letter for letter, (ref, _, _) in letters.items() if ref in planted],
                "passage_texts": {letter: text for letter, (_, _, text) in letters.items()},
                "unit_text": "\n".join(text for _, text in blocks), "strip_patterns": list(settings["strip_patterns"]),
                "estimated_tokens": core.estimate_tokens(main, blocks[0][1] if blocks else ""),
                "question_id": core.sha256_text(prompt["version"] + "\n" + prompt["system"] + "\n" + main)}
    question.update(more or {})
    question["too_large"] = question["estimated_tokens"] > core.prompt_budget(settings, prompt["system"])
    return question

def narrow_question(question_type, unit_ref, blocks, settings, passages=(), more=None):
    """The narrow questions of the checks (map-table-columns, align-symbols, read-formula-
    from-prose, check-rule): same assembly, same budget, same validators."""
    return assemble_question(question_type, unit_ref, blocks, list(passages), core.load_prompt(question_type), settings, more=more)

def judge_question(source, corner, candidates, world, settings):
    """The judge question of one unit (or documentation passage) and one target corner."""
    is_chunk = source.get("heading_chain") is not None
    question_type = {("canon", False): "judge-unit-to-canon", ("doc", False): "judge-unit-to-doc",
                     ("canon", True): "judge-doc-to-canon", ("model", True): "judge-doc-to-model"}[(corner, is_chunk)]
    pool = world["targets"][corner]
    shortlist = [c["target_ref"] for c in candidates]
    record = next((r for r in world["search_records"] if r["unit_ref"] == source["ref"] and r["target_corner"] == corner), {})
    decoys = choose_decoys(source["ref"], shortlist, set(record.get("ranked_anywhere", ())), pool, world["anchors"],
                           1 if len(shortlist) < 6 else 2)
    words = [word for words in world["representations"][source["ref"]]["fields"].values() for word in words]
    passages = [(ref, passage_label(pool[ref]), passage_text(pool[ref], settings, words)) for ref in shortlist + decoys]
    unit_text = core.cut_text(source["text"], int(settings["max_unit_chars"])) if is_chunk else cut_code(source["text"], int(settings["max_unit_chars"]))
    blocks = [("UNIT (%s)" % passage_label(source), unit_text)]
    about = world["documented_by"].get(source["ref"]) or world["documented_by"].get(source.get("parent_ref") or "")
    if about and not is_chunk:
        blocks.append(("WHAT THE PACKAGE SAYS ABOUT IT", core.cut_text(re.sub(r"(?m)^\s*#' ?", "", about["text"]), 1200)))
    return assemble_question(question_type, source["ref"], blocks, passages, core.load_prompt(question_type),
                             settings, planted=decoys, more={"target_corner": corner})

# ---------------------------------------------------------------- validators: code decides what is usable
class Rejected(Exception):
    """An answer that cannot be used. The message is one reason from the fixed plain list."""


def check_quote(quote, text, required=True):
    """A quotation must be verbatim after white-space normalisation, contiguous, without an
    ellipsis, and of a sensible length."""
    quote = core.normalise_text(quote if isinstance(quote, str) else "")
    if not quote and not required:
        return
    flat = core.normalise_text(text)
    if len(quote) < 4 or len(quote) > 400 or "..." in quote or "\u2026" in quote or quote not in flat:
        raise Rejected(core.REJECTION_REASONS[2])

JUDGE_RELATIONS = {"judge-unit-to-canon": ("implements", "partly implements", "deviates from", "merely related"),
                   "judge-unit-to-doc": ("describes", "consistent with", "inconsistent with"),
                   "judge-doc-to-canon": ("consistent with", "inconsistent with", "merely related"),
                   "judge-doc-to-model": ("describes", "inconsistent with")}

def validate_judge(question, answer):
    """The answer to a judge question: its shape, the letters it names, its quotations, planted passages, self-contradiction."""
    allowed = {"matches", "none_reason", "states_nothing_checkable"}
    if not isinstance(answer.get("matches"), list) or set(answer) - allowed:
        raise Rejected(core.REJECTION_REASONS[0])
    for match in answer["matches"]:
        confidence = match.get("confidence") if isinstance(match, dict) else None
        if not isinstance(match, dict) or set(match) - {"letter", "relation", "confidence", "quote_from_passage", "quote_from_unit"} \
                or match.get("relation") not in JUDGE_RELATIONS[question["question_type"]] \
                or isinstance(confidence, bool) or not isinstance(confidence, int) or not 0 <= confidence <= 100:
            raise Rejected(core.REJECTION_REASONS[0])
        if match.get("letter") not in question["letters"]:
            raise Rejected(core.REJECTION_REASONS[1])
    for match in answer["matches"]:
        check_quote(match.get("quote_from_passage"), question["passage_texts"][match["letter"]])
        check_quote(match.get("quote_from_unit"), question["unit_text"])
    if any(match["letter"] in question["planted"] and match["relation"] != "merely related" for match in answer["matches"]):
        raise Rejected(core.REJECTION_REASONS[3])
    if answer["matches"] and str(answer.get("none_reason") or "").strip():
        raise Rejected(core.REJECTION_REASONS[4])

def validate_narrow(question, answer):
    """The narrow question types. Each check mirrors what the prompt allows."""
    kind = question["question_type"]
    if kind == "second-opinion":
        if not isinstance(answer.get("differences"), list):
            raise Rejected(core.REJECTION_REASONS[0])
        for difference in answer["differences"]:
            if not isinstance(difference, dict) or difference.get("letter") not in question["letters"]:
                raise Rejected(core.REJECTION_REASONS[1])
            check_quote(difference.get("quote_from_passage"), question["passage_texts"][difference["letter"]])
            check_quote(difference.get("quote_from_unit"), question["unit_text"])
    elif kind == "interpret-code":
        words = answer.get("interpretation")
        if not isinstance(words, str) or not 3 <= len(words.split()) <= 150:
            raise Rejected(core.REJECTION_REASONS[0])
        check_quote(answer.get("quote_from_unit"), question["code_text"])     # it has to rest on code that is there
    elif kind == "check-rule":
        if answer.get("outcome") not in ("applied", "applied differently", "not applied"):
            raise Rejected(core.REJECTION_REASONS[0])
        check_quote(answer.get("quote_from_passage"), question["rule_text"])
        check_quote(answer.get("quote_from_unit"), question["code_text"], required=answer["outcome"] != "not applied")
    elif kind == "align-symbols":
        pairs = answer.get("alignment")
        if not isinstance(pairs, list) or not isinstance(answer.get("cannot_align", False), bool):
            raise Rejected(core.REJECTION_REASONS[0])
        codes, equations = [p.get("code") for p in pairs if isinstance(p, dict)], [p.get("equation") for p in pairs if isinstance(p, dict)]
        if len(codes) != len(pairs) or set(codes) - set(question["code_symbols"]) or set(equations) - set(question["equation_symbols"]):
            raise Rejected(core.REJECTION_REASONS[1])
        if len(set(codes)) != len(codes) or len(set(equations)) != len(equations) or (pairs and answer.get("cannot_align")):
            raise Rejected(core.REJECTION_REASONS[4])
    elif kind == "read-formula-from-prose":
        formula = answer.get("formula")
        if not isinstance(formula, str) or (not formula.strip()) != bool(answer.get("no_formula_stated")):
            raise Rejected(core.REJECTION_REASONS[4] if isinstance(formula, str) else core.REJECTION_REASONS[0])
        if formula.strip():
            try:
                tree = reading.parse_formula(formula, question["notation"])
            except reading.NotReadable:
                raise Rejected(core.REJECTION_REASONS[0])
            text = question["unit_text"].lower()
            if any(symbol.lower().replace("_", " ") not in text and symbol.lower() not in text for symbol in core.expr_symbols(tree)):
                raise Rejected(core.REJECTION_REASONS[2])
    elif kind == "map-table-columns":
        if answer.get("table") == "NONE":
            return
        if answer.get("table") not in question["letters"]:
            raise Rejected(core.REJECTION_REASONS[1])
        columns, key = answer.get("columns"), answer.get("key")
        if not isinstance(columns, list) or not isinstance(key, dict) or not columns:
            raise Rejected(core.REJECTION_REASONS[0])
        ours, theirs = question["package_header"], question["other_headers"][answer["table"]]
        pairs = [(c.get("package"), c.get("other")) for c in columns if isinstance(c, dict)]
        if len(pairs) != len(columns) or any(a not in ours or b not in theirs for a, b in pairs + [(key.get("package"), key.get("other"))]):
            raise Rejected(core.REJECTION_REASONS[1])
        if len({a for a, _ in pairs}) != len(pairs) or len({b for _, b in pairs}) != len(pairs):
            raise Rejected(core.REJECTION_REASONS[4])

def validate_answer(question, text):
    """The validators, in a fixed order: remove any thought block and code fence; take the last
    balanced JSON object; parse it strictly; check its shape, its letters, its quotations,
    the planted passages and self-contradiction. Returns ("accepted", answer) or
    ("rejected: <one reason from the fixed list>", None). Enforces: R3"""
    try:
        if not isinstance(text, str) or not text.strip():
            raise Rejected(core.REJECTION_REASONS[0])
        for pattern in question.get("strip_patterns", ()):
            text = re.sub(pattern, "", text)
        text = re.sub(r"```[a-zA-Z]*", "", text)
        found = core.last_json_object(text)
        try:
            answer = core.strict_json(found) if found else None
        except ValueError:
            answer = None
        if not isinstance(answer, dict):
            raise Rejected(core.REJECTION_REASONS[0])
        if question["question_type"] == "slice-rules":
            core.validate_slice_rules(question, answer, Rejected)
        elif question["question_type"] == "package-plan":
            core.validate_package_plan(question, answer, Rejected)
        elif question["question_type"] == "match-concepts":
            validate_concepts(question, answer)
        elif question["question_type"] == "trace-gap":
            validate_trace(question, answer)
        elif question["question_type"] in JUDGE_RELATIONS:
            validate_judge(question, answer)
        else:
            validate_narrow(question, answer)
    except Rejected as problem:
        return "rejected: %s" % problem, None
    return "accepted", answer

_NOT_SHOWN = re.compile(r"\b(%s)\b" % "|".join(("sev" "er(e|ity)", "crit" "ical", "maj" "or", "min" "or", "err" "ors?", "find" "ings?",
                        "mater" "ial(ity)?", "(high|medium|low)[- ](risk|priority|impact|rating)", "non-?compl" "ian(t|ce)", "breach")), re.I)

def shown_ai_text(text):
    """Text written by the model passes the same plain-language filter as the tool's own wording:
    if it rates seriousness or uses a policy term, a fixed sentence is shown instead and the
    full text stays in the audit records."""
    text = core.normalise_text(str(text or ""))[:300]
    return core.AI_WORDING_NOT_SHOWN if _NOT_SHOWN.search(text) else text

# ---------------------------------------------------------------- steps 07 and 09: judge-links
def needs_second_opinion(settings, target):
    """`unchecked_only`: a second, oppositely framed question is asked for accepted links that
    no deterministic check will back, that is, passages without a formula or a table."""
    mode = settings["second_opinion"]
    backed = bool(target.get("table")) or bool((target.get("equation") or {}).get("readable"))
    return mode == "all" or (mode == "unchecked_only" and not backed and target.get("heading_chain") is not None)

# ---------------------------------------------------------------- what each piece of code does, in plain words
INTERPRETED_KINDS = (core.KIND_FUNCTION, core.KIND_FORMULA, core.KIND_TOPLEVEL, core.KIND_TEST)

def package_outline(units, package):
    """The whole package in a few lines, as every interpretation question sees it: its name and
    title, then for each file the functions defined there with their arguments and the first
    line of their documentation, and the stored data. Also returns the documentation block of
    each unit, by the reference of the unit it documents."""
    described = {u["roxygen"]["documents_ref"]: u for u in units if u.get("roxygen") and u["roxygen"].get("documents_ref")}
    title = next((row["value"] for row in package.get("rows", []) if row.get("item") == "Title"), "")
    lines, by_file = ["Package %s %s: %s" % (package.get("name", ""), package.get("version", ""), title)], {}
    for unit in units:
        if unit["kind"] == core.KIND_FUNCTION and not unit["inside"]:
            tags = described[unit["ref"]]["roxygen"]["tags"] if unit["ref"] in described else []
            says = next((tag["text"].split("\n")[0] for tag in tags if tag["tag"] in ("title", "description")), "")
            formals = ", ".join(name for name, _ in (unit.get("code") or {}).get("formals", []))
            by_file.setdefault(unit["file"], []).append("%s(%s)%s" % (unit["name"], formals, " - " + says if says else ""))
        elif unit["kind"] in (core.KIND_TABLE, core.KIND_OBJECT):
            by_file.setdefault(unit["file"], []).append("stored data %s" % unit["name"])
    return "\n".join(lines + ["%s: %s" % (file, "; ".join(by_file[file])) for file in sorted(by_file)]), described

def interpret_question(unit, functions, outline, described, settings):
    """The question about one piece of code. The piece is shown whole; around it goes what a
    person would look up to understand it: the function a statement sits inside, the
    documentation the package gives, what calls it and what it calls, the stored data it reads,
    and the outline of the whole package cut around the piece's own name."""
    inside = functions.get(unit["inside"]) if unit["inside"] else None
    about, owner = [], inside or unit
    if inside is not None:
        about.append("This piece is one statement inside the function %s. The whole function:\n%s" % (inside["name"], cut_code(inside["text"], settings["max_unit_chars"])))
    if owner["ref"] in described:
        about.append("What the package's own documentation says:\n%s" % core.cut_text(described[owner["ref"]]["text"], settings["max_passage_chars"]))
    code = owner.get("code") or {}
    callers = sorted(name for name, other in functions.items() if owner["name"] and owner["name"] in (other.get("code") or {}).get("calls", []))
    facts = [("It is called by", callers), ("Within this package it calls", sorted(c for c in code.get("calls", []) if c in functions)),
             ("It reads the stored data", sorted({r["object"] for r in code.get("reads_data", [])}))]
    about.extend("%s: %s." % (says, ", ".join(names)) for says, names in facts if names)
    about.append("The whole package:\n%s" % core.cut_text(outline, 4000, keep_words=(owner["name"],)))
    where = "%s%s, %s%s" % (unit["kind"], " " + unit["name"] if unit["name"] else "", unit["file"],
                           " lines %d-%d" % tuple(unit["lines"]) if unit.get("lines") else "")
    blocks = [("THE PIECE OF CODE (%s)" % where, cut_code(unit["text"], settings["max_unit_chars"])),
              ("WHERE IT SITS IN THE PACKAGE", "\n\n".join(about))]
    return narrow_question("interpret-code", unit["ref"], blocks, settings, more={"code_text": unit["text"]})

def interpret_code(ctx):
    """Step 07a, interpret-code. One question per function, formula statement, top-level
    statement and test block: what does this piece do, given where it sits in the package?
    Accepted answers fill the column "LLM Interpretation" of Chunks_Model. An interpretation
    is an aid to reading and nothing more: it gives no status, raises no flagged item and takes
    no part in the coverage identity, and a question that fails leaves a plain note. Enforces: R3"""
    units = ctx.read("model_units")
    if not ctx.settings.get("interpret_code", True):
        return core.StepResult(messages=["Interpreting the code is switched off (setting interpret_code)."])
    outline, described = package_outline(units, (ctx.read("package_info") or [{}])[0])
    functions = {u["name"]: u for u in units if u["kind"] == core.KIND_FUNCTION and not u["inside"]}
    questions, records = {}, []
    for unit in units:
        if unit["kind"] in INTERPRETED_KINDS and unit["text"].strip():
            question = interpret_question(unit, functions, outline, described, ctx.settings)
            if question["too_large"]:
                records.append({"unit_ref": unit["ref"], "interpretation": "", "question_id": "", "note": "Not asked: the question was too large to ask."})
            else:
                questions[question["question_id"]] = question
    answers = ctx.ask(list(questions.values())) if questions and ctx.ask else {}
    for question_id in sorted(questions, key=lambda key: questions[key]["unit_ref"]):
        final = answers.get(question_id)
        if final is not None and final["outcome"] == "accepted":
            records.append({"unit_ref": questions[question_id]["unit_ref"], "interpretation": final["answer"]["interpretation"].strip(),
                            "quote_from_unit": final["answer"]["quote_from_unit"], "question_id": question_id, "note": ""})
        else:
            reason = (final or {}).get("outcome", "failed: no answer was obtained").split(": ", 1)[-1]
            records.append({"unit_ref": questions[question_id]["unit_ref"], "interpretation": "", "question_id": question_id,
                            "note": "The AI's answer could not be used: %s." % reason})
    done = sum(1 for record in records if record["interpretation"])
    return core.StepResult({"interpretations": sorted(records, key=lambda record: record["unit_ref"])},
                             counts={"pieces_of_code": len(records), "interpreted": done},
                             messages=["%d of %d pieces of code were given an interpretation by the AI." % (done, len(records))])

def judge_links(ctx):
    """Steps 07 and 09, judge-links. Builds one question per unit and corner, asks them
    through ask(), and turns ACCEPTED answers into `corresponds` edges with the relation
    word, the confidence, both quotations and the proposal reason. Rejected answers leave
    the unit without a link and are recorded for account-coverage. Enforces: R3, R4"""
    search_pass, settings = int(ctx.options.get("pass", 1)), ctx.settings
    world = build_world(ctx, search_pass)
    world["search_records"] = [r for r in ctx.read("search_records") if r["search_pass"] == search_pass]
    world["documented_by"] = {u["roxygen"]["documents_ref"]: u for u in world["units"] if u.get("roxygen") and u["roxygen"]["documents_ref"]}
    sources = dict(world["targets"]["model"], **{c["ref"]: c for c in world["doc"]})
    grouped = {}
    for candidate in ctx.read("candidates"):
        if candidate["search_pass"] == search_pass:
            grouped.setdefault((candidate["unit_ref"], candidate["target_corner"]), []).append(candidate)
    questions, skipped = {}, []
    for (unit_ref, corner) in sorted(grouped):
        question = judge_question(sources[unit_ref], corner, grouped[(unit_ref, corner)], world, settings)
        if question["too_large"]:
            skipped.append({"unit_ref": unit_ref, "target_corner": corner, "reason": "the question was too large to ask"})
        else:
            questions[question["question_id"]] = question
    answers = ctx.ask(list(questions.values())) if questions else {}
    edges, problems, records, doc_judgements, follow_ups = [], list(skipped), [], [], {}
    for question_id in sorted(questions):
        question, final = questions[question_id], answers.get(question_id)
        unit_ref, corner = question["unit_ref"], question["target_corner"]
        shown = len(question["letters"])
        if final is None or final["outcome"] != "accepted":
            reason = (final or {}).get("outcome", "failed: no answer was obtained").split(": ", 1)[-1]
            problems.append({"unit_ref": unit_ref, "target_corner": corner, "reason": reason, "question_id": question_id})
            records.append({"unit_ref": unit_ref, "target_corner": corner, "search_pass": search_pass,
                            "note": "%d passages were shown to the AI. Its answer could not be used: %s." % (shown, reason)})
            continue
        answer = final["answer"]
        if answer.get("states_nothing_checkable"):
            doc_judgements.append({"unit_ref": unit_ref, "states_nothing_checkable": True, "question_id": question_id})
        accepted = [m for m in answer["matches"] if m["letter"] not in question["planted"]]
        reasons = {c["target_ref"]: c["reason"] for c in grouped[(unit_ref, corner)]}
        for match in accepted:
            target = question["letters"][match["letter"]]
            how = core.HOW_AI.format(confidence=match["confidence"])
            provenance = core.to_plain(ctx.provenance)
            provenance.update(prompt_hash=question_id, response_hash=final["response_hash"])
            edges.append(dict(core.to_plain(core.Edge(unit_ref, target, "corresponds", how, ctx.provenance,
                         relation=core.RELATION_WORDING[match["relation"]], confidence=match["confidence"],
                         evidence={"how_text": "%s. %s" % (how, reasons.get(target, "")), "question_id": question_id,
                                   "quote_from_passage": match["quote_from_passage"], "quote_from_unit": match["quote_from_unit"],
                                   "search_pass": search_pass})), provenance=provenance))
            linking = core.RELATION_WORDING[match["relation"]] in core.LINKING_RELATIONS
            if linking and match["relation"] not in ("deviates from", "inconsistent with") and needs_second_opinion(settings, world["targets"][corner][target]):
                follow_ups.setdefault((unit_ref, corner), []).append(target)
        linked = [m for m in accepted if core.RELATION_WORDING[m["relation"]] in core.LINKING_RELATIONS]
        note = "" if linked else "%d passages were shown to the AI. None accepted: %s" % (
            shown, "the AI's reason was \u201c%s\u201d" % shown_ai_text(answer.get("none_reason")) if answer.get("none_reason")
            else "the AI only saw passages on the same topic")
        records.append({"unit_ref": unit_ref, "target_corner": corner, "search_pass": search_pass, "note": note})
    opinions = second_opinions(ctx, follow_ups, sources, world) if follow_ups else []
    ledger = ledger_records(ctx.read("graph_ledger"), edges)
    return core.StepResult({"graph_ledger": ledger, "judgement_problems": problems, "search_records": records,
                              "doc_judgements": doc_judgements, "second_opinions": opinions},
                             {"questions": len(questions), "links recorded": len(edges), "answers not usable": len(problems) - len(skipped),
                              "questions too large to ask": len(skipped)}, [])

def second_opinions(ctx, follow_ups, sources, world):
    """The oppositely framed question ("identify any difference ...") for links no check can
    back. A named difference with verbatim quotations is recorded; account-coverage turns
    it into a flagged item of the category "AI answers disagree"."""
    questions = {}
    for (unit_ref, corner), targets in sorted(follow_ups.items()):
        source, pool = sources[unit_ref], world["targets"][corner]
        text = core.cut_text(source["text"], int(ctx.settings["max_unit_chars"]))
        passages = [(ref, passage_label(pool[ref]), passage_text(pool[ref], ctx.settings)) for ref in sorted(set(targets))]
        question = assemble_question("second-opinion", unit_ref, [("UNIT (%s)" % passage_label(source), text)], passages,
                                     core.load_prompt("second-opinion"), ctx.settings, more={"target_corner": corner})
        if not question["too_large"]:
            questions[question["question_id"]] = question
    answers, opinions = ctx.ask(list(questions.values())), []
    for question_id in sorted(questions):
        final, question = answers.get(question_id), questions[question_id]
        if final and final["outcome"] == "accepted":
            for difference in final["answer"]["differences"]:
                opinions.append({"unit_ref": question["unit_ref"], "target_ref": question["letters"][difference["letter"]],
                                 "quote_from_passage": difference["quote_from_passage"], "quote_from_unit": difference["quote_from_unit"],
                                 "what_differs": shown_ai_text(difference.get("what_differs")), "question_id": question_id})
    return opinions


# ================================================================================================
# ---------------------------------------------------------------- the checks, the statuses and the flagged items
SKILL_VERSIONS = {"check-mathematics": "0.0.1", "check-values": "0.0.1", "check-rules": "0.0.1",
                  "check-package-docs": "0.0.1", "account-coverage": "0.0.1"}

class EngineFault(Exception):
    """The run contradicts itself (the coverage identity does not hold). This is a defect in
    the tool, not in the model under review, and the message says which part failed."""

# ---------------------------------------------------------------- the one rule for comparing values
VALUE_RULE_TEXT = (
    "Values are compared under one rule, and there is no hidden tolerance. Percent, basis points and scientific "
    "notation are first converted to plain decimals. A value agrees at stated precision when rounding it to the "
    "number of decimals of the methodology's value gives the methodology's value, under round-half-even or "
    "round-half-up. When a documentation value is written with fewer decimals than the methodology's value, the "
    "comparison is made at that coarser precision and reported as such. Otherwise the values differ, and both are "
    "shown with their citations.")

def compare_values(stated, other, exact=False):
    """The value rule of plan 2.8. `stated` is the methodology's number and `other` the value
    compared with it, both as {"value", "decimals", ...} from shared.parse_number. `exact`
    marks a stored or coded value, which has no written precision of its own.
    Returns (outcome, note): agrees | agrees at stated precision | agrees at the coarser
    precision | differs; the note names the rounding convention when the two differ."""
    m, v = Decimal(stated["value"]), Decimal(other["value"])
    if m == v:
        return "agrees", ""
    if not exact and other["decimals"] < stated["decimals"]:
        step = Decimal(1).scaleb(-int(other["decimals"]))
        if v in (m.quantize(step, rounding=ROUND_HALF_EVEN), m.quantize(step, rounding=ROUND_HALF_UP)):
            return "agrees at the coarser precision", ""
        return "differs", ""
    step = Decimal(1).scaleb(-int(stated["decimals"]))
    even, up = v.quantize(step, rounding=ROUND_HALF_EVEN), v.quantize(step, rounding=ROUND_HALF_UP)
    if m in (even, up):
        return "agrees at stated precision", "" if even == up else "under round-half-even" if m == even else "under round-half-up"
    return "differs", ""

def number_of(text):
    """The single number a table cell or a default holds, or None. A cell that writes one value
    twice - "0.15 (15%)" - holds one number, not none."""
    found = core.find_numbers(str(text))
    if found and all(Decimal(n["value"]) == Decimal(found[0]["value"]) for n in found):
        return found[0]
    return core.parse_number(str(text).strip()) if not found else None

# ---------------------------------------------------------------- expression tree -> SymPy, node by node
class NotEvaluable(Exception):
    """The formula uses an operation the tool cannot evaluate. The message names it."""

def to_sympy(expr):
    """One explicit table from the tool's tree to SymPy objects. Numbers become exact rationals.
    No string is ever handed to SymPy's parser. Enforces: R7"""
    import sympy
    args = [to_sympy(arg) for arg in expr.args] if expr.op not in ("num", "sym") else []
    if expr.op == "num":
        return sympy.Rational(str(Decimal(expr.value)))
    if expr.op == "sym":
        return sympy.Symbol(expr.name, real=True)
    simple = {"add": lambda a, b: a + b, "sub": lambda a, b: a - b, "mul": lambda a, b: a * b,
              "div": lambda a, b: a / b, "pow": lambda a, b: a ** b, "neg": lambda a: -a}
    if expr.op in simple:
        return simple[expr.op](*args)
    if expr.op == "cmp":
        relation = {"<": sympy.Lt, "<=": sympy.Le, ">": sympy.Gt, ">=": sympy.Ge, "==": sympy.Eq, "!=": sympy.Ne}[expr.name]
        return relation(*args)
    if expr.op == "piecewise" or (expr.op == "call" and expr.name == "piecewise"):
        return sympy.Piecewise((args[1], args[0]), (args[2], True))
    calls = {"exp": sympy.exp, "log": sympy.log, "sqrt": sympy.sqrt, "abs": sympy.Abs, "max": sympy.Max, "min": sympy.Min,
             "floor": sympy.floor, "ceiling": sympy.ceiling, "log1p": lambda x: sympy.log(1 + x), "expm1": lambda x: sympy.exp(x) - 1,
             "normal_cdf": lambda x: (1 + sympy.erf(x / sympy.sqrt(2))) / 2,
             "normal_inverse": lambda x: sympy.sqrt(2) * sympy.erfinv(2 * x - 1)}
    if expr.op == "call" and expr.name in calls:
        return calls[expr.name](*args)
    raise NotEvaluable("'%s'" % (expr.name or expr.op))

def symbolic_step(code_tree, stated_tree, settings):
    """Bounded symbolic work: expand the difference always; simplify fully only for small
    trees; nothing for very large ones. "agrees" only when the difference is exactly zero;
    a difference that does not reduce proves nothing and the numeric step decides."""
    size = len(list(core.expr_walk(code_tree))) + len(list(core.expr_walk(stated_tree)))
    if size > 400:
        return "skipped (size)"
    try:
        import sympy
        difference = sympy.expand(to_sympy(code_tree) - to_sympy(stated_tree))
        if difference != 0 and size <= 60:
            difference = sympy.simplify(difference)
        return "agrees" if difference == 0 else "not shown equal"
    except Exception:                                   # SymPy could not handle it: the numeric step decides
        return "not shown equal"

# ---------------------------------------------------------------- the tool's own numeric evaluator
def evaluate(expr, values):
    """The value of a tree at one point. Returns None where the point is outside the domain
    of a function used (such a point is not valid and is not counted). Enforces: R7"""
    from scipy import special
    op = expr.op
    if op == "num":
        return float(Decimal(expr.value))
    if op == "sym":
        return values[expr.name]
    args = [evaluate(arg, values) for arg in expr.args]
    if any(arg is None for arg in args):
        return None
    try:
        if op in ("add", "sub", "mul", "neg"):
            return {"add": lambda: args[0] + args[1], "sub": lambda: args[0] - args[1], "mul": lambda: args[0] * args[1], "neg": lambda: -args[0]}[op]()
        if op == "div":
            return args[0] / args[1] if args[1] != 0 else None
        if op == "pow":
            result = args[0] ** args[1] if not (args[0] == 0 and args[1] < 0) else None
            return result if isinstance(result, float) or isinstance(result, int) else None
        if op == "cmp":
            return {"<": args[0] < args[1], "<=": args[0] <= args[1], ">": args[0] > args[1], ">=": args[0] >= args[1],
                    "==": args[0] == args[1], "!=": args[0] != args[1]}[expr.name]
        if op == "piecewise" or (op == "call" and expr.name == "piecewise"):
            return args[1] if args[0] else args[2]
        name, x = expr.name, args[0]
        if name in ("log", "log1p"):
            return (math.log(x) if x > 0 else None) if name == "log" else (math.log1p(x) if x > -1 else None)
        if name == "sqrt":
            return math.sqrt(x) if x >= 0 else None
        if name == "normal_inverse":
            return float(special.ndtri(x)) if 0 < x < 1 else None
        plain = {"exp": math.exp, "expm1": math.expm1, "abs": abs, "floor": math.floor, "ceiling": math.ceil,
                 "normal_cdf": lambda value: float(special.ndtr(value))}
        if name in plain:
            return float(plain[name](x))
        if name in ("max", "min"):
            return max(args) if name == "max" else min(args)
        if name == "round":
            return float(round(x, int(args[1]) if len(args) > 1 else 0))
    except (OverflowError, ValueError, ZeroDivisionError):
        return None
    raise NotEvaluable("'%s'" % (expr.name or op))

def unit_interval_symbols(tree):
    """Symbols that sit under a square root, a logarithm or the inverse normal function are
    drawn from (0, 1), which keeps most sample points inside the domain."""
    narrow = set()
    for node in core.expr_walk(tree):
        if node.op == "call" and node.name in ("sqrt", "log", "normal_inverse", "log1p"):
            narrow.update(core.expr_symbols(node))
    return narrow

def sample_points(trees, symbols, settings, data_values):
    """The sample points: a fixed, recorded seed; ranges (0, 1), (1, 10) and (-10, 10) in turn;
    symbols with a narrow domain always from (0, 1); symbols supplied by a stored table take
    its values row by row; and, on purpose, points at, just below and just above every
    constant that appears in a comparison, a maximum or a minimum. Enforces: R5"""
    generator = random.Random(int(settings["numeric_seed"]))
    narrow = set().union(*[unit_interval_symbols(tree) for tree in trees])
    ranges = ((0.0, 1.0), (1.0, 10.0), (-10.0, 10.0))
    points = []
    for number in range(int(settings["numeric_points"])):
        low, high = ranges[number % 3]
        point = {}
        for symbol in symbols:
            if symbol in data_values:
                point[symbol] = data_values[symbol][number % len(data_values[symbol])]
            else:
                point[symbol] = generator.uniform(0.0, 1.0) if symbol in narrow else generator.uniform(low, high)
        points.append(point)
    thresholds = set()
    for tree in trees:
        for node in core.expr_walk(tree):
            if node.op in ("cmp", "piecewise") or (node.op == "call" and node.name in ("max", "min", "piecewise")):
                thresholds.update(float(Decimal(arg.value)) for arg in node.args if arg.op == "num")
    for threshold in sorted(thresholds):
        for symbol in symbols:
            for value in (threshold, threshold * 0.999 - 1e-9, threshold * 1.001 + 1e-9):
                point = dict(points[len(points) % max(1, int(settings["numeric_points"]))])
                point[symbol] = value
                points.append(point)
    return points + crossing_points(trees, symbols, points[:3], narrow)

def crossing_points(trees, symbols, bases, narrow):
    """Where a threshold sits on a scaled quantity, as in min(0.2, 0.01 * n), the two sides meet
    far outside the generic ranges (here at n = 20). For every maximum, minimum and comparison,
    and for every symbol in it, a wide grid is scanned for a change of sign of the difference
    of the two sides; the crossing is then narrowed by bisection, and points at it and just
    below and above it are added. Without them a changed cap could pass as agreeing."""
    wide = [-1e6, -1e4, -1e3, -100.0, -10.0, -1.0, -0.1, -0.01, 0.0, 0.01, 0.1, 1.0, 10.0, 100.0, 1e3, 1e4, 1e6]
    unit = [0.001, 0.01, 0.1, 0.25, 0.5, 0.75, 0.9, 0.99, 0.999]
    pairs = []
    for tree in trees:
        for node in core.expr_walk(tree):
            if node.op == "cmp" or (node.op == "call" and node.name in ("max", "min")):
                pairs += [(node.args[i], node.args[j]) for i in range(len(node.args)) for j in range(i + 1, len(node.args))]
    found = []
    for left, right in pairs:
        for symbol in sorted((set(core.expr_symbols(left)) | set(core.expr_symbols(right))) & set(symbols)):
            for base in bases:
                def gap(value):
                    try:
                        a, b = evaluate(left, dict(base, **{symbol: value})), evaluate(right, dict(base, **{symbol: value}))
                    except NotEvaluable:
                        return None
                    return None if a is None or b is None or isinstance(a, bool) or isinstance(b, bool) else a - b
                grid = unit if symbol in narrow else wide
                for low, high in zip(grid, grid[1:]):
                    at_low, at_high = gap(low), gap(high)
                    if at_low is None or at_high is None or at_low == 0 or (at_low > 0) == (at_high > 0):
                        continue
                    for _ in range(60):                  # bisection: deterministic, and bounded
                        middle = (low + high) / 2.0
                        at_middle = gap(middle)
                        if at_middle is None:
                            break
                        low, high, at_low = (middle, high, at_middle) if (at_middle > 0) == (at_low > 0) else (low, middle, at_low)
                    crossing = (low + high) / 2.0
                    found += [dict(base, **{symbol: value}) for value in (crossing, crossing - max(1e-6, abs(crossing) * 1e-3), crossing + max(1e-6, abs(crossing) * 1e-3))]
    return found

def numeric_step(code_tree, stated_tree, symbols, settings, data_values):
    """Evaluate both trees at every sample point. Returns (result text, valid points,
    counterexample or None). Any disagreement beyond the relative tolerance is a difference,
    recorded with the inputs and both results. Enforces: R4"""
    valid, tolerance = 0, float(settings["relative_tolerance"])
    for point in sample_points([code_tree, stated_tree], symbols, settings, data_values):
        ours, theirs = evaluate(code_tree, point), evaluate(stated_tree, point)
        if ours is None or theirs is None or isinstance(ours, bool) or isinstance(theirs, bool):
            continue
        if math.isnan(ours) or math.isnan(theirs) or math.isinf(ours) or math.isinf(theirs):
            continue
        valid += 1
        if abs(ours - theirs) > tolerance * max(1.0, abs(ours), abs(theirs)):
            return "differs", valid, {"inputs": {s: point[s] for s in symbols}, "code": ours, "methodology": theirs}
    if valid < int(settings["min_valid_points"]):
        return "not run", valid, None
    return "agrees (%d points)" % valid, valid, None

# ---------------------------------------------------------------- symbol alignment (fixed before comparing)
def align(code_symbols, stated_symbols, bridge_by_term):
    """Align code symbols with the symbols of the stated formula, by code only: identical
    normalised names; the same name apart from capital letters when that is unambiguous;
    then the bridge vocabulary (both described by the same words). One-to-one. Returns
    (pairs [(code, stated, matched by)], code symbols left, stated symbols left)."""
    pairs, left_code, left_stated = [], list(code_symbols), list(stated_symbols)
    def take(code, stated, how):
        pairs.append((code, stated, how))
        left_code.remove(code)
        left_stated.remove(stated)
    for code in list(left_code):
        if code in left_stated:
            take(code, code, "name")
    for code in list(left_code):
        same = [s for s in left_stated if s.lower() == code.lower()]
        rivals = [c for c in left_code if c.lower() == code.lower()]
        if len(same) == 1 and len(rivals) == 1:
            take(code, same[0], "name, apart from capital letters")
    def described(symbol):
        return {word for entry in bridge_by_term.get(symbol, []) for word in entry["words"]}
    for code in list(left_code):
        mine = described(code) | set(split_words(code.replace("$", " "), set()))
        scores = []
        for stated in left_stated:
            theirs = described(stated)
            if mine and theirs:
                scores.append((len(mine & theirs) / len(mine | theirs), stated))
        scores.sort(key=lambda item: (-item[0], item[1]))
        if scores and scores[0][0] >= 0.5 and (len(scores) == 1 or scores[1][0] < scores[0][0]):
            take(code, scores[0][1], "vocabulary")
    return pairs, left_code, left_stated

def rename(tree, mapping, substitute):
    """The code tree written in the stated formula's symbols; defaults put in as numbers."""
    if tree.op == "sym":
        if tree.name in substitute:
            return core.Expr("num", value=substitute[tree.name])
        return core.Expr("sym", name=mapping.get(tree.name, tree.name))
    return core.Expr(tree.op, name=tree.name, value=tree.value, args=tuple(rename(arg, mapping, substitute) for arg in tree.args))

def right_side(tree):
    """The right-hand side of an equation; an expression that is no equation is returned whole."""
    return tree.args[1] if tree.op == "eq" else tree

def prepare_alignment(code_tree, stated_tree, defaults, bridge_by_term):
    """Everything about one comparison that is settled before any number is computed: the
    aligned pairs, the defaults put in (only where the stated formula shows that very
    number), and what is still left for the align-symbols question."""
    code_side, stated_side = right_side(code_tree), right_side(stated_tree)
    pairs, left_code, left_stated = align(sorted(core.expr_symbols(code_side)), sorted(core.expr_symbols(stated_side)), bridge_by_term)
    constants = {node.value for node in core.expr_walk(stated_side) if node.op == "num"}
    substitute = {}
    for symbol in list(left_code):
        number = number_of(defaults.get(symbol, "")) if defaults.get(symbol) else None
        if number and number["value"] in constants:
            substitute[symbol] = number["value"]
            left_code.remove(symbol)
    return {"code": code_side, "stated": stated_side, "pairs": pairs, "left_code": left_code, "left_stated": left_stated,
            "substitute": substitute}

def compare_formulas(prepared, settings, data_values):
    """Symbolic step, then numeric step, on an alignment that is already fixed. Returns the
    outcome fields of a MathCheck. Enforces: R3, R4"""
    if prepared["left_code"] or prepared["left_stated"]:
        return {"outcome": core.CHECK_UNDECIDED, "undecided_reason": core.UNDECIDED_REASONS[2], "symbolic": "not run",
                "numeric": "not run", "counterexample": None, "points_valid": 0}
    mapping = {code: stated for code, stated, _ in prepared["pairs"]}
    code_tree = rename(prepared["code"], mapping, prepared["substitute"])
    symbols = sorted(set(core.expr_symbols(prepared["stated"])) | set(core.expr_symbols(code_tree)))
    values = {mapping.get(symbol, symbol): rows for symbol, rows in data_values.items()}
    result = {"symbolic": symbolic_step(code_tree, prepared["stated"], settings), "numeric": "not run",
              "counterexample": None, "points_valid": 0, "undecided_reason": None}
    if result["symbolic"] == "agrees":
        return dict(result, outcome="agrees")
    try:
        numeric, valid, counterexample = numeric_step(code_tree, prepared["stated"], symbols, settings, values)
    except NotEvaluable:
        return dict(result, outcome=core.CHECK_UNDECIDED, undecided_reason=core.UNDECIDED_REASONS[4])
    result.update(numeric=numeric, points_valid=valid, counterexample=counterexample)
    if numeric == "differs":
        return dict(result, outcome="differs")
    if numeric == "not run":
        return dict(result, outcome=core.CHECK_UNDECIDED, undecided_reason=core.UNDECIDED_REASONS[5])
    return dict(result, outcome="agrees")

# ---------------------------------------------------------------- what every check step reads
CODE_RELATIONS = ("Implements", "Partly implements", "Differs from")
PROSE_FORMULA_PHRASES = ("product of", "sum of", "multiplied by", "divided by", "ratio of", "square root of", " times the ")

def load_world(ctx, everything=False):   # everything: kept for callers; every unit is read either way
    """Units, chunks, the graph and the latest link of every linked pair, read once per step."""
    units, canon, doc = (ctx.read(kind) for kind in ("model_units", "chunks_canon", "chunks_doc"))
    ledger = ctx.read("graph_ledger")
    world = {"units": units, "canon": canon, "doc": doc, "ledger": ledger, "graph": load_graph(ledger),
             "by_ref": {record["ref"]: record for record in units + canon + doc}, "links": {}, "children": {},
             "tables": {table["unit_ref"]: table for table in ctx.read("parameter_tables")}, "bridge": {}}
    for record in ledger:
        if record["record_type"] == "edge" and record["kind"] == "corresponds":
            world["links"][(record["source"], record["target"])] = record      # a later edge has the last word
    for unit in units:
        if unit.get("parent_ref"):
            world["children"].setdefault(unit["parent_ref"], []).append(unit)
    for entry in ctx.read("bridge_vocabulary"):
        world["bridge"].setdefault(entry["term"], []).append(entry)
    return world

def linked(world, ref, prefix, relations=core.LINKING_RELATIONS):
    """References in one corner ("C-", "D-", "M-") that `ref` is linked to, either way round."""
    found = [t for (s, t), e in world["links"].items() if s == ref and t.startswith(prefix) and e["relation"] in relations]
    found += [s for (s, t), e in world["links"].items() if t == ref and s.startswith(prefix) and e["relation"] in relations]
    return sorted(set(found))

def check_edge(ctx, world, source, target, how, sentence, relation=None):
    """A check never rewrites a link: it appends a new edge on the same pair, which then has
    the last word in the workbook and in the status rules."""
    previous = world["links"].get((source, target), {})
    text = (previous.get("evidence", {}).get("how_text", "") + "\n" + sentence).strip()
    return core.Edge(source, target, "corresponds", how, ctx.provenance, relation=relation or previous.get("relation", ""),
                       confidence=previous.get("confidence"), evidence={"how_text": text, "check": True})

def number_text(value):
    """A computed number for a cell: at most six significant digits, never in Python's own notation."""
    return core.plain_number(float("%.6g" % value)) if isinstance(value, float) else str(value)

# ---------------------------------------------------------------- step 11: check-mathematics
def stated_formula(chunk, prose_answers):
    """The formula a passage states, as (tree, source, reason it cannot be used). A passage
    without any formula gives (None, "", "") and needs no mathematical check."""
    equation = chunk.get("equation")
    if equation:
        if equation["readable"]:
            source = "inline notation" if equation["source_form"] == "inline" else "structured markup"
            return core.expr_from_dict(equation["expression"]), source, ""
        image = equation["source_form"] == "image" or bool(equation.get("image_sha256"))
        return None, "", core.UNDECIDED_REASONS[0] if image else core.UNDECIDED_REASONS[1]
    answer = prose_answers.get(chunk["ref"])
    if answer is not None:
        return (answer, "read by AI from prose", "") if answer != "unusable" else (None, "", core.UNDECIDED_REASONS[1])
    return None, "", ""

def code_forms(unit, world):
    """The formulas a unit offers for comparison, as (tree, defaults, reference it came from).
    A function offers its composed form and the forms of the statements inside it; a
    statement offers its own expression and its composed form."""
    forms = []
    members = [unit] + (world["children"].get(unit["ref"], []) if unit["kind"] == core.KIND_FUNCTION else [])
    owner = unit if unit["kind"] == core.KIND_FUNCTION else world["by_ref"].get(unit.get("parent_ref") or "", unit)
    defaults = {core.normalise_symbol(name): default for name, default in ((owner.get("code") or {}).get("formals") or ()) if default}
    for member in members:
        code = member.get("code") or {}
        for form in ("expression", "composed"):
            if code.get(form) and code[form] not in [f[3] for f in forms]:
                forms.append((core.expr_from_dict(code[form]), defaults, member["ref"], code[form]))
    return [(tree, defaults, ref) for tree, defaults, ref, _ in forms]

def stored_values(world):
    """Numeric columns of stored tables, as {"object$column": [values]}, for the numeric step."""
    values = {}
    for table in world["tables"].values():
        for position, column in enumerate(table["header"]):
            cells = [number_of(row[position]) for row in table["rows"]]
            if cells and all(cells):
                values["%s$%s" % (table["object_name"], column)] = [float(Decimal(cell["value"])) for cell in cells]
    return values

def math_sentence(target_ref, result, pairs, substitute, ours="the code", theirs="the methodology"):
    """The words of one comparison, as shown in the Math check cell."""
    shown = ", ".join("%s = %s" % (code, stated) for code, stated, _ in pairs if code != stated)
    shown += (", " if shown and substitute else "") + ", ".join("%s = %s (its default)" % item for item in sorted(substitute.items()))
    with_text = " (with %s)" % shown if shown else ""
    if result["outcome"] == "agrees":
        how = core.HOW_SYMBOLIC if result["symbolic"] == "agrees" else core.HOW_NUMERIC_AGREES.format(n=result["points_valid"])
        return "%s: %s%s" % (target_ref, how, with_text)
    if result["outcome"] == "differs":
        example = result["counterexample"]
        inputs = ", ".join("%s = %s" % (name, number_text(value)) for name, value in sorted(example["inputs"].items()))
        return "%s: %s. With %s: %s gives %s, %s gives %s%s" % (target_ref, core.HOW_NUMERIC_DIFFERS, inputs or "no inputs", ours,
                                                                number_text(example["code"]), theirs, number_text(example["methodology"]), with_text)
    return "%s: %s: %s" % (target_ref, core.CHECK_UNDECIDED, result["undecided_reason"])

def combine(results):
    """Several forms of one unit against one formula: any form that agrees settles it; else
    any form that differs (with its counterexample); else it could not be decided."""
    for wanted in ("agrees", "differs"):
        for result in results:
            if result["outcome"] == wanted:
                return result
    return results[0]

def check_mathematics(ctx):
    """Step 11, check-mathematics. Pairs: every function or formula statement linked to
    a passage that states a formula; every roxygen \\deqn formula against the function it
    documents; every documentation equation against the methodology equation it is linked
    to. Alignment is settled first (by code, then by the validated align-symbols answer)
    and only then are the two sides compared. Enforces: R3, R4"""
    world, settings = load_world(ctx), ctx.settings
    notation, values = reading.load_notation(), stored_values(world)
    pairs = [(world["by_ref"][s], world["by_ref"][t]) for (s, t), e in sorted(world["links"].items())
             if s.startswith("M-") and t.startswith("C-") and e["relation"] in CODE_RELATIONS
             and world["by_ref"][s]["kind"] in (core.KIND_FUNCTION, core.KIND_FORMULA)]
    prose = {}
    for _, chunk in pairs:
        if not chunk.get("equation") and chunk["kind"] == "Paragraph" and any(p in chunk["text"].lower() for p in PROSE_FORMULA_PHRASES):
            prose[chunk["ref"]] = narrow_question("read-formula-from-prose", chunk["ref"], [("PARAGRAPH", chunk["text"][:3000])],
                                                                settings, more={"notation": notation})
    answers = ctx.ask([q for q in prose.values() if not q["too_large"]]) if prose else {}
    prose_answers = {}
    for ref, question in prose.items():
        final = answers.get(question["question_id"])
        if final and final["outcome"] == "accepted":
            if final["answer"]["formula"].strip():
                prose_answers[ref] = reading.parse_formula(final["answer"]["formula"], notation)
        else:
            prose_answers[ref] = "unusable"
    jobs = []                                            # (unit, target, stated tree, source, [prepared forms], reason)
    for unit, chunk in pairs:
        tree, source, reason = stated_formula(chunk, prose_answers)
        if tree is None and not reason:
            continue
        forms = code_forms(unit, world) if tree is not None else []
        prepared = [dict(prepare_alignment(code, tree, defaults, world["bridge"]), through=ref) for code, defaults, ref in forms]
        jobs.append((unit, chunk, tree, source, prepared, reason or (core.UNDECIDED_REASONS[3] if not prepared else "")))
    for unit in world["units"]:                          # roxygen formulas against the function they document
        target = world["by_ref"].get((unit.get("roxygen") or {}).get("documents_ref") or "")
        for formula in (unit.get("roxygen") or {}).get("formulas", ()):
            if target and formula["readable"] and target["kind"] == core.KIND_FUNCTION:
                tree = core.expr_from_dict(formula["expression"])
                prepared = [dict(prepare_alignment(code, tree, defaults, world["bridge"]), through=ref) for code, defaults, ref in code_forms(target, world)]
                jobs.append((target, unit, tree, "roxygen", prepared, "" if prepared else core.UNDECIDED_REASONS[3]))
    for (s, t), edge in sorted(world["links"].items()):  # documentation equations against methodology equations
        left, right = world["by_ref"][s], world["by_ref"][t]
        if s.startswith("D-") and t.startswith("C-") and (left.get("equation") or {}).get("readable") and right.get("equation"):
            tree, source, reason = stated_formula(right, {})
            prepared = [dict(prepare_alignment(core.expr_from_dict(left["equation"]["expression"]), tree, {}, world["bridge"]), through=s)] if tree else []
            jobs.append((left, right, tree, source, prepared, reason))
    questions = {}
    for unit, chunk, tree, source, prepared, reason in jobs:
        for form in prepared:
            if form["left_code"] and form["left_stated"]:
                question = narrow_question(
                    "align-symbols", unit["ref"], [("CODE SYMBOLS", ", ".join(form["left_code"])), ("EQUATION SYMBOLS", ", ".join(form["left_stated"]))],
                    settings, more={"code_symbols": list(form["left_code"]), "equation_symbols": list(form["left_stated"])})
                form["question_id"] = question["question_id"]
                questions[question["question_id"]] = question
    answers = ctx.ask(list(questions.values())) if questions else {}
    checks, by_pair = [], {}
    for unit, chunk, tree, source, prepared, reason in jobs:
        results = []
        for form in prepared:
            final = answers.get(form.get("question_id", ""))
            if final and final["outcome"] == "accepted":
                for pair in final["answer"]["alignment"]:
                    form["pairs"].append((pair["code"], pair["equation"], "AI (validated)"))
                    form["left_code"].remove(pair["code"])
                    form["left_stated"].remove(pair["equation"])
            results.append(dict(compare_formulas(form, settings, values), form=form))
        if results:
            best = combine(results)
            form = best.pop("form")
        else:
            best = {"outcome": core.CHECK_UNDECIDED, "undecided_reason": reason, "symbolic": "not run", "numeric": "not run",
                    "counterexample": None, "points_valid": 0}
            form = {"code": None, "pairs": [], "substitute": {}, "through": unit["ref"]}
        record = dict(best, check_id="", unit_ref=unit["ref"], target_ref=chunk["ref"], formula_source=source,
                      code_formula=core.expr_to_text(form["code"]) if form["code"] is not None else "",
                      methodology_formula=core.expr_to_text(right_side(tree)) if tree is not None else "",
                      alignment=[list(pair) for pair in form["pairs"]], seed=int(settings["numeric_seed"]), through=form["through"])
        ours = "the documentation" if unit.get("heading_chain") is not None else "the code"
        theirs = "the formula in the roxygen block" if source == "roxygen" else "the methodology"
        record["sentence"] = math_sentence(chunk["ref"], record, form["pairs"], form["substitute"], ours, theirs)
        by_pair[(unit["ref"], chunk["ref"])] = record
        checks.append(record)
    for record in checks:                                # a statement takes the function-level result for the same passage
        unit = world["by_ref"][record["unit_ref"]]
        parent = by_pair.get((unit.get("parent_ref"), record["target_ref"]))
        if unit["kind"] == core.KIND_FORMULA and record["outcome"] != "agrees" and parent:
            record.update(outcome=parent["outcome"], undecided_reason=parent["undecided_reason"], counterexample=parent["counterexample"],
                          inherited_from=parent["unit_ref"], sentence="%s (taken from the check of the whole function %s)" % (parent["sentence"], parent["unit_ref"]))
    edges = []
    for number, record in enumerate(checks, start=1):
        record["check_id"] = "MC-%04d" % number
        if record["outcome"] != core.CHECK_UNDECIDED and record["formula_source"] != "roxygen":
            how = core.HOW_NUMERIC_DIFFERS if record["outcome"] == "differs" else core.HOW_SYMBOLIC if record["symbolic"] == "agrees" \
                else core.HOW_NUMERIC_AGREES.format(n=record["points_valid"])
            edges.append(check_edge(ctx, world, record["unit_ref"], record["target_ref"], how, record["sentence"],
                                    relation="Differs from" if record["outcome"] == "differs" else None))
    counts = {"pairs checked": len(checks)}
    for outcome in ("agrees", "differs", core.CHECK_UNDECIDED):
        counts[outcome] = sum(1 for record in checks if record["outcome"] == outcome)
    return core.StepResult({"math_checks": checks, "graph_ledger": ledger_records(world["ledger"], edges)}, counts, [])

# ---------------------------------------------------------------- tables: reconciliation cell by cell
def quote(text, citation=""):
    """Input text is always shown visibly quoted, with its citation (same form as in verifier5)."""
    inner = core.normalise_text(text or "").replace("\u201c", '"').replace("\u201d", '"')
    return "\u201c%s\u201d%s" % (inner, " (%s)" % citation if citation else "")

def plain_key(text):
    """A row key or header as compared: lower case, single spaces, underscores as spaces."""
    return re.sub(r"[\s_]+", " ", str(text)).strip().lower()

def header_words(header, stop):
    """The index words of a column header, without its unit in brackets."""
    return frozenset(split_words(re.sub(r"\(.*?\)|%", " ", header), stop))

def map_columns(ours, theirs, stop):
    """Columns matched by header: equal words after normalisation. Returns [(ours, theirs)]."""
    pairs, used = [], set()
    for column in ours["header"]:
        same = [other for other in theirs["header"] if other not in used and header_words(other, stop)
                and header_words(other, stop) == header_words(column, stop)]
        if len(same) == 1:
            pairs.append((column, same[0]))
            used.add(same[0])
    return pairs

def as_grid(table):
    """The one layout rule that is supported: a table in long form (two key columns and one
    value column) is turned into a grid, so that it can be laid over a printed grid."""
    kinds = table.get("column_types") or []
    if len(table["header"]) != 3 or len(kinds) != 3 or kinds[2] == "text" or "text" not in kinds[:2]:
        return None
    columns = list(dict.fromkeys(row[1] for row in table["rows"]))
    rows = {}
    for row in table["rows"]:
        rows.setdefault(row[0], {})[row[1]] = row[2]
    return {"header": [table["header"][0]] + columns, "row_key": [table["header"][0]],
            "rows": [[key] + [rows[key].get(column, "") for column in columns] for key in rows]}

def reconcile(ours, theirs, stop, exact, ai_map=None):
    """Compare two tables cell by cell under the value rule. `theirs` holds the values as
    stated (the methodology, or the documentation when it is compared with the package);
    `ours` holds the values compared with them. Columns are matched by header, else by
    values, else by the validated AI answer `ai_map`; rows by normalised key. Returns the
    fields of a ValueCheck."""
    pairs = [(a, b, "by header") for a, b in map_columns(ours, theirs, stop)]
    if len(pairs) < 2 and as_grid(ours) and len(map_columns(as_grid(ours), theirs, stop)) >= 2:
        ours = as_grid(ours)
        pairs = [(a, b, "by header (long form laid over the grid)") for a, b in map_columns(ours, theirs, stop)]
    if ai_map and len(pairs) < 2:
        pairs = [(a, b, "by AI (validated)") for a, b in ai_map["columns"]]
    key_name = (ours.get("row_key") or [""])[0]
    key_pair = next((p for p in pairs if p[0] == key_name), None) or (("(row number)", "(row number)", "") if not pairs else pairs[0])
    if ai_map and ai_map.get("key"):
        key_pair = (ai_map["key"][0], ai_map["key"][1], "by AI (validated)")
    position = lambda table, name: table["header"].index(name) if name in table["header"] else None
    def keyed(table, name):
        index = position(table, name)
        return {plain_key(row[index]) if index is not None else str(n): row for n, row in enumerate(table["rows"], start=1)}
    our_rows, their_rows = keyed(ours, key_pair[0]), keyed(theirs, key_pair[1])
    if len(pairs) < 2 and not (their_rows.keys() & our_rows.keys()):
        return {"column_map": [], "key_map": "", "cells": [], "only_in_package": [], "only_in_other": [], "outcome": "could not be compared"}
    if len(pairs) < 2:                                   # no usable headers: match columns by their values
        pairs = [key_pair] + columns_by_values(ours, theirs, our_rows, their_rows, key_pair, exact)
    cells = []
    for ours_name, theirs_name, _ in pairs:
        if (ours_name, theirs_name) == key_pair[:2]:
            continue
        percent = "%" in theirs_name and "%" not in ours_name
        for key in sorted(our_rows.keys() & their_rows.keys()):
            mine, stated = our_rows[key][position(ours, ours_name)], their_rows[key][position(theirs, theirs_name)]
            stated_number = core.parse_number(stated.strip().rstrip("%"), "%") if percent and number_of(stated) else number_of(stated)
            my_number = number_of(mine)
            if stated_number and my_number:
                outcome, note = compare_values(stated_number, my_number, exact)
            else:
                outcome, note = ("agrees" if plain_key(mine) == plain_key(stated) else "differs"), ""
            cells.append({"row": our_rows[key][position(ours, key_pair[0])] if position(ours, key_pair[0]) is not None else key,
                          "column": ours_name, "value": mine, "stated": stated, "outcome": (outcome + " " + note).strip()})
    only_ours = [our_rows[k][position(ours, key_pair[0]) or 0] for k in sorted(our_rows.keys() - their_rows.keys())]
    only_theirs = [their_rows[k][position(theirs, key_pair[1]) or 0] for k in sorted(their_rows.keys() - our_rows.keys())]
    differs = any(cell["outcome"].startswith("differs") for cell in cells) or only_ours or only_theirs
    return {"column_map": [list(p) for p in pairs], "key_map": "%s = %s" % key_pair[:2], "cells": cells, "only_in_package": only_ours,
            "only_in_other": only_theirs, "outcome": "could not be compared" if not cells else "differs" if differs else "agrees"}

def columns_by_values(ours, theirs, our_rows, their_rows, key_pair, exact):
    """Match columns whose values agree with exactly one column on the other side for most keys."""
    pairs, common = [], sorted(our_rows.keys() & their_rows.keys())
    for i, ours_name in enumerate(ours["header"]):
        if ours_name == key_pair[0]:
            continue
        scores = []
        for j, theirs_name in enumerate(theirs["header"]):
            agree = sum(1 for key in common if number_of(our_rows[key][i]) and number_of(their_rows[key][j]) and
                        compare_values(number_of(their_rows[key][j]), number_of(our_rows[key][i]), exact)[0] != "differs")
            scores.append((agree, theirs_name))
        good = [name for agree, name in scores if common and agree * 2 > len(common)]
        if len(good) == 1:
            pairs.append((ours_name, good[0], "by values"))
    return pairs

def table_sentences(check, ours_label, theirs_label, target_ref):
    """The words of one table comparison, as shown in the workbook."""
    lines = ["row %s, column %s: %s %s, %s %s (%s row %s): %s" % (c["row"], c["column"], ours_label, c["value"], theirs_label,
             c["stated"], target_ref, c["row"], c["outcome"]) for c in check["cells"] if not c["outcome"].startswith("agrees") or " " in c["outcome"][7:]]
    lines += ["row %s is in %s only" % (row, ours_label) for row in check["only_in_package"]]
    lines += ["row %s is in %s only (%s)" % (row, theirs_label, target_ref) for row in check["only_in_other"]]
    if check["outcome"] == "could not be compared":
        return ["The table could not be laid over %s: no columns or row keys could be matched." % target_ref]
    agreed = sum(1 for c in check["cells"] if c["outcome"].startswith("agrees"))
    how = ", ".join(sorted({p[2] for p in check["column_map"] if p[2]}))
    return ["%d of %d values agree with %s (columns matched %s; rows identified by %s)." % (
        agreed, len(check["cells"]), target_ref, how or "by position", check["key_map"])] + lines

# ---------------------------------------------------------------- step 12: check-values
def numbers_with_context(text, stop):
    """Every number of a text with the words around it (four before, three after). Two numbers
    are only ever compared when they are attached to the same words."""
    found = []
    for number in core.find_numbers(text or ""):
        before = split_words(text[max(0, number["position"] - 60):number["position"]], stop)[-4:]
        after = split_words(text[number["position"] + len(number["as_written"]):][:45], stop)[:3]
        found.append(dict(number, context=set(before + after)))
    return found

def prose_value_lines(text, stated_texts, stop, trivial, label):
    """Numbers in a documentation passage or a roxygen block against the numbers of the
    passages it is linked to. A number that is stated somewhere in the linked passages
    agrees; a number attached to the same words as a DIFFERENT stated number differs; a
    number without a counterpart is not compared and nothing is claimed about it."""
    stated = [dict(n, ref=ref) for ref, passage in stated_texts for n in numbers_with_context(passage, stop)]
    lines, differs, mine = [], False, numbers_with_context(text, stop)
    taken = [s for s in stated if any(compare_values(s, number)[0] != "differs" for number in mine)]   # stated numbers that found their equal
    for number in mine:
        if number["value"] in trivial:
            continue
        outcomes = [(compare_values(s, number), s) for s in stated]
        agreeing = [(o, s) for o, s in outcomes if o[0] != "differs"]
        if agreeing:
            (outcome, note), s = agreeing[0]
            lines.append("%s %s, methodology %s (%s): %s" % (label, number["as_written"], s["as_written"], s["ref"], (outcome + " " + note).strip()))
            continue
        close = sorted(((len(number["context"] & s["context"]), s["ref"], s) for s in stated if s not in taken and len(number["context"] & s["context"]) >= 2),
                       key=lambda item: (-item[0], item[1]))
        if close:
            differs = True
            lines.append("%s %s, methodology %s (%s): differs" % (label, number["as_written"], close[0][2]["as_written"], close[0][1]))
        else:
            lines.append("%s %s: not compared (no counterpart in the linked passages)" % (label, number["as_written"]))
    return lines, differs

def check_values(ctx):
    """Step 12, check-values: parameter tables against the tables they are linked to
    (a table that the judge did not link is still matched by its shape); documentation
    tables against methodology tables; numbers written in linked code; numbers in
    documentation passages and roxygen text. All under compare_values. Enforces: R3"""
    world, settings = load_world(ctx), ctx.settings
    stop = load_word_lists()["stop"]
    trivial, checks, edges = set(settings["trivial_numbers"]), [], []
    chunk_table = lambda chunk: dict(chunk["table"], row_key=[chunk["table"]["header"][0]] if chunk["table"]["header"] else [])
    for unit in world["units"]:                          # 1. stored tables
        table = world["tables"].get(unit["ref"])
        if not table:
            continue
        targets = [ref for ref in linked(world, unit["ref"], "C-") + linked(world, unit["ref"], "D-") if world["by_ref"][ref].get("table")]
        if not any(ref.startswith("C-") for ref in targets):
            shapes = sorted(((table_shape_score(table, c["table"], stop), c["ref"]) for c in world["canon"] if c.get("table")), reverse=True)
            if shapes and shapes[0][0] >= 2.0 and (len(shapes) == 1 or shapes[1][0] < shapes[0][0]):
                targets.append(shapes[0][1])
                edges.append(core.Edge(unit["ref"], shapes[0][1], "corresponds", core.HOW_TABLE, ctx.provenance, relation="Implements",
                                         evidence={"how_text": core.HOW_TABLE + "."}))
        for target in sorted(set(targets)):
            result = reconcile(table, chunk_table(world["by_ref"][target]), stop, exact=True)
            labels = ("package", "methodology" if target.startswith("C-") else "documentation")
            checks.append(dict(result, unit_ref=unit["ref"], target_ref=target, what="parameter table", labels=list(labels),
                               sentences=table_sentences(result, labels[0], labels[1], target)))
    pending = {}                                         # tables that code could not lay over each other: ask which columns correspond
    for check in checks:
        if check["outcome"] == "could not be compared" and ctx.ask is not None:
            unit, other = world["by_ref"][check["unit_ref"]], world["by_ref"][check["target_ref"]]
            table = world["tables"][unit["ref"]]
            shown = "\n".join(["; ".join(table["header"])] + ["; ".join(row) for row in table["rows"][:3]])
            question = narrow_question("map-table-columns", unit["ref"], [("PACKAGE TABLE (%s)" % unit["name"], shown)], settings,
                                                     passages=[(other["ref"], passage_label(other), passage_text(other, settings))],
                                                     more={"package_header": list(table["header"])})
            question["other_headers"] = {letter: list(world["by_ref"][ref]["table"]["header"]) for letter, ref in question["letters"].items()}
            pending[question["question_id"]] = (check, question)
    answers = ctx.ask([q for _, q in pending.values() if not q["too_large"]]) if pending else {}
    for question_id, (check, question) in sorted(pending.items()):
        final = answers.get(question_id)
        if final and final["outcome"] == "accepted" and final["answer"].get("table") != "NONE":
            answer = final["answer"]
            ai_map = {"columns": [(c["package"], c["other"]) for c in answer["columns"]], "key": (answer["key"]["package"], answer["key"]["other"])}
            result = reconcile(world["tables"][check["unit_ref"]], chunk_table(world["by_ref"][check["target_ref"]]), stop, exact=True, ai_map=ai_map)
            check.update(result, sentences=table_sentences(result, check["labels"][0], check["labels"][1], check["target_ref"]))
    for (source, target), edge in sorted(world["links"].items()):          # 2. documentation tables against methodology tables
        left, right = world["by_ref"][source], world["by_ref"][target]
        if source.startswith("D-") and target.startswith("C-") and left.get("table") and right.get("table") and edge["relation"] in core.LINKING_RELATIONS:
            result = reconcile(chunk_table(left), chunk_table(right), stop, exact=False)
            checks.append(dict(result, unit_ref=source, target_ref=target, what="documentation table", labels=["documentation", "methodology"],
                               sentences=table_sentences(result, "documentation", "methodology", target)))
    for unit in world["units"]:                          # 3. numbers written in linked code
        code = unit.get("code") or {}
        has_children = unit["kind"] == core.KIND_FUNCTION and any(c["kind"] == core.KIND_FORMULA for c in world["children"].get(unit["ref"], []))
        numbers = [n for n in code.get("numbers", ()) if n["value"] not in trivial]
        if unit["kind"] not in (core.KIND_FUNCTION, core.KIND_FORMULA) or has_children or not numbers or not linked(world, unit["ref"], "C-"):
            continue
        family = [unit["ref"], unit.get("parent_ref")] + [c["ref"] for c in world["children"].get(unit.get("parent_ref") or "", [])]
        passages = sorted({ref for member in family if member for ref in linked(world, member, "C-")})
        cited = sorted({e["target"] for ref in passages for e in world["graph"]["out"].get(ref, []) if e["kind"] == "cross_reference"})
        stated = [dict(n, ref=ref, value=n["value"].lstrip("-")) for ref in passages + cited for n in core.find_numbers(world["by_ref"][ref]["text"])]
        cells, missing = [], False                       # a sign belongs to the formula, which the mathematical check covers
        for number in numbers:
            match = next((s for s in stated if compare_values(s, number, exact=True)[0] != "differs"), None)
            missing = missing or match is None
            cells.append("%s: stated in %s as %s" % (number["as_written"], match["ref"], quote(match["as_written"])) if match
                         else "%s: not stated in the linked passages" % number["as_written"])
        checks.append({"unit_ref": unit["ref"], "target_ref": ", ".join(passages), "what": "number in code", "cells": cells, "column_map": [],
                       "key_map": "", "only_in_package": [], "only_in_other": [], "outcome": "not stated" if missing else "agrees", "sentences": cells})
    sources = [(c, "documentation", "number in documentation", [(r, world["by_ref"][r]["text"]) for r in linked(world, c["ref"], "C-")])
               for c in world["doc"] if c["kind"] == "Paragraph"]
    for unit in world["units"]:                          # 4. numbers in prose: documentation and roxygen text
        target = (unit.get("roxygen") or {}).get("documents_ref")
        if target and world["by_ref"][target]["kind"] == core.KIND_FUNCTION:
            family = [target] + [c["ref"] for c in world["children"].get(target, [])]
            passages = sorted({ref for member in family for ref in linked(world, member, "C-")})
            text = " ".join(tag["text"] for tag in unit["roxygen"]["tags"] if tag["tag"] in ("description", "details", "param", "return"))
            sources.append((dict(unit, text=text), "package documentation", "number in roxygen", [(r, world["by_ref"][r]["text"]) for r in passages]))
    for source, label, what, stated_texts in sources:
        if stated_texts:
            lines, differs = prose_value_lines(source["text"], stated_texts, stop, trivial, label)
            if lines:
                checks.append({"unit_ref": source["ref"], "target_ref": ", ".join(ref for ref, _ in stated_texts), "what": what, "cells": lines,
                               "column_map": [], "key_map": "", "only_in_package": [], "only_in_other": [],
                               "outcome": "differs" if differs else "agrees", "sentences": lines})
    for number, check in enumerate(checks, start=1):
        check.update(check_id="VC-%04d" % number, rule_text=VALUE_RULE_TEXT)
        if check["what"] in ("parameter table", "documentation table") and check["outcome"] in ("agrees", "differs") and check["target_ref"].startswith("C-"):
            how = core.HOW_VALUE_DIFFERS if check["outcome"] == "differs" else core.HOW_VALUE_AGREES
            edges.append(check_edge(ctx, world, check["unit_ref"], check["target_ref"], how, check["sentences"][0],
                                    relation="Differs from" if check["outcome"] == "differs" else None))
    counts = {"checks": len(checks), "with a difference": sum(1 for c in checks if c["outcome"] in ("differs", "not stated"))}
    return core.StepResult({"value_checks": checks, "graph_ledger": ledger_records(world["ledger"], edges)}, counts, [])

# ---------------------------------------------------------------- step 13: check-rules
FLOOR_PHRASES = ("at least", "no less than", "not less than", "never below", "not below", "never taken below", "taken below",
                 "minimum of", "floor of", "floored at", "subject to a minimum")
CAP_PHRASES = ("at most", "no more than", "not more than", "not exceed", "never above", "maximum of", "cap of", "capped at",
               "subject to a maximum")

def stated_rules(chunk):
    """Floors and caps a passage states: (kind, number, sentence), recognised by generic phrases."""
    rules = []
    for sentence in re.split(r"(?<=[.;])\s+", chunk["text"]):
        lowered = sentence.lower()
        for kind, phrases in (("floor", FLOOR_PHRASES), ("cap", CAP_PHRASES)):
            phrase = next((p for p in phrases if p in lowered), None)
            numbers = core.find_numbers(sentence)
            if phrase and numbers:
                after = [n for n in numbers if n["position"] >= lowered.find(phrase)]
                rules.append((kind, (after or numbers)[0], sentence.strip()))
                break
    return rules

def rule_in_trees(kind, number, members):
    """Code looks first: a maximum (for a floor), a minimum (for a cap), a piecewise expression
    or a comparison that holds the stated number, written as a number or held as the default
    of an argument. Returns where it was seen, or ""."""
    wanted = ("max",) if kind == "floor" else ("min",)
    defaults = {core.normalise_symbol(name): number_of(default) for member in members
                for name, default in ((member.get("code") or {}).get("formals") or ()) if default and number_of(default)}
    def holds_number(arg):
        held = {"value": arg.value, "decimals": 0} if arg.op == "num" else defaults.get(arg.name) if arg.op == "sym" else None
        return bool(held) and compare_values(number, held, exact=True)[0] != "differs"
    for member in members:
        code = member.get("code") or {}
        for form in ("expression", "composed"):
            for node in core.expr_walk(core.expr_from_dict(code[form])) if code.get(form) else ():
                holds = node.op in ("cmp", "piecewise") or (node.op == "call" and node.name in wanted + ("piecewise",))
                if holds and any(holds_number(arg) for arg in node.args):
                    name = "maximum" if node.name == "max" else "minimum" if node.name == "min" else "condition"
                    return "%s with %s (%s, line %s)" % (name, number["as_written"], member["ref"], (member.get("lines") or ["?"])[0])
    return ""

def check_rules(ctx):
    """Step 13, check-rules. For every top-level function (and every formula statement
    outside a function) linked to a passage that states a floor or a cap: code inspection
    first; only when the code shows no such node is the check-rule question asked."""
    world, settings = load_world(ctx), ctx.settings
    pending, checks = [], []
    for unit in world["units"]:
        if unit["kind"] not in (core.KIND_FUNCTION, core.KIND_FORMULA) or unit.get("parent_ref"):
            continue
        members = [unit] + world["children"].get(unit["ref"], [])
        passages = sorted({ref for member in members for ref in linked(world, member["ref"], "C-", CODE_RELATIONS)})
        for ref in passages:
            for kind, number, sentence in stated_rules(world["by_ref"][ref]):
                where = rule_in_trees(kind, number, members)
                record = {"unit_ref": unit["ref"], "target_ref": ref, "rule_as_stated": sentence, "rule_kind": kind, "rule_number": number["as_written"],
                          "located_by": "code inspection", "outcome": "applied" if where else "", "where_in_code": where}
                if not where:
                    code_text = cut_code(unit["text"], int(settings["max_unit_chars"]))
                    question = narrow_question("check-rule", unit["ref"], [("UNIT (%s)" % passage_label(unit), code_text),
                                                             ("RULE AS STATED", sentence)], settings, more={"rule_text": sentence, "code_text": code_text})
                    pending.append((record, question))
                checks.append(record)
    answers = ctx.ask([q for _, q in pending if not q["too_large"]]) if pending else {}
    for record, question in pending:
        final = answers.get(question["question_id"])
        record["located_by"] = "AI (validated)"
        if final and final["outcome"] == "accepted":
            record.update(outcome=final["answer"]["outcome"], where_in_code=final["answer"].get("quote_from_unit", ""))
        else:
            record.update(outcome="could not be decided", where_in_code="")
    for number, record in enumerate(checks, start=1):
        record["check_id"] = "RC-%04d" % number
        what = "%s of %s stated in %s" % (record["rule_kind"].capitalize(), record["rule_number"], record["target_ref"])
        shown = {"applied": "present in the code", "not applied": "not present in the code", "applied differently": "applied differently in the code",
                 "could not be decided": "could not be decided (the AI's answer could not be used)"}[record["outcome"]]
        where = record["where_in_code"] if record["located_by"] == "code inspection" else quote(record["where_in_code"]) if record["where_in_code"] else ""
        record["sentence"] = "%s: %s%s; located by %s" % (what, shown, " (%s)" % where if where else "", record["located_by"])
    return core.StepResult({"rule_checks": checks}, {"rules checked": len(checks), "not applied as stated": sum(1 for r in checks if r["outcome"] in ("not applied", "applied differently"))}, [])

# ---------------------------------------------------------------- step 14: check-package-docs (no AI)
def check_package_docs(ctx):
    """Step 14, check-package-docs. Roxygen blocks and help pages against the code they
    document: arguments, stated defaults, stated values, @export against NAMESPACE, usage
    against the signature, page in step with its source block, block and page present for
    exported objects, a data block for every stored object, example code that parses."""
    world, checks = load_world(ctx), []
    trivial = set(ctx.settings["trivial_numbers"])
    def add(unit_ref, about_ref, check, outcome, detail):
        checks.append({"unit_ref": unit_ref, "about_ref": about_ref, "check": check, "outcome": outcome, "detail": detail})
    blocks, pages = {}, {}
    for unit in world["units"]:
        if unit.get("roxygen") and unit["roxygen"]["documents_ref"]:
            blocks[unit["roxygen"]["documents_ref"]] = unit
        if unit.get("helppage"):
            pages[unit["helppage"]["rd_name"]] = unit
    for target_ref, block in sorted(blocks.items()):
        target, tags = world["by_ref"][target_ref], block["roxygen"]["tags"]
        tag_names = {tag["tag"] for tag in tags}
        if target.get("data"):
            described = set(re.findall(r"\\item\{([^{}]+)\}", block["text"]))
            missing = [name for name, _ in target["data"]["columns"] if name not in described and name != "(row name)"]
            add(target_ref, block["ref"], "columns described", "differs" if missing else "agrees",
                "The data block does not describe: %s" % ", ".join(missing) if missing else "Every column is described in %s" % block["ref"])
            continue
        if target["kind"] != core.KIND_FUNCTION:
            continue
        formals = [name for name, _ in target["code"]["formals"]]
        documented = [name.strip() for tag in tags if tag["tag"] == "param" for name in tag["name"].split(",")]
        if tag_names & {"inheritParams", "rdname", "describeIn"}:
            add(block["ref"], target_ref, "arguments", "not applicable", "Arguments are documented elsewhere (inherited); not compared")
        else:
            missing = [name for name in formals if name not in documented and name != "..."]
            extra = [name for name in documented if name not in formals]
            parts = ["@param %s is documented; %s has no such argument" % (name, target["name"]) for name in extra]
            parts += ["argument %s of %s is not documented" % (name, target["name"]) for name in missing]
            add(block["ref"], target_ref, "arguments", "differs" if parts else "agrees",
                "; ".join(parts) + " (arguments: %s)" % ", ".join(formals) if parts else "Every argument is documented")
        defaults = dict(target["code"]["formals"])
        for tag in tags:
            stated = re.search(r"[Dd]efaults?(?:\s+(?:to|is|of|value))?:?\s+([-\w.%\"']+)", tag["text"]) if tag["tag"] == "param" else None
            if stated and defaults.get(tag["name"]):
                written, actual = stated.group(1).strip("\"'."), defaults[tag["name"]].strip("\"'")
                same = (compare_values(number_of(written), number_of(actual), exact=True)[0] != "differs") if number_of(written) and number_of(actual) \
                    else written == actual
                add(block["ref"], target_ref, "defaults", "agrees" if same else "differs",
                    "@param %s states the default %s; the code has %s" % (tag["name"], written, actual))
        exported = target["code"]["exported"]
        if exported is not None and ("export" in tag_names) != bool(exported):
            add(block["ref"], target_ref, "export", "differs", "The block has @export but NAMESPACE does not export %s" % target["name"]
                if "export" in tag_names else "NAMESPACE exports %s but the block has no @export" % target["name"])
        members = [target] + world["children"].get(target_ref, [])
        used = [n for member in members for n in (member.get("code") or {}).get("numbers", ())] + [number_of(d) for d in defaults.values() if d and number_of(d)]
        text = " ".join(tag["text"] for tag in tags if tag["tag"] in ("description", "details", "return"))
        absent = [n["as_written"] for n in core.find_numbers(re.sub(r"\\d?eqn\{.*?\}\}?", " ", text)) if n["value"] not in trivial
                  and not any(compare_values(n, u, exact=True)[0] != "differs" for u in used)]
        if core.find_numbers(text):
            add(block["ref"], target_ref, "stated values", "differs" if absent else "agrees",
                "The block states %s; the code of %s does not use %s" % (", ".join(absent), target["name"], "this value" if len(absent) == 1 else "these values")
                if absent else "Every value stated in the block is used in the code")
        for tag in tags:
            if tag["tag"] == "examples" and any(isinstance(r, tuple) for r in reading.parse_r_source(tag["text"])):
                add(block["ref"], target_ref, "examples", "differs", "The example code could not be parsed")
    for unit in world["units"]:
        page = unit.get("helppage")
        if page:
            target = next((u for u in world["units"] if u["kind"] == core.KIND_FUNCTION and u["name"] == page["rd_name"] and not u["inside"]), None)
            if target:
                usage = re.search(r"\(([^)]*)\)", page["usage"])
                shown = [part.split("=")[0].strip() for part in usage.group(1).split(",")] if usage and usage.group(1).strip() else []
                formals = [name for name, _ in target["code"]["formals"]]
                add(unit["ref"], target["ref"], "usage", "agrees" if shown == formals else "differs",
                    "Usage shows (%s); the function has (%s)" % (", ".join(shown), ", ".join(formals)))
            if page["generated_from_ref"] is None:
                add(unit["ref"], target["ref"] if target else None, "page in step", "not applicable", "A hand-written page (no roxygen source block was found)")
            elif page["in_step_with_source"] is False:
                add(unit["ref"], page["generated_from_ref"], "page in step", "differs", "Out of step with its roxygen source %s: the arguments differ" % page["generated_from_ref"])
            else:
                add(unit["ref"], page["generated_from_ref"], "page in step", "agrees", "In step with its roxygen source %s" % page["generated_from_ref"])
        code = unit.get("code") or {}
        if unit["kind"] == core.KIND_FUNCTION and not unit["inside"]:
            if code.get("exported") and unit["ref"] not in blocks:
                add(unit["ref"], None, "block present", "differs", "Exported function %s has no roxygen block" % unit["name"])
            elif code.get("exported") and unit["name"] not in pages and any(u.get("helppage") for u in world["units"]):
                add(unit["ref"], None, "page present", "differs", "Exported function %s has no help page" % unit["name"])
            elif unit["ref"] not in blocks:
                add(unit["ref"], None, "block present", "not applicable", "Internal function, no block")
        if unit.get("data") and unit["data"]["assessable"] and unit["ref"] not in blocks and unit["file"].startswith("data/"):
            add(unit["ref"], None, "data block present", "differs", "Stored object %s has no roxygen data block" % unit["name"])
    for number, check in enumerate(checks, start=1):
        check["check_id"] = "PD-%04d" % number
    return core.StepResult({"package_doc_checks": checks}, {"checks": len(checks), "differ": sum(1 for c in checks if c["outcome"] == "differs")}, [])

# ---------------------------------------------------------------- step 15: account-coverage
NEXT_STEPS = {
    core.CAT_CODE_DIFFERS: "Compare the code with the cited passage, starting from the inputs shown under What was observed.",
    core.CAT_CODE_NOT_TRACED: "Decide whether this code implements a part of the methodology; if so, name the passage.",
    core.CAT_MATH_UNDECIDED: "Compare the formula in the code with the cited passage by hand; the tool could not decide it.",
    core.CAT_VALUE_DIFFERS: "Compare the listed rows of the package table with the cited table of the methodology.",
    core.CAT_NUMBER_NOT_TRACED: "Find where the methodology states this number, or confirm that it needs no statement.",
    core.CAT_DATA_NOT_TRACED: "Decide which table of the methodology this stored object corresponds to, if any.",
    core.CAT_DATA_NOT_DESCRIBED: "Check whether the stored object and each of its columns should be described in the package.",
    core.CAT_PKGDOC_VS_CODE: "Compare the roxygen block or help page with the function it documents.",
    core.CAT_PKGDOC_VS_CANON: "Compare the value stated in the package documentation with the cited passage.",
    core.CAT_DOC_NOT_TRACED: "Decide which passage of the methodology this statement of the documentation rests on, if any.",
    core.CAT_DOC_VS_CANON: "Compare the statement in the documentation with the cited passage of the methodology.",
    core.CAT_DOC_VS_CODE: "Compare the statement in the documentation with the cited unit of the package.",
    core.CAT_NOT_READ: "Review this item by hand; the tool could not read or assess it.",
    core.CAT_AI_UNUSABLE: "Review this unit by hand, or run the tool again; the AI's answer could not be used.",
    core.CAT_AI_DISAGREE: "Read both quotations and decide whether the unit and the passage state the same thing."}

STATUS_RULES = (      # applied top to bottom; the first rule that fits decides. The manual's table is generated from this list.
    ("model", "could not be read or assessed", core.ST_NOT_ASSESSED, "A file that could not be read, compiled code, or a stored object that is not compared."),
    ("model", "test block", core.ST_UNIT_TEST, "A test_that block. Which function it tests is shown in the column Unit test."),
    ("model", "package file without code", core.ST_SUPPORTING, "DESCRIPTION, NAMESPACE and other files that hold no code."),
    ("model", "example code in a vignette", core.ST_SUPPORTING, "Code inside a vignette; it illustrates the package and is not part of the model."),
    ("model", "its own checks", core.ST_DIFFERS + " / " + core.ST_UNDECIDED, "A roxygen block or help page whose own deterministic checks differ or are undecided."),
    ("model", "documents no single object", core.ST_SUPPORTING, "A roxygen block or help page that documents no single object of the package."),
    ("model", "takes the tracing of what it documents", "the status of the documented object", "A roxygen block or help page whose own checks pass."),
    ("model", "a check or the judge reports a difference", core.ST_DIFFERS, "Linked, and a check, the judge or the second question reports a difference."),
    ("model", "a required check is undecided", core.ST_UNDECIDED, "Linked, and a required check could not be decided."),
    ("model", "linked and all required checks agree", core.ST_TRACED, "Linked to the methodology, and every required check agrees."),
    ("model", "covered by the check of the whole function", core.ST_TRACED, "A statement without a link of its own, inside a function whose agreeing check ran through it."),
    ("model", "supporting code by syntax", core.ST_SUPPORTING, "No link, and no arithmetic and no number other than the trivial ones; the reason is shown."),
    ("model", "vignette prose without checkable statements", core.ST_SUPPORTING, "Vignette text that states no number and no formula."),
    ("model", "not linked", core.ST_NOT_TRACED, "Anything else without a link to the methodology."),
    ("doc", "content cannot be read", core.ST_NOT_ASSESSED, "A figure, an equation that could not be read, or a part of a file that could not be read."),
    ("doc", "states nothing checkable", core.ST_NARRATIVE, "No formula, number, rule or definition, by code or by the judge's answer."),
    ("doc", "a check or the judge reports a difference", core.ST_DIFFERS, "Linked, and a check, the judge or the second question reports a difference."),
    ("doc", "a required check is undecided", core.ST_UNDECIDED, "Linked, and a required check could not be decided."),
    ("doc", "linked and consistent", core.ST_TRACED, "Linked to the methodology, directly or through a linked unit of the package, and consistent."),
    ("doc", "checkable and not linked", core.ST_NOT_TRACED, "States something checkable, and nothing was linked to it."))

def concerns_of(record):
    """Which of the four Concerns a unit falls under, from its kind."""
    if record.get("heading_chain") is not None:
        return core.CONCERNS[3] if record["corner"] == "doc" else core.CONCERNS[0]
    if record.get("data") or record["kind"] in (core.KIND_TABLE, core.KIND_OBJECT):
        return core.CONCERNS[1]
    return core.CONCERNS[2] if record["kind"] in (core.KIND_ROXYGEN, core.KIND_HELP, core.KIND_VIGNETTE) else core.CONCERNS[0]

def citation(record):
    """How a unit or a passage is cited next to a quotation."""
    if record.get("heading_chain") is not None:
        place = " > ".join(record["heading_chain"][-2:]) or record["source_file"]
        return "%s, %s, %s%s" % (record["ref"], place, record["kind"].lower(), " %d" % record["para_no"] if record.get("para_no") else "")
    lines = " lines %d-%d" % tuple(record["lines"]) if record.get("lines") else ""
    return "%s, %s%s" % (record["ref"], record["file"], lines)

def render_path(record, world):
    """The supporting path of an item, hop by hop, each hop with its citation (plan 2.5)."""
    graph, by_ref = world["graph"], world["by_ref"]
    if record.get("heading_chain") is not None:
        hops = ["%s %s in %s" % (record["ref"], record["kind"].lower(), " > ".join(record["heading_chain"][-2:]) or record["source_file"])]
    else:
        hops = ["%s %s `%s` (%s)" % (record["ref"], record["kind"].lower(), record["name"], citation(record).split(", ", 1)[1])]
    parent = by_ref.get(record.get("parent_ref") or "")
    if parent:
        hops.append("inside %s %s %s" % (parent["kind"].lower(), parent["ref"], parent["name"]))
    for owner in [record] + ([parent] if parent else []):
        for edge in graph["out"].get(owner["ref"], []):
            if edge["kind"] == "reads_data":
                where = "; ".join("%s %s" % (k, edge["evidence"][k]) for k in ("row_key", "column") if edge["evidence"].get(k))
                hops.append("reads %s %s %s%s" % (by_ref[edge["target"]]["kind"].lower(), edge["target"], by_ref[edge["target"]]["name"], " [%s]" % where if where else ""))
    for prefix, verb in (("C-", "corresponds to"), ("M-", "corresponds to"), ("D-", "described by documentation")):
        for ref in linked(world, record["ref"], prefix)[:2]:
            edge = world["links"].get((record["ref"], ref)) or world["links"].get((ref, record["ref"]))
            hops.append("%s %s (%s): %s" % (verb, ref, edge["relation"], edge["how"]))
    for edge in graph["in"].get(record["ref"], []):
        if edge["kind"] == "documents":
            hops.append("documented by %s %s" % (by_ref[edge["source"]]["kind"].lower(), edge["source"]))
    return [hops[0]] + ["  > " + hop for hop in dict.fromkeys(hops[1:8])]

def gather(ctx, world):
    """All check records and AI notes, grouped by the unit they belong to."""
    facts = {}
    def at(ref):
        return facts.setdefault(ref, {"math": [], "values": [], "rules": [], "pkgdoc": [], "problems": [], "opinions": [], "unresolved": []})
    for record in ctx.read("math_checks"):
        at(record["target_ref"] if record["formula_source"] == "roxygen" else record["unit_ref"])["math"].append(record)
        if record["formula_source"] == "roxygen":
            at(record["unit_ref"])["math"].append(dict(record, shown_only=True))
    for kind, slot in (("value_checks", "values"), ("rule_checks", "rules"), ("package_doc_checks", "pkgdoc"),
                       ("second_opinions", "opinions"), ("unresolved_references", "unresolved")):
        for record in ctx.read(kind):
            at(record["unit_ref"])[slot].append(record)
    seen = set()
    for record in ctx.read("judgement_problems"):
        if (record["unit_ref"], record["target_corner"]) not in seen:
            seen.add((record["unit_ref"], record["target_corner"]))
            at(record["unit_ref"])["problems"].append(record)
    facts["states_nothing"] = {r["unit_ref"] for r in ctx.read("doc_judgements") if r.get("states_nothing_checkable")}
    return facts

def ai_judged_difference(world, ref, prefix):
    """Passages the AI judge itself called deviating or inconsistent (not a check's edge)."""
    return [t for (s, t), e in sorted(world["links"].items()) if s == ref and t.startswith(prefix) and e["relation"] == "Differs from"
            and not e["evidence"].get("check")]

def model_unit_outcome(unit, world, facts, raised):
    """The status rules for one model unit, top to bottom (plan 2.9). Returns (status, name of
    the deciding rule, reason shown). Differences and undecided checks are raised as it goes."""
    ref, kind, code, mine = unit["ref"], unit["kind"], unit.get("code") or {}, facts.get(unit["ref"], {})
    def note(category, line, check_id=""):
        entry = raised.setdefault((ref, category), {"lines": [], "checks": []})
        entry["lines"].append(line)
        entry["checks"].append(check_id)
    if kind in (core.KIND_NOT_READ, core.KIND_COMPILED) or (unit.get("data") and not unit["data"]["assessable"]):
        reason = unit.get("read_problem") or "This stored object was not compared: %s." % (unit.get("data") or {}).get("not_assessable_reason", "it could not be read")
        note(core.CAT_NOT_READ, reason)
        return core.ST_NOT_ASSESSED, "could not be read or assessed", reason
    if kind == core.KIND_TEST:
        return core.ST_UNIT_TEST, "test block", "a test block"
    if kind == core.KIND_OTHER:
        return core.ST_SUPPORTING, "package file without code", "a package file without code (%s)" % unit["name"]
    if unit["file"].startswith(("vignettes/", "inst/doc/")) and kind != core.KIND_VIGNETTE:
        return core.ST_SUPPORTING, "example code in a vignette", "example code in a vignette; it is not part of the model"
    differs, undecided = False, False
    for record in mine.get("math", []):
        if record.get("shown_only"):
            continue
        if record["outcome"] == "differs":
            differs = True
            note(core.CAT_PKGDOC_VS_CODE if record["formula_source"] == "roxygen" else core.CAT_CODE_DIFFERS, record["sentence"], record["check_id"])
        elif record["outcome"] == core.CHECK_UNDECIDED:
            undecided = True
            note(core.CAT_MATH_UNDECIDED, record["sentence"], record["check_id"])
    for target in ai_judged_difference(world, ref, "C-"):
        differs = True
        note(core.CAT_CODE_DIFFERS, "The AI judged that this unit deviates from %s: %s" % (target, world["links"][(ref, target)]["evidence"]["how_text"].split("\n")[0]))
    for record in mine.get("values", []):
        category = {"parameter table": core.CAT_VALUE_DIFFERS, "number in code": core.CAT_NUMBER_NOT_TRACED,
                    "number in roxygen": core.CAT_PKGDOC_VS_CANON}.get(record["what"], core.CAT_VALUE_DIFFERS)
        if record["what"] == "parameter table" and record["target_ref"].startswith("D-"):
            category = core.CAT_DOC_VS_CODE
        if record["outcome"] in ("differs", "not stated"):
            differs = True
            for line in [s for s in record["sentences"] if "differs" in s or "not stated" in s or "only" in s] or record["sentences"][:1]:
                note(category, line, record["check_id"])
        elif record["outcome"] == "could not be compared":
            undecided = True
            note(core.CAT_NOT_READ, record["sentences"][0], record["check_id"])
    for record in mine.get("rules", []):
        if record["outcome"] in ("not applied", "applied differently"):
            differs = True
            note(core.CAT_CODE_DIFFERS, record["sentence"] + ". The rule as stated: " + quote(record["rule_as_stated"], record["target_ref"]), record["check_id"])
        elif record["outcome"] == "could not be decided":
            undecided = True
            note(core.CAT_AI_UNUSABLE, record["sentence"], record["check_id"])
    for record in mine.get("pkgdoc", []):
        if record["outcome"] == "differs":
            differs = True
            note(core.CAT_DATA_NOT_DESCRIBED if unit.get("data") else core.CAT_PKGDOC_VS_CODE, record["detail"], record["check_id"])
    for record in mine.get("opinions", []):
        differs = True
        note(core.CAT_AI_DISAGREE, "A second, oppositely framed question named a difference from %s: %s against %s" % (
            record["target_ref"], quote(record["quote_from_unit"], ref), quote(record["quote_from_passage"], record["target_ref"])))
    if kind in (core.KIND_ROXYGEN, core.KIND_HELP):
        return ("own checks", differs, undecided)
    is_linked = bool(linked(world, ref, "C-"))
    if is_linked or (differs and not is_linked and (code.get("plumbing") or unit.get("data"))):
        if differs:
            return core.ST_DIFFERS, "a check or the judge reports a difference", "see the flagged item(s)"
        if undecided:
            return core.ST_UNDECIDED, "a required check is undecided", "see the flagged item(s)"
        return core.ST_TRACED, "linked and all required checks agree", "linked to %s" % ", ".join(linked(world, ref, "C-"))
    parent = world["by_ref"].get(unit.get("parent_ref") or "")
    parent_math = [r for r in facts.get(parent["ref"], {}).get("math", []) if r["outcome"] == "agrees"] if parent else []
    if kind == core.KIND_FORMULA and parent and not differs and not undecided and any(
            r["through"] == ref or (r["through"] == parent["ref"] and not unit["inside"] == "") for r in parent_math):
        return core.ST_TRACED, "covered by the check of the whole function", "covered by the check of function %s, which agrees" % parent["ref"]
    parent_plumbing = bool(parent and (parent.get("code") or {}).get("plumbing"))
    if (code.get("plumbing") or parent_plumbing) and not differs and not undecided:
        return core.ST_SUPPORTING, "supporting code by syntax", code.get("plumbing_reason") or "a statement inside supporting code"
    if kind == core.KIND_VIGNETTE and not core.find_numbers(unit["text"]) and "=" not in unit["text"]:
        return core.ST_SUPPORTING, "vignette prose without checkable statements", "vignette prose that states no number and no formula"
    problem = next((p for p in mine.get("problems", []) if p["target_corner"] == "canon"), None)
    if problem:
        note(core.CAT_AI_UNUSABLE, "The AI's answer about the methodology could not be used: %s." % problem["reason"])
    else:
        category = core.CAT_DATA_NOT_TRACED if unit.get("data") else core.CAT_DOC_NOT_TRACED if kind == core.KIND_VIGNETTE else core.CAT_CODE_NOT_TRACED
        note(category, "No passage of the methodology was linked to this unit.")
    return core.ST_NOT_TRACED, "not linked", "no link to the methodology"

def doc_unit_outcome(chunk, world, facts, raised):
    """The status rules for one documentation unit, top to bottom (plan 2.9)."""
    ref, mine = chunk["ref"], facts.get(chunk["ref"], {})
    def note(category, line, check_id=""):
        entry = raised.setdefault((ref, category), {"lines": [], "checks": []})
        entry["lines"].append(line)
        entry["checks"].append(check_id)
    unreadable = chunk["kind"] == "Figure" or chunk.get("not_read_reason") or (chunk["kind"] == "Equation" and not chunk["equation"]["readable"])
    if unreadable:
        reason = chunk.get("not_read_reason") or ("the content of a figure cannot be read" if chunk["kind"] == "Figure" else chunk["equation"]["not_readable_reason"])
        note(core.CAT_NOT_READ, "This %s was not assessed: %s." % (chunk["kind"].lower(), reason))
        return core.ST_NOT_ASSESSED, "content cannot be read", reason
    to_canon, to_model = linked(world, ref, "C-"), linked(world, ref, "M-")
    if not to_canon and (not chunk.get("checkable") or (ref in facts["states_nothing"] and not to_model)):
        return core.ST_NARRATIVE, "states nothing checkable", "no formula, number, rule or definition"
    differs, undecided = False, False
    for target in ai_judged_difference(world, ref, "C-"):
        differs = True
        note(core.CAT_DOC_VS_CANON, "The AI judged this passage inconsistent with %s." % target)
    for target in ai_judged_difference(world, ref, "M-"):
        differs = True
        note(core.CAT_DOC_VS_CODE, "The AI judged this passage inconsistent with %s." % target)
    for record in mine.get("values", []):
        if record["outcome"] == "differs":
            differs = True
            for line in [s for s in record["sentences"] if "differs" in s or "only" in s]:
                note(core.CAT_DOC_VS_CANON, line, record["check_id"])
    for unit_ref in to_model:                            # a package table that differs from this documentation table
        for record in facts.get(unit_ref, {}).get("values", []):
            if record["target_ref"] == ref and record["outcome"] == "differs":
                differs = True
                note(core.CAT_DOC_VS_CODE, record["sentences"][0], record["check_id"])
    for record in mine.get("math", []):
        if record["outcome"] == "differs":
            differs = True
            note(core.CAT_DOC_VS_CANON, record["sentence"], record["check_id"])
        elif record["outcome"] == core.CHECK_UNDECIDED:
            undecided = True
            note(core.CAT_MATH_UNDECIDED, record["sentence"], record["check_id"])
    for record in mine.get("opinions", []):
        differs = True
        note(core.CAT_AI_DISAGREE, "A second, oppositely framed question named a difference from %s: %s against %s" % (
            record["target_ref"], quote(record["quote_from_unit"], ref), quote(record["quote_from_passage"], record["target_ref"])))
    through = [u for u in to_model if linked(world, u, "C-")]
    if to_canon or through:
        if differs:
            return core.ST_DIFFERS, "a check or the judge reports a difference", "see the flagged item(s)"
        if undecided:
            return core.ST_UNDECIDED, "a required check is undecided", "see the flagged item(s)"
        return core.ST_TRACED, "linked and consistent", "linked to %s" % ", ".join(to_canon or ["the methodology through " + through[0]])
    problem = next((p for p in mine.get("problems", [])), None)
    note(core.CAT_AI_UNUSABLE if problem else core.CAT_DOC_NOT_TRACED, "The AI's answer could not be used: %s." % problem["reason"] if problem
         else "This passage states something checkable, and no passage of the methodology was linked to it.")
    return core.ST_NOT_TRACED, "checkable and not linked", "no link to the methodology"

def lines_or(lines, fallback):
    """Several lines for one cell without repeats, or the fallback sentence when there is none."""
    return "\n".join(dict.fromkeys(lines)) if lines else fallback

def model_cells(unit, world, facts):
    """The assessment cells of a model unit, shown on its row of the map and its status on Chunks_Model. Every cell shows real
    content or "Not applicable" for this kind of unit."""
    mine, na = facts.get(unit["ref"], {}), core.NOT_APPLICABLE
    is_code = unit["kind"] in (core.KIND_FUNCTION, core.KIND_FORMULA)
    math = [r for r in mine.get("math", [])]
    pairs = ["%s: `%s` (%s)" % (stated, c, how) for r in math for c, stated, how in r["alignment"]]
    tables = [s for r in mine.get("values", []) if r["what"] == "parameter table" for s in r["sentences"]]
    tests = ["Tested by %s (%s)" % (e["target"], citation(world["by_ref"][e["target"]]).split(", ", 1)[1])
             for e in world["graph"]["out"].get(unit["ref"], []) if e["kind"] == "tested_by"]
    documented = [r["detail"] for r in mine.get("pkgdoc", [])]
    block = next((e["source"] for e in world["graph"]["in"].get(unit["ref"], []) if e["kind"] == "documents"), None)
    if block:
        documented += [r["detail"] for r in facts.get(block, {}).get("pkgdoc", [])]
    documented += ["Described in %s" % ref for ref in linked(world, unit["ref"], "D-")]
    ai_notes = ["AI (second question), about %s: %s" % (r["target_ref"], quote(r["what_differs"])) for r in mine.get("opinions", [])]
    return {"math_check": lines_or([r["sentence"] for r in math], "No linked passage states a formula" if is_code else na),
            "parameter_completeness": lines_or(tables or pairs, na if not unit.get("data") else "No table was linked or matched"),
            "logic_consistency": lines_or([r["sentence"] for r in mine.get("rules", [])], "No linked passage states a floor or a cap" if is_code else na),
            "documentation_consistency": lines_or(documented, na),
            "hard_coded_numbers": lines_or([s for r in mine.get("values", []) if r["what"] == "number in code" for s in r["sentences"]],
                                           "No number other than the trivial ones" if is_code else na),
            "unit_test": lines_or(tests, "No test block calls this function") if unit["kind"] == core.KIND_FUNCTION else na,
            "quality_notes_ai": lines_or(ai_notes, "No note")}

def doc_cells(chunk, world, facts, duplicates):
    """The assessment cells of a documentation unit, shown beside the passage on Chunks_Doc."""
    mine, na = facts.get(chunk["ref"], {}), core.NOT_APPLICABLE
    relations = ["%s %s" % (world["links"][(chunk["ref"], ref)]["relation"], ref) for ref in linked(world, chunk["ref"], "C-") if (chunk["ref"], ref) in world["links"]]
    relations += ["%s: %s" % (ref, (world["links"].get((ref, chunk["ref"])) or world["links"].get((chunk["ref"], ref)))["relation"]) for ref in linked(world, chunk["ref"], "M-")]
    relations += ["A second, oppositely framed question named a difference from %s" % r["target_ref"] for r in mine.get("opinions", [])]
    quality = [r["note"] for r in mine.get("unresolved", [])]
    if chunk.get("numbering_reconstructed"):
        quality.append("Section numbering was reconstructed by counting; Word does not store it")
    if chunk["ref"] in duplicates:
        quality.append("The same text also occurs as %s" % duplicates[chunk["ref"]])
    if chunk["kind"] == "Equation" and not chunk["equation"]["readable"]:
        quality.append("The equation could not be read: %s" % chunk["equation"]["not_readable_reason"])
    return {"value_check": lines_or([s for r in mine.get("values", []) for s in r["sentences"]], "No number to compare" if chunk.get("checkable") else na),
            "math_check": lines_or([r["sentence"] for r in mine.get("math", [])], na),
            "logic_consistency": lines_or(relations, na), "quality_notes": lines_or(quality, "No note"),
            "parameter_note_ai": lines_or(["AI-derived alignment of symbols with %s: %s = %s" % (r["target_ref"], ours, theirs)
                                           for r in mine.get("math", []) for ours, theirs, how in r["alignment"] if how.startswith("AI")], na)}

def build_items(raised, world, run_label):
    """One flagged item per unit and category; several observations of one category are listed
    inside one item. Item ids are given once, in a fixed order: concerns, unit, category. An
    item says what was observed and suggests a neutral next step; it rates nothing. Enforces: R1"""
    order = sorted(raised, key=lambda key: (core.CONCERNS.index(concerns_of(world["by_ref"][key[0]])), key[0], core.CATEGORIES.index(key[1])))
    items = []
    for number, (ref, category) in enumerate(order, start=1):
        record, entry = world["by_ref"][ref], raised[(ref, category)]
        canon = [world["by_ref"][r] for r in linked(world, ref, "C-")[:2]]
        if record.get("heading_chain") is not None:
            doc_side, code_side = [record], [world["by_ref"][r] for r in linked(world, ref, "M-")[:1]]
        else:
            code_side, doc_side = [record], [world["by_ref"][r] for r in linked(world, ref, "D-")[:1]]
        name = "`%s`" % record["name"] if record.get("name") else " > ".join(record.get("heading_chain", [])[-1:]) or ref
        path = render_path(record, world)
        items.append(core.FlaggedItem(
            item_id="RI-%s-%05d" % (run_label, number), category=category, concerns=concerns_of(record), unit_refs=tuple([ref] + entry.get("also", [])),
            item="%s: %s %s" % (category, record["kind"].lower(), name),
            observed="\n".join(dict.fromkeys(entry["lines"])) + "\n\nSupporting path:\n" + "\n".join(path),
            methodology_says="\n".join(quote(c["text"][:500], citation(c)) for c in canon),
            code_does="\n".join(quote(c["text"][:700], citation(c)) for c in code_side),
            documentation_says="\n".join(quote(c["text"][:500], citation(c)) for c in doc_side),
            suggested_next_step=NEXT_STEPS[category], evidence_path=tuple(path), from_checks=tuple(c for c in dict.fromkeys(entry["checks"]) if c)))
    return items

def check_identity(units, doc, statuses, items, world, package_info):
    """The four-part identity of plan 2.9 (part 4 is completed by the workbook builder). Any
    violation stops the run and says that this is a defect in the tool. Enforces: R2"""
    defect = "%s This is a defect in the tool, not in the model under review."
    expected = [u["ref"] for u in units] + [c["ref"] for c in doc]
    if sorted(s["unit_ref"] for s in statuses) != sorted(expected):
        raise EngineFault(defect % "Part 1 of the coverage identity does not hold: not every unit has exactly one status record.")
    named = {ref for item in items for ref in item.unit_refs}
    not_clean = {s["unit_ref"] for s in statuses if not s["clean"]}
    if not_clean != (named & set(expected)) or any(not set(item.unit_refs) & set(world["by_ref"]) for item in items):
        raise EngineFault(defect % "Part 2 of the coverage identity does not hold: units that are not clean and flagged items do not match (%s)."
                         % ", ".join(sorted(not_clean ^ (named & set(expected)))[:5]))
    for entry in (package_info[0]["files"] if package_info else []):
        members = [u for u in units if u["file"] == entry["file"]]
        covered = {n for u in members if u.get("lines") for n in range(u["lines"][0], u["lines"][1] + 1)}
        if not members or set(entry.get("nonblank_lines", [])) - covered:
            raise EngineFault(defect % ("Part 3 of the coverage identity does not hold: %s is not fully covered by units." % entry["file"]))

def account_coverage(ctx):
    """Step 15, account-coverage: one status per unit by the ordered rules, the cells of
    the assessment columns, one flagged item per unit and category, the identity, and the
    totals that the workbook builder must reproduce by counting its rows. Enforces: R1, R2"""
    world = load_world(ctx)
    facts, raised, statuses = gather(ctx, world), {}, []
    outcomes = {u["ref"]: model_unit_outcome(u, world, facts, raised) for u in world["units"]}
    for unit in world["units"]:                          # roxygen blocks and help pages take the tracing of what they document
        if outcomes[unit["ref"]][0] != "own checks":
            continue
        _, differs, undecided = outcomes[unit["ref"]]
        target_ref = (unit.get("roxygen") or {}).get("documents_ref") or next(
            (e["target"] for e in world["graph"]["out"].get(unit["ref"], []) if e["kind"] == "documents"), None)
        target = outcomes.get(target_ref)
        if differs or undecided:
            outcomes[unit["ref"]] = (core.ST_DIFFERS if differs else core.ST_UNDECIDED, "its own checks", "see the flagged item(s)")
        elif target is None or target[0] == "own checks":
            outcomes[unit["ref"]] = (core.ST_SUPPORTING, "documents no single object", "documents no single object of the package")
        else:
            outcomes[unit["ref"]] = (target[0], "takes the tracing of what it documents", "as %s, which it documents" % target_ref)
            for (ref, category), entry in raised.items():
                if ref == target_ref and target[0] in core.NOT_CLEAN_STATUSES:
                    entry.setdefault("also", []).append(unit["ref"])
    by_hash, duplicates = {}, {}
    for chunk in world["doc"]:
        if chunk["kind"] == "Paragraph" and chunk["content_hash"] in by_hash:
            duplicates[chunk["ref"]] = by_hash[chunk["content_hash"]]
        by_hash.setdefault(chunk["content_hash"], chunk["ref"])
    doc_outcomes = {c["ref"]: doc_unit_outcome(c, world, facts, raised) for c in world["doc"]}
    for chunk in world["canon"]:                         # what the methodology shows only as a picture is never skipped
        unread = chunk["kind"] == "Equation" and not chunk["equation"]["readable"]
        if unread or chunk["kind"] == "Figure" or chunk.get("not_read_reason"):
            reason = chunk["equation"]["not_readable_reason"] if unread else chunk.get("not_read_reason") or "the content of a figure cannot be read"
            raised[(chunk["ref"], core.CAT_NOT_READ)] = {"checks": [], "lines": [
                "This %s of the methodology could not be read (%s), so nothing can be checked against it by the tool." % (chunk["kind"].lower(), reason)]}
    items = build_items(raised, world, ctx.options["run"]["run_id"].replace("Run_", ""))
    ids_by_unit = {}
    for item in items:
        for ref in item.unit_refs:
            ids_by_unit.setdefault(ref, []).append(item.item_id)
    for record, outcome, corner in [(u, outcomes[u["ref"]], "model") for u in world["units"]] + [(c, doc_outcomes[c["ref"]], "doc") for c in world["doc"]]:
        cells = model_cells(record, world, facts) if corner == "model" else doc_cells(record, world, facts, duplicates)
        statuses.append({"unit_ref": record["ref"], "corner": corner, "status": outcome[0], "clean": outcome[0] in core.CLEAN_STATUSES,
                         "decided_by_rule": outcome[1], "reason_shown": outcome[2], "item_ids": ids_by_unit.get(record["ref"], []), "cells": cells})
    all_units, all_doc = ctx.read("model_units"), ctx.read("chunks_doc")
    check_identity(all_units, all_doc, statuses, items, world, ctx.read("package_info"))
    coverage = {}
    for corner in ("model", "doc"):
        mine = [s for s in statuses if s["corner"] == corner]
        by_status = {status: sum(1 for s in mine if s["status"] == status) for status in core.CLEAN_STATUSES + core.NOT_CLEAN_STATUSES}
        coverage[corner] = {"total": len(mine), "by_status": by_status, "needs_attention": sum(1 for s in mine if not s["clean"])}
    pointed = {t for (s, t), e in world["links"].items() if t.startswith("C-") and e["relation"] in core.LINKING_RELATIONS}
    coverage["canon"] = {"total": len(world["canon"]), "pointed_to": len(pointed)}
    return core.StepResult({"unit_status": statuses, "flagged_items": items, "coverage": [coverage]},
                             {"units": len(statuses), "units needing attention": sum(1 for s in statuses if not s["clean"]), "flagged items": len(items)}, [])
