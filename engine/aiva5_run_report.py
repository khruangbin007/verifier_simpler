"""
AIVA 0.0.1 - aiva5_run_report.py - running a review and reporting it. For Reviewer 5.

WHAT THIS FILE DOES
  It is the only file that touches the outside world: the project folders, the audit
  store, the language model behind chat(), Output.xlsx, the Word report, the record of
  determinations and the verification of an evidence pack. The four bundles below it
  only receive a context and return records.

WHAT IT TAKES IN AND PRODUCES
  In: a Projects folder, a model ID, the three input folders, the analyst's chat()
  function and the live values (endpoint, token, user id) held in memory.
  Out: one run folder = Outputs/ (Output.xlsx, Validation_Report.docx) + _audit/.

WHICH SHEETS SHOW ITS RESULTS
  All eight. This file lays them out from references/workbook_layout.yaml; the words in
  the cells are written by the bundle that knows what they mean.

DESIGN RULES ENFORCED HERE (function names in brackets)
  R3  a failed or rejected call never stops a run           [ask_one, make_asker]
  R5  results never depend on thread timing; replay         [run_batch, replay_chat]
  R6  inputs are never modified; a run writes only inside its own folder [open_run, AuditStore]
  R8  the access token never persists                       [LiveValues.redact, make_settings]
  R10 plain language outward                                [plain_cell]
  R11 only functions named in STEP_FUNCTIONS can run        [load_pipeline]
  R12 build on local disk, copy whole files, few files      [AuditStore.sync, copy_whole]

HOW TO SANITY-CHECK IT
  Run `python -m unittest engine/tests/test_aiva5_run_report.py`. In the notebook run the
  appendix cell "Reviewer 5 sanity check" with the stand-in chat(), then open Output.xlsx:
  eight sheets, yellow reviewer columns, identity rows on Model_Package_Info; search the
  run folder for your test token: it must not occur.
"""
import concurrent.futures
import datetime
import getpass
import gzip
import inspect
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import time
from dataclasses import dataclass, field, replace

import yaml

import aiva0_shared as shared
import aiva1_documents
import aiva2_package
import aiva3_mapping
import aiva4_checks

SKILL_VERSIONS = {"prepare-run": "0.0.1", "confirm-outline": "0.0.1", "await-determinations": "0.0.1",
                  "record-determinations": "0.0.1", "build-report": "0.0.1"}
ENGINE_DIR = os.path.dirname(os.path.abspath(__file__))
REFERENCES_DIR = os.path.join(ENGINE_DIR, "references")

class RunPaused(Exception):
    """The run stopped on purpose and can be resumed. The message tells the person what to do."""

# ---------------------------------------------------------------- settings (allow-list)
DEFAULT_SETTINGS = {
    "k_candidates": 12, "concurrency_limit": 4, "token_cap": 40000, "answer_reserve": 1500,
    "thinking_reserve": 0, "safety_margin": 0.15, "prompt_target_tokens": 6000,
    "max_attempts": 3, "breaker_after_failures": 8, "retry_wait_seconds": 2.0,
    "token_lifetime_minutes": 14.0, "token_wait": "wait", "foreground_minutes": 0.0,
    "sync_every_calls": 100, "llm_file_roll_mb": 25.0, "require_outline_confirmation": True,
    "second_opinion": "unchecked_only", "judge_supporting_code": False,
    "max_parameter_cells": 5000, "max_parameter_columns": 50, "protect_sheets": True,
    "system_prompt_prefix": "", "strip_patterns": [r"(?s)<think>.*?</think>", r"(?s)<thought>.*?</thought>",
                                                   r"(?s)<\|channel\|>thought.*?<\|channel\|>"],
    "numeric_points": 200, "numeric_seed": 20260917, "min_valid_points": 50,
    "relative_tolerance": 1e-9, "trivial_numbers": ["0", "1", "2", "-1", "10", "100"],
    "bm25_k1": 1.2, "bm25_b": 0.75, "anchor_max_share": 0.10, "walk_restart": 0.25,
    "walk_rounds": 30, "heading_anchor_cap": 0.5, "rrf_constant": 60, "reserved_places": 2,
    "max_unit_chars": 3000, "max_passage_chars": 1100, "max_file_mb": 200.0, "reviewer_id": "", "reviewer_role": "", "read_pictures": True, "interpret_code": True,
    "signals": ["fields", "bridge", "references", "anchors", "signatures", "propagation"]}

def make_settings(overrides=None):
    """The settings of a run. Only names on the allow-list above exist, so a new setting
    can never leak into the manifest by default, and no setting can hold the token.
    Enforces: R8"""
    settings = json.loads(json.dumps(DEFAULT_SETTINGS))
    for name, value in (overrides or {}).items():
        if name not in DEFAULT_SETTINGS:
            raise ValueError("'%s' is not a setting AIVA knows" % name)
        settings[name] = value
    return settings

# ---------------------------------------------------------------- paths and project setup
MODEL_ID_RE = re.compile(r"^[A-Za-z0-9_-]{1,24}$")
INPUT_FOLDERS = (
    ("methodology", "1_Methodology", "Put the canonical methodology here (XML, also inside a .txt "
     "file; .mhtml; .docx; .pdf). Several files are read in file-name order."),
    ("package", "2_Model_Package", "Put the single R package tarball (.tar.gz) here."),
    ("documentation", "3_Model_Documentation", "Put the model documentation here (.docx is "
     "preferred; .pdf, .mhtml and XML are read too). Several files are read in file-name order."))
LONGEST_AUDIT_NAME = "package_doc_checks.jsonl"
PATH_BUDGET = 100

@dataclass
class RunPaths:
    """Where one run lives. local_dir is scratch space on the driver; audit_dir is the copy
    in the Workspace that people see."""
    projects_dir: str; model_id: str; project_date: str; project_dir: str; inputs_dir: str
    run_id: str; run_dir: str; outputs_dir: str; audit_dir: str; local_dir: str

def check_model_id(model_id):
    """Return "" when the model ID is usable, otherwise a plain sentence saying why not."""
    if MODEL_ID_RE.match(model_id or ""):
        return ""
    return ("The model ID may hold at most 24 characters from letters, digits, hyphen and "
            "underscore. Please shorten or change '%s'." % model_id)

def setup_project(projects_dir, model_id, project_date=""):
    """Create the project skeleton and say what is still missing. Inputs are never touched."""
    problem = check_model_id(model_id)
    if problem:
        raise ValueError(problem)
    project_date = project_date or datetime.date.today().isoformat()
    project_dir = os.path.join(projects_dir, model_id, project_date)
    missing = []
    for _, folder, readme in INPUT_FOLDERS:
        path = os.path.join(project_dir, "Inputs", folder)
        os.makedirs(path, exist_ok=True)
        readme_path = os.path.join(path, "README.txt")
        if not os.path.exists(readme_path):
            with open(readme_path, "w", encoding="utf-8") as handle:
                handle.write(readme + "\n")
        if not [n for n in os.listdir(path) if n != "README.txt" and not n.startswith(".")]:
            missing.append("Inputs/%s is still empty. %s" % (folder, readme))
    return project_dir, missing

def new_run_id(project_dir, now=None):
    """Run_<date>_<HHMM>, with a letter added when that folder already exists."""
    now = now or datetime.datetime.now()
    base = now.strftime("Run_%Y-%m-%d_%H%M")
    for suffix in [""] + list("bcdefghijklmnopqrstuvwxyz"):
        if not os.path.exists(os.path.join(project_dir, base + suffix)):
            return base + suffix
    raise ValueError("Too many runs were started in the same minute; please wait a minute.")

def pick_scratch_root(preferred=""):
    """The first folder on the driver AIVA can actually write in, tried in order.

    A run is built on the driver's own disk and copied whole into the Workspace afterwards
    (R12), so this folder is needed before anything else can happen. A cluster is shared, and
    a scratch folder made by one user cannot be written into by another; the folders tried
    here therefore carry the user's own name. Writing is tested, not assumed, because a
    folder can exist and still refuse."""
    user = re.sub(r"[^A-Za-z0-9_.-]", "_", getpass.getuser() or "user")
    refused = []
    for root in ([preferred] if preferred else []) + [
            os.path.join(tempfile.gettempdir(), "aiva_scratch_" + user),
            os.path.join("/local_disk0", "aiva_scratch_" + user)]:
        try:
            os.makedirs(root, exist_ok=True)
            probe = os.path.join(root, ".aiva_write_test")
            with open(probe, "w") as handle:
                handle.write("x")
            os.remove(probe)
            return root
        except OSError as problem:
            refused.append("%s (%s)" % (root, problem.strerror or problem))
    raise PermissionError(
        "AIVA builds each run on the driver's own disk and copies the finished files into the "
        "Workspace, so it needs one folder it may write in. These were refused: %s. Put a "
        "folder you can write in into the 'Scratch folder' widget, or ask for one on this "
        "cluster." % "; ".join(refused))

def open_run(projects_dir, model_id, project_date="", run_id="", scratch_root="", now=None):
    """Create or re-open a run folder and its local scratch folder. Enforces: R6"""
    project_dir, _ = setup_project(projects_dir, model_id, project_date)
    project_date = os.path.basename(project_dir)
    run_id = run_id or new_run_id(project_dir, now)
    run_dir = os.path.join(project_dir, run_id)
    scratch_root = pick_scratch_root(scratch_root)
    place = shared.sha256_text(os.path.abspath(run_dir))[:8]       # two Projects folders never share scratch space
    local_dir = os.path.join(scratch_root, "%s_%s_%s_%s" % (model_id, project_date, run_id, place))
    paths = RunPaths(projects_dir, model_id, project_date, project_dir,
                     os.path.join(project_dir, "Inputs"), run_id, run_dir,
                     os.path.join(run_dir, "Outputs"), os.path.join(run_dir, "_audit"), local_dir)
    longest = os.path.join(paths.audit_dir, LONGEST_AUDIT_NAME)
    relative = os.path.relpath(longest, os.path.dirname(os.path.abspath(projects_dir)))
    if len(relative) > PATH_BUDGET:
        raise ValueError("The folder path is %d characters long and the limit is %d, so that "
                         "Excel can still open downloaded files. Please use a shorter model ID "
                         "or Projects folder." % (len(relative), PATH_BUDGET))
    for folder in (paths.outputs_dir, paths.audit_dir, paths.local_dir):
        os.makedirs(folder, exist_ok=True)
    return paths

def list_input_files(paths):
    """The input files of a project, by corner, in file-name order."""
    inputs = {"glossary": None, "tag_rules": None}
    for corner, folder, _ in INPUT_FOLDERS:
        path = os.path.join(paths.inputs_dir, folder)
        names = sorted(n for n in os.listdir(path) if n != "README.txt" and not n.startswith("."))
        inputs[corner] = [os.path.join(path, n) for n in names]
    for key, name in (("glossary", "glossary.xlsx"), ("tag_rules", "tag_rules.yaml")):
        if os.path.exists(os.path.join(paths.inputs_dir, name)):
            inputs[key] = os.path.join(paths.inputs_dir, name)
    return inputs

# ---------------------------------------------------------------- live values (the token)
@dataclass
class LiveValues:
    """The three values chat() reads when it is CALLED: endpoint, token, user id. They live
    in memory only. `generation` goes up whenever a different token arrives, which is how
    paused workers learn that a fresh one was pasted. Enforces: R8"""
    values: dict = field(default_factory=dict); generation: int = 0; set_at: float = 0.0
    recent_tokens: list = field(default_factory=list)
    lock: threading.Lock = field(default_factory=threading.Lock)

    def update(self, llm_endpoint="", llm_token="", llm_user_id=""):
        """Take the latest widget values; a different token starts a new generation."""
        with self.lock:
            if llm_token and llm_token != self.values.get("llm_token"):
                self.generation += 1
                self.set_at = time.time()
                self.recent_tokens = ([llm_token] + self.recent_tokens)[:5]
            self.values = {"llm_endpoint": llm_endpoint, "llm_token": llm_token, "llm_user_id": llm_user_id}

    def get(self, name):
        """One live value, read at the moment chat() is called."""
        with self.lock:
            return self.values.get(name, "")

    def token_age_minutes(self):
        """Minutes since the current token was pasted."""
        return (time.time() - self.set_at) / 60.0 if self.set_at else 0.0

    def redact(self, text):
        """Remove the current and recent token strings from any text before it is kept."""
        for token in list(self.recent_tokens):
            if token:
                text = text.replace(token, "[token removed]")
        return text

# ---------------------------------------------------------------- the audit store
AUDIT_OBJECTS = ("run_manifest", "package_info", "coverage")

def copy_whole(source, target):
    """Copy one whole file: to a temporary name, then replace; a plain copy if the file
    system does not support replace (probe P-5). Enforces: R12"""
    os.makedirs(os.path.dirname(target), exist_ok=True)
    temporary = target + ".part"
    shutil.copyfile(source, temporary)
    try:
        os.replace(temporary, target)
    except OSError:
        shutil.copyfile(source, target)
        os.remove(temporary)

@dataclass
class AuditStore:
    """All audit files are read and written on local disk; sync() copies changed files
    whole into _audit/ in the Workspace. JSON Lines, one record per line, sorted keys."""
    local_dir: str; remote_dir: str; roll_bytes: int = 25 * 1024 * 1024
    synced: dict = field(default_factory=dict)

    def path(self, kind):
        """The local path of one kind of audit file."""
        extension = ".json" if kind in AUDIT_OBJECTS else ".jsonl"
        return os.path.join(self.local_dir, kind + extension)

    def read(self, kind):
        """Every record of one kind, in the order written."""
        if kind == "llm_calls":
            return self.read_calls()
        if not os.path.exists(self.path(kind)):
            return []
        with open(self.path(kind), encoding="utf-8") as handle:
            if kind in AUDIT_OBJECTS:
                return [json.load(handle)]
            return [json.loads(line) for line in handle if line.strip()]

    def append(self, kind, records):
        """Append records of one kind; the three single-object kinds are rewritten whole."""
        if kind in AUDIT_OBJECTS:
            with open(self.path(kind), "w", encoding="utf-8") as handle:
                handle.write(json.dumps(shared.to_plain(records[-1]), sort_keys=True, indent=1, ensure_ascii=False))
            return
        with open(self.path(kind), "a", encoding="utf-8") as handle:
            for record in records:
                handle.write(shared.canonical_json(record) + "\n")

    def call_files(self):
        """The gzip files of call records, in order."""
        names = sorted(n for n in os.listdir(self.local_dir) if re.fullmatch(r"llm_calls_\d{3}\.jsonl\.gz", n))
        return [os.path.join(self.local_dir, n) for n in names]

    def append_calls(self, records):
        """Call records go to gzip files that roll over at the size limit. The gzip header
        carries no time stamp, so the same records always give the same bytes."""
        files = self.call_files()
        current = files[-1] if files else os.path.join(self.local_dir, "llm_calls_001.jsonl.gz")
        if files and os.path.getsize(current) >= self.roll_bytes:
            current = os.path.join(self.local_dir, "llm_calls_%03d.jsonl.gz" % (len(files) + 1))
        with open(current, "ab") as raw:
            with gzip.GzipFile(filename="", mode="ab", fileobj=raw, mtime=0) as handle:
                for record in records:
                    handle.write((shared.canonical_json(record) + "\n").encode("utf-8"))

    def read_calls(self):
        """Every call record of the run, in the order written."""
        records = []
        for path in self.call_files():
            with gzip.open(path, "rt", encoding="utf-8") as handle:
                records.extend(json.loads(line) for line in handle if line.strip())
        return records

    def sync(self):
        """Copy every file that changed since the last sync, whole. Returns what was copied."""
        copied = []
        for folder, _, names in os.walk(self.local_dir):
            for name in sorted(names):
                source = os.path.join(folder, name)
                relative = os.path.relpath(source, self.local_dir)
                stamp = (os.path.getsize(source), os.stat(source).st_mtime_ns)
                if relative.startswith("work") or self.synced.get(relative) == stamp:
                    continue
                copy_whole(source, os.path.join(self.remote_dir, relative))
                self.synced[relative] = stamp
                copied.append(relative)
        return copied

    def restore(self):
        """On resume: copy the run folder's audit files back to local disk first."""
        if os.path.isdir(self.remote_dir) and not os.path.exists(self.path("run_manifest")):
            shutil.copytree(self.remote_dir, self.local_dir, dirs_exist_ok=True)

def open_store(paths, settings):
    """The audit store of a run, restored from the run folder when local scratch is empty."""
    store = AuditStore(paths.local_dir, paths.audit_dir, int(settings["llm_file_roll_mb"] * 1024 * 1024))
    store.restore()
    return store

# ---------------------------------------------------------------- the wrapper around chat()
AUTH_WORDS = ("401", "403", "unauthor", "expired", "forbidden", "invalid token", "authentication", "credential")
OVERLOAD_WORDS = ("429", "502", "503", "504", "overload", "rate limit", "too many", "timeout", "timed out", "busy")

def classify_failure(text):
    """Sort a failure into a class by the words it contains (probe P-13 matches these lists to
    what the real gateway returns). call_chat adds one more class of its own, "truncated"."""
    lowered = (text or "").lower()
    if any(word in lowered for word in AUTH_WORDS):
        return "authentication"
    if any(word in lowered for word in OVERLOAD_WORDS):
        return "overload"
    return "other"

class ChatOutcome(tuple):
    """What one call to chat() came to. It unpacks as the three values the rest of AIVA has
    always read - (answer text or None, failure class or "", what was seen with tokens
    removed) - and carries the gateway's own record of the call as .meta, so that adding
    metadata did not change a single existing call site."""

    def __new__(cls, answer, failure, seen, meta=None):
        outcome = super().__new__(cls, (answer, failure, seen))
        outcome.meta = dict(meta or {})
        return outcome

    answer = property(lambda self: self[0])
    failure = property(lambda self: self[1])
    seen = property(lambda self: self[2])

# Fields of the gateway's reply that are worth keeping in the audit record. The reply also
# echoes the prompts back (query, defaultprompt, source) and those are dropped: AIVA already
# stores the prompts it sent, and an echo would double the size of every call record.
META_FIELDS = ("chat_id", "thread_id", "prompt_id", "datetime", "user_id", "intent")
RESPONSE_META_FIELDS = ("id", "model", "created", "system_fingerprint", "service_tier")
USAGE_FIELDS = ("prompt_tokens", "completion_tokens", "total_tokens")

def first_choice(response):
    """The first choice of an OpenAI-shaped reply, or an empty dictionary."""
    nested = response.get("response") if isinstance(response.get("response"), dict) else {}
    choices = nested.get("choices")
    return choices[0] if isinstance(choices, list) and choices and isinstance(choices[0], dict) else {}

def chat_metadata(response):
    """What the gateway said about the call itself: which conversation it belonged to, which
    model answered, how many tokens it took and why it stopped. Kept for the audit record so
    that a run can be accounted for afterwards; none of it reaches a status or a check."""
    meta = {name: response[name] for name in META_FIELDS if response.get(name) not in (None, "")}
    nested = response.get("response") if isinstance(response.get("response"), dict) else {}
    meta.update({name: nested[name] for name in RESPONSE_META_FIELDS if nested.get(name) not in (None, "")})
    usage = nested.get("usage") if isinstance(nested.get("usage"), dict) else {}
    meta.update({name: usage[name] for name in USAGE_FIELDS if isinstance(usage.get(name), int)})
    choice = first_choice(response)
    if choice.get("finish_reason"):
        meta["finish_reason"] = choice["finish_reason"]
    return meta

def answer_text_of(response):
    """The generated text. The gateway puts it under "answer"; if that key is missing or
    empty, the same text is read from the OpenAI-shaped part of the reply, so that a
    gateway that only fills one of the two is still usable."""
    if isinstance(response.get("answer"), str) and response["answer"].strip():
        return response["answer"]
    content = first_choice(response).get("message", {})
    content = content.get("content") if isinstance(content, dict) else None
    return content if isinstance(content, str) and content.strip() else None

def summarise_response(response, live):
    """A failed reply, shortened for the audit record: the fields that say what went wrong,
    without the echoed prompts. Falls back to the whole reply when it has no known shape."""
    if not isinstance(response, dict):
        return live.redact(json.dumps(response, default=str)[:2000])
    keep = ("status", "code", "message", "detail", "reason", "answer", "finish_reason")
    kept = {name: response[name] for name in keep if name in response}
    kept.update(chat_metadata(response))
    if not kept:
        kept = {name: value for name, value in response.items()
                if name not in ("query", "defaultprompt", "source", "documents", "history")}
    return live.redact(json.dumps(kept, default=str)[:2000])

def accepts_history(chat):
    """Whether the analyst's chat() takes the third `history` argument. The real gateway cell
    has the signature chat(SystemPrompt, MainPrompt, history=[]); the stand-in and the older
    cell take two arguments. AIVA works with either and never depends on which."""
    try:
        parameters = inspect.signature(chat).parameters
    except (TypeError, ValueError):
        return False
    if any(p.kind == inspect.Parameter.VAR_KEYWORD for p in parameters.values()):
        return "history" in parameters
    positional = [p for p in parameters.values()
                  if p.kind in (inspect.Parameter.POSITIONAL_ONLY, inspect.Parameter.POSITIONAL_OR_KEYWORD)]
    return "history" in parameters or len(positional) >= 3

def call_chat(chat, system_prompt, main_prompt, live):
    """Call chat() once and bring every way a failure can surface (an exception; a dictionary
    that carries a status or code; a dictionary with no answer in either place; an answer the
    model was cut off in the middle of) into one shape. Returns a ChatOutcome, which unpacks
    as (answer, failure, seen).

    History is always sent empty. Every question AIVA asks is self-contained and is asked in
    its own call, so nothing the model said about an earlier unit may colour the next one;
    an empty history is also what makes a question repeatable. Enforces: R5"""
    try:
        response = chat(system_prompt, main_prompt, []) if accepts_history(chat) else chat(system_prompt, main_prompt)
    except Exception as problem:                     # shape 1: chat() raised
        seen = live.redact("%s: %s" % (type(problem).__name__, problem))
        return ChatOutcome(None, classify_failure(seen), seen)
    if isinstance(response, str):                    # a gateway that returns the text alone
        response = {"answer": response}
    if not isinstance(response, dict):
        seen = live.redact(json.dumps(response, default=str)[:2000])
        return ChatOutcome(None, classify_failure(seen), seen)
    meta, text = chat_metadata(response), answer_text_of(response)
    if text is None:                                 # shapes 2 and 3: a reply with no answer
        seen = summarise_response(response, live)
        return ChatOutcome(None, classify_failure(seen), seen, meta)
    if meta.get("finish_reason") == "length":        # shape 4: the answer stops in mid-air
        seen = "the model was cut off at the token limit; the answer is incomplete"
        return ChatOutcome(None, "truncated", seen, meta)
    return ChatOutcome(live.redact(text), "", "", meta)

@dataclass
class AskState:
    """What the workers of one run share: the pause and stop switches, the consecutive-
    failure count of the circuit breaker, and the generation of the token that failed."""
    control: dict = field(default_factory=lambda: {"pause": False, "stop": False})
    consecutive_failures: int = 0; failed_generation: int = -1; waiting_for_token: bool = False
    deadline: float = 0.0; calls_made: int = 0
    lock: threading.Lock = field(default_factory=threading.Lock)

def wait_until_allowed(state, live, settings, sleep):
    """Called before every call. Blocks while the run is paused or while a fresh token is
    awaited; raises RunPaused when the run must stop by itself (mode C, stop switch)."""
    while True:
        if state.control.get("stop"):
            raise RunPaused("The run was stopped on request. Run the cell again to resume.")
        if state.deadline and time.time() > state.deadline:
            raise RunPaused("The time box of this foreground run is over. Paste a fresh token "
                            "and run the cell again; no call will be repeated.")
        too_old = live.set_at and live.token_age_minutes() >= settings["token_lifetime_minutes"]
        needs_token = state.failed_generation == live.generation or too_old
        state.waiting_for_token = bool(needs_token)
        if not needs_token and not state.control.get("pause"):
            return
        if needs_token and settings["token_wait"] == "stop":
            raise RunPaused("Paste a fresh token, then run the cell again; no call will be repeated.")
        sleep(0.05)

def ask_one(question, chat, live, settings, validate, state, sleep):
    """Ask one question until it has a final outcome. Every attempt is recorded. An
    authentication-like failure pauses all workers until a fresh token arrives and does
    not count against the attempts. Enforces: R3"""
    records, attempt = [], 0
    system_prompt = settings["system_prompt_prefix"] + question["system_prompt"]
    while True:
        wait_until_allowed(state, live, settings, sleep)
        generation = live.generation
        started, clock = datetime.datetime.now().isoformat(timespec="seconds"), time.time()
        with state.lock:
            state.calls_made += 1
        outcome_of_call = call_chat(chat, system_prompt, question["main_prompt"], live)
        answer_text, failure, seen = outcome_of_call
        attempt += 1
        outcome, answer = ("failed: " + failure, None) if failure else validate(question, answer_text)
        with state.lock:
            state.consecutive_failures = state.consecutive_failures + 1 if failure else 0
            breaker_open = state.consecutive_failures >= settings["breaker_after_failures"]
            if failure == "authentication":
                state.failed_generation = generation
        final = outcome == "accepted" or (attempt >= settings["max_attempts"] and failure != "authentication")
        records.append({
            "question_id": question["question_id"], "question_type": question["question_type"],
            "unit_ref": question.get("unit_ref", ""), "attempt": len(records) + 1,
            "letters": question.get("letters", {}), "planted": question.get("planted", []),
            "prompt_hash": question["question_id"], "response_hash": shared.sha256_text(answer_text or seen),
            "system_prompt": live.redact(system_prompt), "main_prompt": live.redact(question["main_prompt"]),
            "response_text": answer_text if answer_text is not None else seen, "outcome": outcome,
            "answer": answer, "final": final, "started_at": started,
            "seconds": round(time.time() - clock, 3), "estimated_tokens": question.get("estimated_tokens", 0),
            "gateway": outcome_of_call.meta})
        if final:
            return records
        if failure == "authentication":
            attempt -= 1
        elif breaker_open:
            raise RunPaused("Many calls in a row failed, so the run paused itself to avoid wasting "
                            "calls. Check the gateway, then run the cell again to resume.")
        elif failure:
            sleep(settings["retry_wait_seconds"] * attempt)

def run_batch(batch, chat, live, settings, validate, state, sleep):
    """Ask one batch of questions with a thread pool. Results are gathered in the order of
    the questions, never in the order threads finished. Enforces: R5"""
    finished, paused = {}, None
    with concurrent.futures.ThreadPoolExecutor(max_workers=max(1, int(settings["concurrency_limit"]))) as pool:
        futures = [(q["question_id"], pool.submit(ask_one, q, chat, live, settings, validate, state, sleep))
                   for q in batch]
        for question_id, future in futures:
            try:
                finished[question_id] = future.result()
            except RunPaused as pause:
                paused = paused or pause
                state.control["stop"] = True           # let the other workers end after their call
    if paused:
        state.control["stop"] = False
    return finished, paused

def make_asker(chat, live, store, settings, validate, state=None, sleep=time.sleep):
    """Build ask(), the only place chat() is ever called. ask(questions) returns a dictionary
    question id -> final record. A question that already has a final record in this run is
    never asked again, which is what makes resume safe. Enforces: R3, R5"""
    state = state or AskState()
    if settings["foreground_minutes"]:
        state.deadline = time.time() + 60.0 * settings["foreground_minutes"]

    def ask(questions):
        done = {r["question_id"]: r for r in store.read_calls() if r.get("final")}
        unique = {q["question_id"]: q for q in questions}
        todo = [unique[qid] for qid in sorted(unique) if qid not in done]
        size = max(1, int(settings["sync_every_calls"]))
        for start in range(0, len(todo), size):
            finished, paused = run_batch(todo[start:start + size], chat, live, settings, validate, state, sleep)
            records = [r for qid in sorted(finished) for r in finished[qid]]
            store.append_calls(records)
            store.sync()
            done.update({r["question_id"]: r for r in records if r["final"]})
            if paused:
                raise paused
        return {qid: done[qid] for qid in unique if qid in done}
    ask.state = state
    return ask

def replay_chat(call_records):
    """A chat() that answers from the recorded answers of an earlier run. Running the
    pipeline with it must reproduce the same graph version id, statuses and items, which
    is what reproducibility means for a model that samples its answers. Enforces: R5"""
    recorded = {}
    for record in call_records:
        recorded.setdefault((record["system_prompt"], record["main_prompt"]), []).append(record["response_text"])
    position = {}
    lock = threading.Lock()

    def chat(system_prompt, main_prompt):
        key = (system_prompt, main_prompt)
        with lock:
            index = position.get(key, 0)
            position[key] = index + 1
        answers = recorded.get(key)
        if not answers:
            return {"status": "no recorded answer for this prompt"}
        return {"answer": answers[min(index, len(answers) - 1)]}
    return chat

# ---------------------------------------------------------------- the pipeline runner
CHAT_STEPS = ("interpret-code", "judge-links", "check-mathematics", "check-values", "check-rules")
REPEATABLE_STEPS = ("record-determinations", "build-report")
HUMAN_MESSAGES = {
    "confirm-outline": "Waiting for a person: check Level and Section (heading chain) on Chunks_Canon "
                       "against the methodology's own outline, then run cell 10 to confirm.",
    "await-determinations": "Waiting for a person: download Output.xlsx from Outputs/, fill the four "
                            "yellow columns on Flagged_Items, upload it into Outputs/ and run cell 14."}

def load_pipeline(engine_dir=ENGINE_DIR):
    """Read pipeline.yaml and refuse anything that is not a known skill and function. Enforces: R11"""
    with open(os.path.join(engine_dir, "pipeline.yaml"), encoding="utf-8") as handle:
        pipeline = yaml.safe_load(handle)
    declared = dict(SKILL_VERSIONS)
    for bundle in (aiva1_documents, aiva2_package, aiva3_mapping, aiva4_checks):
        declared.update(bundle.SKILL_VERSIONS)
    for step in pipeline["steps"]:
        if not step.get("human") and step["function"] not in STEP_FUNCTIONS:
            raise ValueError("pipeline.yaml names '%s', which is not in the list of functions "
                             "allowed to run" % step["function"])
        with open(os.path.join(engine_dir, "skills", step["skill"], "SKILL.md"), encoding="utf-8") as handle:
            front_matter = yaml.safe_load(handle.read().split("---")[1])
        step["skill_version"] = str(front_matter["metadata"]["version"])
        if declared.get(step["skill"]) != step["skill_version"]:
            raise ValueError("The contract of skill '%s' is at version %s but the code declares %s: "
                             "contract and code are out of step, so the run does not start."
                             % (step["skill"], step["skill_version"], declared.get(step["skill"])))
    return pipeline

def update_manifest(store, changes):
    """Change fields of the run manifest and write it back."""
    manifest = (store.read("run_manifest") or [{}])[0]
    manifest.update(changes)
    store.append("run_manifest", [manifest])
    return manifest

def confirm_outline(paths, settings, reviewer):
    """Record that a person has checked the outline of the methodology (notebook cell 10)."""
    store = open_store(paths, settings)
    update_manifest(store, {"outline_confirmed_by": reviewer or "not named",
                            "outline_confirmed_at": datetime.datetime.now().isoformat(timespec="seconds")})
    store.sync()
    return "Recorded: the outline was confirmed by %s." % (reviewer or "a person who gave no name")

def human_step_open(step, store, settings, determinations):
    """Is this human step still waiting for its person?"""
    if step["skill"] == "confirm-outline":
        manifest = (store.read("run_manifest") or [{}])[0]
        return settings["require_outline_confirmation"] and not manifest.get("outline_confirmed_by")
    return not determinations and not store.read("determinations")

def run_pipeline(paths, settings, chat=None, live=None, determinations=False, stop_after="",
                 state=None, keep_alive=None, sleep=time.sleep):
    """Run, or resume, the pipeline. Each finished step leaves a step record; called again,
    the run continues at the first step without one. A human step stops the run and says
    what the person should do. Returns {"state", "message", "steps_run"}."""
    store = open_store(paths, settings)
    pipeline = load_pipeline()
    live = live or LiveValues()
    done = {record["step_id"] for record in store.read("step_records")}
    steps_run = []
    if determinations and "15" not in done:
        return {"state": "not ready", "steps_run": [], "message":
                "Determinations can be recorded once the run has reached its flagged items."}
    for step in pipeline["steps"]:
        if stop_after and step["id"] > stop_after:
            break
        if step["skill"] in REPEATABLE_STEPS:
            if not determinations:
                continue                     # these two run each time the determinations cell is run
        elif step["id"] in done:
            continue
        if step.get("human"):
            if human_step_open(step, store, settings, determinations):
                rebuild_outputs(store, paths, settings, HUMAN_MESSAGES[step["skill"]])
                return {"state": "waiting for a person", "message": HUMAN_MESSAGES[step["skill"]], "steps_run": steps_run}
            record_step(store, step, shared.StepResult(), 0.0)
            continue
        try:
            run_step(step, store, paths, settings, chat, live, state, sleep)
        except RunPaused as pause:
            store.sync()
            return {"state": "paused", "message": str(pause), "steps_run": steps_run}
        steps_run.append(step["skill"])
        if keep_alive:
            keep_alive()
        if stop_after and step["id"] == stop_after:
            break
    message = "Every step has run." if not stop_after else "Stopped after step %s as asked." % stop_after
    rebuild_outputs(store, paths, settings, message)
    return {"state": "finished", "message": message, "steps_run": steps_run}

def run_step(step, store, paths, settings, chat, live, state, sleep):
    """Build the context, call the step function, write what it returns, record the step,
    rebuild the outputs and sync. Steps never touch the store themselves."""
    notes = []
    ask = None
    if step["skill"] in CHAT_STEPS and chat is not None:
        ask = make_asker(chat, live, store, settings, aiva3_mapping.validate_answer, state, sleep)
    provenance = shared.Provenance(paths.run_id, step["id"], step["skill"], step["skill_version"],
                                   created_at=datetime.datetime.now().isoformat(timespec="seconds"))
    options = dict(step.get("with") or {})
    options.update({"inputs": list_input_files(paths), "references_dir": REFERENCES_DIR, "paths": paths,
                    "run": {"model_id": paths.model_id, "project_date": paths.project_date, "run_id": paths.run_id}})
    work_dir = os.path.join(paths.local_dir, "work")
    os.makedirs(work_dir, exist_ok=True)
    context = shared.StepContext(settings, options, store.read, ask, work_dir, notes.append, provenance)
    function = STEP_FUNCTIONS[step["function"]]
    started = time.time()
    result = function(context)
    for kind in sorted(result.records):
        store.append(kind, result.records[kind])
    result.messages = list(result.messages) + notes
    record_step(store, step, result, time.time() - started)
    rebuild_outputs(store, paths, settings, "")
    store.sync()

def record_step(store, step, result, seconds):
    """Leave the step record that makes a finished step visible and resume possible."""
    produced = {kind: len(records) for kind, records in sorted(result.records.items())}
    store.append("step_records", [{
        "step_id": step["id"], "skill": step["skill"], "skill_version": step["skill_version"],
        "function": step.get("function", "a person"), "produced": produced, "counts": result.counts,
        "messages": result.messages, "finished_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "seconds": round(seconds, 3)}])

def log_line(store, text):
    """Technical text (exception messages, Python names) belongs in run_log.txt only."""
    with open(os.path.join(store.local_dir, "run_log.txt"), "a", encoding="utf-8") as handle:
        handle.write("%s  %s\n" % (datetime.datetime.now().isoformat(timespec="seconds"), text))

# ---------------------------------------------------------------- step 01: prepare-run
PACKAGES_RECORDED = ("PyYAML", "openpyxl", "python-docx", "numpy", "scipy", "sympy", "rdata",
                     "pdfplumber", "pypdf", "pyreadr")

def fingerprint_file(path, corner, inputs_dir):
    """Name, corner, size, SHA-256 and content identifier of one input file."""
    with open(path, "rb") as handle:
        data = handle.read()
    return {"corner": corner, "file": os.path.relpath(path, inputs_dir).replace(os.sep, "/"),
            "bytes": len(data), "sha256": shared.sha256_bytes(data), "swhid": shared.swhid_content(data)}

def engine_file_hashes():
    """SHA-256 of every file that makes up the engine (code, pipeline, skills, references), so
    that an evidence pack names exactly the code that produced it (the same list as in
    docs/release_manifest.json)."""
    import glob
    found = {}
    for pattern in ("*.py", "pipeline.yaml", "requirements.txt", "skills/*/SKILL.md", "references/**/*"):
        for path in sorted(glob.glob(os.path.join(ENGINE_DIR, pattern), recursive=True)):
            if os.path.isfile(path):
                found["engine/" + os.path.relpath(path, ENGINE_DIR).replace(os.sep, "/")] = file_sha256(path)
    return found

def prepare_run(ctx):
    """Step 01. Fingerprint every input, record the environment and the allow-listed
    settings, and say what changed since the previous run of the same project."""
    import importlib.metadata
    paths, inputs = ctx.options["paths"], ctx.options["inputs"]
    fingerprints = []
    for corner, _, _ in INPUT_FOLDERS:
        fingerprints.extend(fingerprint_file(path, corner, paths.inputs_dir) for path in inputs[corner])
    for key in ("glossary", "tag_rules"):
        if inputs[key]:
            fingerprints.append(fingerprint_file(inputs[key], key, paths.inputs_dir))
    versions = {"python": sys.version.split()[0]}
    for name in PACKAGES_RECORDED:
        try:
            versions[name] = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:
            versions[name] = "not installed"
    previous, changes = previous_run_inputs(paths), []
    if previous is not None:
        before = {f["file"]: f["sha256"] for f in previous["inputs"]}
        now = {f["file"]: f["sha256"] for f in fingerprints}
        changes = ["%s: changed" % n for n in sorted(now) if n in before and before[n] != now[n]]
        changes += ["%s: new" % n for n in sorted(now) if n not in before]
        changes += ["%s: no longer present" % n for n in sorted(before) if n not in now]
    manifest = dict(ctx.options["run"], engine_version=shared.ENGINE_VERSION, versions=versions, engine_files=engine_file_hashes(),
                    settings=ctx.settings, inputs=fingerprints, changes_since_previous_run=changes,
                    previous_run=previous["run_id"] if previous else "",
                    started_at=datetime.datetime.now().isoformat(timespec="seconds"))
    return shared.StepResult({"run_manifest": [manifest]}, {"input files": len(fingerprints)},
                             ["%d input files fingerprinted." % len(fingerprints)])

def previous_run_inputs(paths):
    """The manifest of the latest earlier run of this project, or None."""
    runs = sorted(n for n in os.listdir(paths.project_dir) if n.startswith("Run_") and n < paths.run_id)
    for run in reversed(runs):
        manifest_path = os.path.join(paths.project_dir, run, "_audit", "run_manifest.json")
        if os.path.exists(manifest_path):
            with open(manifest_path, encoding="utf-8") as handle:
                return json.load(handle)
    return None

# ---------------------------------------------------------------- Output.xlsx
CELL_WITHHELD = "This text could not be shown in plain words; the technical text is in the audit records."
CUT_NOTE = " ... (cut here; the full text is in the audit files)"
PYTHON_TRACES = re.compile(r"Traceback|\b\w+(Err" r"or|Exception)\b|<class |object at 0x|\bnan\b|\baiva\d_\w+|"
                           r"[:=(\[]\s*None\b|\{'|\['|^None$")
COVERAGE_ROWS = (("model", "Model units (one row each on Mapping_Model_to_Canon_and_Doc)"),
                 ("doc", "Documentation units (one row each on Mapping_Doc_to_Canon_and_Model)"),
                 ("canon", "Methodology passages (for information only)"))
NEEDS_ATTENTION_MEANS = ("Needs attention = units whose status is not clean; each has an entry on Flagged_Items.")

def quoted(text, citation=""):
    """Text taken from an input or from the AI is always shown visibly quoted, with its
    citation. The wording rules apply to AIVA's own words, not to quotations. Enforces: R1"""
    inner = (text or "").replace("\u201c", '"').replace("\u201d", '"')
    return "\u201c%s\u201d%s" % (inner, " (%s)" % citation if citation else "")

def plain_cell(value, input_text, store):
    """The last gate before a cell is written. AIVA's own words must be free of Python
    traces and of words the wording rule rejects; otherwise the text goes to run_log.txt
    and the cell gets one fixed, plain sentence. Enforces: R1, R10"""
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value == int(value) else shared.plain_number(value)
    text = str(value)
    if not input_text:
        own_words = re.sub(r"\u201c.*?\u201d", "", text, flags=re.S)
        if PYTHON_TRACES.search(own_words) or shared.has_banned_wording(own_words):
            log_line(store, "cell text withheld: " + text)
            text = CELL_WITHHELD
    return text if len(text) <= 32000 else text[:31900] + CUT_NOTE

def lines_by_ref(pairs):
    """Several values in one cell: each on its own line, prefixed with its reference."""
    return "\n".join("%s: %s" % (ref, value) for ref, value in pairs if value not in (None, ""))

def texts_by_ref(pairs):
    """Several quoted texts in one cell, separated by a line of dashes."""
    return "\n----------\n".join("%s: %s" % (ref, text) for ref, text in pairs)

def run_identity(store, paths):
    """What ties a workbook to its run: also written into the workbook's properties."""
    items = store.read("flagged_items")
    ledger, decisions = store.read("graph_ledger"), store.read("determinations")
    return {"model_id": paths.model_id, "date_initiated": paths.project_date, "run_id": paths.run_id,
            "engine_version": shared.ENGINE_VERSION,
            "item_list_hash": shared.sha256_text("\n".join(i["item_id"] for i in items)) if items else "",
            "graph_version_id": "G-" + shared.chain_head(ledger)[:12] if ledger else "",
            "determinations_fingerprint": "DR-" + shared.chain_head(decisions)[:12] if decisions else ""}

def rows_package_info(store, paths, settings, progress):
    """The rows of Model_Package_Info: identity, inputs, what was read, repairs, how values and formulas are compared, AI calls."""
    identity, rows = run_identity(store, paths), []
    def add(group, item, value):
        rows.append({"group": group, "item": item, "value": value})
    add("Identity", "Model ID", identity["model_id"])
    add("Identity", "Date initiated", identity["date_initiated"])
    add("Identity", "Run", identity["run_id"])
    add("Identity", "Run progress", progress)
    add("Identity", "Engine version", identity["engine_version"])
    add("Identity", "Graph version id", identity["graph_version_id"] or shared.NOT_RUN_YET)
    add("Identity", "Determinations record fingerprint", identity["determinations_fingerprint"] or "No determination recorded yet")
    add("Identity", "Item list fingerprint", identity["item_list_hash"] or shared.NOT_RUN_YET)
    manifest = (store.read("run_manifest") or [{}])[0]
    if manifest.get("engine_files"):
        add("Identity", "Engine files fingerprint", shared.sha256_text(shared.canonical_json(manifest["engine_files"]))[:16] +
            " (%d files; the full list is in the run manifest)" % len(manifest["engine_files"]))
    for entry in manifest.get("inputs", []):
        add("Inputs", entry["file"], "SHA-256 %s (%d bytes)" % (entry["sha256"], entry["bytes"]))
    for change in manifest.get("changes_since_previous_run", []):
        add("Inputs", "Changed since run %s" % manifest.get("previous_run", ""), change)
    if manifest.get("outline_confirmed_by"):
        add("Review", "Methodology outline confirmed by", manifest["outline_confirmed_by"])
    for info in store.read("package_info"):
        for row in info.get("rows", []):
            add(row["group"], row["item"], row["value"])
    for row in store.read("info_rows"):
        add(row["group"], row["item"], row["value"])
    repairs = {}
    for repair in store.read("read_repairs"):
        repairs.setdefault(repair["file"], []).append(repair["kind"])
    for name in sorted(repairs):
        kinds = sorted(set(repairs[name]))
        add("Repairs made while reading", name, "; ".join("%s (%d)" % (k, repairs[name].count(k)) for k in kinds))
    add("How values are compared", "The value-comparison rule", aiva4_checks.VALUE_RULE_TEXT)
    add("How values are compared", "Numbers treated as trivial", ", ".join(settings["trivial_numbers"]))
    add("How formulas are compared", "Sample points and seed",
        "%d points, seed %d, relative tolerance %s" % (settings["numeric_points"], settings["numeric_seed"],
                                                       shared.plain_number(settings["relative_tolerance"], 3)))
    for label, value in call_statistics(store):
        add("AI calls", label, value)
    return rows

def chunk_note(chunk):
    """What a reader should know about one chunk: unreadable, how an equation was read, reconstructed numbering."""
    notes = ["Could not be read: %s" % chunk["not_read_reason"]] if chunk.get("not_read_reason") else []
    equation = chunk.get("equation") or {}
    if chunk["kind"] == "Equation" and not equation.get("readable"):
        notes.append("Could not be read: %s" % (equation.get("not_readable_reason") or "unknown form"))
    if chunk["kind"] == "Equation" and equation.get("readable"):
        notes.append("Read as: %s" % equation.get("linear", ""))
    if chunk["kind"] == "Figure":
        notes.append("The content of a figure is never read; it is raised for manual review")
    if chunk.get("numbering_reconstructed"):
        notes.append("Heading numbering was reconstructed by counting")
    return "; ".join(notes)

def rows_chunks(chunks):
    """The rows of Chunks_Canon and Chunks_Doc."""
    rows = []
    for chunk in chunks:
        rows.append({"ref": chunk["ref"], "level": chunk["level"], "section": " > ".join(chunk["heading_chain"]),
                     "para_no": chunk.get("para_label") or chunk["para_no"],
                     "kind": chunk["kind"], "text": chunk["text"],
                     "source_file": chunk["source_file"], "refs_out": "; ".join(chunk["refs_out"]),
                     "checkable": chunk.get("checkable"), "reading_note": chunk_note(chunk)})
    return rows

def unit_expression(unit):
    """A unit's formula or arguments as shown on Chunks_Model."""
    code, data = unit.get("code") or {}, unit.get("data") or {}
    if unit["kind"] == shared.KIND_FUNCTION:
        formals = ["%s = %s" % (n, d) if d not in (None, "") else n for n, d in code.get("formals", [])]
        return "arguments: " + (", ".join(formals) or "none")
    if code.get("expression"):
        return shared.expr_to_text(shared.expr_from_dict(code["expression"]))
    if data:
        return "rows identified by: %s; %s" % (", ".join(data.get("row_keys", [])) or "row number",
                                               " x ".join(str(d) for d in data.get("dims", [])))
    return ""

def rows_model_units(units, interpretations=()):
    """The rows of Chunks_Model. What the AI said a piece of code does is shown as a quotation
    (in \u201c \u201d), because the words are the model's and not AIVA's own. Enforces: R10"""
    rows, said = [], {record["unit_ref"]: record for record in interpretations}
    for unit in units:
        code, told = unit.get("code") or {}, said.get(unit["ref"], {})
        lines = "%d-%d" % tuple(unit["lines"]) if unit.get("lines") else ""
        rows.append({"ref": unit["ref"], "kind": unit["kind"], "file": unit["file"], "lines": lines,
                     "name": unit["name"], "inside": unit["inside"], "text": unit["text"],
                     "expression": unit_expression(unit), "exported": code.get("exported"),
                     "numbers": "; ".join(n["as_written"] for n in code.get("numbers", [])),
                     "reading_note": unit.get("read_problem") or "", "llm_interpretation": told.get("note") or (
                         "\u201c%s\u201d" % re.sub("[\u201c\u201d]", '"', told["interpretation"]) if told.get("interpretation") else ""),
                     "where": "%s%s" % (unit["file"], " lines " + lines if lines else "")})
    return rows

def link_columns(unit_ref, prefix, corner_letter, links, searches, texts):
    """The block of columns that shows what one unit was linked to in one corner."""
    mine = [e for e in links.get(unit_ref, []) if e["target"].startswith(corner_letter)]
    latest = {}
    for edge in mine:                                 # a later edge on the same pair has the last word
        latest[edge["target"]] = edge
    shown = [latest[ref] for ref in sorted(latest)]
    search = searches.get((unit_ref, prefix), {})
    return {prefix + "_refs": "\n".join(e["target"] for e in shown),
            prefix + "_relation": lines_by_ref((e["target"], e["relation"]) for e in shown),
            prefix + "_how": lines_by_ref((e["target"], e["evidence"].get("how_text", "")) for e in shown),
            prefix + "_searched": search.get("searched_text", "") or (shared.NOT_RUN_YET if not searches else ""),
            prefix + "_why_not": "" if shown else search.get("note", ""),
            prefix + "_text": texts_by_ref((e["target"], texts.get(e["target"], "")) for e in shown)}

def rows_mapping(store, corner):
    """One row per model unit (corner "model") or per documentation unit (corner "doc").
    The words in the assessment cells were written by aiva3 and aiva4; this only lays out."""
    units = store.read("model_units") if corner == "model" else store.read("chunks_doc")
    texts = {r["ref"]: r["text"] for kind in ("chunks_canon", "chunks_doc", "model_units") for r in store.read(kind)}
    links, searches = {}, {}
    for record in store.read("graph_ledger"):
        if record.get("record_type") == "edge" and record["kind"] == "corresponds":
            links.setdefault(record["source"], []).append(record)
            links.setdefault(record["target"], []).append(dict(record, target=record["source"]))   # the same link, seen from the other side
    corner_names = {"canon": "canon", "doc": "doc", "model": "model"}
    for record in store.read("search_records"):
        key = (record["unit_ref"], corner_names[record["target_corner"]])
        merged = dict(searches.get(key, {}))
        merged.update({k: v for k, v in record.items() if v})
        searches[key] = merged
    statuses = {r["unit_ref"]: r for r in store.read("unit_status")}
    base_rows = rows_model_units(units) if corner == "model" else rows_chunks(units)
    rows = []
    for row in base_rows:
        row.update(link_columns(row["ref"], "canon", "C-", links, searches, texts))
        other = ("doc", "D-") if corner == "model" else ("model", "M-")
        row.update(link_columns(row["ref"], other[0], other[1], links, searches, texts))
        status = statuses.get(row["ref"])
        row["status"] = status["status"] if status else shared.NOT_RUN_YET
        row["item_ids"] = ("\n".join(status["item_ids"]) or "No item") if status else ""
        row["cells"] = status["cells"] if status else {}
        rows.append(row)
    return rows

def rows_coverage(model_rows, doc_rows, store):
    """Counted from the rows actually written: the second, independent route of the
    coverage identity (part 4). Enforces: R2"""
    rows = []
    for corner, label in COVERAGE_ROWS[:2]:
        written = model_rows if corner == "model" else doc_rows
        row = {"corner": label, "total": len(written), "how_to_read": NEEDS_ATTENTION_MEANS}
        for status in shared.CLEAN_STATUSES + shared.NOT_CLEAN_STATUSES:
            row[status] = sum(1 for r in written if r["status"] == status)
        row["needs_attention"] = sum(row[status] for status in shared.NOT_CLEAN_STATUSES)
        rows.append(row)
    canon = store.read("chunks_canon")
    pointed = {e["target"] for e in store.read("graph_ledger")
               if e.get("record_type") == "edge" and e["kind"] == "corresponds" and e["relation"] in shared.LINKING_RELATIONS}
    count = sum(1 for c in canon if c["ref"] in pointed)
    rows.append({"corner": COVERAGE_ROWS[2][1], "total": len(canon), "how_to_read":
                 "%d of %d passages are pointed to by at least one link. The others are listed in the "
                 "report, for information only." % (count, len(canon))})
    return rows

def latest_determinations(store):
    """The last recorded determination of every item."""
    latest = {}
    for record in store.read("determinations"):
        latest[record["item_id"]] = record
    return latest

def rows_flagged(store):
    """The rows of Flagged_Items with the latest determination of each item."""
    latest, rows = latest_determinations(store), []
    for item in store.read("flagged_items"):
        decided = latest.get(item["item_id"])
        active = decided is not None and decided["decision"] != shared.DECISION_WITHDRAWN
        row = dict(item, unit_refs="\n".join(item["unit_refs"]),
                   status=decided["decision"] if active else shared.ITEM_OPEN,
                   last_decision_recorded=decided["recorded_at"] if decided else "")
        for name in ("decision", "reviewer", "role", "rationale"):
            row[name] = decided[name] if active else ""
        rows.append(row)
    return rows

def sheet_rows(store, paths, settings, progress):
    """The rows of all eight sheets, by sheet name."""
    model_rows, doc_rows = rows_mapping(store, "model"), rows_mapping(store, "doc")
    for row in model_rows + doc_rows:
        for name, text in row.pop("cells").items():
            row[name] = text
    return {"Model_Package_Info": rows_package_info(store, paths, settings, progress),
            "Chunks_Canon": rows_chunks(store.read("chunks_canon")),
            "Chunks_Doc": rows_chunks(store.read("chunks_doc")),
            "Chunks_Model": rows_model_units(store.read("model_units"), store.read("interpretations")),
            "Mapping_Model_to_Canon_and_Doc": model_rows, "Mapping_Doc_to_Canon_and_Model": doc_rows,
            "Mapping_Coverage": rows_coverage(model_rows, doc_rows, store), "Flagged_Items": rows_flagged(store)}

def check_written_totals(rows, store):
    """Identity part 4: what account-coverage counted must equal what the workbook holds."""
    coverage = store.read("coverage")
    if not coverage:
        return
    for position, corner in enumerate(("model", "doc")):
        counted, written = coverage[0][corner], rows["Mapping_Coverage"][position]
        same = counted["total"] == written["total"] and all(
            counted["by_status"].get(s, 0) == written[s] for s in shared.CLEAN_STATUSES + shared.NOT_CLEAN_STATUSES)
        if not same:
            raise aiva4_checks.AivaDefect(
                "Part 4 of the coverage identity does not hold: the totals counted for the %s corner differ "
                "from the rows written to the workbook. This is a defect in AIVA, not in the model under review." % corner)

def load_layout():
    """The workbook layout from references/workbook_layout.yaml."""
    with open(os.path.join(REFERENCES_DIR, "workbook_layout.yaml"), encoding="utf-8") as handle:
        return yaml.safe_load(handle)

def write_sheet(sheet, sheet_layout, rows, colours, settings, store):
    """One generic writer for all eight sheets: header row and first column frozen, filter on
    the header, wrapped text, no merged cells, reviewer columns yellow and unlocked."""
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    columns = sheet_layout["columns"]
    wrap = Alignment(wrap_text=True, vertical="top")
    for number, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=number, value=column["header"])
        cell.font, cell.alignment = Font(bold=True), wrap
        cell.fill = PatternFill("solid", start_color=colours[column["group"]])
        sheet.column_dimensions[get_column_letter(number)].width = column.get("width", 20)
    for row_number, row in enumerate(rows, start=2):
        for number, column in enumerate(columns, start=1):
            value = row.get(column["field"])
            if value in (None, "") and column["group"] == "assessments" and sheet_layout["name"].startswith(("Mapping_Model", "Mapping_Doc")):
                value = shared.NOT_RUN_YET
            cell = sheet.cell(row=row_number, column=number, value=plain_cell(value, column.get("input_text"), store))
            cell.alignment = wrap
            if column["group"] == "reviewer_input":
                cell.fill = PatternFill("solid", start_color=colours["reviewer_input"])
                cell.protection = Protection(locked=False)
    last = get_column_letter(len(columns))
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = "A1:%s%d" % (last, max(1, len(rows) + 1))
    if sheet_layout["name"] == "Flagged_Items" and rows:
        decision_column = get_column_letter([c["field"] for c in columns].index("decision") + 1)
        choice = DataValidation(type="list", formula1='"%s"' % ",".join(shared.DECISION_WORDS), allow_blank=True)
        sheet.add_data_validation(choice)
        choice.add("%s2:%s%d" % (decision_column, decision_column, len(rows) + 1))
    if settings["protect_sheets"]:
        sheet.protection.sheet = True
        sheet.protection.autoFilter = False          # False = not locked: filtering stays possible
        sheet.protection.formatColumns = False

def build_workbook(store, paths, settings, progress, target):
    """Build Output.xlsx on local disk from the audit records. All eight sheets always
    exist; a sheet whose step has not run shows its header only."""
    import openpyxl
    layout = load_layout()
    rows = sheet_rows(store, paths, settings, progress)
    check_written_totals(rows, store)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet_layout in layout["sheets"]:
        sheet = workbook.create_sheet(sheet_layout["name"])
        write_sheet(sheet, sheet_layout, rows[sheet_layout["name"]], layout["colours"], settings, store)
    workbook.properties.title = "AIVA Output"
    workbook.properties.description = shared.canonical_json(run_identity(store, paths))
    workbook.save(target)
    return rows

# ---------------------------------------------------------------- rebuilding the outputs
def file_sha256(path):
    """SHA-256 of a file's bytes."""
    with open(path, "rb") as handle:
        return shared.sha256_bytes(handle.read())

def progress_text(store, waiting_message):
    """Where the run stands, in one or two plain sentences."""
    records = store.read("step_records")
    if not records:
        return "The run has been opened; no step has finished yet."
    last = records[-1]
    text = "Step %s (%s) is the last finished step." % (last["step_id"], last["skill"])
    return text + (" " + waiting_message if waiting_message else "")

def rebuild_outputs(store, paths, settings, waiting_message):
    """Rebuild Output.xlsx (and the report once flagged items exist) on local disk and copy
    them whole into Outputs/. The guard: a workbook in Outputs/ that differs from the last
    one AIVA wrote is a reviewer's work in progress and is never overwritten before it has
    been read in. Afterwards Outputs/ holds exactly the two files. Enforces: R6, R12"""
    work = os.path.join(store.local_dir, "work")
    os.makedirs(work, exist_ok=True)
    progress = progress_text(store, waiting_message)
    manifest = (store.read("run_manifest") or [{}])[0]
    target = os.path.join(paths.outputs_dir, "Output.xlsx")
    guarded = (os.path.exists(target) and manifest.get("last_workbook_sha256")
               and file_sha256(target) not in (manifest["last_workbook_sha256"], manifest.get("last_upload_sha256")))
    local_workbook = os.path.join(work, "Output.xlsx")
    build_workbook(store, paths, settings, progress, local_workbook)
    if guarded:
        log_line(store, "Output.xlsx in Outputs/ was edited and not yet read in: left untouched")
    else:
        copy_whole(local_workbook, target)
        if manifest:
            update_manifest(store, {"last_workbook_sha256": file_sha256(target)})
    if store.read("flagged_items"):
        local_report = os.path.join(work, "Validation_Report.docx")
        build_report_file(store, paths, settings, local_report)
        copy_whole(local_report, os.path.join(paths.outputs_dir, "Validation_Report.docx"))
    build_run_summary(store, paths, progress, os.path.join(store.local_dir, "Run_Summary.docx"))
    store.sync()
    return not guarded

# ---------------------------------------------------------------- Run_Summary.docx
def docx_table(document, header, rows):
    """A plain table in a Word document, header row in bold."""
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = str(text)
    return table

def build_run_summary(store, paths, progress, target):
    """Where the run stands, in plain words; refreshed after every step."""
    import docx
    document = docx.Document()
    document.add_heading("AIVA run summary", level=1)
    document.add_paragraph("Model ID %s, date initiated %s, run %s." % (paths.model_id, paths.project_date, paths.run_id))
    document.add_heading("Where the run stands", level=2)
    document.add_paragraph(progress)
    document.add_heading("What each finished step produced", level=2)
    rows = []
    for record in store.read("step_records"):
        produced = ", ".join("%s: %d" % (k, v) for k, v in sorted(record["counts"].items())) or "nothing to count"
        rows.append((record["step_id"], record["skill"], produced, " ".join(record["messages"])))
    docx_table(document, ("Step", "Skill", "Produced", "Notes"), rows)
    document.save(target)

REPORT_SCOPE = (
    "AIVA compares three things: the methodology, the model code and data in the R package, and the model "
    "documentation. It reads all three, links what corresponds, checks formulas, values and stated rules by "
    "code, and raises what it could not line up for a person to decide. It does not run the model, does not "
    "judge whether the methodology is sound, and rates nothing: every flagged item is a question for a person.")

def call_statistics(store):
    """The AI call statistics shown on Model_Package_Info and in the report's annex."""
    calls = store.read_calls()
    if not calls:
        return [("Questions asked", "No question has been asked yet")]
    final = [c for c in calls if c.get("final")]
    rejected = {}
    for call in final:
        if call["outcome"].startswith("rejected"):
            rejected[call["outcome"]] = rejected.get(call["outcome"], 0) + 1
    with_planted = [c for c in final if c.get("planted")]
    seconds = sorted(c["seconds"] for c in calls)
    rows = [("Questions asked", len(final)), ("Attempts made", len(calls)),
            ("Answers accepted", sum(1 for c in final if c["outcome"] == "accepted")),
            ("Questions that stayed without an answer", sum(1 for c in final if c["outcome"].startswith("failed")))]
    rows += [("Answers %s" % reason, count) for reason, count in sorted(rejected.items())]
    planted = sum(1 for c in with_planted if c["outcome"] == "rejected: " + shared.REJECTION_REASONS[3])
    rows.append(("Answers that accepted a planted control passage", "%d of %d questions that held one" % (planted, len(with_planted))))
    rows.append(("Median seconds per call", shared.plain_number(seconds[len(seconds) // 2], 3)))
    return rows

def call_plan(paths, settings, step_id, seconds_per_call=0.0):
    """The call plan of one AI step, obtained by really building every question of the step
    (building is deterministic and cheap) without asking any. Shown before the step starts."""
    store, collected = open_store(paths, settings), []
    step = next(s for s in load_pipeline()["steps"] if s["id"] == step_id)
    def collect(questions):
        collected.extend(questions)
        return {}
    options = dict(step.get("with") or {})
    options.update({"inputs": list_input_files(paths), "references_dir": REFERENCES_DIR, "paths": paths,
                    "run": {"model_id": paths.model_id, "project_date": paths.project_date, "run_id": paths.run_id}})
    provenance = shared.Provenance(paths.run_id, step["id"], step["skill"], step["skill_version"])
    STEP_FUNCTIONS[step["function"]](shared.StepContext(settings, options, store.read, collect, paths.local_dir, lambda text: None, provenance))
    by_type = {}
    for question in collected:
        by_type[question["question_type"]] = by_type.get(question["question_type"], 0) + 1
    largest = max([q["estimated_tokens"] for q in collected] or [0])
    minutes = len(collected) * seconds_per_call / max(1, int(settings["concurrency_limit"])) / 60.0
    return {"step": step["skill"], "questions": len(collected), "by_type": by_type, "largest_estimated_tokens": largest,
            "token_cap": settings["token_cap"], "expected_minutes": round(minutes, 1),
            "note": "Questions that depend on earlier answers of the same step are not in this count."}

# ---------------------------------------------------------------- determinations
def find_uploads(paths, identity):
    """Every .xlsx in Outputs/ whose embedded identity matches this run, whatever it is called
    (the behaviour of an upload onto an existing name is not documented). Returns the
    matching files, newest last, and plain messages about files that were refused."""
    import openpyxl
    matching, refused = [], []
    for name in sorted(os.listdir(paths.outputs_dir)):
        path = os.path.join(paths.outputs_dir, name)
        if name.lower().endswith(".xls"):
            refused.append("%s is in the old .xls format; save it as .xlsx and upload it again." % name)
        if not name.lower().endswith(".xlsx"):
            continue
        try:
            found = json.loads(openpyxl.load_workbook(path, read_only=True).properties.description or "{}")
        except Exception:
            refused.append("%s could not be opened as a workbook; nothing was read from it." % name)
            continue
        if all(found.get(key) == identity[key] for key in ("model_id", "date_initiated", "run_id", "item_list_hash")):
            matching.append(path)
        else:
            refused.append("%s belongs to another run or another list of items; nothing was read from it." % name)
    return sorted(matching, key=os.path.getmtime), refused

def read_yellow_cells(path, known_ids):
    """The four yellow cells of every row of Flagged_Items, found by item id and never by row
    position, because reviewers sort and filter. Returns ({item id: values}, messages)."""
    import openpyxl
    sheet = openpyxl.load_workbook(path, read_only=True)["Flagged_Items"]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(cell or "") for cell in rows[0]]
    places = {name: header.index(name) for name in ("Item id", "Decision", "Reviewer", "Role", "Rationale")}
    found, messages = {}, []
    for row in rows[1:]:
        item_id = str(row[places["Item id"]] or "").strip()
        values = {name.lower(): str(row[places[name]] or "").strip() for name in ("Decision", "Reviewer", "Role", "Rationale")}
        if not item_id:
            continue
        if item_id not in known_ids:
            messages.append("The row with item id %s belongs to no flagged item of this run and was not read." % item_id)
        elif item_id in found:
            messages.append("Item %s occurs on two rows; neither was recorded." % item_id)
            found[item_id] = None
        else:
            found[item_id] = values
    return {item_id: values for item_id, values in found.items() if values is not None}, messages

def record_determinations(ctx):
    """Step 17, skill record-determinations: find the uploaded workbook by its identity, keep a
    copy of its bytes in _audit/uploads/, read and validate the yellow cells, and append one
    record for every item whose four values changed. Earlier records are never changed; a
    cleared decision is recorded as Withdrawn. The record is a hash chain. Enforces: R4, R12"""
    paths, settings = ctx.options["paths"], ctx.settings
    manifest = (ctx.read("run_manifest") or [{}])[0]
    items = {item["item_id"] for item in ctx.read("flagged_items")}
    identity = {"model_id": paths.model_id, "date_initiated": paths.project_date, "run_id": paths.run_id,
                "item_list_hash": shared.sha256_text("\n".join(i["item_id"] for i in ctx.read("flagged_items"))) if items else ""}
    uploads, messages = find_uploads(paths, identity)
    edited = [p for p in uploads if file_sha256(p) != manifest.get("last_workbook_sha256")]
    if not edited:
        return shared.StepResult({}, {"determinations recorded": 0}, messages + ["No edited workbook of this run was found in Outputs/."])
    chosen = edited[-1]
    digest = file_sha256(chosen)
    os.makedirs(os.path.join(paths.audit_dir, "uploads"), exist_ok=True)
    for path in edited:
        copy_whole(path, os.path.join(paths.audit_dir, "uploads", "%s_%s" % (file_sha256(path)[:12], os.path.basename(path))))
    cells, notes = read_yellow_cells(chosen, items)
    messages += notes
    latest, records = {}, []
    for record in ctx.read("determinations"):
        latest[record["item_id"]] = record
    now = datetime.datetime.now().isoformat(timespec="seconds")
    for item_id in sorted(cells):
        values = cells[item_id]
        previous = latest.get(item_id)
        decision = next((word for word in shared.DECISION_WORDS if word.lower() == values["decision"].lower()), "")
        if values["decision"] and not decision:
            messages.append("Item %s: the decision must be one of: %s. Nothing was recorded for it." % (item_id, ", ".join(shared.DECISION_WORDS)))
            continue
        if decision and not (values["reviewer"] and values["role"] and values["rationale"]):
            messages.append("Item %s: a decision needs a reviewer, a role and a rationale. Nothing was recorded for it." % item_id)
            continue
        if not decision:
            if previous and previous["decision"] != shared.DECISION_WITHDRAWN:
                records.append(shared.Determination(item_id, shared.DECISION_WITHDRAWN, previous["reviewer"], previous["role"],
                                                    "The recorded decision was cleared in the workbook.", now, settings["reviewer_id"], digest))
            continue
        same = previous and all(previous[key] == value for key, value in dict(values, decision=decision).items())
        if not same:
            records.append(shared.Determination(item_id, decision, values["reviewer"], values["role"], values["rationale"], now,
                                                settings["reviewer_id"], digest))
    for path in uploads:                                 # tidy Outputs/ back to exactly the two files AIVA writes
        if os.path.basename(path) != "Output.xlsx":
            os.remove(path)
    chained = shared.chain_records(shared.chain_head(ctx.read("determinations")), [shared.to_plain(r) for r in records])
    manifest.update(last_upload_sha256=digest)
    return shared.StepResult({"determinations": chained, "run_manifest": [manifest]}, {"determinations recorded": len(chained)},
                             messages + ["Text typed outside the four yellow columns is ignored."])

# ---------------------------------------------------------------- Validation_Report.docx
def build_report_file(store, paths, settings, target):
    """The report, in the eight parts of plan 2.11, built from the same records as the
    workbook. Items are ordered by item id and never ranked. Enforces: R1, R10"""
    import docx
    from docx.enum.section import WD_ORIENT
    identity, document = run_identity(store, paths), docx.Document()
    items, latest = store.read("flagged_items"), latest_determinations(store)
    decided = {k: v for k, v in latest.items() if v["decision"] != shared.DECISION_WITHDRAWN}
    info = (store.read("package_info") or [{}])[0]
    document.add_heading("AIVA validation report", level=1)
    open_count = sum(1 for item in items if item["item_id"] not in decided)
    document.add_paragraph("Every flagged item has a recorded determination." if items and not open_count
                           else "%d of %d flagged items are still open." % (open_count, len(items)))
    docx_table(document, ("Field", "Value"), [
        ("Model ID", identity["model_id"]), ("Date initiated", identity["date_initiated"]), ("Run", identity["run_id"]),
        ("Package", "%s %s" % (info.get("name", ""), info.get("version", ""))), ("Engine version", identity["engine_version"]),
        ("Graph version id", identity["graph_version_id"]), ("Item list fingerprint", identity["item_list_hash"][:16]),
        ("Determinations record fingerprint", identity["determinations_fingerprint"] or "No determination recorded yet")])
    manifest = (store.read("run_manifest") or [{}])[0]
    document.add_heading("1. What was reviewed", level=2)
    docx_table(document, ("Input file", "SHA-256", "Bytes"), [(e["file"], e["sha256"], e["bytes"]) for e in manifest.get("inputs", [])])
    for change in manifest.get("changes_since_previous_run", []):
        document.add_paragraph("Changed since run %s: %s" % (manifest.get("previous_run", ""), change))
    document.add_heading("2. What AIVA did and did not assess", level=2)
    document.add_paragraph(REPORT_SCOPE)
    limits = [(s["unit_ref"], s["reason_shown"]) for s in store.read("unit_status") if s["status"] == shared.ST_NOT_ASSESSED]
    limits += [("Repair", "%s: %s" % (r["file"], r["kind"])) for r in store.read("read_repairs")][:40]
    docx_table(document, ("Unit or file", "Limit of this run"), limits or [("None", "Nothing was left unread or unassessed")])
    document.add_heading("3. Coverage", level=2)
    coverage = (store.read("coverage") or [{}])[0]
    rows = [(label, coverage.get(corner, {}).get("total", ""), coverage.get(corner, {}).get("needs_attention", "for information only"))
            for corner, label in COVERAGE_ROWS]
    docx_table(document, ("Corner", "Units", "Needs attention"), rows)
    document.add_paragraph(NEEDS_ATTENTION_MEANS)
    document.add_heading("4. Flagged items by concern and category", level=2)
    counts = {}
    for item in items:
        counts[(item["concerns"], item["category"])] = counts.get((item["concerns"], item["category"]), 0) + 1
    docx_table(document, ("Concerns", "Category", "Items"), [(c, k, n) for (c, k), n in sorted(counts.items())])
    section = document.add_section()
    section.orientation, section.page_width, section.page_height = WD_ORIENT.LANDSCAPE, section.page_height, section.page_width
    document.add_heading("5. Every flagged item", level=2)
    for item in items:
        document.add_heading("%s - %s" % (item["item_id"], item["category"]), level=3)
        document.add_paragraph("Concerns: %s. Units: %s." % (item["concerns"], ", ".join(item["unit_refs"])))
        document.add_paragraph(item["observed"][:3000])
        docx_table(document, ("Methodology says", "Code does", "Documentation says"),
                   [(item["methodology_says"][:1500], item["code_does"][:1500], item["documentation_says"][:1500])])
        document.add_paragraph("Suggested next step: %s" % item["suggested_next_step"])
        record = latest.get(item["item_id"])
        if record:
            by = "" if record["recorded_by"] in ("", record["reviewer"]) else ", recorded by %s for reviewer %s" % (record["recorded_by"], record["reviewer"])
            document.add_paragraph("Determination: %s, by %s (%s) at %s%s. Rationale: %s" % (
                record["decision"], record["reviewer"], record["role"], record["recorded_at"], by, quoted(record["rationale"])))
        else:
            document.add_paragraph("Determination: none recorded yet (the item is open).")
    document.add_heading("6. Methodology passages that nothing points to (for information)", level=2)
    pointed = {e["target"] for e in store.read("graph_ledger") if e.get("record_type") == "edge" and e["kind"] == "corresponds"
               and e["relation"] in shared.LINKING_RELATIONS}
    unpointed = [(c["ref"], " > ".join(c["heading_chain"][-2:]), c["kind"]) for c in store.read("chunks_canon") if c["ref"] not in pointed]
    docx_table(document, ("Passage", "Section", "Kind"), unpointed or [("None", "", "")])
    document.add_heading("7. How to re-verify this pack", level=2)
    document.add_paragraph("Open the notebook, enter the model ID, the date initiated and this run, and run the cell \"Verify this evidence "
                           "pack\". It re-hashes the inputs, re-reads them, verifies both hash chains and re-checks the coverage identity.")
    document.add_heading("Annex. AI call statistics", level=2)
    docx_table(document, ("Measure", "Value"), call_statistics(store))
    document.save(target)

def build_report(ctx):
    """Step 18, skill build-report: the exports of the graph for anyone who wants to load it
    elsewhere (nodes.csv, edges.csv, graph.graphml). The two output files themselves are
    rebuilt by the runner after every step."""
    import csv
    paths, ledger = ctx.options["paths"], ctx.read("graph_ledger")
    folder = os.path.join(paths.audit_dir, "exports")
    os.makedirs(folder, exist_ok=True)
    nodes = [r for r in ledger if r["record_type"] == "node"]
    edges = [r for r in ledger if r["record_type"] == "edge"]
    with open(os.path.join(folder, "nodes.csv"), "w", newline="", encoding="utf-8") as handle:
        csv.writer(handle).writerows([("ref", "kind", "corner")] + [(n["ref"], n["node_kind"], n["corner"]) for n in nodes])
    with open(os.path.join(folder, "edges.csv"), "w", newline="", encoding="utf-8") as handle:
        csv.writer(handle).writerows([("source", "target", "kind", "relation", "how")] +
                                     [(e["source"], e["target"], e["kind"], e.get("relation", ""), e["how"]) for e in edges])
    escape = lambda text: str(text).replace("&", "&amp;").replace("<", "&lt;").replace('"', "&quot;")
    lines = ['<?xml version="1.0" encoding="UTF-8"?>', '<graphml xmlns="http://graphml.graphdrawing.org/xmlns">',
             '<key id="kind" for="all" attr.name="kind" attr.type="string"/>', '<graph edgedefault="directed">']
    lines += ['<node id="%s"><data key="kind">%s</data></node>' % (escape(n["ref"]), escape(n["node_kind"])) for n in nodes]
    lines += ['<edge source="%s" target="%s"><data key="kind">%s</data></edge>' % (escape(e["source"]), escape(e["target"]), escape(e["kind"])) for e in edges]
    with open(os.path.join(folder, "graph.graphml"), "w", encoding="utf-8") as handle:
        handle.write("\n".join(lines + ["</graph>", "</graphml>"]))
    return shared.StepResult({}, {"nodes exported": len(nodes), "edges exported": len(edges)}, [])

# ---------------------------------------------------------------- verify this evidence pack
def verify_evidence_pack(paths, settings, live=None):
    """Works from a run folder and the Inputs folder alone. Returns rows (what was checked,
    "Confirmed" or "Not confirmed", detail). Re-reading the inputs is reperformance: every
    content hash is computed again from the input files and compared with the record."""
    store, rows = AuditStore(paths.audit_dir, paths.audit_dir), []        # read the pack itself, not the scratch copy of this driver
    def line(what, good, detail=""):
        rows.append((what, "Confirmed" if good else "Not confirmed", detail))
    manifest = (store.read("run_manifest") or [{}])[0]
    changed = [e["file"] for e in manifest.get("inputs", []) if not os.path.exists(os.path.join(paths.inputs_dir, e["file"]))
               or file_sha256(os.path.join(paths.inputs_dir, e["file"])) != e["sha256"]]
    line("Input files have the recorded fingerprints", not changed, ", ".join(changed))
    installed, recorded = engine_file_hashes(), manifest.get("engine_files", {})
    other = sorted(name for name in set(installed) | set(recorded) if installed.get(name) != recorded.get(name))
    line("The engine files that produced this run are the ones installed here", not other, ", ".join(other[:5]))
    options = {"inputs": list_input_files(paths), "references_dir": REFERENCES_DIR}
    context = shared.StepContext(settings, options, lambda kind: [], None, paths.local_dir, lambda text: None)
    for kind, function in (("chunks_canon", aiva1_documents.read_methodology), ("chunks_doc", aiva1_documents.read_documentation),
                           ("model_units", aiva2_package.read_package)):
        fresh = {shared.to_plain(r)["ref"]: shared.to_plain(r)["content_hash"] for r in function(context).records.get(kind, [])}
        recorded = {r["ref"]: r["content_hash"] for r in store.read(kind)}
        differing = sorted(ref for ref in set(fresh) | set(recorded) if fresh.get(ref) != recorded.get(ref))
        line("Re-reading the inputs gives the recorded content hashes (%s)" % kind, not differing, ", ".join(differing[:5]))
    ledger_ok, position, _ = shared.verify_chain(store.read("graph_ledger"), aiva3_mapping.LEDGER_VOLATILE)
    line("The graph ledger chain verifies", ledger_ok, "" if ledger_ok else "record %d no longer verifies" % (position + 1))
    decisions_ok, position, _ = shared.verify_chain(store.read("determinations"))
    line("The determinations chain verifies", decisions_ok, "" if decisions_ok else "record %d no longer verifies" % (position + 1))
    statuses, items = store.read("unit_status"), store.read("flagged_items")
    named = {ref for item in items for ref in item["unit_refs"]}
    not_clean = {s["unit_ref"] for s in statuses if not s["clean"]}
    expected = {r["ref"] for kind in ("model_units", "chunks_doc") for r in store.read(kind)}
    line("Every unit has one status; units that are not clean and flagged items match",
         {s["unit_ref"] for s in statuses} == expected and len(statuses) == len(expected) and not_clean == named & expected)
    known = expected | {r["ref"] for r in store.read("chunks_canon")}
    line("Every citation on Flagged_Items resolves to a unit", all(ref in known for ref in named))
    identity = run_identity(store, paths)
    try:
        import openpyxl
        found = json.loads(openpyxl.load_workbook(os.path.join(paths.outputs_dir, "Output.xlsx"), read_only=True).properties.description)
        line("The workbook carries this run's ids and fingerprints", all(found.get(k) == v for k, v in identity.items()))
    except Exception:
        line("The workbook carries this run's ids and fingerprints", False, "Output.xlsx could not be opened")
    secrets = [value for value in (list(live.recent_tokens) if live else []) if value]
    leaked = []
    for folder, _, names in os.walk(paths.run_dir):
        for name in names:
            with open(os.path.join(folder, name), "rb") as handle:
                data = handle.read()
            leaked += [name for value in secrets if value.encode("utf-8") in data]
    line("No access token was written into the run folder", not leaked, ", ".join(sorted(set(leaked))))
    return rows

STEP_FUNCTIONS = {        # every function that pipeline.yaml is allowed to name. Enforces: R11
    "aiva1_documents.read_methodology": aiva1_documents.read_methodology,
    "aiva1_documents.read_documentation": aiva1_documents.read_documentation,
    "aiva2_package.read_package": aiva2_package.read_package,
    "aiva3_mapping.build_graph": aiva3_mapping.build_graph,
    "aiva3_mapping.find_candidates": aiva3_mapping.find_candidates,
    "aiva3_mapping.judge_links": aiva3_mapping.judge_links, "aiva3_mapping.interpret_code": aiva3_mapping.interpret_code,
    "aiva4_checks.check_mathematics": aiva4_checks.check_mathematics,
    "aiva4_checks.check_values": aiva4_checks.check_values,
    "aiva4_checks.check_rules": aiva4_checks.check_rules,
    "aiva4_checks.check_package_docs": aiva4_checks.check_package_docs,
    "aiva4_checks.account_coverage": aiva4_checks.account_coverage,
    "aiva5_run_report.prepare_run": prepare_run,
    "aiva5_run_report.record_determinations": record_determinations,
    "aiva5_run_report.build_report": build_report}
