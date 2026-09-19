"""build_manual.py - assembles docs/AIVA_User_Manual.md and .docx from docs/manual_src/ (plan, Phase 11).

Everything that can be generated from the source is generated here, so that the manual cannot
drift from the code: the code index, the vocabulary, sheets and columns, skills, the map from
design rule to enforcing function, settings, dependencies, the status rules, prompts, tests and
line counts. A chapter asks for a generated part with a marker such as {{settings}}.

    python tools/build_manual.py            (then: python tools/check_docs.py)
"""
import ast
import glob
import os
import re
import sys

import yaml

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ENGINE = os.path.join(ROOT, "engine")
for folder in (ENGINE, os.path.join(ROOT, "tools")):
    if folder not in sys.path:
        sys.path.insert(0, folder)

import aiva0_shared as shared          # noqa: E402
import aiva4_checks                    # noqa: E402
import aiva5_run_report as run         # noqa: E402
import count_lines                     # noqa: E402

BUNDLES = ("aiva0_shared", "aiva0r_reading", "aiva1_documents", "aiva2_package", "aiva3_mapping", "aiva4_checks", "aiva5_run_report")
RULES = ["R%d" % n for n in range(1, 14)]
SETTING_NOTES = {
    "agentic_reading": "Whether a reading step may ask the model what the tags of a file whose shape AIVA does not know are for. \"off\" asks nothing and reads as the built-in rules read; \"rules\" asks one question per file the rules are unsure about and applies the answer under everything the rules already know.",
    "k_candidates": "How many passages are shown to the judge for one unit and corner.", "concurrency_limit": "How many questions are asked at the same time.",
    "token_cap": "The gateway's limit for one call, in tokens.", "answer_reserve": "Tokens kept free for the answer.",
    "thinking_reserve": "Tokens kept free for a model that writes out its reasoning first.", "safety_margin": "Share of the remaining room left unused, because tokens are estimated.",
    "prompt_target_tokens": "The size a prompt should stay below even when the cap allows more.", "max_attempts": "Attempts per question before it ends as failed.",
    "breaker_after_failures": "Failed calls in a row after which the run pauses itself.", "retry_wait_seconds": "Waiting time before a retry; it grows with every attempt.",
    "token_lifetime_minutes": "Age at which a token is treated as run out and a fresh one is awaited.", "token_wait": "wait: workers wait for a fresh token (modes A and B). stop: the run stops and is resumed (mode C).",
    "foreground_minutes": "Time box of a foreground run; 0 means none.", "sync_every_calls": "Call records are written and copied to the Workspace after this many questions.",
    "llm_file_roll_mb": "Size at which a new file of call records is started.", "require_outline_confirmation": "The AI steps wait until a person has confirmed the outline of the methodology.",
    "second_opinion": "When the second, oppositely framed question is asked: unchecked_only, all or none.", "judge_supporting_code": "Also send supporting code by syntax to the search and the judge.",
    "max_parameter_cells": "A stored object with more cells is profiled and not compared.", "max_parameter_columns": "A stored object with more columns is profiled and not compared.",
    "protect_sheets": "Lock every cell except the yellow ones (filtering stays allowed; no password).", "system_prompt_prefix": "Text put in front of every system prompt, for example a switch that turns written-out reasoning off.",
    "strip_patterns": "Patterns of thought blocks that are removed from an answer before it is read.", "numeric_points": "Sample points of the numerical check.",
    "numeric_seed": "The seed of the sample points; recorded in every check.", "min_valid_points": "Fewer valid points than this leave a check undecided.",
    "relative_tolerance": "Two results agree when they differ by less than this, relative to their size.", "trivial_numbers": "Numbers that are not looked up in the methodology.",
    "bm25_k1": "Text ranking: how fast repeated words stop counting.", "bm25_b": "Text ranking: how much long passages are scaled down.",
    "anchor_max_share": "An anchor that more than this share of all units mention is dropped.", "walk_restart": "Restart probability of the walk over units and anchors.",
    "walk_rounds": "Rounds of the walk; fixed, so that it is deterministic.", "heading_anchor_cap": "The most a shared section title can weigh.",
    "rrf_constant": "The constant of reciprocal rank fusion.", "reserved_places": "Places of the shortlist kept for candidates that only the anchors or propagation found.",
    "max_unit_chars": "A longer unit is cut around its formula lines before it is shown to the model.", "max_passage_chars": "A longer passage is cut around the matched words.",
    "max_file_mb": "A larger input file is not read and becomes a not-read unit.", "reviewer_id": "Who runs the notebook; recorded with confirmations and determinations.",
    "reviewer_role": "The role of that person.",
    "interpret_code": "Ask the AI to say in plain words what each function, formula statement, top-level statement and test block does, shown with where it sits in the whole package. Fills the column LLM Interpretation on Chunks_Model. One question per piece of code; switch it off to save the calls.",
    "read_pictures": "Read the words inside pictures by OCR when the optional package rapidocr-onnxruntime is installed. The words are shown under the Figure as a machine reading; a Figure still ends for manual review.",
    "signals": "The search signals in use; the ablation ladder of tools/recall_at_k.py switches them off one by one."}
COLUMN_NOTES = {
    "Math check": "Result of comparing the unit's formula with every linked formula: the aligned symbols, then agrees, differs with a counterexample, or could not be decided with its reason.",
    "Parameter completeness": "For a stored table: the cell-by-cell comparison with the linked table. For code: which symbol of the formula each code symbol stands for.",
    "Logic consistency": "Floors, caps and thresholds stated in linked passages, and whether the code applies them. On the documentation sheet: the judged relations.",
    "Documentation consistency": "Results of the checks of the roxygen block and help page, and where the model documentation describes the unit.",
    "Hard-coded numbers": "Every non-trivial number in the code, and where the linked passages state it.", "Unit test": "Which test block calls the function. For information only; it never raises an item.",
    "Quality notes (AI)": "Text written by the model, labelled as such and filtered.", "Value check": "Numbers and tables of the documentation compared with the methodology under the value rule.",
    "Parameter note (AI)": "A column mapping or symbol alignment proposed by the model, where one was asked for.", "Documentation quality notes": "Deterministic notes: references that resolve to nothing, reconstructed numbering, unreadable parts, repeated paragraphs.",
    "Overall status": "The one final status of the unit (chapter 6).", "Flagged item(s)": "Ids of the rows on Flagged_Items that name this unit."}


def table(header, rows):
    clean = lambda cell: str(cell).replace("|", "/").replace("\n", " ")
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    return "\n".join(lines + ["| " + " | ".join(clean(cell) for cell in row) + " |" for row in rows])


def first_sentence(text):
    text = " ".join((text or "").split())
    return re.split(r"(?<=[.!?])\s", text, maxsplit=1)[0] if text else ""


def part_status_rules():
    out = []
    for corner, title in (("model", "Units of the package"), ("doc", "Units of the documentation")):
        rows = [(n, rule, status, when) for n, (c, rule, status, when) in enumerate([r for r in aiva4_checks.STATUS_RULES if r[0] == corner], start=1)]
        out.append("**%s** (the first rule that fits decides)\n\n%s" % (title, table(("Order", "Rule", "Status", "When it applies"), rows)))
    return "\n\n".join(out)


def part_columns():
    out = []
    for sheet in run.load_layout()["sheets"]:
        rows = [(column["header"], column["group"].replace("_", " "), COLUMN_NOTES.get(column["header"], "")) for column in sheet["columns"]]
        out.append("**%s**\n\n%s" % (sheet["name"], table(("Column", "Colour group", "What it shows"), rows)))
    return "\n\n".join(out)


def skill_front_matter(name):
    with open(os.path.join(ENGINE, "skills", name, "SKILL.md"), encoding="utf-8") as handle:
        return yaml.safe_load(handle.read().split("---")[1])


def part_skills():
    rows = []
    for step in run.load_pipeline()["steps"]:
        front = skill_front_matter(step["skill"])
        rows.append((step["id"], step["skill"], front["metadata"]["version"], step.get("function", "a person"), front["description"]))
    return table(("Step", "Skill", "Version", "Carried out by", "What it does"), rows)


def functions_of(bundle):
    with open(os.path.join(ENGINE, bundle + ".py"), encoding="utf-8") as handle:
        tree = ast.parse(handle.read())
    found = []
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.ClassDef)):
            found.append((node.name, node.lineno, ast.get_docstring(node) or "", isinstance(node, ast.ClassDef)))
            if isinstance(node, ast.ClassDef):
                found += [("%s.%s" % (node.name, sub.name), sub.lineno, ast.get_docstring(sub) or "", False) for sub in node.body if isinstance(sub, ast.FunctionDef)]
    return found


def rule_map():
    found = {rule: [] for rule in RULES}
    for bundle in BUNDLES:
        for name, _, docstring, _ in functions_of(bundle):
            for tagged in re.findall(r"Enforces: ((?:R\d+(?:, )?)+)", docstring):
                for rule in tagged.split(", "):
                    found.setdefault(rule, []).append("%s.%s" % (bundle, name))
    return found


def part_rules_to_functions():
    return table(("Rule", "Enforced in"), [(rule, ", ".join("`%s`" % f for f in functions) or "layout tests") for rule, functions in rule_map().items()])


def part_settings():
    return table(("Setting", "Default", "Meaning"), [("`%s`" % name, default, SETTING_NOTES.get(name, "")) for name, default in run.DEFAULT_SETTINGS.items()])


def part_prompts():
    rows = []
    for path in sorted(glob.glob(os.path.join(ENGINE, "references", "prompts", "*.txt"))):
        with open(path, encoding="utf-8") as handle:
            text = handle.read()
        task = re.search(r"TASK: (.*)", text)
        rows.append((os.path.basename(path)[:-4], text.split("\n", 1)[0], task.group(1) if task else ""))
    return table(("Question type", "Version", "The one task of the prompt"), rows)


def part_tests():
    rows = []
    for path in sorted(glob.glob(os.path.join(ENGINE, "tests", "test_*.py"))):
        with open(path, encoding="utf-8") as handle:
            tree = ast.parse(handle.read())
        count = sum(1 for node in ast.walk(tree) if isinstance(node, ast.FunctionDef) and node.name.startswith("test_"))
        rows.append(("`%s`" % os.path.basename(path), count, first_sentence(ast.get_docstring(tree))))
    return table(("Test file", "Tests", "What it covers"), rows)


def part_line_counts():
    rows = []
    for name, budget in count_lines.BUDGETS.items():
        counted = count_lines.count(os.path.join(ENGINE, name))
        rows.append((name, counted["total"], budget, "%d%%" % round(100 * counted["share"])))
    note = ("**Line counts.** The plan asks for at least 30 percent of each file to be docstrings, comments and overview. The files are below that "
            "share; the numbers are reported here as they are.")
    return note + "\n\n" + table(("File", "Lines", "Budget", "Docstrings and comments"), rows)


def part_dependencies():
    with open(os.path.join(ENGINE, "requirements.txt"), encoding="utf-8") as handle:
        lines = [line.strip() for line in handle if line.strip()]
    rows, group = [], "required"
    for line in lines:
        if line.startswith("#"):
            group = "optional" if "Optional" in line else group
            if line.startswith("# pyreadr"):
                rows.append(("pyreadr", "optional, not used in 0.0.1", "a second, independent reader of stored data"))
            continue
        rows.append((line, group, ""))
    return table(("Package", "Needed", "Note"), rows)


def part_vocabulary():
    parts = [("Statuses that are clean", [(s,) for s in shared.CLEAN_STATUSES]), ("Statuses that need attention", [(s,) for s in shared.NOT_CLEAN_STATUSES]),
             ("Relations, as shown", [(shown, asked) for asked, shown in shared.RELATION_WORDING.items()]),
             ("How a link was established", [(text,) for text in (shared.HOW_PARSED, shared.HOW_AI, shared.HOW_SYMBOLIC, shared.HOW_NUMERIC_AGREES, shared.HOW_NUMERIC_DIFFERS,
                                                                  shared.HOW_VALUE_AGREES, shared.HOW_VALUE_DIFFERS, shared.HOW_TABLE, shared.HOW_PERSON)]),
             ("Kinds of model unit", [(kind,) for kind in shared.UNIT_KINDS]), ("Kinds of chunk", [(kind,) for kind in shared.CHUNK_KINDS]),
             ("Reasons for a check that could not be decided", [(reason,) for reason in shared.UNDECIDED_REASONS]),
             ("Reasons why an answer of the model could not be used", [(reason,) for reason in shared.REJECTION_REASONS]),
             ("Fixed cell texts", [(shared.CHECK_UNDECIDED,), (shared.NOT_APPLICABLE,), (shared.NOT_RUN_YET,), (shared.AI_WORDING_NOT_SHOWN,)])]
    out = []
    for title, rows in parts:
        header = ("Shown in the workbook", "Word allowed in a prompt") if title.startswith("Relations") else ("Wording",)
        out.append("**%s**\n\n%s" % (title, table(header, rows)))
    return "\n\n".join(out)


def part_categories():
    rows = [(category, aiva4_checks.NEXT_STEPS[category]) for category in shared.CATEGORIES]
    words = ", ".join("*%s*" % word for word in shared.DECISION_WORDS)
    return ("**Concerns:** " + ", ".join(shared.CONCERNS) + ".\n\n" + table(("Category", "Suggested next step, as shown"), rows) +
            "\n\n**Decision words:** %s. A cleared decision is recorded as *%s*; an item without an active decision is *%s*." % (words, shared.DECISION_WITHDRAWN, shared.ITEM_OPEN))


def part_code_index():
    out = []
    for bundle in BUNDLES:
        rows = [("`%s`" % name, line, "class" if is_class else "function", first_sentence(doc)) for name, line, doc, is_class in functions_of(bundle)]
        out.append("**%s.py**\n\n%s" % (bundle, table(("Name", "Line", "Kind", "What it does"), rows)))
    return "\n\n".join(out)


PARTS = {"status_rules": part_status_rules, "columns": part_columns, "skills": part_skills, "rules_to_functions": part_rules_to_functions,
         "settings": part_settings, "prompts": part_prompts, "tests": part_tests, "line_counts": part_line_counts, "dependencies": part_dependencies,
         "vocabulary": part_vocabulary, "categories": part_categories, "code_index": part_code_index}


def assemble():
    chapters = []
    for path in sorted(glob.glob(os.path.join(ROOT, "docs", "manual_src", "*.md"))):
        with open(path, encoding="utf-8") as handle:
            chapters.append(handle.read().strip())
    text = "# AIVA 0.0.1 - User Manual\n\nThis manual is assembled by `tools/build_manual.py`; its tables are generated from the source.\n\n" + "\n\n".join(chapters)
    return re.sub(r"\{\{(\w+)\}\}", lambda found: PARTS[found.group(1)](), text) + "\n"


def add_runs(paragraph, text):
    """Inline markup of a paragraph: **bold**, *italic* and `code`."""
    for piece in re.split(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*\*)", text):
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            paragraph.add_run(piece[1:-1]).font.name = "Consolas"
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        elif piece:
            paragraph.add_run(piece)


def to_docx(markdown, target):
    import docx
    from docx.shared import Pt
    document = docx.Document()
    lines, position = re.sub(r"<!--.*?-->", "", markdown).split("\n"), 0
    while position < len(lines):
        line = lines[position]
        if line.startswith("```"):
            block = []
            position += 1
            while position < len(lines) and not lines[position].startswith("```"):
                block.append(lines[position])
                position += 1
            for run_ in [document.add_paragraph().add_run("\n".join(block))]:
                run_.font.name, run_.font.size = "Consolas", Pt(8.5)
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            document.add_heading(line.lstrip("# "), level=min(level, 4))
        elif line.startswith("|") and position + 1 < len(lines) and lines[position + 1].startswith("|---"):
            rows = []
            while position < len(lines) and lines[position].startswith("|"):
                if not lines[position].startswith("|---"):
                    rows.append([cell.strip() for cell in lines[position].strip().strip("|").split("|")])
                position += 1
            grid = document.add_table(rows=0, cols=len(rows[0]))
            grid.style = "Table Grid"
            for number, row in enumerate(rows):
                for cell, text in zip(grid.add_row().cells, row + [""] * (len(rows[0]) - len(row))):
                    cell.text = ""
                    add_runs(cell.paragraphs[0], text)
                    for run_ in cell.paragraphs[0].runs:
                        run_.font.size = Pt(8.5)
                        run_.bold = run_.bold or number == 0
            continue
        elif re.match(r"^(- |\d+\. )", line):
            style = "List Bullet" if line.startswith("- ") else "List Number"
            add_runs(document.add_paragraph(style=style), re.sub(r"^(- |\d+\. )", "", line))
        elif line.strip():
            add_runs(document.add_paragraph(), line.strip())
        position += 1
    document.save(target)


def main():
    markdown = assemble()
    target = os.path.join(ROOT, "docs", "AIVA_User_Manual.md")
    with open(target, "w", encoding="utf-8") as handle:
        handle.write(markdown)
    to_docx(markdown, os.path.join(ROOT, "docs", "AIVA_User_Manual.docx"))
    print("%s (%d lines) and AIVA_User_Manual.docx" % (target, markdown.count("\n")))


if __name__ == "__main__":
    main()
