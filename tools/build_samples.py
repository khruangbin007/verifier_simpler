"""build_samples.py - builds the sample projects under engine/tests/sample_projects/.

Usage: python tools/build_samples.py [A_minimal F_capital F_capital_known D_dosing]
The generated input files are checked in, so the tests never depend on this tool; run it only
when a sample has to change. Everything a sample contains is written down in this file, so a
reviewer can read what the engine is tested against. Samples are invented; none is real.
"""
import io
import os
import sys
import tarfile
import warnings

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAMPLES = os.path.join(ROOT, "engine", "tests", "sample_projects")
FOLDERS = {"canon": "1_Methodology", "package": "2_Model_Package", "doc": "3_Model_Documentation"}


# ------------------------------------------------------------------ writers
def write_bytes(sample, corner, name, data):
    folder = os.path.join(SAMPLES, sample, "Inputs", FOLDERS[corner])
    os.makedirs(folder, exist_ok=True)
    with open(os.path.join(folder, name), "wb") as handle:
        handle.write(data if isinstance(data, bytes) else data.encode("utf-8"))


def tarball(package_name, files):
    """A .tar.gz with one top folder, fixed times and owners, so the same content gives the same bytes."""
    raw = io.BytesIO()
    with tarfile.open(fileobj=raw, mode="w", format=tarfile.GNU_FORMAT) as archive:
        for path in sorted(files):
            data = files[path] if isinstance(files[path], bytes) else files[path].encode("utf-8")
            member = tarfile.TarInfo("%s/%s" % (package_name, path))
            member.size, member.mtime, member.mode = len(data), 1767225600, 0o644
            archive.addfile(member, io.BytesIO(data))
    import gzip
    packed = io.BytesIO()
    with gzip.GzipFile(filename="", mode="wb", fileobj=packed, mtime=0) as handle:
        handle.write(raw.getvalue())
    return packed.getvalue()


def r_data(objects, single=False):
    """Stored R data written without R, by the writer of the `rdata` package (fixture route c)."""
    import rdata
    path = os.path.join(ROOT, "tmp_rdata.bin")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        if single:
            rdata.write_rds(path, objects)
        else:
            rdata.write_rda(path, objects)
    with open(path, "rb") as handle:
        data = handle.read()
    os.remove(path)
    return data


def docx_file(parts):
    """parts: ("h", level, text) | ("p", text) | ("li", text) | ("table", caption, rows) | ("eq", omml xml)
    | ("image", caption)"""
    import docx
    from docx.oxml import parse_xml
    document = docx.Document()
    for part in parts:
        if part[0] == "h":
            document.add_heading(part[2], level=part[1])
        elif part[0] == "p":
            document.add_paragraph(part[1])
        elif part[0] == "li":
            document.add_paragraph(part[1], style="List Bullet")
        elif part[0] == "table":
            document.add_paragraph(part[1], style="Caption")
            table = document.add_table(rows=len(part[2]), cols=len(part[2][0]))
            table.style = "Table Grid"
            for row, cells in zip(table.rows, part[2]):
                for cell, text in zip(row.cells, cells):
                    cell.text = text
        elif part[0] == "eq":
            paragraph = document.add_paragraph()
            paragraph._p.append(parse_xml(part[1]))
        elif part[0] == "image":
            document.add_picture(io.BytesIO(tiny_png()))
            document.add_paragraph(part[1], style="Caption")
    raw = io.BytesIO()
    document.save(raw)
    return raw.getvalue()


def pdf_file(parts):
    """parts as for docx_file; headings are bold and larger, tables are drawn with a grid."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    styles, story, raw = getSampleStyleSheet(), [], io.BytesIO()
    for part in parts:
        if part[0] == "h":
            story.append(Paragraph(part[2], styles["Heading%d" % min(3, part[1])]))
        elif part[0] in ("p", "li"):
            story.append(Paragraph(part[1], styles["BodyText"]))
        elif part[0] == "table":
            story.append(Paragraph(part[1], styles["BodyText"]))
            story.append(Table(part[2], style=TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)])))
        story.append(Spacer(1, 8))
    SimpleDocTemplate(raw, pagesize=A4, invariant=1).build(story)
    return raw.getvalue()


def tiny_png():
    import base64
    return base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def omml(inner):
    return ('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">%s</m:oMath>' % inner)


def run(text):
    return "<m:r><m:t>%s</m:t></m:r>" % text


# ------------------------------------------------------------------ sample A_minimal (neutral flavour)
A_METHODOLOGY = """<?xml version="1.0" encoding="UTF-8"?>
<document>
 <title>Parcel pricing method</title>
 <section><title>1. Purpose and scope</title>
  <p>This method sets out how the price of a parcel is determined from its size, its weight and its destination zone.</p>
  <p>It applies to every parcel handed in at a counter. Pallets and letters are outside its scope.</p>
  <p>The method is reviewed once a year by the pricing committee.</p>
 </section>
 <section><title>2. Definitions</title>
  <list>
   <item>Actual weight (AW) is the weight of the parcel in kilograms as shown by the scale.</item>
   <item>Dimensional weight (DW) is a weight derived from the volume of the parcel.</item>
   <item>Chargeable weight (CW) is the weight on which the price is based.</item>
   <item>The zone is the destination area of the parcel: Local, National or Abroad.</item>
  </list>
 </section>
 <section><title>3. Weights</title>
  <section><title>3.1 Dimensional weight</title>
   <p>The dimensional weight is obtained from the length L, the width W and the height H of the parcel, measured in centimetres.</p>
   <equation><math xmlns="http://www.w3.org/1998/Math/MathML"><mi>DW</mi><mo>=</mo><mfrac><mrow><mi>L</mi><mo>&#xD7;</mo><mi>W</mi><mo>&#xD7;</mo><mi>H</mi></mrow><mn>5000</mn></mfrac></math></equation>
   <p>The divisor of 5000 is fixed and is not varied by zone.</p>
  </section>
  <section><title>3.2 Chargeable weight</title>
   <p>The chargeable weight is the larger of the actual weight and the dimensional weight: CW = max(AW, DW).</p>
   <p>The chargeable weight is never taken below 0.5 kilograms.</p>
   <p>Weights are not rounded before the price is calculated.</p>
  </section>
 </section>
 <section><title>4. Price</title>
  <section><title>4.1 Zone rates</title>
   <p>Each zone has a base rate and a rate per kilogram, as set out in Table 1.</p>
   <table><caption>Table 1. Rates by zone</caption>
    <tr><th>Zone</th><th>Base rate</th><th>Rate per kg</th></tr>
    <tr><td>Local</td><td>3.20</td><td>0.85</td></tr>
    <tr><td>National</td><td>4.90</td><td>1.10</td></tr>
    <tr><td>Abroad</td><td>9.50</td><td>2.75</td></tr>
   </table>
   <p>Rates are stated in the currency of the counter and exclude tax.</p>
  </section>
  <section><title>4.2 Price before surcharge</title>
   <p>The price before surcharge is the base rate plus the rate per kilogram times the chargeable weight: P = B + R * CW.</p>
   <p>B and R are taken from Table 1 for the zone of the parcel.</p>
  </section>
  <section><title>4.3 Fuel surcharge</title>
   <p>A fuel surcharge of 12.5% is added to the price before surcharge: S = P * (1 + 0.125).</p>
   <p>The surcharge rate is set by the pricing committee and is the same for all zones.</p>
  </section>
  <section><title>4.4 Volume discount</title>
   <p>A customer who hands in n parcels at once receives a discount rate of 1% per parcel, capped at 20%: D = min(0.2, 0.01 * n).</p>
   <p>The discount is applied to the price after surcharge.</p>
   <figure><img src="discount_curve.png" alt="Discount rate against number of parcels"/><caption>Figure 1. Discount rate by number of parcels</caption></figure>
  </section>
  <section><title>4.5 Minimum price</title>
   <p>The final price is never below 4.50, whatever the zone and the discount.</p>
   <p>Final prices are rounded to 2 decimals.</p>
  </section>
 </section>
 <section><title>5. Governance</title>
  <p>Changes to Table 1 are approved by the pricing committee and recorded in the change log.</p>
  <p>Questions about this method are answered by the pricing office.</p>
 </section>
</document>
"""

A_PACKAGE = {
    "DESCRIPTION": "Package: parcelcost\nTitle: Parcel Pricing\nVersion: 1.2.0\nDescription: Prices parcels by zone\n    and chargeable weight.\nLicense: MIT\nEncoding: UTF-8\n",
    "NAMESPACE": "export(dim_weight)\nexport(chargeable_weight)\nexport(parcel_price)\nexport(volume_discount)\n",
    "R/weights.R": """#' Dimensional weight
#'
#' Weight derived from the volume of a parcel: \\deqn{DW = \\frac{L \\cdot W \\cdot H}{5000}}
#' @param l length in centimetres
#' @param w width in centimetres
#' @param h height in centimetres
#' @param divisor the fixed divisor, 5000
#' @export
dim_weight <- function(l, w, h, divisor = 5000) {
  stopifnot(is.numeric(l), is.numeric(w), is.numeric(h))
  l * w * h / divisor
}

#' Chargeable weight
#'
#' The larger of actual weight and dimensional weight, never below 0.5 kilograms.
#' @param aw actual weight in kilograms
#' @param dw dimensional weight in kilograms
#' @export
chargeable_weight <- function(aw, dw) {
  cw <- pmax(aw, dw)
  pmax(cw, 0.5)
}
""",
    "R/price.R": """#' Price of a parcel
#'
#' Base rate plus rate per kilogram times chargeable weight, with the fuel surcharge of 12.5%,
#' the volume discount and the minimum price of 4.50.
#' @param zone destination zone: Local, National or Abroad
#' @param cw chargeable weight in kilograms
#' @param n number of parcels handed in at once
#' @export
parcel_price <- function(zone, cw, n = 1) {
  if (!zone %in% zone_rates$zone) stop("unknown zone")
  base <- zone_rates[zone_rates$zone == zone, "base_rate"]
  rate <- zone_rates[zone_rates$zone == zone, "rate_per_kg"]
  price <- base + rate * cw
  with_fuel <- price * (1 + 0.125)
  discounted <- with_fuel * (1 - volume_discount(n))
  round(pmax(discounted, 4.50), 2)
}

#' Volume discount
#'
#' Discount rate of 1% per parcel, capped at 20%.
#' @param n number of parcels handed in at once
#' @export
volume_discount <- function(n) {
  pmin(0.2, 0.01 * n)
}

# Helper used when printing a price list.
format_price <- function(x) {
  paste0(format(x, nsmall = 2), " per parcel")
}
""",
    "R/data.R": """#' Rates by zone
#'
#' Base rate and rate per kilogram for each destination zone.
#' @format A data frame with 3 rows and 3 columns:
#' \\describe{
#'   \\item{zone}{destination zone}
#'   \\item{base_rate}{base rate per parcel}
#'   \\item{rate_per_kg}{rate per kilogram of chargeable weight}
#' }
"zone_rates"
""",
    "man/dim_weight.Rd": """% Generated by roxygen2: do not edit by hand
% Please edit documentation in R/weights.R
\\name{dim_weight}
\\alias{dim_weight}
\\title{Dimensional weight}
\\usage{
dim_weight(l, w, h, divisor = 5000)
}
\\arguments{
\\item{l}{length in centimetres}
\\item{w}{width in centimetres}
\\item{h}{height in centimetres}
\\item{divisor}{the fixed divisor, 5000}
}
\\description{
Weight derived from the volume of a parcel.
}
""",
    "man/chargeable_weight.Rd": """% Generated by roxygen2: do not edit by hand
% Please edit documentation in R/weights.R
\\name{chargeable_weight}
\\alias{chargeable_weight}
\\title{Chargeable weight}
\\usage{
chargeable_weight(aw, dw)
}
\\arguments{
\\item{aw}{actual weight in kilograms}
\\item{dw}{dimensional weight in kilograms}
}
""",
    "man/parcel_price.Rd": """% Generated by roxygen2: do not edit by hand
% Please edit documentation in R/price.R
\\name{parcel_price}
\\alias{parcel_price}
\\title{Price of a parcel}
\\usage{
parcel_price(zone, cw, n = 1)
}
\\arguments{
\\item{zone}{destination zone: Local, National or Abroad}
\\item{cw}{chargeable weight in kilograms}
\\item{n}{number of parcels handed in at once}
}
""",
    "man/volume_discount.Rd": """% Generated by roxygen2: do not edit by hand
% Please edit documentation in R/price.R
\\name{volume_discount}
\\alias{volume_discount}
\\title{Volume discount}
\\usage{
volume_discount(n)
}
\\arguments{
\\item{n}{number of parcels handed in at once}
}
""",
    "tests/testthat.R": "library(testthat)\nlibrary(parcelcost)\ntest_check(\"parcelcost\")\n",
    "tests/testthat/test-price.R": """test_that("dimensional weight follows the divisor", {
  expect_equal(dim_weight(50, 40, 30), 12)
})

test_that("the minimum price holds", {
  expect_equal(parcel_price("Local", 0.5), 4.50)
})
""",
}

A_DOCUMENTATION = [
    ("h", 1, "1 Overview"),
    ("p", "This document describes how the parcelcost package implements the parcel pricing method."),
    ("p", "The package is maintained by the pricing office and released once a year."),
    ("h", 1, "2 Weights"),
    ("p", "The function dim_weight derives the dimensional weight from length, width and height with the fixed divisor of 5000."),
    ("p", "The function chargeable_weight returns the larger of actual weight and dimensional weight and never returns less than 0.5 kilograms."),
    ("h", 1, "3 Price"),
    ("p", "The function parcel_price looks up the base rate and the rate per kilogram of the zone in the table zone_rates."),
    ("table", "Table A. Rates held in the package", [["Zone", "Base rate", "Rate per kg"], ["Local", "3.20", "0.85"],
                                                     ["National", "4.90", "1.10"], ["Abroad", "9.50", "2.75"]]),
    ("p", "The price before surcharge equals the base rate plus the rate per kilogram times the chargeable weight."),
    ("p", "A fuel surcharge of 12.5% is added to the price before surcharge."),
    ("p", "The volume discount is 1% per parcel and is capped at 20%."),
    ("p", "The final price is never below 4.50 and is rounded to 2 decimals."),
    ("h", 1, "4 Testing"),
    ("p", "Unit tests cover the dimensional weight and the minimum price."),
    ("p", "The tests are run before every release."),
    ("h", 1, "5 Contacts"),
    ("p", "Questions are answered by the pricing office."),
    ("image", "Figure A. Release calendar"),
]


def build_a_minimal():
    import pandas
    sample = "A_minimal"
    write_bytes(sample, "canon", "parcel_pricing_method.xml", A_METHODOLOGY)
    rates = pandas.DataFrame({"zone": ["Local", "National", "Abroad"], "base_rate": [3.20, 4.90, 9.50],
                              "rate_per_kg": [0.85, 1.10, 2.75]})
    files = dict(A_PACKAGE)
    files["data/zone_rates.rda"] = r_data({"zone_rates": rates})
    write_bytes(sample, "package", "parcelcost_1.2.0.tar.gz", tarball("parcelcost", files))
    write_bytes(sample, "doc", "parcelcost_documentation.docx", docx_file(A_DOCUMENTATION))



# ------------------------------------------------------------------ sample F_capital (finance flavour, invented)
MATHML_COND_PD = ('<math xmlns="http://www.w3.org/1998/Math/MathML"><mi>PDc</mi><mo>=</mo><mi>N</mi><mo>&#x2061;</mo><mfenced><mfrac>'
                  '<mrow><msup><mi>N</mi><mrow><mo>-</mo><mn>1</mn></mrow></msup><mo>&#x2061;</mo><mfenced><mi>PD</mi></mfenced><mo>+</mo>'
                  '<msqrt><mi>&#x3C1;</mi></msqrt><mo>&#x2062;</mo><msup><mi>N</mi><mrow><mo>-</mo><mn>1</mn></mrow></msup><mo>&#x2061;</mo>'
                  '<mfenced><mn>0.999</mn></mfenced></mrow><msqrt><mrow><mn>1</mn><mo>-</mo><mi>&#x3C1;</mi></mrow></msqrt></mfrac></mfenced></math>')

F_METHODOLOGY = """<?xml version="1.0" encoding="UTF-8"?>
<methodology>
 <title>Capital requirement methodology</title>
 <part><title>A. Framework</title>
  <section><title>1. Scope</title>
   <p>This methodology sets out how the capital requirement for unexpected losses is determined for each exposure.</p>
   <p>It applies to all exposures of the lending book. Trading positions are outside its scope.</p>
  </section>
  <section><title>2. Definitions</title>
   <p>The probability of default (PD) is the likelihood that an obligor defaults within one year.</p>
   <p>The loss given default (LGD) is the share of the exposure that is lost when the obligor defaults.</p>
   <p>The exposure at default (EAD) is the amount outstanding when the obligor defaults.</p>
   <p>The requirement uses the normal distribution function N and its inverse, where &rho; denotes the asset correlation.</p>
   <table><caption>Table 1. Symbols</caption>
    <tr><th>Symbol</th><th>Description</th></tr>
    <tr><td>PDc</td><td>conditional probability of default</td></tr>
    <tr><td>K</td><td>capital requirement per unit of exposure</td></tr>
    <tr><td>M</td><td>effective maturity in years</td></tr>
   </table>
  </section>
  <section><title>3. Risk components</title>
   <section><title>3.1 Probability of default</title>
    <p>The probability of default is estimated from internal ratings and is reviewed every year.</p>
    <section><title>3.1.1 Floor</title>
     <p>The probability of default is never taken below 0.03%.</p>
     <p>The floor applies before any other step of the calculation.</p>
    </section>
   </section>
   <section><title>3.2 Conditional probability of default</title>
    <p>The conditional probability of default is evaluated at the 99.9th percentile of the systematic factor.</p>
    <equation>""" + MATHML_COND_PD + """</equation>
    <p>The confidence level of 0.999 is fixed and is set out again in Annex A.</p>
   </section>
   <section><title>3.3 Asset correlation</title>
    <p>The asset correlation falls with the probability of default: R = 0.12 * (1 - exp(-50 * PD)) / (1 - exp(-50)) + 0.24 * (1 - (1 - exp(-50 * PD)) / (1 - exp(-50))).</p>
    <p>The correlation therefore lies between 0.12 and 0.24.</p>
   </section>
   <section><title>3.4 Loss given default</title>
    <p>The loss given default is at least the floor of its segment, as set out in Table 3.</p>
    <table><caption>Table 3. Floors by segment</caption>
     <tr><th>Segment</th><th>LGD floor</th></tr>
     <tr><td>Retail</td><td>0.15</td></tr>
     <tr><td>Corporate</td><td>0.25</td></tr>
     <tr><td>Bank</td><td>0.45</td></tr>
    </table>
   </section>
  </section>
 </part>
 <part><title>B. Capital</title>
  <section><title>4. Capital requirement</title>
   <section><title>4.1 Requirement per unit of exposure</title>
    <p>The unexpected loss rate is the conditional probability of default less the probability of default: UL = PDc - PD.</p>
    <p>K is the product of LGD and UL.</p>
   </section>
   <section><title>4.2 Risk-weighted amount</title>
    <equation><img src="rwa_formula.png" alt="Formula for the risk-weighted amount"/></equation>
    <p>The requirement is scaled by 12.5 and multiplied by the exposure at default.</p>
   </section>
   <section><title>4.3 Effective maturity</title>
    <p>The effective maturity is at least 1 year.</p>
    <p>The effective maturity is capped at 5 years.</p>
   </section>
  </section>
 </part>
 <annex>
  <heading>Annex A Parameters</heading>
  <p>This annex restates the fixed parameters of the methodology.</p>
  <heading>I. Confidence level</heading>
  <p>The confidence level is 0.999 for every segment.</p>
  <heading>II. Floors</heading>
  <p>Floors are reviewed once a year by the model owner.</p>
  <heading>(a) Probability of default</heading>
  <p>The floor of the probability of default is 0.03%, as stated in section 3.1.1.</p>
  <heading>(b) Loss given default</heading>
  <p>The floors of the loss given default are those of Table 3.</p>
 </annex>
</methodology>
"""

def rd_page(name, title, arguments, source):
    usage = "%s(%s)" % (name, ", ".join(a if d is None else "%s = %s" % (a, d) for a, d, _ in arguments))
    items = "\n".join("\\item{%s}{%s}" % (a, text) for a, _, text in arguments)
    return ("%% Generated by roxygen2: do not edit by hand\n%% Please edit documentation in %s\n\\name{%s}\n\\alias{%s}\n\\title{%s}\n"
            "\\usage{\n%s\n}\n\\arguments{\n%s\n}\n" % (source, name, name, title, usage, items))

F_R_PD = """#' Floor of the probability of default
#'
#' The probability of default is never taken below 0.03%.
#' @param pd probability of default
#' @param floor the floor, defaults to 0.0003
#' @export
floor_pd <- function(pd, floor = 0.0003) {
  pmax(pd, floor)
}

#' Conditional probability of default
#'
#' Evaluated at the 99.9th percentile of the systematic factor:
#' \\deqn{PDc = \\Phi\\left(\\frac{\\Phi^{-1}(PD) + \\sqrt{\\rho} \\cdot \\Phi^{-1}(0.999)}{\\sqrt{1 - \\rho}}\\right)}
#' @param pd probability of default
#' @param rho asset correlation
#' @param q confidence level, defaults to 0.999
#' @export
cond_pd <- function(pd, rho, q = 0.999) {
  a <- qnorm(pd)
  b <- sqrt(rho) * qnorm(q)
  pnorm((a + b) / sqrt(1 - rho))
}
"""
F_R_CORRELATION = """#' Asset correlation
#'
#' The asset correlation falls with the probability of default and lies between 0.12 and 0.24.
#' @param pd probability of default
#' @export
asset_correlation <- function(pd) {
  w <- (1 - exp(-50 * pd)) / (1 - exp(-50))
  0.12 * w + 0.24 * (1 - w)
}
"""
F_R_CAPITAL = """#' Capital requirement per unit of exposure
#'
#' The product of the loss given default, floored by segment, and the unexpected loss rate.
#' @param pd probability of default
#' @param lgd loss given default
#' @param segment segment of the exposure: Retail, Corporate or Bank
#' @export
capital_k <- function(pd, lgd, segment) {
  check_inputs(pd, lgd)
  lgd_f <- pmax(lgd, lgd_floors[lgd_floors$segment == segment, "lgd_floor"])
  pdc <- cond_pd(floor_pd(pd), asset_correlation(pd))
  ul <- pdc - pd
  lgd_f * ul
}

#' Risk-weighted amount
#'
#' The requirement scaled by 12.5 and multiplied by the exposure at default.
#' @param k capital requirement per unit of exposure
#' @param ead exposure at default
#' @export
rwa <- function(k, ead) {
  k * 12.5 * ead
}

#' Effective maturity
#'
#' At least 1 year and capped at 5 years.
#' @param m maturity in years
#' @export
effective_maturity <- function(m) {
  pmin(pmax(m, 1), 5)
}
"""
F_R_UTILS = """# Argument checks shared by the exported functions.
check_inputs <- function(pd, lgd) {
  if (!is.numeric(pd)) stop("pd must be numeric")
  if (!is.numeric(lgd)) stop("lgd must be numeric")
  invisible(TRUE)
}

`%||%` <- function(a, b) if (is.null(a)) b else a
"""
F_R_DATA = """#' Floors of the loss given default by segment
#'
#' @format A data frame with 3 rows and 2 columns:
#' \\describe{
#'   \\item{segment}{segment of the exposure}
#'   \\item{lgd_floor}{floor of the loss given default}
#' }
#' @source Table 3 of the methodology
"lgd_floors"
"""
F_VIGNETTE = """---
title: "Using capreq"
---

This vignette shows how the capital requirement of one exposure is obtained.

```{r}
library(capreq)
k <- capital_k(0.01, 0.40, "Corporate")
rwa(k, 1000)
```

The floors by segment are stored in the object lgd_floors.
"""

def f_package(known):
    """The files of package capreq; `known` plants the six seeded differences of the brief."""
    import pandas
    pd_source, capital, data = F_R_PD, F_R_CAPITAL, F_R_DATA
    floors = pandas.DataFrame({"segment": ["Retail", "Corporate", "Bank"], "lgd_floor": [0.15, 0.25, 0.45]})
    if known:
        pd_source = pd_source.replace("pnorm((a + b) / sqrt(1 - rho))", "pnorm((a - b) / sqrt(1 - rho))")      # 1. a flipped sign
        pd_source = pd_source.replace("  pmax(pd, floor)\n", "  pd\n")                                           # 3. a removed floor
        pd_source = pd_source.replace("#' @param rho asset correlation", "#' @param rho_a asset correlation")    # 5. roxygen out of step
        floors.loc[0, "lgd_floor"] = 0.10                                                                        # 4. a changed cell
    correlation = F_R_CORRELATION.replace("0.12 * w", "0.13 * w") if known else F_R_CORRELATION                  # 2. a changed constant
    confidence = pandas.DataFrame({"parameter": ["confidence_level"], "value": [0.999]})
    files = {
        "DESCRIPTION": "Package: capreq\nTitle: Capital Requirement\nVersion: 0.9.1\nDescription: Capital requirement for unexpected losses.\nLicense: MIT\nEncoding: UTF-8\nVignetteBuilder: knitr\n",
        "NAMESPACE": "\n".join("export(%s)" % n for n in ("floor_pd", "cond_pd", "asset_correlation", "capital_k", "rwa", "effective_maturity")) + "\n",
        "R/pd.R": pd_source, "R/correlation.R": correlation, "R/capital.R": capital, "R/utils.R": F_R_UTILS, "R/data.R": data,
        "data/lgd_floors.rda": r_data({"lgd_floors": floors}), "inst/extdata/confidence.rds": r_data(confidence, single=True),
        "man/floor_pd.Rd": rd_page("floor_pd", "Floor of the probability of default", [("pd", None, "probability of default"), ("floor", "0.0003", "the floor, defaults to 0.0003")], "R/pd.R"),
        # the page of cond_pd is stale on purpose: it was generated before the argument q was added
        "man/cond_pd.Rd": rd_page("cond_pd", "Conditional probability of default", [("pd", None, "probability of default"), ("rho", None, "asset correlation")], "R/pd.R"),
        "man/asset_correlation.Rd": rd_page("asset_correlation", "Asset correlation", [("pd", None, "probability of default")], "R/correlation.R"),
        "man/capital_k.Rd": rd_page("capital_k", "Capital requirement per unit of exposure", [("pd", None, "probability of default"), ("lgd", None, "loss given default"), ("segment", None, "segment of the exposure: Retail, Corporate or Bank")], "R/capital.R"),
        "man/rwa.Rd": rd_page("rwa", "Risk-weighted amount", [("k", None, "capital requirement per unit of exposure"), ("ead", None, "exposure at default")], "R/capital.R"),
        "man/effective_maturity.Rd": rd_page("effective_maturity", "Effective maturity", [("m", None, "maturity in years")], "R/capital.R"),
        "tests/testthat.R": "library(testthat)\nlibrary(capreq)\ntest_check(\"capreq\")\n",
        "tests/testthat/test-pd.R": "test_that(\"the floor of the probability of default holds\", {\n  expect_equal(floor_pd(0.0001), 0.0003)\n})\n\n"
                                    "test_that(\"the conditional probability exceeds the probability\", {\n  expect_true(cond_pd(0.01, 0.2) > 0.01)\n})\n",
        "vignettes/capreq.Rmd": F_VIGNETTE}
    return files

def f_documentation(known):
    level = "99.5th percentile" if known else "99.9th percentile"                                              # 6. a stale value in the documentation
    cond = omml(run("PDc=N") + "<m:d><m:e><m:f><m:num>" + "<m:sSup><m:e>" + run("N") + "</m:e><m:sup>" + run("-1") + "</m:sup></m:sSup><m:d><m:e>" + run("PD")
                + "</m:e></m:d>" + run("+") + "<m:rad><m:radPr><m:degHide m:val=\"1\"/></m:radPr><m:deg/><m:e>" + run("&#x3C1;") + "</m:e></m:rad>" + run("&#xD7;")
                + "<m:sSup><m:e>" + run("N") + "</m:e><m:sup>" + run("-1") + "</m:sup></m:sSup><m:d><m:e>" + run("0.999") + "</m:e></m:d></m:num><m:den>"
                + "<m:rad><m:radPr><m:degHide m:val=\"1\"/></m:radPr><m:deg/><m:e>" + run("1-&#x3C1;") + "</m:e></m:rad></m:den></m:f></m:e></m:d>")
    return [
        ("h", 1, "1 Purpose"),
        ("p", "This document describes how the package capreq implements the capital requirement methodology."),
        ("p", "It is written for model validators and for the developers who maintain the package."),
        ("h", 1, "2 Risk components"),
        ("h", 2, "2.1 Probability of default"),
        ("p", "The function floor_pd applies the floor of 0.03% to the probability of default."),
        ("h", 2, "2.2 Conditional probability of default"),
        ("p", "The function cond_pd evaluates the conditional probability of default at the %s of the systematic factor." % level),
        ("eq", cond),
        ("h", 2, "2.3 Asset correlation"),
        ("p", "The function asset_correlation returns a correlation between 0.12 and 0.24 that falls with the probability of default."),
        ("h", 2, "2.4 Loss given default"),
        ("p", "The floors of the loss given default are held in the object lgd_floors and are shown in Table D1."),
        ("table", "Table D1. Floors by segment", [["Segment", "LGD floor"], ["Retail", "0.15"], ["Corporate", "0.25"], ["Bank", "0.45"]]),
        ("h", 1, "3 Capital"),
        ("p", "The function capital_k multiplies the floored loss given default by the unexpected loss rate."),
        ("p", "The function rwa scales the requirement by 12.5 and multiplies it by the exposure at default."),
        ("p", "The function effective_maturity keeps the maturity between 1 year and 5 years."),
        ("h", 1, "4 Change history"),
        ("p", "The package is released after each yearly review of the methodology."),
        ("image", "Figure D1. Release process"),
    ]

F_USER_GUIDE = [
    ("h", 1, "1 Installing the package"),
    ("p", "The package capreq is installed from the internal repository and needs no other package."),
    ("h", 1, "2 Calculating a requirement"),
    ("p", "Call capital_k with the probability of default, the loss given default and the segment of the exposure."),
    ("p", "The risk-weighted amount is obtained with rwa, which scales the requirement by 12.5."),
    ("table", "Floors used by the package", [["Segment", "LGD floor"], ["Retail", "0.15"], ["Corporate", "0.25"], ["Bank", "0.45"]]),
]

def build_f_capital(sample="F_capital", known=False):
    write_bytes(sample, "canon", "capital_methodology.txt", F_METHODOLOGY)          # XML inside a .txt file, on purpose
    write_bytes(sample, "package", "capreq_0.9.1.tar.gz", tarball("capreq", f_package(known)))
    write_bytes(sample, "doc", "capreq_model_documentation.docx", docx_file(f_documentation(known)))
    write_bytes(sample, "doc", "capreq_user_guide.pdf", pdf_file(F_USER_GUIDE))

def build_f_capital_known():
    build_f_capital("F_capital_known", known=True)

# ------------------------------------------------------------------ sample D_dosing (another field; keeps the engine honest about rule R9)
D_OMML = ('<m:oMath><m:r><m:t>D=W&#215;R</m:t></m:r></m:oMath>')
D_HTML = """<html xmlns:m="http://schemas.microsoft.com/office/2004/12/omml"><head><title>Dosing method</title><style>p{margin:0}</style></head><body>
<h1>1 Purpose</h1>
<p>This method sets out how the daily dose of the medicine is determined for an adult patient.
<p>It applies to oral treatment only.
<h1>2 Dose by body weight</h1>
<p>The body weight (W) is measured in kilograms, and the dose rate (R) is 15 milligrams per kilogram.
<p><!--[if gte msEquation 12]>""" + D_OMML + """<![endif]--><![if !msEquation]><img src="image001.png" alt="Dose formula"><![endif]></p>
<p>The daily dose is never above 4000 milligrams.
<h1>3 Adjustment for kidney function</h1>
<p>The dose is multiplied by the factor of the clearance band of the patient, as set out in Table 1.
<table><caption>Table 1. Factors by clearance band</caption>
<tr><th>Band</th><th>Factor</th></tr><tr><td>Normal</td><td>1.00</td></tr><tr><td>Reduced</td><td>0.75</td></tr><tr><td>Low</td><td>0.50</td></tr></table>
<p><img src="image002.png" alt="Adjusted dose formula"></p>
<p>The adjusted dose is rounded to 10 milligrams.
</body></html>
"""

def mhtml_file(page, images):
    boundary = "----=_NextPart_AIVA_SAMPLE"
    parts = ["MIME-Version: 1.0\nContent-Type: multipart/related; boundary=\"%s\"\n" % boundary,
             "--%s\nContent-Location: file:///C:/method.htm\nContent-Transfer-Encoding: 8bit\nContent-Type: text/html; charset=\"utf-8\"\n\n%s" % (boundary, page)]
    import base64
    for name in images:
        parts.append("--%s\nContent-Location: %s\nContent-Transfer-Encoding: base64\nContent-Type: image/png\n\n%s\n"
                     % (boundary, name, base64.b64encode(tiny_png() + name.encode("ascii")).decode("ascii")))
    return "\n".join(parts) + "\n--%s--\n" % boundary

D_R = """#' Dose by body weight
#'
#' The body weight times the dose rate, never above 4000 milligrams.
#' @param weight body weight in kilograms
#' @param rate dose rate in milligrams per kilogram, defaults to 15
#' @export
dose_by_weight <- function(weight, rate = 15) {
  dose <- weight * rate
  pmin(dose, 4000)
}

#' Dose adjusted for kidney function
#'
#' The dose times the factor of the clearance band, rounded to 10 milligrams.
#' @param dose daily dose in milligrams
#' @param band clearance band: Normal, Reduced or Low
#' @export
adjusted_dose <- function(dose, band) {
  factor <- band_factors[band_factors$band == band, "factor"]
  round(dose * factor / 10) * 10
}
"""
D_GUIDE = [
    ("h", 1, "1 Overview"),
    ("p", "The package dosecalc implements the dosing method for adult patients."),
    ("h", 1, "2 Functions"),
    ("p", "The function dose_by_weight multiplies the body weight by the dose rate of 15 milligrams per kilogram and never returns more than 4000 milligrams."),
    ("p", "The function adjusted_dose multiplies the dose by the factor of the clearance band."),
    ("table", "Factors held in the package", [["Band", "Factor"], ["Normal", "1.00"], ["Reduced", "0.75"], ["Low", "0.50"]]),
]

def build_d_dosing():
    import pandas
    sample = "D_dosing"
    write_bytes(sample, "canon", "dosing_method.mhtml", mhtml_file(D_HTML, ["image001.png", "image002.png"]))
    factors = pandas.DataFrame({"band": ["Normal", "Reduced", "Low"], "factor": [1.0, 0.75, 0.5]})
    files = {"DESCRIPTION": "Package: dosecalc\nTitle: Dose Calculation\nVersion: 0.2.0\nLicense: MIT\nEncoding: UTF-8\n",
             "NAMESPACE": "export(dose_by_weight)\nexport(adjusted_dose)\n", "R/dose.R": D_R,
             "R/data.R": "#' Factors by clearance band\n#'\n#' @format A data frame with 3 rows and 2 columns:\n#' \\describe{\n#'   \\item{band}{clearance band}\n#'   \\item{factor}{factor applied to the dose}\n#' }\n\"band_factors\"\n",
             "inst/extdata/band_factors.rds": r_data(factors, single=True),
             "man/dose_by_weight.Rd": rd_page("dose_by_weight", "Dose by body weight", [("weight", None, "body weight in kilograms"), ("rate", "15", "dose rate in milligrams per kilogram, defaults to 15")], "R/dose.R"),
             "man/adjusted_dose.Rd": rd_page("adjusted_dose", "Dose adjusted for kidney function", [("dose", None, "daily dose in milligrams"), ("band", None, "clearance band: Normal, Reduced or Low")], "R/dose.R")}
    write_bytes(sample, "package", "dosecalc_0.2.0.tar.gz", tarball("dosecalc", files))
    write_bytes(sample, "doc", "dosecalc_guide.pdf", pdf_file(D_GUIDE))


BUILDERS = {"A_minimal": build_a_minimal, "F_capital": build_f_capital, "F_capital_known": build_f_capital_known, "D_dosing": build_d_dosing}

if __name__ == "__main__":
    for name in (sys.argv[1:] or list(BUILDERS)):
        BUILDERS[name]()
        print("built", name)
